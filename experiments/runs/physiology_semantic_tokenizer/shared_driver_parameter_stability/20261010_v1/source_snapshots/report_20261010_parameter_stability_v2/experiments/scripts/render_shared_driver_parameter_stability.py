#!/usr/bin/env python3
"""Chinese PPT/PDF from retained parameter-stability evidence; bitmap figures."""
from __future__ import annotations
import argparse
import json
from pathlib import Path
import subprocess
import sys
import math
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.colors import Normalize
from PIL import Image,ImageFont
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
import fitz

ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT))
from src.inference.shared_driver_stability import PARAMETERS,REFERENCE,BOUNDS,boundary_distance
DS={'eeg_fnirs_single_trial':'Single-Trial','simultaneous_eeg_nirs':'Simultaneous','visual_cognitive_motivation':'Visual'}
NAME={'tau':'τ','neurovascular_gain':'β','kappa':'κ'}
MODEL={'P0':'P0 固定','P1_tau':'P1：τ','P1_beta':'P1：β','P2':'P2：τ、β','P3':'P3：τ、β、κ'}
METHOD={'N0':'基线 N0','C0a':'分段 C0a','C0b':'分段／单初态先验','C1':'连续 90 s','W1':'工作协方差','Hselected':'开发选择的共享','combined':'两项组合','Hfull':'完全共享',
        'H0.25':'共享 λ=0.25','H1.0':'共享 λ=1','H4.0':'共享 λ=4','oracle':'已知驱动／初态','fixed':'固定参考值'}
for strength in (.25,1,4):
    METHOD[f'H{strength:g}_C0a']=f'共享 λ={strength:g}'
    METHOD[f'H{strength:g}_C1']=f'共享 λ={strength:g}＋连续'
COLOR={'N0':'#315B8A','C0a':'#315B8A','C0b':'#A56A3F','C1':'#008879','W1':'#AC5A91','Hselected':'#BD8120','combined':'#6650AD','Hfull':'#8C949C','oracle':'#222222','fixed':'#A4AAB0'}
FONT='/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'


def read(path):return json.loads(Path(path).read_text())


def finite(x):return x is not None and np.isfinite(x)


def fmt(x,d=3):return f'{x:.{d}f}' if finite(x) else '—'


def pct(x,d=1):return f'{100*x:.{d}f}%' if finite(x) else '—'


def clean(x):
    if isinstance(x,dict):return {str(k):clean(v) for k,v in x.items()}
    if isinstance(x,(list,tuple,np.ndarray)):return [clean(v) for v in x]
    if isinstance(x,np.generic):return clean(x.item())
    if isinstance(x,float) and not np.isfinite(x):return None
    return x


def setup():
    font_manager.fontManager.addfont(FONT)
    plt.rcParams.update({'font.family':[font_manager.FontProperties(fname=FONT).get_name(),'DejaVu Sans'],
        'font.size':17,'axes.titlesize':19,'axes.labelsize':17,'xtick.labelsize':15,'ytick.labelsize':15,
        'legend.fontsize':14,'axes.unicode_minus':False,'axes.spines.top':False,'axes.spines.right':False,
        'axes.grid':True,'grid.alpha':.17,'savefig.facecolor':'white','figure.facecolor':'white'})


def save(fig,out,name):
    path=out/'figures'/(name+'.png');path.parent.mkdir(parents=True,exist_ok=True)
    fig.savefig(path,dpi=220,bbox_inches='tight');plt.close(fig);return path


class Deck:
    def __init__(self,out):
        self.out=out;self.prs=Presentation();self.prs.slide_width=Inches(13.333);self.prs.slide_height=Inches(7.5)
        self.sources=[];self.figures=[];self.text_boxes=[];self.figure_pages={}

    def text(self,slide,x,y,w,h,content,size=19,bold=False,color='#24364A'):
        # Check visible CJK text at the final slide size, not only in notes.
        font=ImageFont.truetype(FONT,round(size*4));lines=0
        for line in str(content).split('\n'):
            occupied=0.;lines+=1
            for ch in line:
                advance=font.getlength(ch)/4
                if occupied+advance>(w-.06)*72:lines+=1;occupied=0.
                occupied+=advance
        needed=lines*size*1.16+max(0,str(content).count('\n'))*4
        if needed>(h-.025)*72:
            raise ValueError(f'text overflow ({needed:.0f}>{h*72:.0f} pt): {str(content)[:90]}')
        box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=box.text_frame
        tf.word_wrap=True;tf.margin_top=tf.margin_bottom=Inches(.01);tf.margin_left=tf.margin_right=Inches(.02)
        tf.vertical_anchor=MSO_ANCHOR.TOP
        for j,line in enumerate(str(content).split('\n')):
            p=tf.paragraphs[0] if j==0 else tf.add_paragraph();p.text=line;p.font.name='Noto Sans CJK SC';p.font.size=Pt(size)
            p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color.lstrip('#'));p.line_spacing=1.16;p.space_after=Pt(4)
        self.text_boxes.append(dict(slide=len(self.prs.slides),text=str(content),font_size=size,bounds=[x,y,w,h]))
        return box

    def base(self,title,subtitle,source):
        slide=self.prs.slides.add_slide(self.prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor(248,250,253)
        self.text(slide,.5,.25,12.25,.62,title,26,True)
        self.text(slide,.52,1.01,12.2,.53,subtitle,14,color='#536A80')
        self.text(slide,.52,7.1,11.8,.23,'证据：'+source,8.5,color='#66788A')
        self.text(slide,12.45,7.03,.5,.3,str(len(self.prs.slides)),11,color='#66788A')
        self.sources.append(dict(slide=len(self.prs.slides),title=title,source=source))
        return slide

    def image(self,slide,path,x=.6,y=1.6,w=12.1,h=4.65):
        iw,ih=Image.open(path).size;scale=min(w/iw,h/ih);ww,hh=iw*scale,ih*scale
        slide.shapes.add_picture(str(path),Inches(x+(w-ww)/2),Inches(y+(h-hh)/2),width=Inches(ww),height=Inches(hh))
        self.figures.append(str(path))
        page=len(self.prs.slides);self.figure_pages[page]=self.figure_pages.get(page,0)+1

    def figure(self,title,subtitle,path,interpretation,source):
        slide=self.base(title,subtitle,source);self.image(slide,path)
        self.text(slide,.65,6.35,12.0,.64,interpretation,17.5)
        return slide

    def methods(self,title,subtitle,blocks,source,conclusion=None):
        slide=self.base(title,subtitle,source)
        positions=[(.65,1.68),(6.82,1.68),(.65,4.03),(6.82,4.03)]
        for (heading,body),(x,y) in zip(blocks,positions):
            self.text(slide,x,y,5.8,.4,heading,21,True,color='#087F8C')
            self.text(slide,x,y+.52,5.8,1.63,body,19)
        if conclusion:self.text(slide,.65,6.53,12,.42,conclusion,17,bold=True)
        return slide

    def table(self,title,subtitle,headers,rows,widths,interpretation,source):
        slide=self.base(title,subtitle,source)
        # The full parameter sets are defined on the model page; short table
        # labels avoid wrapping a three-parameter model across cell boundaries.
        labels={MODEL['P1_tau']:'P1τ',MODEL['P1_beta']:'P1β',MODEL['P2']:'P2',MODEL['P3']:'P3',
            'P1_tau':'P1τ','P1_beta':'P1β'}
        rows=[[labels.get(str(value),value) for value in row] for row in rows]
        size=14.5 if len(headers)>7 else 17;measure=ImageFont.truetype(FONT,round(size*4));heights=[]
        for row in [headers]+rows:
            line_counts=[]
            for value,width in zip(row,widths):
                lines=1;occupied=0.
                for ch in str(value):
                    advance=measure.getlength(ch)/4
                    if ch=='\n' or occupied+advance>(width-.09)*72:lines+=1;occupied=0.
                    if ch!='\n':occupied+=advance
                line_counts.append(lines)
            heights.append(max(.43,(max(line_counts)*size*1.25+10)/72))
        total=sum(widths);height=sum(heights)
        if height>4.65:raise ValueError(f'table overflow ({height:.2f} in): {title}; split this table')
        shape=slide.shapes.add_table(len(rows)+1,len(headers),Inches(.57),Inches(1.65),Inches(total),Inches(height))
        table=shape.table
        for j,w in enumerate(widths):table.columns[j].width=Inches(w)
        for i,h in enumerate(heights):table.rows[i].height=Inches(h)
        for i,row in enumerate([headers]+rows):
            for j,value in enumerate(row):
                cell=table.cell(i,j);cell.text=str(value);cell.margin_left=Inches(.04);cell.margin_right=Inches(.03)
                cell.margin_top=Inches(.055);cell.margin_bottom=Inches(.035)
                cell.vertical_anchor=MSO_ANCHOR.TOP
                cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string('24445F' if i==0 else ('FFFFFF' if i%2 else 'EDF3F8'))
                for p in cell.text_frame.paragraphs:
                    p.font.name='Noto Sans CJK SC';p.font.size=Pt(size)
                    p.font.color.rgb=RGBColor(255,255,255) if i==0 else RGBColor(36,54,74);p.font.bold=i==0
                    p.line_spacing=Pt(size*1.25);p.space_before=Pt(0);p.space_after=Pt(0)
                self.text_boxes.append(dict(slide=len(self.prs.slides),text=str(value),font_size=size,kind='table_cell',
                    bounds=[.57+sum(widths[:j]),1.65+sum(heights[:i]),widths[j],heights[i]]))
        self.text(slide,.65,6.38,12,.62,interpretation,17.5)
        return slide


def primary(frame,ds):
    kind='within_record_disjoint_blocks' if ds=='simultaneous_eeg_nirs' else 'independent_native_record'
    return frame[(frame.dataset==ds)&(frame.repeat_kind==kind)]


def best_group(groups,ds,model):
    g=primary(groups,ds);g=g[(g.partition=='evaluation')&(g.model==model)&(g.variant!='Hfull')]
    return g.sort_values(['effective_pair_rate','repeat_median','variant'],ascending=[False,True,True]).iloc[0]


def synthetic_fig(out,summary,model):
    rows=pd.DataFrame(summary['synthetic']);g=rows[(rows.model==model)&(rows.level==2)]
    names=[n for n in PARAMETERS if n in set(g.parameter)];fig,axes=plt.subplots(1,len(names),figsize=(13,4.9),squeeze=False)
    variants=['C0a','C0b','C1','oracle','fixed']
    for ax,name in zip(axes[0],names):
        own=g[g.parameter==name].set_index('variant');x=np.arange(len(variants))
        med=[100*own.loc[v,'median_relative_error'] for v in variants];p90=[100*own.loc[v,'p90_relative_error'] for v in variants]
        ax.bar(x,med,color=[COLOR.get(v,'#A56A3F') for v in variants],width=.65)
        ax.scatter(x,p90,color='#24364A',marker='D',s=55,zorder=3,label='P90')
        ax.axhline(10,color='#B64745',ls='--',lw=1.5,label='中位误差门槛 10%')
        ax.set_xticks(x,['分段','减先验','连续','oracle','固定'],rotation=30,ha='right');ax.set_title(NAME[name]);ax.set_ylabel('相对误差（%）')
    axes[0,0].legend(loc='upper left',fontsize=12);fig.tight_layout();return save(fig,out,'synthetic_'+model)


def repeat_scatter(out,parameters,ds,model,variant):
    g=primary(parameters[(parameters.partition=='evaluation')&(parameters.model==model)&(parameters.variant==variant)],ds)
    sites=sorted(s for s in set(g.site) if s.startswith('prefrontal'));site=sites[0];g=g[g.site==site]
    names=[n for n in PARAMETERS if n in set(g.parameter)];fig,axes=plt.subplots(1,len(names),figsize=(13,4.7),squeeze=False)
    for ax,name in zip(axes[0],names):
        d=g[g.parameter==name];bounds=BOUNDS[PARAMETERS.index(name)];values=d[['A','B']].to_numpy()
        near=(boundary_distance(values,bounds)<.05).any(1);good=d.accepted.to_numpy(bool)
        for mask,color,marker,label in [(~near&good,'#008879','o','内点／数值通过'),(near&good,'#CB8235','^','至少一次近界'),(~good,'#B64745','x','未通过数值检查')]:
            if mask.any():ax.scatter(values[mask,0],values[mask,1],s=65,c=color,marker=marker,label=label,alpha=.8)
        ax.plot(bounds,bounds,'k--',lw=1.2);ax.set_xscale('log');ax.set_yscale('log');ax.set_xlim(bounds[0]*.9,bounds[1]*1.1);ax.set_ylim(bounds[0]*.9,bounds[1]*1.1)
        ticks={'tau':[.5,1,2,4,8],'neurovascular_gain':[.1,1,10],'kappa':[.2,.5,1.5]}[name]
        ax.set_xticks(ticks,[f'{v:g}' for v in ticks]);ax.set_yticks(ticks,[f'{v:g}' for v in ticks]);ax.minorticks_off()
        ax.set_xlabel('片段 A 的 '+NAME[name]);ax.set_ylabel('片段 B 的 '+NAME[name]);ax.set_title(f'{NAME[name]}｜{len(d)} 名被试')
    axes[0,0].legend(fontsize=11,loc='best');fig.tight_layout()
    return save(fig,out,'retest_'+ds+'_'+model+'_'+variant),site


def sharing_scatter_fig(out,parameters,ds,site,model):
    from matplotlib.lines import Line2D
    g=primary(parameters[(parameters.partition=='evaluation')&(parameters.site==site)&(parameters.model==model)],ds)
    names=[n for n in PARAMETERS if n in set(g.parameter)]
    fig,axes=plt.subplots(1,len(names),figsize=(13,5),squeeze=False);comparison=[]
    for ax,name in zip(axes[0],names):
        own=g[g.parameter==name];bounds=BOUNDS[PARAMETERS.index(name)]
        baseline=own[own.variant=='N0'].set_index('subject');shared=own[own.variant=='Hselected'].set_index('subject')
        for subject in baseline.index.intersection(shared.index):
            a=baseline.loc[subject];b=shared.loc[subject]
            ax.plot([a.A,b.A],[a.B,b.B],c='#BAC4CD',lw=.8,alpha=.5,zorder=1)
        for variant,color,marker in [('N0',COLOR['N0'],'o'),('Hselected',COLOR['Hselected'],'s')]:
            d=own[own.variant==variant];v=d[['A','B']].to_numpy();near=(boundary_distance(v,bounds)<.05).any(axis=1)
            for mask,face in [(~near,color),(near,'none')]:
                ax.scatter(v[mask,0],v[mask,1],s=58,facecolors=face,edgecolors=color,marker=marker,lw=1.4,zorder=3)
            failed=~d.accepted.to_numpy(bool);ax.scatter(v[failed,0],v[failed,1],c='#111111',marker='x',s=48,lw=1.1,zorder=4)
        ax.plot(bounds,bounds,'k--',lw=1.1);ax.set_xscale('log');ax.set_yscale('log')
        ax.set_xlim(bounds[0]*.9,bounds[1]*1.1);ax.set_ylim(bounds[0]*.9,bounds[1]*1.1)
        ticks={'tau':[.5,1,2,4,8],'neurovascular_gain':[.1,1,10],'kappa':[.2,.5,1.5]}[name]
        ax.set_xticks(ticks,[f'{v:g}' for v in ticks]);ax.set_yticks(ticks,[f'{v:g}' for v in ticks]);ax.minorticks_off()
        unit='（s）' if name=='tau' else '（1/s）' if name=='kappa' else '（有效增益）'
        ax.set_xlabel('A 的 '+NAME[name]+unit);ax.set_ylabel('B 的 '+NAME[name]+unit);ax.set_title(NAME[name]+f'｜{len(baseline)}名被试')
        comparison.append(NAME[name]+pct(baseline.repeat.median())+'→'+pct(shared.repeat.median()))
    handles=[Line2D([],[],marker='o',ls='',color=COLOR['N0'],label='N0'),
        Line2D([],[],marker='s',ls='',color=COLOR['Hselected'],label='冻结共享'),
        Line2D([],[],marker='o',ls='',mfc='none',mec='#333333',label='空心：至少一次近界'),
        Line2D([],[],marker='x',ls='',color='#111111',label='数值未通过')]
    fig.legend(handles=handles,loc='upper center',ncol=4,fontsize=12);fig.tight_layout(rect=(0,0,1,.89))
    return save(fig,out,'sharing_retest_'+ds+'_'+site+'_'+model),'；'.join(comparison)


def continuity_fig(out,pairs):
    frame=pairs[(pairs.partition=='development')&pairs.variant.isin(['N0','C0b','C1'])]
    fig,axes=plt.subplots(1,3,figsize=(13,4.8),sharey=True)
    for ax,(ds,label) in zip(axes,DS.items()):
        g=frame[(frame.dataset==ds)&(frame.model=='P2')];variants=['N0','C0b','C1'];x=np.arange(3)
        subject=g.groupby(['subject','variant']).repeat_error.mean().unstack('variant')
        for _,row in subject.iterrows():ax.plot(x,[100*row.get(v,np.nan) for v in variants],c='#B8C5D1',lw=1,alpha=.75)
        med=[100*subject[v].median() for v in variants];ax.plot(x,med,'o-',c='#008879',lw=3,ms=8,label='被试中位数')
        ax.axhline(20,c='#B64745',ls='--',lw=1.5);ax.set_xticks(x,['C0a','C0b','C1']);ax.set_title(label);ax.set_xlabel('每种方法均观察相同 90 s')
    axes[0].set_ylabel('参数向量最差重复差异（%）');axes[0].legend();fig.tight_layout();return save(fig,out,'continuity')


def fidelity_tradeoff_fig(out,pairs):
    frame=pairs[(pairs.partition=='development')&(pairs.model=='P2')]
    fig,axes=plt.subplots(1,3,figsize=(13,4.7),sharey=True)
    colors={'C1':COLOR['C1'],'W1':COLOR['W1'],'H4.0':COLOR['Hselected']}
    for ax,(ds,label) in zip(axes,DS.items()):
        g=frame[frame.dataset==ds];baseline=g[g.variant=='N0'].set_index('pair_id');x=np.arange(3)
        for j,variant in enumerate(['C1','W1','H4.0']):
            own=g[g.variant==variant].set_index('pair_id');delta=[]
            for component in ['EEG','HbO','HbR']:
                values=own[component+'_nrmse']-baseline[component+'_nrmse']
                delta.append(values.groupby(own.subject).mean().mean())
            ax.bar(x+(j-1)*.24,delta,.22,color=colors[variant],label=METHOD[variant])
        ax.axhline(0,c='#333333',lw=1);ax.axhline(.01,c='#B64745',ls='--',lw=1.4)
        ax.set_xticks(x,['EEG','HbO','HbR']);ax.set_title(label);ax.set_xlabel('相同原观测坐标的评分模态')
    axes[0].set_ylabel('平均 NRMSE 变化：候选−N0');axes[0].legend(fontsize=11)
    fig.tight_layout();return save(fig,out,'fidelity_tradeoffs')


def tradeoff_fig(out,groups,ds):
    from matplotlib.lines import Line2D
    frame=primary(groups,ds);frame=frame[(frame.partition=='evaluation')&(frame.model!='P0')&(frame.variant!='Hfull')]
    fig,ax=plt.subplots(figsize=(13,4.8));markers={'P1_tau':'o','P1_beta':'s','P2':'^','P3':'D'}
    for _,row in frame.iterrows():
        x=100*row.repeat_median;y=100*row.near_05_pair_rate
        baseline=row.variant=='N0';color=COLOR.get(row.variant,'#333333')
        ax.scatter(x,y,s=80+700*row.effective_pair_rate,
            facecolors='none' if baseline else color,marker=markers[row.model],edgecolors=color,linewidths=3 if baseline else 1.1)
    ax.axvline(20,c='#B64745',ls='--');ax.set_xlabel('被试层面重复差异中位数（%，越小越好）');ax.set_ylabel('至少一次近界的配对比例（%，越小越好）');ax.set_ylim(-4,105)
    models=[Line2D([],[],marker=marker,color='none',markerfacecolor='#718096',markeredgecolor='#718096',markersize=8,label=MODEL[model]) for model,marker in markers.items()]
    methods=[Line2D([],[],marker='o',color='none',markerfacecolor='none' if variant=='N0' else COLOR[variant],
        markeredgecolor=COLOR[variant],markersize=8,label=METHOD[variant]) for variant in sorted(set(frame.variant))]
    first=ax.legend(handles=models,title='形状：参数层次',loc='upper left',bbox_to_anchor=(1.02,1),fontsize=12,title_fontsize=13);ax.add_artist(first)
    ax.legend(handles=methods,title='颜色／填充：方法',loc='upper left',bbox_to_anchor=(1.02,.44),fontsize=12,title_fontsize=13)
    ax.set_xlim(-3,max(50,float(frame.repeat_median.max())*100+15));fig.subplots_adjust(left=.09,right=.76,bottom=.19,top=.94)
    return save(fig,out,'tradeoff_'+ds)


def information_fig(out,run):
    rows=[]
    for path in (run/'development').glob('*.json'):
        d=read(path)
        if d.get('model') not in ('P1_tau','P1_beta','P2','P3'):continue
        for f in d.get('fits',[]):
            info=f.get('information',{})
            if 'conditional_data' not in info:continue
            conditional=np.trace(info['conditional_data']['matrix'])
            if conditional<=0:continue
            rows.append(dict(dataset=d['dataset'],model=d['model'],data=np.trace(info['data_only']['matrix'])/conditional,
                             regularized=np.trace(info['nuisance_regularized']['matrix'])/conditional))
    frame=pd.DataFrame(rows);fig,axes=plt.subplots(1,3,figsize=(13,4.9),sharey=True)
    for ax,(ds,label) in zip(axes,DS.items()):
        models=['P1_tau','P1_beta','P2','P3'];g=frame[frame.dataset==ds].groupby('model').median(numeric_only=True)
        x=np.arange(4);ax.bar(x-.17,[g.loc[m,'data'] for m in models],.32,color='#087F8C',label='消去补偿量：仅数据')
        ax.bar(x+.17,[g.loc[m,'regularized'] for m in models],.32,color='#BD8120',label='加入原正则后的曲率')
        ax.axhline(1,c='#333333',ls='--',lw=1.5);ax.set_yscale('log');ax.set_xticks(x,['τ','β','τ/β','τ/β/κ']);ax.set_title(label);ax.set_xlabel('释放的参数')
    axes[0].set_ylabel('矩阵 trace / 条件数据 trace');axes[0].legend(fontsize=11);fig.tight_layout()
    shares=frame[frame.model=='P3'].groupby('dataset').data.median()
    return save(fig,out,'information'),shares


def gradient_fig(out,run):
    rows=[]
    for path in (run/'development').glob('*.json'):
        d=read(path)
        if d.get('model')!='P3':continue
        for f in d.get('fits',[]):
            if 'raw_parameter_gradient' not in f:continue
            for name,b,g,pg in zip(f['parameter_names'],f['log_parameter_boundary_distance'],f['raw_parameter_gradient'],f['projected_raw_parameter_gradient']):
                rows.append(dict(parameter=name,b=b,gradient=g,kkt=abs(pg)))
    frame=pd.DataFrame(rows);fig,axes=plt.subplots(1,3,figsize=(13,4.7),sharey=True)
    for ax,name in zip(axes,PARAMETERS):
        g=frame[frame.parameter==name];ax.scatter(g.b,np.maximum(g.kkt,1e-10),s=23,alpha=.5,c='#315B8A')
        ax.axvline(.05,c='#CB8235',ls='--');ax.axhline(.001,c='#B64745',ls='--');ax.set_yscale('log');ax.set_xlim(-.015,.51)
        ax.set_xlabel('原对数搜索范围内的边界距离 b');ax.set_title(NAME[name])
    axes[0].set_ylabel('原参数坐标投影梯度绝对值');fig.tight_layout();return save(fig,out,'boundary_kkt')


def inward_gradient_fig(out,run):
    rows=[]
    for path in (run/'development').glob('*__P3__N0.json'):
        d=read(path)
        for f in d['fits']:
            for j,name in enumerate(PARAMETERS):
                b=f['log_parameter_boundary_distance'][j]
                if b>=.1:continue
                fraction=np.log(f['parameters'][name]/BOUNDS[j,0])/np.log(BOUNDS[j,1]/BOUNDS[j,0])
                sign=1 if fraction<.5 else -1
                rows.append(dict(parameter=name,b=b,inward=sign*f['raw_parameter_gradient'][j]))
    frame=pd.DataFrame(rows);fig,axes=plt.subplots(1,3,figsize=(13,4.7))
    for ax,name in zip(axes,PARAMETERS):
        g=frame[frame.parameter==name];ax.scatter(g.b,g.inward,s=32,alpha=.6,c='#315B8A')
        ax.axhline(0,c='#333333',ls='--');ax.axvline(.05,c='#CB8235',ls=':')
        ax.set_yscale('symlog',linthresh=.001);ax.set_xlim(-.002,.102);ax.set_title(NAME[name]);ax.set_xlabel('距最近边界的对数比例 b')
    axes[0].set_ylabel('向内方向导数（原目标 / 参数单位）');fig.tight_layout()
    return save(fig,out,'inward_gradient')


def intersection_sensitivity_fig(out,profiles,objective):
    fig,axes=plt.subplots(2,2,figsize=(13,5.5));models=['P1_tau','P1_beta','P2','P3']
    for ax,model in zip(axes.ravel(),models):
        own=profiles[(profiles.model==model)&(profiles.objective==objective)]
        count=np.zeros((3,3));den=np.zeros((3,3))
        for i,inner in enumerate([.01,.05,.10]):
            for j,tolerance in enumerate([.005,.01,.05]):
                g=own[(own.inner_fraction==inner)&(own.tolerance==tolerance)]
                count[i,j]=g.exists.sum();den[i,j]=len(g)
        im=ax.imshow(count/np.maximum(den,1),vmin=0,vmax=1,cmap='YlGnBu',aspect='auto')
        for i in range(3):
            for j in range(3):ax.text(j,i,f'{int(count[i,j])}/{int(den[i,j])}',ha='center',va='center',fontsize=16,color='white' if count[i,j]/max(den[i,j],1)>.55 else '#182D40')
        ax.set_xticks([0,1,2],['0.5%','1%','5%']);ax.set_yticks([0,1,2],['1%','5%','10%']);ax.set_title(MODEL[model])
        ax.set_xlabel('每记录目标容差 / 固定P0目标');ax.set_ylabel('离边界的最小比例');ax.grid(False)
    fig.subplots_adjust(left=.08,right=.88,bottom=.12,top=.91,wspace=.31,hspace=.85)
    cax=fig.add_axes([.91,.16,.02,.69]);bar=fig.colorbar(im,cax=cax);bar.set_label('找到共同内点的开发组比例')
    return save(fig,out,'intersection_sensitivity_'+('obs' if objective.startswith('observation') else 'current'))


def profile_fig(out,run,ds,objective='regularized'):
    files=sorted((run/'profiles').glob(ds+'*__P1_tau__'+objective+'.json'))
    if not files:return None,None
    # The median fixed-P0 reconstruction selects the illustrative subject;
    # neither parameter values nor successful intersections select the example.
    candidates=[]
    for path in files:
        d=read(path);ref=read(run/'development'/(d['pair_id']+'__P0__N0.json'))
        candidates.append((float(np.mean([f['nrmse'] for f in ref['fits']])),path))
    center=np.median([x[0] for x in candidates]);path=min(candidates,key=lambda x:abs(x[0]-center))[1]
    first=read(path);pair=first['pair_id'];fig,axes=plt.subplots(1,2,figsize=(13,5.0))
    for ax,model,name in zip(axes,['P1_tau','P1_beta'],['tau','neurovascular_gain']):
        d=read(run/'profiles'/(pair+'__'+model+'__'+objective+'.json'));values=np.array([r['values'][0] for r in d['rows']])
        costs=np.array([[f.get('objective',np.nan) for f in r['fits']] for r in d['rows']]);den=np.asarray(d['intersections'][0]['reference_objectives'])
        minima=np.nanmin(costs,axis=0)
        for side,ref in enumerate(d.get('independent_reference',[])):minima[side]=min(minima[side],ref['objective'])
        relative=100*(costs-minima)/den
        for side,color,label in [(0,'#315B8A','片段 A'),(1,'#CA7A31','片段 B')]:
            ax.plot(values,relative[:,side],marker='o',ls='-' if side==0 else '--',c=color,label=label,ms=4,lw=1.7)
            failed=np.array([not row['fits'][side].get('accepted',False) for row in d['rows']])
            ax.scatter(values[failed],relative[failed,side],c='#111111',marker='x',s=44,linewidths=1.4,zorder=4,
                label='未通过数值检查' if side==0 else None)
        ax.axhline(1,c='#008879',ls='--',label='参考目标的 1%');ax.axhline(5,c='#7A8D98',ls=':',label='参考目标的 5%')
        bounds=np.array(d['bounds'])[0];inside=np.exp(np.log(bounds[0])+np.array([.05,.95])*np.log(bounds[1]/bounds[0]))
        for v in inside:ax.axvline(v,c='#AC5A91',ls=':',lw=1.4)
        ax.set_xscale('log');ax.set_yscale('symlog',linthresh=1);ax.set_ylim(bottom=-.1);ax.set_title(NAME[name]+' 的完整补偿量剖面')
        ax.set_xlabel(NAME[name]+('（秒）' if name=='tau' else '（冻结坐标中的有效增益）'));ax.set_ylabel('Δ目标 / 冻结 P0 目标（%）')
    handles,labels=axes[0].get_legend_handles_labels();fig.legend(handles,labels,loc='upper center',ncol=3,fontsize=12)
    fig.tight_layout(rect=(0,0,1,.83));return save(fig,out,'profiles_'+ds+'_'+objective),pair


def waveform_fig(out,run,pair,model,variant,side):
    co=read(run/'coordinates.json')[pair['dataset']+'__'+pair['site']]
    with np.load(run/'prepared'/(pair['id']+'.npz')) as a:
        targets=[np.column_stack((a[f'eeg_{s}'].reshape(-1,30)@co['pc']*co['eeg_factor'],a[f'hb_{s}'].reshape(-1,2)*co['hb_factor'])) for s in range(2)]
    with np.load(run/'fits'/(pair['id']+'__'+model+'__'+variant+'.npz')) as a:best=[a[f'prediction_{s}'].reshape(-1,3) for s in range(2)]
    with np.load(run/'fits'/(pair['id']+'__'+model+'__N0.npz')) as a:base=[a[f'prediction_{s}'].reshape(-1,3) for s in range(2)]
    fig,axes=plt.subplots(3,1,figsize=(13,5.2),sharex=True);clock=np.arange(360)*.25;s=side
    for j,label in enumerate(['EEG','HbO','HbR']):
        ax=axes[j];ax.plot(clock,targets[s][:,j]/co['sd'][j],c='#263238',lw=1.9,label='观测')
        ax.plot(clock,base[s][:,j]/co['sd'][j],c='#315B8A',ls='--',lw=1.5,label='N0 完整预测')
        if variant!='N0':ax.plot(clock,best[s][:,j]/co['sd'][j],c='#008879',ls='-.',lw=1.6,label=METHOD.get(variant,variant)+'完整预测')
        for boundary in (30,60):ax.axvline(boundary,c='#A0AAB2',ls=':',lw=1)
        ax.set_ylabel(label+' / SD')
        if j==2:ax.set_xlabel('该 90 s 片段内时间（秒）')
    axes[0].legend(fontsize=12,loc='upper left',ncol=3);fig.tight_layout()
    metrics=[]
    for j,label in enumerate(['EEG','HbO','HbR']):
        observed=targets[s][:,j]/co['sd'][j];predicted=best[s][:,j]/co['sd'][j]
        oi=int(np.argmax(abs(observed)));pi=int(np.argmax(abs(predicted)))
        metrics.append(dict(component=label,nrmse=float(np.sqrt(np.mean((predicted-observed)**2))),
            correlation=float(np.corrcoef(observed,predicted)[0,1]),bias=float(np.mean(predicted-observed)),
            observed_extreme=float(observed[oi]),observed_time=float(clock[oi]),predicted_extreme=float(predicted[pi]),predicted_time=float(clock[pi])))
    return save(fig,out,'waveform_'+pair['dataset']+'_'+model+'_'+str(side)),metrics


def build(run,out):
    setup();summary=read(run/'summary.json');plan=read(run/'pair_plan.json');selection=read(run/'frozen_selection.json')
    if summary['execution']!='completed':raise ValueError('analysis or numerical verification is incomplete')
    groups=pd.DataFrame(summary['groups']);pairs=pd.read_csv(run/'pair_metrics.csv');parameters=pd.read_csv(run/'parameter_estimates.csv')
    deck=Deck(out)
    deck.methods('当前数据的生理参数：稳定性与可达性',
        '固定六状态、单驱动 SSM｜E1–E6 实测与已知真值校准｜2026-10-10',[
        ('本轮回答什么','是否存在两次记录都能接受的共同内点？求解是否充分？连续性、权重或共享能否改善独立重复？'),
        ('数据和分工','三公开数据集；18 名开发被试决定坐标与超参数，53 名候选评价被试用于冻结后的比较。'),
        ('怎样定义“最好”','在预先冻结的候选中比较内点、重复差异、逐模态拟合及参数转移，同时要求保留真实个体差异。'),
        ('结论口径',summary.get('headline_conclusion','重建误差、参数重复性和生理真实性分别判断。实测没有个体参数真值，固定或完全共享不能证明个体恢复。'))],
        str(run.relative_to(ROOT))+'/{summary.json,pair_plan.json}')
    assessment=summary.get('evaluation_assessment',[])
    if assessment:
        overview=[]
        for ds,label in DS.items():
            own=[r for r in assessment if r['dataset']==ds]
            r=min(own,key=lambda x:(-x['effective_pairs']/x['planned_pairs'],x['best_repeat_median'],x['model']))
            overview.append([label,'同记录块' if ds=='simultaneous_eeg_nirs' else '独立记录',
                MODEL[r['model']],METHOD[r['descriptive_best_variant']],pct(r['best_repeat_median']),
                pct(r['best_near_05_pair_rate']),f"{r['effective_pairs']}/{r['planned_pairs']}"])
        deck.table('先看实测结论：稳定、内点与拟合需同时成立',
            '跨参数层次的描述性排名：先取有效配对率最高，再取重复差异中位最低；评价结果不用于重新选择方法',
            ['数据集','重复来源','参数层次','观察最佳方法','重复中位','近界配对','有效配对'],overview,
            [1.9,1.6,1.8,2.7,1.65,1.6,1.4],
            '重复差异＝各自由参数2|A−B|/(A+B)的最大值；近界＝任一参数距原log边界<5%。有效配对还要求重复≤20%、数值及拟合/转移通过。',
            'summary.json：evaluation_assessment；完整指标与逐模型对照见后续方法/结果页')
    inventory=[]
    for ds,label in DS.items():
        ps=[p for p in plan['pairs'] if p['dataset']==ds]
        primary_pairs=[p for p in ps if p['partition']=='evaluation' and p['repeat_kind']=='independent_native_record']
        auxiliary_pairs=[p for p in ps if p['partition']=='evaluation' and p['repeat_kind']=='within_record_disjoint_blocks']
        inventory.append([label,len(plan['development_subjects'][ds]),len(plan['candidate_subjects'][ds]),
            sum(p['partition']=='development' for p in ps),sum(p['partition']=='evaluation' for p in ps),
            f"{len({p['subject'] for p in primary_pairs})}/{len(primary_pairs)}",f"{len({p['subject'] for p in auxiliary_pairs})}/{len(auxiliary_pairs)}"])
    primary_subjects=len({(p['dataset'],p['subject']) for p in plan['pairs'] if p['partition']=='evaluation' and p['repeat_kind']=='independent_native_record'})
    deck.table('71 名被试对应多组区域配对，不能当作更多被试',
        '“组”＝同一被试、同一 ROI／montage 的 A/B 配对；最后两列只统计评价身份，格式为人数/配对组数',
        ['数据集','开发人数','候选人数','开发组','评价组','评价独立人/组','评价同记录人/组'],inventory,[2.0,1.25,1.25,1.25,1.25,2.2,2.3],
        f'53只是候选身份数；独立原生记录的主评价共{primary_subjects}人。Simultaneous 全部为同记录块，只作辅助稳定性证据。',
        'pair_plan.json；canonical native record/event metadata')
    deck.methods('先冻结身份，再建立同一个观测坐标',
        '原 outer-fold 变换不沿用；候选评价被试不参与 PCA、尺度、工作协方差或共享参数拟合',[
        ('开发身份','直接复用原 diagnostic_plan 的 QC 分层名单。开发选择没有读取拟合误差或参数，三个数据集各 6 人。'),
        ('重复配对','匹配任务、条件标记计数比例、时长、Hb 通道和 EEG montage。不能确认 session 的拼接段归入同记录辅助评价。'),
        ('区域依赖','一名被试可以提供多个区域；Visual 的两个 Probe 是空间位置，不能配成两次重复。区间重采样以被试为单位。'),
        ('确认性边界','这些公开被试历史上曾被查看。本轮是冻结评价，不是全新的独立确认队列；未访问其他协议的受保护 split。')],
        'pair_plan.json；coordinates.json；docs/DATA_CONTRACT.md')
    deck.methods('输入是处理后的观测，不是生理真值',
        '每个真实连续 90 s 片段划为三块；A/B 处理相同；4 Hz 输出，120 点＝30 s',[
        ('EEG 路径','6 个真实邻近通道 × 5 个频带功率 → 对数 → 每块前 5 s 参考 → 开发 PCA 第一分量 → 开发尺度。'),
        ('Hb 路径','原生相对 HbO/HbR → 不作运动校正 → 原 30 s 滤波／重采样／前 5 s 参考。HbO 与 HbR 共用一个正比例因子。'),
        ('保持什么','原生通道、时间锚点、真实有限支持和单位来源保留。每段初始参考不是静息状态真值，也不是个体光学标定。'),
        ('评分坐标','所有方法都在相同的冻结观测坐标中评分。NRMSE 分母来自开发 SD，不能用本窗口标准差或加权损失替换。')],
        'prepared/*.json；coordinates.json；native_feature_operators；fit_feature_coordinate')
    deck.methods('六个状态的职责与真实计算路径',
        'r 在 0.25 s 网格上自由估计、区间零阶保持；RK4 积分其余五个状态',[
        ('神经血流驱动','状态为 [r,s,f,v,p,q]。r 是冻结坐标的有效驱动；s=df/dt，单位s⁻¹。f、v、p、q是归一化流量、体积及模型坐标的总/脱氧Hb。'),
        ('血流部分','ds/dt = βr − κs − γ(f−1)\ndf/dt = s\ndv/dt = (f−v^(1/α))/τ'),
        ('Hb 部分','dp/dt = [f−v^(1/α−1)p]/τ\ndq/dt = [fE(f)/E₀−v^(1/α−1)q]/τ\nE(f)=1−(1−E₀)^(1/f)。'),
        ('从状态到观测','EEG 读出为 r；HbO 为 P₀(p−1)−Q₀(q−1)，HbR 为 Q₀(q−1)。再通过同一分块观测处理算子。')],
        'shared_driver_rk4._slope/_integrate；native_model_operator')
    deck.table('逐层释放参数，联合模型按整个向量判断',
        '固定 α=E₀=0.32（无量纲）、γ=0.32 s⁻²；P₀=1、Q₀=0.35为模型Hb尺度，均未经个体标定',
        ['层次','自由生理参数','原搜索范围','解释与作用'],[
        ['P0','无','τ=2、β=1、κ=0.64','固定参数重建对照'],
        ['P1-τ','τ','0.5–8 s','条件于观测坐标的转运时间常数'],
        ['P1-β','β','0.1–10','冻结观测坐标中的有效驱动增益'],
        ['P2','τ、β','两者原范围不变','检查过去固定伴随参数的补偿'],
        ['P3','τ、β、κ','κ：0.2–1.5 s⁻¹','P2 信息与合成校准通过后才加入']],
        [1.0,1.8,3.0,5.7],
        '“搜索范围”不是人群正常区间。P2/P3 只要有一个自由参数漂移或近界，就不能宣称整个生理向量更稳定。',
        'resolved_config.yaml；dimension_decision.json')
    deck.methods('拟合目标含工程正则，不能冒充噪声似然',
        '原目标 Q＝观测平方误差＋驱动二阶平滑＋初态先验＋log-flow 软锚定',[
        ('观测项','Qobs = Σ[(预测−观测)/开发 SD]²。W1 另乘冻结工作精度矩阵；比较方法时仍报告原坐标 NRMSE。'),
        ('原正则','驱动：0.01·Σ(Δ²r)²/dt³。初态：100·||x₀−[0,1,1,1,1]||²。flow：dt·Σlog²(f)/log²(2)。'),
        ('新增共享','z=log θ；仅 H 臂增加 λΣ(z−μ)²/Ω。每组只加一次参数先验，不随三段窗口数量重复添加。'),
        ('数学合法域与数值','要求 f,v,p,q>0，Q₀q≤P₀p，氧提取合法；RK4 中间点也检查。基线每起点 600B 次单 trial 前向预算，B 为段数。')],
        'fit_nonlinear_shared_parameters；objective_components；resolved_config.yaml')
    deck.methods('连续性实验只改变初态重置与先验数量',
        '三臂看见同一 90 s 观测，使用同一组三块 30 s 处理算子',[
        ('C0a：原分段','3 段各有自由驱动及 5 个初态；生理参数跨三段共享。三段初态都施加原先验。'),
        ('C0b：减少先验','仍允许在 30、60 s 重置初态；只约束第一段初态。C0b−C0a 主要反映先验数量变化。'),
        ('C1：保持连续','整段只有 5 个起点初态；五个 ODE 状态跨30、60 s连续。r仍按网格自由更新，平滑不跨这两个原边界额外加项。'),
        ('正确比较方向','C1 对 C0b 减少自由度，其最优训练目标不应更小。需要看到重复性或转移收益，并保持原始拟合在容差内。')],
        'record_arrays；curvature_breaks；development_methods/*.json')
    deck.methods('把重复、拟合、转移和 ICC 分别计算',
        'A/B片段分别初始化和估计，只共用冻结的总体对象；独立原生记录与同记录块分层报告',[
        ('重复差异','d_j = 2|θA,j−θB,j|/(|θA,j|+|θB,j|)，范围0–200%。向量取各自由参数最大值；先被试内平均区域，再报被试中位数/P90。'),
        ('原始拟合','NRMSE_m = sqrt(mean_t[(ŷm−ym)²])/SDdev,m。先两记录平均、再被试内区域平均、最后被试等权平均。'),
        ('参数转移','固定 A 的生理参数去拟合 B，只重新估计 B 的驱动和初态；反向同做。它是条件重建，不是无条件预测。'),
        ('ICC 与区间','在 log 参数上按数据集、同一 ROI/montage 分别算绝对一致性 ICC(A,1)。95% 区间按被试 bootstrap，冻结训练对象不重拟合。')],
        'metric_contract in summary.json；pair_metrics.csv；parameter_estimates.csv')
    deck.methods('有效内点配对率保留所有登记配对作为分母',
        '单独离开边界、相关性变高或重复差异变小，都不等于生理参数恢复',[
        ('内点定义','b=min[(logθ−logL)/log(U/L), (logU−logθ)/log(U/L)]。主定义 b≥5%；同时报告精确贴边及 1%/10% 敏感性。'),
        ('数值与重复','A/B 两次均通过收敛及原参数投影梯度检查；每个自由参数都在内区；最差分量重复差异 ≤20%。'),
        ('拟合与转移','每条记录、每个模态的 NRMSE 恶化均 ≤0.01；转移也逐记录、逐模态通过。任一失败使该配对不通过。'),
        ('防止假稳定','固定和完全共享的个体差异为零，不计为有效个体恢复。合成恢复还检查误差中位≤10%、P90≤25%、斜率0.8–1.2。')],
        'result_metrics；resolved_config.yaml；pair_metrics.csv')
    deck.methods('合成校准保留同一采样与处理，噪声为已知设定',
        '每名虚拟被试的两次记录共享参数真值，但驱动相位、初态和观测噪声独立生成',[
        ('多个参数真值','参考值的log坐标加入SD=0、0.15、0.45的差异；非零差异层另设少量2.5%近界真值。未释放参数保持参考值。'),
        ('驱动与初态','驱动为4个随机相位正弦叠加，角频率0.08/0.25/0.6/1.4 rad/s，SD=0.025。s₀的SD为0.005；四个正初态的log SD为0.015。'),
        ('观测与噪声','90 s、4 Hz、相同分块处理算子。三个通道工作尺度均0.025，再加独立高斯噪声SD=0.00125，即归一化0.05；不宣称等同实测噪声。'),
        ('真实 oracle','生成用8子步，拟合用4子步。oracle只估生理参数，并获知真实驱动与初态；普通拟合独立重估补偿量，不使用个体真值。')],
        'synthetic_pair；oracle_fit；resolved_config.yaml；synthetic/*.json')
    generation=[]
    for model in ('P1_tau','P1_beta','P2','P3'):
        paths=list((run/'synthetic').glob(model+'__*.json'));ok=sum(read(p).get('status')=='completed' for p in paths)
        generation.append([MODEL[model],len(paths),ok,len(paths)-ok,2*ok])
    deck.table('多真值校准检验个体差异，生成失败单列',
        '每层次：32 名虚拟被试 × 3 种差异水平 × 2 次独立记录；参数真值仅供生成、评分及 oracle',
        ['参数层次','计划被试条件','可生成条件','生成/任务失败','可评分记录'],generation,[2.1,2.3,2.3,2.3,2.5],
        '零／小／大 log 差异 SD 为 0、0.15、0.45。近上界 β 的指定驱动可违反数学合法域；成功子集误差不代表整个参数盒可恢复。',
        'synthetic_manifest.json；synthetic/*.json；synthetic_summary.json')
    for model in ('P1_tau','P1_beta','P2','P3'):
        fig=synthetic_fig(out,summary,model)
        recovered=[r for r in summary['synthetic'] if r['model']==model and r['level']==2 and r['variant'] in ('C0a','C1')]
        names=[name for name in PARAMETERS if any(r['parameter']==name for r in recovered)]
        comparison='；'.join(NAME[name]+': '+pct(next(r['median_relative_error'] for r in recovered if r['parameter']==name and r['variant']=='C0a'))+'→'+pct(next(r['median_relative_error'] for r in recovered if r['parameter']==name and r['variant']=='C1')) for name in names)
        deck.figure(MODEL[model]+'：理想条件下检验误差与假稳定',
            '大个体差异条件；柱＝相对误差中位数，黑菱形＝P90；固定参考值是无个体差异的对照',fig,
            '分段→连续的误差中位数：'+comparison+'。正确生成模型下的恢复不能替代实测生理验证。',
            'synthetic_estimates.csv；synthetic_summary.json')
    syn=pd.DataFrame(summary['synthetic'])
    for level,description in [(0,'零个体差异'),(1,'较小个体差异'),(2,'较大个体差异')]:
        rows=[]
        for model in ('P1_tau','P1_beta','P2','P3'):
            for name in PARAMETERS:
                g=syn[(syn.model==model)&(syn.parameter==name)&(syn.level==level)]
                if not len(g):continue
                a=g[g.variant=='C0a'].iloc[0];b=g[g.variant=='C1'].iloc[0]
                rows.append([MODEL[model],NAME[name],pct(a.median_relative_error),pct(b.median_relative_error),
                    fmt(a.log_recovery_slope,2),fmt(b.log_recovery_slope,2),pct(a.convergence)])
        deck.table('合成校准：'+description+'下的恢复与收敛',
            '误差对每条记录的已知参数真值评分；斜率为被试两次均值的 log 估计对 log 真值回归',
            ['层次','参数','分段误差','连续误差','分段斜率','连续斜率','分段数值通过'],rows,
            [2.0,.8,1.8,1.8,1.8,1.8,2.0],
            '数值通过率以每层每水平计划64条记录为分母。零差异时斜率无定义；恒定参数在这一条件表现好不能证明差异恢复。',
            'synthetic_summary.json；synthetic_estimates.csv')
    calibration_rows=[]
    for r in summary.get('synthetic_information_calibration',{}).get('rows',[]):
        if r['level']!=2 or r['variant'] not in ('C0a','C1'):continue
        calibration_rows.append([r['model'].replace('P1_tau','P1τ').replace('P1_beta','P1β'),NAME[r['parameter']],METHOD[r['variant']],
            r['accepted_inner_subjects'],f"{r['mean_actual_squared_repeat']:.2e}",f"{r['mean_linearized_repeat_variance']:.2e}",fmt(r['mean_standardized_squared_repeat'],2)])
    for j in range(0,len(calibration_rows),7):
        deck.table('合成信息近似能否解释实际重复波动',
            '大差异条件；仅纳入两次均通过数值检查且全参数在5%内区的被试；d＝logθA−logθB',
            ['模型','参数','方法','可用人数','平均d²','平均工作V','平均d²/V'],calibration_rows[j:j+7],
            [1.1,.65,2.1,1.3,2.35,2.35,2.4],
            'V=0.05²·diag(I⁻¹A+I⁻¹B)，I为有效数据曲率。理想局部线性近似预期平均d²/V约1；正则偏差与非线性可使它失准。',
            'summary.json：synthetic_information_calibration；synthetic/*.json')
    baseline=[]
    for ds,label in DS.items():
        for model in ('P1_tau','P1_beta','P2','P3'):
            g=groups[(groups.partition=='development')&(groups.dataset==ds)&(groups.model==model)&(groups.variant=='N0')]
            own=pairs[(pairs.partition=='development')&(pairs.dataset==ds)&(pairs.model==model)&(pairs.variant=='N0')]
            subject=own.groupby('subject').mean(numeric_only=True)
            baseline.append([label,MODEL[model],len(own),pct(subject.repeat_error.median()),pct(own.near_05.mean()),pct(own.accepted.mean())])
    for j in (0,6):
        deck.table('实测基线：释放更多参数不保证更稳定',
            '开发集描述；近界比例按 A/B 至少一次近界的配对计，数值通过要求两次都通过',
            ['数据集','自由参数','登记组','重复中位','近界配对','数值通过'],baseline[j:j+6],[2,2.1,1.2,2.1,2.1,2.0],
            '重复统计先在被试内平均区域，失败仍保留在登记分母。Simultaneous 是同记录辅助结果，不与独立记录混称重测。',
            'pair_metrics.csv；development/*.json')
    infofig,shares=information_fig(out,run)
    deck.figure('参数信息会被驱动与初态的补偿抵消',
        'A 为 log 参数灵敏度，B 为补偿量灵敏度；条件信息 AᵀA，有效数据项 Aᵀ(I−BB⁺)A',infofig,
        f'柱为开发记录 trace 比值中位数，虚线为条件参照。P3 三数据集仅保留 {pct(shares.min())}–{pct(shares.max())}；正则后曲率不是额外观测信息，也不是校准 Fisher 信息。',
        'development/*.json：information；fit_nonlinear_shared_parameters')
    deck.figure('近界解与求解不充分是两种不同问题',
        '每点＝开发记录中的一个 P3 参数；竖线 b=5%，横线原参数投影梯度 10⁻³',gradient_fig(out,run),
        '边界上的向外下降方向满足约束 KKT；向内部仍能下降才支持求解不足。图底部 10⁻¹⁰ 是零值显示下限，不是新的收敛门槛。',
        'development/*.json：raw_parameter_gradient / projected_raw_parameter_gradient')
    deck.figure('向内下降还是向外下降，需要看梯度符号',
        '每点＝P3 开发记录的一个近界参数（b<10%）；靠下界取 +∂Q/∂θ，靠上界取 −∂Q/∂θ',inward_gradient_fig(out,run),
        '负值表示向内仍可下降；正值表示局部向外更有利。各参数纵轴独立缩放；b>0 的近界点仍属内点，不能直接当成精确约束边界。',
        'development/*__P3__N0.json：raw_parameter_gradient；original parameter bounds')
    directions=[]
    for ds,label in DS.items():
        candidates=[]
        for path in sorted((run/'development').glob(ds+'*__P3__N0.json')):
            d=read(path)
            if d['pair_id'] not in plan['profile_panel']:continue
            ref=read(run/'development'/(d['pair_id']+'__P0__N0.json'))
            candidates.append((float(np.mean([f['nrmse'] for f in ref['fits']])),d))
        center=np.median([v for v,_ in candidates]);d=min(candidates,key=lambda item:abs(item[0]-center))[1]
        for side,f in zip(['A','B'],d['fits']):
            eigen,vectors=np.linalg.eigh(f['information']['data_only']['matrix']);v=vectors[:,0]
            if v[np.argmax(abs(v))]<0:v=-v
            directions.append([label,side,fmt(eigen[0]/eigen[-1],4),fmt(v[0],2),fmt(v[1],2),fmt(v[2],2)])
    deck.table('最弱信息方向是参数的组合，不能逐项孤立解释',
        '每个数据集取剖面面板内 P0 误差最接近中位数的被试；矩阵只含消去补偿量后的数据项',
        ['数据集','记录','最小/最大特征值','Δlogτ','Δlogβ','Δlogκ'],directions,[2.0,1.0,3.1,1.8,1.8,1.8],
        '后三列为最小特征值单位特征向量；最大绝对分量定为正。沿此组合改变参数，局部数据约束最弱；局部曲率不能替代完整剖面。',
        'development/*.json：information.data_only.matrix；pair_plan.profile_panel')
    deck.methods('共同内点检验重新优化所有补偿量',
        'QᵖA(θ)=minηA Q(A,θ,ηA)，B 同理；η 包含该记录全部驱动与自由初态',[
        ('剖面与条件切片','每个固定参数节点重新拟合补偿量；P2/P3 在完整向量节点固定所有自由生理参数。固定原驱动的切片不用于可达性判断。'),
        ('两类目标','原目标保留驱动、初态和 flow 正则，去掉新增参数共享。纯观测目标去掉这些正则，只保留原数学合法域。'),
        ('共同容许集合','每条记录的目标增量 ≤0.5%、1% 或5% 的冻结 P0 目标；每模态 NRMSE 增量≤0.01；共同参数还须处于内区。'),
        ('有限搜索边界','双向 warm start，单参数先9节点再加密至25；联合模型做有限网格。找到可行节点是存在性证据，未找到不证明不存在。')],
        'profiles/*.json；profile_pair；profile_intersections')
    deck.methods('共享参数与进入内区是两笔不同的代价',
        '以下最小值都是当前有限搜索中的已找到值；不宣称全局最优',[
        ('先分别选择参数','独立参照＝min QᵖA＋min QᵖB。两记录可以各选自己的参数；原目标参照同时保留 N0 的更低值。'),
        ('再要求一个共同向量','共同最小值＝minθ[QᵖA(θ)＋QᵖB(θ)]。θ 相同，驱动和初态仍为两份记录各自拟合。'),
        ('两项差值怎么计算','Δshare＝共同最小值−独立参照。\nΔinner＝仅在5%内区的共同最小值−全搜索范围的共同最小值。'),
        ('分母、限制与跨度','代价表除以 Q0,A＋Q0,B；容差仍逐记录计算，并逐模态查误差。可行跨度是最宽参数方向在原 log 搜索范围中的占比，不是置信区间。')],
        'profile_intersections；summary.json：delta_share_reference_fraction / delta_inner_reference_fraction')
    optimizer_rows=[]
    for item in summary.get('profile_optimization',[]):
        n=item['record_node_fits'];accepted=item['accepted_record_node_fits']
        optimizer_rows.append([MODEL[item['model']],
            '原工程目标' if item['objective']=='current_without_parameter_shrinkage' else '纯观测目标',
            f"{item['completed_groups']}/{item['registered_groups']}",accepted,n-accepted,pct(accepted/n)])
    if optimizer_rows:
        deck.table('先检查每个剖面节点的补偿量是否优化充分',
            '一个节点拟合＝一个参数向量下的一条记录；同组节点相互依赖，不能当作独立样本',
            ['模型','目标','完成组数','数值通过节点','未通过节点','通过比例'],optimizer_rows,[1.65,2.35,1.7,2.4,2.2,2.3],
            '未通过节点保留目标与失败信息，但不作为共同可接受内点证据。纯观测目标的未通过节点会限制“未找到”结论的强度。',
            'summary.json：profile_optimization；profiles/*.json')
    for ds,label in DS.items():
        for objective,desc in [('regularized','原工程目标'),('obs','纯观测目标')]:
            fig,pair=profile_fig(out,run,ds,objective)
            if fig is None:continue
            deck.figure(label+'：'+desc+'的 A/B 参数剖面',
                '例子按 P0 重建误差最接近开发中位数选取；每节点独立优化补偿量；纵轴为对称 log 刻度',fig,
                '蓝实线/橙虚线为A/B，叉号为未通过数值检查的节点；横线是1%/5%容差，竖线是5%内区。两图纵轴独立缩放；曲线相近不等于唯一辨识。',
                f'profiles/{pair}__P1_*__{objective}.json')
    prof=pd.DataFrame(summary['profile_intersections'])
    for ds,label in DS.items():
        rows=[]
        for model in ('P1_tau','P1_beta','P2','P3'):
            g=prof[(prof.dataset==ds)&(prof.model==model)&(prof.inner_fraction==.05)&(prof.tolerance==.01)]
            current=g[g.objective=='current_without_parameter_shrinkage'];obs=g[g.objective=='observation_only_legal_domain']
            widths=[max(w) for w,n in zip(current.log_range_width,current.accepted_nodes) if n>=2 and all(finite(x) for x in w)]
            rows.append([model.replace('P1_tau','P1τ').replace('P1_beta','P1β'),f'{int(current.exists.sum())}/{len(current)}',f'{int(obs.exists.sum())}/{len(obs)}',
                pct(current.delta_share_reference_fraction.median(),2),pct(current.delta_inner_reference_fraction.median(),2),pct(np.median(widths)) if widths else '节点不足',
                int((current.sampled_classification=='acceptable_shared_nodes_only_outside_interior').sum()),int((current.accepted_nodes==1).sum())])
        deck.table(label+'：共同内点是否被实际找到',
            '主判据：5% 内区、每记录目标容差1%、每模态 NRMSE 容差0.01，两个补偿量优化均通过数值检查',
            ['层次','原目标可行','观测项可行','共享代价','入内代价','可行跨度','仅近界组','单点组'],rows,[1.2,1.6,1.6,1.6,1.6,1.7,1.45,1.4],
            '两项代价均除以两记录P0目标之和；总成本小仍须逐记录逐模态通过。仅近界＝可行节点均在内区外；单点不能确定宽度。',
            'profile_intersections.json；profiles/*.json')
    for objective,label in [('current_without_parameter_shrinkage','原目标'),('observation_only_legal_domain','纯观测目标')]:
        deck.figure(label+'：共同内点结论随容差怎样变化',
            '每格＝找到可行共同内点的组数/18个开发面板组；颜色使用统一0–1刻度，不是显著性或概率',
            intersection_sensitivity_fig(out,prof,objective),
            '向右放宽目标容差，向下扩大近界排除区；每格仍要求两个补偿量优化均通过、每记录每模态 NRMSE 增量≤0.01。',
            'summary.json：profile_intersections；pair_plan.profile_panel')
    numerical=pd.DataFrame(summary['numerical_reference']);numrows=[]
    for model,g in numerical.groupby('model'):
        numrows.append([MODEL[model],g.pair_id.nunique(),pct(g.objective_gain_reference_fraction.median(),2),
            pct(g.objective_gain_reference_fraction.max(),2),pct((g.accepted_nodes/g.total_nodes).mean()),
            int(g[g.trigger_N2].pair_id.nunique())])
    deck.table('E2：更充分搜索究竟改善了多少原目标',
        'N1＝基线与剖面搜索中已找到的最好值；只比较同一原目标，保留未通过数值检查的节点记录',
        ['层次','开发组数','目标改善中位','最大改善','节点数值通过','更好内点组'],numrows,[1.7,1.65,2.2,2.1,2.2,1.7],
        '改善以各记录冻结P0目标为分母。“更好内点”要求原近界解被内点改善超过1%；达到25%开发组才触发新求解器对照。',
        'solver_reference.csv；followup_decisions.json；profiles/*.json')
    continuity_data=pairs[(pairs.partition=='development')&(pairs.dataset=='simultaneous_eeg_nirs')&(pairs.model=='P2')]
    c0=continuity_data[continuity_data.variant=='N0'].groupby('subject').mean(numeric_only=True).repeat_error.median()
    c1=continuity_data[continuity_data.variant=='C1'].groupby('subject').mean(numeric_only=True).repeat_error.median()
    c1_pairs=continuity_data[continuity_data.variant=='C1'];preserved=int((c1_pairs.raw_fit_pass&c1_pairs.transfer_pass).sum())
    deck.figure('连续性降低部分重复误差，但仍须保住拟合',
        'P2 开发结果；灰线＝一名被试的区域平均，绿色＝被试中位数；三臂观测完全相同',continuity_fig(out,pairs),
        f'Simultaneous 重复中位{pct(c0)}→{pct(c1)}，但拟合与转移同时保住的配对为{preserved}/{len(c1_pairs)}。20%虚线为工程门槛；C0b改变的是初态先验数量。',
        'pair_metrics.csv：development / N0,C0b,C1')
    cal=read(run/'development_calibration.json');covrows=[]
    for key,v in cal.items():
        eig=np.linalg.eigvalsh(v['covariance']['covariance']);ds=key.split('__')[0];site=key[len(ds)+2:]
        short=site.replace('prefrontal','PF').replace('posterior','P').replace('motor','M').replace('__montage','')
        covrows.append([DS[ds]+' '+short,fmt(v['covariance']['ar1'],3),fmt(min(eig),3),fmt(max(eig),3),v['covariance']['n_samples']])
    for j in (0,6):
        deck.table('W1 使用冻结工作误差协方差',
            '源于开发数据内三折时间留出残差；每折隐藏各30 s块的连续10 s，三个模态一起隐藏',
            ['开发坐标','AR(1)','最小特征值','最大特征值','残差点数'],covrows[j:j+6],[4.2,1.4,2.1,2.1,1.7],
            '三模态协方差 trace 归一化为3，平均精度再归一化为1。强时间相关可能含模型失配；W1 加权损失不与 W0 直接排名。',
            'development_calibration.json；crossfit/*.json')
    deck.figure('重复收益需核查代价：加权臂的 EEG 误差上升',
        'P2 开发结果；柱为每名被试区域平均后的平均配对差；纵轴统一，正值表示原始重建变差',
        fidelity_tradeoff_fig(out,pairs),
        '虚线为0.01容差的数值参照；真实通过规则逐记录、逐模态执行，不能用总体平均抵消局部恶化。W1 加权损失未参与此比较。',
        'pair_metrics.csv：development / P2 / C1,W1,H4.0 versus N0')
    deck.methods('E5 的共享强度只由开发留出记录选择',
        'λ∈{0,0.25,1,4}；同分选更弱共享，完全共享只作端点对照',[
        ('训练中心 μ','对开发记录的原目标作 pooled 拟合，各记录驱动/初态独立。该中心是工程最优参照，不是正常生理人群均值。'),
        ('尺度 Ω','总体冻结尺度来自开发 A/B 的 log 参数交叉矩，并设0.01下限；留出记录选择时只使用源记录侧的工作尺度。'),
        ('选强度而不偷看','仅用 A 建共享对象和 A 的估计，固定参数去重建 B；再反向做。最终冻结对象才使用全部开发记录。'),
        ('数值与收缩边界','未通过数值检查的总体/源侧中心不参与非零强度选择。减少方差、向中心收缩本身都不是个体差异恢复成功。')],
        'development_calibration.json；frozen_selection.json')
    centers=[]
    for model in ('P1_tau','P1_beta','P2','P3'):
        working=[v['models'][model] for v in cal.values()]
        near=0
        for w in working:
            names=list(w['mean']);values=np.exp([w['mean'][n] for n in names]);bounds=np.asarray([BOUNDS[PARAMETERS.index(n)] for n in names])
            near+=int((boundary_distance(values,bounds)<.05).any())
        centers.append([MODEL[model],len(working),sum(w['pooled_fit']['accepted'] for w in working),
            sum(w['pooled_fit']['accepted'] and all(f['pooled_fit']['accepted'] for f in w['record_folds'].values()) for w in working),near])
    deck.table('共享中心本身也必须通过数值与边界检查',
        '每个坐标分别建立共享对象；不同数据集、区域、montage 不强行合并为同一个总体',
        ['模型','坐标数','总体中心通过','总体及两折均通过','中心有参数近界'],centers,[1.9,1.6,2.7,3.4,2.9],
        '未通过数值检查的中心禁止用于非零强度选择。完全共享仍作为端点诊断保留，其稳定性和重建误差不能证明个体恢复。',
        'development_calibration.json：models.*.pooled_fit / record_folds / mean')
    selrows=[]
    for model,entry in selection['models'].items():
        singles='、'.join(METHOD[v] for v in entry['singles']) or '无'
        strengths=list(entry['sharing_strengths'].values())
        selrows.append([MODEL[model],singles,'是' if entry['combine'] else '否',str({str(v):strengths.count(v) for v in sorted(set(strengths))})])
    screen=[]
    for model,entry in selection['models'].items():
        for r in entry['candidates']:
            screen.append([MODEL[model],METHOD[r['variant']],f"{100*r['repeat_gain']:+.1f}",
                pct(r['raw_pass_fraction']),pct(r['transfer_pass_fraction']),pct(r['convergence_pair_fraction']),
                '保留' if r['retain'] else '不保留'])
    for j in (0,6):
        deck.table('单项改进必须同时保住拟合与转移表现',
            '开发筛选：每名被试区域先平均，再对18名被试等权；重复收益＝N0差异−候选差异，正值较好',
            ['层次','方法','重复收益pp','拟合通过','转移通过','数值通过','决定'],screen[j:j+6],
            [1.75,2.8,1.55,1.55,1.55,1.55,1.35],
            '保留要求重复收益中位>1个百分点，拟合及转移通过率均≥80%。这只是开发筛选；数值通过率≥95%等绝对资格另判。',
            'frozen_selection.json：models.*.candidates')
    deck.table('方法在评价前冻结，评价后不再调参',
        '每个参数层次最多两项单因素改进及其组合；基线与完全共享端点同时评价',
        ['参数层次','保留的单项','组合','各坐标 λ 的数量分布'],selrows,[1.7,5.2,1.0,3.6],
        'λ分布来自开发内部选择；只有保留共享臂的层次在评价中使用它。保留不等于生理资格通过；后文“最佳”是冻结候选中的描述排名。',
        'frozen_selection.json：evaluation_results_read=0')
    selected=summary.get('selected_synthetic') or {}
    if selected.get('estimates'):
        extra=pd.DataFrame(selected['estimates']);vectors=pd.DataFrame(summary['synthetic_vector_recovery'])
        training=selected.get('training_centers',[]);accepted_centers=sum(t['accepted'] for t in training)
        deck.methods('共享后的稳定性必须接受独立真值校准',
            '共享强度沿用开发筛选保留的值；合成训练身份和评分身份分离，不把生成真值送入拟合器',[
            ('训练对象从哪里来','每个差异水平使用另外8名虚拟被试、各2次独立记录，拟合共享中心与对角工作尺度；原32名仅用于评分。'),
            ('中心数值检查',f'共{len(training)}个合成共享中心，{accepted_centers}个通过数值检查。未通过中心对应的结果保留为失败诊断，不能据其宣称共享校准成功。'),
            ('准确性与响应分别评价','每条记录取全部自由参数中最大的相对误差，再报告中位数/P90；参数响应斜率由被试重复平均估计与真值的log回归得到。'),
            ('零差异与非零差异','零差异下斜率和真值方差比没有定义；小、大差异下需同时看误差、0.8–1.2响应斜率、方差压缩与数值失败分母。')],
            'synthetic_selected_training/*.json；synthetic_selected_summary.json；synthetic_selected_estimates.csv')
        for model in sorted(set(extra.model)):
            names=[n for n in PARAMETERS if n in set(extra[extra.model==model].parameter)]
            for level,description in [(0,'零差异'),(1,'小差异'),(2,'大差异')]:
                center_ok=next(t['accepted'] for t in training if t['model']==model and t['level']==level)
                rows=[]
                baseline=syn[(syn.model==model)&(syn.variant=='C0a')&(syn.level==level)]
                sources=[('C0a',baseline)]+[(v,extra[(extra.model==model)&(extra.variant==v)&(extra.level==level)]) for v in sorted(set(extra[extra.model==model].variant))]
                for variant,g in sources:
                    own=g.set_index('parameter');v=vectors[(vectors.model==model)&(vectors.variant==variant)&(vectors.level==level)].iloc[0]
                    ratios=[own.loc[n,'between_log_variance']/own.loc[n,'truth_log_variance'] for n in names if own.loc[n,'truth_log_variance']>1e-10]
                    slopes=[fmt(own.loc[n,'log_recovery_slope'],2) if n in names else '固定' for n in PARAMETERS]
                    rows.append([METHOD.get(variant,variant.replace('_C0a','')),*slopes,pct(v.median_worst_parameter_error),pct(v.p90_worst_parameter_error),
                        f'{min(ratios):.2f}–{max(ratios):.2f}' if ratios else '无定义',pct(v.all_parameter_convergence)])
                deck.table(MODEL[model]+'：共享后的多真值恢复（'+description+'）',
                    '每个差异水平另用8名虚拟训练被试建立共享中心和尺度；评分仍用原32名独立虚拟被试',
                    ['方法','τ斜率','β斜率','κ斜率','误差中位','误差P90','方差比范围','数值通过'],rows,[2.1,1,1,1,1.5,1.5,2.4,1.8],
                    ('该水平共享中心未收敛，共享行仅作失败诊断。' if not center_ok else '')+
                    '误差取每记录最差参数；方差比＝估计/真值的被试间log方差。斜率0.8–1.2才满足响应目标。',
                    'synthetic_selected_summary.json；synthetic_vector_recovery in summary.json')
    best_rows=[]
    for ds,label in DS.items():
        table_rows=[]
        for model in ('P1_tau','P1_beta','P2','P3'):
            winner=best_group(groups,ds,model);best_rows.append(winner)
            baseline=primary(groups,ds);baseline=baseline[(baseline.partition=='evaluation')&(baseline.model==model)&(baseline.variant=='N0')].iloc[0]
            for row in [baseline,winner]:
                table_rows.append([model.replace('P1_tau','P1τ').replace('P1_beta','P1β'),METHOD.get(row.variant,row.variant),
                    pct(row.repeat_median),pct(row.repeat_p90),pct(row.near_05_pair_rate),
                    fmt(row.HbO_nrmse),fmt(row.HbR_nrmse),f'{int(row.effective_pairs)}/{int(row.planned_pairs)}'])
        deck.table(label+'：冻结候选中可达到的最好结果',
            ('同记录不重叠块的辅助评价' if ds=='simultaneous_eeg_nirs' else '不同原生记录的主评价')+'；各模型第一行基线，第二行按有效配对率优先排名的候选',
            ['模型','方法','重复中位','重复P90','近界配对','HbO误差','HbR误差','有效配对'],table_rows,
            [.75,2.45,1.4,1.35,1.4,1.35,1.35,1.45],
            '近界＝A/B至少一次有自由参数进入5%近界区；误差＝开发SD归一化NRMSE。有效配对还要求数值、重复、逐记录拟合及转移全部通过。',
            'summary.json：groups；pair_metrics.csv')
        deck.figure(label+'：稳定、内点与有效配对的折中',
            '每点＝一个冻结候选；点面积随有效配对率增加，形状对应自由参数层次；完全共享不参与个体排名',
            tradeoff_fig(out,groups,ds),
            '左下方较好，但仅靠图上位置仍不足以证明生理恢复；需同时核对原始拟合、参数转移、合成响应与未加参数先验的剖面。',
            'summary.json：evaluation groups')
        if ds=='visual_cognitive_motivation':
            auxiliary=groups[(groups.partition=='evaluation')&(groups.dataset==ds)
                &(groups.repeat_kind=='within_record_disjoint_blocks')&(~groups.model.eq('P0'))&(~groups.variant.eq('Hfull'))]
            rows=[[MODEL[r.model],METHOD[r.variant],int(r.independent_subjects),int(r.planned_pairs),
                pct(r.repeat_median),pct(r.near_05_pair_rate),f'{int(r.effective_pairs)}/{int(r.planned_pairs)}'] for _,r in auxiliary.iterrows()]
            deck.table('Visual 的同记录块保留为独立的辅助结果',
                '这些候选身份没有满足配对条件的两条独立原生记录；不把区域或窗口当作更多被试',
                ['模型','方法','人数','登记组','重复中位','近界配对','有效配对'],rows,[1.75,2.85,1.05,1.2,1.75,1.75,1.7],
                '本页与前页使用同一冻结坐标、方法和评分规则，但评价的是短时块稳定性。不得把这三名被试加入独立记录ICC的分母。',
                'summary.json：evaluation / Visual / within_record_disjoint_blocks')
    coordinate_groups=pd.DataFrame(summary.get('coordinate_groups',[]))
    if len(coordinate_groups):
        for ds,label in DS.items():
            rows=[]
            for model in ('P1_tau','P1_beta','P2','P3'):
                g=primary(coordinate_groups[(coordinate_groups.partition=='evaluation')&(coordinate_groups.model==model)
                    &(coordinate_groups.variant!='Hfull')],ds)
                r=g.sort_values(['effective_pair_rate','repeat_median','site','variant'],ascending=[False,True,True,True]).iloc[0]
                site=r.site.replace('prefrontal','前额').replace('motor','运动').replace('posterior','后部').replace('__montage',' M')
                rows.append([MODEL[model],site,METHOD[r.variant],int(r.independent_subjects),pct(r.repeat_median),
                    pct(r.near_05_pair_rate),f'{int(r.effective_pairs)}/{int(r.planned_pairs)}',pct(r.convergence_record_rate)])
            deck.table(label+'：区域平均是否掩盖了局部较好结果',
                '仅作事后描述：各参数层次按有效配对率优先、重复差异其次，列出观察最佳坐标及冻结方法',
                ['层次','区域/坐标','方法','人数','重复中位','近界配对','有效配对','记录数值通过'],rows,
                [1.65,1.9,2.6,.75,1.4,1.4,1.35,1.45],
                '同一坐标的被试才可直接比较；小人数及观察排序均限制推广。局部较好不能替代完整参数向量、真值响应与独立确认。',
                'summary.json：coordinate_groups；pair_metrics.csv')
    for ds,label in DS.items():
        winner=best_group(groups,ds,'P2');fig,site=repeat_scatter(out,parameters,ds,'P2',winner.variant)
        deck.figure(label+'：独立估计的 τ、β 是否一致',
            f'{site}；方法：{METHOD[winner.variant]}。每点＝该坐标的一名评价被试，黑虚线为A=B；对数坐标保留原搜索范围',fig,
            '三角表示至少一次近界，叉号表示数值检查未通过。P2 的判断要求 τ 与 β 同时达标，不能只选择散点更集中的参数。',
            'parameter_estimates.csv；frozen_selection.json')
    sensitivity=[]
    for ds,label in DS.items():
        for model in ('P1_tau','P1_beta','P2','P3'):
            r=best_group(groups,ds,model)
            sensitivity.append([label,model.replace('P1_tau','P1τ').replace('P1_beta','P1β'),METHOD[r.variant],
                pct(r.exact_boundary_pair_rate),pct(r.near_01_pair_rate),pct(r.near_05_pair_rate),pct(r.near_10_pair_rate)])
    for j in (0,6):
        deck.table('近界判断不依赖“是否恰好等于边界”',
            '冻结候选中的观察最佳结果；A/B任一次、任一自由参数满足条件即记该配对近界',
            ['数据集','模型','方法','精确贴边','1%近界','5%近界','10%近界'],sensitivity[j:j+6],
            [2.0,1.0,3.0,1.6,1.6,1.6,1.6],
            '精确贴边按 b<10⁻⁸ 的浮点容差判断。5%是本轮主定义；改变操作阈值不会把搜索范围变成人群正常区间。',
            'summary.json：groups；pair_metrics.csv')
    differences=pd.DataFrame(summary.get('paired_method_differences',[]))
    if len(differences):
        for model,entry in selection['models'].items():
            for variant in entry['singles']+(['combined'] if entry['combine'] else []):
                rows=[]
                for ds,label in DS.items():
                    g=primary(differences[(differences.partition=='evaluation')&(differences.model==model)&(differences.variant==variant)],ds)
                    for metric,name in [('repeat_error','重复差异'),('HbO_nrmse','HbO误差'),('HbR_nrmse','HbR误差')]:
                        r=g[g.metric==metric].iloc[0];lo,hi=r['interval']
                        rows.append([label,name,r.available_subjects,fmt(r.mean_subject_difference,3),f'[{fmt(lo)}, {fmt(hi)}]'])
                deck.table(MODEL[model]+'：候选相对基线的配对差值',
                    METHOD[variant]+'；差值＝候选−N0，负值较好；先区域平均，再对被试均值差作配对 bootstrap',
                    ['数据集','指标','被试数','平均差值','95%区间'],rows,[2.1,2.1,1.5,2.3,4.5],
                    '重复差异使用比例值，0.01为1个百分点；误差使用开发SD单位。区间条件于冻结训练对象，未作多重比较校正，不宣称确认性优势。',
                    'summary.json：paired_method_differences；pair_metrics.csv')
    distributions=pd.DataFrame(summary['parameter_distributions'])
    for model,entry in selection['models'].items():
        if 'Hselected' not in entry['singles']:continue
        for key,strength in entry['sharing_strengths'].items():
            if strength==0:continue
            ds,site=key.split('__',1);fig,comparison=sharing_scatter_fig(out,parameters,ds,site,model)
            deck.figure(DS[ds]+'：共享改变了哪些 A/B 参数估计',
                f'{site}，λ={strength:g}；每个点是一名被试的一种方法；灰线连接同一被试，黑虚线A=B；原搜索范围统一',fig,
                '各分量重复差异中位：'+comparison+'。空心标记按当前分量判断；整个向量仍要求所有参数一起通过。',
                'parameter_estimates.csv；frozen_selection.json；evaluation/*.json')
    information_budget=pd.DataFrame(summary.get('sharing_information_budget',[]));budget_rows=[]
    if len(information_budget):
        for ds,label in DS.items():
            for _,r in primary(information_budget,ds).iterrows():
                site=r.site.replace('prefrontal','前额').replace('motor','运动').replace('__montage',' M')
                budget_rows.append([label+' '+site,NAME[r.parameter],r.strength,int(r.available_subjects),
                    pct(r.data_fraction),pct(r.original_regularization_fraction),pct(r.added_parameter_prior_fraction)])
        for start in range(0,len(budget_rows),5):
            deck.table('共享后的局部约束有多少来自新增先验',
                '仅接受记录、λ>0坐标；各log参数方向的有效Gauss–Newton对角曲率，先记录平均，再被试等权',
                ['数据集/坐标','参数','λ','人数','数据份额','原正则增量','新增共享份额'],budget_rows[start:start+5],
                [3.3,.7,.7,.8,1.85,2.2,2.95],
                '总曲率＝原正则下有效曲率＋λ/Ω；原正则增量＝前者−纯数据有效曲率。这些局部份额不表示全局辨识概率或生理真实性。',
                'summary.json：sharing_information_budget；evaluation/*.json；development_calibration.json')
    shrinkage_rows=[]
    for model,entry in selection['models'].items():
        if 'Hselected' not in entry['singles']:continue
        for key,strength in entry['sharing_strengths'].items():
            if strength==0:continue
            ds,site=key.split('__',1)
            chosen=primary(distributions[(distributions.partition=='evaluation')&(distributions.site==site)&(distributions.model==model)],ds)
            baseline=chosen[chosen.variant=='N0'].set_index('parameter')
            own=chosen[chosen.variant=='Hselected'].set_index('parameter')
            for name,r in own.iterrows():
                a=baseline.loc[name];ratio=r.between_log_variance/a.between_log_variance if a.between_log_variance>1e-12 else np.nan
                estimates=primary(parameters[(parameters.partition=='evaluation')&(parameters.site==site)&(parameters.model==model)&(parameters.parameter==name)],ds)
                old=estimates[estimates.variant=='N0'].repeat.median();new=estimates[estimates.variant=='Hselected'].repeat.median()
                short=site.replace('prefrontal','前额').replace('motor','运动').replace('__montage',' M')
                shrinkage_rows.append([DS[ds]+' '+short,NAME[name],strength,r.available_subjects,fmt(ratio,2),pct(old),pct(new)])
    for j in range(0,len(shrinkage_rows),5):
        deck.table('共享是否压低了实测中的个体差异',
            '仅显示冻结λ>0的坐标；每个参数单列；方差来自被试A/B的log参数平均，包含未收敛的有限估计',
            ['数据集/坐标','参数','λ','人数','方差比','原重复差异','共享后差异'],shrinkage_rows[j:j+5],
            [3.4,.65,.75,.85,1.4,2.3,2.95],
            '方差比＝共享后/N0，重复差异为该坐标被试中位数。方差变小本身不证明真实差异被保留；需与多真值恢复斜率共同判断。',
            'summary.json：parameter_distributions；parameter_estimates.csv；frozen_selection.json')
    for ds,label in DS.items():
        rows=[];nullrows=[]
        for model in ('P1_tau','P1_beta','P2','P3'):
            win=best_group(groups,ds,model)
            g=primary(distributions[(distributions.partition=='evaluation')&(distributions.model==model)],ds)
            site=sorted(s for s in set(g.site) if s.startswith('prefrontal'))[0]
            for _,r in g[(g.site==site)&(g.variant==win.variant)].iterrows():
                shared=g[(g.site==site)&(g.variant=='Hfull')&(g.parameter==r.parameter)].iloc[0]
                rows.append([model.replace('P1_tau','P1τ').replace('P1_beta','P1β'),METHOD[win.variant],NAME[r.parameter],fmt(r.median_A),fmt(r.median_B),
                    f'{fmt(r.geometric_AB_p10)}–{fmt(r.geometric_AB_p90)}',fmt(shared.median_geometric_AB)])
        deck.table(label+'：目前能给出的具体参数数值',
            '固定前额坐标；A/B 中位数分别计算；区间为被试 A/B 几何均值的 P10–P90，包含未收敛有限解',
            ['模型','方法','参数','A 中位','B 中位','个体 P10–P90','完全共享值'],rows,[1.1,2.4,.7,1.2,1.2,3.4,2.4],
            'τ 单位秒，κ 单位 s⁻¹；β 是冻结坐标中的有效增益。这里是条件拟合分布，不是正常区间、精度区间或真实生理范围。',
            'summary.json：parameter_distributions；parameter_estimates.csv')
        for model in ('P0','P1_tau','P1_beta','P2','P3'):
            variants=['N0'] if model=='P0' else ['N0','Hfull']
            for variant in variants:
                r=primary(groups[(groups.partition=='evaluation')&(groups.model==model)&(groups.variant==variant)],ds).iloc[0]
                nullrows.append([MODEL[model],METHOD[variant],fmt(r.EEG_nrmse),fmt(r.HbO_nrmse),fmt(r.HbR_nrmse),
                    fmt(r.HbO_transfer),fmt(r.HbR_transfer)])
        deck.table(label+'：固定参数也能重建到什么程度',
            '开发SD归一化误差；先记录平均、被试内区域平均、再被试等权平均；转移重新估计目标记录驱动与初态',
            ['模型','方法','EEG','HbO','HbR','转移HbO','转移HbR'],nullrows,[1.8,2.0,1.65,1.65,1.65,1.8,1.8],
            '完全共享的重复差异为零但不提供个体差异。误差接近不等于参数准确，也可能反映补偿量可以解释大部分观测。',
            'summary.json：groups；evaluation/*.json')
    iccs=pd.DataFrame(summary['icc'])
    for ds,label in DS.items():
        rows=[]
        for model in ('P1_tau','P1_beta','P2','P3'):
            win=best_group(groups,ds,model);g=primary(iccs[(iccs.partition=='evaluation')&(iccs.model==model)&(iccs.variant==win.variant)],ds)
            sites=sorted(s for s in set(g.site) if s.startswith('prefrontal'));g=g[g.site==sites[0]]
            for _,r in g.iterrows():
                interval=r['interval'];rows.append([model,NAME[r.parameter],METHOD[r.variant],f'{r.paired_accepted_subjects}/{r.planned_subjects}',fmt(r.icc,2),f'[{fmt(interval[0],2)}, {fmt(interval[1],2)}]'])
        deck.table(label+'：ICC 需结合配对数和区间阅读',
            '仅展示固定前额坐标；log 参数绝对一致性 ICC(A,1)，95% 被试百分位 bootstrap 区间',
            ['模型','参数','方法','可用/登记','ICC','95%区间'],rows,[1.0,.85,4.1,1.5,1.2,2.9],
            'ICC 使用两次均通过数值检查的配对；失败分母另由有效配对率保留。区间条件于冻结训练对象，未宣称多候选显著性优势。',
            'summary.json：icc；parameter_estimates.csv')
    # Median baseline error, fixed prefrontal coordinate, fixed P2 example rule.
    for ds,label in DS.items():
        g=primary(pairs[(pairs.partition=='evaluation')&(pairs.model=='P2')&(pairs.variant=='N0')],ds)
        sites=sorted(s for s in set(g.site) if s.startswith('prefrontal'));g=g[g.site==sites[0]]
        center=g[['HbO_nrmse','HbR_nrmse']].mean(axis=1).median();idx=(g[['HbO_nrmse','HbR_nrmse']].mean(axis=1)-center).abs().idxmin()
        ident=g.loc[idx,'pair_id'];pair=next(p for p in plan['pairs'] if p['id']==ident);win=best_group(groups,ds,'P2')
        detail=read(run/'evaluation'/(ident+'__P2__'+win.variant+'.json'))
        wave_rows=[]
        for side in range(2):
            fig,metrics=waveform_fig(out,run,pair,'P2',win.variant,side);f=detail['fits'][side]
            worst=max(metrics,key=lambda m:m['nrmse'])
            text=f"误差最大的是{worst['component']}（NRMSE={fmt(worst['nrmse'])}）；观测主极值{fmt(worst['observed_extreme'],2)}SD@{fmt(worst['observed_time'],1)}s，预测主极值{fmt(worst['predicted_extreme'],2)}SD@{fmt(worst['predicted_time'],1)}s。"
            deck.figure(label+'：典型片段 '+('A' if side==0 else 'B')+' 的观测与完整重建',
                '例子：前额坐标、基线Hb误差最接近评价中位数；灰线为30 s边界；每个模态独立纵轴缩放',fig,
                text,
                f'evaluation/{ident}__P2__{win.variant}.json；fits/*.npz')
            for row in metrics:
                wave_rows.append([('A' if side==0 else 'B')+' '+row['component'],fmt(row['nrmse']),fmt(row['correlation'],2),fmt(row['bias']),
                    f"{fmt(row['observed_extreme'],2)} @ {fmt(row['observed_time'],1)}",f"{fmt(row['predicted_extreme'],2)} @ {fmt(row['predicted_time'],1)}"])
        deck.table(label+'：核验典型曲线的幅度、偏移与时间',
            '同一对示例；主极值＝绝对幅值最大的有符号点，格式为开发SD单位的幅度 @ 片段内秒数',
            ['片段/模态','NRMSE','相关r','平均偏差','观测主极值@秒','预测主极值@秒'],wave_rows,[1.8,1.4,1.4,1.8,3.05,3.05],
            '平均偏差＝预测−观测，单位为开发SD。相关只描述波形同步程度，不能替代幅度误差；两次记录的误差也不能证明参数为生理真值。',
            f'evaluation/{ident}__P2__{win.variant}.json；fits/*.npz；prepared/*.npz')
    findings=summary.get('final_findings',[
        ('可得到什么','可报告冻结观测坐标下的最佳条件拟合参数、剖面容许范围和重复差异；这些结果不等同直接测得的个体生理常数。'),
        ('不能用什么替代','更低重建误差、强共享后的稳定性、或者固定参数的零差异，都不能替代个体参数的辨识与真实差异恢复。'),
        ('后续用途','保留通过检查的条件重建与状态诊断；个体参数作为研究性摘要时须同时携带近界、收敛、重复与先验依赖信息。'),
        ('后续优先级','依据共同内点剖面和原始拟合代价决定继续拟合还是增加独立观测约束；不以更漂亮的参数分布作为成功标准。')])
    if len(coordinate_groups):
        visual=primary(coordinate_groups[(coordinate_groups.partition=='evaluation')&(coordinate_groups.model=='P1_tau')
            &(coordinate_groups.variant=='N0')],'visual_cognitive_motivation')
        local=visual.sort_values(['effective_pair_rate','repeat_median'],ascending=[False,True]).iloc[0]
        site=local.site.replace('prefrontal','前额').replace('motor','运动').replace('__montage',' M')
        findings[2]=('Visual：局部τ是待验证候选',
            f'{site}单τ的重复差异中位{pct(local.repeat_median)}、有效配对{int(local.effective_pairs)}/{int(local.planned_pairs)}，近界{pct(local.near_05_pair_rate)}。仅{int(local.independent_subjects)}人；需独立复验，不能外推为普遍个体生理精度。')
    deck.methods('对当前数据最有用的结论与下一步',
        '最佳数值结果、可重复有效参数和个体生理真实性是三个不同层次',findings,
        'summary.json；profile_intersections.json；verification.json')
    ver=summary.get('mathematical_and_numerical_verification') or read(run/'verification.json')
    screen=summary['postfit_numerical_screen']
    deck.methods('执行、数值与证据边界',
        '所有正式计算在用户级 systemd 下执行；失败与软件恢复记录保留，报告不替代 owning evidence',[
        ('资源与恢复','52 个物理核；拟合48个单线程进程，原生数据准备16进程。任务逐项保存，可恢复；两个软件问题的原失败记录保留。'),
        ('独立核验',f"核验{ver['record_checks']}条结果；拟合器通过解全部在8子步合法。其中{len(screen['excluded_records'])}条因积分精度不足排除资格；余{screen['qualified_records']}条最大差{screen['qualified_max_substep_difference_development_sd']:.2g}开发SD。评分重算差为{ver['metric_max_abs_error']:.2g}。"),
        ('未执行与条件分支',summary.get('unexecuted_steps_text','高维自由度只在前一层信息通过后增加；未重新训练 tokenizer，未执行其他协议的受保护评价或生理资格确认。')),
        ('保留证据','配置、身份配对、开发变换、逐项结果、剖面、评分表、运行日志及代码快照均在同一 run 下。PPT/PDF 是沟通导出。')],
        'resources.json；software_recovery_*.json；source_snapshots/；verification.json')
    if screen['excluded_records']:
        rows=[[DS[r['dataset']],r['subject']+' '+r['site'],MODEL[r['model']],r['variant'],
            'A' if r['side']==0 else 'B',fmt(r['max_substep_difference_development_sd'],5)] for r in screen['excluded_records']]
        deck.table('两条对照记录因积分精度不足而不计为合格解',
            '4→8子步最大逐点变化超过0.001开发SD；原始拟合、原始失败核验及登记分母均保留',
            ['数据集','被试/区域','模型','对照臂','片段','重积分最大变化'],rows,[1.9,3.0,2.1,1.3,1.0,3.2],
            f"受影响的冻结评价记录为{screen['excluded_evaluation_records']}条。原始 verification_pass 仍为False；最终汇总完成了失败归类，没有放宽精度门槛。",
            'verification.json；verify/*.json；summary.json：postfit_numerical_screen')
    reference_slide=deck.methods('方法出处与解释限制',
        '文献提供分析方法；当前项目的数值和结论只以本轮保留证据为准',[
        ('完整剖面','Raue et al., Bioinformatics (2009)：固定目标参数后重新优化其余未知量。本轮目标含工程正则，故不直接套用似然置信区间。'),
        ('层级共享','Friston et al., Frontiers in Systems Neuroscience (2015)：群体共享可能改善估计，也会引入先验依赖；本轮另查恢复斜率和差异压缩。'),
        ('有限样本限制','PLOS Computational Biology, doi:10.1371/journal.pcbi.1011417：非线性ODE模型中的渐近阈值需校准。本轮剖面仅作工程容差诊断。'),
        ('数据出处','Single-Trial：Shin et al. EEG+NIRS；Simultaneous：Shin et al. Cognitive Tasks；Visual：Kyushu visual cognitive motivation 原生发布说明。')],
        'docs/DATASETS_DESCRIPTION.md；Raue 2009；Friston 2015；PLOS pcbi.1011417')
    reference_urls={'Raue et al.':'https://doi.org/10.1093/bioinformatics/btp358',
        'Friston et al.':'https://www.frontiersin.org/journals/systems-neuroscience/articles/10.3389/fnsys.2015.00164/full',
        'PLOS Computational Biology':'https://journals.plos.org/ploscompbiol/article?id=10.1371/journal.pcbi.1011417'}
    for shape in reference_slide.shapes:
        if not shape.has_text_frame:continue
        for prefix,url in reference_urls.items():
            if shape.text.startswith(prefix):
                for paragraph in shape.text_frame.paragraphs:
                    for run_text in paragraph.runs:run_text.hyperlink.address=url
    pptx=out/'PARAMETER_STABILITY.pptx';deck.prs.save(pptx)
    (out/'slide_sources.json').write_text(json.dumps(clean(dict(schema='parameter_stability_slide_sources_v1',
        evidence_root=str(run.relative_to(ROOT)),implementation_root=str(ROOT),slides=deck.sources,method_references=reference_urls)),ensure_ascii=False,indent=2)+'\n')
    (out/'layout_checks.json').write_text(json.dumps(clean(dict(text_boxes=deck.text_boxes,figures=sorted(set(deck.figures)))),ensure_ascii=False,indent=2)+'\n')
    return deck,pptx


def export_and_check(deck,pptx,out,run):
    profile='/tmp/lo_parameter_stability_'+out.name
    result=subprocess.run(['libreoffice','-env:UserInstallation=file://'+profile,'--headless','--convert-to','pdf','--outdir',str(out),str(pptx)],capture_output=True,text=True,check=True)
    (out/'pdf_export.log').write_text(result.stdout+result.stderr)
    pdf=pptx.with_suffix('.pdf');doc=fitz.open(pdf);pages=[];renders=out/'preview';renders.mkdir(exist_ok=True)
    for i,page in enumerate(doc):
        blocks=page.get_text('dict')['blocks'];texts=[b for b in blocks if b['type']==0]
        outside=[]
        for block in texts:
            x0,y0,x1,y1=block['bbox']
            if x0<-.1 or y0<-.1 or x1>page.rect.width+.1 or y1>page.rect.height+.1:outside.append(block['bbox'])
        images=page.get_images(full=True)
        pages.append(dict(page=i+1,text_characters=len(page.get_text()),image_objects=len(images),
            expected_bitmap_figures=deck.figure_pages.get(i+1,0),out_of_page_text=outside))
        page.get_pixmap(matrix=fitz.Matrix(1.25,1.25),alpha=False).save(renders/f'page_{i+1:02d}.png')
    # Contact sheets support a whole-deck clipping/legibility inspection.
    thumbnails=[]
    for i in range(len(doc)):
        im=Image.open(renders/f'page_{i+1:02d}.png').convert('RGB');im.thumbnail((533,300));thumbnails.append(im)
    for start in range(0,len(thumbnails),12):
        sheet=Image.new('RGB',(533*3,300*4),'#DDE3EA')
        for k,im in enumerate(thumbnails[start:start+12]):sheet.paste(im,((k%3)*533,(k//3)*300))
        sheet.save(renders/f'contact_{start//12+1:02d}.png')
    import zipfile
    with zipfile.ZipFile(pptx) as z:media=[p for p in z.namelist() if p.startswith('ppt/media/')]
    snapshot=run/'source_snapshots'/('report_'+out.name)/'experiments/scripts/render_shared_driver_parameter_stability.py'
    snapshot.parent.mkdir(parents=True,exist_ok=True)
    if snapshot.exists() and snapshot.read_bytes()!=Path(__file__).read_bytes():
        raise ValueError('retained renderer snapshot differs; use a new versioned export directory')
    snapshot.write_bytes(Path(__file__).read_bytes())
    verification=dict(status='completed',slides=len(deck.prs.slides),pdf_pages=len(doc),pages=pages,
        searchable_text_all_pages=all(p['text_characters']>30 for p in pages),
        bitmap_media_only=all(p.lower().endswith('.png') for p in media),unique_bitmap_figures=len(set(deck.figures)),
        every_figure_page_has_images=all(p['image_objects']>=p['expected_bitmap_figures'] for p in pages),
        figures_as_pdf_image_objects=sum(p['image_objects'] for p in pages),out_of_page_text_count=sum(len(p['out_of_page_text']) for p in pages),
        pptx=str(pptx),pdf=str(pdf),owning_evidence=str(run/'summary.json'),renderer_source_snapshot=str(snapshot))
    (out/'export_validation.json').write_text(json.dumps(verification,ensure_ascii=False,indent=2)+'\n')
    if len(doc)!=len(deck.prs.slides) or not verification['searchable_text_all_pages'] or not verification['bitmap_media_only'] or not verification['every_figure_page_has_images'] or verification['out_of_page_text_count']:
        raise ValueError('PPT/PDF export validation failed')
    return verification


def main():
    parser=argparse.ArgumentParser();parser.add_argument('--run-dir',required=True);parser.add_argument('--output-dir',required=True)
    args=parser.parse_args();run=Path(args.run_dir).resolve();out=Path(args.output_dir).resolve();out.mkdir(parents=True,exist_ok=True)
    if (out/'export_validation.json').exists():raise ValueError('retained report export exists; use a new versioned output directory')
    deck,pptx=build(run,out);result=export_and_check(deck,pptx,out,run)
    print(json.dumps({k:result[k] for k in ['slides','pdf_pages','unique_bitmap_figures','pptx','pdf']},ensure_ascii=False))


if __name__=='__main__':main()
