#!/usr/bin/env python3
"""Static spatial illustrations from one explicit JSON + NPZ export.

Rendering reads one explicit export without discovering runs or reading raw
data, fitting a model, or estimating uncertainty. ``--demo`` uses the committed
nonlinear forward model. ``--whole-head-demo`` separately generates synthetic
signals on public drawing metadata and exports correlation, oracle residual
OLS and conditional-history OLS relations. Both save the export consumed by
``--bundle``. Observations use [time, Hb channel, HbO/HbR] and [time, EEG channel];
Hb masks are paired. ``validate_bundle`` defines the two versioned schemas.
"""
from __future__ import annotations

import argparse
from dataclasses import asdict
import json
from pathlib import Path
import sys

import numpy as np

PROJECT_ROOT = Path(__file__).resolve().parents[2]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))
from src.inference.shared_driver_attribution import hb_diagnostic_coordinates


SCHEMA = "shared_driver_spatial_visualization_v1"
WHOLE_HEAD_SCHEMA = "shared_driver_spatial_visualization_v2"
FONT_PATH = Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc")
DPI = 240
BLUE = "#2775AD"
ORANGE = "#D67A26"
INK = "#18324D"
GRAY = "#929BA4"
BACKGROUND = "#FAFCFE"
REQUIRED_ARRAYS = {
    "time", "hb_observation", "eeg_observation", "visible_hb", "score_hb",
    "visible_eeg", "score_eeg",
}
OPTIONAL_ARRAYS = {
    "hb_truth", "prediction_hb", "prediction_valid", "drivers_k1", "drivers_k2",
    "hb_k1", "hb_k2", "hb_spatial_permuted", "equivalence_total_a",
    "equivalence_total_b", "equivalence_components_a", "equivalence_components_b",
}


def _array(value, shape, name, *, finite=True):
    value = np.asarray(value)
    if value.shape != shape or value.dtype.kind not in "fiu":
        raise ValueError(f"{name} must be numeric with shape {shape}")
    if np.isinf(value).any() or (finite and not np.isfinite(value).all()):
        raise ValueError(f"{name} contains nonfinite values")
    return value


def _text(value, name):
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"{name} must be a nonempty string")


def _keys(value, required, optional, name):
    if not isinstance(value, dict) or not required <= value.keys() or value.keys() - required - optional:
        raise ValueError(f"invalid keys in {name}; required={sorted(required)}, optional={sorted(optional)}")


def validate_bundle(manifest, arrays):
    """Validate the complete export, without reading paths or calculating fits.

    J=2..4 paired Hb channels, E=1..12 EEG channels, one or two local modes.
    Missing observation values may be NaN only outside visible/score masks.
    Forward examples and equivalence constructs are optional and synthetic only.
    Prediction arrays are optional and use an explicit paired validity mask;
    invalid entries must be NaN, so no line bridges failed/missing segments.
    """
    if isinstance(manifest, dict) and manifest.get("schema") == WHOLE_HEAD_SCHEMA:
        return _validate_whole_head_bundle(manifest, arrays)
    required = {"schema", "kind", "arrays_file", "title", "units", "names", "geometry",
                "W", "W_semantics", "coordinate_operator", "snapshot_times_s", "provenance"}
    _keys(manifest, required, {"demo_status", "spatial_permutation"}, "manifest")
    if manifest["schema"] != SCHEMA:
        raise ValueError("unsupported spatial visualization schema")
    if manifest["kind"] not in {"synthetic_demo", "measured_export"}:
        raise ValueError("kind must be synthetic_demo or measured_export")
    for key in ("title", "W_semantics"):
        _text(manifest[key], key)
    filename = manifest["arrays_file"]
    if not isinstance(filename, str) or Path(filename).name != filename or not filename.endswith(".npz"):
        raise ValueError("arrays_file must name one local NPZ file, without directories")
    _keys(manifest["units"], {"time", "geometry", "hb", "eeg", "driver"}, set(), "units")
    for key, unit in manifest["units"].items():
        _text(unit, f"units.{key}")
        if any(token in unit.lower() for token in ("zscore", "z-score", "standardized", "标准差单位")):
            raise ValueError("export physical/model coordinates, not independently standardized channels")
    if manifest["units"]["time"] != "s":
        raise ValueError("time unit must be s")
    if not isinstance(manifest["provenance"], dict) or not manifest["provenance"]:
        raise ValueError("provenance must document the explicit export source and processing")
    _keys(manifest["names"], {"hb", "eeg", "modes", "sources", "detectors"}, set(), "names")
    names = manifest["names"]
    for key, group in names.items():
        if not isinstance(group, list) or not group or len(set(group)) != len(group):
            raise ValueError(f"names.{key} must contain unique nonempty strings")
        for name in group:
            _text(name, f"names.{key}")
    j, e, s, d = (len(names[key]) for key in ("hb", "eeg", "sources", "detectors"))
    k = len(names["modes"])
    if not 2 <= j <= 4 or not 1 <= e <= 12 or k not in (1, 2):
        raise ValueError("expected J=2..4 Hb channels, E=1..12 EEG channels and K=1 or 2 modes")
    geometry = manifest["geometry"]
    _keys(geometry, {"hb_xy", "eeg_xy", "source_xy", "detector_xy", "hb_pairs", "neighbour_edges"}, set(), "geometry")
    for key, count in (("hb_xy", j), ("eeg_xy", e), ("source_xy", s), ("detector_xy", d)):
        _array(geometry[key], (count, 2), f"geometry.{key}")
    pairs = np.asarray(geometry["hb_pairs"])
    if pairs.shape != (j, 2) or pairs.dtype.kind not in "iu" or np.any(pairs < 0):
        raise ValueError("hb_pairs must be integer [J,2] source/detector indices")
    if np.any(pairs[:, 0] >= s) or np.any(pairs[:, 1] >= d) or len(set(map(tuple, pairs))) != j:
        raise ValueError("hb_pairs must be unique valid source/detector pairs")
    edges = geometry["neighbour_edges"]
    if not isinstance(edges, list):
        raise ValueError("neighbour_edges must be an explicit list")
    for edge in edges:
        edge = np.asarray(edge)
        if edge.shape != (2,) or edge.dtype.kind not in "iu" or np.any(edge < 0) or np.any(edge >= j) or edge[0] == edge[1]:
            raise ValueError("neighbour_edges must index two different Hb channels")
    weights = _array(manifest["W"], (j, k), "W")
    if np.any(weights < 0) or not np.allclose(weights.sum(axis=1), 1, atol=1e-12, rtol=0):
        raise ValueError("W must be nonnegative with each row summing to one")
    operator = manifest["coordinate_operator"]
    _keys(operator, {"total", "oxygenation"}, set(), "coordinate_operator")
    if operator != {"total": [1, 1], "oxygenation": [.35, -.65]}:
        raise ValueError("this renderer requires T=O+R and X=0.35O-0.65R in a common Hb unit")
    if not isinstance(arrays, dict) or not REQUIRED_ARRAYS <= arrays.keys() or arrays.keys() - REQUIRED_ARRAYS - OPTIONAL_ARRAYS:
        raise ValueError("unexpected or missing NPZ array keys")
    time = np.asarray(arrays["time"])
    if time.ndim != 1 or len(time) < 3 or time.dtype.kind not in "fiu" or not np.isfinite(time).all() or np.any(np.diff(time) <= 0):
        raise ValueError("time must be finite and strictly increasing [T], T>=3")
    t = len(time)
    for suffix, shape in (("hb", (t, j, 2)), ("eeg", (t, e))):
        observation = _array(arrays[f"{suffix}_observation"], shape, f"{suffix}_observation", finite=False)
        masks = []
        for prefix in ("visible", "score"):
            mask = np.asarray(arrays[f"{prefix}_{suffix}"])
            if mask.shape != shape or mask.dtype != np.bool_:
                raise ValueError(f"{prefix}_{suffix} must be boolean {shape}")
            if suffix == "hb" and not np.array_equal(mask[..., 0], mask[..., 1]):
                raise ValueError("HbO/HbR masks must hide/show complete pairs")
            masks.append(mask)
        if np.any(masks[0] & masks[1]):
            raise ValueError("visible and score masks must be disjoint")
        if not np.isfinite(observation[masks[0] | masks[1]]).all():
            raise ValueError("visible/scored observations must be finite")
    for name in ("hb_truth", "hb_k1", "hb_k2", "hb_spatial_permuted", "equivalence_total_a", "equivalence_total_b"):
        if name in arrays:
            _array(arrays[name], (t, j, 2), name)
    for name, count in (("drivers_k1", 1), ("drivers_k2", 2)):
        if name in arrays:
            _array(arrays[name], (t, count), name)
    if ("prediction_hb" in arrays) != ("prediction_valid" in arrays):
        raise ValueError("prediction_hb and prediction_valid must be exported together")
    if "prediction_hb" in arrays:
        pred = _array(arrays["prediction_hb"], (t, j, 2), "prediction_hb", finite=False)
        valid = np.asarray(arrays["prediction_valid"])
        if valid.shape != pred.shape or valid.dtype != np.bool_ or not np.array_equal(valid[..., 0], valid[..., 1]):
            raise ValueError("prediction_valid must be paired boolean [T,J,2]")
        if not np.isfinite(pred[valid]).all() or not np.isnan(pred[~valid]).all():
            raise ValueError("valid predictions must be finite; invalid predictions must be NaN")
    forward = {"drivers_k1", "drivers_k2", "hb_k1", "hb_k2"}
    if arrays.keys() & forward and not forward <= arrays.keys():
        raise ValueError("export all four K1/K2 forward arrays together")
    equivalent = {"equivalence_total_a", "equivalence_total_b", "equivalence_components_a", "equivalence_components_b"}
    if arrays.keys() & equivalent and not equivalent <= arrays.keys():
        raise ValueError("export all four equivalence arrays together")
    for name in ("equivalence_components_a", "equivalence_components_b"):
        if name in arrays:
            _array(arrays[name], (t, j, 2, 2), name)
    if equivalent <= arrays.keys():
        for suffix in ("a", "b"):
            if not np.allclose(arrays[f"equivalence_components_{suffix}"].sum(axis=2), arrays[f"equivalence_total_{suffix}"], atol=1e-12, rtol=1e-12):
                raise ValueError("equivalence components do not sum to their total")
    synthetic_arrays = forward | equivalent | {"hb_spatial_permuted"}
    if arrays.keys() & synthetic_arrays and k != 2:
        raise ValueError("the two-core synthetic examples require two named local modes")
    if manifest["kind"] != "synthetic_demo" and arrays.keys() & synthetic_arrays:
        raise ValueError("synthetic forward/equivalence examples cannot be labeled measured_export")
    if manifest["kind"] == "measured_export" and "hb_truth" in arrays:
        raise ValueError("measured_export cannot label a conditional estimate as generator truth")
    snapshots = _array(manifest["snapshot_times_s"], (3,), "snapshot_times_s")
    if np.any(np.diff(snapshots) <= 0) or snapshots[0] < time[0] or snapshots[-1] > time[-1]:
        raise ValueError("snapshot_times_s must be three increasing times inside support")
    if "spatial_permutation" in manifest:
        permutation = np.asarray(manifest["spatial_permutation"])
        if permutation.shape != (j,) or permutation.dtype.kind not in "iu" or set(permutation) != set(range(j)):
            raise ValueError("spatial_permutation must be a permutation of Hb indices")
    if "hb_spatial_permuted" in arrays and ("spatial_permutation" not in manifest or not forward <= arrays.keys()):
        raise ValueError("permuted forward example requires its permutation and the K1/K2 example")
    for status in manifest.get("demo_status", []):
        _keys(status, {"channel", "interval_s", "label"}, set(), "demo_status item")
        if manifest["kind"] != "synthetic_demo" or type(status["channel"]) is not int or not 0 <= status["channel"] < j:
            raise ValueError("demo_status applies only to a synthetic Hb channel")
        interval = _array(status["interval_s"], (2,), "demo_status.interval_s")
        if not time[0] <= interval[0] < interval[1] <= time[-1]:
            raise ValueError("demo status interval outside time support")
        _text(status["label"], "demo_status.label")


def load_bundle(manifest_path):
    """Read only the named manifest and its sibling NPZ (pickle disabled)."""
    path = Path(manifest_path).resolve(strict=True)
    manifest = json.loads(path.read_text(encoding="utf-8"))
    filename = manifest.get("arrays_file")
    if not isinstance(filename, str) or Path(filename).name != filename or not filename.endswith(".npz"):
        raise ValueError("arrays_file must be a sibling NPZ filename")
    array_path = path.parent / filename
    if array_path.is_symlink():
        raise ValueError("array export must be a regular sibling file, not a symlink")
    with np.load(array_path, allow_pickle=False) as exported:
        arrays = {name: exported[name] for name in exported.files}
    validate_bundle(manifest, arrays)
    return manifest, arrays


def _fresh(path):
    path = Path(path)
    path.mkdir(parents=True, exist_ok=False)
    return path


def create_demo_bundle(output, seed=20261016):
    """Save one reproducible toy forward generation; output must be fresh."""
    if isinstance(seed, bool) or not isinstance(seed, int) or seed < 0:
        raise ValueError("seed must be a nonnegative integer")
    from src.inference.shared_driver_reconstruction import nonlinear_driver_forward
    from src.inference.t3a_balloon_robust_ssm import BalloonParameters

    parameters = BalloonParameters()
    initial = np.array([0., 1., 1., 1., 1.])
    dt = .25
    time = np.arange(0., 60. + dt / 2, dt)
    r1 = .09 * np.exp(-.5 * ((time - 12.) / 2.8) ** 2) + .055 * np.exp(-.5 * ((time - 34.) / 3.5) ** 2)
    r2 = .025 * np.exp(-.5 * ((time - 14.) / 3.) ** 2) + .11 * np.exp(-.5 * ((time - 43.) / 2.5) ** 2)
    local = np.stack([
        nonlinear_driver_forward(r, initial, parameters, dt, substeps=4, derivative=False, hemodynamics='inlet_balance', viscoelastic_s=0., numerical_backend='python')["canonical_prediction"][:, 1:]
        for r in (r1, r2)
    ], axis=1)
    weights = np.array([[.85, .15], [.65, .35], [.35, .65], [.15, .85]])
    truth = np.einsum("jk,tkc->tjc", weights, local)
    one = np.broadcast_to(local[:, :1], truth.shape).copy()
    permutation = np.arange(4)[::-1]
    permuted = np.einsum("jk,tkc->tjc", weights[permutation], local)
    rng = np.random.default_rng(seed)
    noise_sd = np.array([.0015, .0007])
    observed = truth + rng.normal(size=truth.shape) * noise_sd
    eeg_weights = np.column_stack((np.linspace(.9, .1, 8), np.linspace(.1, .9, 8)))
    eeg = np.column_stack((r1, r2)) @ eeg_weights.T + rng.normal(0., .003, size=(len(time), 8))
    visible_hb = np.ones_like(observed, dtype=bool)
    visible_hb[:, 2] = False
    score_hb = np.zeros_like(visible_hb)
    score_hb[:, 2] = True
    # Deliberately missing input segment, not an actual solver failure.
    missing = (time >= 24) & (time <= 29)
    visible_hb[missing, 3] = False
    observed[missing, 3] = np.nan
    repeated = np.repeat(local[:, :1], 2, axis=1)
    weight_b = np.array([[.2, .8], [.4, .6], [.6, .4], [.8, .2]])
    component_a = weights[None, :, :, None] * repeated[:, None, :, :]
    component_b = weight_b[None, :, :, None] * repeated[:, None, :, :]
    arrays = dict(
        time=time, hb_observation=observed, eeg_observation=eeg,
        visible_hb=visible_hb, score_hb=score_hb,
        visible_eeg=np.ones_like(eeg, dtype=bool), score_eeg=np.zeros_like(eeg, dtype=bool),
        hb_truth=truth, drivers_k1=r1[:, None], drivers_k2=np.column_stack((r1, r2)),
        hb_k1=one, hb_k2=truth, hb_spatial_permuted=permuted,
        equivalence_total_a=component_a.sum(axis=2), equivalence_total_b=component_b.sum(axis=2),
        equivalence_components_a=component_a, equivalence_components_b=component_b,
    )
    manifest = dict(
        schema=SCHEMA, kind="synthetic_demo", arrays_file="arrays.npz",
        title="粗空间结构：静态科研可视化合成预期例",
        units=dict(time="s", geometry="假想二维模板单位", hb="相对 Hb 模型单位", eeg="EEG 模型单位", driver="驱动模型单位"),
        names=dict(hb=[f"H{i}" for i in range(1, 5)], eeg=[f"E{i}" for i in range(1, 9)], modes=["局部模式 1", "局部模式 2"], sources=["S1", "S2"], detectors=["D1", "D2", "D3"]),
        geometry=dict(
            hb_xy=[[-1.5, .9], [-.5, 1.], [.5, 1.], [1.5, .9]],
            eeg_xy=[[-1.7, .1], [-1.2, 1.7], [-.5, 2.1], [0., .2], [.5, 2.1], [1.2, 1.7], [1.7, .1], [0., -1.]],
            source_xy=[[-1., 1.4], [1., 1.4]], detector_xy=[[-2., .4], [0., .6], [2., .4]],
            hb_pairs=[[0, 0], [0, 1], [1, 1], [1, 2]], neighbour_edges=[[0, 1], [1, 2], [2, 3]],
        ),
        W=weights.tolist(), W_semantics="已知生成权重；不是估计的脑贡献",
        coordinate_operator=dict(total=[1, 1], oxygenation=[.35, -.65]), snapshot_times_s=[18., 36., 48.],
        spatial_permutation=permutation.tolist(),
        demo_status=[dict(channel=3, interval_s=[24., 29.], label="示意缺失状态；不是实际求解失败")],
        provenance=dict(
            source="deterministic toy forward generation; no measured data or optimizer",
            seed=seed, generator="src.inference.shared_driver_reconstruction.nonlinear_driver_forward",
            parameters=asdict(parameters), initial_physical_s_f_v_p_q=initial.tolist(),
            clock=dict(start_s=0., stop_s=60., dt_s=dt, samples=len(time), substeps=4, driver_hold="left zero-order hold"),
            drivers=dict(r1="0.09 exp(-0.5((t-12)/2.8)^2)+0.055 exp(-0.5((t-34)/3.5)^2)", r2="0.025 exp(-0.5((t-14)/3)^2)+0.11 exp(-0.5((t-43)/2.5)^2)", unit="驱动模型单位"),
            observation="independent nonlinear cores first; mix Hb outputs by W; no per-channel z-score, baseline subtraction or temporal smoothing",
            eeg_weights=eeg_weights.tolist(), eeg_noise_sd=.003, hb_noise_sd=noise_sd.tolist(), noise="independent Gaussian; seed fixed",
            k1="one r1 core copied to every channel; K2 has two independent forward solves using identical parameters/initial states",
            oracle_permutation="hold both true drivers, states and physics fixed; permute only Hb rows of W; no refit or null-performance test",
            equivalence=dict(cores="two identical copies of the r1 nonlinear trajectory and initial state", W_A=weights.tolist(), W_B=weight_b.tolist(), constraint="both nonnegative and row-stochastic; sum components in Hb observation coordinates"),
            masks="H3 pair held for scoring; H4 24..29s neither visible nor scored; masks are illustrative and no fit is performed",
        ),
    )
    validate_bundle(manifest, arrays)
    output = _fresh(output)
    np.savez_compressed(output / "arrays.npz", **arrays)
    path = output / "manifest.json"
    path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return path


def _plotting():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib import font_manager
    if not FONT_PATH.is_file():
        raise FileNotFoundError(f"required Chinese font not found: {FONT_PATH}")
    font_manager.fontManager.addfont(str(FONT_PATH))
    family = font_manager.FontProperties(fname=str(FONT_PATH)).get_name()
    plt.rcParams.update({
        "font.family": family, "font.size": 11, "axes.titlesize": 12, "axes.labelsize": 11,
        "xtick.labelsize": 10, "ytick.labelsize": 10, "legend.fontsize": 10,
        "figure.facecolor": BACKGROUND, "axes.facecolor": "white", "text.color": INK,
        "axes.labelcolor": INK, "axes.edgecolor": "#BBC8D4", "xtick.color": INK,
        "ytick.color": INK, "axes.spines.top": False, "axes.spines.right": False,
        "axes.unicode_minus": False, "savefig.facecolor": BACKGROUND,
    })
    return plt


def _badge(manifest):
    if manifest["schema"] == WHOLE_HEAD_SCHEMA:
        layout = "REFED公开布局" if manifest["provenance"].get("geometry_source") == "src/data/assets/refed_paper_layout_v1.json" else "显式二维布局"
        return f"{layout} × 合成信号｜非实测结果" if manifest["kind"] == "synthetic_demo" else f"{layout} × 显式实测导出｜未重新拟合"
    return "合成演示 · 非实测结果" if manifest["kind"] == "synthetic_demo" else "显式实测导出 · 未重新拟合"


def _canvas(plt, manifest, title, footer, *, rows=1, cols=1, figsize=(13, 7.2), **kwargs):
    fig, axes = plt.subplots(rows, cols, figsize=figsize, squeeze=False, **kwargs)
    fig.suptitle(title, x=.055, y=.975, ha="left", fontsize=18, fontweight="bold", color=INK)
    fig.text(.055, .92, _badge(manifest), fontsize=12, color=ORANGE, fontweight="bold")
    fig.text(.055, .035, footer, fontsize=11, va="bottom", linespacing=1.5)
    fig.subplots_adjust(left=.075, right=.96, top=.84, bottom=.16, wspace=.32, hspace=.5)
    return fig, axes


def _clean(ax):
    ax.grid(axis="y", color="#E8EDF2", lw=.6)
    ax.set_axisbelow(True)


def _geometry(ax, m, *, optodes=False, color_values=None, norm=None, cmap=None):
    g, n = m["geometry"], m["names"]
    hb = np.asarray(g["hb_xy"])
    if color_values is None and not optodes:
        eeg = np.asarray(g["eeg_xy"])
        ax.scatter(eeg[:, 0], eeg[:, 1], facecolors="white", edgecolors=GRAY, marker="^", s=60, zorder=2, label="EEG 位置")
        for label, xy in zip(n["eeg"], eeg):
            ax.annotate(label, xy, xytext=(3, 5), textcoords="offset points", fontsize=10, color="#647382")
    if optodes:
        source, detector = np.asarray(g["source_xy"]), np.asarray(g["detector_xy"])
        for index, (s, d) in enumerate(g["hb_pairs"]):
            ax.plot([source[s, 0], detector[d, 0]], [source[s, 1], detector[d, 1]], "-", color="#A7B6C6", lw=2)
        ax.scatter(*source.T, s=110, c=ORANGE, marker="s", label="光源 S", zorder=4)
        ax.scatter(*detector.T, s=110, c=BLUE, marker="D", label="探测器 D", zorder=4)
        for labels, coords, offset in ((n["sources"], source, (5, 5)), (n["detectors"], detector, (0, -20))):
            for label, xy in zip(labels, coords):
                ax.annotate(label, xy, xytext=offset, textcoords="offset points", fontsize=10, ha="center" if offset[0] == 0 else "left")
    elif color_values is None:
        for a, b in g["neighbour_edges"]:
            ax.plot(hb[[a, b], 0], hb[[a, b], 1], ":", color=BLUE, lw=2, zorder=1)
    if color_values is None:
        sc = ax.scatter(*hb.T, c=INK, s=70, marker="o", edgecolors="white", linewidths=1.2, zorder=5)
    else:
        color_values = np.asarray(color_values)
        valid = np.isfinite(color_values)
        sc = ax.scatter(*hb[valid].T, c=color_values[valid], cmap=cmap, norm=norm, s=150,
                        marker="o", edgecolors="white", linewidths=1.2, zorder=5)
        if (~valid).any():
            ax.scatter(*hb[~valid].T, c=GRAY, marker="x", s=150, linewidths=2, zorder=6)
    for label, xy in zip(n["hb"], hb):
        ax.annotate(label, xy, xytext=(0, -18), textcoords="offset points", ha="center", fontsize=11, color=INK)
    all_xy = np.concatenate([np.asarray(g[k]) for k in ("hb_xy", "eeg_xy", "source_xy", "detector_xy")])
    low, high = all_xy.min(axis=0) - .45, all_xy.max(axis=0) + .45
    ax.set(xlim=(low[0], high[0]), ylim=(low[1], high[1]), aspect="equal")
    ax.set_xticks([])
    ax.set_yticks([])
    for spine in ax.spines.values():
        spine.set_visible(False)
    return sc


def _layout(plt, m, a):
    fig, axes = _canvas(plt, m, "01  观测位置、邻域与共享光极有不同含义",
        "假想二维模板单位；不是 mm 或皮层定位。点与线仅定义观测结构；两类边均不表示神经连接。\nW 是固定观测混合；本例是已知生成权重，不是估计的脑贡献。" if m["kind"] == "synthetic_demo" else
        f"坐标单位：{m['units']['geometry']}。邻域线和共享光极线表示观测结构，不表示神经连接。\nW 含义：{m['W_semantics']}。", cols=3)
    _geometry(axes[0, 0], m)
    axes[0, 0].set_title("A  Hb 邻域（虚线）与 EEG 位置（三角）")
    _geometry(axes[0, 1], m, optodes=True)
    axes[0, 1].set_title("B  共享光源／探测器（实线）")
    axes[0, 1].legend(loc="lower center", bbox_to_anchor=(.5, -.12), ncol=2, frameon=False)
    ax = axes[0, 2]
    im = ax.imshow(m["W"], cmap="Blues", vmin=0, vmax=1, aspect="auto")
    labels = []
    for name, (s, d) in zip(m["names"]["hb"], m["geometry"]["hb_pairs"]):
        labels.append(f"{name} = {m['names']['sources'][s]}–{m['names']['detectors'][d]}")
    ax.set_xticks(range(len(m["names"]["modes"])), m["names"]["modes"])
    ax.set_yticks(range(len(labels)), labels)
    for row, values in enumerate(m["W"]):
        for col, value in enumerate(values):
            ax.text(col, row, f"{value:.2f}", ha="center", va="center", color="white" if value > .6 else INK, fontsize=13)
    ax.set_title("C  Hb × 局部模式权重 W；每行和为 1")
    fig.colorbar(im, ax=ax, fraction=.05, pad=.05, label="无量纲权重")
    return fig, dict(source_arrays=[], observation_units=m["units"]["geometry"], weights=m["W"], edge_semantics=["neighbourhood", "shared_optodes"])


def _waveforms(plt, m, a):
    t, obs, names = a["time"], a["hb_observation"], m["names"]["hb"]
    truth = a.get("hb_truth")
    j = len(names)
    footer = f"每行一个 Hb 观测对；每列一种色团；所有通道共用该色团的纵轴范围。单位：{m['units']['hb']}。\n"
    footer += "蓝实线＝无噪声生成真值；黑点线＝带噪合成观测；灰区＝观测示意缺失，生成真值仍已知。未逐通道标准化。" if truth is not None and m["kind"] == "synthetic_demo" else "黑点线＝导出观测；空白保留缺失；可见／评分角色见图 05。未逐通道标准化。"
    fig, axes = _canvas(plt, m, "02  同一单位下保留各位置的幅度与时序差异", footer, rows=j, cols=2, sharex=True)
    fig.subplots_adjust(hspace=.22, wspace=.2)
    limits = []
    for c in range(2):
        values = obs[..., c].ravel()
        if truth is not None:
            values = np.r_[values, truth[..., c].ravel()]
        finite = values[np.isfinite(values)]
        lo, hi = (float(finite.min()), float(finite.max())) if len(finite) else (-1., 1.)
        pad = max((hi-lo)*.12, 1e-6)
        limits.append([lo-pad, hi+pad])
    for row, name in enumerate(names):
        for c, chrom in enumerate(("HbO", "HbR")):
            ax = axes[row, c]
            ax.plot(t, obs[:, row, c], color="#111111", lw=.7, marker=".", markersize=1.8, markevery=3, alpha=.65)
            if truth is not None:
                ax.plot(t, truth[:, row, c], color=BLUE, lw=1.8)
            for status in m.get("demo_status", []):
                if status["channel"] == row:
                    ax.axvspan(*status["interval_s"], color="#D9DDE1", alpha=.8)
            ax.set_ylim(limits[c])
            ax.set_ylabel(name)
            ax.axhline(0, color="#B8C5D0", lw=.6)
            _clean(ax)
            if row == 0:
                ax.set_title(f"{chrom}（{m['units']['hb']}）")
            if row == j-1:
                ax.set_xlabel("时间 / s")
    return fig, dict(source_arrays=["hb_observation"] + (["hb_truth"] if truth is not None else []), y_limits_by_chromophore=limits, channelwise_standardization=False)


def _dynamics(plt, m, a):
    t = a["time"]
    has_perm = "hb_spatial_permuted" in a
    fig, axes = _canvas(plt, m, "03  两个独立局部动态可改变空间波形差异",
        "先分别求解非线性血流动力学单元，再按 W 混合 Hb 输出；物理参数和初态相同。生成条件对照，不是拟合优劣。\n底行是已知驱动前向对照：固定真值驱动与全部动力学，只置换 Hb 空间对应；不是拟合或置换检验性能。" if has_perm else
        "先分别求解非线性血流动力学单元，再按 W 混合 Hb 输出；物理参数和初态相同。生成条件对照，不是拟合优劣。",
        rows=3 if has_perm else 2, cols=2, sharex=True)
    fig.subplots_adjust(hspace=.55, wspace=.25)
    driver_values = np.concatenate([a[key].ravel() for key in ("drivers_k1", "drivers_k2")])
    driver_lo, driver_hi = float(driver_values.min()), float(driver_values.max())
    driver_pad = max((driver_hi-driver_lo)*.1, 1e-8)
    for col, k in enumerate((1, 2)):
        ax = axes[0, col]
        for index in range(k):
            ax.plot(t, a[f"drivers_k{k}"][:, index], color=(BLUE, ORANGE)[index], ls=("-", "--")[index], lw=2, label=f"r{index+1}")
        ax.set_title(f"K_gen={k}：已知生成驱动")
        ax.set_ylabel(m["units"]["driver"])
        ax.set_ylim(driver_lo-driver_pad, driver_hi+driver_pad)
        ax.legend(frameon=False, loc="upper right")
        _clean(ax)
    rows = [0, len(m["names"]["hb"])-1]
    comparison = np.concatenate([a[key][..., 0].ravel() for key in ("hb_k1", "hb_k2", "hb_spatial_permuted") if key in a])
    comparison_lo, comparison_hi = float(comparison.min()), float(comparison.max())
    comparison_pad = max((comparison_hi-comparison_lo)*.12, 1e-8)
    limits = [comparison_lo-comparison_pad, comparison_hi+comparison_pad]
    for col, channel in enumerate(rows):
        ax = axes[1, col]
        for k, color, style in ((1, BLUE, "-"), (2, ORANGE, "--")):
            ax.plot(t, a[f"hb_k{k}"][:, channel, 0], color=color, ls=style, lw=1.8, label=f"K_gen={k}")
        ax.set(title=f"{m['names']['hb'][channel]} 的 HbO：单／双动态", ylabel=m["units"]["hb"], ylim=limits)
        ax.legend(frameon=False, loc="upper right", bbox_to_anchor=(1, 1.15), ncol=2)
        _clean(ax)
        if has_perm:
            ax = axes[2, col]
            ax.plot(t, a["hb_k2"][:, channel, 0], color=BLUE, lw=1.8, label="原空间对应")
            ax.plot(t, a["hb_spatial_permuted"][:, channel, 0], color=ORANGE, ls="--", lw=1.8, label="置换空间对应")
            ax.set(title=f"{m['names']['hb'][channel]} 的 HbO：已知驱动前向对照", ylabel=m["units"]["hb"], ylim=limits)
            ax.legend(frameon=False, loc="upper right", bbox_to_anchor=(1, 1.15), ncol=2)
            _clean(ax)
        axes[-1, col].set_xlabel("时间 / s")
    return fig, dict(source_arrays=["drivers_k1", "drivers_k2", "hb_k1", "hb_k2"] + (["hb_spatial_permuted"] if has_perm else []), representative_channels=rows, comparison="generating conditions; no estimation", oracle_spatial_permutation=m.get("spatial_permutation"), hb_shared_limits=limits, driver_shared_limits=[driver_lo-driver_pad, driver_hi+driver_pad])


def _coordinates(plt, m, a):
    from matplotlib.colors import Normalize
    source = "hb_truth" if "hb_truth" in a else "hb_observation"
    y, time = a[source], a["time"]
    coordinates = hb_diagnostic_coordinates(y)
    total, exchange = coordinates[..., 0], coordinates[..., 1]
    indices = [int(np.argmin(np.abs(time-point))) for point in m["snapshot_times_s"]]
    source_label = "无噪声生成真值（不是观测或预测）" if source == "hb_truth" else "导出观测（不是预测）"
    geometry_note = "位置为假想二维模板单位，不是 mm 或皮层定位。" if m["kind"] == "synthetic_demo" else f"位置单位：{m['units']['geometry']}。"
    fig, axes = _canvas(plt, m, "04  总 Hb 与氧合方向是同单位的观测坐标",
        f"T = HbO + HbR；X = 0.35 HbO − 0.65 HbR。单位均为{m['units']['hb']}；它们不是血流或血氧饱和度。\n绘制{source_label}。每个圆点是一个 Hb 位置；不作空间插值。\n每行色条固定；灰色／灰叉为缺失，白色为近零。{geometry_note}", rows=2, cols=4)
    fig.subplots_adjust(left=.06, right=.94, wspace=.33, hspace=.45)
    maxima = []
    cmap = plt.get_cmap("RdBu_r").copy()
    cmap.set_bad("#D9DDE1")
    for row, (values, label) in enumerate(((total, "T：总 Hb"), (exchange, "X：氧合方向"))):
        finite = values[np.isfinite(values)]
        vmax = max(float(np.max(np.abs(finite))) if len(finite) else 0., 1e-8)
        maxima.append(vmax)
        norm = Normalize(-vmax, vmax)
        for col, index in enumerate(indices):
            im = _geometry(axes[row, col], m, color_values=values[index], norm=norm, cmap=cmap)
            axes[row, col].set_title(f"{label}\nt = {time[index]:g} s")
        ax = axes[row, 3]
        im = ax.pcolormesh(time, np.arange(len(m["names"]["hb"])), np.ma.masked_invalid(values.T), shading="nearest", cmap=cmap, norm=norm, rasterized=True)
        ax.set_yticks(range(len(m["names"]["hb"])), m["names"]["hb"])
        ax.set_xlabel("时间 / s")
        ax.set_title(f"{label}：完整时间支持")
        for index in indices:
            ax.axvline(time[index], color=INK, lw=.9, ls=":")
        fig.colorbar(im, ax=list(axes[row]), fraction=.018, pad=.025, label=m["units"]["hb"])
    return fig, dict(source_arrays=[source], coordinate_operator=m["coordinate_operator"], actual_snapshot_times_s=time[indices].tolist(), shared_symmetric_color_limits=[[-v, v] for v in maxima], spatial_interpolation=False)


def _visibility(plt, m, a):
    from matplotlib.colors import BoundaryNorm, ListedColormap
    from matplotlib.patches import Patch
    time, names = a["time"], m["names"]
    visible = np.column_stack((a["visible_eeg"], a["visible_hb"][..., 0]))
    score = np.column_stack((a["score_eeg"], a["score_hb"][..., 0]))
    code = np.where(visible, 1, np.where(score, 2, 0))
    fig, axes = _canvas(plt, m, "05  输入可见性、隐藏评分目标与预测状态分开呈现",
        "每列是一个时间点；每行是 EEG 通道或完整 HbO/HbR 观测对。可见与评分互斥；灰色表示两者均不可用。\n此演示掩码未用于拟合。隐藏目标只用于评分；生成真值只是参考。预测与数值失败应由未来实验明确导出。" if m["kind"] == "synthetic_demo" else
        "每列是一个时间点；每行是 EEG 通道或完整 HbO/HbR 观测对。可见与评分互斥；灰色表示两者均不可用。\n灰虚线是隐藏评分目标；它不是预测。预测无效区间保留空白，不跨缺失或失败片段连线。", rows=2, cols=2)
    grid = axes[0, 0].get_gridspec()
    for ax in axes[0]:
        ax.remove()
    mask_ax = fig.add_subplot(grid[0, :])
    cmap = ListedColormap(["#D9DDE1", BLUE, "#F0C18E"])
    mask_ax.pcolormesh(time, np.arange(code.shape[1]), code.T, shading="nearest", cmap=cmap, norm=BoundaryNorm([-.5, .5, 1.5, 2.5], 3))
    labels = names["eeg"] + [name+" (O/R)" for name in names["hb"]]
    mask_ax.set_yticks(range(len(labels)), labels)
    mask_ax.invert_yaxis()
    mask_ax.set_xlabel("时间 / s")
    mask_ax.legend(handles=[Patch(facecolor=BLUE, label="模型可见输入"), Patch(facecolor="#F0C18E", label="隐藏评分目标"), Patch(facecolor="#D9DDE1", label="缺失／不可用")], ncol=3, frameon=False, loc="lower center", bbox_to_anchor=(.5, 1.01))
    hidden = np.flatnonzero(a["score_hb"][..., 0].any(axis=0))
    channel = int(hidden[0]) if len(hidden) else 0
    for c, ax in enumerate(axes[1]):
        target = np.where(a["score_hb"][:, channel, c], a["hb_observation"][:, channel, c], np.nan)
        ax.plot(time, target, color=GRAY, ls="--", lw=1.5, label="隐藏目标（仅评分）")
        if "hb_truth" in a:
            ax.plot(time, a["hb_truth"][:, channel, c], color=BLUE, lw=1.4, label="生成真值参考" if m["kind"] == "synthetic_demo" else "导出真值参考")
        if "prediction_hb" in a:
            prediction = np.where(a["prediction_valid"][:, channel, c], a["prediction_hb"][:, channel, c], np.nan)
            ax.plot(time, prediction, color=ORANGE, lw=1.8, marker="o", markevery=15, markersize=2.5, label="导出预测")
        for status in m.get("demo_status", []):
            mask_ax.axvspan(*status["interval_s"], ymin=0, ymax=.04, color=GRAY, alpha=.7)
        title = f"{names['hb'][channel]} 的 {('HbO', 'HbR')[c]}"
        if "prediction_hb" not in a:
            title += "  |  预测待实验产生"
        ax.set(title=title, xlabel="时间 / s", ylabel=m["units"]["hb"])
        if "prediction_hb" not in a:
            ax.title.set_color(ORANGE)
        ax.legend(frameon=False, loc="upper left" if c == 0 else "lower left")
        _clean(ax)
    if m.get("demo_status"):
        labels = [f"{names['hb'][status['channel']]}：{status['label']}" for status in m["demo_status"]]
        mask_ax.text(.995, -.26, "；".join(labels), transform=mask_ax.transAxes, ha="right", fontsize=10, color="#667481")
    return fig, dict(source_arrays=["visible_hb", "score_hb", "visible_eeg", "score_eeg", "hb_observation"] + (["prediction_hb", "prediction_valid"] if "prediction_hb" in a else []), plotted_hidden_channel=channel, prediction_status="exported" if "prediction_hb" in a else "pending", visible_feature_cells=int(visible.sum()), score_feature_cells=int(score.sum()), visible_scalar_count=int(a["visible_eeg"].sum()+a["visible_hb"].sum()), score_scalar_count=int(a["score_eeg"].sum()+a["score_hb"].sum()), demo_status=m.get("demo_status", []))


def _equivalence(plt, m, a):
    time = a["time"]
    total_a, total_b = a["equivalence_total_a"], a["equivalence_total_b"]
    components_a, components_b = a["equivalence_components_a"], a["equivalence_components_b"]
    maximum = float(np.max(np.abs(total_a-total_b)))
    component_difference = float(np.max(np.abs(components_a-components_b)))
    fig, axes = _canvas(plt, m, "06  重复血流动力学单元可产生相同观测、不同分配",
        f"构造的观测等价例；不是实测参数剖面或置信区间。两单元的驱动、参数、初态完全相同；W_A／W_B 非负且行和为 1。\n全时间 × 通道 × 色团 max |总量 A − B| = {maximum:.3e}；max |分量 A − B| = {component_difference:.3e}（{m['units']['hb']}）。", rows=2, cols=2, sharex=True)
    channel = 0
    for c, chrom in enumerate(("HbO", "HbR")):
        ax = axes[0, c]
        ax.plot(time, total_a[:, channel, c], color="#111111", lw=2.2, label="总观测 A")
        ax.plot(time, total_b[:, channel, c], color=ORANGE, ls="--", lw=1.7, label="总观测 B")
        ax.set(title=f"{m['names']['hb'][channel]} 的 {chrom}：总观测重合", ylabel=m["units"]["hb"])
        ax.legend(frameon=False)
        _clean(ax)
        ax = axes[1, c]
        for k, color in enumerate((BLUE, ORANGE)):
            ax.plot(time, components_a[:, channel, k, c], color=color, ls="-", lw=1.9, label=f"模式 {k+1}，分配 A")
            ax.plot(time, components_b[:, channel, k, c], color=color, ls="--", lw=1.5, marker=("o", "s")[k], markersize=2.7, markevery=18, label=f"模式 {k+1}，分配 B")
        ax.set(title=f"{chrom}：加权分量改变，总和保持", xlabel="时间 / s", ylabel=m["units"]["hb"])
        ax.legend(frameon=False, ncol=2)
        _clean(ax)
    return fig, dict(source_arrays=["equivalence_total_a", "equivalence_total_b", "equivalence_components_a", "equivalence_components_b"], max_abs_total_difference=maximum, max_abs_component_difference=component_difference, plotted_channel=channel, interpretation="constructed observational equivalence of duplicate nonlinear cores; not a profile or confidence interval")


def correlation_matrix(signals, min_count=20):
    """Pairwise-complete Pearson r and actual pair counts, without imputation."""
    values = np.asarray(signals, dtype=float)
    if values.ndim != 2 or values.shape[1] < 1 or min_count < 3:
        raise ValueError("signals must be [time,node]; min_count must be >=3")
    finite = np.isfinite(values)
    observed = finite.astype(float)
    filled = np.where(finite, values, 0.)
    count = observed.T @ observed
    total = filled.T @ observed
    squared = (filled * filled).T @ observed
    cross = filled.T @ filled
    with np.errstate(invalid="ignore", divide="ignore"):
        centered = cross - total * total.T / count
        variance = squared - total * total / count
        denominator = np.sqrt(np.maximum(variance, 0.) * np.maximum(variance.T, 0.))
        result = centered / denominator
    valid = (count >= min_count) & (variance > 0.) & (variance.T > 0.)
    result[~valid] = np.nan
    result = np.clip(result, -1., 1.)
    return result, count.astype(np.int64)


def residual_correlation(signals, controls, min_count=20):
    """OLS residual Pearson r after a common listwise-complete oracle design.

    The intercept is included; all nodes use precisely the same row support.
    This is conditional linear association and is not a causal estimator.
    """
    values = np.asarray(signals, dtype=float)
    controls = np.asarray(controls, dtype=float)
    if values.ndim != 2 or controls.ndim != 2 or len(values) != len(controls):
        raise ValueError("signals and controls must be two-dimensional with equal time")
    keep = np.isfinite(values).all(axis=1) & np.isfinite(controls).all(axis=1)
    design = np.column_stack((np.ones(keep.sum()), controls[keep]))
    if keep.sum() <= design.shape[1] + 1:
        raise ValueError("insufficient common complete rows for residual correlation")
    residual = np.full_like(values, np.nan)
    residual[keep] = values[keep] - design @ np.linalg.lstsq(design, values[keep], rcond=None)[0]
    return correlation_matrix(residual, min_count=min_count)


def lag_prediction_gain(signals, covariates, p=2):
    """Same-sample nested OLS log(RSS_restricted/RSS_full), row source.

    The full design has intercept and lag 1..p of every node and known
    covariate. A restricted design removes only the nominated source's lags;
    target history and all other node/covariate histories stay. Self relations
    are undefined (NaN). Returns gain, full RSS, restricted RSS, row count.
    """
    values = np.asarray(signals, dtype=float)
    covariates = np.asarray(covariates, dtype=float)
    if (values.ndim != 2 or covariates.ndim != 2 or len(values) != len(covariates)
            or isinstance(p, bool) or not isinstance(p, int) or p < 1 or len(values) <= p):
        raise ValueError("invalid node/covariate arrays or lag order")
    t, nodes = values.shape
    history = np.column_stack([values[p-lag:t-lag] for lag in range(1, p+1)])
    oracle = np.column_stack([covariates[p-lag:t-lag] for lag in range(1, p+1)])
    design = np.column_stack((np.ones(t-p), history, oracle))
    target = values[p:]
    keep = np.isfinite(design).all(axis=1) & np.isfinite(target).all(axis=1)
    design, target = design[keep], target[keep]
    if len(design) <= design.shape[1] + 1:
        raise ValueError("insufficient shared rows for the full lag regression")
    error = target - design @ np.linalg.lstsq(design, target, rcond=None)[0]
    rss_full = np.sum(error * error, axis=0)
    if np.any(rss_full <= np.finfo(float).tiny):
        raise ValueError("degenerate zero full-model RSS")
    gain = np.full((nodes, nodes), np.nan)
    rss_restricted = np.full_like(gain, np.nan)
    for source in range(nodes):
        retain = np.ones(design.shape[1], dtype=bool)
        retain[1 + source + nodes * np.arange(p)] = False
        restricted = design[:, retain]
        residual = target - restricted @ np.linalg.lstsq(restricted, target, rcond=None)[0]
        rss = np.sum(residual * residual, axis=0)
        difference = rss - rss_full
        tolerance = 256 * np.finfo(float).eps * np.maximum(rss, rss_full)
        if np.any(difference < -tolerance):
            raise ArithmeticError("nested OLS RSS ordering failed beyond rounding tolerance")
        # Only machine-rounding violations are set to the exact nested bound.
        rss = np.where(difference < 0., rss_full, rss)
        gain[source] = np.log(rss / rss_full)
        rss_restricted[source] = rss
    np.fill_diagonal(gain, np.nan)
    np.fill_diagonal(rss_restricted, np.nan)
    return gain, rss_full, rss_restricted, int(keep.sum())


def intervention_response(transition, n_steps=20, amplitude=1.):
    """Exact same-noise impulse response [source,next step,target].

    Transition is [target,source]. At time zero a source state is set to its
    factual value plus amplitude. The returned support is k=1..n_steps, so the
    instantaneous intervention pulse is excluded. For this linear generator,
    A**k times the pulse exactly equals paired counterfactual replay.
    """
    matrix = np.asarray(transition, dtype=float)
    if (matrix.ndim != 2 or matrix.shape[0] != matrix.shape[1] or not np.isfinite(matrix).all()
            or isinstance(n_steps, bool) or not isinstance(n_steps, int) or n_steps < 1
            or not np.isfinite(amplitude)):
        raise ValueError("invalid transition, response length or amplitude")
    responses = []
    power = np.eye(len(matrix))
    for _ in range(n_steps):
        power = matrix @ power
        responses.append(amplitude * power.T)
    return np.stack(responses, axis=1)


def build_whole_head_demo(seed=20261016):
    """Build an independent synthetic relation example, retaining all layout nodes."""
    if isinstance(seed, bool) or not isinstance(seed, int) or seed < 0:
        raise ValueError("seed must be a nonnegative integer")
    layout_path = PROJECT_ROOT / "src/data/assets/refed_paper_layout_v1.json"
    layout = json.loads(layout_path.read_text(encoding="utf-8"))
    if (layout.get("schema") != "refed_paper_layout_v1" or len(layout["eeg"]) != 64
            or len(layout["channels"]) != 51 or len(layout["sources"]) != 18
            or len(layout["detectors"]) != 18):
        raise ValueError("expected complete versioned 64-EEG/51-Hb/18-S/18-D public layout")
    normalization = layout["normalization"]
    center = np.asarray(normalization["center_page_pt"], dtype=float)
    radius = float(normalization["radius_page_pt"])
    if (center.shape != (2,) or radius <= 0 or not np.isfinite(center).all()
            or not np.isfinite(radius) or normalization["x_positive"] != "figure_right"
            or normalization["y_positive"] != "figure_anterior"):
        raise ValueError("invalid public drawing coordinate normalization")

    def positions(entries):
        return (np.asarray([row["center_page_pt"] for row in entries]) - center) / radius * [1., -1.]

    source_xy, detector_xy, eeg_xy = (positions(layout[key]) for key in ("sources", "detectors", "eeg"))
    pairs = np.asarray([[row["source"]-1, row["detector"]-1] for row in layout["channels"]], dtype=int)
    hb_xy = (source_xy[pairs[:, 0]] + detector_xy[pairs[:, 1]]) / 2
    names = dict(hb=[f"H{row['channel']:02d}" for row in layout["channels"]],
                 eeg=[row["channel_name"] for row in layout["eeg"]],
                 modes=["synthetic structural state"],
                 sources=[f"S{row['index']:02d}" for row in layout["sources"]],
                 detectors=[f"D{row['index']:02d}" for row in layout["detectors"]])
    # Deterministic geometry-only selection; never based on correlation or fit.
    selected = [int(np.argmax(hb_xy[:, 1]))]
    for _ in range(11):
        distance = np.min(np.sum((hb_xy[:, None] - hb_xy[selected][None]) ** 2, axis=2), axis=1)
        distance[selected] = -1.
        selected.append(int(np.argmax(distance)))
    selected = np.asarray(selected, dtype=int)
    roles = {label: int(selected[i]) for i, label in enumerate("ABCDEFGHIJKL")}
    j, e, s, d = 51, 64, 18, 18
    dt, samples, burn, u_history = .5, 600, 128, 20
    total = samples + burn
    rng = np.random.default_rng(seed)
    transition = np.diag(.32 + .14 * (hb_xy[:, 1] + 1.) / 2.)
    for label in ("A", "B", "E", "F", "G", "H"):
        transition[roles[label], roles[label]] = 0.
    transition[roles["C"], roles["C"]] = .45
    transition[roles["D"], roles["D"]] = .40
    transition[roles["C"], roles["A"]] = .62
    transition[roles["D"], roles["C"]] = -.58
    transition[roles["J"], roles["I"]] = .34
    transition[roles["K"], roles["J"]] = .28
    u_loading = .05 * np.cos(2 * np.arctan2(hb_xy[:, 1], hb_xy[:, 0]))
    for label, loading in dict(A=.70, B=.65, E=.52, F=.50, G=-.48, H=-.46).items():
        u_loading[roles[label]] = loading
    common = np.zeros(total)
    source_noise = np.zeros((total, s))
    detector_noise = np.zeros((total, d))
    state = np.zeros((total, j))
    exchange = np.zeros((total, j))
    neural = np.zeros((total, 3))
    state_sd = np.full(j, .19)
    state_sd[[roles["A"], roles["B"]]] = .28
    for step in range(1, total):
        common[step] = .90 * common[step-1] + rng.normal(0., .45)
        source_noise[step] = .78 * source_noise[step-1] + rng.normal(0., .13, s)
        detector_noise[step] = .76 * detector_noise[step-1] + rng.normal(0., .13, d)
        state[step] = transition @ state[step-1] + u_loading * common[step] + rng.normal(0., state_sd)
        exchange[step] = .68 * exchange[step-1] + rng.normal(0., .13, j)
        neural[step] = .72 * neural[step-1] + rng.normal(0., .35, 3)
    hardware = source_noise[:, pairs[:, 0]] + detector_noise[:, pairs[:, 1]]
    observed_total = state + hardware
    hb = np.stack((.65 * observed_total + exchange, .35 * observed_total - exchange), axis=-1)
    hb_clean = np.stack((.65 * state + exchange, .35 * state - exchange), axis=-1)
    eeg_loading = np.column_stack((eeg_xy[:, 0], eeg_xy[:, 1], 1. - np.sum(eeg_xy**2, axis=1)))
    eeg = neural @ eeg_loading.T + .14 * common[:, None] + rng.normal(0., .18, (total, e))
    time = np.arange(samples) * dt
    observed, truth = hb[burn:], hb_clean[burn:]
    covariates = np.column_stack((common[burn:], source_noise[burn:], detector_noise[burn:]))
    partial_controls = np.column_stack((
        *[common[burn-lag:total-lag] for lag in range(u_history+1)],
        source_noise[burn:], detector_noise[burn:],
    ))
    hbo, hbr = observed[..., 0], observed[..., 1]
    hbt = hbo + hbr
    arrays = dict(time=time, hb_observation=observed, eeg_observation=eeg[burn:],
                  visible_hb=np.ones_like(observed, dtype=bool), score_hb=np.zeros_like(observed, dtype=bool),
                  visible_eeg=np.ones((samples, e), dtype=bool), score_eeg=np.zeros((samples, e), dtype=bool),
                  hb_truth=truth, hbo=hbo, hbr=hbr, hbt=hbt, structural_hbt=state[burn:],
                  hardware_hbt=hardware[burn:], exchange=exchange[burn:],
                  covariates=covariates, partial_controls=partial_controls,
                  selected_indices=selected, known_A=transition)
    for key, signal in (("corr_hbo", hbo), ("corr_hbr", hbr), ("corr_hbt", hbt),
                        ("corr_hbt_no_hardware", state[burn:])):
        arrays[key], arrays[key + "_count"] = correlation_matrix(signal)
    arrays["partial_hbt"], arrays["partial_hbt_count"] = residual_correlation(hbt, partial_controls)
    arrays["conditional_lag_gain"], arrays["lag_full_rss"], arrays["lag_restricted_rss"], n_lag = lag_prediction_gain(
        hbt[:, selected], covariates, p=2)
    direct = transition[np.ix_(selected, selected)].T.copy()
    np.fill_diagonal(direct, 0.)
    arrays["known_direct"] = direct
    response = intervention_response(transition, n_steps=20)
    arrays["do_response"] = response[selected][:, :, selected]
    arrays["known_total_effect"] = arrays["do_response"].mean(axis=1)
    arrays["do_time"] = np.arange(1, 21) * dt
    shared_source = pairs[:, None, 0] == pairs[None, :, 0]
    shared_detector = pairs[:, None, 1] == pairs[None, :, 1]
    pair_class = shared_source.astype(np.int8) + 2 * shared_detector.astype(np.int8)
    np.fill_diagonal(pair_class, -1)
    arrays["optode_pair_class"] = pair_class
    upper = np.triu(np.ones((j, j), dtype=bool), k=1)
    group_counts, group_mean, group_nohw = [], [], []
    for code in range(4):
        mask = upper & (pair_class == code)
        group_counts.append(int(mask.sum()))
        group_mean.append(float(np.mean(np.abs(arrays["corr_hbt"][mask]))) if mask.any() else np.nan)
        group_nohw.append(float(np.mean(np.abs(arrays["corr_hbt_no_hardware"][mask]))) if mask.any() else np.nan)
    arrays["optode_group_counts"] = np.asarray(group_counts, dtype=np.int64)
    arrays["optode_group_mean_abs_corr_hbt"] = np.asarray(group_mean)
    arrays["optode_group_mean_abs_corr_no_hardware"] = np.asarray(group_nohw)
    for lag_samples, suffix in ((0, "zero"), (4, "lag2s")):
        eeg_support = arrays["eeg_observation"] if lag_samples == 0 else arrays["eeg_observation"][:-lag_samples]
        hb_support = hbt if lag_samples == 0 else hbt[lag_samples:]
        joint, counts = correlation_matrix(np.column_stack((eeg_support, hb_support)))
        key = "crossmodal_corr_eeg_hbt_" + suffix
        arrays[key], arrays[key + "_count"] = joint[:e, e:], counts[:e, e:]
    neighbours = []
    for left in range(j):
        # Geometry proximity is only a display label, never a brain connection.
        distances = np.sum((hb_xy - hb_xy[left]) ** 2, axis=1)
        for right in np.argsort(distances, kind="stable")[1:3]:
            neighbours.append(tuple(sorted((left, int(right)))))
    support = dict(start_s=float(time[0]), stop_s=float(time[-1]), samples=samples,
                   dt_s=dt, aggregation="one synthetic seed; time samples are not independent subjects")
    oracle_names = ["U"] + names["sources"] + names["detectors"]
    partial_names = [f"U_lag{lag}" for lag in range(u_history+1)] + names["sources"] + names["detectors"]
    specs = {}

    def specification(key, label, metric, unit, node_indices, *, directed=False,
                      controls=None, estimator="", evidence_type="synthetic_statistical_association", **extra):
        specs[key] = dict(key=key, label=label, metric=metric, unit=unit,
                          source_node_group="hb", target_node_group="hb",
                          node_indices=list(map(int, node_indices)),
                          node_names=[names["hb"][index] for index in node_indices],
                          directed=directed, symmetric=not directed, matrix_axis="row_source_column_target",
                          support=support.copy(), controls=[] if controls is None else controls,
                          estimator=estimator, evidence_type=evidence_type, display_top_k=20,
                          display_selection="finite off-diagonal entries ranked by absolute value, stable ties; display only",
                          **extra)

    for key, label in (("corr_hbo", "HbO Pearson相关"), ("corr_hbr", "HbR Pearson相关"), ("corr_hbt", "HbT Pearson相关")):
        specification(key, label, "Pearson r", "1", range(j), estimator="pairwise-complete Pearson correlation, no temporal filtering or channel normalization",
                      count_key=key+"_count", min_count=20)
    specification("partial_hbt", "控制已知共因与光极噪声后的HbT残差相关", "oracle residual Pearson r", "1", range(j),
                  controls=partial_names, estimator="intercept OLS per node on one common complete support, then Pearson residual correlation",
                  count_key="partial_hbt_count", min_count=20,
                  control_time_support_s=[-u_history*dt, float(time[-1])],
                  warning="oracle generator covariates are unavailable in ordinary measured data; residual association is not causality")
    specification("conditional_lag_gain", "条件滞后预测增益", "log(RSS_restricted/RSS_full)", "natural log ratio", selected,
                  directed=True, controls=oracle_names,
                  estimator="same-sample multivariate nested OLS; remove source lags only, retain own/other selected-node histories and all oracle covariate histories",
                  lag_order=2, lag_seconds=[dt, 2*dt], regression_rows=n_lag,
                  warning="nonnegative predictive gain has no excitation/inhibition sign; no p-values, confidence intervals or significance claim")
    specs["conditional_lag_gain"]["support"].update(start_s=float(time[2]), samples=n_lag,
        history_support_s=[float(time[0]), float(time[-2])])
    specification("known_direct", "生成器已知直接系数", "off-diagonal A[target,source]", "model unit / model unit per step", selected,
                  directed=True, estimator="read the explicitly declared structural transition matrix; no estimation",
                  evidence_type="generator_structure", warning="causal meaning only inside this synthetic generator")
    specs["known_direct"]["support"] = dict(kind="time-invariant declared generator coefficient", dt_s=dt)
    specification("known_total_effect", "生成器干预总效应", "mean signed response at next steps 1..20 to a +1 source-state pulse", "target model unit / +1 source model unit", selected,
                  directed=True, estimator="exact A**k impulse propagation, equivalent to same-initial-state/same-noise counterfactual replay",
                  evidence_type="generator_intervention", intervention_amplitude=1., response_steps=20,
                  response_time_s=[dt, 20*dt], response_array="do_response", response_time_array="do_time",
                  warning="effects are between synthetic states assigned to Hb drawing positions, not causal interactions between source/detector hardware; includes indirect paths and excludes the instantaneous source pulse; not estimated from observed data")
    specs["known_total_effect"]["support"] = dict(kind="counterfactual response after the intervention", start_s=dt,
        stop_s=20*dt, samples=20, dt_s=dt, intervention_time_s=0., baseline="same factual state and same external forcing/noise")
    for lag_samples, suffix, label in ((0, "zero", "EEG与HbT零时差相关"), (4, "lag2s", "EEG(t)与HbT(t+2秒)固定时差相关")):
        key = "crossmodal_corr_eeg_hbt_" + suffix
        specs[key] = dict(
            key=key, label=label, metric="pairwise-complete Pearson r at a fixed declared time offset", unit="1",
            source_node_group="eeg", target_node_group="hb", source_node_indices=list(range(e)), target_node_indices=list(range(j)),
            source_node_names=names["eeg"], target_node_names=names["hb"],
            directed=False, symmetric=False, matrix_axis="row_source_column_target",
            support=dict(source_time_s=[0., float(time[samples-lag_samples-1])],
                         target_time_s=[float(time[lag_samples]), float(time[-1])],
                         aligned_samples=samples-lag_samples, dt_s=dt, target_offset_s=lag_samples*dt),
            controls=[], estimator="Pearson correlation of EEG[t] with HbT[t+offset], no regression controls; offset fixed before generation, no maximum-over-lag selection",
            evidence_type="synthetic_crossmodal_association", count_key=key+"_count", min_count=20,
            display_top_k=20, display_selection="finite entries ranked by absolute correlation; display only",
            warning="descriptive cross-correlation; a positive Hb time offset does not establish EEG-to-Hb causal direction",
        )
    example_nodes = {label: dict(selected_index=i, hb_index=int(selected[i]), name=names["hb"][selected[i]])
                     for i, label in enumerate("ABCDEFGHIJKL")}
    manifest = dict(
        schema=WHOLE_HEAD_SCHEMA, view="whole_head", kind="synthetic_demo", arrays_file="arrays.npz",
        title="全头皮布局与跨光极关系：独立合成机制演示",
        units=dict(time="s", geometry="normalized public figure drawing radius", hb="signed synthetic Hb model unit", eeg="synthetic EEG proxy unit", driver="synthetic state unit"),
        names=names,
        geometry=dict(hb_xy=hb_xy.tolist(), eeg_xy=eeg_xy.tolist(), source_xy=source_xy.tolist(),
                      detector_xy=detector_xy.tolist(), hb_pairs=pairs.tolist(), neighbour_edges=[list(edge) for edge in sorted(set(neighbours))]),
        head=dict(center_xy=[0., 0.], radius=1., x_positive="figure_right", y_positive="anterior",
                  coordinate_system="REFED Figure 2 common two-dimensional drawing, normalized isotropically by its head outline",
                  orientation_note="left/right refer only to the published drawing, not individual MRI, anatomical registration or cortical location",
                  reference=dict(name=layout["reference"]["channel_name"], xy=positions([layout["reference"]])[0].tolist(), role="reference_only")),
        coordinate_operator=dict(total=[1, 1], oxygenation=[.35, -.65]),
        snapshot_times_s=[60., 150.], selected_hb_indices=selected.tolist(),
        selection_rule="start at maximum anterior coordinate, then greedily maximize distance to the nearest selected Hb midpoint; stable native-order ties; 12 nodes; no signal-based selection",
        example_nodes=example_nodes, relation_specs=specs,
        optode_groups=dict(codes={"0": "non_shared", "1": "shared_source_only", "2": "shared_detector_only", "3": "shared_both"},
                           diagonal_code=-1, unit="unordered distinct Hb-channel pair", counts=group_counts,
                           classification="mutually exclusive; both is its own group (zero for unique source-detector pairs)",
                           statistic="mean absolute full-window Pearson HbT correlation across unordered pairs; no uncertainty or independent-subject claim"),
        provenance=dict(
            geometry_source="src/data/assets/refed_paper_layout_v1.json", geometry_asset_schema=layout["schema"],
            geometry_publication=layout["source"],
            geometry_transform="(page_point-center_page_pt)/radius_page_pt * [1,-1]; Hb is source-detector drawing midpoint",
            signal_source="independent synthetic data generated here; no REFED measurements, native CSV, caches, or measured runs read",
            generator="stable known linear structural state and observation generator; relation teaching only; not Balloon and not a physiology model",
            seed=seed, clock=dict(start_s=0., stop_s=float(time[-1]), dt_s=dt, samples=samples, burn_in_samples=burn),
            structural_equation="x[t]=A@x[t-1]+b*U[t]+independent Gaussian innovation; A axis [target,source]; U[t]=0.90U[t-1]+N(0,0.45^2)",
            observation_equation="hardware_j=source_noise[S_j]+detector_noise[D_j]; HbO=.65*(x+hardware)+exchange; HbR=.35*(x+hardware)-exchange; HbT=x+hardware",
            exchange_equation="exchange[t]=.68 exchange[t-1]+N(0,.13^2), independent by Hb node",
            hardware_equation="source noise AR .78; detector noise AR .76; independent innovations SD .13 per optical node; no proven device or shallow-physiology interpretation",
            eeg_equation="three independent AR(.72) synthetic processes mixed by [x_position,y_position,1-r^2], plus .14U and independent N(0,.18^2); no EEG connectivity or causal claims",
            initial_states="all zero before 128-sample burn-in", state_innovation_sd=state_sd.tolist(), U_loading=u_loading.tolist(),
            covariate_columns=oracle_names, partial_control_columns=partial_names,
            normalization="none; all signs and scales retained in common model coordinates",
            stability=dict(spectral_radius=float(np.max(np.abs(np.linalg.eigvals(transition)))),
                           reason="off-diagonal graph is acyclic; eigenvalues are self-AR coefficients, each below one"),
            truth_kind="synthetic generator states and noiseless synthetic observation, never hidden measured targets",
            analysis="all 600 generated observation samples are visible; no training, selection partition, SSM fitting, held-out scoring, p-value or confidence interval",
            interpretation="A/B share U but neither directly nor indirectly causes the other; A→C is positive, C→D negative, so A→D is indirect and negative; separate I→J→K is positive",
        ),
    )
    return manifest, arrays




def create_whole_head_demo_bundle(output, seed=20261016):
    """Save the separate whole-head synthetic example into a fresh directory."""
    manifest, arrays = build_whole_head_demo(seed=seed)
    validate_bundle(manifest, arrays)
    output = _fresh(output)
    np.savez_compressed(output / "arrays.npz", **arrays)
    path = output / "manifest.json"
    path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2)+"\n", encoding="utf-8")
    return path


def _relation_specs(manifest):
    specs = manifest.get("relation_specs", {})
    return specs if isinstance(specs, dict) else {item["key"]: item for item in specs}


def _validate_whole_head_bundle(m, a):
    """Validate an already exported whole-head display; never estimate relations."""
    if m.get("view") != "whole_head" or m.get("kind") not in {"synthetic_demo", "measured_export"}:
        raise ValueError("v2 requires view=whole_head and an explicit export kind")
    filename = m.get("arrays_file")
    if not isinstance(filename, str) or Path(filename).name != filename or not filename.endswith(".npz"):
        raise ValueError("arrays_file must be a sibling NPZ filename")
    for key in ("title",):
        _text(m.get(key), key)
    if not isinstance(m.get("provenance"), dict) or not m["provenance"]:
        raise ValueError("whole-head source provenance is required")
    if m.get("units", {}).get("time") != "s":
        raise ValueError("time unit must be s")
    for key, unit in m["units"].items():
        _text(unit, f"units.{key}")
        if any(token in unit.lower() for token in ("zscore", "z-score", "standardized", "标准差单位")):
            raise ValueError("whole-head fields must retain their common exported coordinates")
    names, g = m["names"], m["geometry"]
    for key in ("hb", "eeg", "sources", "detectors"):
        group = names[key]
        if not isinstance(group, list) or not group or len(set(group)) != len(group):
            raise ValueError(f"names.{key} must contain unique names")
        for name in group:
            _text(name, f"names.{key}")
    j, e, s, d = (len(names[key]) for key in ("hb", "eeg", "sources", "detectors"))
    if not 1 <= j <= 128 or not 1 <= e <= 128:
        raise ValueError("whole-head visualization supports 1..128 Hb and EEG nodes")
    for key, count in (("hb_xy", j), ("eeg_xy", e), ("source_xy", s), ("detector_xy", d)):
        _array(g[key], (count, 2), f"geometry.{key}")
    pairs = np.asarray(g["hb_pairs"])
    if pairs.shape != (j, 2) or pairs.dtype.kind not in "iu" or np.any(pairs < 0) or np.any(pairs[:, 0] >= s) or np.any(pairs[:, 1] >= d) or len(set(map(tuple, pairs))) != j:
        raise ValueError("hb_pairs must contain unique valid source/detector indices")
    head = m["head"]
    _array(head["center_xy"], (2,), "head.center_xy")
    if not np.isfinite(head["radius"]) or head["radius"] <= 0:
        raise ValueError("head radius must be finite and positive")
    if head.get("x_positive") != "figure_right" or head.get("y_positive") not in {"anterior", "figure_anterior"}:
        raise ValueError("whole-head coordinates must declare figure-right and anterior orientation")
    reference = m.get("reference", head.get("reference", g.get("reference")))
    if reference is not None:
        if not isinstance(reference, dict) or reference.get("role") != "reference_only":
            raise ValueError("a separate reference must declare reference_only")
        _text(reference.get("name"), "reference.name")
        if reference["name"].upper() in {name.upper() for name in names["eeg"]}:
            raise ValueError("the reference_only electrode must not be an EEG input channel")
        _array(reference["xy"], (2,), "reference.xy")
    if not isinstance(a, dict) or not REQUIRED_ARRAYS <= a.keys():
        raise ValueError("missing time, observations or visibility masks")
    time = np.asarray(a["time"])
    if time.ndim != 1 or time.dtype.kind not in "fiu" or len(time) < 3 or not np.isfinite(time).all() or np.any(np.diff(time) <= 0):
        raise ValueError("time must be finite and strictly increasing")
    for suffix, shape in (("hb", (len(time), j, 2)), ("eeg", (len(time), e))):
        observed = _array(a[f"{suffix}_observation"], shape, f"{suffix}_observation", finite=False)
        masks = []
        for prefix in ("visible", "score"):
            mask = np.asarray(a[f"{prefix}_{suffix}"])
            if mask.shape != shape or mask.dtype != np.bool_:
                raise ValueError(f"{prefix}_{suffix} must be boolean {shape}")
            if suffix == "hb" and not np.array_equal(mask[..., 0], mask[..., 1]):
                raise ValueError("Hb visibility/scoring must use complete pairs")
            masks.append(mask)
        if np.any(masks[0] & masks[1]) or not np.isfinite(observed[masks[0] | masks[1]]).all():
            raise ValueError("visible/scored samples must be finite and masks disjoint")
    for key, values in a.items():
        if np.asarray(values).dtype.kind not in "fiub" or np.isinf(values).any():
            raise ValueError(f"array {key} must be numeric/bool without infinity")
    for key, values in (("hbo", a["hb_observation"][..., 0]), ("hbr", a["hb_observation"][..., 1]), ("hbt", a["hb_observation"].sum(axis=-1))):
        if key in a:
            _array(a[key], (len(time), j), key, finite=False)
            if not np.allclose(a[key], values, equal_nan=True):
                raise ValueError(f"{key} does not agree with exported observations")
    snapshots = np.asarray(m["snapshot_times_s"])
    if snapshots.shape != (2,) or not np.isfinite(snapshots).all() or np.any(np.diff(snapshots) <= 0) or snapshots[0] < time[0] or snapshots[-1] > time[-1]:
        raise ValueError("whole-head snapshots require two increasing times in support")
    specs = _relation_specs(m)
    if not isinstance(specs, dict):
        raise ValueError("relation_specs must explicitly describe each exported relation")
    causal = {"known_direct", "known_total_effect", "do_response", "known_A", "hb_truth", "structural_hbt"}
    if a.keys() & causal and m["kind"] != "synthetic_demo":
        raise ValueError("generator causal truth is synthetic only; measured exports cannot inherit it")
    for key, spec in specs.items():
        if key not in a or not isinstance(spec, dict):
            raise ValueError(f"relation {key} must name one explicitly exported array")
        crossmodal = spec.get("source_node_group") == "eeg"
        groups = (("source_node_indices", e), ("target_node_indices", j)) if crossmodal else (("node_indices", j),)
        checked = []
        for field, upper in groups:
            indices = np.asarray(spec[field])
            if indices.ndim != 1 or indices.dtype.kind not in "iu" or not len(indices) or len(set(indices.tolist())) != len(indices) or np.any(indices < 0) or np.any(indices >= upper):
                raise ValueError("relation node indices must be unique valid channel indices")
            checked.append(indices)
        shape = (len(checked[0]), len(checked[-1]))
        values = _array(a[key], shape, key, finite=False)
        if crossmodal and (spec.get("target_node_group") != "hb" or not np.array_equal(checked[0], np.arange(e)) or not np.array_equal(checked[1], np.arange(j))):
            raise ValueError("crossmodal arrays must use the complete native EEG/Hb node order")
        if spec.get("matrix_axis") != "row_source_column_target":
            raise ValueError("relations must declare row_source_column_target")
        for field in ("label", "metric", "unit", "support", "estimator"):
            if field not in spec:
                raise ValueError(f"relation {key} must declare {field}")
        if not isinstance(spec.get("directed"), bool) or not isinstance(spec.get("symmetric"), bool):
            raise ValueError("relation directed/symmetric declarations must be boolean")
        if "count_key" in spec:
            counts = np.asarray(a.get(spec["count_key"]))
            support = spec["support"] if isinstance(spec["support"], dict) else {}
            upper = min(len(time), support.get("aligned_samples", support.get("samples", len(time))))
            if counts.shape != shape or counts.dtype.kind not in "iu" or np.any(counts < 0) or np.any(counts > upper):
                raise ValueError("relation counts must be matching integers within declared time support")
            if "min_count" in spec and np.any(np.isfinite(values) & (counts < spec["min_count"])):
                raise ValueError("finite relations require at least the declared minimum pair count")
        if spec["symmetric"] and not np.allclose(values, values.T, equal_nan=True, atol=1e-9):
            raise ValueError("symmetric relation matrix is not symmetric")
        if key.startswith(("corr_", "crossmodal_corr_")) or key == "partial_hbt":
            if np.nanmax(np.abs(values)) > 1 + 1e-8:
                raise ValueError("correlation must lie in [-1,1]")
        if key == "conditional_lag_gain" and np.nanmin(values) < -1e-9:
            raise ValueError("nested conditional predictive gain must be nonnegative")
        if key in {"known_direct", "known_total_effect"} and (m["kind"] != "synthetic_demo" or spec.get("evidence_type") not in ({"generator_structure", "generator_intervention"} if key == "known_direct" else {"generator_intervention"})):
            raise ValueError("causal arrays require synthetic generator_intervention evidence")
    for key in ("corr_hbo", "corr_hbr", "corr_hbt", "partial_hbt", "conditional_lag_gain", "known_direct", "known_total_effect"):
        if key in a and key not in specs:
            raise ValueError(f"relation {key} lacks its explicit evidence description")
    if {"corr_hbt", "partial_hbt"} <= specs.keys() and not np.array_equal(specs["corr_hbt"]["node_indices"], specs["partial_hbt"]["node_indices"]):
        raise ValueError("paired Pearson/partial maps must use identical ordered Hb indices")


def interpolate_head_field(xy, values, head, grid_size=180, max_edge=.48, max_distance=.22):
    """Local linear interpolation only: head ∩ hull ∩ short triangles ∩ near sensor.

    Distances are fractions of the metadata head radius. No extrapolation or
    spatial smoothing is performed. An insufficient geometry remains point-only.
    """
    from scipy.spatial import Delaunay, cKDTree, QhullError
    xy, values = np.asarray(xy, float), np.asarray(values, float)
    if xy.ndim != 2 or xy.shape[1] != 2 or values.shape != (len(xy),):
        raise ValueError("xy must be [N,2] and values [N]")
    if type(grid_size) is not int or grid_size < 5 or max_edge <= 0 or max_distance <= 0:
        raise ValueError("grid size and interpolation distance limits must be positive")
    center, radius = np.asarray(head["center_xy"], float), float(head["radius"])
    gx, gy = np.meshgrid(np.linspace(center[0]-radius, center[0]+radius, grid_size), np.linspace(center[1]-radius, center[1]+radius, grid_size))
    result = dict(x=gx, y=gy, z=np.full(gx.shape, np.nan), mask=np.ones(gx.shape, bool), method="points_only", max_edge_radius=float(max_edge), max_distance_radius=float(max_distance))
    good = np.isfinite(values) & np.isfinite(xy).all(axis=1)
    points, vals = xy[good], values[good]
    if len(points) < 3 or np.linalg.matrix_rank(points-points[0]) < 2:
        return result
    try:
        tri = Delaunay(points)
    except QhullError:
        return result
    queries = np.column_stack((gx.ravel(), gy.ravel()))
    simplex = tri.find_simplex(queries)
    vertices = points[tri.simplices]
    longest = np.max(np.stack([np.linalg.norm(vertices[:, i]-vertices[:, (i+1) % 3], axis=1) for i in range(3)]), axis=0)
    valid = (simplex >= 0) & (np.linalg.norm(queries-center, axis=1) < radius)
    valid &= cKDTree(points).query(queries)[0] <= max_distance*radius
    valid &= longest[np.maximum(simplex, 0)] <= max_edge*radius
    if valid.any():
        si, q = simplex[valid], queries[valid]
        bary = np.einsum("nij,nj->ni", tri.transform[si, :2], q-tri.transform[si, 2])
        weights = np.column_stack((bary, 1-bary.sum(axis=1)))
        flat = result["z"].ravel()
        flat[valid] = np.sum(vals[tri.simplices[si]]*weights, axis=1)
        result["mask"] = ~valid.reshape(gx.shape)
        result["method"] = "local_linear_strict_hull_head_short_triangles_near_sensor"
    return result


def _whole_canvas(plt, m, title, footer, rows=1, cols=1, figsize=(14, 9.2), **kwargs):
    import fitz
    fig, axes = _canvas(plt, m, title, footer, rows=rows, cols=cols, figsize=figsize, **kwargs)
    fig.texts[-1].set_text(_wrap_pdf(fitz.Font(fontfile=str(FONT_PATH)), footer, figsize[0]*72*.90, 10.5))
    fig.texts[-1].set_fontsize(10.5)
    fig.subplots_adjust(left=.055, right=.91, top=.845, bottom=.19, wspace=.25, hspace=.36)
    return fig, axes


def _signed_cmap():
    from matplotlib import colormaps
    from matplotlib.colors import ListedColormap
    colors = colormaps["RdBu_r"](np.linspace(0., 1., 257))
    colors[128] = [1., 1., 1., 1.]
    cmap = ListedColormap(colors, name="signed_white_zero")
    cmap.set_bad("#D6D9DD")
    return cmap


def _head(ax, m):
    from matplotlib.patches import Circle, Polygon, Arc
    center = np.asarray(m["head"]["center_xy"])
    r = m["head"]["radius"]
    ax.add_patch(Circle(center, r, facecolor="#F7F8FA", edgecolor="#778493", lw=1.4, zorder=0))
    ax.add_patch(Circle(center, r, facecolor="none", edgecolor="#D2D8DE", hatch="///", lw=0, zorder=.1))
    ax.add_patch(Polygon(center+np.array([[-.09, .995], [0, 1.12], [.09, .995]])*r, closed=False, fill=False, edgecolor=INK, lw=1.4, zorder=8))
    for sign in (-1, 1):
        ax.add_patch(Arc(center+np.array([sign*1.01, 0])*r, .16*r, .38*r, angle=0, theta1=80 if sign < 0 else -100, theta2=280 if sign < 0 else 100, color="#778493", lw=1.3, zorder=8))
    for label, point in (("前 / 鼻", (0, 1.2)), ("后", (0, -1.15)), ("图左", (-1.17, -.4)), ("图右", (1.17, -.4))):
        p = center+np.array(point)*r
        ax.text(*p, label, ha="center", va="center", fontsize=8.5, color="#647382", zorder=10)
    ax.set(xlim=(center[0]-1.3*r, center[0]+1.3*r), ylim=(center[1]-1.25*r, center[1]+1.28*r), aspect="equal")
    ax.set_axis_off()


def _whole_nodes(ax, m, indices=None, *, labels=True, highlight=None, size=26, color=INK):
    xy = np.asarray(m["geometry"]["hb_xy"])
    ax.scatter(*xy.T, s=size, c=color, edgecolors="white", linewidths=.4, zorder=5)
    indices = np.arange(len(xy)) if indices is None else np.asarray(indices)
    if labels:
        for i in indices:
            ax.text(*xy[i], str(i+1), fontsize=7.5, ha="center", va="center", color=INK, zorder=7,
                    bbox=dict(boxstyle="round,pad=.08", facecolor="white", edgecolor="none", alpha=.8))
    if highlight is not None:
        ax.scatter(*xy[np.asarray(highlight)].T, s=75, facecolors="none", edgecolors=ORANGE, lw=1.5, zorder=8)


def _top_edges(values, directed=False, count=20):
    values = np.asarray(values)
    rows, cols = np.where(np.isfinite(values) & (np.abs(values) > 1e-12))
    edges = [(int(i), int(j)) for i, j in zip(rows, cols) if i != j and (directed or i < j)]
    return sorted(edges, key=lambda pair: (-abs(float(values[pair])), pair))[:count]


def _network(ax, m, values, indices, *, directed=False, edges=None, title="", limit=1., positive_only=False):
    from matplotlib.patches import FancyArrowPatch
    from matplotlib.colors import Normalize
    _head(ax, m)
    xy = np.asarray(m["geometry"]["hb_xy"])[indices]
    edges = _top_edges(values, directed) if edges is None else edges
    cmap = __import__("matplotlib").colormaps["YlOrRd"] if positive_only else _signed_cmap()
    norm = Normalize(0 if positive_only else -limit, limit)
    for i, j in edges:
        value = float(values[i, j])
        if not np.isfinite(value):
            continue
        width = .6+2.7*min(abs(value)/max(limit, 1e-12), 1.)
        color = cmap(norm(value))
        style = "-" if value >= 0 else "--"
        if directed:
            patch = FancyArrowPatch(xy[i], xy[j], arrowstyle="-|>", mutation_scale=10+width, linewidth=width, color=color, linestyle=style, shrinkA=7, shrinkB=7, connectionstyle="arc3,rad=.10", zorder=3, alpha=.86)
            ax.add_patch(patch)
        else:
            ax.plot(xy[[i, j], 0], xy[[i, j], 1], color=color, lw=width, ls=style, alpha=.85, zorder=3)
    _whole_nodes(ax, m, indices=indices, labels=True, size=16, color="#ABB4BE")
    missing = ~np.isfinite(values).any(axis=1)
    if missing.any():
        ax.scatter(*xy[missing].T, color=GRAY, marker="x", s=70, zorder=9)
    possible = len(indices)*(len(indices)-1)//(1 if directed else 2)
    valid = np.isfinite(values) & ~np.eye(len(indices), dtype=bool)
    valid_count = int(valid.sum()) if directed else int(np.triu(valid,1).sum())
    ax.set_title(f"{title}\n显示 {len(edges)} / 有效 {valid_count}；可能 {possible} 条", fontsize=11)
    return dict(displayed_edges=[list(pair) for pair in edges], displayed_edge_count=len(edges), valid_edge_count=valid_count, possible_edge_count=possible, node_indices=np.asarray(indices).tolist(), threshold_is_display_only=True)


def _matrix(ax, m, values, indices, title, *, limit=1., positive_only=False):
    cmap = __import__("matplotlib").colormaps["YlOrRd"].copy() if positive_only else _signed_cmap()
    cmap.set_bad("#D6D9DD")
    im = ax.imshow(np.ma.masked_invalid(values), cmap=cmap, vmin=0 if positive_only else -limit, vmax=limit, interpolation="nearest", aspect="equal")
    n = len(indices)
    ticks = np.arange(n) if n <= 15 else np.unique(np.r_[0, np.arange(4, n-3, 5), n-1])
    labels = [str(int(indices[i])+1) for i in ticks]
    ax.set_xticks(ticks, labels, fontsize=8)
    ax.set_yticks(ticks, labels, fontsize=8)
    ax.set(xlabel="目标 Hb 编号（列）", ylabel="来源 Hb 编号（行）", title=title)
    return im


def _whole_layout(plt, m, a):
    reference = m.get("reference", m["head"].get("reference", m["geometry"].get("reference")))
    n_eeg = len(m["names"]["eeg"])
    geometry_note = "公开论文坐标按原图头圈中心与半径归一化，并翻转 PDF 的 Y 轴。" if m["provenance"].get("geometry_source") == "src/data/assets/refed_paper_layout_v1.json" else f"坐标与头圈来自显式导出，单位：{m['units']['geometry']}。"
    reference_note = f"{reference['name']} 仅为独立参考，不计入 {n_eeg} 路 EEG。" if reference else "未导出独立参考位置。"
    footer = geometry_note+"鼻朝上，图左／图右不替代个体解剖侧别。\n"+reference_note+"Hb 点与源探测配对见 CSV；斜线为未填色底图，不是个体 MRI 或精确皮层定位。"
    fig, axes = _whole_canvas(plt, m, "01  全头视图显示传感器布局与覆盖空白", footer, cols=2, figsize=(14, 9.4))
    _head(axes[0, 0], m)
    xy = np.asarray(m["geometry"]["eeg_xy"])
    axes[0, 0].scatter(*xy.T, s=35, c="#D8E7F0", edgecolors=BLUE, lw=.7, zorder=4)
    for label, point in zip(m["names"]["eeg"], xy):
        axes[0, 0].annotate(label, point, xytext=(0, 5), textcoords="offset points", ha="center", fontsize=8, color=INK)
    if reference:
        axes[0, 0].scatter(*reference["xy"], marker="*", c=ORANGE, s=100, zorder=7)
        axes[0, 0].annotate(f"{reference['name']} 参考", reference["xy"], xytext=(0, -15), textcoords="offset points", ha="center", fontsize=8, color=ORANGE, zorder=8)
    axes[0, 0].set_title(f"{len(xy)} 路 EEG：完整通道名"+("；星号为独立参考" if reference else ""))
    _head(axes[0, 1], m)
    _whole_nodes(axes[0, 1], m, size=100)
    axes[0, 1].set_title(f"{len(m['names']['hb'])} 个 Hb 位置：编号与完整 CSV 配对表对应")
    return fig, dict(source_arrays=[], coordinate_source=m["provenance"].get("geometry_source"), head_limits_from_metadata=True, eeg_inputs=len(xy), hb_nodes=len(m["names"]["hb"]), reference_only=reference is not None)


def _whole_hardware(plt, m, a):
    fig, axes = _whole_canvas(plt, m, "02  源探测路径与共享光极构成硬件关系", "左：真实 S–D 配对路径，橙方为光源 S，蓝菱为探测器 D；高亮第一条 Hb 路径及共享光极的路径。\n中：全部 Hb 节点，边仅表示与高亮 Hb 共用 S（橙实线）或 D（蓝虚线）；不是神经连接。右：完整共享矩阵，0=不同光极，1=同 S，2=同 D。", cols=3, figsize=(15, 8.7))
    g = m["geometry"]
    source, detector, hb = (np.asarray(g[key]) for key in ("source_xy", "detector_xy", "hb_xy"))
    pairs = np.asarray(g["hb_pairs"])
    same_s = pairs[:, 0, None] == pairs[None, :, 0]
    same_d = pairs[:, 1, None] == pairs[None, :, 1]
    groups = same_s.astype(int)+2*same_d.astype(int)
    np.fill_diagonal(groups, 0)
    for ax in axes[0, :2]:
        _head(ax, m)
    for i, (s, d) in enumerate(pairs):
        highlighted = same_s[0, i] or same_d[0, i]
        axes[0, 0].plot([source[s, 0], detector[d, 0]], [source[s, 1], detector[d, 1]], color=ORANGE if highlighted else "#BFC8D0", lw=2.2 if highlighted else .7, alpha=1 if highlighted else .6, zorder=2)
    axes[0, 0].scatter(*source.T, marker="s", c=ORANGE, edgecolors="white", s=50, zorder=5, label="光源 S")
    axes[0, 0].scatter(*detector.T, marker="D", c=BLUE, edgecolors="white", s=45, zorder=5, label="探测器 D")
    for prefix, xy in (("S", source), ("D", detector)):
        for i, point in enumerate(xy):
            axes[0, 0].annotate(f"{prefix}{i+1}", point, xytext=(3, 4 if prefix == 'S' else -9), textcoords="offset points", fontsize=6.5, zorder=7)
    axes[0, 0].scatter(*hb.T, c=INK, s=8, zorder=4)
    axes[0, 0].legend(frameon=False, ncol=2, loc="lower center", bbox_to_anchor=(.5, -.08))
    axes[0, 0].set_title(f"{len(source)} 光源 × {len(detector)} 探测器；{len(pairs)} 条测量路径")
    for i in range(1, len(hb)):
        if same_s[0, i] or same_d[0, i]:
            axes[0, 1].plot(hb[[0, i], 0], hb[[0, i], 1], color=ORANGE if same_s[0, i] else BLUE, ls="-" if same_s[0, i] else "--", lw=2.5, zorder=3)
    _whole_nodes(axes[0, 1], m, highlight=[0])
    axes[0, 1].set_title("Hb 1 的共享光极邻接；全部 Hb 仍在全头位置")
    from matplotlib.colors import ListedColormap, BoundaryNorm
    im = axes[0, 2].imshow(groups, cmap=ListedColormap(["white", ORANGE, BLUE]), norm=BoundaryNorm([-.5,.5,1.5,2.5],3), interpolation="nearest")
    axes[0, 2].set(title="所有 Hb 对的共享光极矩阵", xlabel="Hb 编号（列）", ylabel="Hb 编号（行）")
    ticks = [0, 9, 19, 29, 39, 50] if len(hb) == 51 else np.arange(0, len(hb), max(1,len(hb)//5))
    axes[0, 2].set_xticks(ticks, np.asarray(ticks)+1)
    axes[0, 2].set_yticks(ticks, np.asarray(ticks)+1)
    fig.colorbar(im, ax=axes[0, 2], fraction=.045, pad=.035, ticks=[0,1,2], label="共享类别；对角排除")
    return fig, dict(source_arrays=[], highlighted_hb=0, shared_source_edges=int(np.triu(same_s,1).sum()), shared_detector_edges=int(np.triu(same_d,1).sum()), hardware_only=True)


def _whole_heatmaps(plt, m, a):
    from matplotlib.colors import Normalize
    time = a["time"]
    values = [a.get("hbo", a["hb_observation"][..., 0]), a.get("hbr", a["hb_observation"][..., 1]), a.get("hbt", a["hb_observation"].sum(axis=-1))]
    indices = [int(np.argmin(abs(time-t))) for t in m["snapshot_times_s"]]
    hb_unit = "合成相对 Hb 单位" if m["kind"] == "synthetic_demo" else m["units"]["hb"]
    footer = f"列为两个预先固定时点；行分别为 HbO、HbR 和 HbT=HbO+HbR，单位：{hb_unit}。每行跨时点固定同一色条，白色=0。\n仅在线性三角网凸包∩头圈内局部插值；排除边长>0.48 头半径的三角和距最近传感器>0.22 半径的点。斜线=无覆盖，灰叉=缺失。"
    fig, axes = _whole_canvas(plt, m, "03  全头 Hb 热图保留实际覆盖，空白区域不外推", footer, rows=3, cols=2, figsize=(12.8, 13.1))
    fig.subplots_adjust(hspace=.19, wspace=.15, bottom=.145)
    xy = np.asarray(m["geometry"]["hb_xy"])
    records = []
    cmap = _signed_cmap()
    cmap.set_bad((0., 0., 0., 0.))
    for row, (label, vals) in enumerate(zip(("HbO", "HbR", "HbT"), values)):
        finite = vals[np.isfinite(vals)]
        limit = max(float(np.max(np.abs(finite))) if len(finite) else 0., 1e-8)
        norm = Normalize(-limit, limit)
        for col, ti in enumerate(indices):
            ax = axes[row, col]
            _head(ax, m)
            field = interpolate_head_field(xy, vals[ti], m["head"])
            im = ax.pcolormesh(field["x"], field["y"], np.ma.array(field["z"], mask=field["mask"]), cmap=cmap, norm=norm, shading="auto", zorder=1)
            valid = np.isfinite(vals[ti])
            ax.scatter(*xy[valid].T, c=vals[ti, valid], cmap=cmap, norm=norm, s=22, edgecolors=INK, lw=.4, zorder=4)
            if (~valid).any():
                ax.scatter(*xy[~valid].T, marker="x", c=GRAY, s=60, zorder=5)
            if field["method"] == "points_only":
                ax.text(.5, .1, "有效位置不足／共线：仅显示测量点", transform=ax.transAxes, ha="center", fontsize=9, color=ORANGE)
            ax.set_title(f"{label} | t={time[ti]:g} s", fontsize=12)
            records.append(dict(chromophore=label, time_s=float(time[ti]), color_limits=[-limit,limit], interpolation=field["method"], unobserved_grid_cells=int(field["mask"].sum())))
        fig.colorbar(im, ax=list(axes[row]), fraction=.018, pad=.025, label=hb_unit)
    return fig, dict(source_arrays=["hb_observation"], snapshots=records, spatial_interpolation="strict local linear; no extrapolation", thresholds_radius=dict(max_edge=.48,max_distance=.22))


def _whole_correlations(plt, m, a):
    specs = _relation_specs(m)
    keys = [key for key in ("corr_hbo", "corr_hbr", "corr_hbt") if key in specs]
    fig, axes = _whole_canvas(plt, m, "04  跨光极相关在全头位置与完整矩阵中同时呈现", "Pearson r 使用导出声明的完整时间支持；每个节点是一条 Hb 测量，不是独立被试。橙红实线=正相关，蓝虚线=负相关，色条统一为 [−1,1]。\n地图仅显示 |r| 最大的至多 20 条边，筛选只影响显示；右侧与 NPZ 保留全部 Hb 对。相关不提供方向或因果证据。", rows=2, cols=max(1,len(keys)), figsize=(15, 10.2))
    records = {}
    for col, key in enumerate(keys):
        spec, value = specs[key], a[key]
        ids = np.asarray(spec["node_indices"])
        records[key] = _network(axes[0,col], m, value, ids, title=spec["label"])
        im = _matrix(axes[1,col], m, value, ids, "完整 Pearson r 矩阵")
    if keys:
        fig.colorbar(im, ax=list(axes[1]), fraction=.018, pad=.025, label="Pearson r；无量纲")
    else:
        for ax in axes.ravel():
            ax.text(.5,.5,"未导出相关关系；不在渲染时估计",ha="center",va="center",transform=ax.transAxes)
            ax.set_axis_off()
    return fig, dict(source_arrays=keys, relations=records, r_limits=[-1,1], relation_estimation_performed=False)


def _whole_partial(plt, m, a):
    specs = _relation_specs(m)
    keys = [key for key in ("corr_hbt", "partial_hbt") if key in specs]
    if m["kind"] == "synthetic_demo":
        footer = "左为 HbT Pearson，右为 oracle 部分相关：控制已知 U 当前及过去 10 s（21 点）、18 路 source noise 与 18 路 detector noise。\n按两矩阵的 max(|r|) 选择同一组至多 20 条边；完整矩阵均保留。控制变量是合成生成器真值；不是实测可得校正，也不是因果检验。"
        if len(keys) == 2:
            nodes = m["example_nodes"]
            first, second = nodes["A"]["hb_index"], nodes["B"]["hb_index"]
            footer += f"\n预先设定共因对 A/B：Pearson={a['corr_hbt'][first,second]:.3f}，oracle 部分相关={a['partial_hbt'][first,second]:.3f}；该差异不等于因果效应。"
    else:
        controls = ", ".join(map(str, specs.get("partial_hbt", {}).get("controls", []))) or "未导出"
        footer = f"左为导出 HbT Pearson，右为按导出控制变量计算的残差相关。控制变量：{controls}。\n地图共享显示边集与色标；只读取保留关系，不在渲染时重新拟合。部分相关不是因果检验。"
    fig, axes = _whole_canvas(plt, m, "05  控制变量前后的相关关系需要并排核对", footer, rows=2, cols=2, figsize=(14, 10.1))
    records={}
    if len(keys)==2:
        union = np.maximum(abs(a[keys[0]]), abs(a[keys[1]]))
        edges = _top_edges(union)
        for col,key in enumerate(keys):
            ids = np.asarray(specs[key]["node_indices"])
            records[key] = _network(axes[0,col],m,a[key],ids,edges=edges,title="HbT Pearson" if col==0 else ("Oracle 控制后的部分相关" if m["kind"] == "synthetic_demo" else "导出控制变量下的部分相关"))
            im = _matrix(axes[1,col],m,a[key],ids,"完整矩阵；同一 r 色标")
        fig.colorbar(im,ax=list(axes[1]),fraction=.018,pad=.025,label="相关系数；无量纲")
    else:
        for ax in axes.ravel():
            ax.text(.5,.5,"未导出配对相关／部分相关；不重新拟合",ha="center",va="center",transform=ax.transAxes)
            ax.set_axis_off()
    return fig,dict(source_arrays=keys,relations=records,controls="oracle known common/source/detector noise",relation_estimation_performed=False)


def _whole_lag(plt,m,a):
    key="conditional_lag_gain"
    spec=_relation_specs(m)[key]
    ids=np.asarray(spec["node_indices"])
    value=a[key]
    limit=max(float(np.nanmax(value)),1e-8)
    order = spec.get("lag_order", "导出声明的")
    dt = spec.get("support", {}).get("dt_s", "导出声明的")
    lag_seconds = spec.get("lag_seconds", [])
    history = f"{max(lag_seconds):g} s" if lag_seconds else "见导出时间支持"
    controls = "已知 U 与 36 路光极噪声的历史" if m["kind"] == "synthetic_demo" else ", ".join(map(str,spec.get("controls",[]))) or "无额外协变量"
    footer = f"地图箭头：来源的过去 → 目标的现在；行=来源，列=目标。{order} 个滞后、dt={dt} s，历史支持 {history}；颜色／宽度只表示非负增益。\n同样本嵌套 OLS：gain=log(RSS_reduced/RSS_full)，约简模型仅去掉来源历史，保留目标／其他节点历史；额外控制：{controls}。\n增益为样本内描述，未估计显著性或置信区间，不能解释成兴奋／抑制或实测因果。"
    fig,axes=_whole_canvas(plt,m,"06  条件滞后增益给出方向预测，不能代替干预因果",footer,cols=2,figsize=(14,8.8))
    record=_network(axes[0,0],m,value,ids,directed=True,title=f"{len(ids)} 个预先选定 Hb 节点；其余节点仍保留",limit=limit,positive_only=True)
    im=_matrix(axes[0,1],m,value,ids,"完整条件滞后预测矩阵",limit=limit,positive_only=True)
    fig.colorbar(im,ax=axes[0,1],fraction=.045,pad=.04,label="log RSS 比；无量纲，越大预测增益越大")
    return fig,dict(source_arrays=[key],relation=record,limits=[0,limit],matrix_axis="row_source_column_target",interpretation="conditional temporal prediction; not measured causal evidence")


def _whole_causal(plt,m,a):
    keys=("known_direct","known_total_effect")
    specs=_relation_specs(m)
    fig,axes=_whole_canvas(plt,m,"07  仅已知合成机制允许区分直接结构与总干预效应", "生成器真值：线性时序系统，非 Balloon 模型。左：直接跨节点系数 A[target,source] 转为行来源、列目标；不含自回归。\n右：同初态同噪声，把来源状态在 t=0 置高 1 单位；取 0.5..10 s 共 20 点平均响应差，含间接路径。总效应矩阵保留自身持续响应，地图排除自环。\n因果节点是放在 Hb 位置的合成状态；正实线／负虚线不代表实测兴奋／抑制。两列单位不同，分别固定对称色标。",rows=2,cols=2,figsize=(14,10.4))
    records={}
    for col,key in enumerate(keys):
        ids=np.asarray(specs[key]["node_indices"])
        value=a[key]
        limit=max(float(np.nanmax(abs(value))),1e-8)
        label="直接结构系数" if col==0 else "do 脉冲后 10 s 平均总效应"
        records[key]=_network(axes[0,col],m,value,ids,directed=True,title=label,limit=limit)
        im=_matrix(axes[1,col],m,value,ids,label+"：完整矩阵",limit=limit)
        fig.colorbar(im,ax=axes[1,col],fraction=.045,pad=.035,label="合成响应 / 单位干预" if col else "每步结构系数")
    return fig,dict(source_arrays=list(keys),relations=records,evidence_type="generator_intervention",generator_truth_only=True)


def _whole_counterexample(plt,m,a):
    specs=_relation_specs(m)
    ids=np.asarray(specs["known_direct"]["node_indices"])
    case=m.get("counterexample",{})
    local=np.asarray(case.get("selected_local_indices",[0,1,2,3]),int)
    if local.shape!=(4,):
        local=np.arange(4)
    nodes=ids[local]
    labels=[f"{letter}=Hb{int(node)+1}" for letter,node in zip("ABCD",nodes)]
    corr=a["corr_hbt"]
    direct=a["known_direct"]
    total=a["known_total_effect"]
    ab=float(corr[nodes[0],nodes[1]])
    ab_do=float(total[local[0],local[1]])
    ac=float(total[local[0],local[2]])
    ad=float(total[local[0],local[3]])
    fig,axes=_whole_canvas(plt,m,"08  共同原因可产生相关；方向传播必须用已知干预核验", f"预先设计的共因反例：A 与 B 同受 U 驱动，无 A→B 结构路径。Pearson r(A,B)={ab:.3f}，do(A)→B 的 10 s 平均效应={ab_do:.4f}。\n同一生成器包含 A→C→D：do(A)→C={ac:.4f}，do(A)→D={ad:.4f}。相关、条件预测、直接结构、总效应回答不同问题；未运行实测拟合或置信区间。",rows=2,cols=2,figsize=(14,10.1))
    _head(axes[0,0],m)
    _whole_nodes(axes[0,0],m,indices=nodes,highlight=nodes,size=18,color="#BEC5CE")
    xy=np.asarray(m["geometry"]["hb_xy"])
    for label,node in zip(labels,nodes):
        axes[0,0].annotate(label,xy[node],xytext=(5,9),textcoords="offset points",fontsize=9,color=ORANGE,zorder=10)
    for src,tgt,color,style in ((nodes[0],nodes[2],ORANGE,"-"),(nodes[2],nodes[3],BLUE,"--")):
        axes[0,0].annotate("",xy=xy[tgt],xytext=xy[src],arrowprops=dict(arrowstyle="->",color=color,linestyle=style,lw=2.3,connectionstyle="arc3,rad=.18"),zorder=8)
    axes[0,0].plot(xy[nodes[:2],0],xy[nodes[:2],1],"-",color=ORANGE,lw=1.1,alpha=.6,zorder=3)
    axes[0,0].set_title("无箭头细线=A/B 相关；箭头=真值路径，正实／负虚")
    axes[0,0].text(.5,-.06,"已知非空间共同驱动：U→A，U→B；无 A→B 路径",transform=axes[0,0].transAxes,fontsize=8,ha="center",va="top")
    time=a["time"]
    signal=a.get("hbt",a["hb_observation"].sum(axis=-1))
    for i,(node,label,color) in enumerate(zip(nodes[:2],labels[:2],(BLUE,ORANGE))):
        axes[0,1].plot(time,signal[:,node],color=color,lw=1.2,ls="-" if i==0 else "--",label=label)
    axes[0,1].set(title="共同驱动下 A/B 波形可相关",xlabel="时间 / s",ylabel=m["units"]["hb"])
    axes[0,1].legend(frameon=False)
    response=a["do_response"][local[0]]
    do_time=a.get("do_time",np.arange(response.shape[0])*.5)
    for i,(index,label,color) in enumerate(zip(local[1:],labels[1:],(BLUE,ORANGE,"#7A59A2"))):
        axes[1,0].plot(do_time,response[:,index],color=color,ls=("--","-",":")[i],lw=2,label=label)
    axes[1,0].axhline(0,color=GRAY,lw=.7)
    axes[1,0].set(title="固定噪声 do(A)+1：B 无响应，C/D 沿路径传播",xlabel="干预后时间 / s",ylabel="响应差 / 合成 Hb 单位")
    axes[1,0].legend(frameon=False)
    ax=axes[1,1]
    ax.set_axis_off()
    rows=[labels[1],labels[2],labels[3]]
    text="来源固定为 "+labels[0]+"\n\n目标               相关 r       直接系数       总效应\n"
    for li,node,label in zip(local[1:],nodes[1:],rows):
        text+=f"{label:12s}   {corr[nodes[0],node]: .3f}       {direct[local[0],li]: .3f}         {total[local[0],li]: .4f}\n\n"
    ax.text(.02,.95,text,va="top",fontsize=11,linespacing=1.4,transform=ax.transAxes)
    ax.text(.02,.1,"三个目标为预先设定机制例，非按效果挑选。\n一次合成实现不等于独立被试证据；无 CI／p 值。",va="bottom",fontsize=10,color="#647382",transform=ax.transAxes)
    for ax in (axes[0,1],axes[1,0]):
        _clean(ax)
    return fig,dict(source_arrays=["corr_hbt","known_direct","known_total_effect","do_response","hb_observation"],predeclared_case_nodes=nodes.tolist(),pearson_ab=ab,do_ab=ab_do,do_ac=ac,do_ad=ad,example_selection="predeclared generator mechanism; not performance selected")


def _whole_specs(m,a):
    n = m["names"]
    reference = m.get("reference",m["head"].get("reference",m["geometry"].get("reference")))
    reference_note = f"{reference['name']} 独立标为参考。" if reference else "未导出独立参考位置。"
    partial_explanation = "方法：比较导出的 HbT 相关与控制已知共同原因、source noise 和 detector noise 后的部分相关。两图共享显示边集和色标。合成真值控制变量不是实测可得信息；本页不宣称因果或统计显著性。" if m["kind"] == "synthetic_demo" else "方法：比较显式导出的 HbT Pearson 与控制变量下的残差相关，变量及支持以导出说明为准。两图共享显示边集和色标；不重新拟合，不把条件相关解释成因果。未导出时显示未估计状态。"
    specs=[
        ("01_whole_head_layout","全头传感器布局与覆盖",_whole_layout,f"方法：读取显式导出的共同二维几何；{len(n['eeg'])} EEG、{len(n['hb'])} Hb、{len(n['sources'])} 光源、{len(n['detectors'])} 探测器位于同一固定头圈。{reference_note}图示二维位置不能解释成个体 MRI 或精确皮层定位。"),
        ("02_optode_measurement_structure","光极测量结构与硬件共享",_whole_hardware,"方法：按导出的源探测配对绘制路径；共享硬件矩阵依据配对计算。解读：共用 S 或 D 是硬件结构，不是脑网络连接。完整配对与位置见 hb_pairs.csv。"),
        ("03_whole_head_hb_heatmaps","局部覆盖内的全头 Hb 热图",_whole_heatmaps,"方法：HbO、HbR 与 HbT=HbO+HbR 在两个固定时点显示；每色团全时间支持固定色标。局部线性插值同时约束头圈、凸包、三角边长和最近传感器距离。斜线保留无覆盖区域，缺失值以灰叉表示，不跨大缺口填色。"),
        ("04_whole_head_correlations","跨光极相关：地图与完整矩阵",_whole_correlations,"方法：读取显式导出的 Pearson 矩阵，统一 r 范围为 −1 到 1。地图只显示至多 20 条绝对值最大边；矩阵和保留数组包含全部关系。正实线、负虚线无方向，相关不能解释为因果。"),
        ("05_common_cause_and_hardware","控制变量前后的相关比较",_whole_partial,partial_explanation),
    ]
    rel=_relation_specs(m)
    if "conditional_lag_gain" in rel:
        spec = rel["conditional_lag_gain"]
        specs.append(("06_conditional_lag_prediction","条件历史提供方向预测",_whole_lag,f"方法：读取同样本嵌套 OLS 的 log(RSS_reduced/RSS_full)。来源历史阶数为 {spec.get('lag_order','导出声明的阶数')}；时间支持与控制变量以可见图注和导出说明为准。其他选定节点历史保留。行=来源，列=目标；预测增益非负，不表示兴奋／抑制，也不是干预因果证据。"))
    if {"known_direct","known_total_effect"}<=rel.keys():
        specs.append(("07_generator_causal_effects","生成器直接结构与总干预效应",_whole_causal,"方法：仅对已知合成线性生成器展示真值。直接系数描述单步结构；同初态同噪声单位脉冲后的 10 s 平均响应描述总效应，包含间接路径。该因果解释依赖已知生成机制，不能迁移为实测结论。"))
        if "do_response" in a and "corr_hbt" in rel:
            specs.append(("08_common_cause_counterexample","共因反例与已知路径传播",_whole_counterexample,"方法：预先设定 A/B 共同原因且无 A→B 路径，同时设置 A→C→D。地图标明各位置，波形显示相关，固定噪声干预响应核验传播。所有数字从保留数组读取或汇总；案例不是后验效果选择，无被试推断或置信区间。"))
    if {"crossmodal_corr_eeg_hbt_zero", "crossmodal_corr_eeg_hbt_lag2s"} <= rel.keys():
        specs.append(("09_crossmodal_whole_head", "EEG 与 Hb 的全头跨模态关系", _whole_crossmodal, f"方法：读取 EEG×HbT 在两个预先固定时差下的导出相关矩阵，时差及支持见可见图注。三角和圆点分别对应 EEG 与 Hb，全部节点保留；至多 20 条边只作显示。固定时差相关不建立传播或因果方向，完整 {len(n['eeg'])}×{len(n['hb'])} 矩阵保留全部配对。"))
    return specs


def _figure_text_audit(fig):
    from matplotlib.text import Text
    fig.canvas.draw()
    renderer, bounds = fig.canvas.get_renderer(), fig.bbox
    # Locators may retain a tick beyond the view interval (for example 1.4
    # when ylim ends at 1.22). Matplotlib does not draw those tick labels.
    # Exclude only those labels; ordinary off-canvas text still fails.
    undrawn_tick_labels = set()
    for ax in fig.axes:
        for axis in (ax.xaxis, ax.yaxis):
            low, high = sorted(axis.get_view_interval())
            for ticks, locations in ((axis.get_major_ticks(), axis.get_majorticklocs()),
                                     (axis.get_minor_ticks(), axis.get_minorticklocs())):
                for tick, location in zip(ticks, locations):
                    if location < low-1e-10 or location > high+1e-10:
                        undrawn_tick_labels.update((tick.label1, tick.label2))
    count, outside = 0, []
    for artist in fig.findobj(Text):
        if artist in undrawn_tick_labels or not artist.get_visible() or not artist.get_text():
            continue
        count += 1
        rect = artist.get_window_extent(renderer)
        if rect.x0 < -.5 or rect.y0 < -.5 or rect.x1 > bounds.width+.5 or rect.y1 > bounds.height+.5:
            outside.append(dict(text=artist.get_text(), bbox_fraction=[rect.x0/bounds.width,rect.y0/bounds.height,rect.x1/bounds.width,rect.y1/bounds.height]))
    if outside:
        raise RuntimeError(f"figure has text outside its image boundary: {outside}")
    return dict(visible_text_artists=count, text_artists_outside_figure=outside)


def _render_whole_head(manifest_path,output,m,a,*,fresh):
    import csv
    output=_fresh(output) if fresh else Path(output)
    (output/"figures").mkdir(exist_ok=False)
    plt=_plotting()
    figures=[]
    for name,title,function,explanation in _whole_specs(m,a):
        fig,meta=function(plt,m,a)
        path=Path("figures")/f"{name}.png"
        size=fig.get_size_inches().tolist()
        meta["layout_audit"] = _figure_text_audit(fig)
        fig.savefig(output/path,dpi=DPI)
        plt.close(fig)
        figures.append(dict(name=name,title=title,path=str(path),dpi=DPI,size_inches=size,pixels=[int(round(v*DPI)) for v in size],explanation=explanation,metadata=meta))
    with (output/"hb_pairs.csv").open("w",encoding="utf-8",newline="") as stream:
        writer=csv.writer(stream)
        writer.writerow(["hb","source","detector","x","y","coordinate_semantics"])
        for name,(s,d),xy in zip(m["names"]["hb"],m["geometry"]["hb_pairs"],m["geometry"]["hb_xy"]):
            writer.writerow([name,m["names"]["sources"][s],m["names"]["detectors"][d],*xy,"public paper normalized 2D; not individual MRI"])
    pages=_atlas(output,m,figures)
    audit=dict(schema=WHOLE_HEAD_SCHEMA,kind=m["kind"],view="whole_head",source_manifest=str(Path(manifest_path).resolve()),source_arrays_file=m["arrays_file"],output=str(output.resolve()),figure_count=len(figures),page_count=len(pages),figure_dpi=DPI,pdf_figures="embedded PNG bitmaps; body text searchable",fitting_performed=False,relation_estimation_performed=False,rendering_scope="read exported relations only; demo relation estimation occurs separately during bundle creation",spatial_interpolation="local linear within coverage only",uncertainty_estimated=False,coordinate_source=m["provenance"].get("geometry_source"),head=m["head"],figures=figures,pages=pages)
    (output/"audit.json").write_text(json.dumps(audit,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    return audit


def _whole_crossmodal(plt,m,a):
    keys=("crossmodal_corr_eeg_hbt_zero","crossmodal_corr_eeg_hbt_lag2s")
    specs = _relation_specs(m)
    offsets = [specs[key]["support"].get("target_offset_s", 0.) for key in keys]
    n_eeg, n_hb = len(m["names"]["eeg"]), len(m["names"]["hb"])
    footer = f"三角为 {n_eeg} EEG，圆点为 {n_hb} Hb；共用导出头圈。实线正相关／虚线负相关；无因果箭头，每图只显示 |r| 前 20 条。\n两列分别为 EEG(t) 对 HbT(t+{offsets[0]:g} s)、HbT(t+{offsets[1]:g} s)，时差来自显式导出，不逐对挑最大时延。固定时差相关不证明传播，保留全部 {n_eeg}×{n_hb} 配对。"
    fig,axes=_whole_canvas(plt,m,"09  EEG 与 Hb 的跨模态关系也保留全头位置和完整矩阵",footer,rows=2,cols=2,figsize=(15,11.4))
    eeg=np.asarray(m["geometry"]["eeg_xy"])
    hb=np.asarray(m["geometry"]["hb_xy"])
    records={}
    from matplotlib.colors import Normalize
    cmap=_signed_cmap()
    norm=Normalize(-1,1)
    for col,key in enumerate(keys):
        value=a[key]
        ax=axes[0,col]
        _head(ax,m)
        finite=np.argwhere(np.isfinite(value))
        pairs=sorted([tuple(map(int,pair)) for pair in finite],key=lambda pair:(-abs(float(value[pair])),pair))[:20]
        for i,j in pairs:
            v=float(value[i,j])
            ax.plot([eeg[i,0],hb[j,0]],[eeg[i,1],hb[j,1]],color=cmap(norm(v)),ls="-" if v>=0 else "--",lw=.6+2.4*abs(v),alpha=.8,zorder=2)
        ax.scatter(*eeg.T,marker="^",s=20,facecolor="white",edgecolor=BLUE,lw=.6,zorder=4,label="EEG")
        ax.scatter(*hb.T,marker="o",s=18,facecolor=ORANGE,edgecolor="white",lw=.4,zorder=5,label="Hb")
        for i in sorted({pair[0] for pair in pairs}):
            ax.annotate(m["names"]["eeg"][i],eeg[i],xytext=(0,5),textcoords="offset points",ha="center",fontsize=6.5,zorder=8)
        for j in sorted({pair[1] for pair in pairs}):
            ax.annotate(str(j+1),hb[j],xytext=(0,-8),textcoords="offset points",ha="center",fontsize=7,zorder=8)
        ax.set_title(f"EEG(t) 对 HbT(t+{offsets[col]:g} s)"+f"\n显示 {len(pairs)}/{value.size} 对；不是方向／因果证据",fontsize=11)
        ax.legend(frameon=False,ncol=2,loc="lower center",bbox_to_anchor=(.5,-.04))
        im=axes[1,col].imshow(value,cmap=cmap,vmin=-1,vmax=1,aspect="auto",interpolation="nearest")
        axes[1,col].set(xlabel="Hb 目标编号（列）",ylabel="EEG 来源通道（行）",title="完整 EEG × HbT 相关矩阵")
        ht=np.unique(np.r_[0,np.arange(9,n_hb-3,10),n_hb-1])
        et=np.unique(np.r_[np.arange(0,len(eeg),max(1,len(eeg)//12)),len(eeg)-1])
        axes[1,col].set_xticks(ht,np.asarray(ht)+1,fontsize=8)
        axes[1,col].set_yticks(et,[m["names"]["eeg"][i] for i in et],fontsize=7)
        records[key]=dict(displayed_edges=[list(pair) for pair in pairs],displayed_edge_count=len(pairs),valid_edge_count=int(np.isfinite(value).sum()),possible_edge_count=int(value.size),threshold_is_display_only=True,all_eeg_nodes=len(eeg),all_hb_nodes=len(hb),causal_direction=False)
    fig.colorbar(im,ax=list(axes[1]),fraction=.018,pad=.025,label="Pearson r；无量纲")
    return fig,dict(source_arrays=list(keys),relations=records,r_limits=[-1,1],fixed_time_offset_s=offsets,interpretation="fixed-offset crossmodal association; not causal propagation")


def _wrap_pdf(font, text, width, size):
    lines = []
    for paragraph in text.split("\n"):
        line = ""
        for char in paragraph:
            if line and font.text_length(line+char, fontsize=size) > width:
                lines.append(line)
                line = char
            else:
                line += char
        lines.append(line)
    return "\n".join(lines)


def _atlas(output, m, figures):
    import fitz
    doc = fitz.open()
    font = fitz.Font(fontfile=str(FONT_PATH))
    page_audit = []
    for index, entry in enumerate(figures, start=1):
        width_in, height_in = entry["size_inches"]
        image_width, image_height = width_in*72, height_in*72
        page_width = image_width+72
        text = _wrap_pdf(font, entry["explanation"], image_width, 11.5)
        caption_top = 72+image_height+25
        page_height = max(756, caption_top+18*len(text.split("\n"))+48)
        page = doc.new_page(width=page_width, height=page_height)
        writer = fitz.TextWriter(page.rect)
        writer.append((36, 31), entry["title"], font=font, fontsize=19)
        writer.append((36, 57), _badge(m), font=font, fontsize=12)
        rect = fitz.Rect(36, 72, 36+image_width, 72+image_height)
        page.insert_image(rect, filename=str(output/entry["path"]))
        for line_index, line in enumerate(text.split("\n")):
            writer.append((36, caption_top + 18*line_index), line, font=font, fontsize=11.5)
        writer.append((36, page_height-23), f"来源：单一 JSON / NPZ 显式导出；图 {index}/{len(figures)}。每条曲线均来自保留数组。", font=font, fontsize=10)
        writer.write_text(page, color=(.095, .196, .302))
        page_audit.append(dict(page=index, figure=entry["name"], image_objects=len(page.get_images(full=True)), searchable_text_characters=len(page.get_text())))
    path = output / "FIGURE_ATLAS.pdf"
    # Embed only used CJK glyphs; the full font collection dwarfs the figures.
    # fontTools handles this Noto collection reliably (the native CFF subsetter
    # used by MuPDF is not suitable for all members of the collection).
    doc.subset_fonts(fallback=True)
    doc.save(path, garbage=4, deflate=True)
    doc.close()
    with fitz.open(path) as checked:
        for page, record in zip(checked, page_audit):
            if len(page.get_images(full=True)) != 1 or not page.get_text().strip():
                raise RuntimeError("atlas must contain one embedded bitmap per page and searchable text")
            record["image_objects"] = len(page.get_images(full=True))
            images = []
            for info in page.get_image_info():
                rect = fitz.Rect(info["bbox"])
                dpi_x = float(info["width"] * 72 / rect.width)
                dpi_y = float(info["height"] * 72 / rect.height)
                if not 200 <= dpi_x <= 300 or not 200 <= dpi_y <= 300:
                    raise RuntimeError("embedded figure resolution is outside 200–300 dpi")
                images.append(dict(pixels=[info["width"], info["height"]], bbox=list(rect), dpi_x=dpi_x, dpi_y=dpi_y))
            record["images"] = images
            span_count, outside = 0, []
            for block in page.get_text("dict")["blocks"]:
                if block["type"] != 0:
                    continue
                for line in block["lines"]:
                    for span in line["spans"]:
                        span_count += 1
                        rect = fitz.Rect(span["bbox"])
                        if rect.x0 < -.5 or rect.y0 < -.5 or rect.x1 > page.rect.width+.5 or rect.y1 > page.rect.height+.5:
                            outside.append(dict(text=span["text"], bbox=list(rect)))
            record["text_span_count"] = span_count
            record["text_spans_outside_page"] = outside
            if outside:
                raise RuntimeError(f"text spans outside PDF page {record['page']}")
    return page_audit


def _render(manifest_path, output, *, fresh):
    m, a = load_bundle(manifest_path)
    if m["schema"] == WHOLE_HEAD_SCHEMA:
        return _render_whole_head(manifest_path, output, m, a, fresh=fresh)
    output = _fresh(output) if fresh else Path(output)
    figures_dir = output / "figures"
    figures_dir.mkdir(exist_ok=False)
    plt = _plotting()
    specs = [
        ("01_observation_structure", "观测位置与混合结构", _layout, "方法：显式 Hb/EEG 坐标、邻域边、光源—探测器配对与固定 W 分开绘制。解读：邻近与共享光极是两种观测关系；图中没有估计脑源或神经连接。"),
        ("02_spatial_waveforms", "空间波形与共同单位", _waveforms, "方法：每行一个 Hb 对，每列一种色团；同一色团跨通道固定纵轴，不逐通道 z-score。解读：比较同一单位下的幅度、峰值时间与缺失段；生成真值与观测均不能被解释成算法预测。"),
    ]
    if "hb_k1" in a:
        specs.append(("03_local_dynamics", "局部动态生成条件与已知驱动前向对照", _dynamics, "方法：分别求解非线性血流动力学单元后混合其 Hb 输出；双动态的第二驱动包含独立晚脉冲。解读：K_gen 是生成条件，不是拟合优劣。底行只置换固定 W 的 Hb 对应，保持已知驱动不变；不是实测置换检验。"))
    specs.extend([
        ("04_hb_coordinates", "总 Hb 与氧合方向的空间表达", _coordinates, "方法：逐时间点按 T=HbO+HbR、X=0.35HbO−0.65HbR 作线性坐标变换；快照时点在导出中预先声明，色条按完整时间支持固定。解读：圆点是观测位置，不作空间插值；T/X 不能直接解释成血流或血氧饱和度。"),
        ("05_visibility_and_pending", "输入、评分目标与未产生的预测", _visibility, "方法：显式读取成对 Hb 与 EEG 可见／评分掩码，二者互斥；缺失和无效预测保持空白。解读：灰虚线是隐藏目标，生成真值只是参考；没有导出预测时显示待产生，不生成数值性能或成功率。"),
    ])
    if "equivalence_total_a" in a:
        specs.append(("06_observational_equivalence", "相同总观测不保证唯一局部分配", _equivalence, "方法：重复两个驱动、参数和初态完全相同的合法非线性血流动力学单元，改变非负且行和为一的 W；对所有保留时间、通道和色团计算最大差。解读：总观测不变而分量改变，是一个构造的观测等价负例；不代表实测参数剖面、置信区间或任意非线性旋转。"))
    figures = []
    for name, title, function, explanation in specs:
        fig, meta = function(plt, m, a)
        path = Path("figures") / f"{name}.png"
        fig.savefig(output/path, dpi=DPI)
        plt.close(fig)
        figures.append(dict(name=name, title=title, path=str(path), dpi=DPI, size_inches=[13, 7.2], pixels=[3120, 1728], explanation=explanation, metadata=meta))
    pages = _atlas(output, m, figures)
    audit = dict(schema=SCHEMA, kind=m["kind"], source_manifest=str(Path(manifest_path).resolve()),
        source_arrays_file=m["arrays_file"], output=str(output.resolve()), figure_count=len(figures), page_count=len(pages),
        figure_dpi=DPI, pdf_figures="embedded PNG bitmaps; body text searchable", fitting_performed=False,
        spatial_interpolation=False, uncertainty_estimated=False, figures=figures, pages=pages)
    (output/"audit.json").write_text(json.dumps(audit, ensure_ascii=False, indent=2)+"\n", encoding="utf-8")
    return audit


def render_bundle(manifest_path, output):
    """Render an explicit validated export into a fresh output directory."""
    return _render(manifest_path, output, fresh=True)


def _experiment_evidence(directory, *, audit_arrays=True):
    """Read one explicit completed run; independently recompute its endpoints."""
    import hashlib
    from collections import Counter, defaultdict

    directory = Path(directory).resolve()
    read = lambda name: json.loads((directory/name).read_text())
    summary, plan = read("summary.json"), read("plan.json")
    if summary.get("execution") != "completed" or summary.get("terminal") != plan["planned_cells"]:
        raise ValueError("a complete planned denominator is required for a report")
    coordinate = read("coordinate.json")
    calibration = {p.stem: json.loads(p.read_text()) for p in (directory/"calibration").glob("*.json")}
    identities = {r["id"]: r for r in plan["rows"]}
    targets, cells, groups = {}, [], defaultdict(list)
    max_error, checked = 0., 0
    for spec in plan["cells"]:
        path = directory/"cells"/(spec["id"]+".json")
        row = json.loads(path.read_text())
        if any(row.get(k) != spec[k] for k in ("id", "record_id", "arm", "method", "mode", "pairing")):
            raise ValueError("cell identity differs from the immutable plan")
        if row["status"] == "missing":
            raise ValueError("missing terminal cell")
        if row["status"] == "completed" and (not row["success"] or not np.isfinite(row["nmse"])):
            raise ValueError("invalid completed score")
        if audit_arrays and row["status"] == "completed":
            identity = row["record_id"]
            if identity not in targets:
                key = hashlib.sha256(identity.encode()).hexdigest()[:24]
                with np.load(directory/"prepared"/(key+".npz"), allow_pickle=False) as z:
                    targets[identity] = z["hb"][:, 0].copy()/coordinate["hb_scale"]
            with np.load(path.with_suffix(".npz"), allow_pickle=False) as z:
                if "target_prediction" in z:
                    pred = z["target_prediction"]
                else:
                    arm = calibration[row["arm"]]["arm"]
                    offset = arm["eeg_channels"]*5
                    pred = z["prediction"][:, offset:offset+2]
                score = slice(48, 72) if row["mode"] == "center_pair" else slice(20, 120)
                sd = np.asarray(coordinate["sd"])[60:62]
                value = float(np.mean(((pred[score]-targets[identity][score])/sd)**2))
            error = abs(value-row["nmse"])
            max_error = max(max_error, error)
            if not np.isclose(value, row["nmse"], rtol=1e-10, atol=1e-11):
                raise ValueError(f"independent scoring mismatch: {row['id']}")
            checked += 1
        compact = {k: row.get(k) for k in (
            "id", "record_id", "subject", "arm", "mode", "pairing", "method", "scenario",
            "status", "success", "nmse", "nrmse", "component_nmse", "correlation",
            "amplitude_rms_ratio", "eeg_weight", "nfev", "integration_max_training_sd",
            "known_truth_diagnostics", "error", "convergence_reason")}
        cells.append(compact)
        key = "|".join(row[k] for k in ("scenario", "arm", "mode", "pairing", "method"))
        groups[key].append(row)
    counts = dict(Counter(r["status"] for r in cells))
    if counts != summary["status_counts"]:
        raise ValueError("status counts differ from the retained summary")
    group_error, contrast_error, interval_error = 0., 0., 0.
    for key, rows in groups.items():
        subjects = defaultdict(list)
        for row in rows:
            if row["status"] not in ("unavailable_wrong_EEG", "unavailable_preparation"):
                subjects[row["subject"]].append(row["nmse"] if row["status"] == "completed" else 100.)
        actual = {s: float(np.mean(v)) for s, v in subjects.items()}
        declared = summary["groups"][key]["subject_nmse_with_failure_penalty"]
        if set(actual) != set(declared):
            raise ValueError("subject aggregation identities differ")
        for s in actual:
            group_error = max(group_error, abs(actual[s]-declared[s]))
        if group_error > 1e-10:
            raise ValueError("independent subject aggregation mismatch")
    for contrast in summary["comparisons"]:
        a = {r["record_id"]: r for r in groups[contrast["candidate"]]}
        b = {r["record_id"]: r for r in groups[contrast["reference"]]}
        paired = defaultdict(lambda: ([], []))
        for identity in sorted(set(a) & set(b)):
            left, right = a[identity], b[identity]
            if any(r["status"] in ("unavailable_wrong_EEG", "unavailable_preparation") for r in (left, right)):
                continue
            x, y = paired[left["subject"]]
            x.append(left["nmse"] if left["status"] == "completed" else 100.)
            y.append(right["nmse"] if right["status"] == "completed" else 100.)
        differences = np.asarray([np.sqrt(np.mean(paired[s][1]))-np.sqrt(np.mean(paired[s][0]))
                                  for s in sorted(paired)])
        if len(differences):
            contrast_error = max(contrast_error, abs(float(differences.mean())-contrast["mean_nrmse_improvement"]))
            draw = np.random.default_rng(20261016).integers(0, len(differences), (1000, len(differences)))
            interval = np.quantile(differences[draw].mean(axis=1), [.025, .975])
            interval_error = max(interval_error, float(np.max(abs(interval-contrast["interval_95"]))))
            if contrast_error > 1e-10 or interval_error > 1e-10:
                raise ValueError("independent paired-bootstrap mismatch")
    checks = dict(planned=plan["planned_cells"], terminal=len(cells), status_counts=counts,
        rescored_arrays=checked, max_nmse_error=max_error, max_subject_nmse_error=group_error,
        max_paired_gain_error=contrast_error, max_interval_error=interval_error,
        independent_scoring="retained arrays; frozen train SD; per-window NMSE then within-subject mean and square root",
        bootstrap="1000 subject resamples, seed 20261016; common available canonical windows; failure NMSE 100")
    return dict(root=directory, summary=summary, plan=plan, coordinate=coordinate,
        calibration=calibration, cells=cells, identities=identities, checks=checks)


def _report_number(value, digits=4):
    return "不可估计" if value is None else f"{value:.{digits}f}"


def _report_group(evidence, arm="K1_true_H4_E12", mode="heldout_pair", pairing="paired", method="SSM"):
    scenario = evidence["cells"][0]["scenario"]
    return evidence["summary"]["groups"]["|".join((scenario, arm, mode, pairing, method))]


def _report_contrast(evidence, comparison, arm="K1_true_H4_E12", mode="heldout_pair", pairing="paired"):
    matches = [c for c in evidence["summary"]["comparisons"]
               if c["comparison"] == comparison and c["candidate"].split("|")[1:4] == [arm, mode, pairing]]
    if len(matches) != 1:
        raise ValueError(f"contrast must be unique: {comparison}, {arm}, {mode}, {pairing}")
    return matches[0]


class _SpatialReport:
    """Local PPT and searchable PDF export; every figure/formula is a PNG."""
    def __init__(self, output):
        import fitz
        from PIL import ImageFont
        from pptx import Presentation
        from pptx.util import Inches
        self.output = _fresh(output)
        for name in ("figures", "formulas", "tables", "preview"):
            (self.output/name).mkdir()
        self.deck = Presentation()
        self.deck.slide_width, self.deck.slide_height = Inches(13.333), Inches(7.5)
        self.font = ImageFont.truetype(str(FONT_PATH), 80)
        self.pdf_font = fitz.Font(fontfile=str(FONT_PATH))
        self.pages, self.figures, self.formulas, self.layout = [], [], [], []

    def text(self, slide, content, x, y, w, h, size=20, bold=False, color="18324D"):
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
        from PIL import ImageFont
        font = ImageFont.truetype(str(FONT_PATH), round(size*4))
        lines = 0
        for line in str(content).split("\n"):
            used = ""; lines += 1
            for char in line:
                if used and font.getlength(used+char)/4 > w*72:
                    lines += 1; used = char
                else:
                    used += char
        if lines*size*1.18 > h*72:
            raise ValueError(f"slide {len(self.deck.slides)} text exceeds box: {content[:80]}")
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = box.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        for i, line in enumerate(str(content).split("\n")):
            p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            p.text, p.line_spacing = line, 1.1
            p.space_after = p.space_before = Pt(0)
            for run in p.runs:
                run.font.name = "Noto Sans CJK SC"
                run.font.size, run.font.bold = Pt(size), bold
                run.font.color.rgb = RGBColor.from_string(color)
        self.layout.append(dict(page=len(self.deck.slides), text=content, x=x, y=y, w=w, h=h, size=size))

    def page(self, title, subtitle, source):
        slide = self.deck.slides.add_slide(self.deck.slide_layouts[6])
        self.text(slide, title, .55, .30, 12.2, .60, 26, True)
        self.text(slide, subtitle, .58, 1.02, 12.1, .48, 14, color="526879")
        self.text(slide, f"局部空间实验  |  Single-Trial  |  2026-10-20  |  {len(self.deck.slides):02}",
                  .58, 7.13, 12.1, .22, 10, color="526879")
        slide.notes_slide.notes_text_frame.text = "Owning evidence: "+str(source)
        self.pages.append(dict(page=len(self.deck.slides), title=title, subtitle=subtitle, source=str(source), paragraphs=[], images=[]))
        return slide

    def body(self, title, subtitle, paragraphs, source, *, size=21):
        slide = self.page(title, subtitle, source)
        step = 5.05/len(paragraphs)
        for i, paragraph in enumerate(paragraphs):
            self.text(slide, paragraph, .8, 1.72+i*step, 11.7, step-.18, size)
        self.pages[-1]["paragraphs"] = list(paragraphs)
        return slide

    def picture(self, slide, path, x, y, w, h):
        from PIL import Image
        from pptx.util import Inches
        frame_w, frame_h = w, h
        with Image.open(path) as im:
            ratio = im.width/im.height
            if Path(path).parent.name == "formulas":
                w, h = min(w, im.width/300), min(h, im.height/300)
        width, height = min(w, h*ratio), min(h, w/ratio)
        slide.shapes.add_picture(str(path), Inches(x+(frame_w-width)/2), Inches(y+(frame_h-height)/2),
                                width=Inches(width), height=Inches(height))
        self.pages[-1]["images"].append(str(Path(path).relative_to(self.output)))

    def figure(self, title, subtitle, fig, caption, source, key):
        import matplotlib.pyplot as plt
        path = self.output/"figures"/(key+".png")
        _figure_text_audit(fig)
        fig.savefig(path, dpi=240, facecolor="white")
        plt.close(fig)
        self.figures.append(path)
        slide = self.page(title, subtitle, source)
        self.picture(slide, path, .62, 1.62, 12.05, 3.95)
        self.text(slide, caption, .8, 5.79, 11.7, 1.06, 17)
        self.pages[-1]["paragraphs"] = [caption]
        return slide

    def equation(self, title, latex, key, explanation, source):
        import fitz, subprocess
        directory = self.output/"formulas"
        tex = directory/(key+".tex")
        tex.write_text("\\documentclass[12pt]{article}\n\\usepackage{amsmath,amssymb}\n"
            "\\usepackage[paperwidth=14in,paperheight=5in,margin=.2in]{geometry}\n"
            "\\pagestyle{empty}\\begin{document}\\fontsize{23}{29}\\selectfont\\[\n"
            +latex+"\n\\]\\end{document}\n")
        result = subprocess.run(["pdflatex", "-interaction=nonstopmode", "-halt-on-error", tex.name],
                                cwd=directory, capture_output=True, text=True)
        if result.returncode:
            raise ValueError(result.stdout[-2000:])
        with fitz.open(tex.with_suffix(".pdf")) as doc:
            boxes = [fitz.Rect(b[:4]) for b in doc[0].get_text("blocks")]
            clip = boxes[0]
            for box in boxes[1:]:
                clip |= box
            clip += (-8, -8, 8, 8)
            path = tex.with_suffix(".png")
            doc[0].get_pixmap(clip=clip, dpi=300).save(path)
        self.formulas.append(path)
        slide = self.page(title, "LaTeX 公式位图；符号、单位和时间支持在本页解释。", source)
        self.picture(slide, path, .75, 1.62, 11.8, 2.25)
        self.text(slide, explanation, .8, 4.10, 11.7, 2.65, 19)
        self.pages[-1]["paragraphs"] = [explanation]

    def finish(self, checks, appendix):
        import fitz, html, subprocess, zipfile
        from PIL import Image

        ppt = self.output/"LOCAL_SPATIAL_REPORT.pptx"
        self.deck.save(ppt)
        (self.output/"slide_sources.json").write_text(json.dumps(self.pages, ensure_ascii=False, indent=2)+"\n")
        (self.output/"text_layout.json").write_text(json.dumps(self.layout, ensure_ascii=False, indent=2)+"\n")
        profile = (self.output/"libreoffice_profile").resolve().as_uri()
        conversion = subprocess.run(["libreoffice", "-env:UserInstallation="+profile, "--headless",
            "--convert-to", "pdf", "--outdir", str(self.output), str(ppt)],
            capture_output=True, text=True, timeout=180)
        (self.output/"pdf_export.log").write_text(conversion.stdout+"\n"+conversion.stderr)
        if conversion.returncode or not ppt.with_suffix(".pdf").exists():
            raise ValueError("PPT PDF export failed")

        document = fitz.open()
        html_parts = ['<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>局部空间实验技术报告</title>',
            '<style>body{max-width:1100px;margin:40px auto;font:18px/1.65 sans-serif;color:#18324d}'
            'img{max-width:100%;height:auto}section{margin:3em 0}small{overflow-wrap:anywhere}'
            'table{border-collapse:collapse}td,th{padding:7px;border:1px solid #bbc}</style>',
            '<h1>局部多通道与空间动态维数：技术报告</h1>',
            '<p>2026-10-20；Single-Trial 公开开发面板。本文为保留运行的通信导出，运行摘要与逐单元记录是事实来源。</p>']

        def put(page, text, x, y, width, size=12.5, color=(.095, .196, .302)):
            lines = _wrap_pdf(self.pdf_font, str(text), width, size).split("\n")
            writer = fitz.TextWriter(page.rect)
            for i, line in enumerate(lines):
                writer.append((x, y+i*size*1.4), line, font=self.pdf_font, fontsize=size)
            writer.write_text(page, color=color)
            return y+len(lines)*size*1.4

        for entry in self.pages:
            page = document.new_page(width=842, height=595)
            put(page, entry["title"], 38, 37, 766, 20)
            y = put(page, entry["subtitle"], 38, 65, 766, 11.5)+12
            html_parts.append("<section><h2>"+html.escape(entry["title"])+"</h2><p>"+html.escape(entry["subtitle"])+"</p>")
            for relative in entry["images"]:
                path = self.output/relative
                with Image.open(path) as im:
                    width, height = im.size
                scale = min(758/width, 266/height)
                if path.parent.name == "formulas":
                    scale = min(scale, 72/300)
                w, h = width*scale, height*scale
                page.insert_image(fitz.Rect(42+(758-w)/2, y, 42+(758+w)/2, y+h), filename=str(path))
                y += h+20
                html_parts.append('<img src="'+html.escape(relative)+'" alt="'+html.escape(entry["title"])+'">')
            for paragraph in entry["paragraphs"]:
                required = len(_wrap_pdf(self.pdf_font, paragraph, 758, 13).split("\n"))*18.2+14
                if y+required > 520:
                    page = document.new_page(width=842, height=595)
                    y = put(page, entry["title"]+"（续）", 38, 37, 766, 18)+18
                y = put(page, paragraph, 42, y, 758, 13)+14
                html_parts.append("<p>"+html.escape(paragraph).replace("\n", "<br>")+"</p>")
            put(page, "来源："+entry["source"], 38, 556, 766, 7.3)
            html_parts.append("<small>来源："+html.escape(entry["source"])+"</small></section>")
        for title, paragraphs in appendix:
            page = document.new_page(width=842, height=595)
            y = put(page, title, 38, 37, 766, 20)+20
            html_parts.append("<section><h2>"+html.escape(title)+"</h2>")
            for paragraph in paragraphs:
                required = len(_wrap_pdf(self.pdf_font, paragraph, 758, 12.5).split("\n"))*17.5+14
                if y+required > 550:
                    page = document.new_page(width=842, height=595)
                    y = put(page, title+"（续）", 38, 37, 766, 18)+18
                y = put(page, paragraph, 42, y, 758, 12.5)+14
                html_parts.append("<p>"+html.escape(paragraph).replace("\n", "<br>")+"</p>")
            html_parts.append("</section>")
        document.subset_fonts(fallback=True)
        document.save(self.output/"TECHNICAL_REPORT.pdf", garbage=4, deflate=True)
        document.close()
        html_parts.append("</html>")
        (self.output/"TECHNICAL_REPORT.html").write_text("\n".join(html_parts))

        atlas = fitz.open()
        for path in self.figures+self.formulas:
            with Image.open(path) as im:
                width, height = im.size
            dpi = 300 if path.parent.name == "formulas" else 240
            page = atlas.new_page(width=width*72/dpi+36, height=height*72/dpi+36)
            page.insert_image(fitz.Rect(18, 18, page.rect.width-18, page.rect.height-18), filename=str(path))
        atlas.save(self.output/"FIGURE_ATLAS.pdf", deflate=True)
        atlas.close()
        exports = {}
        for filename in ("LOCAL_SPATIAL_REPORT.pdf", "TECHNICAL_REPORT.pdf", "FIGURE_ATLAS.pdf"):
            with fitz.open(self.output/filename) as doc:
                outside = []
                for i, page in enumerate(doc):
                    for block in page.get_text("dict")["blocks"]:
                        if block["type"] != 0:
                            continue
                        for line in block["lines"]:
                            for span in line["spans"]:
                                box = fitz.Rect(span["bbox"])
                                if box.x0 < -1 or box.y0 < -1 or box.x1 > page.rect.width+1 or box.y1 > page.rect.height+1:
                                    outside.append(dict(page=i+1, text=span["text"], bbox=list(box)))
                exports[filename] = dict(pages=len(doc), image_objects=len({x[0] for page in doc for x in page.get_images()}),
                    displayed_images=sum(len(page.get_image_info()) for page in doc),
                    selectable_characters=sum(len(page.get_text()) for page in doc), text_outside_page=outside)
                if outside:
                    raise ValueError(f"PDF text clipping: {filename}: {outside[:2]}")
                if filename != "FIGURE_ATLAS.pdf" and exports[filename]["selectable_characters"] < 1000:
                    raise ValueError("report body text must remain searchable")
                if filename == "LOCAL_SPATIAL_REPORT.pdf":
                    if len(doc) != len(self.deck.slides) or exports[filename]["image_objects"] < len(self.figures)+len(self.formulas):
                        raise ValueError("PPT page/image count mismatch")
                    for i, page in enumerate(doc):
                        page.get_pixmap(dpi=100, alpha=False).save(self.output/"preview"/f"page_{i+1:02}.png")
        audit = dict(schema="shared_driver_spatial_report_v1", output=str(self.output.resolve()),
            kind="completed_public_experiment", page_count=len(self.deck.slides), figure_count=len(self.figures),
            formula_count=len(self.formulas), bitmap_figures=True, figure_dpi=240, formula_dpi=300,
            independent_evidence_checks=checks, exports=exports)
        (self.output/"delivery_checks.json").write_text(json.dumps(audit, ensure_ascii=False, indent=2)+"\n")
        (self.output/"render_shared_driver_spatial.py").write_bytes(Path(__file__).read_bytes())
        return audit


def _report_csv(output, name, records):
    import csv
    if not records:
        return
    keys = list(dict.fromkeys(k for row in records for k in row))
    with (Path(output)/"tables"/(name+".csv")).open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=keys)
        writer.writeheader()
        writer.writerows({k: json.dumps(v, ensure_ascii=False) if isinstance(v, (list, dict)) else v
                          for k, v in row.items()} for row in records)


def _report_forest(contrasts, labels):
    plt = _plotting()
    fig, ax = plt.subplots(figsize=(12.2, 3.9))
    for i, (row, label) in enumerate(zip(contrasts, labels)):
        value, interval = row["mean_nrmse_improvement"], row["interval_95"]
        if value is None:
            ax.text(0, i, "不可估计", va="center")
            continue
        ax.errorbar(value, i, xerr=[[value-interval[0]], [interval[1]-value]], color=BLUE,
                    marker="o", markersize=6, capsize=4, linewidth=2)
        ax.annotate(f"{value:+.4f}  [{interval[0]:+.4f}, {interval[1]:+.4f}]",
                    (value, i), xytext=(8, 8), textcoords="offset points", fontsize=10)
    ax.set_yticks(range(len(labels)), labels)
    ax.invert_yaxis()
    ax.axvline(0, color=GRAY, linestyle="--")
    ax.set_xlabel("归一化误差收益（正值有利于候选）")
    ax.grid(axis="x", alpha=.18)
    ax.margins(x=.35, y=.25)
    fig.subplots_adjust(left=.26, right=.96, bottom=.20, top=.92)
    return fig


def _report_examples(evidence):
    rows = [r for r in evidence["cells"] if r["arm"] == "K1_true_H4_E12" and r["mode"] == "heldout_pair"
            and r["pairing"] == "paired" and r["method"] == "SSM"]
    rows.sort(key=lambda r: (r["nrmse"] if r["status"] == "completed" else 10., r["record_id"]))
    selected, used = [], set()
    for label, rank in (("较好", 0), ("居中", (len(rows)-1)/2), ("困难", len(rows)-1)):
        choices = sorted(enumerate(rows), key=lambda pair: (abs(pair[0]-rank), pair[1]["record_id"]))
        index, row = next((i, r) for i, r in choices if r["subject"] not in used)
        used.add(row["subject"])
        selected.append(dict(label=label, rank=index+1, denominator=len(rows), **row))
    return selected


def _report_waveforms(evidence, example, mode="heldout_pair"):
    import hashlib
    plt = _plotting()
    identity = example["record_id"]
    key = hashlib.sha256(identity.encode()).hexdigest()[:24]
    with np.load(evidence["root"]/"prepared"/(key+".npz"), allow_pickle=False) as z:
        target = z["hb"][:, 0].copy()/evidence["coordinate"]["hb_scale"]
    sd = np.asarray(evidence["coordinate"]["sd"])[60:62]
    definitions = [("K1_true_H4_E12", "SSM", "K1", BLUE, "--"),
        ("K2_true_H4_E12", "SSM", "K2", ORANGE, "-."),
        ("K1_true_H4_E12", "joint_ridge", "联合 ridge", "#278C79", ":")]
    if mode == "center_pair":
        definitions.append(("K1_true_H4_E12", "own_context_interpolation", "自身插值", GRAY, "--"))
    fig, axes = plt.subplots(1, 2, figsize=(12.2, 3.9), sharex=True)
    time = np.arange(120)*.25-5
    for i, ax in enumerate(axes):
        ax.plot(time, target[:, i]/sd[i], color=INK, linewidth=2, label="处理后实测（评分目标）")
        ax.axvline(0, color=GRAY, linestyle=":", linewidth=1)
        if mode == "center_pair":
            ax.axvspan(7, 13, color="#ECE3D3", alpha=.55)
        ax.set(title=("HbO", "HbR")[i], xlabel="事件相对时间（秒）", ylabel="幅值 / 本色团训练 SD")
        ax.grid(alpha=.15)
    values = []
    for arm, method, label, color, style in definitions:
        row = next(r for r in evidence["cells"] if r["record_id"] == identity and r["arm"] == arm
                   and r["mode"] == mode and r["pairing"] == "paired" and r["method"] == method)
        values.append(f"{label} {_report_number(row['nrmse'])}" if row["status"] == "completed" else f"{label} 失败")
        if row["status"] != "completed":
            continue
        with np.load(evidence["root"]/"cells"/(row["id"]+".npz"), allow_pickle=False) as z:
            prediction = z["target_prediction"] if "target_prediction" in z else z["prediction"][:, 60:62]
            if method == "own_context_interpolation":
                prediction = prediction.copy()
                prediction[:48] = np.nan
                prediction[72:] = np.nan
            for i, ax in enumerate(axes):
                ax.plot(time, prediction[:, i]/sd[i], color=color, linestyle=style, linewidth=1.7, label=label)
    axes[0].legend(loc="best", fontsize=10)
    for ax in axes:
        ax.set_xlim(-5, 25)
    fig.subplots_adjust(left=.07, right=.98, bottom=.18, top=.89, wspace=.26)
    return fig, "；".join(values)


def _report_waveform_explanation(evidence, example, mode="heldout_pair"):
    import hashlib
    identity = example["record_id"]
    key = hashlib.sha256(identity.encode()).hexdigest()[:24]
    with np.load(evidence["root"]/"prepared"/(key+".npz"), allow_pickle=False) as z:
        target = z["hb"][:, 0].copy()/evidence["coordinate"]["hb_scale"]
    sd = np.asarray(evidence["coordinate"]["sd"])[60:62]
    indices = np.arange(48,72) if mode == "center_pair" else np.arange(20,120)
    paragraphs = []
    signed = lambda value: f"{value:+.2f}".replace("-", "−")
    for component, name in enumerate(("HbO", "HbR")):
        truth = target[indices,component]/sd[component]
        peak = indices[int(np.argmax(truth))]*.25-5
        parts = [f"{name} 在评分区间由 {signed(truth[0])} 变为 {signed(truth[-1])} 个训练 SD；最大值时刻为事件后 {peak:.2f} 秒。"]
        for arm,method,label in [("K1_true_H4_E12","SSM","K1"),("K2_true_H4_E12","SSM","K2"),
                                 ("K1_true_H4_E12","joint_ridge","联合 ridge")]:
            row = next(r for r in evidence["cells"] if r["record_id"]==identity and r["arm"]==arm
                       and r["method"]==method and r["mode"]==mode and r["pairing"]=="paired")
            if row["status"] != "completed":
                parts.append(label+" 数值失败，未绘制其预测。")
            else:
                with np.load(evidence["root"]/"cells"/(row["id"]+".npz"), allow_pickle=False) as z:
                    prediction = z["target_prediction"] if "target_prediction" in z else z["prediction"][:,60:62]
                    p = prediction[indices,component]/sd[component]
                bias = float(np.mean(p-truth))
                delay = indices[int(np.argmax(p))]*.25-5-peak
                parts.append(f"{label}：相关 {_report_number(row['correlation'][component],2)}，幅比 {_report_number(row['amplitude_rms_ratio'][component],2)}，均差 {signed(bias)} SD，峰差 {signed(delay)} 秒。")
        paragraphs.append("".join(parts))
    paragraphs.append("幅比是预测 RMS 除以目标 RMS，1 表示幅度相同；均差是预测减实测的平均差，正值为整体高估。峰差是预测最大值时刻减实测最大值时刻，正值表示更晚，不是生理延迟估计。样例与色团纵轴独立定标以显示细节；同一子图内所有方法共享刻度。")
    return paragraphs


def render_experiment_report(measured_run, synthetic_run, output):
    """Export the frozen spatial experiment without fitting or raw-data access."""
    from collections import Counter
    measured, synthetic_root = Path(measured_run).resolve(), Path(synthetic_run).resolve()
    if measured == synthetic_root:
        raise ValueError("measured and synthetic evidence must have separate identities")
    config = json.loads((measured/"resolved_config.json").read_text())
    if (config["schema"] != "shared_driver_spatial_v1"
            or config["data"]["datasets"] != ["eeg_fnirs_single_trial"]
            or config["data"]["protected_data"] != "forbidden"
            or json.loads((measured/"run_identity.json").read_text())["family"] != "public"):
        raise ValueError("report requires the registered Single-Trial spatial experiment")
    if json.loads((synthetic_root/"run_identity.json").read_text())["family"] != "synthetic":
        raise ValueError("synthetic controls must retain their generator identity")
    parent = json.loads((synthetic_root/"summary.json").read_text())
    if parent.get("execution") != "completed" or parent.get("calibration_oracle") is not False:
        raise ValueError("complete learned-calibration synthetic controls are required")
    main = _experiment_evidence(measured)
    synthetic = {name: _experiment_evidence(synthetic_root/"scenarios"/name)
                 for name in config["synthetic"]["scenarios"]}
    report = _SpatialReport(output)
    plt = _plotting()
    source = str(measured/"summary.json")
    protocol = "experiments/configs/physiology_semantic_tokenizer/shared_driver_spatial_v1.yaml"
    model_source = "src/inference/shared_driver_spatial.py; src/inference/shared_driver_rk4.py"
    value = lambda arm="K1_true_H4_E12", mode="heldout_pair", method="SSM": _report_group(
        main, arm=arm, mode=mode, method=method)["mean_subject_nrmse_with_failure_penalty"]
    k1, k2 = value(), value("K2_true_H4_E12")
    joint, spatial, template = (value(method=x) for x in ("joint_ridge", "spatial_ridge", "training_template"))
    # This dated narrative describes this evidence generation. A different run
    # needs a fresh interpretation, rather than silently reusing its prose.
    expected = [1.1716097413841338, 2.1478277094932774, .9970527304437269]
    if (not np.allclose([k1,k2,joint], expected, rtol=0, atol=1e-10)
            or main["checks"]["status_counts"] != {
                "completed":11042, "failed_numerical":67, "unavailable_wrong_EEG":735}):
        raise ValueError("the dated 2026-10-20 narrative requires its original evidence generation")
    _report_csv(report.output, "measured_groups", [dict(group=k, **v) for k, v in main["summary"]["groups"].items()])
    _report_csv(report.output, "measured_comparisons", main["summary"]["comparisons"])
    _report_csv(report.output, "measured_cells", [{k:v for k,v in r.items() if k != "known_truth_diagnostics"} for r in main["cells"]])
    for name, evidence in synthetic.items():
        _report_csv(report.output, "synthetic_"+name, [dict(group=k, **v) for k,v in evidence["summary"]["groups"].items()])
    _report_csv(report.output, "calibration_selection", [dict(arm=arm, selected_weight=c["selected_eeg_weight"],
        training_planned=c["training_subset"]["planned"], training_succeeded=c["training_subset"]["succeeded"], **v)
        for arm,c in main["calibration"].items() for v in c["candidates"]])

    report.body("实测隐藏预测仍由线性空间基线领先", "EEG：脑电；HbO/HbR：氧合/去氧血红蛋白；ridge：带正则的线性岭回归。", [
        f"主指标是整对隐藏 Hb 的归一化均方根误差（NRMSE），按训练标准差（SD）评分。联合 ridge {joint:.4f}、空间 ridge {spatial:.4f}、单核心 K1 {k1:.4f}、双核心 K2 {k2:.4f}；越低越好。",
        "29 人按 17／6／6 人划分训练、选择和评价；评价含 36 个窗口。统计独立单元是被试，通道、时间点和重复窗口都不增加独立人数。",
        "K2 的主结果含预定失败惩罚；34/36 窗成功。K1 的五个空间／数量候选均选 EEG 权重 0，不能把它们的表现解释为 EEG 的增量。",
        "完整合成支持匹配双源条件下的恢复能力；实测没有确认正确几何、EEG 配对或两套独立轨迹的稳定额外收益。"], source)
    report.body("三个问题使用各自的对照回答", "先区分观测数量、空间对应与动态容量，再解释结果。", [
        "观测数量：同一目标及同一 K1 动态下，比较四对与两对 Hb、十二路与六路 EEG；子面板由最大面板切片，训练尺度不重新拟合。",
        "空间对应：同样信号、相同参数数目与选择预算，比较真实、置换和 K1 均匀几何；置换改变通道与几何权重的对应。",
        "动态容量：守恒状态空间模型简称 SSM。K2 含两套独立驱动及初态；另用同一读出、固定 EEG 权重 1 的轨迹相等限制模型，区分读出家族和动态自由度。",
        "跨模态增量：正确 EEG 与 Hb-only、任务和条件匹配的错配 EEG 比较，并与获得相同可见输入的 ridge 比较。"], protocol)
    report.body("评价只覆盖六名公开开发被试", "受试者划分先于信号读取；每人两个记录、每记录三个不重叠窗口。", [
        "训练：17 人、102 窗；选择：6 人、36 窗；评价：6 人、36 窗。划分种子固定为 20261016，没有根据评价误差调整窗口或传感器。",
        "评价包含 18 个运动想象窗（左手 8、右手 10）与 18 个心算／基线窗（心算 4、基线 14）。本组结果不是任务分类性能。",
        "36 个评价窗中，29 个具有同被试、同任务／条件、同传感器布局（montage）的非重叠错配 EEG；其余 7 个保留缺失，不用不匹配对象补齐。",
        "这些被试已有公开开发历史；换分组没有形成新的未查看确认集。本版没有读取其他数据集、保护分区或个体解剖信息。"], str(measured/"plan.json"))

    patch = main["plan"]["rows"][0]["patch"]
    fig, axes = plt.subplots(1, 2, figsize=(12.2, 3.9), gridspec_kw={"width_ratios":[1.25,1]})
    e, h = np.asarray(patch["eeg_coordinates"]), np.asarray(patch["hb_coordinates"])
    axes[0].scatter(e[:,0], e[:,1], s=34, color=BLUE, marker="o", label="EEG")
    axes[0].scatter(h[1:,0], h[1:,1], s=80, color="#278C79", marker="s", label="Hb 供体")
    axes[0].scatter(h[0,0], h[0,1], s=130, color=ORANGE, marker="*", label="隐藏 Hb 目标")
    for i, name in enumerate(patch["eeg_channel_names"]):
        axes[0].annotate(name, e[i,:2], xytext=(4,5), textcoords="offset points", fontsize=10)
    positions = [(-.77,1.12),(-.77,1.03),(.43,1.14),(.49,1.02)]
    for i, name in enumerate(patch["hb_pair_names"]):
        axes[0].annotate(f"H{i}: {name}", h[i,:2], xytext=positions[i], textcoords="data",
                         fontsize=10, arrowprops=dict(arrowstyle="-",color=GRAY,linewidth=.7))
    axes[0].set(xlabel="模板 x（无单位）", ylabel="模板 y（无单位）", title="三维模板坐标的 x-y 投影")
    axes[0].set_ylim(.10,1.18)
    axes[0].set_xlim(-.98,.99)
    axes[0].legend(loc="lower right", fontsize=9)
    weights = np.asarray(main["calibration"]["K2_true_H4_E12"]["design"]["hb_loadings"])
    im = axes[1].imshow(weights, vmin=0, vmax=1, cmap="Blues", aspect="auto")
    axes[1].set_yticks(range(4), ["H0 目标", "H1 供体", "H2 供体", "H3 供体"])
    axes[1].set_xticks([0,1], ["局部模式 1", "局部模式 2"])
    axes[1].set_title("冻结 Hb 混合权重；每行和为 1")
    for i in range(4):
        for j in range(2):
            axes[1].text(j,i,f"{weights[i,j]:.3f}",ha="center",va="center",color="white" if weights[i,j]>.6 else INK)
    fig.colorbar(im,ax=axes[1],label="混合权重（无单位）")
    fig.subplots_adjust(left=.06,right=.96,bottom=.2,top=.86,wspace=.37)
    report.figure("目标不共享光极，供体仍有硬件依赖", "每个点是传感器位置；模板几何不等于个体皮层定位。", fig,
        "目标为 FpzFp1。目标与三个供体不共享已核验光极；H1/H2 共享探测器，H2/H3 共享光源。四对 Hb 不代表四个独立组织来源。右图给出实际冻结的 K2 空间混合权重。", str(measured/"plan.json"), "geometry")
    report.body("测量、变换和评分坐标明确分开", "原生 EEG 为微伏；Single-Trial fNIRS 输入是双波长光强。", [
        "EEG 使用原生 200 Hz、双侧乳突参考的公开信号，保留眼动等伪迹。五频带为 1–4、4–8、8–13、13–30、30–45 Hz，构造 4 Hz 的参考相对对数功率（log-power）。",
        "Hb 从原生 760/850 nm 光强构造光密度，再用修正 Beer–Lambert 定律（MBLL）转换为相对 HbO/HbR；原生样本按互不交叠的 bin 聚合，不跨隐藏 Hb 区间滤波。",
        "每窗 30 秒，从事件前 5 秒至事件后 25 秒；前 5 秒定义参考，输出为 120 个位置。模型对自身轨迹施加同一参考算子。",
        "EEG 每频带跨通道共用训练均方根（RMS），Hb 所有通道／色团共用一个训练 RMS。评分另用逐特征训练 SD；这些相对坐标不是绝对浓度标定或噪声似然。"], "src/data/shared_driver_spatial.py; src/inference/observation_baselines.py")
    report.body("隐藏整对 Hb 时，连参考段也只用于评分", "所有任务为离线条件重建或补全，不是实时预测。", [
        "整对隐藏：目标 HbO/HbR 的整个 30 秒窗口都不进入拟合，包括前 5 秒参考；目标参考只定义评分坐标。其他 Hb 供体和完整 EEG 可见。",
        "中心隐藏：仅隐藏事件后 7–13 秒的目标 Hb，对该 6 秒区间评分；目标自身的其他时间可见。完整重建：目标可见，评价事件后 0–25 秒。",
        "整对隐藏和完整重建都在事件后 0–25 秒评分。隐藏先于输入组装和训练坐标变换；目标均值、幅度或参考不得用于校正预测。",
        "训练模板和 ridge 始终排除目标对；因此只有整对隐藏时才与 SSM 同输入。中心任务另设目标自身上下文插值；完整可见时该插值复制目标，零误差没有预测意义。"], protocol)
    report.equation("K1 与 K2 只改变局部动态维数",
        r"r_k(t)=\sum_{m=1}^{16}\Phi_m(t)c_{mk},\qquad "
        r"\dot{\mathbf x}_k=F(\mathbf x_k,r_k;\theta),\qquad "
        r"\mathbf x_k=(s_k,f_k,v_k,p_k,q_k)^{\mathsf T}",
        "state_dimensions",
        "每个模式使用 16 个预定余弦时间系数，并拟合 5 个初态；K1 每窗 21 个未知量，K2 每窗 42 个。驱动 r 是条件模型输入，不能当作已测得神经活动。\n"
        "s 为血流调节信号，f 为相对流量，v 为相对血容量，p 为归一化总 Hb 状态，q 为归一化去氧 Hb 状态。血管参数共享且冻结；没有每通道自由生理参数、额外 Hb 时间残差或网络连接。", model_source)
    report.equation("每个核心独立传播同一守恒动力学",
        r"\begin{aligned}"
        r"\dot s&=r-\kappa s-\gamma(f-1),&\dot f&=s,\\"
        r"\tau\dot v&=f-v^{1/\alpha},&\tau\dot p&=f-v^{1/\alpha}p/v,\\"
        r"\tau\dot q&=fE(f)/E_0-v^{1/\alpha}q/v,&E(f)&=1-(1-E_0)^{1/f}."
        r"\end{aligned}", "balloon",
        "冻结取值：通过时间常数为 2 秒，阻尼系数为 0.64，反馈系数为 0.32，流出指数参数为 0.32，静息提取率为 0.32；有效神经血管增益为 1。这些是条件模型假设，不是逐被试生理估计。\n"
        "流量、血容量及 Hb 状态保持正值，并检查绝对 HbO/HbR 的物理关系。RK4 在每个 0.25 秒间隔内积分四步；用八步重放检查输出误差。总 Hb 状态与血容量没有被预先强制相等。", model_source)
    report.equation("Hb 在独立传播之后混合，EEG 使用统计读出",
        r"\begin{aligned}"
        r"h(\mathbf x_k)&=\begin{bmatrix}P_0(p_k-1)-Q_0(q_k-1)\\Q_0(q_k-1)\end{bmatrix},\\"
        r"\widehat{\mathbf H}_j&=\mathcal P_H\!\left[\sum_k W_{jk}h(\mathbf x_k)\right],\qquad "
        r"\widehat E_{cb}=\mathcal P_E\!\left[\sum_k G_{ck}a_{bk}r_k\right]."
        r"\end{aligned}", "observation",
        "Hb 两行依次为 HbO 和 HbR；静息总 Hb 常数为 1，去氧 Hb 常数为 0.35。W 为非负、逐行和为 1 的冻结几何矩阵；K1 的每个 Hb 权重都是 1，多个读出不能独立产生不同峰形。\n"
        "G 是固定高斯空间核，每列按欧氏范数归一；核宽由正的传感器到锚点距离中位数确定。a 是每模式五个训练拟合的带符号频带系数。它们映射 log-power，不能称为 EEG 电压 leadfield 或脑源定位。", "src/inference/observation_baselines.py; "+model_source)
    report.body("训练对象、选择对象和当前窗未知量不同", "训练与选择完成后，评价过程不再重新校准读出或尺度。", [
        "训练被试：拟合共同坐标与逐特征 SD；用所有训练 Hb 做条件拟合，得到驱动参考，再回归五个带符号频带系数／模式。该驱动参考不是神经真值。",
        "选择被试：分别评价 EEG 权重 0、0.1、1，以隐藏 Hb 的被试等权误差选择；数值失败使用预定惩罚，平局优先较小权重。",
        "评价被试：冻结训练尺度、空间矩阵、频带系数、基线和选定权重；当前窗口只拟合驱动系数及核心初态。训练目标通道可用于训练，这不是未见传感器泛化。",
        "空间 ridge 输入其他 Hb；联合 ridge 输入其他 Hb 加 EEG；训练模板是训练窗平均目标。ridge 使用相同训练分区与正则，独立于 SSM 拟合成功与否。"], str(measured/"calibration"))
    report.equation("数据项按可见数量归一，正则保持固定",
        r"J=\lambda_E\frac{T}{n_E}\sum_{\mathcal V_E}\left(\frac{\widehat E-E}{\sigma_E}\right)^2"
        r"+\frac{T}{n_H}\sum_{\mathcal V_H}\left(\frac{\widehat H-H}{\sigma_H}\right)^2"
        r"+J_{\rm curve}+J_{\rm init}+J_{\rm flow}",
        "objective",
        "T 为 120；n 是该模态实际可见标量数量，σ 为冻结训练 SD。EEG 权重为零时，该模态在变换与初始化之前就被忽略。通道增加不会自动提高某模态的总数据权重。\n"
        "曲率系数 0.01、物理初态系数 100、对数流量先验系数 1；流量对数标准差为 0.693147。最多 180 次目标求值，梯度阈值为十万分之一。本目标是工程准则，训练 SD 不是传感器噪声标准差。", protocol)
    report.equation("先计算窗口误差，再按被试等权汇总",
        r"M_w=\frac{1}{2|I|}\sum_{c\in\{O,R\}}\sum_{t\in I}"
        r"\left(\frac{\widehat H_{wtc}-H_{wtc}}{\sigma_c}\right)^2,\qquad "
        r"N_s=\sqrt{\frac{1}{|W_s|}\sum_{w\in W_s}M_w},\qquad "
        r"\overline N=\frac1S\sum_s N_s",
        "metric",
        "I 是本任务预定评分时间，W 是同一被试的可用窗口，S 是独立被试数。HbO/HbR 在窗口内等权；先平均归一化均方误差，再开方，再平均被试，不能改成全部时间点直接池化。\n"
        "数值失败的窗口均方误差固定为 100。不可用错配供体保留在计划分母，但不进入共同可用配对收益；成功子集另外展示，不能替代含失败惩罚的主结果。", "experiments/scripts/evaluate_shared_driver_spatial.py:score_prediction/subject_endpoint")
    report.equation("收益为参考减候选，区间按人重采样",
        r"\Delta=\frac1S\sum_s\left(N_{s,\rm reference}-N_{s,\rm candidate}\right),\qquad "
        r"CI_{95}=\operatorname{quantile}_{.025,.975}\!\left(\Delta^{*(1)},\ldots,\Delta^{*(1000)}\right)",
        "paired",
        "正收益表示候选误差更低；单位是训练 SD 归一化误差的绝对差，不是百分比。先取双方共同可用的同身份窗口，再进行被试内聚合。\n"
        "评价仅有 6 名被试，自助重采样（bootstrap）每次抽取整人，训练与选择冻结。区间为描述性逐项 95% 区间，未做多重比较校正；跨零表示方向不确定，排除零也不等于新的确认性证据。", "experiments/scripts/evaluate_shared_driver_spatial.py:compare_cells/paired_bootstrap")
    scenarios = config["synthetic"]["scenarios"]
    scenario_labels = ["单源", "双源", "相同双源", "无信息 EEG", "几何失配", "共同 Hb 干扰"]
    fig, ax = plt.subplots(figsize=(12.2, 3.9))
    x = np.arange(len(scenarios))
    for offset, arm, method, label, color in [
        (-.25,"K1_true_H4_E12","SSM","K1",BLUE),
        (0,"K2_true_H4_E12","SSM","K2",ORANGE),
        (.25,"K1_true_H4_E12","joint_ridge","联合 ridge","#278C79")]:
        vals = [_report_group(synthetic[name],arm=arm,method=method)["mean_subject_nrmse_with_failure_penalty"] for name in scenarios]
        bars = ax.bar(x+offset,vals,.24,label=label,color=color)
        for bar,val in zip(bars,vals):
            ax.text(bar.get_x()+bar.get_width()/2,val+.015,f"{val:.3f}",ha="center",fontsize=10)
    ax.set_xticks(x,scenario_labels)
    ax.set(ylabel="隐藏 Hb 的平均 NRMSE（训练 SD）",ylim=(0,1.22))
    ax.legend(ncol=3,loc="upper left");ax.grid(axis="y",alpha=.15)
    fig.subplots_adjust(left=.07,right=.98,bottom=.18,top=.94)
    report.figure("合成双源可以恢复，失配仍会破坏预测", "每场景独立训练／选择／评价；评价为 4 个生成身份、每身份 2 条记录。", fig,
        "各场景 2,632 个评价单元均成功。双源 K2 为 0.0201，优于 K1 的 0.8127 和联合 ridge 的 0.2048；几何失配时 K2 为 1.0084。相同双源表示时间秩塌缩，EEG 读出家族并不等于 K1。", str(synthetic_root/"summary.json"), "synthetic_primary")
    report.body("无信息 EEG 和共同干扰限制机制解释", "已知生成器状态是合成真值；实测没有对应潜变量真值。", [
        "无信息 EEG 场景使用独立生成的同谱／同空间结构 EEG。Hb 供体仍含双源信息，因此 K2 可以在主要依赖 Hb 的条件下恢复隐藏目标，不能据此声称 EEG 有效。",
        "共同 Hb 干扰为向 HbO/HbR 同向添加的已知成分；K2 误差 0.0519，K1 为 0.7104。更低的总波形误差可以伴随对干扰的吸收。",
        "几何失配条件下，生成器使用冻结置换的空间映射，真实几何候选因假设不匹配而恶化。该反例检验模型脆弱性，不是一个新的生理机制。",
        "完整合成采用训练 Hb 条件驱动来校准 EEG，已经不使用上一轮冒烟检查的已知读出；两轮配置、时长与随机样本不同，数值不直接作前后改善比较。"], str(synthetic_root/"summary.json"))
    recovery = []
    for name in scenarios:
        arm = "K1_true_H4_E12" if name == "single_source" else "K2_true_H4_E12"
        rows = [r for r in synthetic[name]["cells"] if r["arm"]==arm and r["mode"]=="heldout_pair"
                and r["pairing"]=="paired" and r["method"]=="SSM" and r["status"]=="completed"]
        subject = {}
        for row in rows:
            truth = row["known_truth_diagnostics"]
            values = [float(np.mean([v["reference_centered_nrmse"] for v in truth["drivers"]]))]
            values += [float(np.mean([v["coordinates"][state]["reference_centered_nrmse"] for v in truth["states"]]))
                       for state in ("f","q")]
            subject.setdefault(row["subject"],[]).append(values)
        medians = np.median([np.mean(v,axis=0) for v in subject.values()],axis=0)
        recovery.append(dict(scenario=name,driver=float(medians[0]),flow=float(medians[1]),deoxy_state=float(medians[2]),
            subjects=len(subject),records=len(rows),pooling="mean modes within record, mean two records within identity, median four identities",
            denominator="own reference-centered generator trajectory RMS"))
    _report_csv(report.output,"synthetic_latent_recovery",recovery)
    fig, ax = plt.subplots(figsize=(12.2,3.9))
    matrix=np.asarray([[r[k] for k in ("driver","flow","deoxy_state")] for r in recovery])
    im=ax.imshow(matrix.T,aspect="auto",cmap="YlOrRd",vmin=0,vmax=max(1.,float(matrix.max())))
    ax.set_xticks(range(6),scenario_labels);ax.set_yticks(range(3),["驱动","相对流量状态","去氧 Hb 状态"])
    for i in range(6):
        for j in range(3):ax.text(i,j,f"{matrix[i,j]:.3f}",ha="center",va="center",fontsize=11)
    fig.colorbar(im,ax=ax,label="参考中心化误差 / 真值轨迹 RMS")
    fig.subplots_adjust(left=.16,right=.95,bottom=.17,top=.93)
    report.figure("隐藏波形误差与潜变量恢复分别判断","每格为四个生成身份的中位数；颜色是归一化误差，不是相关系数。",fig,
        "先平均记录内模式，再平均身份内两记录，最后取四身份中位数。无信息 EEG 和共同干扰下，驱动误差约 0.45，隐藏 Hb 误差仍可很低。此处分母是真值轨迹的参考中心化 RMS，区别于 Hb 评分的训练 SD。",str(synthetic_root/"scenarios"),"latent_recovery")
    arm_labels={"K1_true_H4_E12":"K1 真 · 4Hb/12EEG","K1_sham_H4_E12":"K1 置换 · 4Hb/12EEG",
        "K1_uniform_H4_E12":"K1 均匀 · 4Hb/12EEG","K2_true_H4_E12":"K2 真 · 4Hb/12EEG",
        "K2_sham_H4_E12":"K2 置换 · 4Hb/12EEG","K1_true_H2_E12":"K1 真 · 2Hb/12EEG",
        "K1_true_H4_E6":"K1 真 · 4Hb/6EEG"}
    arms=sorted(main["calibration"],key=lambda a:list(arm_labels).index(a))
    selection=np.array([[c["subject_mean_nmse"] for c in main["calibration"][a]["candidates"]] for a in arms])
    fig,ax=plt.subplots(figsize=(12.2,3.9))
    im=ax.imshow(selection,aspect="auto",cmap="YlOrRd",vmin=0,vmax=4)
    ax.set_yticks(range(7),[arm_labels[a] for a in arms])
    ax.set_xticks(range(3),["EEG 权重 0","EEG 权重 0.1","EEG 权重 1"])
    for i,a in enumerate(arms):
        c=main["calibration"][a]
        for j,row in enumerate(c["candidates"]):
            star=" ★" if row["eeg_weight"]==c["selected_eeg_weight"] else ""
            ax.text(j,i,f"{row['subject_mean_nmse']:.4f}；成功 {row['succeeded']}/36{star}",
                    ha="center",va="center",fontsize=11,color="white" if row["subject_mean_nmse"]>2.5 else INK)
    fig.colorbar(im,ax=ax,label="选择集平均 NMSE（含失败惩罚）")
    fig.subplots_adjust(left=.23,right=.95,bottom=.15,top=.94)
    report.figure("五个 K1 候选均选择忽略 EEG","星号标出冻结选择；选择使用平均 NMSE，主报告使用被试平均 NRMSE。",fig,
        "K2 真几何选择权重 1，置换几何选择 0.1。真几何的权重 0 和 0.1 各有一次选择窗失败，惩罚影响了选择；不能将选中较大权重解释为更强生理耦合。",str(measured/"calibration"),"selection")
    report.body("成功、数值失败与缺失对照分别保留","完整计划为 11,844 单元；一个单元是窗口×模型×隐藏方式×配对×方法。",[
        "评价共有 11,042 个成功单元、67 个数值失败单元、735 个错配 EEG 不可用单元。735 个不可用单元来自 7 个缺少匹配供体的窗口及其全部相关比较。",
        "67 个数值失败均达到 180 次目标求值预算，不能作为缺失删除。K2 置换模型占 40 个，K2 真几何占 21 个，K1 两 Hb 面板占 6 个；它们不是独立被试计数。",
        "训练条件拟合：每个 K1 候选为 102/102 成功；K2 真几何为 98/102，置换为 96/102。EEG 读出使用成功训练子集，原始观测基线仍使用全部训练窗口。",
        "没有追加评价调参、重试困难窗口或扩大优化预算。正式运行使用 24 进程、单线程数值库和独立 user systemd 服务；失败与资源先导记录均保留。"],source)


    model_defs=[("训练模板","K1_true_H4_E12","training_template",GRAY),
        ("空间 ridge","K1_true_H4_E12","spatial_ridge","#775DA6"),
        ("联合 ridge","K1_true_H4_E12","joint_ridge","#278C79"),
        ("K1","K1_true_H4_E12","SSM",BLUE),("K2","K2_true_H4_E12","SSM",ORANGE)]
    fig,ax=plt.subplots(figsize=(12.2,3.9))
    highest=0.
    for i,(label,arm,method,color) in enumerate(model_defs):
        g=_report_group(main,arm=arm,method=method)
        val=g["mean_subject_nrmse_with_failure_penalty"]
        ax.bar(i,val,color=color,width=.58,label=label)
        ax.text(i,val+.045,f"{val:.4f}\n成功 {g['succeeded']}/36",ha="center",fontsize=10)
        subject=np.asarray(list(g["subject_nrmse_with_failure_penalty"].values()))
        highest=max(highest,float(subject.max()))
        ax.scatter(i+np.linspace(-.17,.17,len(subject)),subject,color=INK,s=20,alpha=.8,marker="o")
    ax.set_xticks(range(5),[d[0] for d in model_defs])
    ax.set(ylabel="被试平均 NRMSE（含失败惩罚）",ylim=(0,max(3.,highest*1.16)))
    ax.grid(axis="y",alpha=.15);fig.subplots_adjust(left=.08,right=.98,bottom=.15,top=.95)
    report.figure("线性空间基线优于两个守恒候选","柱为六名被试等权均值；黑点各代表一名被试，不是窗口。",fig,
        f"联合 ridge {joint:.4f}、空间 ridge {spatial:.4f}、K1 {k1:.4f}、K2 {k2:.4f}。所有方法预测同一隐藏目标；K2 的两个失败窗按冻结 NMSE 惩罚进入主结果。",source,"measured_primary")
    cs=[_report_contrast(main,"model_vs_joint_ridge"),
        _report_contrast(main,"model_vs_joint_ridge","K2_true_H4_E12"),
        _report_contrast(main,"model_vs_training_template"),
        _report_contrast(main,"K2_vs_K1","K2_true_H4_E12")]
    report.figure("K1 超过模板，但没有超过强线性基线","共同可用 36 窗、6 人；误差收益及描述性 95% 被试 bootstrap 区间。",_report_forest(cs,
        ["K1 相对联合 ridge","K2 相对联合 ridge","K1 相对训练模板","K2 相对 K1"]),
        "K1 相对联合 ridge 的收益为 −0.1746；K2 相对 K1 为 −0.9762。负值有利于参考方法。区间未做多重比较校正，结论限定于冻结面板和算法预算。",source,"baseline_contrasts")
    cs=[_report_contrast(main,"paired_vs_Hb_only","K2_true_H4_E12"),
        _report_contrast(main,"paired_vs_wrong_EEG","K2_true_H4_E12"),
        _report_contrast(main,"paired_vs_Hb_only"),
        _report_contrast(main,"paired_vs_wrong_EEG")]
    report.figure("没有确认正确 EEG 配对的稳定增量","Hb-only 对照有 36 共同可用窗；错配 EEG 对照有 29 共同可用窗；均为 6 人。",_report_forest(cs,
        ["K2：正确相对 Hb-only","K2：正确相对错配 EEG","K1：正确相对 Hb-only","K1：正确相对错配 EEG"]),
        "K2 相对 Hb-only 收益 0.0647，区间 [−1.4459, 1.4178]；相对错配 EEG 为 0.5023，区间 [−0.0005, 1.4851]，均跨零。K1 权重为 0，因此配对对照在共同窗口上严格相同。",source,"eeg_contrasts")
    cs=[_report_contrast(main,"true_vs_sham_K2","K2_true_H4_E12"),
        _report_contrast(main,"true_vs_sham_K1"),
        _report_contrast(main,"true_vs_uniform_K1"),
        _report_contrast(main,"multi_vs_small_Hb"),
        _report_contrast(main,"multi_vs_small_EEG")]
    report.figure("几何与通道数量的额外收益仍不确定","每项固定目标、时间支持和训练尺度；零收益需要结合选定策略解释。",_report_forest(cs,
        ["K2：真几何相对置换","K1：真几何相对置换","K1：真几何相对均匀","四对 Hb 相对两对","十二 EEG 相对六 EEG"]),
        "K2 真实几何收益 0.4534，但区间跨零。增加 Hb 的收益 0.0729，区间 [−0.0194, 0.1830]。K1 的几何及 EEG 数量对照为零，是忽略 EEG 后的相同模型，不能泛化为几何信息无价值。",source,"geometry_count")
    cs=[_report_contrast(main,"K2_free_vs_tied_fixed1","K2_true_H4_E12"),
        _report_contrast(main,"K2_vs_K1","K2_true_H4_E12")]
    report.figure("同一读出下，独立轨迹也没有稳定收益","固定 EEG 权重 1；两核自由模型与轨迹及初态相等的限制模型配对。",_report_forest(cs,
        ["同读出：自由 K2 相对相等限制","不同读出家族：K2 相对 K1"]),
        "严格限制模型保留 K2 的加载两列之和，并将曲率、初态和流量先验各乘 2。解除轨迹与初态相等约束的收益为 −0.9769，区间 [−2.3024, 0.0005]；不支持为当前预测任务增加动态复杂度。",source,"tied")
    fig,axes=plt.subplots(1,2,figsize=(12.2,3.9))
    modes=["full","heldout_pair","center_pair"];mode_names=["完整可见","整对隐藏","中心隐藏"]
    for i,arm in enumerate(["K1_true_H4_E12","K2_true_H4_E12"]):
        gs=[_report_group(main,arm=arm,mode=m) for m in modes]
        axes[0].plot(range(3),[g["mean_subject_nrmse_with_failure_penalty"] for g in gs],marker=("o","s")[i],color=(BLUE,ORANGE)[i],label=("K1","K2")[i])
        axes[1].plot(range(3),[g["success_subset_mean_subject_nrmse"] for g in gs],marker=("o","s")[i],color=(BLUE,ORANGE)[i],label=("K1","K2")[i])
    for ax,title in zip(axes,["主结果：含数值失败惩罚","成功子集：仅作补充诊断"]):
        ax.set_xticks(range(3),mode_names);ax.set(title=title,ylabel="被试平均 NRMSE",ylim=(0,2.5));ax.grid(alpha=.15);ax.legend()
    fig.subplots_adjust(left=.07,right=.97,bottom=.18,top=.87,wspace=.25)
    report.figure("更低的成功子集拟合误差不能掩盖失败","中心任务只评 6 秒；其他任务评事件后 25 秒，跨任务数值不能解释为同一难度。",fig,
        "K2 完整可见成功子集误差 0.6920，含失败后为 1.2893；整对隐藏成功子集为 1.2734，含失败后为 2.1478。两种汇总的窗口集合不同，成功子集不是新的配对优势证据。",source,"failure_sensitivity")
    center_defs=[("自身上下文插值","K1_true_H4_E12","own_context_interpolation",GRAY),
        ("联合 ridge","K1_true_H4_E12","joint_ridge","#278C79"),("K1","K1_true_H4_E12","SSM",BLUE),
        ("K2","K2_true_H4_E12","SSM",ORANGE)]
    fig,ax=plt.subplots(figsize=(12.2,3.9))
    for i,(name,arm,method,color) in enumerate(center_defs):
        g=_report_group(main,arm=arm,mode="center_pair",method=method)
        v=g["mean_subject_nrmse_with_failure_penalty"];ax.bar(i,v,color=color,width=.6)
        ax.text(i,v+.03,f"{v:.4f}\n成功 {g['succeeded']}/36",ha="center")
    ax.set_xticks(range(4),[x[0] for x in center_defs]);ax.set(ylabel="中心隐藏 NRMSE（含失败惩罚）",ylim=(0,1.65))
    ax.grid(axis="y",alpha=.15);fig.subplots_adjust(left=.07,right=.98,bottom=.15,top=.94)
    report.figure("中心补全仍以自身上下文插值最强","目标事件后 7–13 秒隐藏；目标区间外的时间可见。",fig,
        "自身上下文插值为 0.4910，联合 ridge 为 0.7686，K1 为 0.8570，K2 为 1.2575。插值只使用可见目标上下文；ridge 全程排除目标，因此不能把中心任务当作所有方法同输入的主比较。",source,"center_prediction")
    fig,axes=plt.subplots(1,3,figsize=(12.2,3.9),sharex=True,sharey=True)
    for ax,(label,arm,method,color) in zip(axes,[model_defs[3],model_defs[4],model_defs[2]]):
        rows=[r for r in main["cells"] if r["arm"]==arm and r["method"]==method and r["mode"]=="heldout_pair"
              and r["pairing"]=="paired" and r["status"]=="completed"]
        xy=np.asarray([r["correlation"] for r in rows],float)
        ax.scatter(xy[:,0],xy[:,1],color=color,alpha=.65,s=25)
        ax.axhline(0,color=GRAY,linestyle="--");ax.axvline(0,color=GRAY,linestyle="--")
        ax.set(xlim=(-1.05,1.05),ylim=(-1.05,1.05),xlabel="HbO：预测与目标的 Pearson 相关",
               title=f"{label}；成功 {len(rows)}/36")
    axes[0].set_ylabel("HbR：预测与目标的 Pearson 相关")
    fig.subplots_adjust(left=.07,right=.98,bottom=.2,top=.86,wspace=.12)
    report.figure("误差、波形相关与成分归属不能互相替代","每点是一个成功窗口，同一人的多个点相互依赖；两条零线只表示相关符号。",fig,
        "相关在各色团的评分时间内单独计算，不受正比例幅值变化影响；较高相关仍可有幅度偏差。K2 两个失败窗不出现在本图，但保留在主风险中。相关也不能给潜在成分指定生理来源。",source,"correlations")
    examples=_report_examples(main)
    _report_csv(report.output,"selected_examples",[{k:v for k,v in r.items() if k!="known_truth_diagnostics"} for r in examples])
    for index,example in enumerate(examples):
        identity=main["identities"][example["record_id"]]
        fig,metrics=_report_waveforms(main,example)
        caption=f"按 K1 主任务误差选择第 {example['rank']}/{example['denominator']} 窗，三例来自不同被试。目标全程隐藏，事件时刻为 0；各色团用固定训练 SD 表示幅值。窗口 NRMSE：{metrics}。"
        report.figure(f"{example['label']}样例：同输入预测的波形差异",
            f"{identity['subject']} / {identity['record']} / 事件 {identity['event_id']} / {identity['task']} / {identity['condition']}",
            fig,caption,example["record_id"],f"example_{index+1}")
        report.body(f"{example['label']}样例：峰形、幅度与评分的对应",
                    "与前页使用同一窗口和同一评分区间；幅度比与相关分别解释。",
                    _report_waveform_explanation(main, example), example["record_id"], size=18)
    fig,metrics=_report_waveforms(main,examples[1],mode="center_pair")
    report.figure("同一居中样例：局部隐藏与整对隐藏不同","保持样例不变，但模型按中心隐藏任务独立重新拟合；阴影为评分区间。",fig,
        f"阴影为事件后 7–13 秒，其余目标可见；灰色插值只显示隐藏段。NRMSE：{metrics}。各任务独立重拟合，跨任务差异不能解释为同输入条件下的预测收益。",
        examples[1]["record_id"],"center_example")
    report.body("结构与优化都是可能的限制","以下是与实现相容的解释，尚未通过新增干预实验区分。",[
        "K1 的四个 Hb 读出共享同一对波形，只增加联合约束，无法独立表示不同位置的峰形。困难样例中，HbR 的均值和幅度偏差明显；这与结构约束相容，但不能据此确定真实成因。",
        "K2 放开第二套轨迹后，在匹配合成中有效，在几何失配中失效。实测失败又集中于冻结求值预算；模型失配与优化不足都可能限制当前结果。",
        "EEG log-power 是带伪迹的统计特征，训练 Hb 条件驱动也不是神经真值。没有额外独立来源锚点，不能将拟合加载符号、空间模式或峰时差赋予唯一生理含义。",
        "当前最稳妥的用途是保留空间 ridge 作为强预测基线、保留 K1 作为受约束的条件参考。来源解释、个体参数估计及下游 tokenizer 价值都需要各自独立证据。"],protocol)
    report.body("结论：先保留可预测信息，暂不增加机制主张","结果对应固定前额面板、固定生理参数、当前求解预算和公开开发人群。",[
        "空间 Hb 的直接线性信息是当前最有用的预测基线。K1 超过训练模板，但未超过同输入联合 ridge；K2 的复杂度和数值失败没有换来稳定隐藏预测收益。",
        "K1 自动选择忽略 EEG；K2 的正确 EEG 相对 Hb-only 和错配对照的区间跨零。当前结果不建立额外的跨模态耦合资格。",
        "真实几何与置换几何未形成稳定差异；两套独立轨迹与初态也未超过同读出的相等限制模型。不能据此命名两个皮层源或推断连接方向。",
        "后续如继续 K2，应先在训练／合成支持上解决预算内收敛和失配问题，再定义新的冻结比较；本轮没有根据评价结果重调模型，也没有自动启动下一轮实验。"],source)
    report.body("交付与独立核验","PPT 与 PDF 正文可选择；图表和公式均嵌入 PNG，保留 LaTeX 源码。",[
        "独立复算全部 11,042 个成功实测单元和 15,792 个合成单元的误差；被试聚合、共同窗口配对与 1,000 次 bootstrap 区间均与保留摘要一致。",
        "完整表格位于 tables：逐评价单元、各组汇总、所有配对比较、权重选择、合成潜变量恢复及样例选择。逐页来源在 slide_sources.json，原运行保持唯一事实源。",
        "数值失败及错配缺失均保留；资源先导的解释器启动错误单独保留，正式实测使用新的 v2 目录及仓库环境，没有混用失败先导的准备文件。",
        "配套技术报告提供本套页面的完整文字、方法补充和读取入口；delivery_checks.json 记录数值复算、图像对象、可选文本及页面边界检查。"],str(report.output/"delivery_checks.json"))
    checks={"measured":main["checks"],"synthetic":{name:e["checks"] for name,e in synthetic.items()}}
    appendix=[
        ("技术补充：输入与复现",[
            "冻结配置："+str(measured/"resolved_config.json")+"；拟合源码快照："+str(measured/"source_snapshot")+
            "。运行入口 evaluate_shared_driver_spatial.py 依次执行 plan、prepare、calibrate、measured、summarize；完整合成另用 synthetic 阶段。",
            "PPT 中的训练 SD 是经过共同训练 RMS 缩放后的逐特征标准差。Hb 的共同正尺度保留通道和色团幅比；各样例图以对应色团的同一训练 SD 显示，未逐窗或逐帧归一。",
            "空间核为模板坐标下的高斯距离核，默认宽度为正距离中位数；EEG 列按欧氏范数归一，Hb 行按和归一。K2 锚点固定为目标与最远的已选供体。sham 固定目标 Hb 行，只置换其他 Hb 行及 EEG 行。",
            "signed EEG 系数的 ridge 为 0.1 倍训练 Gram 平均对角线，不新增逐通道自由系数、截距或评价期再定标。基础 ridge 也仅在训练数据拟合；不会继承 SSM 的数值失败。",
            "错配 EEG 的供体只提供 EEG，不读取其 Hb 作为目标输入。隐藏目标及参考对输入不变、零 EEG 权重忽略、K2 同轨迹等价、解析导数和积分精度均有 fixture 测试。"]),
        ("技术补充：统计与解释限制",[
            "主输出是被试内平均窗口 NMSE 后开方，再平均被试。选择阶段使用被试等权 NMSE，以便按预定失败惩罚选择；它与报告 NRMSE 的开方顺序不同，不是同一个统计量。",
            "区间固定训练和选择，仅反映六名评价被试在当前面板的重采样不确定性；不是覆盖训练不确定性的总体置信声明。全部比较均为开发性描述，没有对报告中的多项区间进行家族错误率校正。",
            "评价任务分布不平衡：心算 4 窗、基线 14 窗、左手想象 8 窗、右手想象 10 窗。没有把这些条件当作独立被试，也没有基于条件差异另行调整参数。",
            "数值失败均达到冻结目标求值预算。当前负结果同时反映模型、优化和预算，不证明所有双源模型都无效；增加预算或改变参数属于后续版本实验。",
            "合成驱动和状态的真值来自生成器；实测目标是真实测量经声明变换后的隐藏 Hb，实测驱动和状态只是条件估计。模板几何、未共享光极及预测增益均不能排除组织重叠、浅层或系统性干扰。"]),
        ("文件与复核入口",[
            "主实测："+str(measured)+"；完整合成："+str(synthetic_root)+"。",
            "原始数据说明：data/EEG+NIRS Single-Trial/Open access dataset for simultaneous EEG and NIRS Brain-Computer Interfaces (BCIs).html。数据单位、时钟和来源合同：docs/DATA_CONTRACT.md。",
            "报告图表的底层数值保存在 tables/*.csv；PPT/PDF 的每页证据路径见 slide_sources.json。误差为参考减候选时正值有利于候选；所有成功率与缺失均有原计划分母。",
            "报告不是新的状态源；项目状态继续由 research_state/registry.json 及其生成视图持有。保留旧冒烟、资源先导和正式运行，不将通信导出替代原始运行证据。"])]
    return report.finish(checks,appendix)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    source = parser.add_mutually_exclusive_group(required=True)
    source.add_argument("--demo", action="store_true", help="generate a toy forward-only bundle; no measured data")
    source.add_argument("--whole-head-demo", action="store_true", help="public 64-EEG/51-Hb geometry with independent synthetic relations")
    source.add_argument("--bundle", type=Path, help="one explicit visualization manifest.json")
    source.add_argument("--experiment-report", action="store_true", help="2026-10-20 PPT and technical PDF from completed spatial experiment runs")
    parser.add_argument("--measured-run", type=Path)
    parser.add_argument("--synthetic-run", type=Path)
    parser.add_argument("--output", type=Path, required=True, help="fresh output directory; existing paths are refused")
    parser.add_argument("--seed", type=int, default=20261016, help="demo Gaussian-noise seed")
    args = parser.parse_args(argv)
    if args.experiment_report:
        if args.measured_run is None or args.synthetic_run is None:
            parser.error("--experiment-report requires --measured-run and --synthetic-run")
        audit = render_experiment_report(args.measured_run, args.synthetic_run, args.output)
    elif args.demo or args.whole_head_demo:
        output = _fresh(args.output)
        creator = create_whole_head_demo_bundle if args.whole_head_demo else create_demo_bundle
        manifest_path = creator(output/"bundle", seed=args.seed)
        audit = _render(manifest_path, output, fresh=False)
    else:
        audit = render_bundle(args.bundle, args.output)
    print(json.dumps({key: audit[key] for key in ("output", "kind", "figure_count", "page_count")}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
