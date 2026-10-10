#!/usr/bin/env python3
"""Chinese PPT + bitmap-figure PDF export from completed observation evidence."""
from __future__ import annotations

import argparse
import json
from pathlib import Path
import subprocess
import sys

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager
from PIL import Image, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
import fitz

from src.inference.observation_semantics import block_interval, lagged
from src.inference.shared_driver_attribution import hb_diagnostic_coordinates, predict_standardized_ridge

DS = {'eeg_fnirs_single_trial': 'Single-Trial', 'simultaneous_eeg_nirs': 'Simultaneous',
      'visual_cognitive_motivation': 'Visual'}
COLORS = dict(observed='#263238', prediction='#087F8C', physical='#315B8A', component='#D67B32',
              absent='#7E8790', correct='#087F8C', mismatched='#C74D4A', missing_at_inference='#315B8A')
ARM_LABELS = dict(EOG_real='真实 EOG', EOG_wrong_subject='错被试 EOG', EOG_time_shift='时移 EOG',
    training_mean='训练基线', EOG_instant='仅即时事件', shared_optode='共享 optode',
    disjoint_distance_matched='非共享／匹配距离', disjoint_far='非共享／远端', training_template='条件时间模板',
    absent='无锚点', correct='正确锚点', mismatched='错配锚点', missing_at_inference='推理缺失')
PAIR_LABELS = {'tau/Hb_gain': 'τ / 增益', 'tau/extra_HbT': 'τ / 额外 HbT', 'tau/lag': 'τ / 时移',
               'Hb_gain/extra_HbT': '增益 / 额外 HbT', 'Hb_gain/lag': '增益 / 时移', 'extra_HbT/lag': '额外 HbT / 时移'}


def read_json(path):
    return json.loads(Path(path).read_text())


def rgb(value):
    return RGBColor.from_string(value.lstrip('#'))


def fmt(value, digits=3):
    return f'{value:.{digits}f}'


def ci_text(row, digits=3):
    return f"{row['mean']:.{digits}f} [{row['low']:.{digits}f}, {row['high']:.{digits}f}]"


def select(frame, **filters):
    for key, value in filters.items():
        frame = frame[frame[key] == value]
    if len(frame) != 1:
        raise ValueError(f'expected one evidence row for {filters}; got {len(frame)}')
    return frame.iloc[0]


def setup_plots():
    font = Path('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc')
    font_manager.fontManager.addfont(str(font))
    plt.rcParams.update({'font.family': font_manager.FontProperties(fname=str(font)).get_name(),
        'font.size': 12, 'axes.titlesize': 14, 'axes.labelsize': 12, 'xtick.labelsize': 11,
        'ytick.labelsize': 11, 'legend.fontsize': 10.5, 'axes.unicode_minus': False,
        'axes.spines.top': False, 'axes.spines.right': False, 'axes.grid': True,
        'grid.alpha': .18, 'figure.facecolor': 'white', 'savefig.facecolor': 'white'})


class Deck:
    def __init__(self, out, run, refinement):
        self.out = out
        self.run = run
        self.refinement = refinement
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(13.333), Inches(7.5)
        self.sources = []
        self.figure_paths = []

    def text(self, slide, x, y, w, h, content, size=18, color='#26384A', bold=False):
        # Conservative CJK-aware layout before export; retain editable text.
        requested = size
        for _ in range(24):
            font = ImageFont.truetype('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc', max(1, round(size*4)))
            lines = 0
            for paragraph in content.split('\n'):
                width = 0.
                lines += 1
                for char in paragraph:
                    advance = font.getlength(char)/4
                    if width+advance > (w-.055)*72:
                        lines += 1
                        width = 0.
                    width += advance
            height = lines*size*1.12+max(0, len(content.split('\n'))-1)*7
            if height <= (h-.035)*72:
                break
            size -= .25
        if size < max(7.5, requested-5):
            raise ValueError(f'text too dense for slide: {content[:60]}')
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = box.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = Inches(.02)
        frame.margin_top = frame.margin_bottom = Inches(.015)
        frame.vertical_anchor = MSO_ANCHOR.TOP
        for i, line in enumerate(content.split('\n')):
            p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            p.text = line
            p.font.name = 'Noto Sans CJK SC'
            p.font.size = Pt(size)
            p.font.bold = bold
            p.font.color.rgb = rgb(color)
            p.space_after = Pt(7 if i < len(content.split('\n'))-1 else 0)
            p.line_spacing = 1.12
        return box

    def base(self, title, subtitle, source):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = rgb('#F7F9FC')
        self.text(slide, .5, .3, 12.2, .65, title, size=26, bold=True)
        self.text(slide, .52, 1.04, 12.2, .5, subtitle, size=13, color='#526477')
        self.text(slide, .52, 7.14, 11.8, .25, '证据：'+source, size=8.5, color='#65758A')
        self.text(slide, 12.25, 7.1, .55, .3, f'{len(self.prs.slides):02d}', size=11, color='#65758A')
        self.sources.append(dict(slide=len(self.prs.slides), title=title, source=source))
        return slide

    def picture(self, slide, path, x=.5, y=1.6, w=8.45, h=4.85):
        iw, ih = Image.open(path).size
        scale = min(w/iw, h/ih)
        ww, hh = iw*scale, ih*scale
        slide.shapes.add_picture(str(path), Inches(x+(w-ww)/2), Inches(y+(h-hh)/2), width=Inches(ww), height=Inches(hh))

    def page(self, title, subtitle, figure, paragraphs, conclusion, source):
        slide = self.base(title, subtitle, source)
        if figure is not None:
            self.picture(slide, figure)
            x, w, size = 9.2, 3.55, 16.5
        else:
            x, w, size = .65, 12.0, 20
        y = 1.72
        for label, body in paragraphs:
            self.text(slide, x, y, w, .4, label, size=size, bold=True, color='#087F8C')
            height = 1.13 if figure is not None else .72
            self.text(slide, x, y+.42, w, height, body, size=size-1)
            y += 1.55 if figure is not None else 1.48
        self.text(slide, .62, 6.63, 12.05, .42, conclusion, size=15, color='#173D59', bold=True)
        return slide

    def table(self, title, subtitle, headers, rows, widths, conclusion, source, note=''):
        slide = self.base(title, subtitle, source)
        table = slide.shapes.add_table(len(rows)+1, len(headers), Inches(.6), Inches(1.75),
                                      Inches(12.1), Inches(4.1 if not note else 3.75)).table
        for c, width in enumerate(widths):
            table.columns[c].width = Inches(width)
        for r, row in enumerate([headers]+rows):
            for c, value in enumerate(row):
                cell = table.cell(r, c)
                cell.text = str(value)
                cell.margin_left = cell.margin_right = Inches(.09)
                cell.margin_top = cell.margin_bottom = Inches(.045)
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb('#173D59' if r == 0 else '#FFFFFF' if r % 2 else '#EAF0F5')
                for paragraph in cell.text_frame.paragraphs:
                    paragraph.font.name = 'Noto Sans CJK SC'
                    paragraph.font.size = Pt(14 if r else 15)
                    paragraph.line_spacing = 1.05
                    paragraph.space_after = Pt(0)
                    paragraph.font.bold = r == 0
                    paragraph.font.color.rgb = rgb('#FFFFFF' if r == 0 else '#26384A')
        if note:
            self.text(slide, .66, 5.72, 12.0, .75, note, size=14)
        self.text(slide, .65, 6.65, 12., .4, conclusion, size=15, bold=True, color='#173D59')
        return slide

    def figure(self, fig, name, *, tight=True):
        path = self.out/'figures'/f'{name}.png'
        if tight:
            fig.tight_layout(pad=1.2)
        fig.savefig(path, dpi=240)
        plt.close(fig)
        self.figure_paths.append(path)
        return path


def grouped_bar(deck, frame, groups, arms, name, *, value='nrmse', unit='subject', ylabel='NRMSE（训练 SD；越低越好）', percent=False):
    fig, ax = plt.subplots(figsize=(9, 5))
    group_values = list(frame[groups].drop_duplicates())
    offsets = (np.arange(len(arms))-(len(arms)-1)/2)*(.8/len(arms))
    colors = ['#087F8C', '#315B8A', '#B6854B', '#8B96A3']
    for i, arm in enumerate(arms):
        means, lo, hi = [], [], []
        for group in group_values:
            data = frame[(frame[groups] == group) & (frame.arm == arm)].groupby(unit)[value].mean()
            row = block_interval(data)
            mult = 100 if percent else 1
            means.append(row['mean']*mult)
            lo.append((row['mean']-row['low'])*mult)
            hi.append((row['high']-row['mean'])*mult)
        ax.bar(np.arange(len(group_values))+offsets[i], means, width=.75/len(arms),
               color=COLORS.get(arm, colors[i % len(colors)]), label=ARM_LABELS.get(arm, arm),
               yerr=np.array([lo, hi]), capsize=3)
    ax.set_xticks(np.arange(len(group_values)), [DS.get(g, PAIR_LABELS.get(g, g)) for g in group_values],
                  rotation=15 if len(group_values) > 3 else 0)
    ax.set_ylabel(ylabel)
    ax.legend(ncol=2, loc='upper center', bbox_to_anchor=(.5, 1.18), frameon=False)
    return deck.figure(fig, name)


def forest(deck, rows, labels, name, *, xlabel='NRMSE 改善＝对照 − 候选（正数更好）', colors=None):
    fig, ax = plt.subplots(figsize=(9, 5))
    for i, (row, label) in enumerate(zip(rows, labels)):
        ax.errorbar(row['mean'], i, xerr=[[max(0, row['mean']-row['low'])], [max(0, row['high']-row['mean'])]],
            fmt='o', capsize=4, color=colors[i] if colors else '#087F8C', markersize=7)
    ax.set_yticks(np.arange(len(labels)), labels)
    ax.invert_yaxis()
    ax.axvline(0, color='#53606D', linestyle='--')
    ax.set_xlabel(xlabel)
    return deck.figure(fig, name)


def build_report(run, refinement, out):
    if out.exists() and any(out.iterdir()):
        raise ValueError('choose a new versioned report export directory')
    out.mkdir(parents=True, exist_ok=True)
    (out/'figures').mkdir()
    for path in (run, refinement):
        if read_json(path/'summary.json')['execution'] != 'completed' or read_json(path/'verification.json')['status'] != 'passed':
            raise ValueError('completed and verified evidence required')
    setup_plots()
    deck = Deck(out, run, refinement)
    summary = read_json(run/'summary.json')
    a = pd.read_csv(run/'stability_metrics.csv')
    b = pd.read_csv(run/'event_metrics.csv')
    bc = pd.read_csv(run/'event_comparisons.csv')
    c0 = pd.read_csv(run/'hardware_comparisons.csv')
    c = pd.read_csv(refinement/'hardware_metrics.csv')
    cc = pd.read_csv(refinement/'hardware_comparisons.csv')
    topo = read_json(refinement/'hardware_topology.json')
    matched = {r['target'] for r in topo if r['available'] and abs(r['distance_mismatch']) <= .02}
    cm = c[c.target.isin(matched)]
    d0 = pd.read_csv(run/'mechanism_metrics.csv')
    d = pd.read_csv(refinement/'mechanism_metrics.csv')
    dcomp = pd.read_csv(refinement/'mechanism_comparisons.csv')
    ds = list(summary['B']['datasets'])
    gain_b = [select(bc, dataset=x, endpoint='EEG_voltage', contrast='real_vs_template') for x in ds]
    gain_c = select(cc, analysis_panel='distance_caliper', contrast='shared_vs_disjoint')
    mechanism_mean = d.groupby('arm')[['error', 'component_error', 'physical_error', 'false_component_projection']].mean()

    # One complete raster diagram, including arrows and labels.
    fig, ax = plt.subplots(figsize=(9, 4.8))
    ax.axis('off')
    boxes = [(.06, .65, '原生观测\nEEG / Hb / 光强 / EOG', '#E4EDF7'),
             (.43, .65, '独立证据与竞争解释\n稳定性 · 事件 · 硬件 · 合成', '#DCF0EE'),
             (.06, .12, 'Semantic 分支\n模型条件状态与响应', '#E4EDF7'),
             (.58, .12, 'Observation 分支\n测量条件＋额外上下文＋来源未定', '#FFF0DE')]
    for x, y, text, color in boxes:
        ax.text(x, y, text, transform=ax.transAxes, fontsize=16, va='center', ha='left',
                bbox=dict(boxstyle='round,pad=.8', fc=color, ec='none'))
    ax.annotate('', xy=(.43, .66), xytext=(.35, .66), xycoords='axes fraction', arrowprops=dict(arrowstyle='->', lw=2, color='#65758A'))
    for target in ((.20, .28), (.76, .28)):
        ax.annotate('', xy=target, xytext=(.62, .49), xycoords='axes fraction', arrowprops=dict(arrowstyle='->', lw=2, color='#65758A'))
    diagram = deck.figure(fig, '01_framework')
    deck.page('Observation 分支：从拟合到可检验证据',
        'A–D 实验套件｜公开实测开发＋已知真值合成｜2026-10-09', diagram,
        [('电位关联成立', f'EOG 的跨被试电位预测改善 {gain_b[0]["mean"]:.3f} / {gain_b[1]["mean"]:.3f} NRMSE。'),
         ('空间信息不能简化为硬件', f'距离匹配后，共享 donor 仍改善 {gain_c["mean"]:.3f}；非共享 donor 也能预测。'),
         ('来源名称仍需独立观测', 'Hb 来源归属未闭合；错配锚点会明显增加合成错误归属。')],
        '结论附着于具体读出；分支名称本身不提供生理来源资格。',
        'event_comparisons.csv；refinement/hardware_comparisons.csv、mechanism_metrics.csv')

    deck.table('本次回答四个问题，保留一个数据缺口',
        'A 是保留证据的重新分析；B/C 是新的实测读出；D 为合成机制诊断。',
        ['组别', '问题', '证据单元', '允许的结论'],
        [['A', '同样能解释观测的分解，哪些量稳定？', '72 窗／18 人；有限 profile 子集', '条件范围和未定区域'],
         ['B', 'EOG 能预测什么实际变化？', '55 人；440 个原生窗口', '电位、光学或 Hb 的具体关联'],
         ['C', '空间收益是否只来自共享硬件？', 'Single-Trial；14 人评价', '共享与非共享 donor 的对照'],
         ['D', '相似观测下锚点是否保护归属？', '两套各 128 个评价生成身份', '限定生成模型内的机制辨别'],
         ['E', '能否命名浅层或全身生理来源？', '缺独立短间距／外周生理验证集', '本次未执行，不能补写来源']],
        [1.0, 4.0, 3.4, 3.7], '实测结果是公开开发证据，未开启受保护评价或训练新的 tokenizer。',
        'resolved_config.yaml；summary.json；refinement/resolved_config.yaml')

    deck.table('被试隔离、原生时钟和缺失支持均保留',
        'B/C 按原始编号奇数训练、偶数评价；每记录最多八个双时钟不重叠的 30 秒窗。',
        ['数据集', '训练', '评价', '原生信号', '本次用途'],
        [['Single-Trial', '15 人／120 窗', '14 人／112 窗', '200 Hz 电位＋10 Hz 光强', 'B 电位／OD／Hb；C 硬件'],
         ['Simultaneous', '13 人／104 窗', '13 人／104 窗', '200 Hz 电位＋发布 Hb', 'B 电位／Hb'],
         ['Visual', '沿用冻结外层训练尺度', 'A：6 人／24 窗', '保留的 Hb 分解数组', '只做 A 的证据分析'],
         ['REFED', '—', '—', '未新增消费原生数组', '本轮不承担来源验证']],
        [2.25, 2.25, 2.25, 2.7, 2.65],
        '缺失属于具体模态／通道；没有用 QC 注释把有效信号自动置零。',
        'measured_plan.json；prepared/*.json；DATA_CONTRACT.md',
        '55/55 记录、440/440 窗完成。旧 pilot 因一个坏光学通道误丢整窗，修正后以新目录重跑；旧记录保留。')

    deck.page('从训练到评价：哪些量被冻结',
        '训练、当前窗口变换和评分承担不同职责；所有图均为离线特征分析。', None,
        [('仅在训练被试拟合', 'EOG 阈值和尺度、带截距 ridge 系数、Hb 训练 SD，以及按任务条件与相对时间构造的平均模板。'),
         ('当前窗口只作固定变换', 'EEG 电位带通并重采样；Hb 经原有 0.01–0.2 Hz 处理和前 5 秒参考；EOG 检测与 FIR 特征不使用评价目标。'),
         ('评价端点只用于评分', 'B 的 EEG/Hb/OD 目标不用于回归重拟合；C 的目标两路 Hb 全部隐藏，输入为 donor Hb 和已知任务条件。')],
        '原生时钟用冻结的 EEG／fNIRS 双锚点对应；不宣称实时或严格未来预测。',
        'evaluate_observation_semantics.py：prepare_worker、event_analysis、hardware_analysis')

    deck.page('指标与区间：先窗口，再被试',
        '误差和差值方向在全套报告保持一致。', None,
        [('NRMSE 的分母', '每窗：sqrt(mean(((预测 − 目标) / 训练 SD)²))。SD 是训练标准差，RMS 是均方根。EEG 按六通道、Hb 按两色团评分；不以当前窗尺度作分母。'),
         ('改善的方向', 'Δ = 对照 NRMSE − 候选 NRMSE；正数表示候选更好。先同窗／同目标配对，再被试内平均，最后被试等权。'),
         ('95% 区间与多重比较', '2,000 次被试块 bootstrap 百分位区间；D 以生成身份为块。主要检验使用单侧 sign-flip 和族内 Holm，窗口数不当作独立样本量。')],
        '所有区间条件于冻结模型与既定划分；不是重新训练模型的总体确认区间。',
        'observation_semantics.py：block_interval、sign_flip_p、holm；各 comparisons.csv')

    deck.page('A｜可测坐标不等于组织来源',
        '先保留 HbO/HbR 的相对幅度，再变换；ρ 是固定诊断常数。', None,
        [('两个坐标', 'T = ΔHbO + ΔHbR；Xρ = ρΔHbO − (1−ρ)ΔHbR，ρ=0.35。T 是表观总 Hb 变化；Xρ 是氧合分配相关坐标。'),
         ('物理解释的限制', '只有参考氧合比例 S₀=1−ρ 时，Xρ 才能抵消总量变化。当前没有测得个体 S₀；不能称绝对血氧饱和度、脑血容量或代谢率。'),
         ('比例不能直接定源', 'T 与 Xρ 不是等范数正交坐标；物理项与额外项可相关或互相抵消。坐标平方量比值不能解释为来源能量百分比。')],
        '固定 [0.65,0.35] 加载令额外项 X₀.₃₅≈0，这是代数约束，不是来源已辨识。',
        'shared_driver_attribution.py：hb_diagnostic_coordinates；用户计划 §三')

    deck.table('A｜三个集合回答不同的稳定性问题',
        '范围是有限候选的包络，不是生理后验或校准置信区间。',
        ['集合', '变化对象', '接纳规则', '不能据此声称'],
        [['近等价 profile', '固定成分位移后，物理路径重新拟合', '工程目标增幅 ≤5%；同一目标', '已搜索全部可能解释'],
         ['测量／先验敏感性', '增益、时移、加载、先验与训练尺度', '保留预定的收敛变体', '属于同一似然置信集合'],
         ['响应／初态方向', 'BC/HselectedC；初态/响应切线', '复用强化实验的冻结解与 profile', '自由响应参数的完整 profile'],
         ['teacher 字段', '每时刻 lower/upper、稳定掩码和候选数', '范围阈值 0.20 训练坐标 SD；至少两解', '掩码等于真实来源概率']],
        [2.25, 3.4, 3.35, 3.1], '单一候选的范围虽为零，仍不提供稳定解释证据。',
        'stability_metrics.csv；stability/*.npz；stability_worker')

    # A finite profile range and one fully traceable example.
    aprof = a[(a.kind == 'near_equivalent_profile') & (a.arm == 'D-fixed') & (a.coordinate == 'HbT')]
    fig, ax = plt.subplots(figsize=(9, 5))
    fields = [('prediction', '完整预测', COLORS['prediction']), ('physical_prediction', '物理项', COLORS['physical']),
              ('observation_component', '额外项', COLORS['component'])]
    ids = sorted(aprof.id.unique())
    labels = []
    for k, identity in enumerate(ids):
        part = aprof[aprof.id == identity]
        ref = part.iloc[0]
        labels.append(DS[ref.dataset]+'\n'+ref.subject)
        for j, (field, label, color) in enumerate(fields):
            row = part[part.field == field].iloc[0]
            ax.scatter(k+(j-1)*.18, row.rms_range_SD, color=color,
                marker='o' if row.sufficient_candidates else 'x', s=60, label=label if k == 0 else None)
    ax.axhline(.2, linestyle='--', color='#7E8790', label='0.20 SD 工程阈值')
    ax.set_xticks(range(len(ids)), labels)
    ax.set_ylabel('候选范围的时间 RMS / 训练 HbT SD')
    ax.legend(ncol=2)
    fig_a = deck.figure(fig, '02_profile_ranges')
    valid = aprof[(aprof.field == 'observation_component') & aprof.sufficient_candidates]
    deck.page('A｜完整预测接近，额外 HbT 仍可变化',
        '每组三点是一个保留窗口；圆点表示至少两解，叉号表示只有参考解。', fig_a,
        [('分母必须可见', f'固定方向 profile 共 6 窗，仅 {len(valid)} 窗存在接纳的非零位移；其余不能判定“可靠”。'),
         ('幅度分配可能互相抵消', '部分窗口中物理 HbT 与额外 HbT 的变化远大于完整预测变化，支持保留归属范围。'),
         ('阈值不是统计检验', '虚线仅为预定 0.20 训练 SD 的工程稳定性阈值；6 窗也不是随机人群抽样。')],
        '可预测的完整 Hb 变化，仍可能没有唯一的物理／额外项分配。',
        'stability_metrics.csv：near_equivalent_profile、D-fixed、HbT')

    identity = sorted(valid.id)[0]
    archive = np.load(run/'stability'/f'{identity}.npz')
    fig, axes = plt.subplots(3, 1, figsize=(9, 5.1), sharex=True)
    for ax, (field, label, color) in zip(axes, fields):
        values = archive[f'near_equivalent_profile__D-fixed__{field}__candidates']
        coords = hb_diagnostic_coordinates(values)[:, :, 0]/archive['coordinate_scale'][0]
        for j, curve in enumerate(coords):
            ax.plot(np.arange(120)/4, curve, color=color, alpha=.55, linewidth=1.3,
                    label='接纳解释' if j == 0 else None)
        if field == 'prediction':
            ax.plot(np.arange(120)/4, archive['observed_coordinates'][:, 0]/archive['coordinate_scale'][0],
                    color=COLORS['observed'], linewidth=1.7, label='实测目标')
        ax.set_ylabel(label+'\n训练 SD')
        ax.legend(loc='upper right')
    axes[-1].set_xlabel('相对窗口时间（秒）')
    fig_example = deck.figure(fig, '03_profile_waveforms')
    example_rows = aprof[aprof.id == identity].set_index('field')
    deck.page('A｜同一窗口中的相互抵消是可见的',
        '示例选择：有非零接纳位移的窗口中按稳定身份排序取首个；没有按图形美观挑选。', fig_example,
        [('本页身份', identity.replace('eeg_fnirs_single_trial__', '').replace('__prefrontal__', ' / ').replace('__', ' / ')),
         ('同一批解释', f'完整预测范围 {example_rows.loc["prediction", "rms_range_SD"]:.3f} SD；物理项 {example_rows.loc["physical_prediction", "rms_range_SD"]:.3f}；额外项 {example_rows.loc["observation_component", "rms_range_SD"]:.3f}。'),
         ('纵轴独立但单位相同', '每面板均除以同一训练 HbT SD；为读清峰谷，纵轴范围独立。颜色在全套报告一致。')],
        '黑线是参与拟合的实测观测；彩色曲线是模型条件解释，不是潜变量真值。',
        f'stability/{identity}.npz；原 attribution profiles')

    fig, axes = plt.subplots(1, 2, figsize=(9, 5))
    for ax, field, title in zip(axes, ('physical_prediction', 'observation_component'), ('物理项', '额外项')):
        sub = a[(a.kind == 'sensitivity') & (a.arm == 'M-observation') & (a.field == field)]
        wide = sub.pivot(index=['id', 'dataset'], columns='coordinate', values='rms_range_SD').reset_index()
        for dsid, part in wide.groupby('dataset'):
            ax.scatter(part.HbT, part.HbX, s=28, alpha=.7, label=DS[dsid])
        ax.axvline(.2, color='gray', linestyle='--')
        ax.axhline(.2, color='gray', linestyle='--')
        ax.set(xlabel='HbT 范围（训练 SD）', ylabel='Xρ 范围（训练 SD）', title=title)
    axes[0].legend()
    a_sensitivity = deck.figure(fig, '04_sensitivity_coordinates')
    coverage = a[(a.kind == 'sensitivity') & (a.arm == 'M-observation') & (a.field == 'observation_component') & (a.coordinate == 'HbT')]
    deck.page('A｜稳定性取决于具体坐标与候选集合',
        '每点一个窗口；72 窗属于18名被试，重复窗口有依赖关系。', a_sensitivity,
        [('读图方法', '横轴为 HbT 范围，纵轴为 Xρ 范围；左下区域同时小于预定 0.20 SD 阈值。'),
         ('额外 HbT 的覆盖率', f'M-observation 的额外 HbT 在 {int(coverage.stable_coordinate.sum())}/{len(coverage)} 窗低于阈值；这只对当前敏感性集合成立。'),
         ('加载改变也改变坐标', '本集合含 ρ=0.25/0.45，所以 Xρ 不再被固定方向恒等约束为零。不能将某一集合的稳定性推广到全部机制。')],
        '输出每个字段的范围、覆盖率和候选来源，优于一个笼统的“可靠／不可靠”标签。',
        'stability_metrics.csv：sensitivity、M-observation')

    fig, ax = plt.subplots(figsize=(9, 5))
    resp = a[(a.kind == 'response_choice') & (a.field == 'observation_component') & (a.coordinate == 'HbT')]
    for i, (dsid, part) in enumerate(resp.groupby('dataset')):
        values = part[part.sufficient_candidates].rms_range_SD.to_numpy()
        ax.scatter(np.full(len(values), i)+np.linspace(-.17, .17, len(values)), values, s=30, alpha=.75, label=DS[dsid])
        ax.text(i, ax.get_ylim()[0], f'不同解 {len(values)}/{len(part)}', va='bottom', ha='center', fontsize=11)
    ax.set_xticks(range(3), [DS[v] for v in sorted(resp.dataset.unique())])
    ax.axhline(.2, color='gray', linestyle='--')
    ax.set_ylabel('BC → 训练选择响应：额外 HbT 变化（训练 SD）')
    a_response = deck.figure(fig, '05_response_choice')
    deck.page('A｜更换响应解释，也会移动额外成分',
        '只比较已保存的 BC 与 HselectedC；没有重新利用评价数据选择响应。', a_response,
        [('响应选择', '图中每点是两种响应产生不同结果的窗口；完全相同的解计入分母，但不冒充第二个解释。'),
         ('初态与响应方向', '另保留 initial/response 切线引导的位移 profile 及对应包络，可逐字段查阅。'),
         ('适用上限', '响应参数在这些 profile 中仍被冻结；本轮没有新增原生连续状态拟合，也不能排除其竞争解释。')],
        '动态模型选择的影响必须进入 teacher 元数据；不能统一归为测量噪声。',
        'stability_metrics.csv：response_choice、near_equivalent_initial、near_equivalent_response')

    deck.table('A｜已交付可供 teacher 使用的字段',
        '这是条件监督接口，不是已经晋级的生理 teacher。',
        ['字段', '实际含义', '使用限制'],
        [['lower / upper', '同一解释集合内每时刻 HbT / Xρ 的最小／最大值', '集合范围，不是概率区间'],
         ['candidate_count', '参考解加接纳的不同解释数', '仅一个候选时，稳定掩码为假'],
         ['pointwise_stable', '范围不超过 0.20 个训练坐标 SD', '工程阈值，不代表来源唯一'],
         ['sufficient_candidates', '是否存在至少两种接纳解释', '约束性零坐标需单独标记解释'],
         ['来源与比较类型', 'profile、测量敏感性或响应选择', '不能混合为校准生理后验']],
        [3.1, 5.55, 3.45], '本次没有以这些字段重新训练 tokenizer，也没有把 EOG 关联写成 Hb 来源标签。',
        'observation_semantics.py：coordinate_stability；stability/*.json、*.npz')

    deck.page('B｜电位混入与慢 Hb 响应分开检验',
        '输入是实际 HEOG/VEOG；评价目标不再是自由拟合的共同残差。', None,
        [('电位路径', '原生 200 Hz → 0.5–15 Hz 零相位带通 → 50 Hz。signed HEOG/VEOG 的 −0.04、0、+0.04 秒输入，经训练冻结 ridge 预测六路前额 EEG。'),
         ('事件路径', '按训练尺度检测 VEOG 幅度和 HEOG 斜率越阈候选，阈值 3 SD、间隔 0.5 秒。有限脉冲响应（FIR）即时滞后0/0.25/0.5/1秒，延迟滞后2/4/6/8秒。'),
         ('独立目标与对照', '目标为 HbT/Xρ 和原始 OD 一步差分幅度；比较真实、同条件错训练被试和不环绕的 +8 秒时移。Hb/OD 的基线为训练条件时间模板。')],
        '三个机制分开表述：眼电相关电位、眼动伴随光学变化、眼动相关慢 Hb；关联不自动定源。',
        'resolved_config.yaml：events；event_models.json；event_inventory.csv')

    fig_b = grouped_bar(deck, b[b.endpoint == 'EEG_voltage'], 'dataset',
        ['training_mean', 'EOG_real', 'EOG_wrong_subject', 'EOG_time_shift'], '06_eeg_prediction')
    deck.page('B｜真实 EOG 可预测跨被试前额电位',
        '柱为被试等权平均误差；误差棒为冻结模型下的 95% 被试块区间。', fig_b,
        [('Single-Trial：14 人', f'相对训练均值改善 {ci_text(gain_b[0])} NRMSE；112/112 窗配对。'),
         ('Simultaneous：13 人', f'改善 {ci_text(gain_b[1])}；104/104 窗配对。'),
         ('两个 null 同样通过', '真实 EOG 优于错被试与时移；六项主要电位比较经 Holm 后均通过。支持 EOG 相关电信号读出。')],
        '本结果不保证减去预测后得到“纯神经 EEG”，更不能将同一来源名称传播给 Hb。',
        'event_metrics.csv、event_comparisons.csv：EEG_voltage')

    models = read_json(run/'event_models.json')
    fig, axes = plt.subplots(1, 2, figsize=(9, 5))
    matrices = [np.array(models[dsid]['voltage_model']['coefficient']).T for dsid in ds]
    lim = max(abs(m).max() for m in matrices)
    for ax, dsid, matrix in zip(axes, ds, matrices):
        spec = next(s for s in read_json(run/'measured_plan.json')['records'] if s['dataset'] == dsid)
        names = read_json(run/'prepared'/f"{spec['key']}.json")['eeg_channels']
        im = ax.imshow(matrix, vmin=-lim, vmax=lim, cmap='RdBu_r', aspect='auto')
        ax.set_yticks(range(6), names)
        ax.set_xticks(range(6), ['H −40', 'V −40', 'H 0', 'V 0', 'H +40', 'V +40'], rotation=50, ha='right')
        ax.set_title(DS[dsid])
        ax.set_xlabel('EOG 类型／时滞（毫秒）')
        ax.grid(False)
    fig.subplots_adjust(left=.08, right=.82, bottom=.23, top=.89, wspace=.5)
    color_axis = fig.add_axes([.85, .23, .025, .66])
    fig.colorbar(im, cax=color_axis, label='回归系数：µV / 训练输入 SD')
    coef_fig = deck.figure(fig, '07_eog_spatial_coefficients', tight=False)
    deck.page('B｜signed 映射保留通道与方向信息',
        '每格是一个训练冻结系数；H=HEOG，V=VEOG；两图共享颜色尺度。', coef_fig,
        [('比 log-energy 更直接', 'signed 电位映射保留正负方向，可检查不同前额电极的投影；无需把残差当成眼电真值。'),
         ('系数是联合回归结果', '三个相近时滞高度相关，单格系数不能单独解释为电流传播方向或电极敏感度。'),
         ('预测而非拟合说明证据', '这些系数只来自奇数编号训练被试；上一页的偶数编号被试提供独立预测证据。')],
        'EOG 来源锚点支持的是这一电位读出；不是整条 O 表征或全部残差的来源。',
        'event_models.json：voltage_model；prepared/*.json：eeg_channels')

    spec = next(s for s in read_json(run/'measured_plan.json')['records'] if s['dataset'] == ds[0] and s['split'] == 'evaluate')
    meta = read_json(run/'prepared'/f"{spec['key']}.json")
    with np.load(run/'prepared'/f"{spec['key']}.npz") as arrays:
        eeg_observed = arrays['eeg_voltage'][0]
        eog = arrays['eog_voltage'][0]
    predicted = predict_standardized_ridge(models[ds[0]]['voltage_model'], lagged(eog, [-2, 0, 2]))
    fig, axes = plt.subplots(2, 1, figsize=(9, 5), sharex=True)
    t = np.arange(1500)/50
    axes[0].plot(t, eog[:, 0], color='#7E8790', linewidth=1, label='HEOG')
    axes[0].plot(t, eog[:, 1], color='#D67B32', linewidth=1, label='VEOG')
    axes[1].plot(t, eeg_observed[:, 0], color='#263238', linewidth=1.1, label=f'实测 {meta["eeg_channels"][0]}')
    axes[1].plot(t, predicted[:, 0], color='#087F8C', linewidth=1, label='训练冻结 EOG 预测')
    for ax in axes:
        ax.axvspan(8, 21, color='#087F8C', alpha=.07)
        ax.set_ylabel('滤波电位（µV）')
        ax.legend(loc='upper right')
    axes[-1].set_xlabel('窗口时间（秒）；淡色区域为共同评分支持')
    wave_eeg = deck.figure(fig, '07b_native_eog_waveform')
    example_score = select(b, dataset=ds[0], endpoint='EEG_voltage', arm='EOG_real', id=meta['windows'][0]['id'])
    deck.page('B｜EOG 与电位的对应可直接核验',
        '示例固定取 Single-Trial 首名评价被试的首窗；没有按预测成绩选择。', wave_eeg,
        [('上图是输入', 'HEOG 与 VEOG 经过相同的0.5–15 Hz滤波；保留正负电位。两路来自实际辅助电极。'),
         ('下图是独立目标', f'展示 {meta["eeg_channels"][0]} 的电位和冻结映射预测；该窗口六通道总体 NRMSE={example_score.nrmse:.3f}。'),
         ('峰谷与残余', '同步大幅变化可被预测，局部偏差仍存在；两面板纵轴独立，不以视觉重合代替六通道指标。')],
        '目标 EEG 没有参与本窗口的回归拟合；本页是跨被试冻结预测。',
        f'prepared/{spec["key"]}.npz；event_models.json；event_metrics.csv')

    optical_rows, labels = [], []
    for dsid in ds:
        for endpoint in ('HbT', 'HbX', 'OD_jump'):
            part = bc[(bc.dataset == dsid) & (bc.endpoint == endpoint) & (bc.contrast == 'real_vs_template')]
            if len(part):
                optical_rows.append(part.iloc[0])
                labels.append(DS[dsid]+' · '+endpoint)
    fig_hb = forest(deck, optical_rows, labels, '08_eog_hb_optical')
    deck.page('B｜电位证据没有延伸成稳定 Hb 归属',
        '点为“模板误差 − EOG 模型误差”；横线为95%被试块区间，正值才表示改善。', fig_hb,
        [('HbT 并未获益', f'Single-Trial：{ci_text(select(bc, dataset=ds[0], endpoint="HbT", contrast="real_vs_template"))}；Simultaneous 同样没有正面改善。'),
         ('原始光学突变也未获益', f'Single-Trial OD 差分幅度：{ci_text(select(bc, dataset=ds[0], endpoint="OD_jump", contrast="real_vs_template"))}。'),
         ('负结果的实际范围', '它限制此事件检测、FIR、训练划分及目标通道；不证明眼动与所有 Hb 或接触变化完全无关。')],
        '本轮可称“EOG 相关电位”，不能称“已识别眼动导致的 Hb 成分”。',
        'event_comparisons.csv：real_vs_template；各 Hb/OD 端点')

    delayed = bc[bc.contrast == 'delayed_increment']
    delay_fig = forest(deck, [row for _, row in delayed.iterrows()],
        [DS[row.dataset]+' · '+row.endpoint for _, row in delayed.iterrows()], '09_delayed_increment')
    deck.page('B｜加入延迟事件项没有一致增量',
        '对照模型只用 0–1 秒事件滞后；候选再加入 2/4/6/8 秒，训练数据和目标保持相同。', delay_fig,
        [('比较方向', '横轴为“仅即时事件误差 − 即时＋延迟误差”。Single-Trial HbT 的变化偏负，其他端点缺少一致正面证据。'),
         ('不能把延迟当生理证明', 'Hb 使用离线零相位低频处理；一个延迟回归系数不等于直接测得的生理响应时延。'),
         ('仍存在竞争解释', '即使未来观察到事件相关 Hb，也要区分电位混入、伴随运动／接触与延迟生理作用。')],
        '保留具体端点的负结果；没有把“慢变化”升级为自主神经或浅层血流。',
        'event_comparisons.csv：delayed_increment；event_models.json：instant_model')

    power = pd.read_csv(run/'power_chain.csv')
    fig, axes = plt.subplots(1, 2, figsize=(9, 5))
    bands = ['1-4Hz', '4-8Hz', '8-13Hz', '13-30Hz', '30-45Hz']
    for dsid, part in power.groupby('dataset'):
        for ax, metric in zip(axes, ['logpower_change', 'power_cross_term_relative_RMS']):
            means = part.groupby(['subject', 'band'])[metric].mean().groupby('band').mean().reindex(bands)
            ax.plot(range(5), means, marker='o', label=DS[dsid])
            ax.set_xticks(range(5), bands, rotation=30)
    axes[0].set_ylabel('mean[log P(余量) − log P(原始)]')
    axes[1].set_ylabel('功率交叉项 RMS / 原始功率 RMS')
    axes[0].axhline(0, color='gray', linestyle='--')
    axes[0].legend()
    power_fig = deck.figure(fig, '10_power_nonadditivity')
    deck.page('B｜电位可相加，功率与 log-power 不可相加',
        '冻结 EOG 电位预测先从原生 EEG 相减，再经过原有五频带功率链；图为被试等权均值。', power_fig,
        [('左图是变换后的影响', '低频变化通常更明显；纵轴是自然对数差，不是百分比，也不是“去除了多少神经信号”。'),
         ('右图检验非加性', 'P(原始) − P(预测电位) − P(余量) 的 RMS 非零，反映平方运算的交叉项。'),
         ('没有“纯净 EEG”真值', '余量只是减去一个受 EOG 支持的预测之后的信号；未验证其全部为神经来源。')],
        '不要在 log-power 残差上直接贴“眼电幅度”标签。',
        'power_chain.csv；observation_baselines.py：eeg_band_power')

    events = pd.read_csv(run/'event_inventory.csv')
    fig, ax = plt.subplots(figsize=(9, 5))
    sub = events[events.split == 'evaluate']
    for i, dsid in enumerate(ds):
        counts = sub[sub.dataset == dsid].groupby(['subject', 'kind']).size().unstack(fill_value=0)
        for j, label in enumerate(('VEOG 幅度候选', 'HEOG 斜率候选')):
            vals = counts[j].to_numpy()
            x = i*3+j
            ax.scatter(np.full(len(vals), x)+np.linspace(-.15, .15, len(vals)), vals, s=35, alpha=.8)
            ax.plot([x-.2, x+.2], [np.median(vals)]*2, color='black', lw=2)
    ax.set_xticks(range(0, 5), ['Single\nVEOG', 'Single\nHEOG', '', 'Simultaneous\nVEOG', 'Simultaneous\nHEOG'])
    ax.set_ylabel('每名评价被试八个窗口内的候选事件数')
    counts_fig = deck.figure(fig, '11_event_counts')
    deck.page('B｜检测事件数不是独立样本量',
        '每点一名被试，黑横线为中位数；候选事件由信号阈值检测而来。', counts_fig,
        [('完整数量', f'训练＋评价共检测 {len(events):,} 次候选事件；评价部分 {len(sub):,} 次。两类检测可能重叠。'),
         ('事件名称保持谨慎', '未用人工眨眼／眼跳标注核验灵敏度与特异性，因此称 VEOG 幅度／HEOG 斜率越阈候选。'),
         ('统计单位仍是被试', '电位与 Hb 的区间按14或13名评价被试重采样；不把几千次事件当作几千名独立受试者。')],
        '事件级设计增加机制具体性，但不会自动增加独立样本量。',
        'event_inventory.csv；measured_plan.json')

    # C geometry diagram and distance balance.
    geo = read_json(run/'prepared/eeg_fnirs_single_trial__subject_01__session_00.json')['geometry']
    points = np.array([[r[k] for k in ('x', 'y')] for r in geo])
    chosen = next(t for t in topo if t['target'] in matched)
    fig, axes = plt.subplots(1, 2, figsize=(9, 5))
    ax = axes[0]
    ax.scatter(points[:, 0], points[:, 1], color='#C5CDD5', s=22)
    for arm, color in [('shared_optode', '#087F8C'), ('disjoint_distance_matched', '#315B8A')]:
        j = chosen['indices'][arm]
        ax.plot(points[[chosen['target_index'], j], 0], points[[chosen['target_index'], j], 1], color=color, lw=2)
        ax.scatter(*points[j], color=color, s=80, label=ARM_LABELS[arm])
        ax.annotate(geo[j]['base_channel_name'], points[j], xytext=(4, 7), textcoords='offset points', fontsize=10)
    ax.scatter(*points[chosen['target_index']], marker='*', color='black', s=150, label='隐藏目标')
    ax.set(xlabel='模板 x', ylabel='模板 y', title='一个纯几何选择的例子')
    ax.legend(loc='lower left', fontsize=9.5)
    ax.set_aspect('equal')
    x = [r['distances']['shared_optode'] for r in topo]
    y = [r['distances']['disjoint_distance_matched'] for r in topo]
    axes[1].scatter(x, y, c=['#087F8C' if r['target'] in matched else '#B6BDC6' for r in topo], s=45)
    lim = max(max(x), max(y))*1.04
    axes[1].plot([0, lim], [0, lim], '--', color='gray')
    axes[1].set(xlabel='共享 donor 距离', ylabel='非共享 donor 距离', title='绿色：距离差 ≤0.02')
    geom_fig = deck.figure(fig, '12_hardware_geometry')
    deck.page('C｜先匹配几何，再比较硬件共享',
        '左：通道中心的模板投影，不是个体组织深度；右：每点一个目标通道。', geom_fig,
        [('三臂同容量', '各使用一个 donor 的 HbO/HbR，共两个预测输入和一个截距；训练冻结 ridge。'),
         ('主比较限于20通道', '在所有共享／非共享配对中最小化距离差，并取 ≤0.02 模板头单位的20/36个目标。'),
         ('共享不等于纯硬件', '共享 source/detector 也改变组织重叠；模板距离匹配只能减少距离混杂。')],
        '没有 optode 身份的 Simultaneous 不进入本项硬件检验。',
        'refinement/hardware_topology.json；prepared/*.json：geometry')

    fig_c = grouped_bar(deck, cm[cm.pairing == 'real'].assign(panel='20 个距离匹配目标'), 'panel',
        ['training_template', 'shared_optode', 'disjoint_distance_matched', 'disjoint_far'], '13_hardware_prediction')
    gc2 = select(cc, analysis_panel='distance_caliper', contrast='disjoint_distance_matched_vs_template')
    deck.page('C｜距离匹配后仍有共享优势',
        '柱为14名被试等权误差；每人先平均八窗及20目标；2,240个目标窗口配对。', fig_c,
        [('共享对非共享', f'改善 {ci_text(gain_c)} NRMSE；Holm p={gain_c.holm_p:.5f}。'),
         ('非共享也提供信息', f'非共享、距离匹配 donor 相对条件时间模板改善 {ci_text(gc2)}。'),
         ('两类结论并存', '共享关系有额外预测价值，同时空间共同变化不限于共享硬件的通道。')],
        '共享硬件不能解释全部空间收益；但本结果还不能识别浅层或脑内来源。',
        'refinement/hardware_metrics.csv、hardware_comparisons.csv：distance_caliper')

    hardware_models = read_json(refinement/'hardware_models.json')
    target_name = chosen['target']
    target_index = chosen['target_index']
    with np.load(run/'prepared'/f"{spec['key']}.npz") as arrays:
        all_hb = arrays['hb'][0]
    frozen = hardware_models['targets'][target_name]
    condition = meta['windows'][0]['condition']
    template = np.array(hardware_models['templates'][condition])[:, target_index]
    sd_hb = np.array(frozen['sd'])
    fig, axes = plt.subplots(2, 1, figsize=(9, 5), sharex=True)
    predictions = {arm: template+predict_standardized_ridge(frozen['models'][arm], all_hb[:, chosen['indices'][arm]])
                   for arm in ('shared_optode', 'disjoint_distance_matched')}
    for j, ax in enumerate(axes):
        ax.plot(np.arange(120)/4, all_hb[:, target_index, j]/sd_hb[j], color='#263238', linewidth=1.8, label='隐藏实测目标')
        ax.plot(np.arange(120)/4, template[:, j]/sd_hb[j], color='#7E8790', linestyle='--', label='训练条件模板')
        for arm, color in [('shared_optode', '#087F8C'), ('disjoint_distance_matched', '#315B8A')]:
            ax.plot(np.arange(120)/4, predictions[arm][:, j]/sd_hb[j], color=color, label=ARM_LABELS[arm])
        ax.axvspan(8, 21, color='#087F8C', alpha=.07)
        ax.set_ylabel(('HbO', 'HbR')[j]+' / 训练 SD')
        ax.legend(ncol=2, fontsize=9)
    axes[-1].set_xlabel('窗口时间（秒）')
    ccurve = deck.figure(fig, '13b_hidden_hb_waveform')
    chosen_rows = cm[(cm.window_id == meta['windows'][0]['id']) & (cm.target == target_name) & (cm.pairing == 'real')]
    values = {row.arm: row.nrmse for _, row in chosen_rows.iterrows()}
    deck.page('C｜目标完全隐藏，仍可预测部分波形',
        f'例：{spec["subject"]} 首窗、几何顺序首个合格目标 {target_name}；非按成绩选择。', ccurve,
        [('模型能看到什么', '仅一个 donor 的两路完整 Hb、训练冻结参数和任务条件；目标 HbO/HbR 都只用于画图和评分。'),
         ('本窗口的数值', f'模板 {values["training_template"]:.3f}；共享 {values["shared_optode"]:.3f}；非共享匹配 {values["disjoint_distance_matched"]:.3f} NRMSE。'),
         ('幅度与形态分开读', '模型可跟随部分峰谷，幅度与偏移仍有误差；本例不要求逐窗都优于模板，总体结论看配对区间。')],
        '纵轴各除以对应的训练色团 SD；不是原生 µM，也不是来源分量真值。',
        'prepared 数组；refinement/hardware_models.json、hardware_metrics.csv')

    rows, labels = [], []
    for arm in ('shared_optode', 'disjoint_distance_matched', 'disjoint_far'):
        for suffix, label in (('_vs_wrong', '对错被试'), ('_vs_shift', '对时移')):
            rows.append(select(cc, analysis_panel='distance_caliper', contrast=arm+suffix))
            labels.append(ARM_LABELS[arm]+' · '+label)
    cnull = forest(deck, rows, labels, '14_hardware_nulls')
    deck.page('C｜真实同步优于错被试与时移',
        '预测器、训练尺度和目标窗口保持固定；只替换 donor 信号。', cnull,
        [('正值的含义', '真实同步 donor 的误差低于同条件错训练被试及 +8 秒非环绕时移。'),
         ('匹配与缺失分母', '几何主比较为2,240/2,240配对；部分错被试 donor 有缺失，表中保留实际配对数。'),
         ('不能反推组织来源', '同步生理、共享传感器、采样组织重叠都可能贡献这些预测关系；null 只排除当前错配解释。')],
        '空间可预测性是独立约束，组织来源仍需额外测量。',
        'refinement/hardware_comparisons.csv：distance_caliper 的 wrong / shift')

    original_topo = read_json(run/'hardware_topology.json')
    fig, axes = plt.subplots(1, 2, figsize=(9, 5))
    differences = [np.abs([r['distance_mismatch'] for r in original_topo]),
                   np.abs([r['distance_mismatch'] for r in topo])]
    axes[0].boxplot(differences, tick_labels=['最近共享规则', '最小距离差规则'], showfliers=True)
    axes[0].axhline(.02, color='#087F8C', linestyle='--', label='主要比较阈值')
    axes[0].set_ylabel('共享／非共享距离差（模板头单位）')
    axes[0].legend(fontsize=10)
    comparison = [select(c0, contrast='shared_vs_disjoint'), select(cc, analysis_panel='all_targets', contrast='shared_vs_disjoint'), gain_c]
    axes[1].errorbar([r['mean'] for r in comparison], range(3),
        xerr=np.array([[r['mean']-r['low'] for r in comparison], [r['high']-r['mean'] for r in comparison]]), fmt='o', capsize=4, color='#087F8C')
    axes[1].set_yticks(range(3), ['原规则／36目标', '重匹配／36目标', '重匹配／20目标'])
    axes[1].invert_yaxis()
    axes[1].axvline(0, color='gray', linestyle='--')
    axes[1].set_xlabel('非共享误差 − 共享误差')
    balance = deck.figure(fig, '15_geometry_refinement')
    deck.page('C｜首轮距离不平衡，原结果完整保留',
        '几何重匹配不使用 Hb 预测成绩；它是同一公开人群的补充开发诊断。', balance,
        [('首轮的限制', '最近共享 donor 对应的非共享 donor 距离中位约为其1.45倍，不能单独分离硬件与距离。'),
         ('补充后的变化', f'全目标优势由 {comparison[0]["mean"]:.3f} 降为 {comparison[1]["mean"]:.3f}；严格匹配20目标为 {gain_c["mean"]:.3f}。'),
         ('不混写为独立重复', '训练／评价被试未更换；补充结果减少几何混杂，没有新增独立人群证据。')],
        '报告同时呈现原规则和修正规则，避免用更好看的子集覆盖原结果。',
        '两个运行的 hardware_topology.json、hardware_comparisons.csv；v2 合同')

    fig, ax = plt.subplots(figsize=(9, 5))
    real = cm[cm.arm.isin(['shared_optode', 'disjoint_distance_matched'])]
    pivot = real.pivot(index=['subject', 'id'], columns='arm', values=['nrmse', 'synchronous_OD_jump_correlation'])
    dx = pivot['synchronous_OD_jump_correlation']['shared_optode']-pivot['synchronous_OD_jump_correlation']['disjoint_distance_matched']
    dy = pivot['nrmse']['disjoint_distance_matched']-pivot['nrmse']['shared_optode']
    scatter = pd.DataFrame({'optical': dx, 'gain': dy}).groupby('subject').mean()
    ax.scatter(scatter.optical, scatter.gain, s=55, color='#087F8C')
    for subject, row in scatter.iterrows():
        ax.annotate(subject[-2:], (row.optical, row.gain), xytext=(4, 3), textcoords='offset points', fontsize=10)
    ax.axvline(0, color='gray', linestyle='--')
    ax.axhline(0, color='gray', linestyle='--')
    ax.set_xlabel('OD 突变同步相关：共享 − 非共享（每被试均值）')
    ax.set_ylabel('隐藏 Hb 预测改善：非共享 − 共享')
    od_fig = deck.figure(fig, '16_optical_context')
    deck.page('C｜光学共同变化提供背景，不能直接定源',
        '每点一名评价被试；先平均相同目标与窗口，再比较两类 donor。', od_fig,
        [('横轴的实际计算', '原始正光强取 −log，再计算0.1秒一步 OD 差分的双波长 RMS；比较目标／donor 突变幅度相关。'),
         ('它不是接触真值', '光学突变可能包含接触、运动、真实快速变化等；没有同步接触传感器或人工标签。'),
         ('本图只作诊断', '没有预定的来源判别检验；散点趋势不能证明共享预测增益由某一种硬件噪声造成。')],
        '结合原始测量域进行审查，仍须保留多种竞争解释。',
        'refinement/hardware_metrics.csv：synchronous_OD_jump_correlation')

    deck.page('D｜改变机制位置，保持推断条件对称',
        '使用现有非线性六状态 SSM；所有候选获得相同已知驱动与初态。', None,
        [('四种干预位置', 'τ（平均通过时间，秒）改变状态演化；无量纲 Hb 增益改变观测映射；额外 HbT 以 [0.65,0.35] 加到观测；时移改变 Hb 与驱动的时间对应。'),
         ('六组竞争，每组九点候选', '按预定波形、幅度、频谱代价选择非零干预对；匹配只使用生成真值，不使用之后的噪声或推断结果。'),
         ('正确／错误／缺失锚点', '锚点为有噪声的独立增益标定、时间标定及四段额外 HbT 测量；不包含 τ 或机制类别标签。错误锚点来自竞争机制。')],
        '这是已知驱动／初态的能力上界，不能直接迁移为真实数据来源辨识能力。',
        'mechanism_library、mechanism_worker；两个 resolved_config.yaml')

    deck.table('D｜白噪声天花板与慢相关噪声诊断',
        '每套64个训练身份定尺度，128个独立评价身份；每身份六对机制×两真机制×四锚点条件。',
        ['条件', '生成与推断', '无锚点错误率', '解释'],
        [['首轮白噪声', '边际 SD=0.20训练SD；独立时间噪声', f'{100*d0[d0.arm=="absent"].error.mean():.2f}%', '已达天花板，不能显示锚点必要性'],
         ['补充 AR(1)', 'ρ=0.95；相同边际SD；新的身份；精确白化', f'{100*mechanism_mean.loc["absent","error"]:.2f}%', '有效时间信息下降，检验敏感性'],
         ['推断评分', '观测马氏距离＋可用锚点的冻结高斯距离', '目标只用于评分', '每个候选面对相同噪声模型'],
         ['候选集合', '相对最优代价增幅 ≤3.841459', '报告经验覆盖率', '有限似然比集合，不叫生理后验']],
        [2.3, 5.1, 1.85, 2.85], '两套各6,144个评分单元；生成身份才是统计独立单元。',
        'mechanism_calibration.json、mechanism_metrics.csv；refinement 的对应文件')

    match0 = pd.read_csv(run/'mechanism_matching.csv')
    match1 = pd.read_csv(refinement/'mechanism_matching.csv')
    fig, axes = plt.subplots(1, 2, figsize=(9, 5))
    pairs = list(PAIR_LABELS)
    for frame, name, color in [(match0, '白噪声', '#7E8790'), (match1, '慢相关噪声', '#087F8C')]:
        metric = 'whitened_waveform_distance_SD' if 'whitened_waveform_distance_SD' in frame else 'waveform_distance_SD'
        vals = frame.groupby('pair')[metric].median().reindex(pairs)
        axes[0].plot(range(6), vals, marker='o', color=color, label=name)
        ratio = frame.groupby('pair').amplitude_log_ratio.median().reindex(pairs)
        axes[1].plot(range(6), ratio, marker='o', color=color, label=name)
    for ax in axes:
        ax.set_xticks(range(6), [PAIR_LABELS[p] for p in pairs], rotation=45, ha='right')
    axes[0].set_ylabel('白化波形 RMS 距离（训练 SD）')
    axes[1].set_ylabel('|log(干预 RMS1 / RMS2)|')
    axes[0].legend()
    match_fig = deck.figure(fig, '17_mechanism_matching')
    deck.page('D｜“最相似”不等于完全不可区分',
        '每个点为128个生成身份的中位数；幅度与频谱平衡参与匹配，但不是精确等式。', match_fig,
        [('匹配质量必须报告', '左图为考虑噪声协方差后的波形差；右图越接近零，干预幅度越相近。六种竞争难度并不相同。'),
         ('保留天花板结果', '白噪声首轮完全判对，反映固定驱动/初态、有限候选与充分时间信息的有利条件。'),
         ('没有制造来源后验', '训练尺度、噪声和候选均由生成合同规定；有限字典的分类表现不代表复杂真实来源已被辨识。')],
        '原始白噪声结果和补充慢噪声结果各自保留，不能择优汇总。',
        '两个运行的 mechanism_matching.csv；mechanism_library')

    dfig = grouped_bar(deck, d, 'pair', ['absent', 'correct', 'mismatched', 'missing_at_inference'],
        '18_mechanism_errors', value='error', ylabel='机制类别错误率（%；越低越好）', percent=True)
    deck.page('D｜错配锚点比缺失锚点更危险',
        '慢相关噪声；先在每生成身份内平均两种真机制，再跨128身份汇总。', dfig,
        [('总体错误率', f'无锚点 {100*mechanism_mean.loc["absent","error"]:.2f}%；正确 {100*mechanism_mean.loc["correct","error"]:.2f}%；错配 {100*mechanism_mean.loc["mismatched","error"]:.2f}%。'),
         ('正确锚点收益有上限', '无锚点已经很强，正确锚点只带来有限绝对改善；逐对比较还须查看 Holm 校正结果。'),
         ('缺失时如实退回', '缺失锚点与无锚点逐单元完全一致；没有以隐藏真值填补不存在的辅助输入。')],
        '锚点有价值的前提是其来源、时钟、标定和可用性正确；不能盲目提高监督权重。',
        'refinement/mechanism_metrics.csv、mechanism_comparisons.csv')

    fig, axes = plt.subplots(1, 2, figsize=(9, 5))
    regimes = ['absent', 'correct', 'mismatched', 'missing_at_inference']
    for ax, metric, title in zip(axes, ['component_error', 'false_component_projection'],
                                ['额外成分误差', '真实 τ 变化的错误分配']):
        for i, arm in enumerate(regimes):
            vals = d[d.arm == arm].groupby('subject')[metric].mean()
            row = block_interval(vals)
            ax.bar(i, row['mean'], color=COLORS[arm], yerr=[[row['mean']-row['low']], [row['high']-row['mean']]], capsize=4)
        ax.set_xticks(range(4), [ARM_LABELS[v] for v in regimes], rotation=30, ha='right')
        ax.set_title(title)
    axes[0].set_ylabel('成分 RMSE / 训练 Hb SD')
    axes[1].set_ylabel('估计额外项对真实变化的有符号投影')
    protection = deck.figure(fig, '19_teacher_protection')
    deck.page('D｜错误锚点会把动力学变化吸入额外项',
        'τ 干预的真实额外项为零；右图直接检验是否虚构一个额外成分。', protection,
        [('成分误差', f'无锚点 {mechanism_mean.loc["absent","component_error"]:.4f} SD，正确 {mechanism_mean.loc["correct","component_error"]:.4f}，错配 {mechanism_mean.loc["mismatched","component_error"]:.4f}。'),
         ('τ 的虚假归属', f'平均有符号投影：无锚点 {mechanism_mean.loc["absent","false_component_projection"]:.4f}，正确 {mechanism_mean.loc["correct","false_component_projection"]:.4f}，错配 {mechanism_mean.loc["mismatched","false_component_projection"]:.4f}。'),
         ('投影不是能量占比', '计算为 ⟨估计额外项,真实变化⟩ / ||真实变化||²，先按训练 Hb SD 缩放；允许负值，不是来源概率。')],
        '监督目标需要锚点有效性和竞争解释检查；仅重建准确并不足够。',
        'refinement/mechanism_metrics.csv：component_error、false_component_projection')

    fig, axes = plt.subplots(1, 2, figsize=(9, 5))
    for ax, metric, label in zip(axes, ['ambiguous', 'set_coverage'], ['多机制候选集合比例（%）', '真实机制在候选集合内（%）']):
        vals = []
        for arm in regimes:
            data = d[d.arm == arm].groupby('subject')[metric].mean()
            row = block_interval(data)
            vals.append(row)
        ax.bar(range(4), [r['mean']*100 for r in vals], color=[COLORS[v] for v in regimes],
            yerr=np.array([[100*(r['mean']-r['low']) for r in vals], [100*(r['high']-r['mean']) for r in vals]]), capsize=3)
        ax.set_xticks(range(4), [ARM_LABELS[v] for v in regimes], rotation=30, ha='right')
        ax.set_ylabel(label)
    axes[1].axhline(95, color='gray', linestyle='--', label='95% 仅作参考')
    axes[1].set_ylim(0, 105)
    axes[1].legend(fontsize=10)
    uncertainty = deck.figure(fig, '20_ambiguity_coverage')
    deck.page('D｜输出候选集合仍需检查经验覆盖',
        '候选得分距最优不超过3.841459即接纳；先合并同机制的参数候选。', uncertainty,
        [('两种覆盖不要混淆', '本页是“真机制是否在集合内”；精确生成参数的覆盖率另保存在表中，不能相互替代。'),
         ('错误锚点会排除真机制', '错配锚点的集合覆盖下降，说明即使允许多个解释，错误外部约束也可能给出错误而自信的集合。'),
         ('参考线不是资格线', '3.841459 没有在本实验中被校准成95%生理置信区间；虚线只便于读数。')],
        'teacher 应交付候选范围及其成立条件，而不是把工程集合直接称后验。',
        'refinement/mechanism_metrics.csv：ambiguous、set_coverage、parameter_set_coverage')

    # A source example of false attribution, selected by deterministic identity not maximal error.
    subset = d[(d.arm == 'mismatched') & (d.truth == 'tau') & (d.error == 1) & (d.repeat < 3)]
    if not subset.empty:
        r = subset.sort_values('id').iloc[0]
        archive = np.load(refinement/'mechanisms'/f'r{int(r["repeat"]):03d}.npz')
        pair_index = list(PAIR_LABELS).index(r.pair)
        prefix = f'p{pair_index}__tau'
        fig, axes = plt.subplots(3, 1, figsize=(9, 5), sharex=True)
        scale = archive['sd'][0]
        axes[0].plot(np.arange(120)/4, archive[prefix+'__observed'][:, 0]/scale, color='#263238', label='有噪观测 HbO')
        for arm in ('absent', 'correct', 'mismatched'):
            axes[0].plot(np.arange(120)/4, archive[prefix+'__'+arm+'__prediction'][:, 0]/scale, color=COLORS[arm], label=ARM_LABELS[arm])
            axes[1].plot(np.arange(120)/4, archive[prefix+'__'+arm+'__physical'][:, 0]/scale, color=COLORS[arm])
            axes[2].plot(np.arange(120)/4, archive[prefix+'__'+arm+'__component'][:, 0]/scale, color=COLORS[arm])
        axes[1].plot(np.arange(120)/4, archive[prefix+'__truth'][:, 0]/scale, '--', color='black', label='无噪真物理 HbO')
        axes[2].axhline(0, linestyle='--', color='black', label='真额外项＝0')
        for ax, label in zip(axes, ['完整预测', '物理项', '额外项']):
            ax.set_ylabel(label+'\n训练 SD')
        axes[0].legend(ncol=2, fontsize=9)
        axes[1].legend(fontsize=9)
        axes[2].legend(fontsize=9)
        axes[-1].set_xlabel('窗口时间（秒）')
        dcurve = deck.figure(fig, '21_false_attribution_example')
        deck.page('D｜拟合相近时，分量解释可以走错',
            f'例：生成身份 {int(r["repeat"])}，{PAIR_LABELS[r.pair]}；在保留曲线的前3身份中按ID取首个错配失败例。', dcurve,
            [('黑虚线是真值', '合成 τ 改变了物理响应；真实额外项为零。此处真值来自生成器，不是实测组织信号。'),
             ('错配锚点的影响', f'该例被判为 {r.predicted}，额外成分误差 {r.component_error:.3f} SD；有符号投影 {r.false_component_projection:.3f}。'),
             ('例子不代替分母', '本页用于读懂失败机制；总体结果仍以128身份、全部机制对及全部失败单元计算。')],
            '完整预测与成分归属必须分开评分。',
            f'refinement/mechanisms/r{int(r["repeat"]):03d}.npz、mechanism_metrics.csv')

    pass_error = dcomp[(dcomp.metric == 'error') & (dcomp.contrast == 'correct_vs_absent') & (dcomp.holm_p < .05)]
    pass_component = dcomp[(dcomp.metric == 'component_error') & (dcomp.contrast == 'correct_vs_absent') & (dcomp.holm_p < .05)]
    deck.table('统计上能确认什么，哪些仍不确定',
        '显著性来自配对被试／生成身份，而非图中的时间点数。',
        ['问题', '主要数值', '统计与解释'],
        [['B 电位预测', f'改善 {gain_b[0]["mean"]:.3f} / {gain_b[1]["mean"]:.3f}', '真实优于基线和两个 null；六项均通过 Holm'],
         ['B Hb／OD', '没有一致正面增量', '多个端点为负或区间跨零；不能升级来源'],
         ['C 距离匹配', f'{ci_text(gain_c)} NRMSE', '14人配对；仍有组织重叠竞争解释'],
         ['D 正确锚点', f'错误率 {100*mechanism_mean.loc["absent","error"]:.2f}% → {100*mechanism_mean.loc["correct","error"]:.2f}%', f'六机制对中分类改善 {len(pass_error)}/6、成分误差改善 {len(pass_component)}/6 通过既定 Holm 族'],
         ['D 错配锚点', f'错误率 {100*mechanism_mean.loc["mismatched","error"]:.2f}%', '只在生成模型与合成锚点假设内成立']],
        [2.25, 4.2, 5.65], '小的平均改善不应替代成对检验；负结果和不确定性也是本轮结论。',
        'event_comparisons.csv；refinement/hardware_comparisons.csv、mechanism_comparisons.csv')

    deck.table('本轮的解释等级与下游用途',
        '同一 token 的不同字段可以拥有不同证据等级。',
        ['对象', '本轮支持', '适合的后续用途', '不支持的命名'],
        [['HbT / Xρ', '明确的观测坐标、条件范围', '分量级监督与有效性字段', '绝对血氧／代谢率'],
         ['EOG 相关电位', '跨被试和配对 null 的预测证据', 'O 分支的辅助读出或测量上下文', '整条 O 等于眼电／噪声'],
         ['同步空间 Hb', '非共享 donor 的独立预测收益', '空间一致性约束与遗漏上下文', '浅层、全身或脑内的唯一来源'],
         ['合成机制标签', '限定模型内的条件辨别与失败', 'teacher 竞争解释、锚点缺失测试', '已验证的实测组织来源']],
        [2.3, 3.3, 3.7, 2.8], 'O 可以承载有生理意义的额外上下文；不应先强制它与任务标签无关。',
        '本报告对应 A/B/C/D 保留证据；用户计划 §五、§八')

    deck.table('E 未执行；下一步应补独立约束',
        '本轮缺少可访问的独立机制验证集，没有用现有 Hb 导出或合成观测替代。',
        ['尚未解决的问题', '最需要的观测／检验', '为什么现有结果不够'],
        [['共享硬件 vs 组织重叠', '局部短间距、光学接触／运动传感器', '几何匹配不等于敏感组织一致'],
         ['浅层 vs 脑内系统反应', '同步血压、呼气末CO₂及局部短间距', '组织位置与驱动来源是两个问题'],
         ['原生连续初态竞争', '冻结设计的连续记录 C-free/C-chain 实验', '本轮A只复用窗口级候选和profile'],
         ['学生推理缺锚点', '预先定义缺失机制与降级输出', '蒸馏不能消除观测不可辨识性']],
        [3.3, 4.6, 4.2], '建议先把可支持的电位读出和坐标范围接入 teacher 原型，再用新独立观测审查来源。',
        'summary.json：E；用户计划 §七 E；DATA_CONTRACT.md')

    bfailed = int(b.nrmse.isna().sum())
    cfailed = int(c.nrmse.isna().sum())
    deck.table('完整分母、软件检查与交付边界',
        '成功子集的均值不覆盖缺失和失败；本报告引用固定的两个实验运行。',
        ['项目', '计划／可用', '处理方式'],
        [['原生读取', f'{summary["measured"]["planned_records"]}／{summary["measured"]["available_records"]}记录；440／440窗', '按模态和通道保留支持；旧pilot仍在'],
         ['A 保留窗口', '72／72窗口；有限profile另报候选数', '单一候选不判稳定'],
         ['B 评分行', f'{len(b):,}行；非有限误差 {bfailed}', '相同评分时段；主要电位216/216窗'],
         ['C 补充评分行', f'{len(c):,}行；非有限误差 {cfailed}', '主要20目标2,240/2,240；null缺失独立报告'],
         ['D 评分单元', '两套各6,144／6,144；各128身份', '无锚点＝推理缺失逐单元核验']],
        [2.5, 4.15, 5.45], '未进行 protected evaluation、新 tokenizer 训练或额外独立生理采集。',
        'summary.json、verification.json；refinement/summary.json、verification.json')

    deck.page('证据、代码与复现入口',
        '报告是沟通材料；CSV、逐单元记录与冻结配置仍是结果的事实源。', None,
        [('主要运行', 'experiments/runs/physiology_semantic_tokenizer/observation_semantics/20261009_v2/\nA/B 与原始 C/D；resolved_config、来源、supervisor日志、逐记录数组和 verification。'),
         ('有界补充', '同一 suite 的 20261009_refinement_v1/\n几何匹配 C 与独立生成身份的慢相关噪声 D；引用主运行的 prepared 数据，不重读原生文件。'),
         ('软件入口', 'evaluate_observation_semantics.py --stage check / pilot / all / refinement / verify\n配置 observation_semantics_v1.yaml、v2.yaml；针对性测试 test_observation_semantics.py。')],
        '图片均以240 dpi PNG嵌入；PPT正文、表格和PDF文字保留可选择文本。',
        'experiments/README.md；docs/EXPERIMENT_PLAN.md；两个 launch.json')

    pptx_path = out/'observation_semantics_20261009.pptx'
    deck.prs.save(pptx_path)
    (out/'slide_sources.json').write_text(json.dumps(deck.sources, ensure_ascii=False, indent=2)+'\n')
    command = ['libreoffice', '-env:UserInstallation=file:///tmp/lo_observation_semantics_report', '--headless', '--convert-to', 'pdf', '--outdir', str(out), str(pptx_path)]
    process = subprocess.run(command, capture_output=True, text=True, timeout=180)
    (out/'pdf_export.log').write_text(process.stdout+'\n'+process.stderr)
    if process.returncode:
        raise RuntimeError('PDF export failed')
    pdf_path = pptx_path.with_suffix('.pdf')
    doc = fitz.open(pdf_path)
    validation = []
    for i, page in enumerate(doc):
        pix = page.get_pixmap(matrix=fitz.Matrix(1.25, 1.25))
        pix.save(out/'figures'/f'page_{i+1:02d}.png')
        image_count = len(page.get_images(full=True))
        # PowerPoint picture objects are the only scientific figures in the deck.
        expected_images = sum(shape.shape_type == 13 for shape in deck.prs.slides[i].shapes)
        blocks = page.get_text('dict')['blocks']
        out_of_bounds = []
        for block in blocks:
            if block['type'] != 0:
                continue
            rect = fitz.Rect(block['bbox'])
            if rect.x0 < -1 or rect.y0 < -1 or rect.x1 > page.rect.width+1 or rect.y1 > page.rect.height+1:
                out_of_bounds.append(list(rect))
        validation.append(dict(page=i+1, text_characters=len(page.get_text()), image_objects=image_count,
            expected_figure_images=expected_images, text_outside_page=out_of_bounds,
            passed=bool(len(page.get_text()) > 70 and image_count >= expected_images and not out_of_bounds)))
    # A compact visual contact sheet is a QA artifact, not a replacement for slide text.
    thumbs = []
    for i in range(len(doc)):
        im = Image.open(out/'figures'/f'page_{i+1:02d}.png').convert('RGB')
        im.thumbnail((640, 360))
        thumbs.append(im)
    for start in range(0, len(thumbs), 8):
        canvas = Image.new('RGB', (1280, 4*360), 'white')
        for j, im in enumerate(thumbs[start:start+8]):
            canvas.paste(im, ((j % 2)*640, (j//2)*360))
        canvas.save(out/f'contact_sheet_{start//8+1}.png')
    result = dict(status='passed' if all(r['passed'] for r in validation) and len(doc) == len(deck.prs.slides) else 'failed',
        slides=len(deck.prs.slides), pdf_pages=len(doc), scientific_png_figures=len(deck.figure_paths),
        source_runs=[str(run), str(refinement)], pptx=str(pptx_path), pdf=str(pdf_path),
        checks=validation, rendering='LibreOffice to PDF; PyMuPDF rendered page inspection; native WPS not available')
    (out/'presentation_validation.json').write_text(json.dumps(result, ensure_ascii=False, indent=2)+'\n')
    print(json.dumps({k: result[k] for k in ('status', 'slides', 'pdf_pages', 'scientific_png_figures', 'pptx', 'pdf')}, ensure_ascii=False))
    if result['status'] != 'passed':
        raise ValueError('report validation failed')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--run-dir', required=True, type=Path)
    parser.add_argument('--refinement-dir', required=True, type=Path)
    parser.add_argument('--output-dir', required=True, type=Path)
    args = parser.parse_args()
    build_report(args.run_dir.resolve(), args.refinement_dir.resolve(), args.output_dir.resolve())


if __name__ == '__main__':
    main()
