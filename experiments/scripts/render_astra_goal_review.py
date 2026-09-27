#!/usr/bin/env python3
"""Presentation of the retained 2026-09-26 Astra goal; never fit or load raw data.

This is a dated communication export. Research-state and individual run records
remain the owners of experiment status and results. Figures are PNG bitmaps;
slide titles, explanations, tables, and speaker notes remain editable text.
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.patches import FancyBboxPatch
import numpy as np
import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
RUNS = ROOT / "experiments/runs/physiology_semantic_tokenizer/shared_driver_reconstruction"
SESSION = "01a0dc92-38f4-7873-912f-56a07cc23240"
FONT = "Noto Sans CJK SC"
INK, MUTED, BG = "142C3D", "536777", "F4F7FA"
TEAL, BLUE, ORANGE, RED = "008D87", "3979B7", "D58C36", "C35459"
COLORS = ["#" + c for c in [BLUE, TEAL, ORANGE]]
COMP = ["EEG", "HbO", "HbR"]
RUN_NAMES = ["linear_screen", "nonlinear_replay", "nonlinear_fit", "gain_prior",
             "gain_prior_budget", "fixed_roi_tau", "fixed_roi_flow",
             "fixed_roi_initial", "optical_motion_audit", "fixed_roi_optical",
             "conditional_optical_gain"]
EVIDENCE = {f"E{i:02d}": RUNS / f"20260926_{n}_v1" for i, n in enumerate(RUN_NAMES, 1)}
EVIDENCE["E12"] = ROOT / "tests/test_shared_driver_step_control.py"


def run(name):
    return RUNS / f"20260926_{name}_v1"


def read_json(path):
    return json.loads(Path(path).read_text())


def metrics(name):
    return pd.read_csv(run(name) / "metrics.csv")


def select(frame, method, *, kind="measured", mode="full", pipeline=None):
    sub = frame[(frame.kind == kind) & (frame.method == method) & (frame["mode"] == mode)]
    if pipeline is not None:
        sub = sub[sub.optical_pipeline == pipeline]
    return sub


def score(part, success_only=True):
    scored = part[part.status == "completed"] if success_only else part
    columns = ["nmse_" + c for c in COMP]
    out = np.sqrt(scored.groupby(["subject", "session"])[columns].mean()
                  .groupby("subject").mean().mean()).to_numpy()
    passed = int((scored[["nrmse_" + c for c in COMP]] < .5).all(axis=1).sum())
    return out, len(scored), passed


def triple(values):
    return " / ".join(f"{v:.3f}" for v in values)


def curve(name, group, method, trial, mode="full"):
    path = run(name) / "cells" / f"{group}__{method}__{mode}" / "trajectories.npz"
    row = metrics(name)
    row = row[(row.group == group) & (row.method == method) &
              (row["mode"] == mode) & (row.trial == trial)]
    assert len(row) == 1 and row.iloc[0].status == "completed", (path, trial)
    with np.load(path, allow_pickle=False) as a:
        index = np.flatnonzero(a["trial_indices"] == trial)
        assert len(index) == 1
        i = int(index[0])
        out = {k: a[k][i].copy() for k in ["target", "prediction", "states"]}
        out["normalizer"] = a["normalizer"].copy()
        if "truth_states" in a:
            out["truth_states"] = a["truth_states"][i].copy()
    if mode == "full":
        computed = np.sqrt(np.mean((out["target"] - out["prediction"])**2, axis=0)) / out["normalizer"]
        np.testing.assert_allclose(computed, row.iloc[0][["nrmse_"+c for c in COMP]].to_numpy(float), rtol=2e-10)
    out["nrmse"] = row.iloc[0][["nrmse_"+c for c in COMP]].to_numpy(float)
    return out


def median_identity(name, method, pipeline=None):
    part = select(metrics(name), method, pipeline=pipeline).copy()
    part = part[part.status == "completed"]
    part["rank"] = part[["nrmse_"+c for c in COMP]].mean(axis=1)
    row = part.sort_values(["rank", "group", "trial"]).iloc[(len(part)-1)//2]
    return str(row.group), int(row.trial)


def configure_charts():
    font = Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc")
    font_manager.fontManager.addfont(str(font))
    plt.rcParams.update({"font.family": [font_manager.FontProperties(fname=str(font)).get_name(), "DejaVu Sans"],
                         "font.size": 12, "axes.unicode_minus": False,
                         "axes.spines.top": False, "axes.spines.right": False,
                         "axes.labelcolor": "#"+MUTED, "text.color": "#"+INK,
                         "xtick.color": "#"+MUTED, "ytick.color": "#"+MUTED,
                         "axes.edgecolor": "#BCC8D1", "savefig.facecolor": "white"})


def clean(ax):
    ax.grid(axis="y", color="#DFE6EC", lw=.7, zorder=0)
    ax.set_axisbelow(True)


def bars(ax, names, values, *, ylim=.9, threshold=True, ylabel="NRMSE"):
    x = np.arange(len(names))
    for j, c in enumerate(COMP):
        y = np.array(values)[:, j]
        b = ax.bar(x+(j-1)*.23, y, .21, color=COLORS[j], label=c, zorder=3)
        ax.bar_label(b, fmt="%.3f", fontsize=9, padding=3)
    if threshold:
        ax.axhline(.5, color="#"+RED, ls="--", lw=1.3)
        ax.text(.99, .5, " 目标 0.5", transform=ax.get_yaxis_transform(), ha="right", va="bottom", color="#"+RED, fontsize=10)
    ax.set(xticks=x, xticklabels=names, ylim=(0, ylim), ylabel=ylabel)
    clean(ax)


def savefig(fig, path):
    fig.savefig(path, dpi=240, bbox_inches="tight", pad_inches=.16)
    plt.close(fig)
    return path


def make_figures(out):
    dest = out / "figures"
    dest.mkdir(parents=True, exist_ok=True)
    figpaths = {}
    def save(fig, key):
        figpaths[key] = savefig(fig, dest / (key+".png"))
    configure_charts()

    # Exact model topology, rendered as one bitmap diagram.
    fig, ax = plt.subplots(figsize=(11.8, 4.5))
    ax.set(xlim=(0, 12), ylim=(0, 5)); ax.axis("off")
    def box(x,y,w,h,text,color):
        ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle="round,pad=.06,rounding_size=.1",
                                   facecolor=color,edgecolor="none"))
        ax.text(x+w/2,y+h/2,text,ha="center",va="center",fontsize=15,color="white",linespacing=1.6)
    def arrow(x1,y1,x2,y2):
        ax.annotate("",(x2,y2),(x1,y1),arrowprops=dict(arrowstyle="->",color="#7C939F",lw=2))
    box(.15,1.7,1.75,1.2,"共享驱动\nr(t)","#"+INK)
    box(2.7,1.7,1.7,1.2,"血管信号 s\n血流 f","#"+TEAL)
    box(5.3,1.7,2.3,1.2,"血容量 v\n总 Hb p · 脱氧 q","#"+TEAL)
    box(8.6,1.7,2.8,1.2,"光学观测映射\nHbO / HbR","#"+BLUE)
    box(2.7,3.7,2.5,.75,"EEG 功率特征","#"+BLUE)
    arrow(1.95,2.3,2.65,2.3); arrow(4.45,2.3,5.25,2.3); arrow(7.65,2.3,8.55,2.3)
    arrow(1.1,2.95,2.65,4.05)
    ax.text(2.34,1.28,"β：有效驱动增益",ha="center",fontsize=12)
    ax.text(6.43,1.28,"τ：通过时间常数",ha="center",fontsize=12)
    ax.text(9.98,1.28,"幅度尺度 / 混合",ha="center",fontsize=12)
    ax.text(6,.35,"只有一个逐时刻驱动；血流状态由方程与五个初态产生",ha="center",fontsize=15)
    save(fig,"model")

    fig, axes = plt.subplots(1,2,figsize=(11.8,4.5),gridspec_kw={"width_ratios":[1.7,1]})
    linear=metrics("linear_screen"); names=["固定 τ","训练 τ","同目标 oracle"]
    vals=[score(select(linear,m),False)[0] for m in ["fixed_tau","fitted_tau","oracle_lower_bound"]]
    bars(axes[0],names,vals,ylim=.94);axes[0].set_title("线性筛查：有限残差不等于有效状态",pad=17)
    axes[0].legend(loc="upper right",fontsize=10)
    rp=select(metrics("nonlinear_replay"),"fitted_tau");rp=rp[rp.status=="completed"]
    pre=np.sqrt(rp.groupby(["subject","session"])[["linear_nmse_HbR"]].mean().groupby("subject").mean().mean()).iloc[0]
    post=score(rp)[0][2]
    b=axes[1].bar(["线性", "非线性重放"],[pre,post],color=["#"+TEAL,"#"+ORANGE],width=.55)
    axes[1].bar_label(b,fmt="%.3f",padding=5);axes[1].axhline(.5,c="#"+RED,ls="--")
    axes[1].set(ylim=(0,.75),ylabel="HbR NRMSE",title="同一 68 条成功轨迹")
    clean(axes[1]);fig.tight_layout(w_pad=3);save(fig,"linear_replay")

    # All curves are redrawn from retained arrays, with no refit/alignment.
    def curves_plot(key,name,group,trial,methods,labels,*,states=False):
        items=[curve(name,group,m,trial) for m in methods]
        for a in items[1:]:
            np.testing.assert_array_equal(a['target'],items[0]['target'])
            np.testing.assert_array_equal(a['normalizer'],items[0]['normalizer'])
        rows=4 if states else 3
        fig,axes=plt.subplots(rows,1,figsize=(8.4,5.6 if states else 5.1),sharex=True)
        t=np.arange(items[0]['target'].shape[0])*.25
        for j in range(3):
            axes[j].plot(t,items[0]['target'][:,j]/items[0]['normalizer'][j],c="#243B4A",lw=1.75,label="观测目标")
            for k,a in enumerate(items):
                axes[j].plot(t,a['prediction'][:,j]/a['normalizer'][j],c=["#9AA9B2","#"+TEAL,"#"+ORANGE][k],
                             lw=1.5,ls=["--","-",":"][k],label=labels[k])
            axes[j].set_ylabel(COMP[j]+" / SD",fontsize=11);clean(axes[j])
            axes[j].text(.99,.94,"NRMSE " + " / ".join(f"{a['nrmse'][j]:.3f}" for a in items),
                         transform=axes[j].transAxes,ha="right",va="top",fontsize=9,
                         bbox=dict(facecolor="white",edgecolor="none",alpha=.8))
        if states:
            for j,label in zip([2,3,4,5],["f 血流","v 容量","p 总Hb","q 脱氧"]):
                axes[-1].plot(t,items[-1]['states'][:,j],lw=1.3,label=label)
            axes[-1].axhline(1,c="#BCC8D1",ls="--",lw=1)
            axes[-1].set_ylabel("相对状态",fontsize=11);axes[-1].legend(ncol=4,fontsize=8.5,loc="upper right");clean(axes[-1])
        axes[0].legend(ncol=len(items)+1,fontsize=9,loc="lower center",bbox_to_anchor=(.5,1.02))
        axes[-1].set_xlabel("窗内时间（秒）")
        fig.tight_layout(h_pad=.4);save(fig,key)
        return items

    g,t=median_identity("nonlinear_fit","fixed_tau_100")
    curves_plot("smoothing_curves","nonlinear_fit",g,t,["fixed_tau_100","fixed_tau_0p01"],["强平滑","弱平滑"])
    figpaths["smoothing_identity"]=(g,t)

    fig,axes=plt.subplots(1,2,figsize=(11.6,4.5))
    # Known paired synthetic result from the terminal report.
    axes[0].bar(["训练 β\n无幅度先验","训练 β\n有幅度先验"],[.113,.464],color=["#"+TEAL,"#"+ORANGE],width=.55)
    for i,v in enumerate([.113,.464]):axes[0].text(i,v+.015,f"{v:.3f}",ha="center")
    axes[0].set(ylim=(0,.57),ylabel="真实驱动 NRMSE",title="合成：先验引入恢复偏差");clean(axes[0])
    gain_rows=metrics('gain_prior')
    a=select(gain_rows,'fixed_gain_no_amplitude',mode='center_EEG')
    a=a[a.status=='completed']
    b=select(gain_rows,'fixed_gain_amplitude',mode='center_EEG')
    b=b[b.status=='completed']
    common=a[['group','trial']].merge(b[['group','trial']])
    own=pd.read_csv(run('gain_prior')/'report_v1/own_context_rows.csv')
    own=own[(own.kind=='measured')&(own['mode']=='center_EEG')].merge(common,on=['group','trial'])
    compared=[frame.merge(common,on=['group','trial']) for frame in [a,b]]+[own]
    hidden_values=[float(np.sqrt(frame.groupby(['subject','session']).nmse_EEG.mean().groupby('subject').mean().mean())) for frame in compared]
    assert all(len(frame)==71 for frame in compared)
    axes[1].bar(["无幅度先验","有幅度先验","自身插值"],hidden_values,
                color=["#"+ORANGE,"#"+TEAL,"#9AA9B2"],width=.55)
    for i,v in enumerate(hidden_values):axes[1].text(i,v+.13,f"{v:.3f}",ha="center")
    axes[1].set(ylim=(0,6.2),ylabel="隐藏 EEG NRMSE",title="实测：压住振荡，仅接近插值")
    clean(axes[1]);fig.tight_layout(w_pad=3);save(fig,"gain_tradeoff")

    fig,ax=plt.subplots(figsize=(8.5,4.2))
    x=np.arange(2);a=ax.bar(x-.17,[30,18],.32,color="#9AA9B2",label="原预算");b=ax.bar(x+.17,[41,60],.32,color="#"+TEAL,label="热启动延长")
    ax.bar_label(a,padding=4);ax.bar_label(b,padding=4)
    ax.set(xticks=x,xticklabels=["训练增益","训练增益 + 幅度先验"],ylim=(0,82),ylabel="full 成功数 / 72")
    ax.axhline(72,c="#BCC8D1",ls="--");ax.legend();clean(ax);fig.tight_layout();save(fig,"budget")

    fig,axes=plt.subplots(1,2,figsize=(11.8,4.4))
    d=metrics("fixed_roi_tau");vals=[score(select(d,m)) for m in ['fixed_roi_fixed_tau','fixed_roi_trained_tau','fixed_roi_prior_tau']]
    bars(axes[0],["固定 τ","训练 τ","训练 τ + 先验"],[s[0] for s in vals],ylim=.64)
    axes[0].set_title("成功子集误差均低于 0.5",pad=17);axes[0].legend(fontsize=9)
    b=axes[1].bar(["固定 τ","训练 τ","训练 τ + 先验"],[s[1] for s in vals],color=["#9AA9B2","#"+TEAL,"#"+ORANGE],width=.55)
    axes[1].bar_label(b,padding=4);axes[1].set(ylim=(0,82),ylabel="成功数 / 72",title="但训练 τ 后，覆盖率下降")
    clean(axes[1]);fig.tight_layout(w_pad=2.5);save(fig,"fixed_roi")

    flow=metrics("fixed_roi_flow")
    fm=['fixed_roi_trained_tau','trained_tau_logflow01','trained_tau_logflow1','trained_tau_logflow4']
    fig,ax=plt.subplots(figsize=(8.5,4.4)); x=np.arange(4)
    success=[score(select(flow,m))[1] for m in fm]; passed=[score(select(flow,m))[2] for m in fm]
    b=ax.bar(x-.18,success,.34,color="#BCD6DC",label="拟合成功");ax.bar_label(b,padding=3)
    b=ax.bar(x+.18,passed,.34,color="#"+TEAL,label="逐条三项 < 0.5");ax.bar_label(b,padding=3)
    ax.set(xticks=x,xticklabels=["w=0","w=0.1","w=1","w=4"],ylim=(0,85),ylabel="轨迹数 / 72")
    ax.axhline(72,c="#BCC8D1",ls="--");ax.legend(loc="upper left",ncol=2,fontsize=10);clean(ax);fig.tight_layout();save(fig,"flow_coverage")
    g,t=median_identity("fixed_roi_flow","fixed_tau_logflow1")
    curves_plot("flow_curves","fixed_roi_flow",g,t,['fixed_tau_logflow1','trained_tau_logflow1'],['固定 τ · w=1','训练 τ · w=1'],states=True)
    figpaths['flow_identity']=(g,t)

    choices=pd.read_csv(run('fixed_roi_flow')/'report_v2/tau_training_selection.csv')
    fig,axes=plt.subplots(1,3,figsize=(11.7,4.2),sharey=True)
    for ax,sub in zip(axes,['subject_01','subject_09','subject_18']):
        for method,label,color in zip(fm,['w=0','w=0.1','w=1','w=4'],['#AFBCC5','#77A3CA','#'+TEAL,'#'+ORANGE]):
            p=choices[(choices.kind=='measured')&(choices.subject==sub)&(choices.method==method)&(choices.status=='completed')].sort_values('outer')
            ax.plot(p.outer,p.tau,'o-',lw=1.5,ms=5,label=label,color=color)
        ax.set(title=sub.replace('subject_','S'),xticks=range(4),xlabel='训练折（原索引）',ylim=(2.2,5.2));clean(ax)
    axes[0].set_ylabel('训练 τ（秒）');axes[2].legend(fontsize=10,ncol=2)
    fig.tight_layout();save(fig,'tau_folds')

    init=pd.read_csv(run('fixed_roi_initial')/'synthetic_report_v1/tau_relative_error_by_initial_truth.csv')
    fig,ax=plt.subplots(figsize=(8.5,4.2));x=np.arange(3)
    for j,(m,label,color) in enumerate([('trained_tau_logflow1','自由初态','#'+TEAL),('trained_tau_logflow1_tied','绑定 p₀=v₀','#'+ORANGE)]):
        p=init[init.method==m].sort_values('true_initial_log_p_over_v')
        b=ax.bar(x+(j-.5)*.3,p['max']*100,.28,label=label,color=color);ax.bar_label(b,fmt='%.2f%%',padding=3,fontsize=10)
    ax.set(xticks=x,xticklabels=['真实 log(p₀/v₀) = −0.05','真实 p₀=v₀','真实 log(p₀/v₀) = +0.05'],ylabel='τ 最大相对误差（%）',ylim=(0,8));ax.tick_params(axis='x',labelsize=10)
    ax.legend();clean(ax);fig.tight_layout();save(fig,'initial')

    motion=run('optical_motion_audit')
    with np.load(motion/'measured_traces.npz',allow_pickle=False) as a:
        arrays={k:a[k].copy() for k in a.files}
    md=read_json(motion/'measured_summary.json')
    subjects=np.array([row['identity']['subject'] for row in md['rows']])
    # Plot each subject's median; every retained window participates.
    fig,axes=plt.subplots(1,3,figsize=(11.8,4.2),sharey=True)
    t=np.arange(300)/10-5
    for ax,sub in zip(axes,['subject_01','subject_09','subject_18']):
        for p,label,color in [('current','旧运动处理','#'+RED),('no_motion','无运动校正','#'+TEAL),('mne_tddr','MNE TDDR','#'+BLUE)]:
            vals=arrays[p+'__hb_native'][subjects==sub,:,0]
            ax.plot(t,np.median(vals,axis=0),lw=2,label=label,color=color)
        ax.axvline(0,c='#BCC8D1',ls='--');ax.set(title=sub.replace('subject_','S')+' · 24 窗口中位曲线',xlabel='事件相对时间（秒）');clean(ax)
    axes[0].set_ylabel('滤波前 HbO（工程相对单位）');axes[2].legend(fontsize=9)
    fig.tight_layout(w_pad=1.3);save(fig,'motion_measured')

    fig,axes=plt.subplots(1,2,figsize=(11.8,4.1))
    with np.load(motion/'synthetic_traces.npz',allow_pickle=False) as a:
        for p,label,color in [('no_motion','原信号','#'+INK),('current','旧运动处理','#'+RED),('mne_tddr','MNE TDDR','#'+TEAL)]:
            y=a['asymmetric_positive__'+p+'__od'][:,0];axes[0].plot(np.arange(len(y))/10,y,label=label,color=color,lw=1.7)
    axes[0].set(title='无净漂移的非对称周期输入',xlabel='时间（秒）',ylabel='760 nm 光密度 OD');axes[0].legend(fontsize=9);clean(axes[0])
    hrf=pd.read_csv(run('fixed_roi_optical')/'report_v1/background_paired_HRF_gain.csv')
    hrf=hrf[(hrf.context_s==30)&hrf.background.isin(['white only','heart_resp_mayer_white/high'])]
    for j,c in enumerate(['gain_760','gain_850']):
        b=axes[1].bar(np.arange(2)+(j-.5)*.3,hrf[c],.28,label=c[-3:]+' nm',color=COLORS[j]);axes[1].bar_label(b,fmt='%.2f',padding=3)
    axes[1].axhline(1,c='#BCC8D1',ls='--');axes[1].set(xticks=[0,1],xticklabels=['安静背景','含心搏 / 呼吸 / 低频背景'],ylim=(0,1.22),ylabel='已知慢响应的配对增益',title='MNE 对响应的保留依赖背景');axes[1].tick_params(axis='x',labelsize=10)
    axes[1].legend(fontsize=9);clean(axes[1]);fig.tight_layout(w_pad=3);save(fig,'motion_synthetic')

    cond=metrics('conditional_optical_gain');methods=['fixed_tau_logflow1','conditional_fixed_gain','conditional_trained_gain']
    fig,axes=plt.subplots(1,2,figsize=(11.8,4.4),sharey=True)
    for ax,p,title in zip(axes,['no_motion','mne_tddr'],['无运动校正','MNE TDDR']):
        vals=[score(select(cond,m,pipeline=p)) for m in methods]
        names=[f'{chr(65+i)}\n成功 {v[1]}/72\n达标 {v[2]}/72' for i,v in enumerate(vals)]
        bars(ax,names,[v[0] for v in vals],ylim=.63);ax.set_title(title,pad=18)
    axes[0].legend(fontsize=9,loc='upper left',ncol=3);fig.tight_layout(w_pad=2);save(fig,'conditional_scores')

    selections=pd.read_csv(run('conditional_optical_gain')/'report_v1/training_selections.csv')
    syn=selections[selections.kind=='synthetic']
    fig,axes=plt.subplots(1,2,figsize=(11.6,4.3))
    axes[0].plot([.4,2.1],[.4,2.1],ls='--',c='#A4B4C0',label='理想恢复')
    axes[0].scatter(syn.true_g,syn.g,s=55,color='#'+TEAL,alpha=.75,zorder=4)
    axes[0].set(xlabel='增益真值 g',ylabel='训练估计 g',xticks=[.5,1,2],title='12 组合成：增益恢复准确');clean(axes[0])
    vals=[score(select(cond,m,kind='synthetic'))[0] for m in methods]
    bars(axes[1],['A 单位映射','B 固定增益','C 训练增益'],vals,ylim=.17,threshold=False)
    axes[1].set_title('观测重建误差很低');axes[1].legend(fontsize=9)
    fig.tight_layout(w_pad=2.8);save(fig,'conditional_synthetic')

    syn_group='synthetic_g1_slow_r0';syn_items=[curve('conditional_optical_gain',syn_group,m,0) for m in methods]
    fig,axes=plt.subplots(1,2,figsize=(11.8,4.1))
    for ax,j,yl in zip(axes,[0,2],['驱动 r（原坐标）','相对血流 f']):
        tt=np.arange(120)*.25
        ax.plot(tt,syn_items[-1]['truth_states'][:,j],color='#'+INK,lw=2,label='合成真值')
        for k,a in enumerate(syn_items):
            ax.plot(tt,a['states'][:,j],c=['#9AA9B2','#'+TEAL,'#'+ORANGE][k],ls=['--','-',':'][k],lw=1.5,label=chr(65+k))
        ax.set(xlabel='窗内时间（秒）',ylabel=yl);clean(ax)
    axes[0].set_title('动态形状接近，整体水平仍可偏移');axes[1].set_title('低观测误差允许不同内部状态')
    axes[0].legend(ncol=4,fontsize=9);fig.tight_layout(w_pad=2.7);save(fig,'latent_truth')

    g,t=median_identity('conditional_optical_gain',methods[0],pipeline='mne_tddr')
    curves_plot('conditional_median','conditional_optical_gain',g,t,methods,['A 单位映射','B 条件固定 β','C 条件训练 β'])
    figpaths['conditional_identity']=(g,t)
    curves_plot('conditional_failure','conditional_optical_gain','mne_tddr__subject_09_o3',3,methods,['A 单位映射','B 条件固定 β','C 条件训练 β'])

    fig,axes=plt.subplots(1,2,figsize=(11.8,4.3))
    state=pd.read_csv(run('conditional_optical_gain')/'report_v1/full_state_ranges_and_recovery.csv')
    for j,(p,pname) in enumerate([('no_motion','无校正'),('mne_tddr','MNE')]):
        for k,m in enumerate(methods):
            s=state[(state.kind=='measured')&(state.pipeline==p)&(state.method==m)&(state.status=='completed')]
            lo,hi=s.f_min.min(),s.f_max.max();x=j*4+k
            axes[0].plot([x,x],[lo,hi],color=["#9AA9B2","#"+TEAL,"#"+ORANGE][k],lw=8,solid_capstyle='round')
            axes[0].text(x,hi+.06,f'{hi:.2f}',ha='center',fontsize=9)
            axes[0].text(x,lo-.08,f'{lo:.2f}',ha='center',va='top',fontsize=9)
    axes[0].axhline(1,c='#A4B4C0',ls='--');axes[0].set(xticks=[0,1,2,4,5,6],xticklabels=['无校正\nA','B','C','MNE\nA','B','C'],ylabel='成功轨迹血流 f 的全时程范围',title='条件映射使状态幅度缩小');clean(axes[0])
    for p,ls,label in [('no_motion','--','无校正'),('mne_tddr','-','MNE')]:
        s=selections[(selections.kind=='measured')&(selections.pipeline==p)&(selections.status=='completed')]
        for sub,col in zip(['subject_01','subject_09','subject_18'],COLORS):
            ss=s[s.subject==sub].sort_values('fold');axes[1].plot(ss.fold,ss.beta,marker='o',ls=ls,color=col,label=sub.replace('subject_','S')+' '+label)
    axes[1].set(xticks=range(4),xlabel='训练折（原索引）',ylabel='有效增益 β',title='同一被试的 β 仍会跨折剧变');axes[1].legend(ncol=2,fontsize=8.5);clean(axes[1]);fig.tight_layout(w_pad=2.5);save(fig,'states_gain')

    own=pd.read_csv(run('fixed_roi_flow')/'report_v2/paired_own_context.csv')
    own=own[(own.kind=='measured')&(own.method=='trained_tau_logflow1')]
    vals=[];base=[]
    for c,mode in [('EEG','center_EEG'),('HbO','center_fNIRS'),('HbR','center_fNIRS')]:
        row=own[own['mode']==mode].iloc[0];vals.append(row['candidate_NRMSE_'+c]);base.append(row['own_context_NRMSE_'+c])
    fig,axes=plt.subplots(1,2,figsize=(11.8,4.1),gridspec_kw={'width_ratios':[1.4,1]})
    x=np.arange(3)
    for j,(y,label,color) in enumerate([(vals,'共享生理模型','#'+ORANGE),(base,'同模态自身上下文','#9AA9B2')]):
        b=axes[0].bar(x+(j-.5)*.32,y,.3,label=label,color=color);axes[0].bar_label(b,fmt='%.3f',padding=4,fontsize=10)
    axes[0].set(xticks=x,xticklabels=COMP,ylim=(0,4.9),ylabel='隐藏分量 NRMSE',title='旧目标 · flow w=1 的同身份比较');axes[0].legend(fontsize=9);clean(axes[0])
    latest=[score(select(cond,'conditional_trained_gain',mode=mode,pipeline='mne_tddr'))[0][i] for i,mode in enumerate(['center_EEG','center_fNIRS','center_fNIRS'])]
    b=axes[1].bar(COMP,latest,color=COLORS,width=.55);axes[1].bar_label(b,fmt='%.3f',padding=4,fontsize=10)
    axes[1].set(ylim=(0,4.9),ylabel='隐藏分量 NRMSE',title='最终条件增益 C · MNE');clean(axes[1]);fig.tight_layout(w_pad=2.5);save(fig,'hidden')

    fig,axes=plt.subplots(1,2,figsize=(11.8,4.1))
    path=run('conditional_optical_gain')/'training/no_motion__subject_01_o1__conditional_trained_gain/start_0/result.json'
    trace=pd.DataFrame(read_json(path)['trace']);tail=trace.tail(26)
    axes[0].plot(tail.iteration,tail.parameter_value,'o-',ms=3,lw=1.2,c='#'+ORANGE)
    axes[0].set(title='保留的失败训练：增益两侧往复振荡',xlabel='迭代',ylabel='有效增益 β');clean(axes[0])
    axes[1].semilogy(trace.iteration,trace.projected_scaled_gradient_inf_norm,c='#'+BLUE,lw=1.8)
    axes[1].axhline(1e-4,c='#'+RED,ls='--',label='原梯度阈值 10⁻⁴')
    axes[1].set(title='目标下降后，仍未满足原成功条件',xlabel='迭代',ylabel='投影缩放梯度 ∞ 范数');axes[1].legend(fontsize=9);clean(axes[1]);fig.tight_layout(w_pad=3);save(fig,'solver')
    return figpaths


class Deck:
    def __init__(self):
        self.prs=Presentation();self.prs.slide_width=Inches(13.333333);self.prs.slide_height=Inches(7.5)
        self.prs.core_properties.title='2026-09-26 Astra Goal：生理语义 SSM 开发复盘'
        self.prs.core_properties.subject='11组实验、拟合曲线、参数与观测问题、暂停时的开发边界'
        self.prs.core_properties.author='项目开发复盘'
        self.slide_index=[]

    def rect(self,s,x,y,w,h,fill,line=None,radius=False):
        shape=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                 Inches(x),Inches(y),Inches(w),Inches(h))
        shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(fill)
        if line:shape.line.color.rgb=RGBColor.from_string(line)
        else:shape.line.fill.background()
        if radius:
            shape.adjustments[0]=.12
        return shape

    def text(self,s,txt,x,y,w,h,size=20,color=INK,bold=False,align=None):
        box=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
        tf=box.text_frame;tf.clear();tf.word_wrap=True
        tf.margin_left=tf.margin_right=Inches(.01);tf.margin_top=tf.margin_bottom=Inches(.01)
        for i,line in enumerate(txt.split('\n')):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
            p.text=line;p.font.name=FONT;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
            p.space_after=Pt(size*.22);p.line_spacing=1.17
            if align is not None:p.alignment=align
        return box

    def new(self,title,kicker,sources=(),note='',dark=False):
        s=self.prs.slides.add_slide(self.prs.slide_layouts[6])
        s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(INK if dark else BG)
        n=len(self.prs.slides);col='FFFFFF' if dark else INK
        self.text(s,kicker.upper(),.48,.25,11,.28,size=10,color='79D5CE' if dark else TEAL,bold=True)
        self.text(s,title,.48,.70,12.3,.95,size=29,color=col,bold=True)
        self.rect(s,.5,6.99,12.3,.012,'365266' if dark else 'D9E2E8')
        footer='2026.09.26  ASTRA GOAL  /  '+(' · '.join(sources) if sources else '开发复盘')
        self.text(s,footer,.5,7.13,11.7,.20,size=8.5,color='AFC3D0' if dark else MUTED)
        self.text(s,f'{n:02d}',12.2,7.08,.6,.3,size=11,color='AFC3D0' if dark else MUTED,align=PP_ALIGN.RIGHT)
        source_notes='\n'.join(f'{key}: {EVIDENCE[key]}' for key in sources if key in EVIDENCE)
        s.notes_slide.notes_text_frame.text=(
            f'复盘对象：{SESSION}，gpt-6-astra，2026-09-26 15:16–22:23（Asia/Shanghai）。\n'
            '本页是保留证据的沟通摘要，不是新的实验状态或授权记录。所有实验均为原开发面板；本次导出未拟合模型。\n'
            +note+'\n\n证据路径：\n'+source_notes)
        self.slide_index.append(dict(slide=n,title=title,evidence=list(sources)))
        return s

    def image(self,s,path,x,y,w,h):
        with Image.open(path) as im:iw,ih=im.size
        scale=min(w/iw,h/ih);ww,hh=iw*scale,ih*scale
        s.shapes.add_picture(str(path),Inches(x+(w-ww)/2),Inches(y+(h-hh)/2),width=Inches(ww),height=Inches(hh))

    def callout(self,s,heading,body,x=9.25,y=1.8,w=3.5,h=1.45,color=TEAL):
        self.rect(s,x,y,w,h,'FFFFFF',radius=True)
        self.rect(s,x,y,.05,h,color)
        heading_size=18.5 if len(heading)>10 else 21
        self.text(s,heading,x+.18,y+.14,w-.35,.55,size=heading_size,color=color,bold=True)
        self.text(s,body,x+.18,y+.76,w-.35,h-.8,size=16.5,color=INK)

    def band(self,s,txt,y=6.23,color=TEAL):
        self.rect(s,.5,y,12.3,.56,'E4F2F0' if color==TEAL else 'F8EBE7',radius=True)
        self.text(s,txt,.68,y+.085,11.95,.38,size=16.5,color=color,bold=True)

    def table(self,s,headers,rows,x=.55,y=1.9,w=12.2,h=3.5,widths=None,size=16):
        shape=s.shapes.add_table(len(rows)+1,len(headers),Inches(x),Inches(y),Inches(w),Inches(h))
        table=shape.table
        if widths:
            for col,ratio in zip(table.columns,widths):col.width=Inches(w*ratio/sum(widths))
        for r,values in enumerate([headers]+rows):
            for c,value in enumerate(values):
                cell=table.cell(r,c);cell.text=str(value);cell.margin_left=Inches(.13);cell.margin_right=Inches(.09)
                cell.margin_top=Inches(.06);cell.margin_bottom=Inches(.04);cell.vertical_anchor=MSO_ANCHOR.MIDDLE
                cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string(INK if r==0 else ('FFFFFF' if r%2 else 'EDF2F6'))
                for p in cell.text_frame.paragraphs:
                    p.font.name=FONT;p.font.size=Pt(size if r else size-1);p.font.bold=(r==0)
                    p.font.color.rgb=RGBColor.from_string('FFFFFF' if r==0 else INK)
        return shape


def build_slides(figs):
    d=Deck()
    s=d.new('从“曲线能拟合”走向“状态可解释”','ASTRA GOAL / 一天的生理语义 SSM 开发',
            ['E07','E09','E11'],dark=True,
            note='目标会话通过本地只读 thread_goals 与 threads 关联确认：唯一匹配的项目 Astra goal，time_used_seconds=25381。创建15:16:05，最后更新22:23:43；最终状态paused。会话文本确认联合重建为主、遮挡预测独立检验。')
    d.text(s,'2026 年 9 月 26 日\n15:16 — 22:23 · 约 7 小时持续开发',.55,1.83,7.6,1.25,size=25,color='E6EEF3')
    d.text(s,'11 组主要实验  ·  72 个实测验证身份\n合成恢复、曲线核验、参数诊断与观测管道追溯',.55,3.33,11.9,1,size=22,color='BFD2DD')
    for x,num,label in [(.55,'< 0.5','三路汇总误差已可达到'),(4.75,'58 / 72','最终 MNE 条件增益逐条达标'),(8.95,'未证实','稳定的个体生理参数')]:
        d.rect(s,x,4.83,3.83,1.43,'213F51',radius=True)
        d.text(s,num,x+.20,5.03,3.45,.53,size=31,color='75D5CB',bold=True)
        d.text(s,label,x+.20,5.73,3.45,.32,size=15,color='E6EEF3')
    d.text(s,'复盘生成：2026-09-27  |  结果截点：原对话暂停时',.58,6.55,11,.25,size=11,color='9AB4C2')

    s=d.new('低误差已做到；三项目标还没有同时成立','先看结果',['E03','E07','E11'])
    d.table(s,['用户目标','昨天最强的证据','截至暂停时的判断'],[
        ['NRMSE < 0.5','最终 MNE / C：0.293 / 0.386 / 0.330\nEEG / HbO / HbR；72/72 成功','三项汇总已达标\n逐条 58/72 同时达标'],
        ['曲线形状合理','EEG 快速变化改善；多数 Hb 慢趋势可跟随','局部双峰、深谷、恢复时序仍失配'],
        ['参数有生理意义\n并反映被试差异','合成可恢复 τ 与增益；同折多起点一致','实测增益跨折 >5 倍\n绝对尺度与驱动基线仍弱识别']
    ],y=1.9,h=3.65,widths=[2.0,5.2,4.7],size=18)
    d.band(s,'结论：完成了重建能力与失效机制的推进，尚未完成生理语义资格验证。')

    s=d.new('开发沿着四个问题推进，形成 11 组主要实验','一天的路线图',['E01','E03','E07','E09','E11','E12'])
    chapters=[('15:20–16:26','模型是否有表达能力？','① 线性筛查\n② 非线性重放\n③ 非线性重拟合',TEAL),
              ('16:27–18:53','参数为何不稳定？','④ 增益 × 驱动幅度先验\n⑤ 预算热启动\n⑥ 固定 ROI 与连续 τ',BLUE),
              ('18:53–20:29','状态为何走向边界？','⑦ 血流软约束\n⑧ 初态绑定（合成）\n⑨ 光学运动处理归因',ORANGE),
              ('20:29–22:23','观测尺度是否合理？','⑩ none / MNE 观测对照\n⑪ 条件映射与共享增益\n末段：步长控制软件验证',TEAL)]
    for i,(tm,q,body,col) in enumerate(chapters):
        x=.55+i*3.14;d.rect(s,x,1.97,2.95,3.79,'FFFFFF',radius=True)
        d.text(s,tm,x+.16,2.13,2.63,.4,size=18,color=col,bold=True)
        d.text(s,q,x+.16,2.82,2.62,1,size=22,bold=True)
        d.text(s,body,x+.16,4.02,2.62,1.46,size=16.5)
    d.band(s,'时间为北京时间，分支有交叠；停止的分支与未启动的实验同样保留在复盘中。')

    s=d.new('所有主比较围绕同一批 72 个验证身份','如何读后面的数字',['E01','E06','E10','E11'],
            note='同一三被试S01/S09/S18，三个session，每人24 trial；四个outer fold，每折18训练、6验证。相同身份会在多方法/遮挡模式/版本中重复出现，因此评价行数不能当成独立样本数。固定ROI、光学路径变化后，Hb目标和训练SD重新生成；只能在相同目标版本内做配对比较。')
    d.table(s,['比较对象','统一规则'],[
        ['3 被试 × 24 窗口 = 72','S01 / S09 / S18；3 会话、4 折；每折 18 训练 + 6 验证'],
        ['主验收：联合重建','EEG 为 log-power/PCA 特征；HbO/HbR 为相对观测量'],
        ['评分尺度','各分量使用训练折 SD；先 trial → session → subject 等权平均 NMSE，再开根'],
        ['成功分母','每臂保留 72 身份；并列“拟合成功数”“三项逐条 <0.5 数”'],
        ['独立诊断','中心特征遮挡、自身上下文、多起点、真值恢复；示例沿用原 fold / trial 索引']
    ],h=3.65,widths=[3,9],size=17)
    d.band(s,'更换 ROI / 光学管道会改变目标曲线与 SD；跨版本分数不能直接作为同任务排名。',color=RED)

    s=d.new('新增自由度放在驱动、初态和共享参数上','模型设计',['E01','E03','E04','E06','E11'])
    d.image(s,figs['model'],.5,1.55,12.25,4.5)
    d.band(s,'没有给每个血流状态逐点自由修正；本轮是受约束轨迹拟合，尚非完整随机 SSM 后验。')

    s=d.new('①② 线性低误差提示潜力，非线性重放暴露偏差','表达能力 / 线性筛查 + 固定驱动重放',['E01','E02'],
            note='线性12组合成+12组实测，3456评价行含上下文和遮挡/移位控制。训练tau数学有效68/72，小信号有效0/72；oracle小信号有效0/72。重放576条：569完成，7域失败。右图线性HbR与非线性HbR使用共同成功68条身份；没有重拟合。')
    d.image(s,figs['linear_replay'],.52,1.64,12.2,3.85)
    d.text(s,'0 / 72',.67,5.63,2,.45,size=26,color=RED,bold=True)
    d.text(s,'训练 τ 轨迹通过小信号范围检查',2.49,5.68,7.8,.38,size=19)
    d.band(s,'oracle 同目标选参仅是乐观诊断；释放 τ 尚未解决 EEG 过平滑和生理域问题。')

    g,t=figs['smoothing_identity']
    s=d.new('③ 减弱驱动平滑，找回 EEG 的快速变化','非线性重拟合 / 4 臂 × 合成与实测',['E03'],
            note=f'曲线身份：{g}, trial={t}，按固定tau强平滑基线成功full的平均NRMSE排序取中位，其他方法共享身份。图示误差从保留数组独立复算。固定tau强/弱正则共同70条的EEG NRMSE为0.804→0.331；全成功子集弱正则为0.333/0.375/0.352。合成864/864，实测832/864评价收敛。')
    d.image(s,figs['smoothing_curves'],.52,1.65,8.43,4.36)
    d.callout(s,'0.804 → 0.331','同一 70 条成功轨迹\nEEG NRMSE',y=1.86,h=1.68)
    d.callout(s,'代价仍在状态内部','固定 τ 弱平滑：\n最小相对血流 0.049\n29/71 轨迹状态偏离 >50%',y=3.85,h=2.02,color=ORANGE)
    d.band(s,'弱平滑固定 τ：58/72 逐条达标；训练选择 τ：67/72，但仍有 4 条数值失败。')

    s=d.new('④ 增益 × 幅度先验：四臂分开检验两种自由度','约束与参数的权衡',['E04'])
    data=metrics('gain_prior');methods=['fixed_gain_no_amplitude','fixed_gain_amplitude','trained_gain_no_amplitude','trained_gain_amplitude']
    rows=[]
    for label,m in zip(['A 固定 β','B 固定 β + 幅度先验','C 训练 β','D 训练 β + 幅度先验'],methods):
        v,n,k=score(select(data,m));rows.append([label,f'{n}/72',triple(v),f'{k}/72'])
    d.table(s,['方法','成功数','EEG / HbO / HbR NRMSE','三项逐条达标'],rows,y=1.91,h=3.08,widths=[3.6,1.3,4.3,1.9],size=18)
    d.text(s,'训练难点：24 个共享参数训练组仅 8 个成功。\n72 起点中 50 个预算耗尽：35 个无域拒绝，15 个有域拒绝。',.7,5.23,11.7,.8,size=19)
    d.band(s,'有幅度先验可以压住隐藏段振荡，但尚未可靠提高完整面板的拟合与参数可信度。')

    s=d.new('幅度先验提升稳定性，也会压缩真实驱动','④ 的关键负结果',['E04'],
            note='左侧为完整合成full的driver NRMSE算术均值，未事后对齐。右侧为实测固定beta两臂共同成功71条隐藏EEG，A=5.069，B=1.203，自身插值=1.210。B全部72条为1.200，对应插值1.205；本页严格使用共同71条重算，不能混作同一子集。增益偏高18–33%源于有幅度先验的合成D臂。')
    d.image(s,figs['gain_tradeoff'],.52,1.7,12.2,3.9)
    d.text(s,'合成增益偏高 18%–33%；无先验训练增益的误差 <1%。',.7,5.70,11.8,.4,size=20,color=ORANGE,bold=True)
    d.band(s,'更平稳的曲线可以来自先验收缩；它并不自动意味着更准确的生理轨迹。')

    s=d.new('⑤ 延长预算救回慢收敛，但没有修复全部问题','优化归因 / 35 个保留失败起点热启动',['E05'],
            note='每合格起点增加最多90迭代/3600次试次评估，目标、数据、梯度阈值不变。35起点新增22收敛、13失败。可用训练组8→17/24。A/B和864合成行继承；C/D共同成功身份ΔNMSE全部0，收益来自覆盖。没有将继承结果计为独立重复。')
    d.image(s,figs['budget'],.55,1.88,8.3,3.95)
    d.callout(s,'8 → 17 / 24','可用共享参数训练组\n新增 22 个收敛起点',y=1.92,h=1.76)
    d.callout(s,'13 个仍失败','包括慢尾、域边界与增益漂移；\n异常血流范围约 0.018–2.55',y=3.92,h=1.93,color=ORANGE)
    d.band(s,'原本共同成功的预测完全不变；预算扩展的收益主要是覆盖增加。')

    s=d.new('⑥ 固定 Hb 位置，排除通道变化造成的混杂','固定 AF7Fp1 / 连续共享 τ',['E06'],
            note='重新为固定ROI生成训练尺度；不与此前变通道的Hb分数做同目标差值。合成72/72训练起点、648/648验证成功，tau真值1/2/4恢复0.991–1.006/1.995–2.004/3.982–4.028。实测30/72起点、10/24训练组收敛，42失败均有域拒绝。')
    d.image(s,figs['fixed_roi'],.55,1.67,12.15,3.85)
    d.text(s,'合成 τ=1 / 2 / 4 秒准确恢复；实测却有 42/72 训练起点触发域拒绝。',.67,5.70,12,.4,size=18.5)
    d.band(s,'S09 的 τ 训练全部失败；较低的成功子集均值不能代表更好的整体模型。',color=RED)

    s=d.new('⑦ 血流软约束，使共享 τ 恢复完整拟合覆盖','log-flow 权重 0 / 0.1 / 1 / 4 的敏感性',['E07'],
            note='惩罚为sqrt(w*dt)*log(f)/ln2，是工程软锚，不是正常人群范围。六臂2592行含864继承行，合成1296完成，实测1114完成。右侧固定和训练tau比较同为w=1、相同72身份全成功；0→1原29共同成功身份误差略增，主要收益是救回失败。')
    d.image(s,figs['flow_coverage'],.55,1.84,8.38,4.0)
    d.callout(s,'72 / 72 全部收敛','w=1，共享训练 τ\n三项逐条达标 64/72',y=1.89,h=1.8)
    d.callout(s,'τ 自由度有拟合收益','同 w=1 固定 → 训练 τ\n0.330 / 0.423 / 0.380\n→ 0.299 / 0.356 / 0.335',y=3.95,h=1.92)
    d.band(s,'这是旧预处理目标上的进步；后续发现目标含算法漂移，故不能作为生理恢复证明。',color=RED)

    g,t=figs['flow_identity']
    s=d.new('旧目标上的典型曲线：慢趋势能跟随，局部仍失配','⑦ / 原评分坐标除以训练 SD',['E07'],
            note=f'固定tau/w1基线成功full的中位身份：{g}, trial={t}。曲线无平移、无幅度或时间对齐；状态图为共享tau/w1。全72图册见report_v2/FULL_CURVE_ATLAS.pdf。')
    d.image(s,figs['flow_curves'],.55,1.65,8.35,4.52)
    d.callout(s,'典型样例','S01 · fold 3 · trial 3\n按固定 τ 基线中位选取',y=1.93,h=1.72)
    d.callout(s,'保留曲线与状态两条证据','全体成功轨迹血流范围\n0.321–1.702\n静息值 1 不是已证实的正常界',y=3.91,h=2.0,color=ORANGE)
    d.band(s,'图示 EEG 为功率特征，Hb 为相对量；图形贴合与绝对生理幅度是不同问题。')

    s=d.new('参数能稳定求出，仍可能依赖约束强度','⑦ / 跨折 τ 与合成偏差',['E07'])
    d.image(s,figs['tau_folds'],.55,1.65,12.15,3.43)
    d.table(s,['合成约束 w','0','0.1','1','4'],[
        ['τ 平均绝对相对误差','0.349%','0.357%','0.821%','2.659%'],
        ['驱动 NRMSE','0.109','0.108','0.155','0.254']
    ],y=5.16,h=.93,widths=[4,2,2,2,2],size=14)
    d.band(s,'多起点一致 ≠ 个体参数已识别；S09 的 τ 随约束改变，重叠训练折也不是独立重复。')

    s=d.new('⑧ 绑定 p₀=v₀ 不是无条件成立的修复','初态机制 / 仅完成合成阶段',['E08','E07'],
            note='36组初态压力条件，216/216起点和2592/2592验证行成功。图为各真初态层tau最大相对误差，非均值。p0=v0匹配层自由/绑定平均约0.84%。模型方程来源审计说明当前inlet-balance变体与所引文献在独立p0/v0时不等价；只澄清来源，未改正在运行的方程。后续发现旧光学目标漂移，registry将此分支记stopped，原数值manifest保留synthetic_terminal；实测未执行。')
    d.image(s,figs['initial'],.55,1.83,8.4,4.05)
    d.callout(s,'匹配真值：约 0.84%','自由 / 绑定的 τ 平均误差接近；\n均可正确恢复',y=1.90,h=1.86)
    d.callout(s,'错设时：最高约 6.18%','绑定把初态不匹配转移到 τ，\nHb 重建也变差',y=4.00,h=1.82,color=ORANGE)
    d.band(s,'旧目标实测阶段停止；同时澄清 HbT 方程来源及 p₀−v₀ 衰减项的解释边界。')

    s=d.new('⑨ 关键转折：HbO 的主要慢下降来自旧运动处理','光学管道的同窗口归因核查',['E09'],
            note='仅保留三被试同72个AF7Fp1窗口，OD/原生Hb/4Hz Hb共216/216复现通过，最大差9.10e-15，不是全部bitexact。下方比例为逐窗口算法干预差/current下降量的被试内中位：99.91/97.25/96.35%；不等于生理解释率。关闭旧算法后晚期15–25秒减早期0–5秒为负的HbO窗口72→47/72。图为未最终滤波的Hb_native，无重新缩放。')
    d.image(s,figs['motion_measured'],.5,1.7,12.25,3.55)
    d.table(s,['复现核验','下降窗口数','算法差值 / 原下降量（被试中位）'],[
        ['216 / 216 通过','旧处理 72/72 → 无校正 47/72','S01 99.91% · S09 97.25% · S18 96.35%']
    ],y=5.35,h=.76,widths=[2.2,3.3,6.7],size=15)
    d.band(s,'这确认了算法干预造成主要下降；无校正曲线与 MNE 曲线都不能直接称为真实生理。',color=RED)

    s=d.new('合成控制证实旧算法漂移，也限定了 MNE 的适用解释','⑨⑩ / 零信号、周期、趋势、已知响应、运动与背景对照',['E09','E10'],
            note='左图读取asymmetric_positive合成OD保留数组：幅值0.01无净漂移波形经旧算法产生约1.45漂移。右图读取background_paired_HRF_gain.csv的30秒条件，三个seed，white only与heart_resp_mayer_white/high；760/850增益0.0966/0.2545与0.8231/0.8334。该增益是所构造控制中的响应保留指标，不是实测生理效能。')
    d.image(s,figs['motion_synthetic'],.53,1.8,12.18,3.9)
    d.text(s,'因此采用版本化 v4 接口，同时保留 none 与 MNE 两条路径；旧缓存保持原样。',.68,5.77,11.8,.4,size=18)
    d.band(s,'处理结果有背景与时间上下文依赖，不能用“更平滑”或“分数更低”来选真值。')

    s=d.new('⑩ 换到两条新光学路径后，学习 τ 的优势减弱','同一管道内比较固定 / 训练 τ',['E10'])
    data=metrics('fixed_roi_optical');rows=[]
    for pipeline,plabel in [('no_motion','无运动校正'),('mne_tddr','MNE TDDR')]:
        for method,mlabel in [('fixed_tau_logflow1','固定 τ'),('trained_tau_logflow1','训练 τ')]:
            v,n,k=score(select(data,method,pipeline=pipeline));rows.append([plabel,mlabel,f'{n}/72',triple(v),f'{k}/72'])
    d.table(s,['光学路径','τ 方案','成功','EEG / HbO / HbR','逐条达标'],rows,y=1.86,h=2.85,widths=[2.2,1.5,1.1,4.2,1.5],size=18)
    d.text(s,'两条路径各有 59 个共同成功身份：学习 τ 仅小幅改善 HbO，EEG 略差。\nS09 两折共享 τ 训练失败；MNE 下 S01 三折触及 τ 下界 0.5 秒。',.68,5.01,11.9,.93,size=19)
    d.band(s,'初态 HbO 接近零引发极小步长；正性坐标已完成软件验证，尚无实测救回结论。',color=ORANGE)

    s=d.new('⑪ 条件光学映射：把状态幅度假设写到观测层','固定目标 / 固定评分 SD / 三臂对照',['E11'],
            note='71μM是皮层源总Hb基线假设；还假设路径3cm、DPF6、组织敏感度1。B=k inv(Eold)[ln(10)*1e-6*Edecadic*71]只进入预测均值，不再变换target和SD。βref=1/sqrt(det(B))。A→B同时改变映射、β参考以及物理状态惩罚相对贡献，不是纯光学校准效应。B→C才隔离训练增益。三臂tau固定2秒，仍自由五初态与flow权重1。')
    items=[('A','单位 Hb 映射','β = 1','继承原基线及失败',MUTED),('B','条件 2×2 光学映射','β = βref','工程幅度参考',TEAL),('C','相同条件光学映射','训练共享 β','18 条训练后，验证冻结',ORANGE)]
    for i,(tag,head,sub,body,col) in enumerate(items):
        x=.6+i*4.2;d.rect(s,x,1.9,3.97,3.1,'FFFFFF',radius=True)
        d.text(s,tag,x+.18,2.08,.6,.7,size=36,color=col,bold=True)
        d.text(s,head,x+.18,3.02,3.6,.65,size=23,bold=True)
        d.text(s,sub,x+.18,3.82,3.6,.45,size=22,color=col)
        d.text(s,body,x+.18,4.50,3.6,.33,size=16)
    d.text(s,'71 μM、DPF=6、3 cm 光程与组织敏感度=1，均为明确的条件假设。',.7,5.44,11.9,.52,size=20)
    d.band(s,'B → C 隔离训练增益的作用；A → B 同时改变映射与参考增益，不是纯校准效应。')

    s=d.new('条件映射下，合成增益可以准确、稳定恢复','⑪ / 新生成 12 组合成真值',['E11'],
            note='g真值0.5/1/2×slow/mixed×2重复；36/36训练起点、648/648验证成功；12/12增益筛查通过。最大参数相对误差1.11%，多起点最大跨度为真值的0.0026%。不是继承旧真值族，也不是个体物理标定。')
    d.image(s,figs['conditional_synthetic'],.52,1.7,12.2,3.9)
    d.text(s,'最大增益误差 1.11%  ·  36/36 起点成功  ·  648/648 验证成功',.7,5.75,11.9,.43,size=22,color=TEAL,bold=True)
    d.band(s,'合成结果确认该条件模型可恢复增益；仍需检验驱动和状态本身。')

    s=d.new('低观测误差之下，驱动整体水平仍难确定','⑪ / 原坐标潜在轨迹与合成真值',['E11','E07'],
            note='预定合成身份synthetic_g1_slow_r0/trial0，直接读取保留truth_states和拟合states，没有去均值、符号、幅度或时间对齐。C完整合成full driver NRMSE平均约0.730。早期stationary_baseline_null_check表明不同平衡血流水平经基线处理可给出几乎相同零观测。曲线是选定机制示例，不替代全合成汇总。')
    d.image(s,figs['latent_truth'],.55,1.82,12.15,3.75)
    d.text(s,'完整合成的原坐标 driver 平均 NRMSE ≈ 0.730',.7,5.70,11.9,.45,size=22,color=ORANGE,bold=True)
    d.band(s,'去基线后，驱动常量与初态 / 血流水平可以相互补偿；不能以去均值后的好形状替代原误差。',color=ORANGE)

    s=d.new('实测条件增益改善 Hb 拟合，仍未全面改善三路','⑪ / 两条管道各自保留完整 72 分母',['E11'],
            note='success与逐条达标读取原metrics；A/B/C目标、训练SD完全相同。NRMSE为各自成功子集，成功集合不同时不直接归因。B→C共同身份：MNE72条平均ΔNMSE=EEG+0.0265,HbO−0.0147,HbR−0.0172；none共同66条方向相同。实测63/72训练起点收敛，23/24训练组可用；none S01 fold1的三起点全失败导致18模式评价失败。')
    d.image(s,figs['conditional_scores'],.52,1.65,12.2,4.23)
    d.band(s,'MNE：A → C 达标 51→58；B → C 达标仍 58，Hb 变好但 EEG 变差。')

    g,t=figs['conditional_identity']
    s=d.new('典型样例：主要峰谷能够恢复','⑪ / 按 A 基线成功窗口的中位误差选样',['E11'],
            note=f'MNE固定单位映射基线full成功窗口平均分量NRMSE稳定排序取中位：{g}/trial{t}。A/B/C均展示相同身份。坐标为原目标/冻结训练SD，无幅度或相位校正。')
    d.image(s,figs['conditional_median'],.52,1.68,8.42,4.38)
    d.callout(s,'S01 · fold 3 · trial 15','三个方法共享同一目标、\n同一训练评分尺度',y=1.88,h=1.88)
    d.callout(s,'图上可见的进步','EEG 快速起伏可跟随；\nHbO/HbR 的主要慢趋势\n具有一致的时间结构',y=4.01,h=1.92)
    d.band(s,'单例用于解释曲线形状；总体判断仍以全部 72 身份及完整图册为依据。')

    s=d.new('反例：HbO 双峰与末端谷值仍没有恢复','⑪ / 已记录的困难样例，不替换、不隐藏',['E11'],
            note='身份mne_tddr__subject_09_o3/trial3。C的HbO NRMSE≈0.937；同身份no_motion C≈0.993。该样例由原完整图册识别为明显形态失败，不是典型中位，也不将其视为全体均值。')
    d.image(s,figs['conditional_failure'],.52,1.68,8.42,4.38)
    d.callout(s,'S09 · fold 3 · trial 3','C 的 HbO NRMSE\nMNE 0.937 / 无校正 0.993',y=1.88,h=1.88,color=RED)
    d.callout(s,'残差带有结构','双峰、末端深谷未跟随；\n状态幅度更小，仍不足以\n证明曲线形状已验收',y=4.01,h=1.92,color=ORANGE)
    d.band(s,'汇总 NRMSE <0.5 不能替代困难样例的形态检查。',color=RED)

    s=d.new('状态幅度变温和，参数仍缺少跨折稳定性','⑪ / 幅度与可重复性必须分开判断',['E11'])
    d.image(s,figs['states_gain'],.53,1.75,12.2,3.93)
    d.text(s,'C 的血流范围：无校正 0.844–1.076；MNE 0.900–1.054。',.7,5.75,11.85,.38,size=19)
    d.band(s,'同折成功起点 β 最大跨度 0.159%，但 S09 跨折变化 5.81 / 5.42 倍；尚非稳定被试特征。',color=ORANGE)

    s=d.new('独立遮挡检验持续提醒：重建好不等于可预测','跨方向保留的共同负结果',['E07','E11'],
            note='左图只比较旧flow目标上trained_tau_logflow1与自身上下文的同成功身份：EEG69条，Hb72条。右图为最终MNE/C各72成功身份，没有把两张图跨目标配对。遮挡发生在离线处理特征层，不是原始传感器遮挡或未来预测。')
    d.image(s,figs['hidden'],.53,1.8,12.2,3.82)
    d.text(s,'左：EEG 使用 69 个共同成功身份，Hb 使用 72 个；右：三项各 72 个。',.7,5.70,11.8,.42,size=17.5)
    d.band(s,'联合重建是用户确认的主目标；遮挡是独立诊断，尚未支持共享状态的预测优势。')

    s=d.new('暂停前：定位数值振荡，步长控制完成软件验证','求解器开发 / 未完成下一轮合成与实测对照',['E12','E11'],
            note='历史最后一条消息22:23:35报告非零残差振荡软件例：旧方法40次迭代不收敛，新插值步4次达到同梯度阈值。tests/test_shared_driver_step_control.py覆盖目标、预算、物理域、默认路径一致性。计划合同shared_driver_conditional_step_control_v1.yaml仅准备；pilot目录仅resource_probe.json，无launch/manifest/结果，不称已启动完整pilot或已获得实测改进。正性Hb初态候选只完成软件/合成核验，未实测救回。')
    d.image(s,figs['solver'],.53,1.68,12.2,3.65)
    d.text(s,'软件振荡例：旧方法 40 次未收敛 → 新步控制 4 次达到相同标准。\n实际模型的完整对照尚未启动；不能把软件例的加速外推为实测成果。',.7,5.45,11.9,.68,size=18)
    d.band(s,'另外保留 Hb 正性初态坐标候选；保持原目标与收敛标准，尚无实测修复结论。')

    s=d.new('留下的不只是分数，还有可复用的实现与诊断能力','软件与证据交付',['E03','E07','E09','E11','E12'],
            note='实现入口src/inference/shared_driver_reconstruction.py；实验入口experiments/scripts/evaluate_shared_driver_reconstruction.py；呈现入口render_shared_driver_reconstruction.py；光学审计analyze_fnirs_motion_correction.py；预处理src/data/homer2_preprocessing.py；约束/状态解读src/inference/t3a_balloon_robust_ssm.py。历史针对性测试批次会重叠，不相加为独立测试总数。三组全72曲线图册分别flow、optical、conditional_gain。')
    d.table(s,['交付','具体能力'],[
        ['共享驱动求解器','线性筛查、非线性受约束拟合、精确离散敏感度、共享 τ / β 训练'],
        ['约束与数值诊断','驱动幅度、log-flow、初态绑定、正性坐标、预算热启动、可选步长控制'],
        ['观测层修正','固定 ROI；v4 显式 none / MNE；条件光学均值映射；旧版本可复现'],
        ['验证工具','合成真值恢复、遮挡与移位、全失败分母、多起点、4/8 子步积分核查'],
        ['完整证据包','版本合同、冻结源码、持续服务日志、逐组终态、曲线图册与结果审计']
    ],y=1.92,h=3.7,widths=[2.8,9.3],size=18)
    d.band(s,'已有代码仍在工作树；本次只制作复盘，没有重新训练、扩大数据范围或执行受保护评价。')

    s=d.new('下一步应围绕剩余证据缺口，而不是继续堆叠自由度','结论与建议',['E07','E09','E11','E12'])
    d.table(s,['已经解决 / 明确','尚未解决','建议的下一步'],[
        ['模型具备低误差重建能力','困难 Hb 波形的峰谷与时序','固定失败身份，检查残差结构'],
        ['部分失败源于预算 / 域 / 步长','实际模型的新步长策略效果','按已准备合同完成同预算对照'],
        ['旧运动算法确实引入漂移','真实慢响应的处理保持能力','保留多背景 / 连续上下文控制'],
        ['τ / β 在合成中可以恢复','实测光学尺度、基线和参数互相补偿','先检验 β–τ 联合恢复与可辨识性']
    ],y=1.88,h=3.55,widths=[3.6,4.1,4.3],size=18)
    d.text(s,'总体判断：目标取得实质进展，但“低误差 + 合理形状 + 稳定生理参数”尚未同时获得支持。',.7,5.73,11.9,.85,size=22,color=TEAL,bold=True)

    s=d.new('实验清单 1 / 2：模型、先验与数值优化','附录 / 所有数字均为评价行，不是新增独立样本',['E01','E02','E03','E04','E05','E06'])
    d.table(s,['编号 / 方向','合成证据','实测证据','主要结论'],[
        ['E01 线性筛查','12 组 / 1728 行','12 组 / 1728 行','EEG 未达标；训练 τ 小信号有效 0/72'],
        ['E02 非线性重放','288/288 成功','281/288 成功','7 域失败；HbR 误差被线性近似低估'],
        ['E03 非线性拟合','864/864 成功','832/864 成功','弱平滑显著改善 EEG；状态仍有极端值'],
        ['E04 增益 × 幅度先验','864/864 成功','574/864 成功','抑振伴随收缩偏差；训练覆盖低'],
        ['E05 预算热启动','864 行继承','734/864 成功','35 起点救回 22；共同成功预测不变'],
        ['E06 固定 ROI / τ','648/648 成功','379/648 成功','排除位置混杂；共享 τ full 仅29/72成功']
    ],y=1.87,h=4.11,widths=[3.1,2.4,2.4,5.2],size=14.5)
    d.band(s,'同一 trial 在方法、模式和版本中重复出现；每组都先完成软件 / 合成检查，再按合同进入实测。')

    s=d.new('实验清单 2 / 2：状态约束与观测层','附录 / 每页备注含精确证据路径',['E07','E08','E09','E10','E11','E12'],
            note='E07合成1296+实测1296，合计2592，含864继承行。E08合成2592全部新算，旧目标实测未跑。E09为72窗口×3路径只读处理归因，不是拟合。E10合成432继承，实测864新算。E11合成648新算，实测1296包括A基线432继承与B/C新算864。E12仅软件验证和资源探测，暂无完整pilot/实测结果。报告/图册为沟通资产，原实验单元终态和registry仍是事实源。')
    d.table(s,['编号 / 方向','合成证据','实测 / 状态','主要结论'],[
        ['E07 血流软约束','1296/1296 成功','1114/1296 成功','w=1 full72/72；三项逐条64/72'],
        ['E08 初态绑定','2592/2592 成功','停止于合成；实测未跑','真值不匹配时 τ 偏差增大'],
        ['E09 光学运动归因','零 / 趋势 / 周期等控制','72 窗口；216/216 复现','旧算法引入主要 HbO 下降'],
        ['E10 光学路径对照','432 行继承 + 光学控制','786/864 成功','共享 τ 收益弱，失败与触界保留'],
        ['E11 条件映射 / 增益','648/648 成功','1272/1296 成功','MNE C逐条58/72；β跨折不稳定'],
        ['E12 数值候选','软件 / 小型合成核验','新完整对照未启动','正性坐标 + 步长控制待实测复核']
    ],y=1.86,h=4.12,widths=[3.15,2.55,2.8,4.6],size=14.3)
    d.band(s,'各 E 编号对应原 run / 软件源；完整曲线、配置与逐条结果保留在原实验目录。')
    return d


def verify_evidence():
    """Targeted, read-only consistency checks; no access to raw or protected data."""
    registry=read_json(ROOT/'research_state/registry.json')
    records={v['entity']:v for v in registry['records'] if 'entity' in v}
    assert records['main.shared_driver_fixed_roi_initial']['execution']=='stopped'
    for name in RUN_NAMES:
        p=run(name);manifest=read_json(p/'manifest.json')
        if name=='fixed_roi_initial':
            assert manifest['execution']=='synthetic_terminal'
        else:
            assert manifest['execution']=='completed',name
        if name=='optical_motion_audit':
            assert read_json(p/'measured_summary.json')['windows']==72
            continue
        data=metrics(name)
        subset=data[data.kind=='measured']
        if name!='fixed_roi_initial':
            assert len(subset[['subject','trial']].drop_duplicates())==72
        else:
            assert subset.empty
    cd=metrics('conditional_optical_gain')
    v,mne_success,mne_pass=score(select(cd,'conditional_trained_gain',pipeline='mne_tddr'))
    np.testing.assert_allclose(v,[.292583,.386454,.330479],atol=1e-6)
    assert mne_success==72 and mne_pass==58
    f,n,k=score(select(metrics('fixed_roi_flow'),'trained_tau_logflow1'))
    np.testing.assert_allclose(f,[.29934,.35567,.33474],atol=1e-5)
    assert n==72 and k==64
    pilot=RUNS/'20260926_conditional_step_control_pilot_v1'
    assert not (pilot/'manifest.json').exists(), 'New step-control evidence exists; review dated scope before re-export.'
    return dict(primary_mne_nrmse=v.tolist(),primary_mne_success=mne_success,primary_mne_pass=mne_pass,
                flow_nrmse=f.tolist(),flow_success=72,flow_pass=k,
                measured_unique_identities=72,main_experiment_directions=11)


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output',type=Path,required=True,help='New versioned export directory')
    args=parser.parse_args();out=args.output.resolve()
    if out.exists() and any(out.iterdir()):
        raise SystemExit('Use a new empty versioned export directory; retained exports are not overwritten.')
    verification=verify_evidence();out.mkdir(parents=True,exist_ok=True)
    figs=make_figures(out);deck=build_slides(figs)
    path=out/'20260926_Astra_SSM开发复盘_v1.pptx'
    deck.prs.save(path)
    for slide in deck.prs.slides:
        for shape in slide.shapes:
            assert shape.left>=0 and shape.top>=0
            assert shape.left+shape.width<=deck.prs.slide_width+10
            assert shape.top+shape.height<=deck.prs.slide_height+10
    audit=dict(session=SESSION,evidence_date='2026-09-26',created_date='2026-09-27',
               slide_count=len(deck.prs.slides),figures_dpi=240,raw_data_access=False,new_fits=False,
               wps_checked=False,verification=verification,slides=deck.slide_index,
               selection={k:v for k,v in figs.items() if k.endswith('_identity')})
    (out/'export_validation.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(dict(pptx=str(path),slides=len(deck.prs.slides),figures=sum(isinstance(v,Path) for v in figs.values())),ensure_ascii=False))


if __name__=='__main__':
    main()
