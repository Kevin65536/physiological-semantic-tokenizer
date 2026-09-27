#!/usr/bin/env python3
"""Render retained shared-driver linear diagnostic evidence; never fit models."""
from __future__ import annotations

import argparse
import json
from pathlib import Path
import re

import fitz
import markdown
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
import numpy as np
import pandas as pd


COMPONENTS = ["EEG", "HbO", "HbR"]
METHODS = ["fixed_tau", "fitted_tau", "source_prior_tau", "oracle_lower_bound", "own_context"]
LABELS = {"fixed_tau": "固定 τ", "fitted_tau": "训练拟合 τ", "source_prior_tau": "来源先验 τ", "oracle_lower_bound": "全目标 oracle 下界", "own_context": "自身上下文"}


def truth(series):
    return series.astype(str).str.lower().isin(["true", "1", "1.0"])


def table(frame):
    def cell(value):
        if isinstance(value, (float, np.floating)):
            return f"{value:.4f}" if np.isfinite(value) else "—"
        return str(value).replace("|", "/")
    return "| " + " | ".join(frame.columns) + " |\n| " + " | ".join(["---"] * len(frame.columns)) + " |\n" + "\n".join("| " + " | ".join(map(cell, row)) + " |" for row in frame.itertuples(index=False, name=None)) + "\n"


def summarize(data):
    rows = []
    for (kind, method, mode), part in data.groupby(["kind", "method", "mode"], sort=False):
        completed = part.status.isin(["completed", "ok", "success"])
        scored = COMPONENTS if mode == "full" else (["EEG"] if mode in ["center_EEG", "center_EEG_shift", "fNIRS_only"] else ["HbO", "HbR"])
        finite = np.isfinite(part[[f"nrmse_{c}" for c in scored]].to_numpy(float)).all(axis=1)
        valid = completed & truth(part.physical_valid) & truth(part.small_signal_valid) & finite
        passed = valid & part[[f"nrmse_{c}" for c in COMPONENTS]].lt(.5).all(axis=1)
        row = dict(kind=kind, method=method, mode=mode, total=len(part), valid=int(valid.sum()), failed=int((~valid).sum()), all_three_below_05=int(passed.sum()), physical_valid=int(truth(part.physical_valid).sum()), small_signal_valid=int(truth(part.small_signal_valid).sum()))
        if mode != "full":
            row["all_three_below_05"] = "N/A"
        if method == "own_context":
            row.update(valid="N/A", failed="N/A", all_three_below_05="N/A", physical_valid="N/A", small_signal_valid="N/A")
        # Fixed full denominator is retained; finite errors are descriptive,
        # including physically invalid reconstructions, never a success count.
        for component in COMPONENTS:
            values = pd.to_numeric(part[f"nrmse_{component}"], errors="coerce").replace([np.inf, -np.inf], np.nan)
            scoring = part[["subject", "session", f"nmse_{component}", f"nrmse_{component}"]].copy()
            scoring.replace([np.inf, -np.inf], np.nan, inplace=True)
            means = scoring.groupby(["subject", "session"])[[f"nmse_{component}", f"nrmse_{component}"]].mean().groupby("subject").mean().mean()
            row[f"NRMSE_{component}"] = np.sqrt(means[f"nmse_{component}"])
            row[f"NMSE_{component}"] = means[f"nmse_{component}"]
            row[f"mean_trial_NRMSE_{component}"] = means[f"nrmse_{component}"]
            row[f"finite_{component}"] = int(values.notna().sum())
        rows.append(row)
    return pd.DataFrame(rows)


def render(run, out):
    if out.exists() and any(out.iterdir()):
        raise ValueError("Export directory is not empty; choose a new versioned --output")
    data = pd.read_csv(run / "metrics.csv")
    required = {"kind", "group", "subject", "outer", "trial", "session", "method", "mode", "status", "physical_valid", "small_signal_valid", *(f"nrmse_{c}" for c in COMPONENTS), *(f"nmse_{c}" for c in COMPONENTS)}
    if not required.issubset(data.columns):
        raise ValueError(f"Missing metrics columns: {sorted(required - set(data.columns))}")
    if data.empty:
        raise ValueError("No completed evidence rows")
    out.mkdir(parents=True, exist_ok=True)
    (out / "figures").mkdir()
    font = Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc")
    if font.exists():
        font_manager.fontManager.addfont(str(font))
        plt.rcParams["font.family"] = font_manager.FontProperties(fname=str(font)).get_name()
    plt.rcParams.update({"font.size": 9, "axes.unicode_minus": False})
    figures = []

    def save(fig, name, caption):
        fig.tight_layout()
        fig.savefig(out / "figures" / f"{name}.png", dpi=240, bbox_inches="tight")
        plt.close(fig)
        figures.append((name, caption))

    summary = summarize(data)
    summary.to_csv(out / "summary.csv", index=False)
    kinds = list(data.kind.unique())
    fig, axes = plt.subplots(1, len(kinds), figsize=(6 * len(kinds), 4.2), squeeze=False)
    for ax, kind in zip(axes.flat, kinds):
        part = summary[(summary.kind == kind) & (summary["mode"] == "full")]
        x = np.arange(len(part))
        for j, c in enumerate(COMPONENTS):
            ax.bar(x + (j - 1) * .23, part[f"NRMSE_{c}"], .23, label=c)
        ax.axhline(.5, color="red", ls="--", lw=1)
        ax.set_xticks(x, [LABELS.get(m, m) for m in part.method], rotation=25, ha="right")
        ax.set_title(f"{kind}：全观测重建")
        ax.set_ylabel("NRMSE = √被试／会话等权 NMSE")
        ax.legend()
    save(fig, "aggregate_nrmse", "全观测误差包含物理失效结果；按 trial→session→subject 聚合 NMSE 后取根，oracle 不参与方法选择。")

    group_records = []
    for group in data.group.astype(str).unique():
        path = run / "groups" / f"{group}.json"
        if path.exists():
            group_records.append((group, json.loads(path.read_text())))
    measured_groups = [(g, r) for g, r in group_records if g in set(data.loc[data.kind.eq("measured"), "group"].astype(str))]
    records = measured_groups or group_records
    fig, axes = plt.subplots(2, 1, figsize=(10, 7))
    for method in METHODS[:3]:
        values = [record.get("selected_tau", {}).get(method, np.nan) for _, record in records]
        axes[0].plot(range(len(records)), values, "o-", label=LABELS[method])
    axes[0].set_xticks(range(len(records)), [g for g, _ in records], rotation=45, ha="right", fontsize=6)
    axes[0].set_ylabel("τ（秒）")
    axes[0].legend()
    for group, record in records:
        profile = pd.DataFrame(record.get("profiles", []))
        if {"tau", "train_sse"}.issubset(profile.columns):
            profile = profile.sort_values("tau")
            values = profile.train_sse.to_numpy(float)
            axes[1].plot(profile.tau, values - np.nanmin(values), label=group)
    axes[1].set_xscale("log")
    axes[1].set_xlabel("τ（秒）")
    axes[1].set_ylabel("训练 SSE − 本折最小值")
    if records:
        axes[1].legend(fontsize=6, ncol=3)
    save(fig, "tau_profiles", "被试／折参数与训练残差剖面；此 SSE 剖面不是置信区间，折间变化不等同于真实个体差异。")

    # Select examples once by the fixed baseline; candidates cannot pick their
    # own easiest examples. Large or invalid fits remain eligible for display.
    examples = data[data.kind.eq("measured") & data.method.eq("fixed_tau") & data["mode"].eq("full")].copy()
    if examples.empty:
        examples = data[data.method.eq("fixed_tau") & data["mode"].eq("full")].copy()
    examples["rank_error"] = examples[[f"nrmse_{c}" for c in COMPONENTS]].mean(axis=1)
    examples = examples.sort_values(["rank_error", "group", "trial"], kind="stable", na_position="last")
    selected = []
    if len(examples):
        selected = [("median", examples.iloc[(len(examples) - 1) // 2]), ("worst", examples.iloc[-1])]
    for label, row in selected:
        path = run / "predictions" / f"{row.group}.npz"
        with np.load(path, allow_pickle=False) as arrays:
            indices = arrays["trial_indices"]
            matched = np.flatnonzero(indices.astype(str) == str(row.trial))
            if not len(matched):
                raise ValueError(f"Missing example trial {row.trial} in {path}")
            index = int(matched[0])
            target = arrays["target"][index]
            t = np.arange(len(target)) * .25
            fig, axes = plt.subplots(3, 2, figsize=(11, 7.5), sharex=True)
            for j, c in enumerate(COMPONENTS):
                axes[j, 0].plot(t, target[:, j], "k", lw=1.5, label="实测／目标")
                axes[j, 0].set_ylabel(c + "（原评分坐标）")
                for method in METHODS[:4]:
                    key = f"prediction_{method}_full"
                    if key in arrays:
                        axes[j, 0].plot(t, arrays[key][index, :, j], lw=1, label=LABELS[method], alpha=.85)
                state_key = "states_fitted_tau_full"
                if state_key in arrays:
                    states = arrays[state_key][index]
                    for k in [2 * j, 2 * j + 1]:
                        axes[j, 1].plot(t, states[:, k], label=["r", "s", "f", "v", "p", "q"][k])
                axes[j, 1].legend()
            axes[0, 0].legend(fontsize=7)
            axes[-1, 0].set_xlabel("秒")
            axes[-1, 1].set_xlabel("秒")
            fig.suptitle(f"{label}：{row.group} / trial {row.trial}；按固定 τ 误差选样")
            save(fig, f"curves_{label}", f"固定 τ 基线排序的 {label} 样例；右列为拟合 τ 的六状态，失效状态不裁剪。")

    full = summary[summary["mode"].eq("full")]
    text = "# 共享生理驱动重建：线性小信号诊断\n\n"
    text += f"证据来源：`{run}`。本报告包含 {len(data)} 条评价记录；不将本轮线性诊断称为非线性 SSM 完成或生理 teacher 资格。\n\n"
    text += "## 主要结果\n\n"
    for row in full.itertuples():
        text += f"- {row.kind}／{LABELS.get(row.method, row.method)}：EEG、HbO、HbR 聚合 NRMSE 分别为 {row.NRMSE_EEG:.3f}、{row.NRMSE_HbO:.3f}、{row.NRMSE_HbR:.3f}；物理及小信号有效 {row.valid}/{row.total}，失败 {row.failed}，三分量均低于 0.5 且有效 {row.all_three_below_05}/{row.total}。\n"
    text += "\n主 NRMSE 是逐 trial NMSE 在 session 内平均、session 在 subject 内等权平均、subject 等权平均后取平方根，不是逐 trial NRMSE 的算术平均。后者仅在 summary.csv 单列描述。有限误差输出不能代替成功比例；失败、评分分量非有限、物理无效或超出小信号范围均不算生理重建成功。遮挡模式只评分隐藏分量，其他分量空白不算失败，三分量共同阈值记 N/A。自身上下文对照没有生理状态，其物理有效性标为 N/A，不计为生理模型失败。\n\n"
    for kind in kinds:
        subset = full[full.kind.eq(kind)].set_index("method")
        if "fixed_tau" not in subset.index:
            continue
        baseline = subset.loc["fixed_tau"]
        for method in ["fitted_tau", "source_prior_tau"]:
            if method not in subset.index:
                continue
            candidate = subset.loc[method]
            differences = [candidate[f"NRMSE_{c}"] - baseline[f"NRMSE_{c}"] for c in COMPONENTS]
            text += f"{kind}／{LABELS[method]} 相对固定 τ 的三分量聚合 NRMSE 差为 {differences[0]:+.3f}、{differences[1]:+.3f}、{differences[2]:+.3f}（负数表示改善；不是配对置信区间）。"
            passed = all(candidate[f"NRMSE_{c}"] < .5 for c in COMPONENTS)
            text += ("有限输出的三分量聚合误差达到 0.5 以下。" if passed else "有限输出的三分量聚合误差尚未全部达到 0.5 以下。")
            if candidate.failed:
                text += f"仍有 {int(candidate.failed)} 条生理失效或数值失败，不能宣布生理重建达标。"
            text += "\n\n"
    measured_full = full[full.kind.eq("measured")].set_index("method")
    if "fitted_tau" in measured_full.index:
        fit = measured_full.loc["fitted_tau"]
        if fit.NRMSE_HbO < .5 and fit.NRMSE_HbR < .5 and fit.failed:
            text += "**本轮关键结论：两项 Hb 的残差能够低于目标阈值，但共享状态超出已声明的小信号／物理域。曲线贴合没有消除状态异常，因此线性化的生理解释不可信，不能宣布生理重建达标。** 下一步应以受约束非线性动力学检查这些状态是否真实可实现，并保留当前负结果；本报告不把拟合 τ 差异宣布为生理个体差异。\n\n"
    synthetic = data[data.kind.eq("synthetic") & data["mode"].eq("full")]
    if "driver_nrmse" in synthetic and len(synthetic):
        text += "合成驱动恢复（没有事后尺度或时间对齐）：\n\n"
        diagnostics = synthetic.groupby("method")["driver_nrmse"].agg(["count", "mean", "max"]).reset_index()
        text += table(diagnostics) + "\n"
    columns = ["kind", "method", "mode", "total", "valid", "failed", "all_three_below_05", "NRMSE_EEG", "NRMSE_HbO", "NRMSE_HbR"]
    text += "## 全部分母和遮挡对照\n\n" + table(summary[columns])
    text += "\n## 方法与解释边界\n\n"
    text += "仅一个共享驱动和五个血流初态；血流状态没有逐点自由创新。参数只在训练折选择，评分采用冻结训练 SD。全观测重建与中心特征遮挡补全分别报告；离线处理后的补全不是原始传感器遮挡或未来预测。EEG 坐标是 log-power PCA，不是原始电压。\n\n"
    text += "oracle 使用同一目标选择最小 SSE，仅为有限 4 Hz ZOH 驱动、现有测量插值算子和 τ 网格所定义线性族的乐观残差下界，不能参与模型选择、个体参数解释或预测优劣比较。仅当三分量聚合 NMSE 的平均值大于 0.25，才可在这一有限线性族中排除三分量聚合 NRMSE 同时低于 0.5；不能对逐 trial NRMSE 算术平均作此推论。先验 τ 的工程 SSE 惩罚不是校准后验。τ 的离散剖面和折间一致性只用于诊断，不能证明参数唯一可辨识。\n\n"
    text += "[Tak 等（2015）](https://www.fil.ion.ucl.ac.uk/~wpenny/publications/tak-penny15.pdf)区分血流参数与 HbO/HbR 光学权重，并采用参数先验；这不意味着相同先验已获当前数据验证。[Friston 等（2016）](https://pmc.ncbi.nlm.nih.gov/articles/PMC4767224/)提供保留个体差异的层次收缩方法；本轮尚未实现层次模型。[Raue 等（2009）](https://pubmed.ncbi.nlm.nih.gov/19505944/)支持通过重新优化其余参数的剖面分析评估可辨识性，不能用局部曲率或参数不贴边替代。\n\n"
    text += "## 图像核查\n\n样例按固定 τ 全观测平均 NRMSE 稳定排序，选择中位与最差记录，候选方法共享样例。图像用于检查幅度、峰时、双 Hb 联动及状态异常；不能仅凭曲线接近证明生理语义。\n\n"
    for name, caption in figures:
        text += f"![{name}](figures/{name}.png)\n\n{caption}\n\n"
    text += "## 交付验证\n\n全部图形以 240 dpi PNG 嵌入 PDF，正文与表格保持可搜索。程序检查图像对象及正文提取；WPS 滚动与缩放未检查。\n"
    export_pdf(text, out, figures)


def export_pdf(text, out, figures, *, extra_css="", split_images=True):
    """Keep report prose selectable and embed whole figures as bitmap images."""
    (out / "REPORT.md").write_text(text)
    body = markdown.markdown(text, extensions=["tables"])
    body = body.replace("<img ", '<img width="510" ')
    css = "body{font-family:sans-serif;font-size:9pt;line-height:1.4}h1{font-size:20pt}h2{font-size:13pt}table{font-size:6pt;border-collapse:collapse}td,th{border:0.4pt solid #bbb;padding:3pt}img{max-width:100%}"
    def rectfn(number, filled):
        if number > 80:
            raise RuntimeError("PDF pagination failed")
        page = fitz.paper_rect("a4")
        return page, page + (36, 32, -36, -32), None
    # Separate stories prevent a leading image page break from discarding the
    # preceding prose overflow in MuPDF. Each figure still starts a fresh page.
    pdf = fitz.open()
    for section in re.split(r"(?=<h2>|<p><img )" if split_images else r"(?=<h2>)", body):
        if not section.strip():
            continue
        story = fitz.Story(section, user_css=css + extra_css, archive=fitz.Archive(str(out)))
        part = story.write_with_links(rectfn)
        pdf.insert_pdf(part)
        part.close()
    pdf.save(out / "REPORT.pdf", garbage=4, deflate=True)
    validation = {"pages": len(pdf), "image_objects": sum(len(p.get_images()) for p in pdf), "figure_count": len(figures), "searchable_text_chars": sum(len(p.get_text()) for p in pdf), "wps_checked": False, "bitmap_dpi": 240}
    if validation["image_objects"] != len(figures) or validation["searchable_text_chars"] < 100:
        raise RuntimeError(f"PDF verification failed: {validation}")
    # Render representative pages for visual review without rasterizing delivery.
    for i in sorted({0, max(0, len(pdf) - len(figures)), len(pdf) - 1}):
        pdf[i].get_pixmap(matrix=fitz.Matrix(1.3, 1.3)).save(out / f"preview_page_{i+1:02d}.png")
    pdf.close()
    (out / "export_validation.json").write_text(json.dumps(validation, indent=2))
    print(json.dumps(validation))


def render_replay(run, out):
    """Compare a fixed-drive nonlinear replay against its exact parent fits."""
    if out.exists() and any(out.iterdir()):
        raise ValueError("Export directory is not empty; choose a new versioned --output")
    manifest = json.loads((run / "manifest.json").read_text())
    summary = json.loads((run / "summary.json").read_text())
    data = pd.read_csv(run / "metrics.csv")
    parent = Path(manifest["source_run"])
    if not parent.is_absolute():
        parent = Path(__file__).resolve().parents[2] / parent
    parent_data = pd.read_csv(parent / "metrics.csv")
    out.mkdir(parents=True, exist_ok=True)
    (out / "figures").mkdir()
    font = Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc")
    if font.exists():
        font_manager.fontManager.addfont(str(font))
        plt.rcParams["font.family"] = font_manager.FontProperties(fname=str(font)).get_name()
    plt.rcParams.update({"font.size": 9, "axes.unicode_minus": False})
    figures = []

    def save(fig, name, caption):
        fig.tight_layout()
        fig.savefig(out / "figures" / f"{name}.png", dpi=240, bbox_inches="tight")
        plt.close(fig)
        figures.append((name, caption))

    aggregates = summary["aggregates"]
    rows = []
    for record in aggregates:
        counts = record.get("failure_counts", {})
        row = {"kind": record["kind"], "method": record["method"], "expected": record["expected"], "observed": record["observed_rows"], "completed": record["completed"], "domain_fail": counts.get("failed_domain", 0), "integration_fail": counts.get("failed_integration_check", 0), "missing": record["expected"] - record["observed_rows"]}
        for c in COMPONENTS:
            row[f"linear_{c}"] = record.get(f"linear_nrmse_same_subset_{c}", np.nan)
            row[f"nonlinear_{c}"] = record.get(f"nrmse_{c}", np.nan)
        rows.append(row)
    frame = pd.DataFrame(rows)
    frame.to_csv(out / "replay_summary.csv", index=False)
    kinds = list(frame.kind.unique())
    fig, axes = plt.subplots(len(kinds), 3, figsize=(12, 3.8 * len(kinds)), squeeze=False)
    for i, kind in enumerate(kinds):
        subset = frame[frame.kind.eq(kind)]
        for j, c in enumerate(COMPONENTS):
            ax = axes[i, j]
            x = np.arange(len(subset))
            ax.bar(x - .18, subset[f"linear_{c}"], .36, label="线性：同成功子集")
            ax.bar(x + .18, subset[f"nonlinear_{c}"], .36, label="非线性固定驱动重放")
            ax.axhline(.5, color="red", ls="--", lw=1)
            ax.set_xticks(x, [f"{LABELS[m]}\n{n}/{d} 成功" for m, n, d in zip(subset.method, subset.completed, subset.expected)], rotation=25, ha="right", fontsize=7)
            ax.set_title(f"{kind} / {c}")
            ax.set_ylabel("√被试／会话等权 NMSE")
            if i == 0 and j == 0:
                ax.legend(fontsize=7)
    save(fig, "replay_aggregate", "只在非线性域及积分分辨率检查均通过的同一子集比较误差；完整失败分母单独列出，不将缺失条目算作成功。")

    examples = parent_data[parent_data.kind.eq("measured") & parent_data.method.eq("fixed_tau") & parent_data["mode"].eq("full")].copy()
    if examples.empty:
        examples = parent_data[parent_data.method.eq("fixed_tau") & parent_data["mode"].eq("full")].copy()
    examples["rank_error"] = examples[[f"nrmse_{c}" for c in COMPONENTS]].mean(axis=1)
    examples = examples.sort_values(["rank_error", "group", "trial"], kind="stable", na_position="last")
    if examples.empty:
        raise ValueError("Parent fixed-baseline examples unavailable")
    selections = [("median", examples.iloc[(len(examples) - 1) // 2]), ("worst", examples.iloc[-1])]
    selection_notes = []
    for label, selected in selections:
        replay_path = run / "predictions" / f"{selected.group}.npz"
        if not replay_path.exists():
            selection_notes.append(f"父级 {label} 样例 {selected.group}/trial {selected.trial} 的重放数组缺失；保留身份，不换样。")
            continue
        with np.load(parent / "predictions" / f"{selected.group}.npz", allow_pickle=False) as linear, np.load(replay_path, allow_pickle=False) as nonlinear:
            indices = np.flatnonzero(linear["trial_indices"].astype(str) == str(selected.trial))
            replay_indices = np.flatnonzero(nonlinear["trial_indices"].astype(str) == str(selected.trial))
            if len(indices) != 1 or len(replay_indices) != 1:
                raise ValueError("Parent/replay example identity mismatch")
            index, replay_index = int(indices[0]), int(replay_indices[0])
            target = linear["target"][index]
            if not np.array_equal(target, nonlinear["target"][replay_index]):
                raise ValueError("Parent/replay target mismatch")
            t = np.arange(len(target)) * .25
            fig, axes = plt.subplots(3, 2, figsize=(11, 7.5), sharex=True)
            selected_rows = data[data.group.eq(selected.group) & data.trial.eq(selected.trial)]
            statuses = dict(zip(selected_rows.method, selected_rows.status))
            for j, c in enumerate(COMPONENTS):
                axes[j, 0].plot(t, target[:, j], "k", lw=1.5, label="实测／目标")
                for method, color in [("fixed_tau", "#2877ab"), ("fitted_tau", "#dc8030"), ("source_prior_tau", "#358d53")]:
                    key = f"prediction_{method}_full"
                    axes[j, 0].plot(t, linear[key][index, :, j], "--", color=color, lw=1, label=f"{LABELS[method]} 线性")
                    if statuses.get(method) == "completed":
                        axes[j, 0].plot(t, nonlinear[key][replay_index, :, j], color=color, lw=1.2, label=f"{LABELS[method]} 非线性")
                axes[j, 0].set_ylabel(c + "（原评分坐标）")
                state_key = "states_fitted_tau_full"
                for k, color in zip([2 * j, 2 * j + 1], ["#2877ab", "#dc8030"]):
                    name = ["r", "s", "f", "v", "p", "q"][k]
                    axes[j, 1].plot(t, linear[state_key][index, :, k], "--", color=color, label=f"{name} 线性")
                    if statuses.get("fitted_tau") == "completed":
                        axes[j, 1].plot(t, nonlinear[state_key][replay_index, :, k], color=color, label=f"{name} 非线性")
                axes[j, 1].legend(fontsize=7)
            axes[0, 0].legend(fontsize=6, ncol=2)
            for ax in axes[-1]:
                ax.set_xlabel("秒")
            fig.suptitle(f"父级 {label} 样例：{selected.group} / trial {selected.trial}；不重新拟合")
            status_text = "；".join(f"{LABELS.get(m,m)}={s}" for m, s in statuses.items())
            save(fig, f"replay_{label}", f"父级固定 τ 选定的 {label} 样例。虚线=原线性，实线=非线性重放；右侧为拟合 τ 六状态。{status_text}。失败方法不画非线性曲线，不换样。")

    text = "# 非线性固定驱动重放\n\n"
    text += f"本轮来源：`{run}`；父级线性拟合：`{parent}`。保持父级 τ、驱动 r(t) 及血流初态固定，代入非线性动力学，**不是重新拟合**。\n\n"
    text += f"## 主要结果与完整分母\n\n预期 {summary['expected_groups']} 组，完成 {summary['completed_groups']} 组，组失败 {summary['failed_groups']}；方法×trial 成功 {summary['completed_method_trials']}/{summary['expected_method_trials']}。组作业完成不意味着其中所有轨迹通过。\n\n"
    for row in frame.itertuples():
        text += f"- {row.kind}／{LABELS.get(row.method,row.method)}：成功 {row.completed}/{row.expected}；域失败 {row.domain_fail}、积分检查失败 {row.integration_fail}、缺行 {row.missing}。同成功子集 EEG/HbO/HbR NRMSE：线性 {row.linear_EEG:.3f}/{row.linear_HbO:.3f}/{row.linear_HbR:.3f} → 非线性 {row.nonlinear_EEG:.3f}/{row.nonlinear_HbO:.3f}/{row.nonlinear_HbR:.3f}。\n"
    text += "\n" + table(frame[["kind", "method", "expected", "observed", "completed", "domain_fail", "integration_fail", "missing"]])
    text += "\n同成功子集的聚合误差：\n\n" + table(frame[["kind", "method", *[f"{prefix}_{c}" for prefix in ["linear", "nonlinear"] for c in COMPONENTS]]])
    if "integration_max_difference_training_sd" in data:
        precision = data.groupby("kind").integration_max_difference_training_sd.agg(["count", "median", "max"]).reset_index()
        text += "\n积分分辨率核查：下表比较两个子步分辨率的预测差，单位为冻结训练 SD；域失败没有可用积分误差，不填零。\n\n" + table(precision)
        for row in precision.itertuples():
            text += f"{row.kind}：最大差 {row.max:.6g} 个训练 SD，中位差 {row.median:.6g}；有效比较 {row.count} 条。\n\n"
    text += "\n## 结论与限制\n\n误差比较采用同成功子集，按 session、subject 等权聚合 NMSE 后取根；该条件性指标不代表失败轨迹的误差，也不等同于完整面板达标。各方法成功子集可能不同，不能用此表直接作跨方法胜负排名。\n\n"
    for row in frame[frame.kind.eq("measured")].itertuples():
        values = [row.nonlinear_EEG, row.nonlinear_HbO, row.nonlinear_HbR]
        reached = all(np.isfinite(v) and v < .5 for v in values)
        text += f"{LABELS.get(row.method,row.method)}：" + ("成功子集三分量误差均低于 0.5。" if reached else "成功子集三分量误差尚未全部低于 0.5。")
        if row.completed < row.expected:
            text += f"另有 {row.expected-row.completed} 条未通过，不能宣称全分母达标。"
        text += "\n\n"
    text += "本轮只能判断既有线性驱动／初态在非线性动力学中的可实现性与重放误差；不能将重放失败等同于所有非线性模型不可拟合，也不能把成功重放称为参数唯一可辨识或 teacher 资格。下一步若需要降低误差，应另行实现受约束非线性重拟合并保持共同训练／验证分母。oracle 继承全目标拟合驱动，只作诊断，不是可部署候选或选择依据。\n\n"
    oracle = frame[frame.kind.eq("measured") & frame.method.eq("oracle_lower_bound")]
    if len(oracle) and all(oracle.iloc[0][f"nonlinear_{c}"] < .5 for c in COMPONENTS):
        text += "oracle 同成功子集的非线性误差仍全部低于 0.5，提示完整非线性建模存在值得验证的空间，不能据此断言非线性重建不可能。但该驱动及逐 trial τ 由同一目标优化而来，并且仍有完整分母中的失败；它不能作为主方法成功或泛化资格。\n\n"
    text += "## 固定样例与图像核查\n\n中位与最差样例严格沿用父级 fixed_tau/full 的误差排序；失败时保留样例及失败说明，不挑换漂亮曲线。\n\n"
    text += "\n\n".join(selection_notes) + "\n\n"
    for name, caption in figures:
        text += f"![{name}](figures/{name}.png)\n\n{caption}\n\n"
    text += "## 交付验证\n\n全部图形作为 240 dpi PNG 嵌入 PDF；正文与表格可搜索。WPS 滚动／缩放未检查。\n"
    export_pdf(text, out, figures)


def continuation_parent_comparison(run, out, config, data, labels):
    """Compare exactly matched parent identities; newly available fits are separate."""
    manifest = json.loads((run / 'manifest.json').read_text())
    parent = Path(config['continuation']['parent_run'])
    if not parent.is_absolute():
        parent = Path(manifest['project_root']) / parent
    previous = pd.read_csv(parent / 'metrics.csv')
    identity = ['kind','group','subject','session','trial','method','mode']
    if data.duplicated(identity).any() or previous.duplicated(identity).any():
        raise ValueError('Continuation parent comparison has duplicate trial identities')
    joined = data.merge(previous, on=identity, suffixes=('_new','_parent'), how='outer', indicator=True, validate='one_to_one')
    if not joined['_merge'].eq('both').all():
        raise ValueError('Continuation and parent identity sets differ')
    expected = len(config['subjects'])*len(config['outer_folds'])*6
    if expected != 72:
        raise ValueError('Continuation comparison requires the frozen 72-trial panel')
    for (kind, method, mode), part in joined.groupby(['kind','method','mode']):
        if len(part) != expected or len(part.group.unique()) != 12 or not part.groupby('group').size().eq(6).all():
            raise ValueError('Continuation comparison does not contain all 72 identities')
    fixed = joined[joined.method.isin(['fixed_gain_no_amplitude','fixed_gain_amplitude'])]
    if (not fixed.status_new.eq(fixed.status_parent).all()
            or not truth(fixed.converged_new).eq(truth(fixed.converged_parent)).all()):
        raise ValueError('Inherited fixed-gain status changed from parent')
    for prefix in ['nmse_', 'nrmse_']:
        for component in COMPONENTS:
            field = prefix+component
            a = pd.to_numeric(fixed[field+'_new'], errors='coerce').to_numpy()
            b = pd.to_numeric(fixed[field+'_parent'], errors='coerce').to_numpy()
            if not np.array_equal(a, b, equal_nan=True):
                raise ValueError('Inherited fixed-gain score changed from parent')
    joined['success_new'] = joined.status_new.eq('completed') & truth(joined.converged_new)
    joined['success_parent'] = joined.status_parent.eq('completed') & truth(joined.converged_parent)
    transitions, scores = [], []
    for (kind, method, mode), part in joined.groupby(['kind','method','mode'], sort=False):
        groups = {'success_to_success':part.success_new & part.success_parent,
                  'new_success':part.success_new & ~part.success_parent,
                  'lost_success':~part.success_new & part.success_parent,
                  'both_failed':~part.success_new & ~part.success_parent}
        transitions.append(dict(kind=kind, method=method, mode=mode, expected=expected,
            **{name:int(mask.sum()) for name,mask in groups.items()}))
        for subset, mask in groups.items():
            selected = part[mask]
            if subset == 'both_failed': continue
            for component in COMPONENTS:
                field = 'nmse_'+component
                row = dict(kind=kind, method=method, mode=mode, subset=subset,
                    common_or_new_identities=len(selected), expected=expected, component=component)
                for suffix in ['parent','new']:
                    usable = selected if subset == 'success_to_success' else selected[selected['success_'+suffix]]
                    column = field+'_'+suffix
                    value = usable.groupby(['subject','session'])[column].mean().groupby('subject').mean().mean()
                    row[suffix+'_nrmse'] = float(np.sqrt(value))
                if subset == 'success_to_success':
                    delta = selected[['subject','session']].assign(delta=selected[field+'_new']-selected[field+'_parent'])
                    row['delta_nmse'] = delta.groupby(['subject','session']).delta.mean().groupby('subject').mean().mean()
                else:
                    row['delta_nmse'] = np.nan
                if np.isfinite(row['parent_nrmse']) or np.isfinite(row['new_nrmse']): scores.append(row)
    transitions = pd.DataFrame(transitions); scores = pd.DataFrame(scores)
    transitions.to_csv(out / 'continuation_parent_transitions.csv', index=False)
    scores.to_csv(out / 'continuation_parent_paired_scores.csv', index=False)
    body = '\n## 与父运行的同身份验证对照\n\n'
    body += (f'父证据：`{parent.name}/metrics.csv`。按kind/group/subject/session/trial/method/mode一对一核对，'
        '每种数据、方法和模式均保留原12组×6条=72个身份。固定A/B的status、converged及各分量NMSE/NRMSE与父证据逐值一致（含相同缺失值）。'
        '下面将成功状态迁移与误差比较分开；新增成功样本不混入共同成功的改善量。继承的合成与固定臂不是新独立复制。\n\n')
    for (kind, mode), part in transitions.groupby(['kind','mode'], sort=False):
        body += f'### {kind} · {mode} · 成功状态迁移\n\n'
        body += table(pd.DataFrame([{'臂':labels[r.method].split()[0], '原成功→成功':r.success_to_success,
            '新增成功':r.new_success, '失去成功':r.lost_success, '双方失败':r.both_failed, '完整分母':r.expected}
            for r in part.itertuples()])) + '\n\n'
    if len(scores):
        # The key scientific comparison is measured C/D; full CSV also keeps inherited arms.
        relevant = scores[scores.kind.eq('measured') & scores.method.str.startswith('trained_gain_')]
        names = {'success_to_success':'共同成功', 'new_success':'新增成功（无父成功分数）', 'lost_success':'失去成功（无新成功分数）'}
        for (mode, subset), part in relevant.groupby(['mode','subset'], sort=False):
            body += f'### 实测 · {mode} · {names[subset]}\n\n'
            body += table(pd.DataFrame([{'臂':labels[r.method].split()[0], '分量':r.component,
                '身份/计划':f'{r.common_or_new_identities}/{r.expected}', '父NRMSE':r.parent_nrmse,
                '本轮NRMSE':r.new_nrmse, 'ΔNMSE':r.delta_nmse} for r in part.itertuples()])) + '\n\n'
    body += ('NRMSE均按所列相同身份内trial→session→subject等权NMSE后取根；ΔNMSE仅在双方成功时计算。'
        '新增或失去成功只报告对应可用一侧，不拿失败最后有效解补分数，不用不同成功子集的整体均值声称预测改善。\n\n')
    return body


def gain_continuation_details(run, out, config, data, labels, save):
    """Summarize inherited evidence and declared warm-restart traces, without refitting."""
    rows, traces, cells = [], [], []
    for group in data.group.unique():
        kind = data.loc[data.group.eq(group), 'kind'].iloc[0]
        for amplitude in (0, 1):
            method = 'trained_gain_amplitude' if amplitude else 'trained_gain_no_amplitude'
            for start in range(len(config['gain_training']['start_values'])):
                path = run / 'training' / f'{group}__amplitude{amplitude}' / f'start_{start}' / 'result.json'
                result = json.loads(path.read_text()) if path.exists() else {}
                info = result.get('continuation', {})
                inherited = info.get('inherited', False)
                parent_total = result.get('evaluations') if inherited else info.get('parent_evaluations')
                parent_opt = result.get('optimization_evaluations') if inherited else info.get('parent_optimization_evaluations')
                trace = result.get('trace', [])
                first = trace[0] if trace else {}
                parent_objective = info.get('parent_objective')
                initial_objective = first.get('objective', info.get('initial_objective'))
                difference = (initial_objective-parent_objective
                    if initial_objective is not None and parent_objective is not None else np.nan)
                row = dict(kind=kind, group=group, method=method, start=start,
                    action=info.get('action', 'missing_provenance' if result else 'missing'), parent_record=info.get('parent_record'),
                    status=result.get('status', 'missing'), converged=bool(result.get('converged', False)),
                    parent_library_evaluations=parent_total, new_library_evaluations=0 if inherited else info.get('new_evaluations'),
                    cumulative_library_evaluations=parent_total if inherited else info.get('cumulative_evaluations'),
                    parent_optimization_evaluations=parent_opt,
                    new_optimization_evaluations=0 if inherited else info.get('new_optimization_evaluations'),
                    cumulative_optimization_evaluations=parent_opt if inherited else info.get('cumulative_optimization_evaluations'),
                    original_start_gain=result.get('start_gain'), parent_gain=info.get('parent_gain'),
                    initial_gain=first.get('parameter_value'), final_gain=result.get('parameter_value'),
                    parent_objective=parent_objective, initial_objective=initial_objective,
                    first_objective_difference=difference, objective_continuity=info.get('objective_continuity'),
                    continuity_rtol=info.get('objective_continuity_rtol'), continuity_atol=info.get('objective_continuity_atol'),
                    trace_points=len(trace),
                    new_domain_rejections=(None if inherited else result.get('starts', [{}])[0].get('domain_rejections') if result.get('starts') else None),
                    new_fine_replays=0 if inherited else len(result.get('integration_checks', [])))
                rows.append(row)
                for point in trace:
                    traces.append(dict(kind=kind, group=group, method=method, start=start, **point))
        for method, mode in data[data.group.eq(group)][['method','mode']].drop_duplicates().itertuples(index=False, name=None):
            path = run / 'cells' / f'{group}__{method}__{mode}' / 'result.json'
            result = json.loads(path.read_text()) if path.exists() else {}
            info = result.get('continuation', {})
            cells.append(dict(kind=kind, group=group, method=method, mode=mode,
                action=info.get('action', 'new_validation' if result else 'missing'), parent_record=info.get('parent_record')))
    records = pd.DataFrame(rows); trace_data = pd.DataFrame(traces); cell_data = pd.DataFrame(cells)
    records.to_csv(out / 'continuation_training.csv', index=False)
    trace_data.to_csv(out / 'continuation_traces.csv', index=False)
    cell_data.to_csv(out / 'continuation_cells.csv', index=False)
    body = '\n## 预算热启动与继承证据\n\n'
    body += ('保留原始起点谱系不等于从原始0.5/1/2重新初始化：eligible起点使用父运行最后有效driver、初态与β，'
        'LM阻尼重新设定。因此这是warm restart，不是exact optimizer resume；即使预算总量相同也不保证等于一次不中断优化。'
        '所有原身份、失败起点与成功阈值保留，未以热启动后较好的结果增加独立样本数。\n\n')
    action_labels = {'inherited_training':'继承训练', 'restarted_training':'热启动训练',
                     'inherited_cell':'继承验证', 'inherited_validation':'继承验证', 'new_validation':'重算验证'}
    for kind, part in records.groupby('kind', sort=False):
        display = []
        for (method, action), subset in part.groupby(['method','action'], sort=False):
            display.append({'臂':labels[method].split()[0], '来源':action_labels.get(action,action),
                '起点数':len(subset), '收敛成功':int((subset.status.eq('completed') & subset.converged).sum()),
                '新增优化调用':int(subset.new_optimization_evaluations.sum()) if subset.new_optimization_evaluations.notna().all() else '记录不完整',
                '累计优化调用':int(subset.cumulative_optimization_evaluations.sum()) if subset.cumulative_optimization_evaluations.notna().all() else '记录不完整'})
        body += f'### {kind} · 训练来源与调用数\n\n' + table(pd.DataFrame(display)) + '\n\n'
    counts = cell_data.groupby(['kind','action']).size().rename('cells').reset_index()
    body += table(counts.assign(action=counts.action.map(lambda x:action_labels.get(x,x))).rename(
        columns={'kind':'数据','action':'验证来源','cells':'cell数'})) + '\n\n'
    body += ('优化调用按单trial forward计数；CSV另列父/新增/累计的库reported total evaluations，'
        '该total=optimization+库内4子步独立核验（完整组18条）；runner随后另做18条8子步fine replay，不计入该字段。CSV用library_evaluations命名，另列new_fine_replays，绝不能将库调用数称为整个workflow的全部forward调用或墙钟时间。'
        '继承条目新增库调用和新增fine replay均为0，累计库数保持父值；这不抹去父阶段已发生的库调用与细积分历史成本，也不表示它们按本轮3600/90重新运行。\n\n')
    restarted = records[records.action.eq('restarted_training')]
    if len(restarted):
        continuity = restarted.first_objective_difference.abs().dropna()
        passed = int(restarted.objective_continuity.eq(True).sum())
        body += (f'热启动记录{len(restarted)}/{config["continuation"]["expected_measured_restarts"]}；'
            f'首目标与父末目标连续性通过{passed}/{len(restarted)}，可计算差值{len(continuity)}/{len(restarted)}，'
            f'绝对差中位{continuity.median():.3g}、最大{continuity.max():.3g}。'
            '通过标志及容差取自runner记录；首目标由trace首点独立读取，失败或缺失不补作通过。\n\n')
        domain_seen = restarted.new_domain_rejections.gt(0)
        domain_known = restarted.new_domain_rejections.notna()
        domain_failed = domain_seen & ~(restarted.status.eq('completed') & restarted.converged)
        body += (f'父运行“无域/导数拒绝”只用于热启动资格，不保证后续轨迹无拒绝。'
            f'本热启动段域拒绝记录可用{int(domain_known.sum())}/{len(restarted)}，'
            f'出现域拒绝{int(domain_seen.sum())}条，其中最终未成功{int(domain_failed.sum())}条；'
            '逐步计数保留在continuation_traces.csv，不能将父阶段无域拒绝标签沿用为本段结果。\n\n')
        short = lambda g: str(g).replace('subject_', 'S')
        body += table(pd.DataFrame([{'组':short(r.group), '臂':labels[r.method].split()[0],
            '原起点':r.original_start_gain, '首目标−父末':f'{r.first_objective_difference:.3g}',
            '连续':str(r.objective_continuity), '新增/累计优化':f'{r.new_optimization_evaluations}/{r.cumulative_optimization_evaluations}',
            '终态':r.status.replace('completed','成功').replace('failed_numerical','数值失败')}
            for r in restarted.itertuples()])) + '\n\n'
    if len(trace_data):
        fig, axes = plt.subplots(3, 2, figsize=(12, 9), sharex='col')
        for column, method in enumerate(['trained_gain_no_amplitude','trained_gain_amplitude']):
            part = trace_data[trace_data.method.eq(method)]
            groups = list(part.group.unique())
            colors = plt.get_cmap('tab20')
            for group_index, group in enumerate(groups):
                for start, points in part[part.group.eq(group)].groupby('start', sort=True):
                    points = points.sort_values(['evaluations','iteration'], kind='stable')
                    x = points.evaluations.to_numpy()
                    values = [points.objective-points.objective.iloc[0], points.projected_scaled_gradient_inf_norm, points.parameter_value]
                    for ax, value in zip(axes[:,column], values):
                        ax.plot(x, value, color=colors(group_index % 20), ls=['-', '--', ':'][int(start)%3],
                            lw=.9, alpha=.8, label=group.replace('subject_','S') if start == part[part.group.eq(group)].start.min() else None)
            axes[0,column].set_title(labels[method])
            axes[0,column].set_ylabel('目标 − 热启动首目标')
            axes[1,column].set_ylabel('projected scaled gradient')
            axes[1,column].set_yscale('log')
            axes[1,column].axhline(config['gain_training']['gradient_tolerance'], color='red', ls=':', lw=1)
            axes[2,column].set_ylabel('β'); axes[2,column].set_xlabel('本热启动段优化trial forward调用')
            if groups: axes[0,column].legend(fontsize=5, ncol=2)
        save(fig, 'gain_continuation_traces', '每条线是一条热启动谱系：颜色区分训练组，线型区分原起点。上行目标仅减去同条首目标以显示变化；中行保留实际梯度及原收敛阈值，末行为β。横轴为新增优化调用，父/累计调用另见表与CSV；继承谱系没有新trace，不伪造轨迹。')
    else:
        body += '本次记录没有可用热启动trace，不能展示优化轨迹。\n\n'
    return body


def gain_prior_details(run, out, config, data, states, labels, save, selected):
    """Read declared training starts only; no parameter estimation or winner search."""
    candidates, selections = [], []
    training = config["gain_training"]
    for group in data.group.unique():
        prep = json.loads((run / "prepared" / f"{group}.json").read_text())
        spec = prep["spec"]
        for amplitude in (0, 1):
            method = "trained_gain_amplitude" if amplitude else "trained_gain_no_amplitude"
            folder = run / "training" / f"{group}__amplitude{amplitude}"
            selection_path = folder / "selection.json"
            selection = json.loads(selection_path.read_text()) if selection_path.exists() else {}
            chosen = selection.get("selected_training_start")
            group_rows = []
            for index, start_gain in enumerate(training["start_values"]):
                path = folder / f"start_{index}" / "result.json"
                result = json.loads(path.read_text()) if path.exists() else {}
                info = result.get("information", {})
                attempts = result.get("starts", [])
                attempt = attempts[0] if len(attempts) == 1 else {}
                row = dict(kind=spec["kind"], group=group, method=method, start=index,
                    start_gain=start_gain, true_gain=spec.get("true_gain"),
                    selected=index == chosen, status=result.get("status", "missing"),
                    converged=bool(result.get("converged", False)),
                    gain=result.get("parameter_value"), objective=result.get("objective"),
                    boundary=result.get("boundary_status", "N/A"),
                    convergence_reason=result.get("convergence_reason", attempt.get("convergence_reason")),
                    projected_gradient=result.get("projected_scaled_gradient_inf_norm", attempt.get("projected_scaled_gradient_inf_norm")),
                    domain_rejections=attempt.get("domain_rejections"),
                    nondecreasing_rejections=attempt.get("nondecreasing_rejections"),
                    derivative_rejections=attempt.get("derivative_rejections"),
                    evaluations=result.get("optimization_evaluations"),
                    weighted_data_sse=result.get("weighted_data_sse"),
                    driver_amplitude_cost=result.get("driver_amplitude_cost"),
                    parameter_prior_cost=result.get("parameter_prior_cost"),
                    initial_state_cost=result.get("initial_state_penalty_cost"),
                    data_information=info.get("data_only", {}).get("value"),
                    regularized_information=info.get("nuisance_regularized", {}).get("value"),
                    prior_information=info.get("with_parameter_prior"))
                candidates.append(row); group_rows.append(row)
            chosen_rows = [r for r in group_rows if r["selected"]]
            if selection.get("status") == "completed" and len(chosen_rows) != 1:
                raise ValueError("Completed gain selection has no declared selected start")
            chosen_row = chosen_rows[0] if chosen_rows else {}
            completed = [r for r in group_rows if r["status"] == "completed" and r["converged"]]
            if chosen_row and (chosen_row["status"] != "completed" or not chosen_row["converged"]):
                raise ValueError("Gain selection points to nonconverged training start")
            selections.append(dict(kind=spec["kind"], group=group, method=method,
                status=selection.get("status", "missing"), gain=selection.get("parameter_value"),
                true_gain=spec.get("true_gain"), selected_start=chosen,
                completed_starts=len(completed), expected_starts=len(training["start_values"]),
                converged_gain_min=min([r["gain"] for r in completed], default=np.nan),
                converged_gain_max=max([r["gain"] for r in completed], default=np.nan),
                boundary=chosen_row.get("boundary", "N/A"),
                data_information=chosen_row.get("data_information"),
                regularized_information=chosen_row.get("regularized_information"),
                prior_information=chosen_row.get("prior_information")))
    candidates = pd.DataFrame(candidates); selections = pd.DataFrame(selections)
    candidates.to_csv(out / "gain_training_candidates.csv", index=False)
    selections.to_csv(out / "gain_training_selection.csv", index=False)
    short = lambda value: str(value).replace("subject_", "S").replace("synthetic_gain", "G").replace("mixed", "混").replace("slow", "慢")
    body = "\n## 有效增益训练与工程先验\n\n"
    body += (f"每个训练组和可训练臂有{len(training['start_values'])}个独立起点；只在全部起点终态后按收敛目标选择。"
        f"β范围{training['bounds']}、log prior SD={training['prior_log_sd']:.4g}是工程敏感性设置，不是正常生理范围。"
        "幅度先验的σr直接取既有冻结训练EEG normalizer/loading，未重新去均值估计尺度。"
        "其中心为零但不强制driver样本均值为零，也不代表事件前五秒是静息。"
        "下列局部信息在logβ坐标，data-only先消除driver和初态方向；regularized另包含这些变量的惩罚。"
        "加β先验的信息单列，不能将正则带来的信息解释成数据可辨识性或置信区间。\n\n")
    # Synthetic truth permits an explicit bias diagnostic; never infer physical calibration.
    expected_synthetic = len(config["synthetic"]["gains"])*len(config["synthetic"]["spectra"])*config["synthetic"]["replicates"]
    synthetic = selections[selections.kind.eq("synthetic")].copy()
    synthetic["relative_gain_error"] = pd.to_numeric(synthetic.gain, errors="coerce") / pd.to_numeric(synthetic.true_gain, errors="coerce") - 1.
    gain_recovery = []
    for method in ("trained_gain_no_amplitude", "trained_gain_amplitude"):
        part = synthetic[synthetic.method.eq(method)]
        good = part[part.status.eq("completed") & np.isfinite(part.relative_gain_error)]
        gain_recovery.append(dict(method=method, expected=expected_synthetic, terminal=int(part.status.isin(["completed", "failed_training"]).sum()),
            completed=len(good), relative_error_minimum=good.relative_gain_error.min(),
            relative_error_median=good.relative_gain_error.median(), relative_error_maximum=good.relative_gain_error.max()))
    pd.DataFrame(gain_recovery).to_csv(out / "synthetic_gain_recovery.csv", index=False)
    if all(r["terminal"] == expected_synthetic for r in gain_recovery):
        body += "### 合成真值偏差诊断\n\n"
        for r in gain_recovery:
            if not r["completed"]:
                body += f"{labels[r['method']]}：成功训练组0/{r['expected']}，β误差不可估计。"
                continue
            body += (f"{labels[r['method']]}：成功训练组{r['completed']}/{r['expected']}；"
                f"β相对真值误差(估计/真值−1)中位{r['relative_error_median']:+.1%}，"
                f"范围[{r['relative_error_minimum']:+.1%}, {r['relative_error_maximum']:+.1%}]。")
        no_amp = synthetic[synthetic.method.eq("trained_gain_no_amplitude") & synthetic.status.eq("completed")]
        with_amp = synthetic[synthetic.method.eq("trained_gain_amplitude") & synthetic.status.eq("completed")]
        matched = with_amp.merge(no_amp, on="group", suffixes=("_with", "_without"), validate="one_to_one")
        if len(matched):
            delta = matched.relative_gain_error_with - matched.relative_gain_error_without
            direction = "全部上移" if delta.gt(0).all() else "全部下移" if delta.lt(0).all() else "方向不一致"
            body += (f"同组配对{len(matched)}/{expected_synthetic}：加入幅度先验后β偏差{direction}，"
                f"误差变化中位{delta.median()*100:+.1f}个百分点。")
            if delta.gt(0).all():
                body += ("这项负结果表明幅度先验下的有效增益恢复出现系统上偏。"
                    "幅度惩罚压低r而提高β可部分维持血管输入β×r，是与该结果一致、受结果支持的补偿解释；"
                    "仅凭这些拟合不能证明它是唯一机制，也不能把更平稳的状态或更大的正则信息当作无偏生理参数恢复。")
        body += "这仍是给定生成模型与观测坐标的合成诊断，不是独立标定的被试生理估计；失败训练组不进入误差统计，但保留完整分母。\n\n"
    else:
        body += "合成训练选择尚未全部终态，暂不生成完整面板的β真值偏差结论。\n\n"
    steps = config["tensor"]["steps"]
    dt = config["tensor"]["dt_s"]
    amplitude = config["regularization"]["driver_amplitude_weight"]
    baseline_samples = int(round(5. / dt))
    special_squared = steps / baseline_samples
    ordinary_shrink = 1. / (1. + amplitude*dt)
    special_shrink = special_squared / (special_squared + amplitude*dt)
    body += ("简化算子解释：仅保留全观测EEG数据项与幅度ridge，忽略Hb、曲率和初态耦合时，"
        f"本合同基线算子B=I−1wᵀ使用首{baseline_samples}/{steps}点；"
        f"其奇异值为{steps-2}个1、1个√{special_squared:g}和1个0。"
        "由于driver prior SD来自同一EEG训练尺度/loading，各可观测奇异模态的收缩因子为"
        "σ²/(σ²+λ·dt)，这里σ指B的奇异值而非噪声SD。"
        f"本合同λ={amplitude:g}、dt={dt:g}秒，对应普通模态{ordinary_shrink:.3g}、"
        f"特殊模态{special_shrink:.3g}，DC由零中心ridge锚定为0。"
        f"若主要血管输入βr近似保持，普通模态的收缩可伴随β约放大{1./ordinary_shrink:.3g}倍，"
        "与上述合成偏差方向及量级相容；这是简化机制解释，非完整联合模型的精确预测或唯一原因。"
        "实际数据权重使用训练信号SD而非噪声SD，且未对时间协方差白化，故本目标应称标准化数据误差加工程ridge，"
        "不能将其当作校准likelihood，也不能据此断言生理先验本身失效。\n\n")
    for (kind, method), part in selections.groupby(["kind", "method"], sort=False):
        body += f"### {kind} · {labels[method]}\n\n"
        display = pd.DataFrame([{"组":short(r.group), "真β":r.true_gain, "选择β":r.gain,
            "成功起点":f"{r.completed_starts}/{r.expected_starts}", "边界":r.boundary,
            "起点β跨度":f"{r.converged_gain_min:.3g}–{r.converged_gain_max:.3g}"}
            for r in part.itertuples()])
        body += table(display) + "\n\n"
        information = pd.DataFrame([{"组":short(r.group), "数据":r.data_information,
            "加nuisance正则":r.regularized_information, "再加β先验":r.prior_information}
            for r in part.itertuples()])
        body += table(information) + "\n\n"
    body += ("全部起点的目标、β、收敛、边界、梯度、数据SSE和惩罚分解保留在gain_training_candidates.csv；"
        "同组起点跨度仅统计成功起点，失败数没有删除。完整库结果仍由training/start_*/result.json持有。\n\n")
    # Compact candidate tables retain all starts, split by group to avoid wide PDFs.
    for kind, part in candidates.groupby("kind", sort=False):
        body += f"### {kind} · 全部起点目标\n\n"
        display = []
        for (group, method), starts in part.groupby(["group", "method"], sort=False):
            row = {"组":short(group), "臂":labels[method].split()[0]}
            for r in starts.itertuples():
                objective = f"{r.objective:.4g}" if pd.notna(r.objective) else "—"
                status = "成功" if r.status == "completed" and r.converged else "失败"
                row[f"{'原起点' if config.get('continuation') else '起点'}{r.start_gain:g}"] = objective + "/" + status + ("*" if r.selected else "")
            display.append(row)
        body += table(pd.DataFrame(display)) + "\n\n"
    body += "星号标记所选起点；失败起点目标仅为审计信息，不能参与选择。\n\n"
    failed_starts = candidates[~(candidates.status.eq("completed") & candidates.converged)].copy()
    failed_starts["domain_category"] = np.select(
        [failed_starts.domain_rejections.eq(0), failed_starts.domain_rejections.gt(0)],
        ["无域拒绝", "有域拒绝"], default="域记录缺失")
    failure_diagnostics = []
    for (kind, method, category), part in failed_starts.groupby(["kind", "method", "domain_category"], sort=False):
        gradients = pd.to_numeric(part.projected_gradient, errors="coerce").dropna()
        failure_diagnostics.append(dict(kind=kind, method=method, domain_category=category,
            failed_starts=len(part), evaluation_budget=int(part.convergence_reason.eq("evaluation_budget").sum()),
            gradient_scored=len(gradients), gradient_minimum=gradients.min(),
            gradient_median=gradients.median(), gradient_maximum=gradients.max(),
            convergence_threshold=training["gradient_tolerance"]))
    failure_diagnostics = pd.DataFrame(failure_diagnostics)
    failure_diagnostics.to_csv(out / "gain_training_failures.csv", index=False)
    body += "### 训练起点失败的分类\n\n"
    if len(failure_diagnostics):
        for kind, part in failure_diagnostics.groupby("kind", sort=False):
            body += f"{kind}：\n\n" + table(pd.DataFrame([{
                "臂":labels[r.method].split()[0], "域拒绝":r.domain_category,
                "失败数":r.failed_starts, "预算耗尽":r.evaluation_budget,
                "梯度数":r.gradient_scored, "梯度中位":f"{r.gradient_median:.3g}",
                "梯度范围":f"{r.gradient_minimum:.3g}–{r.gradient_maximum:.3g}"}
                for r in part.itertuples()])) + "\n\n"
    else:
        body += "没有失败的训练起点。\n\n"
    batch = 18
    ordinary_steps = max(0, (training["max_evaluations"] - batch) // (2*batch))
    body += (f"梯度为库返回的projected_scaled_gradient_inf_norm，收敛阈值{training['gradient_tolerance']:.3g}；"
        "无梯度或域记录的起点不补零。domain_rejections读取每个独立任务内部starts[0]，表示优化提议曾被拒绝，"
        "不等于最终可行轨迹生理不合理；没有域拒绝也不证明已接近收敛。"
        f"本合同每起点max_evaluations={training['max_evaluations']}，按每组{batch}个trial计forward调用；"
        f"初始求值消耗{batch}次，若每个正常GN步只有一次完整提议与Jacobian复算，则每步{2*batch}次，"
        f"最多{ordinary_steps}个这样的步。max_iterations={training['max_iterations']}只是另一个上限，"
        "不保证能执行到该数；拒绝和额外线搜索会消耗预算。evaluation_budget不能一概解释成域失败或近收敛。\n\n")
    fig, axes = plt.subplots(1, 2, figsize=(12, 4))
    for ax, kind in zip(axes, ["synthetic", "measured"]):
        part = selections[selections.kind.eq(kind)]
        groups = list(part.group.unique())
        for method in part.method.unique():
            values = part[part.method.eq(method)].set_index("group").reindex(groups)
            ax.plot(range(len(groups)), pd.to_numeric(values.gain, errors="coerce"), "o-", label=labels[method])
        if kind == "synthetic" and len(part):
            actual = part.drop_duplicates("group").set_index("group").reindex(groups)
            ax.plot(range(len(groups)), pd.to_numeric(actual.true_gain, errors="coerce"), "k--", label="真β")
        ax.set_xticks(range(len(groups)), [short(g) for g in groups], rotation=55, ha="right", fontsize=6)
        ax.set_ylabel("训练冻结 β"); ax.set_title(kind)
        if len(part): ax.legend(fontsize=7)
    save(fig, "gain_training", "每组训练冻结的有效β；缺失/失败不补值。合成β为预设测试点，非人口生理范围。")
    body += "## 驱动与缺口状态诊断\n\n"
    diagnostics = []
    for (kind, method, mode), part in states.groupby(["kind", "method", "mode"], sort=False):
        good = part[part.converged & part.finite_states]
        denominator = len(data[(data.kind == kind) & (data.method == method) & (data['mode'] == mode)])
        row = dict(kind=kind, method=method, mode=mode, successful=len(good), observed=denominator)
        for field in ["driver_mean", "driver_rms", "driver_center_rms", "driver_difference_rms", "driver_center_difference_rms"]:
            values = pd.to_numeric(good.get(field, pd.Series(dtype=float)), errors="coerce")
            row[field+"_median"] = values.median(); row[field+"_maximum"] = values.max()
        for name in ["r", "s", "f", "v", "p", "q"]:
            for suffix, operation in [("minimum", "min"), ("maximum", "max"), ("center_minimum", "min"), ("center_maximum", "max")]:
                values = pd.to_numeric(good.get(name+"_"+suffix, pd.Series(dtype=float)), errors="coerce")
                row[name+"_"+suffix] = getattr(values, operation)()
        diagnostics.append(row)
    diagnostics = pd.DataFrame(diagnostics)
    diagnostics.to_csv(out / "gain_driver_state_diagnostics.csv", index=False)
    for (kind, mode), part in diagnostics.groupby(["kind", "mode"], sort=False):
        body += f"### {kind} · {mode}\n\n"
        body += table(pd.DataFrame([{"臂":labels[r.method].split()[0], "成功/记录":f"{r.successful}/{r.observed}",
            "r RMS中位":r.driver_rms_median, "缺口r RMS最大":r.driver_center_rms_maximum,
            "缺口Δr RMS最大":r.driver_center_difference_rms_maximum,
            "全路径f最小":r.f_minimum, "全路径f最大":r.f_maximum} for r in part.itertuples()])) + "\n\n"
    body += ("这里的Δr是相邻采样差分，RMS不是导数也不是生理正常范围；高差分RMS提示残余振荡。"
        "中心16点均以相同位置统计；full模式的中心统计不表示数据被隐藏。六状态全路径/中心极值和失败最后有效轨迹详见state_ranges.csv，"
        "成功子集汇总见gain_driver_state_diagnostics.csv。\n\n")
    # Factorial contrasts require the same four successful identities.
    identity = ["kind", "group", "mode", "trial", "subject", "session"]
    common = None
    for letter, method in zip("ABCD", config["methods"]):
        part = data[(data.method == method) & data.status.eq("completed") & truth(data.converged)]
        fields = ["nmse_"+c for c in COMPONENTS]
        part = part.reindex(columns=identity+fields).rename(columns={f:f+"_"+letter for f in fields})
        common = part if common is None else common.merge(part, on=identity, validate="one_to_one")
    factorial = []
    for (kind, mode), part in common.groupby(["kind", "mode"], sort=False):
        expected = len(data[(data.kind == kind) & (data['mode'] == mode) & (data.method == config['methods'][0])])
        for component in COMPONENTS:
            f = "nmse_"+component
            if part[f+"_A"].isna().all(): continue
            for contrast, values in [("B−A", part[f+"_B"]-part[f+"_A"]),
                ("C−A", part[f+"_C"]-part[f+"_A"]), ("D−B", part[f+"_D"]-part[f+"_B"]),
                ("D−C", part[f+"_D"]-part[f+"_C"]),
                ("交互 D−C−B+A", part[f+"_D"]-part[f+"_C"]-part[f+"_B"]+part[f+"_A"])]:
                frame = part[['subject','session']].assign(delta=values)
                value = frame.groupby(['subject','session']).delta.mean().groupby('subject').mean().mean()
                factorial.append(dict(kind=kind, mode=mode, component=component, contrast=contrast,
                    common_success=len(part), observed_reference=expected, delta_nmse=value))
    factorial = pd.DataFrame(factorial)
    factorial.to_csv(out / 'paired_factorial.csv', index=False)
    body += "## 四臂同身份比较\n\n以下全部使用四臂共同成功的相同身份；ΔNMSE按session及subject等权，负值表示该对比误差下降。B−A隔离固定β下幅度先验，D−B隔离有幅度先验时训练β，交互不能由各臂不同成功子集推断。无四臂共同成功的分组没有可估计对比；其失败仍保留在前文完整分母。\n\n"
    if len(factorial):
        for (kind, mode), part in factorial.groupby(['kind','mode'], sort=False):
            body += f"### {kind} · {mode}\n\n" + table(part[['component','contrast','common_success','observed_reference','delta_nmse']].rename(
                columns={'component':'分量','contrast':'比较','common_success':'共同成功','observed_reference':'A记录数','delta_nmse':'ΔNMSE'})) + "\n\n"
    # Hidden EEG states on the exact identities chosen from full-observation A.
    fig, axes = plt.subplots(len(selected), 3, figsize=(12, 3.6*len(selected)), squeeze=False)
    for i, (name, sample) in enumerate(selected):
        group, trial = sample.group, int(sample.trial)
        n = config['tensor']['steps']
        t = np.arange(n)*config['tensor']['dt_s'] - (5 if sample.kind == 'measured' else 0)
        target_drawn = False
        unavailable_notes = []
        for method in config["methods"]:
            path = run / "cells" / f"{group}__{method}__center_EEG" / "trajectories.npz"
            rows = data[(data.group == group) & (data.method == method) & (data['mode'] == "center_EEG") & (data.trial == trial)]
            status = rows.status.iloc[0] if len(rows) else "missing"
            if not path.exists():
                unavailable_notes.append(f"{labels[method].split()[0]}: {status} / 未估计")
                for ax in axes[i]: ax.plot([], [], ":", label=labels[method]+" 未估计")
                continue
            with np.load(path, allow_pickle=False) as arrays:
                matches = np.flatnonzero(arrays["trial_indices"] == trial)
                if not len(matches): raise ValueError("Fixed hidden example missing")
                j = int(matches[0]); n = arrays['prediction'].shape[1]
                t = np.arange(n)*config['tensor']['dt_s'] - (5 if sample.kind == 'measured' else 0)
                style = '-' if status == 'completed' else ':'
                if not target_drawn and 'target' in arrays:
                    axes[i,0].plot(t, arrays['target'][j,:,0], 'k', lw=1.3, label='目标')
                    target_drawn = True
                if not np.isfinite(arrays['states'][j]).any():
                    unavailable_notes.append(f"{labels[method].split()[0]}: {status} / 未估计")
                for ax, value in zip(axes[i], [arrays['prediction'][j,:,0], arrays['states'][j,:,0], arrays['states'][j,:,2]]):
                    suffix = ' 未估计' if not np.isfinite(value).any() else (' 失败迭代' if status != 'completed' else '')
                    ax.plot(t, value, style, lw=.9, label=labels[method]+suffix)
        if unavailable_notes:
            axes[i,1].text(.02, .98, "\n".join(unavailable_notes), transform=axes[i,1].transAxes,
                ha='left', va='top', fontsize=6, bbox=dict(facecolor='white', alpha=.85, edgecolor='none'))
        for ax, ylabel in zip(axes[i], ['EEG', 'driver r', 'flow f']):
            ax.axvspan(t[(n-16)//2], t[(n+16)//2-1], color='grey', alpha=.12)
            ax.set_ylabel(ylabel); ax.set_xlabel('相对事件 / 秒' if sample.kind == 'measured' else '秒')
        axes[i,0].set_title(f'{name}: {short(group)} / trial {trial}')
        axes[i,0].legend(fontsize=6)
    save(fig, 'gain_hidden_fixed_examples', '沿用full模式A基线选出的中位/最差身份，显示中心EEG隐藏时预测、驱动和血流；灰区为隐藏段，失败最后有效轨迹为虚线，未按候选结果重挑。')
    # These identities are fixed in advance, independent of reconstruction scores.
    truth_groups = ["synthetic_gain1_slow_r0", "synthetic_gain1_mixed_r0"]
    truth_methods = ["trained_gain_no_amplitude", "trained_gain_amplitude"]
    fig, axes = plt.subplots(3, 2, figsize=(11, 8), sharex="col")
    truth_notes = []
    for column, group in enumerate(truth_groups):
        trial = 0
        prepared_path = run / "prepared" / f"{group}.npz"
        axes[0,column].set_title(f"{group} / trial {trial}", fontsize=9)
        if not prepared_path.exists():
            for ax in axes[:,column]:
                ax.text(.5, .5, "准备数据缺失；固定身份不替换", ha="center", va="center", transform=ax.transAxes)
            truth_notes.append(f"{group}: prepared缺失")
            continue
        with np.load(prepared_path, allow_pickle=False) as prepared:
            n = prepared["target"].shape[1]
            t = np.arange(n)*config["tensor"]["dt_s"]
            axes[0,column].plot(t, prepared["target"][trial,:,0], "k", lw=1.15, label="含噪EEG目标")
            axes[0,column].plot(t, prepared["clean"][trial,:,0], color=".55", ls=":", lw=1., label="无噪处理EEG")
            axes[1,column].plot(t, prepared["truth"][trial], "k", lw=1.4, label="driver真值")
            axes[2,column].plot(t, prepared["truth_states"][trial,:,2], "k", lw=1.4, label="flow真值")
        for method, color in zip(truth_methods, ["tab:blue", "tab:orange"]):
            rows = data[(data.group == group) & (data.method == method) & (data['mode'] == "full") & (data.trial == trial)]
            status = rows.status.iloc[0] if len(rows) else "missing"
            truth_notes.append(f"{group} / {labels[method]}: {status}")
            path = run / "cells" / f"{group}__{method}__full" / "trajectories.npz"
            if not path.exists():
                for ax in axes[:,column]:
                    ax.plot([], [], color=color, ls=":", label=labels[method]+" 无轨迹")
                continue
            with np.load(path, allow_pickle=False) as arrays:
                matches = np.flatnonzero(arrays["trial_indices"] == trial)
                if not len(matches):
                    raise ValueError("Prespecified synthetic truth example missing trial 0")
                index = int(matches[0])
                values = [arrays["prediction"][index,:,0], arrays["states"][index,:,0], arrays["states"][index,:,2]]
                style = "-" if status == "completed" else ":"
                for ax, value in zip(axes[:,column], values):
                    suffix = " 未估计" if not np.isfinite(value).any() else ("" if status == "completed" else " 失败最后有效轨迹")
                    ax.plot(t, value, color=color, ls=style, lw=1., label=labels[method]+suffix)
        for ax, ylabel in zip(axes[:,column], ["处理坐标 EEG", "driver r", "flow f"]):
            ax.set_ylabel(ylabel); ax.legend(fontsize=6)
        axes[2,column].set_xlabel("秒")
    save(fig, "gain_synthetic_truth_examples", "事先固定β真值1、replicate0的slow/mixed条件，各trial0；首行黑线为含噪处理EEG目标，灰虚线为无噪处理EEG；后二行黑线为无噪driver/flow真值。仅叠加C/D全观测结果，失败最后有效轨迹为虚线，未做去均值、幅度或时间对齐。")
    body += ("## 固定合成真轨迹样例\n\n选样身份固定为synthetic_gain1_slow_r0和synthetic_gain1_mixed_r0，各trial0，"
        "与误差排序或候选效果无关。EEG同时显示含噪观测目标及无噪处理轨迹，r/f显示生成真值；仅对比C/D。"
        "没有重新拟合或事后去均值、幅度、时间对齐。失败、缺失保留原身份，不以较好样例替换。\n\n"
        + "；".join(truth_notes) + "。\n\n")
    if config.get("continuation"):
        body = body.replace("个独立起点；只在全部起点终态后按收敛目标选择。", "条原起点谱系；其中可能继承或从父解热启动，只在全部谱系终态后按同一收敛目标选择。")
        body = body.replace("本合同每起点max_evaluations=", "本轮仅热启动段每起点max_evaluations=")
        body += "上述新预算不适用于继承的终态起点；继承记录保留父运行的实际预算与调用数。合成偏差与固定合成曲线沿用父证据，不能当作本轮独立复现。\n\n"
        body += gain_continuation_details(run, out, config, data, labels, save)
        body += continuation_parent_comparison(run, out, config, data, labels)
    return body


def fixed_roi_training_failure_details(run, out, config, candidates, save):
    """Map local training failures to original trials; never score last-valid fits."""
    rows=[]
    failed=candidates[~(candidates.status.eq('completed')&candidates.converged)]
    for item in failed.itertuples():
        folder=run/'training'/f'{item.group}__{item.method}'/f'start_{item.start}'
        path=folder/'result.json'
        if not path.exists():continue
        result=json.loads(path.read_text());attempt=(result.get('starts') or [{}])[0]
        training=result.get('training_trials',[]);events=attempt.get('evaluation_failures',[])
        by_local={}
        for event in events:
            local=int(event['trial'])
            if not 0<=local<len(training):raise ValueError('Training failure lacks local-to-global trial mapping')
            by_local.setdefault(local,[]).append(event)
        arrays_path=folder/'trajectories.npz'
        saved_states=saved_indices=None
        if arrays_path.exists():
            with np.load(arrays_path,allow_pickle=False) as arrays:
                saved_states=arrays['states'].copy();saved_indices=arrays['trial_indices'].copy()
            if training and not np.array_equal(saved_indices,np.asarray(training)):
                raise ValueError('Failed training trajectory trial order differs from training_trials')
        for local,trial in enumerate(training):
            ev=by_local.get(local,[])
            row=dict(kind=item.kind,group=item.group,subject=item.subject,method=item.method,start=item.start,
                status=item.status,local_trial=local,global_trial=int(trial),evaluation_failures=len(ev),
                forward_failures=sum(not e.get('derivative',False) for e in ev),
                derivative_failures=sum(bool(e.get('derivative',False)) for e in ev),
                errors='; '.join(sorted(set(e.get('error','') for e in ev))),
                start_domain_rejections=attempt.get('domain_rejections'),trajectory=str(arrays_path),
                finite_last_state=False,final_flow_min=np.nan,minimum_time_from_window_start_s=np.nan,
                minimum_time_relative_event_s=np.nan)
            if saved_states is not None and np.isfinite(saved_states[local]).all():
                flow=saved_states[local,:,2];index=int(np.argmin(flow));time=index*config['tensor']['dt_s']
                row.update(finite_last_state=True,final_flow_min=float(flow[index]),minimum_time_from_window_start_s=time,
                    minimum_time_relative_event_s=time-(5. if item.kind=='measured' else 0.))
            rows.append(row)
    frame=pd.DataFrame(rows)
    frame.to_csv(out/'failure_by_trial.csv',index=False)
    body='\n## 失败训练的试次定位\n\n'
    body+=('evaluation_failures中的trial是训练集合局部索引，本表经该起点training_trials映射为原始global trial；'
        '不能把局部序号直接当作原试次。统计只覆盖失败训练起点；forward求值异常与derivative异常分列，'
        '重复求值次数不是独立失败试次数。末有效轨迹不是有效参数估计，不计入验证误差或合成恢复主分数。\n\n')
    if not len(frame):return body+'没有可映射的失败训练轨迹。\n\n'
    counted=frame[frame.evaluation_failures.gt(0)]
    if len(counted):
        summary=[]
        for (subject,trial),part in counted.groupby(['subject','global_trial'],sort=False):
            summary.append({'被试':str(subject).replace('subject_','S'),'原trial':trial,'涉及失败起点':len(part),
                'forward失败':int(part.forward_failures.sum()),'derivative失败':int(part.derivative_failures.sum()),
                '末有效f最小':f'{part.final_flow_min.min():.3g}'})
        body+=table(pd.DataFrame(summary))+'\n\n'
    else:body+='这些失败起点没有记录trial级求值异常；不能据此归因为某个试次域失败。\n\n'
    measured=frame[frame.kind.eq('measured')&frame.finite_last_state].sort_values(
        ['subject','final_flow_min','group','method','start','global_trial'],kind='stable')
    selected=measured.groupby('subject',sort=False).head(1)
    selected.to_csv(out/'failed_training_selected_examples.csv',index=False)
    if len(selected):
        fig,axes=plt.subplots(3,len(selected),figsize=(4.2*len(selected),8),squeeze=False,sharex='col')
        for column,item in enumerate(selected.itertuples()):
            with np.load(item.trajectory,allow_pickle=False) as arrays:
                matches=np.flatnonzero(arrays['trial_indices']==item.global_trial)
                if len(matches)!=1:raise ValueError('Failed training selected global trial is not unique')
                index=int(matches[0]);t=np.arange(arrays['states'].shape[1])*config['tensor']['dt_s']-5.
                for row,component in enumerate([1,2]):
                    axes[row,column].plot(t,arrays['target'][index,:,component],'k',lw=1.2,label='训练目标')
                    axes[row,column].plot(t,arrays['prediction'][index,:,component],color='tab:red',ls='--',lw=1.,label='失败末有效预测')
                    axes[row,column].set_ylabel(['HbO','HbR'][row]);axes[row,column].legend(fontsize=6)
                flow=arrays['states'][index,:,2]
                axes[2,column].plot(t,flow,color='tab:red',ls='--',lw=1.,label='失败末有效flow')
                axes[2,column].axhline(0,color='grey',lw=.7)
                axes[2,column].axvline(item.minimum_time_relative_event_s,color='grey',ls=':',lw=.7)
                axes[2,column].set_ylabel('flow f');axes[2,column].set_xlabel('相对事件 / 秒')
                axes[2,column].text(.03,.05,f'min f={item.final_flow_min:.3g}',transform=axes[2,column].transAxes,fontsize=8)
                axes[0,column].set_title(f"{str(item.subject).replace('subject_','S')} / trial{item.global_trial}\n{item.group} / {'先验τ' if item.method.endswith('prior_tau') else '训练τ'} / start{item.start}",fontsize=8)
                for ax in axes[:,column]:ax.axvline(0,color='grey',ls=':',lw=.6)
        save(fig,'fixed_roi_failed_training_states','每被试从全部失败训练起点的末有效轨迹中，按最小flow固定选一条（相同最小值按身份排序）。这里只展示病态失败迭代，不是有效估计，不计主评分，也未替换成功样例；灰竖线定位最小flow。')
        body+=table(pd.DataFrame([{'被试':str(r.subject).replace('subject_','S'),'组':str(r.group).replace('subject_','S'),
            '原trial':r.global_trial,'起点':r.start,'最小f':f'{r.final_flow_min:.3g}',
            '窗内时间s':r.minimum_time_from_window_start_s,'事件时间s':r.minimum_time_relative_event_s}
            for r in selected.itertuples()]))+'\n\n'
    body+=('共享组目标需要所有训练trial共同可行；上述集中求值异常和接近零的末有效flow支持“个别轨迹阻碍共享组优化”的解释，'
        '但不证明输入试次本身无效，也不证明删除它们即可修复模型。初态、驱动、固定β及未经绝对标定的幅度映射仍可能共同作用。'
        '本轮未删除、替换或重新拟合这些身份；failure_by_trial.csv同时保留没有求值异常但属于失败起点的训练trial。\n\n')
    return body


def fixed_roi_tau_details(run, out, config, data, states, labels, save):
    """Describe shared nonlinear tau estimates; optimizer traces are not profiles."""
    methods = config['methods']; initial = config['schema'] == 'shared_driver_fixed_roi_initial_v1'
    flow = config['schema'] in ('shared_driver_fixed_roi_flow_v1','shared_driver_fixed_roi_initial_v1')
    trainable = [m for m in methods if m not in ('fixed_roi_fixed_tau','fixed_tau_logflow1','fixed_tau_logflow1_tied')]
    options = config['tau_training']; candidates, selections, traces, mapping = [], [], [], []
    for group in data.group.unique():
        prep = json.loads((run/'prepared'/f'{group}.json').read_text()); spec = prep['spec']
        if spec['kind'] == 'measured' and prep['status'] == 'completed':
            projection = prep['metadata']['projection']; roi = prep['metadata']['fixed_roi']
            if projection['fnirs_pair'] != config['fixed_roi']['pair_index'] or roi['pair_name'] != config['fixed_roi']['pair_name']:
                raise ValueError('Fixed ROI metadata differs from declared AF7Fp1 target')
            mapping.append(dict(group=group, pair=roi['pair_name'], pair_index=projection['fnirs_pair'],
                parent_pair_index=roi.get('parent_pair_index'), EEG_factor=projection['eeg_factor'],
                Hb_factor=projection['fnirs_factor'], common_training_MAD=roi.get('common_training_MAD')))
        for method in trainable:
            folder = run/'training'/f'{group}__{method}'
            path = folder/'selection.json'; choice = json.loads(path.read_text()) if path.exists() else {}
            selected = choice.get('selected_training_start'); local = []
            for start, start_tau in enumerate(options['start_values']):
                path = folder/f'start_{start}'/'result.json'; result = json.loads(path.read_text()) if path.exists() else {}
                info = result.get('information', {}); attempt = (result.get('starts') or [{}])[0]
                row = dict(kind=spec['kind'], group=group, subject=spec['subject'], outer=spec['outer'],
                    method=method, start=start, start_tau=start_tau, true_tau=spec.get('true_tau'),
                    true_initial_log_p_over_v=spec.get('true_initial_log_p_over_v'),
                    initial_state_constraint=result.get('initial_state_constraint'),
                    nuisance_free_parameters_per_trial=result.get('nuisance_free_parameters_per_trial'),
                    selected=start == selected, status=result.get('status','missing'), converged=bool(result.get('converged',False)),
                    tau=result.get('parameter_value'), objective=result.get('objective'), boundary=result.get('boundary_status','N/A'),
                    gradient=result.get('projected_scaled_gradient_inf_norm',attempt.get('projected_scaled_gradient_inf_norm')),
                    reason=result.get('convergence_reason',attempt.get('convergence_reason')),
                    optimization_evaluations=result.get('optimization_evaluations'), domain_rejections=attempt.get('domain_rejections'),
                    weighted_data_sse=result.get('weighted_data_sse'), parameter_prior_cost=result.get('parameter_prior_cost'),
                    initial_state_cost=result.get('initial_state_penalty_cost'),
                    flow_prior_cost=result.get('flow_prior_cost'),flow_prior_weight=result.get('flow_prior_weight'),
                    flow_prior_log_sd=result.get('flow_prior_log_sd'),inheritance=json.dumps(result.get('inheritance',{})),
                    derivative_rejections=attempt.get('derivative_rejections'),integration_checks=json.dumps(result.get('integration_checks',[])),
                    data_information=info.get('data_only',{}).get('value'),
                    regularized_information=info.get('nuisance_regularized',{}).get('value'), total_information=info.get('with_parameter_prior'))
                candidates.append(row); local.append(row)
                for point in result.get('trace',[]): traces.append(dict(group=group,kind=spec['kind'],method=method,start=start,**point))
            chosen = [r for r in local if r['selected']]
            if choice.get('status') == 'completed' and (len(chosen) != 1 or chosen[0]['status'] != 'completed' or not chosen[0]['converged']):
                raise ValueError('Tau selection lacks one converged declared start')
            chosen = chosen[0] if chosen else {}
            good = [r for r in local if r['status'] == 'completed' and r['converged']]
            selections.append(dict(kind=spec['kind'], group=group, subject=spec['subject'], outer=spec['outer'], method=method,
                status=choice.get('status','missing'), tau=choice.get('parameter_value'), true_tau=spec.get('true_tau'),
                true_initial_log_p_over_v=spec.get('true_initial_log_p_over_v'),
                successful_starts=len(good), expected_starts=len(options['start_values']), selected_start=selected,
                tau_min=min([r['tau'] for r in good],default=np.nan), tau_max=max([r['tau'] for r in good],default=np.nan),
                **{key:chosen.get(key) for key in ['boundary','objective','gradient','data_information','regularized_information','total_information']}))
    candidates = pd.DataFrame(candidates); selections = pd.DataFrame(selections); traces = pd.DataFrame(traces)
    candidates.to_csv(out/'tau_training_candidates.csv',index=False)
    selections.to_csv(out/'tau_training_selection.csv',index=False)
    traces.to_csv(out/'tau_training_traces.csv',index=False)
    pd.DataFrame(mapping).to_csv(out/'fixed_roi_mapping.csv',index=False)
    short = lambda x:str(x).replace('subject_','S').replace('synthetic_tau','T').replace('slow','慢').replace('mixed','混')
    method_short = labels if flow else {methods[0]:'τ2',methods[1]:'训练τ',methods[2]:'先验τ'}
    body = ('\n## 固定观测位置与解释范围\n\nAF7Fp1（pair0）是预声明的共同帽布局位置，不是个体皮层共定位。'
        'EEG投影保持父规范，Hb目标和训练SD因固定位置而改变；仅本轮同一新目标下重算的τ2作为配对基线。'
        '旧parent的τ和误差只可作为背景，未并入本轮paired scores。目标和预测共用原生观测算子；初始前5秒baseline仍不是潜在静息约束。'
        '所有cell保存的目标/SD与prepared逐值核对，固定ROI元数据另行核对。\n\n')
    if mapping:
        body += table(pd.DataFrame([{'组':short(r['group']),'本轮Hb对':r['pair'],'旧对索引':r['parent_pair_index'],
            'EEG因子':r['EEG_factor'],'新Hb因子':r['Hb_factor']} for r in mapping]))+'\n\n'
    body += ('## 共享τ、多起点与局部信息\n\nτ仅在18条训练轨迹上共享拟合，验证时冻结；'
        f'搜索范围{options["bounds"]}秒，' + ('各新臂无τ先验；log-flow为工程软锚。' if flow else f'先验臂log SD={options["prior_log_sd"]:.4g}，均为工程设置。') +
        '多起点终点散点与迭代trace不是逐点重优化的profile；局部投影信息不是置信区间、校准后验或全局唯一可辨识性证明。'
        '分别列data-only、消除带正则nuisance后的信息、再加入参数先验的信息，不能把正则带来的信息冒充数据支持。\n\n')
    for (kind,method),part in selections.groupby(['kind','method'],sort=False):
        body += f'### {kind} · {method_short[method]}\n\n'
        body += table(pd.DataFrame([{'组':short(r.group),'真τ':r.true_tau,'选择τ':r.tau,
            '成功起点':f'{r.successful_starts}/{r.expected_starts}','起点τ范围':f'{r.tau_min:.3g}–{r.tau_max:.3g}' if r.successful_starts else '—',
            '边界':r.boundary or '—'} for r in part.itertuples()]))+'\n\n'
        body += table(pd.DataFrame([{'组':short(r.group),'训练目标':r.objective,'数据':r.data_information,
            '加nuisance正则':r.regularized_information,'再加τ先验':r.total_information} for r in part.itertuples()]))+'\n\n'
    failures = candidates[~(candidates.status.eq('completed')&candidates.converged)].copy()
    body += f'全部训练起点{len(candidates)}，未成功{len(failures)}；失败起点保留终点目标、梯度、域拒绝及预算原因，不参与参数选择。'
    body += f'收敛阈值{options["gradient_tolerance"]:.3g}；目标/梯度/预算明细见tau_training_candidates.csv。\n\n'
    if len(failures):
        display=failures.groupby(['kind','method','reason'],dropna=False).size().rename('起点数').reset_index()
        display['method']=display.method.map(method_short)
        body+=table(display.rename(columns={'kind':'数据','method':'方法','reason':'停止原因'}))+'\n\n'
    variation=[]
    for (subject,method),part in selections[selections.kind.eq('measured')].groupby(['subject','method'],sort=False):
        good=pd.to_numeric(part.loc[part.status.eq('completed'),'tau'],errors='coerce').dropna()
        variation.append(dict(subject=subject,method=method,successful_folds=len(good),expected_folds=len(config['outer_folds']),
            tau_min=good.min(),tau_max=good.max(),geometric_mean=float(np.exp(np.log(good).mean())),log_tau_sd=np.log(good).std()))
    pd.DataFrame(variation).to_csv(out/'tau_subject_fold_variation.csv',index=False)
    body+='### 被试与折间变化（描述性）\n\n'
    if variation:
        body+=table(pd.DataFrame([{'被试':short(r['subject']),'方法':method_short[r['method']],
            '成功折':f"{r['successful_folds']}/{r['expected_folds']}",'τ范围':f"{r['tau_min']:.3g}–{r['tau_max']:.3g}" if r['successful_folds'] else '—',
            '几何均值':r['geometric_mean'],'折间log SD':r['log_tau_sd']} for r in variation]))+'\n\n'
    body+='折间训练集重叠，四折不能视为四次独立重复；固定位置减少选对混杂，但不消除观测规范、初态或固定β的补偿。没有足够数据支持人口参数分布或唯一的被试生理差异。\n\n'
    fig,axes=plt.subplots(2,2,figsize=(12,8))
    measured=selections[selections.kind.eq('measured')]; groups=list(measured.group.unique())
    for method in trainable:
        part=measured[measured.method.eq(method)].set_index('group').reindex(groups)
        axes[0,0].plot(range(len(groups)),pd.to_numeric(part.tau,errors='coerce'),'o-',label=method_short[method])
        synthetic=selections[selections.kind.eq('synthetic')&selections.method.eq(method)&selections.status.eq('completed')]
        axes[0,1].scatter(synthetic.true_tau,synthetic.tau,label=method_short[method],s=25)
        good=selections[selections.method.eq(method)&selections.status.eq('completed')]
        axes[1,0].scatter(pd.to_numeric(good.data_information),pd.to_numeric(good.regularized_information),label=method_short[method],s=20)
    axes[0,0].set_xticks(range(len(groups)),[short(g) for g in groups],rotation=50,ha='right',fontsize=6)
    axes[0,0].set_ylabel('训练冻结τ / 秒');axes[0,0].set_title('同位置：被试×折');axes[0,0].legend(fontsize=7)
    if not len(measured):
        axes[0,0].text(.5,.5,'实测阶段未纳入本报告',ha='center',va='center',transform=axes[0,0].transAxes)
    axes[0,1].plot(config['synthetic']['taus_s'],config['synthetic']['taus_s'],'k:');axes[0,1].set_xlabel('合成真τ / 秒');axes[0,1].set_ylabel('训练估计τ / 秒');axes[0,1].legend(fontsize=7)
    axes[1,0].set_xscale('symlog',linthresh=1e-8);axes[1,0].set_yscale('symlog',linthresh=1e-8)
    axes[1,0].set_xlabel('data-only局部信息');axes[1,0].set_ylabel('nuisance正则化局部信息');axes[1,0].legend(fontsize=7)
    for (group,method),part in candidates.groupby(['group','method'],sort=False):
        values=pd.to_numeric(part.objective,errors='coerce');minimum=values.min()
        for ok,marker in [(True,'o'),(False,'x')]:
            take=part.converged.eq(ok)&part.objective.notna()
            axes[1,1].scatter(pd.to_numeric(part.loc[take,'tau']),values[take]-minimum,marker=marker,s=14,alpha=.65,
                color=plt.get_cmap('tab10')(trainable.index(method)))
    axes[1,1].set_xlabel('多起点终点τ / 秒');axes[1,1].set_ylabel('终点目标 − 同组同臂最小终点目标');axes[1,1].set_title('圆：收敛；叉：失败。不是profile')
    save(fig,'fixed_roi_tau_parameters','被试/折与合成τ恢复、局部信息和多起点终点；未收敛目标仅供审计。信息与终点散点均不构成profile或置信区间。')
    if len(traces):
        for batch in range(0,len(trainable),2):
            plotted=trainable[batch:batch+2]
            fig,axes=plt.subplots(3,len(plotted),figsize=(12,9),sharex='col',squeeze=False)
            for col,method in enumerate(plotted):
                trace_kind='measured' if 'measured' in data.kind.unique() else 'synthetic'
                part=traces[traces.kind.eq(trace_kind)&traces.method.eq(method)]
                for index,((group,start),points) in enumerate(part.groupby(['group','start'],sort=False)):
                    points=points.sort_values(['evaluations','iteration'],kind='stable');color=plt.get_cmap('tab20')(index//3%20)
                    for ax,field in zip(axes[:,col],['objective','projected_scaled_gradient_inf_norm','parameter_value']):
                        ax.plot(points.evaluations,points[field],color=color,ls=['-','--',':'][int(start)%3],lw=.8,alpha=.7)
                axes[0,col].set_title(method_short[method]);axes[0,col].set_ylabel('训练目标')
                axes[1,col].set_yscale('log');axes[1,col].axhline(options['gradient_tolerance'],color='red',ls=':');axes[1,col].set_ylabel('projected scaled gradient')
                axes[2,col].set_ylabel('τ / 秒');axes[2,col].set_xlabel('单起点优化trial forward调用')
            save(fig,'fixed_roi_tau_traces'+(f'_{batch//2+1}' if len(trainable)>2 else ''),('实测' if 'measured' in data.kind.unique() else '合成')+'训练迭代轨迹：目标、实际梯度和τ；红线为原收敛阈值。各线为起点，不是profile扫描；完整身份见tau_training_traces.csv。')
    # Truth recovery uses the original state coordinates, without alignment.
    recovered=[]
    for (method,mode),part in states[states.kind.eq('synthetic')].groupby(['method','mode'],sort=False):
        good=part[part.converged & part.finite_states]
        row=dict(method=method,mode=mode,completed=len(good),expected=int(((data.kind=='synthetic')&(data.method==method)&(data['mode']==mode)).sum()))
        for name in ['r','s','f','v','p','q']:
            values=pd.to_numeric(good.get('truth_rmse_'+name,pd.Series(dtype=float)),errors='coerce')
            row[name+'_rmse']=float(np.sqrt(np.mean(values**2))) if len(values) else np.nan
        recovered.append(row)
    pd.DataFrame(recovered).to_csv(out/'fixed_roi_synthetic_state_recovery.csv',index=False)
    body+='### 合成状态恢复\n\n原坐标状态RMSE在成功trial内平方平均后取根，不做平移、幅度、符号或时间对齐；不是对静息常数归一化的误差。\n\n'
    for mode in config['modes']:
        part=[r for r in recovered if r['mode']==mode]
        body+=f'{mode}：\n\n'+table(pd.DataFrame([{'方法':method_short[r['method']],'成功/计划':f"{r['completed']}/{r['expected']}",
            **{name:r[name+'_rmse'] for name in ['r','s','f','v','p','q']}} for r in part]))+'\n\n'
    fig,axes=plt.subplots(3,3,figsize=(13,9),sharex='col');example_notes=[]
    for col,tau in enumerate(config['synthetic']['taus_s']):
        group=f'synthetic_tau{tau:g}_slow_r0'+('_logpv+0.00' if initial else '');trial=0;path=run/'prepared'/f'{group}.npz'
        axes[0,col].set_title(f'真τ={tau:g}s / slow / r0 / trial0',fontsize=9)
        if not path.exists():
            for ax in axes[:,col]:ax.text(.5,.5,'固定样例缺失，不替换',transform=ax.transAxes,ha='center')
            continue
        with np.load(path,allow_pickle=False) as a:
            t=np.arange(a['target'].shape[1])*config['tensor']['dt_s']
            axes[0,col].plot(t,a['target'][trial,:,1],color='.65',lw=.8,label='含噪HbO目标')
            for ax,value in zip(axes[:,col],[a['clean'][trial,:,1],a['truth'][trial],a['truth_states'][trial,:,2]]):ax.plot(t,value,'k',lw=1.3,label='生成真值')
        for method in methods:
            path=run/'cells'/f'{group}__{method}__full'/'trajectories.npz'
            row=data[(data.group==group)&(data.method==method)&(data['mode']=='full')&(data.trial==trial)]
            status=row.status.iloc[0] if len(row) else 'missing';example_notes.append(f'{group}/{method_short[method]}:{status}')
            if not path.exists():
                for ax in axes[:,col]:ax.plot([],[],':',label=method_short[method]+' 未估计')
                continue
            with np.load(path,allow_pickle=False) as a:
                indices=np.flatnonzero(a['trial_indices']==trial)
                if not len(indices):raise ValueError('Fixed three-tau example lacks trial0')
                i=int(indices[0]);values=[a['prediction'][i,:,1],a['states'][i,:,0],a['states'][i,:,2]]
                for ax,value in zip(axes[:,col],values):
                    suffix=' 未估计' if not np.isfinite(value).any() else (' 失败迭代' if status!='completed' else '')
                    ax.plot(t,value,'-' if status=='completed' else ':',lw=.9,label=method_short[method]+suffix)
        for ax,name in zip(axes[:,col],['处理HbO','driver r','flow f']):ax.set_ylabel(name);ax.legend(fontsize=5)
        axes[2,col].set_xlabel('秒')
    save(fig,'fixed_roi_three_tau_truth','固定选择slow、replicate0、trial0的τ=1/2/4三组；相同driver/初态/标准噪声配对，展示真值与预声明各臂估计。flow真值随τ不必改变，因为τ不直接进入s/f方程；Hb动态承载下游τ差异。失败保留，不做事后对齐。')
    if initial:body+='该通用三τ图仅示log(p0/v0)=0匹配层；±0.05错设层必须结合后文专门分层图解释。\n\n'
    body+='## 固定三τ真轨迹与隐藏检验边界\n\n三τ图事先按slow/replicate0/trial0身份固定，不按恢复质量选样。'+ '；'.join(example_notes)+'。\n\n'
    body+='full是联合重建，隐藏模式只评分被隐藏分量；参数训练与验证目标分开，不能把full重建当作隐藏恢复证据。中心插补使用两侧上下文，仍不是因果预测或独立被试验证；完整目标仅用于评分，own-context与τ2都在同一新观测坐标比较。\n\n'
    body += fixed_roi_training_failure_details(run, out, config, candidates, save)
    check_path = run/'unit_reexpression_check.json'
    if check_path.exists():
        check = json.loads(check_path.read_text())
        if check.get('status') == 'passed' and check.get('case_count') == 2:
            body += '单位重表达software check为2 passed：同步缩放目标、SD与观测算子保持物理拟合不变；这不是独立幅度标定或生理参数恢复证据。\n\n'
    return body


def decompose_processed_hbt(state, prediction, operator, p_loading):
    """Exact additive posthoc identity; the normalized cross term is signed."""
    volume=operator@(p_loading*(state[:,3]-1.))
    initial_difference=operator@(p_loading*(state[:,4]-state[:,3]))
    total=prediction[:,1]+prediction[:,2]
    error=float(np.max(np.abs(volume+initial_difference-total)))
    if not np.allclose(volume+initial_difference,total,rtol=1e-8,atol=1e-9):
        raise ValueError('Processed HbT decomposition disagrees with saved prediction')
    ms=float(np.mean(total**2));vms=float(np.mean(volume**2));dms=float(np.mean(initial_difference**2))
    cross=float(2*np.mean(volume*initial_difference))
    row=dict(total_rms=np.sqrt(ms),volume_rms=np.sqrt(vms),initial_difference_rms=np.sqrt(dms),
        signed_cross_term=cross,additivity_max_abs_difference=error,
        volume_rms_over_total=np.sqrt(vms/ms) if ms>1e-24 else np.nan,
        initial_difference_rms_over_total=np.sqrt(dms/ms) if ms>1e-24 else np.nan,
        signed_cross_over_total_ms=cross/ms if ms>1e-24 else np.nan)
    return row,volume,initial_difference,total


def flow_full_curve_atlas(run,out,config,data,labels,comparison_methods=None,*,both_states=False,context_label=""):
    """All 72 measured identities, fixed order; no score selection or new fit."""
    methods=comparison_methods or ['fixed_tau_logflow1','trained_tau_logflow1']
    names=[labels[m] for m in methods] if comparison_methods else ['固定τ/w1','共享τ/w1']
    identities=['group','subject','outer','session','trial']
    subset=data[data.kind.eq('measured')&data['mode'].eq('full')&data.method.isin(methods)]
    panels=subset[identities].drop_duplicates().sort_values(['subject','outer','trial'],kind='stable')
    if len(panels)!=72 or len(subset)!=72*len(methods):raise ValueError('Full curve atlas requires all arms and exactly 72 identities')
    if subset.duplicated(identities+['method']).any():raise ValueError('Duplicate atlas identity')
    folder=out/'curve_atlas';folder.mkdir();pdf=fitz.open();index=[]
    cache={}
    for (group,method),part in subset.groupby(['group','method']):
        path=run/'cells'/f'{group}__{method}__full'/'trajectories.npz'
        if path.exists():
            with np.load(path,allow_pickle=False) as a:cache[(group,method)]={k:a[k].copy() for k in ['target','prediction','states','trial_indices']}
    for start in range(0,72,3):
        page_number=start//3+1;separate_states=len(methods)>2
        ncols=3+len(methods) if separate_states else 4
        fig,axes=plt.subplots(3,ncols,figsize=(14 if separate_states else 12,8),sharex=True)
        page_labels=[]
        for row,identity in enumerate(panels.iloc[start:start+3].itertuples(index=False)):
            records=subset[subset.group.eq(identity.group)&subset.trial.eq(identity.trial)].set_index('method')
            t=np.arange(config['tensor']['steps'])*config['tensor']['dt_s']-5.
            text_status=[];actual_drawn=False
            for method,color,style in zip(methods,['.5','tab:blue','tab:orange'],['--','-',':']):
                record=records.loc[method];errors='/'.join(f'{record.get("nrmse_"+c,np.nan):.3f}' for c in COMPONENTS)
                text_status.append(names[methods.index(method)]+f' {record.status}, NRMSE={errors}')
                arrays=cache.get((identity.group,method))
                if arrays is None:continue
                where=np.flatnonzero(arrays['trial_indices']==identity.trial)
                if len(where)!=1:raise ValueError('Atlas trial missing')
                k=int(where[0])
                if not actual_drawn:
                    for col in range(3):axes[row,col].plot(t,arrays['target'][k,:,col],'k',lw=.9,label='实测')
                    actual_drawn=True
                successful=record.status=='completed' and bool(truth(pd.Series([record.converged])).iloc[0])
                if separate_states and not successful:
                    axes[row,3+methods.index(method)].text(.5,.5,'FAILED',ha='center',fontsize=6,transform=axes[row,3+methods.index(method)].transAxes)
                    continue
                for col in range(3):axes[row,col].plot(t,arrays['prediction'][k,:,col],style,color=color,lw=.9,label=names[methods.index(method)])
                state_ax=axes[row,3+methods.index(method)] if separate_states else axes[row,3]
                if method==methods[1] or both_states:
                    if not np.isfinite(arrays['states'][k]).any():state_ax.text(.5,.5,'未估计',ha='center',transform=state_ax.transAxes)
                    for j,name in enumerate(['f','v','p','q'],start=2):state_ax.plot(t,arrays['states'][k,:,j],lw=.85,ls=style if both_states else '-',color=f'C{j-2}' if both_states else None,label=name if separate_states else name+(' 定τ' if method==methods[0] else ' 学τ') if both_states else name)
            identity_label=f'{identity.subject} / fold {identity.outer} / trial {identity.trial}'
            page_labels.append(identity_label+'；'+'；'.join(text_status))
            axes[row,0].set_ylabel(identity_label,fontsize=7)
            for col,name in enumerate(['EEG','HbO','HbR',('两臂状态' if both_states else names[1]+'状态')] if not separate_states else ['EEG','HbO','HbR']+[name+'状态' for name in names]):
                axes[row,col].set_title(name if row==0 else '',fontsize=8)
                axes[row,col].tick_params(labelsize=6);axes[row,col].axvline(0,color='.75',ls=':',lw=.5)
            for state_axis in axes[row,3:]:state_axis.axhline(1,color='.5',ls=':',lw=.7)
            if row==0:
                axes[row,0].legend(fontsize=6,ncol=1);axes[row,3].legend(fontsize=6,ncol=2 if separate_states else 4)
            for ax in axes[row]:ax.set_xlabel('相对事件 / 秒',fontsize=7)
            index.append(dict(page=page_number,**identity._asdict(),**({m+'_status':records.loc[m,'status'] for m in methods} if separate_states else dict(fixed_status=records.loc[methods[0],'status'],trained_status=records.loc[methods[1],'status']))))
        fig.tight_layout(rect=(0,.01,1,.97));png=folder/f'page_{page_number:02d}.png';fig.savefig(png,dpi=240);plt.close(fig)
        page=pdf.new_page(width=864,height=720)
        heading=f'{context_label+" · " if context_label else ""}全72条曲线核验 · {page_number}/24 · {names[0]} 对 {names[1]}'
        if separate_states:heading=f'{context_label} · 全72条曲线 · {page_number}/24 · 三臂分列状态'
        page.insert_htmlbox(fitz.Rect(24,8,840,36),'<b>'+heading+'</b>',css='*{font-size:11pt}')
        page.insert_image(fitz.Rect(24,38,840,582),filename=str(png))
        caption='<br>'.join(page_labels)+f'<br>黑：实测；灰虚：{names[0]}；蓝：{names[1]}。末列为{names[1]} f/v/p/q，基线1是静息参照，不是正常界。NRMSE顺序EEG/HbO/HbR；失败保留，不以图像选择方法。当前Hb为相对观测坐标。'
        if both_states:caption+=' 状态列同色为同一状态，虚线固定τ、实线训练τ；失败的末有效轨迹不代表有效估计。'
        if separate_states:caption='<br>'.join(page_labels)+'<br>黑：实测；灰虚：A；蓝实：B；橙点：C。仅成功轨迹绘线，失败保留状态标签。后三列分别为各臂f/v/p/q，独立纵轴；基线1不是正常界。NRMSE顺序EEG/HbO/HbR。'
        spare,scale=page.insert_htmlbox(fitz.Rect(24,588,840,710),caption,css='*{font-size:9pt;line-height:1.35}')
        if spare<0 or scale<.95:raise ValueError('Atlas caption failed to fit legibly')
    pdf.save(out/'FULL_CURVE_ATLAS.pdf',garbage=4,deflate=True)
    audit=dict(pages=len(pdf),image_objects=sum(len(p.get_images()) for p in pdf),searchable_text_chars=sum(len(p.get_text()) for p in pdf),bitmap_dpi=240,wps_checked=False,identities=72)
    if audit['pages']!=24 or audit['image_objects']!=24 or audit['searchable_text_chars']<1000:raise ValueError('Atlas verification failed')
    (out/'atlas_validation.json').write_text(json.dumps(audit,indent=2));pd.DataFrame(index).to_csv(out/'atlas_index.csv',index=False);pdf.close()
    return f'\n## 全72条曲线核验图册\n\n[FULL_CURVE_ATLAS.pdf](FULL_CURVE_ATLAS.pdf)按固定被试/折/trial顺序展示全部72条实测身份，每页3条，比较{names[0]}与{names[1]}；包含EEG/HbO/HbR、共享状态及逐条误差/状态。图册保留失败、不按效果挑样；静息参照1不是生理正常界。\n\n'


def flow_hbt_decomposition(run,out,config,data,labels,save,selected):
    """Read retained successful full fits; do not change scores or fit anything."""
    import importlib.util
    import sys
    root=run/'source_snapshot'
    name='report_frozen_hbt_core'
    spec=importlib.util.spec_from_file_location(name,root/'src/inference/t3a_balloon_robust_ssm.py')
    core=importlib.util.module_from_spec(spec);sys.modules[name]=core;spec.loader.exec_module(core)
    fixed={k:v for k,v in config['fixed'].items() if k!='kappa'}
    p=core.BalloonParameters(core.BalloonFixedParameters(**fixed),core.BalloonFreeParameters(kappa=config['fixed']['kappa'],tau=config['reference_tau_s']))
    rest=np.array([0.,0.,1.,1.,1.,1.]);base=core.observation_map(rest,p)[1:].sum()
    shifted=rest.copy();shifted[4]+=1.;loading=float(core.observation_map(shifted,p)[1:].sum()-base)
    shifted=rest.copy();shifted[5]+=.1
    if abs(base)>1e-12 or abs(core.observation_map(shifted,p)[1:].sum()-base)>1e-12:
        raise ValueError('Frozen optical map does not permit HbO+HbR=P-loading*(p-1)')
    spec=importlib.util.spec_from_file_location('report_hbt_operators',root/'src/inference/observation_baselines.py')
    owners=importlib.util.module_from_spec(spec);spec.loader.exec_module(owners)
    operators=owners.native_feature_operators(config['tensor']['steps'])
    operator=operators['fnirs']@operators['native_interpolation']
    rows=[];curves={}
    for identity in data[data['mode'].eq('full')&data.status.eq('completed')&truth(data.converged)].itertuples():
        path=run/'cells'/f'{identity.group}__{identity.method}__full'/'trajectories.npz'
        with np.load(path,allow_pickle=False) as arrays:
            index=np.flatnonzero(arrays['trial_indices']==identity.trial)
            if len(index)!=1:raise ValueError('HbT identity missing')
            k=int(index[0]);row,volume,difference,total=decompose_processed_hbt(arrays['states'][k],arrays['prediction'][k],operator,loading)
        rows.append(dict(kind=identity.kind,group=identity.group,subject=identity.subject,session=identity.session,outer=identity.outer,trial=identity.trial,method=identity.method,**row))
        if any(example.group==identity.group and example.trial==identity.trial for _,example in selected):curves[(identity.group,identity.trial,identity.method)]=(volume,difference,total)
    frame=pd.DataFrame(rows);frame.to_csv(out/'flow_hbt_decomposition.csv',index=False)
    parent_check=run/'parent_initial_difference_decomposition.json'
    parent_verified=0
    if parent_check.exists():
        reference=json.loads(parent_check.read_text())
        if reference['parent_run']!=config['parent_run']:raise ValueError('HbT audit references a different parent')
        previous=pd.DataFrame(reference['rows']);identity=['subject','outer','trial','method']
        inherited=frame[frame.kind.eq('measured')&frame.method.isin(config['methods'][:2])]
        paired=inherited.merge(previous,on=identity,how='outer',validate='one_to_one',indicator=True)
        if not paired['_merge'].eq('both').all():raise ValueError('Parent HbT audit identities differ')
        for field,ref in [('volume_rms_over_total','volume_rms_ratio'),('initial_difference_rms_over_total','initial_difference_rms_ratio'),('signed_cross_over_total_ms','cross_energy_ratio')]:
            if not np.allclose(paired[field],paired[ref],rtol=1e-8,atol=1e-9,equal_nan=True):raise ValueError('Parent HbT decomposition mismatch '+field)
        parent_verified=len(paired)
    display=[]
    for (kind,method),part in frame.groupby(['kind','method'],sort=False):
        expected=int((data.kind.eq(kind)&data.method.eq(method)&data['mode'].eq('full')).sum())
        display.append({'数据':kind,'方法':labels[method],'成功/计划':f'{len(part)}/{expected}',
            'V/总RMS中位':part.volume_rms_over_total.median(),'δ/总RMS中位':part.initial_difference_rms_over_total.median(),
            '有符号交叉项中位':part.signed_cross_over_total_ms.median(),'最大相加误差':part.additivity_max_abs_difference.max()})
    display.sort(key=lambda row: (['synthetic','measured'].index(row['数据']),list(labels.values()).index(row['方法'])))
    summary=pd.DataFrame(display);summary.to_csv(out/'flow_hbt_decomposition_summary.csv',index=False)
    displayed=summary.copy();displayed['最大相加误差']=displayed['最大相加误差'].map(lambda value:f'{value:.2e}')
    body=('\n## Posthoc HbT组成诊断\n\n'
        '此处是事后机制诊断，不是新增主评分或独立解释比例。当前p=v+δ，实际冻结光学映射下HbT=HbO+HbR；'
        f'从冻结core逐项核验的p系数为{loading:g}。采用与full拟合相同的L=原生Hb处理算子×模型时钟插值，'
        'V=L[P0(v−1)]、D=L[P0(p−v)]严格相加为已保存预测HbO+HbR（实际增益/坐标系数已纳入p系数）。'
        'D是初态差值衰减项：幅度来自p0−v0，衰减速率f_out/(τv)仍随驱动下的v变化；不是独立于driver的私有源，只有近静息时才近似时间常数τ的指数。不能仅凭范数高就认定非生理。'
        '每trial核验RMS(total)^2=RMS(V)^2+RMS(D)^2+2mean(V·D)。'
        '两项RMS/total与有符号交叉项2mean(V·D)/mean(total²)可能大于1或为负，反映抵消，不能称独立解释百分比；total RMS不超过1e−12时比值记为未定义。'
        '表内分别取逐trial比值中位，不能再把这些中位数代回同一trial恒等式；完整逐条RMS和交叉项见CSV。'
        '仅成功full进入该诊断，失败不以末有效迭代补入；原分母、主评分和预声明方法不变。\n\n'+table(displayed)+'\n\n')
    if parent_verified:
        body+=f'父两基线{parent_verified}条成功full身份与独立[父初态差值分解审计](../parent_initial_difference_decomposition.json)逐项核对通过；继承结果不是新增复制。\n\n'
    for rank,example in selected:
        fig,axes=plt.subplots(2,len(config['methods'])//2,figsize=(13,7),sharex=True)
        for ax,method in zip(axes.flat,config['methods']):
            curve=curves.get((example.group,example.trial,method))
            ax.set_title(labels[method],fontsize=10)
            if curve is None:
                ax.text(.5,.5,'未成功：不作该分解\n固定身份不替换',ha='center',va='center',transform=ax.transAxes)
                continue
            volume,difference,total=curve;t=np.arange(len(total))*config['tensor']['dt_s']-(5 if example.kind=='measured' else 0)
            ax.plot(t,total,'k',label='预测HbT',lw=1.3);ax.plot(t,volume,label='血容量V',lw=1);ax.plot(t,difference,label='初态差异D',lw=1)
            ax.axhline(0,color='.7',ls=':',lw=.6);ax.legend(fontsize=7);ax.set_ylabel('处理后HbT')
        for ax in axes[-1]:ax.set_xlabel('相对事件 / 秒' if example.kind=='measured' else '秒')
        fig.suptitle(f'Posthoc {rank}: {example.group} / trial {int(example.trial)}',fontsize=11)
        save(fig,'flow_hbt_'+rank,'父基线固定身份的各臂预测HbT分解；V+D严格等于总预测，成分不独立，可互相抵消。失败不换样；这不是新主评分或生理正常性判据。')
    return body


def fixed_roi_flow_details(run, out, config, data, states, labels, save, selected):
    """Conditional paired scores and flow diagnostics, without selecting an arm."""
    identities=['kind','group','subject','session','trial','mode']
    transitions=[]; score_rows=[]; identity_rows=[]
    comparisons=[(m,'fixed_roi_fixed_tau') for m in config['methods'][2:]]
    comparisons += [(m,'fixed_roi_trained_tau') for m in config['methods'][3:]]
    comparisons += [('trained_tau_logflow1','fixed_tau_logflow1')]
    initial=config['schema']=='shared_driver_fixed_roi_initial_v1'
    if initial:
        m=config['methods'];comparisons=[(m[2],m[0]),(m[3],m[1]),(m[1],m[0]),(m[3],m[2])]
    for method,baseline in comparisons:
        a=data[data.method.eq(method)];b=data[data.method.eq(baseline)]
        joined=a.merge(b,on=identities,suffixes=('','_base'),how='outer',validate='one_to_one',indicator=True)
        if not joined['_merge'].eq('both').all():raise ValueError('Flow comparison changed identity set')
        ok=joined.status.eq('completed')&truth(joined.converged)
        baseok=joined.status_base.eq('completed')&truth(joined.converged_base)
        joined['transition']=np.select([ok&baseok,ok&~baseok,~ok&baseok],['both_success','new_success','lost_success'],default='both_failed')
        for (kind,mode),part in joined.groupby(['kind','mode'],sort=False):
            row=dict(kind=kind,method=method,baseline=baseline,mode=mode,expected=len(part))
            row.update({key:int(part.transition.eq(key).sum()) for key in ['both_success','new_success','lost_success','both_failed']})
            transitions.append(row)
            for subset in ['both_success','new_success']:
                good=part[part.transition.eq(subset)]
                for component in COMPONENTS:
                    if mode=='center_EEG' and component!='EEG':continue
                    if mode=='center_fNIRS' and component=='EEG':continue
                    field='nmse_'+component
                    if field not in good:continue
                    result=dict(kind=kind,method=method,baseline=baseline,mode=mode,subset=subset,n=len(good),expected=len(part),component=component)
                    def average(values):
                        return values.groupby([good.subject,good.session]).mean().groupby(level=0).mean().mean()
                    result['candidate_nrmse']=np.sqrt(average(good[field]))
                    result['baseline_nrmse']=np.sqrt(average(good[field+'_base'])) if subset=='both_success' else np.nan
                    result['delta_nmse']=average(good[field]-good[field+'_base']) if subset=='both_success' else np.nan
                    score_rows.append(result)
        identity_rows.append(joined[identities+['transition']].assign(method=method,baseline=baseline))
    transition=pd.DataFrame(transitions);scores=pd.DataFrame(score_rows)
    transition.to_csv(out/'flow_paired_transitions.csv',index=False)
    scores.to_csv(out/'flow_paired_scores.csv',index=False)
    pd.concat(identity_rows,ignore_index=True).to_csv(out/'flow_paired_identities.csv',index=False)
    body='\n## Log-flow工程软约束与配对结果\n\n六臂及权重预声明，不按验证成绩选择w。log-flow软约束只是以f=1为中心的工程锚，log SD=ln2不是正常范围、人口分布或校准生理先验。不同w的总目标不可横向排名；多起点仅在同组同臂内选择。继承父运行的两臂及原prepared观测不是新独立复制。\n\n'
    equation_path=run/'model_equation_audit.json'
    if 'measured' in data.kind.unique() and equation_path.exists():
        audit=json.loads(equation_path.read_text())
        if audit['execution_effect']['running_flow_model_changed'] or audit['execution_effect']['frozen_snapshot_changed']:
            raise ValueError('Equation audit says the running/frozen model was changed')
        body+=('### 方程来源与初态自由度\n\n'
            '本轮运行的是合同定义的inlet-balance变体：τ·dp/dt=f−f_out·p/v，'
            '与Tak等2015文中式(3)的τ·dp/dt=(f−f_out)·p/v不同，不能在独立p0/v0初态下视作相同方程。'
            '当前δ=p−v满足dδ/dt=−f_out·δ/(τv)，近静息时衰减尺度约为τ；'
            '因此自由p0/v0允许未知窗前历史，同时提供与τ耦合的衰减Hb成分。原文方程则保持p/v恒定；两者在p0=v0时均保持p=v。'
            '该自由度尚无本实验独立浓度/幅度标定支持，也不能仅凭数学自由度断言生理不可能。'
            '来源审计没有修改本轮方程或冻结代码；当前各臂比较仍使用完全相同的原变体。'
            '详见[model_equation_audit.json](../model_equation_audit.json)及'
            f"[Tak et al. 2015]({audit['primary_source']['url']})。\n\n")
        initial=[]
        for method in config['methods']:
            all_rows=data[data.kind.eq('measured')&data.method.eq(method)&data['mode'].eq('full')]
            good=states[states.kind.eq('measured')&states.method.eq(method)&states['mode'].eq('full')&states.converged&states.finite_states]
            ratio=good.p_initial/good.v_initial;delta=good.p_initial-good.v_initial
            initial.append({'方法':labels[method],'成功/计划':f'{len(good)}/{len(all_rows)}',
                'p0/v0最小':ratio.min(),'p0/v0中位':ratio.median(),'p0/v0最大':ratio.max(),'最大绝对初态差':delta.abs().max()})
        initial=pd.DataFrame(initial);initial.to_csv(out/'flow_initial_pv_audit.csv',index=False)
        body+='实测full成功子集的初态审计；失败完整分母仍保留，上述比值不是独立测得浓度：\n\n'+table(initial)+'\n\n'
    null_path=run/'stationary_baseline_null_check.json'
    if null_path.exists():
        null=json.loads(null_path.read_text())
        if null.get('status')!='passed':raise ValueError('Stationary baseline null check did not pass')
        records=null['records'];maximum=max(r['full_processed_max_abs'] for r in records)
        drift=max(r['rollout_state_max_abs_drift'] for r in records)
        body+=('### 恒定水平负对照\n\n'
            f"独立确定性算子检查（非重新拟合）覆盖f={null['grid']['flows']}、τ={null['grid']['taus_s']}："
            '令外加常数r=γ(c−1)/β、s=0、v=p=c^α、q=vE(c)/E0，'
            f'末轨迹漂移最大{drift:.3g}，完整处理后的观测绝对最大{maximum:.4g}。'
            '不同的数学域有效恒定绝对水平被baseline处理消去，不能由零处理观测证明f≈1。'
            '原初态惩罚已提供绝对水平锚，flow正则再增加工程锚；相同数据残差不表示正则总目标相等。'
            '这是恒定轨迹族的负对照，不是任意动态轨迹的全局等价变换，也不意味着实测τ毫无信息。'
            '证据：[stationary_baseline_null_check.json](../stationary_baseline_null_check.json)。\n\n')
    counts=[]
    for kind,part in data.groupby('kind',sort=False):
        for inherited,arm_names in [(True,config['methods'][:2]),(False,config['methods'][2:])]:
            rows=part[part.method.isin(arm_names)]
            if initial and kind=='synthetic':
                if inherited:continue
                rows=part
            counts.append({'数据':kind,'来源':'父继承' if inherited else '本轮新拟合',
                'cell数':len(rows[['group','method','mode']].drop_duplicates()),'记录trial数':len(rows),
                '成功trial数':int((rows.status.eq('completed')&truth(rows.converged)).sum())})
    body+=table(pd.DataFrame(counts))+'\n\n'
    for (kind,mode),part in transition.groupby(['kind','mode'],sort=False):
        body+=f'### {kind} / {mode} 身份转移\n\n'+table(pd.DataFrame([{'候选':labels[r.method],'基线':labels[r.baseline],'共同成功':r.both_success,'新成功':r.new_success,'失去成功':r.lost_success,'双方失败':r.both_failed,'计划':r.expected} for r in part.itertuples()]))+'\n\n'
    body+='同w=1固定τ对训练τ的误差只在共同成功身份比较；新成功样本另列，不能用不同成功子集的总体均值声称预测改善。\n\n'
    part=scores[scores.baseline.eq('fixed_tau_logflow1')]
    if initial:
        m=config['methods'];part=scores[(scores.method.eq(m[2])&scores.baseline.eq(m[0]))|(scores.method.eq(m[3])&scores.baseline.eq(m[1]))]
    for (kind,mode,method,baseline),group in part.groupby(['kind','mode','method','baseline'],sort=False):
        body+=f'{kind} / {mode} / {labels[method]} vs {labels[baseline]}：\n\n'+table(group[['subset','component','n','expected','candidate_nrmse','baseline_nrmse','delta_nmse']].rename(columns={'subset':'子集','component':'分量','n':'成功数','expected':'计划','candidate_nrmse':'候选','baseline_nrmse':'基线','delta_nmse':'ΔNMSE'}))+'\n\n'
    recovery=[];extremes=[]
    for (kind,method,mode),part in states.groupby(['kind','method','mode'],sort=False):
        good=part[part.converged&part.finite_states]
        row=dict(kind=kind,method=method,mode=mode,completed=len(good),expected=int(((data.kind==kind)&(data.method==method)&(data['mode']==mode)).sum()))
        for name in ['flow_below_half_s','flow_above_two_s','flow_below_01_s','driver_mean','driver_ac_rms']:
            row[name+'_mean']=pd.to_numeric(good.get(name,pd.Series(dtype=float))).mean()
        extremes.append(row)
        if kind=='synthetic':
            rec=dict(method=method,mode=mode,completed=len(good),expected=row['expected'])
            for name in ['driver_dc_error','driver_ac_rmse','initial_rmse','truth_rmse_f']:
                values=pd.to_numeric(good.get(name,pd.Series(dtype=float)))
                rec[name]=np.sqrt(np.mean(values**2)) if len(values) else np.nan
            rec['driver_dc_signed_mean']=pd.to_numeric(good.get('driver_dc_error',pd.Series(dtype=float))).mean()
            rec['flow_signed_mean']=pd.to_numeric(good.get('flow_mean_error',pd.Series(dtype=float))).mean()
            recovery.append(rec)
    pd.DataFrame(extremes).to_csv(out/'flow_duration_summary.csv',index=False)
    pd.DataFrame(recovery).to_csv(out/'flow_synthetic_recovery.csv',index=False)
    candidates=pd.read_csv(out/'tau_training_candidates.csv')
    body+='## 训练求值、域拒绝与正则代价\n\n总目标、数据SSE和flow代价分开保留于tau_training_candidates.csv；不同w不作总目标胜负比较。域拒绝次数、终点梯度和停止原因反映优化路径，不能把预算失败统一称为域不合理。\n\n'
    diagnostics=[]
    for (kind,method),part in candidates.groupby(['kind','method'],sort=False):
        good=part.status.eq('completed')&truth(part.converged)
        gradient=pd.to_numeric(part.gradient,errors='coerce')
        diagnostics.append({'数据':kind,'方法':labels[method],'成功/起点':f'{good.sum()}/{len(part)}',
            '域拒绝':pd.to_numeric(part.domain_rejections,errors='coerce').sum(),
            '梯度中位':gradient.median(),'梯度最大':gradient.max(),
            'flow代价中位':pd.to_numeric(part.flow_prior_cost,errors='coerce').median()})
    body+=table(pd.DataFrame(diagnostics))+'\n\n'
    body+='## Driver DC/AC、初态和极端持续时间\n\nflow低于0.5、高于2和低于0.1的时间仅作预声明描述阈值（采样数×dt），不是正常范围或连续域边界。失败末有效状态仅留在逐条state_ranges，不混入成功条件均值。DC误差、去各自均值后的AC误差分列；后者不是校正预测或主评分。\n\n'
    for mode in data['mode'].unique():
        part=pd.DataFrame(recovery);part=part[part['mode'].eq(mode)] if len(part) else part
        if len(part):body+=f'合成 / {mode}：\n\n'+table(part.drop(columns='mode').replace({'method':labels}).rename(columns={'method':'方法','completed':'成功','expected':'计划','driver_dc_error':'DC RMSE','driver_ac_rmse':'AC RMSE','initial_rmse':'初态RMSE','truth_rmse_f':'f RMSE','driver_dc_signed_mean':'DC均偏差','flow_signed_mean':'f均偏差'}))+'\n\n'
    fig,axes=plt.subplots(1,3,figsize=(13,4),squeeze=False)
    for ax,field,title in zip(axes[0],['flow_below_half_s','flow_above_two_s','flow_below_01_s'],['f<0.5','f>2','f<0.1']):
        for x,method in enumerate(config['methods']):
            p=states[states.method.eq(method)&states['mode'].eq('full')]
            for success,marker in [(True,'o'),(False,'x')]:
                q=p[p.converged.eq(success)&p.finite_states]
                ax.scatter(np.full(len(q),x),q[field],marker=marker,s=12,alpha=.45)
        ax.set_xticks(range(len(labels)),list(labels.values()),rotation=35,ha='right',fontsize=7)
        ax.set_title(title);ax.set_ylabel('路径采样持续时间 / 秒')
    save(fig,'flow_extreme_durations','全观测路径极端持续时间；圆为成功、叉为失败末有效状态。固定阈值只作工程诊断，不是正常生理范围。')
    synthetic=states[states.kind.eq('synthetic')&states['mode'].eq('full')]
    if len(synthetic):
        fig,axes=plt.subplots(1,3,figsize=(13,4))
        for ax,field,title in zip(axes,['driver_dc_error','driver_ac_rmse','flow_mean_error'],['driver DC有符号误差','driver AC RMSE','flow有符号均误差']):
            for mi,method in enumerate(config['methods']):
                part=synthetic[synthetic.method.eq(method)]
                for ok,marker in [(True,'o'),(False,'x')]:
                    good=part[part.converged.eq(ok)&part.finite_states]
                    ax.scatter(mi+np.linspace(-.15,.15,len(good)),good[field],marker=marker,s=13,alpha=.55)
            ax.axhline(0,color='.5',ls=':',lw=.7);ax.set_title(title)
            ax.set_xticks(range(len(labels)),list(labels.values()),rotation=35,ha='right',fontsize=7)
        save(fig,'flow_synthetic_bias','合成全观测driver DC、AC与flow均值偏差；无幅度/时间对齐，所有固定身份保留。圆为成功、叉为失败末有效状态。DC与flow保留有符号误差，AC单列不能替代原坐标评分。')
    if selected:
        fig,axes=plt.subplots(2,len(selected),figsize=(12,7),squeeze=False,sharex='col')
        for col,(rank,example) in enumerate(selected):
            for mi,method in enumerate(config['methods']):
                path=run/'cells'/f'{example.group}__{method}__full'/'trajectories.npz'
                record=data[(data.group==example.group)&(data.trial==example.trial)&(data.method==method)&data['mode'].eq('full')]
                success=len(record) and record.status.iloc[0]=='completed' and truth(record.converged).iloc[0]
                if not path.exists():
                    for ax in axes[:,col]:ax.plot([],[],':',label=labels[method]+' 未估计')
                    continue
                with np.load(path,allow_pickle=False) as arrays:
                    index=np.flatnonzero(arrays['trial_indices']==int(example.trial))
                    if len(index)!=1:raise ValueError('Fixed parent example missing')
                    state=arrays['states'][index[0]];t=np.arange(len(state))*config['tensor']['dt_s']-(5 if example.kind=='measured' else 0)
                    suffix='' if success else (' 失败迭代' if np.isfinite(state).any() else ' 未估计')
                    for ax,k in zip(axes[:,col],[0,2]):ax.plot(t,state[:,k],color=plt.get_cmap('tab10')(mi),ls='-' if success else ':',lw=.9,label=labels[method]+suffix)
            axes[0,col].set_title(f'{rank}: {example.group} / trial {int(example.trial)}',fontsize=9)
            axes[0,col].set_ylabel('driver r');axes[1,col].set_ylabel('flow f');axes[1,col].set_xlabel('相对事件秒' if example.kind=='measured' else '秒')
            axes[0,col].legend(fontsize=7,ncol=2)
        save(fig,'flow_fixed_parent_examples','固定基线median/worst身份的各臂driver/flow横比，无对齐。失败末有效迭代虚线，缺失不换样。')
    return body


def fixed_roi_initial_details(run,out,config,data,states,labels,save):
    """Prespecified truth strata expose tied-model misspecification, not arm selection."""
    truth_by_group={}
    for group in data.group.unique():
        prep=json.loads((run/'prepared'/f'{group}.json').read_text())
        truth_by_group[group]=prep['spec']
    synthetic=data[data.kind.eq('synthetic')].copy()
    synthetic['true_ratio']=synthetic.group.map(lambda g:truth_by_group[g]['true_initial_log_p_over_v'])
    synthetic['true_tau']=synthetic.group.map(lambda g:truth_by_group[g]['true_tau'])
    parameters=pd.read_csv(out/'parameters.csv')
    parameter_map={(r.group,r.method):r.tau for r in parameters.itertuples()}
    synthetic['tau_estimate']=[parameter_map.get((r.group,r.method),np.nan) for r in synthetic.itertuples()]
    def summarize(part, ratio, method, mode, tau=None):
        good=part[part.status.eq('completed')&truth(part.converged)]
        state=states[states.kind.eq('synthetic')&states.method.eq(method)&states['mode'].eq(mode)&states.group.isin(part.group.unique())&states.converged&states.finite_states]
        row=dict(true_initial_log_p_over_v=ratio,true_tau=tau,method=method,mode=mode,expected=len(part),completed=len(good))
        for component in COMPONENTS:
            field='nmse_'+component
            row['nrmse_'+component]=np.sqrt(good.groupby(['subject','session'])[field].mean().groupby('subject').mean().mean()) if field in good else np.nan
        row['tau_rmse']=np.sqrt(np.mean((good.tau_estimate-good.true_tau)**2)) if len(good) else np.nan
        for field in ['truth_rmse_r','truth_rmse_f','truth_rmse_p','truth_rmse_v','driver_dc_error','driver_ac_rmse']:
            values=pd.to_numeric(state.get(field,pd.Series(dtype=float)),errors='coerce')
            row[field]=np.sqrt(np.mean(values**2)) if len(values) else np.nan
        row['initial_logratio_rmse']=np.sqrt(np.mean((state.initial_log_p_over_v-ratio)**2)) if len(state) else np.nan
        row['path_max_abs_p_minus_v']=state.path_max_abs_p_minus_v.max() if len(state) else np.nan
        return row
    strata=[];detailed=[]
    for (ratio,method,mode),part in synthetic.groupby(['true_ratio','method','mode'],sort=True):
        strata.append(summarize(part,ratio,method,mode))
        for tau,cell in part.groupby('true_tau',sort=True):detailed.append(summarize(cell,ratio,method,mode,tau))
    strata=pd.DataFrame(strata);detailed=pd.DataFrame(detailed)
    strata.to_csv(out/'initial_truth_strata.csv',index=False);detailed.to_csv(out/'initial_truth_tau_strata.csv',index=False)
    invariance=states[states.finite_states].copy();invariance['tied']=invariance.method.isin(config['initial_constraint']['tied_methods'])
    invariance[['kind','group','method','mode','trial','status','converged','tied','initial_log_p_over_v','path_max_abs_p_minus_v']].to_csv(out/'initial_path_invariance.csv',index=False)
    if invariance.loc[invariance.tied,'path_max_abs_p_minus_v'].gt(1e-8).any():raise ValueError('Tied p=v invariant violated by saved trajectory')
    body=('\n## 初态约束：预声明真值分层\n\n'
        '四臂均固定flow权重1；比较固定τ/共享τ各自的free与tied初态。tied在降维优化坐标中令p0=v0，'
        '每trial自由度125→124，仍保留原五个物理初态惩罚项（v与p两项都计入）；不是投影旧解，也不改变动力学。'
        '合成36组为本轮全新生成/拟合；仅实测free两臂继承，不构成独立复制。'
        'true log(p0/v0)=0是约束匹配层，±0.05是预声明错设压力层，不是人口正常范围。'
        '必须分别看这三层及τ=1/2/4，不能用合并合成均值掩盖错设偏差。'
        '以下表按初态层分开；initial_truth_tau_strata.csv进一步保留每个τ、模式、方法的误差与完整失败分母。'
        '局部信息来自各自有效降维参数空间，不是置信区间，也不能将减少自由度后的信息变化当作额外数据。\n\n')
    for ratio in config['synthetic']['initial_log_p_over_v']:
        body+=f'### 真log(p0/v0)={ratio:+.2f}：'+('匹配' if ratio==0 else '错设压力')+'\n\n'
        for mode in config['modes']:
            part=strata[strata.true_initial_log_p_over_v.eq(ratio)&strata['mode'].eq(mode)].set_index('method').reindex(config['methods'])
            body+=f'{mode}观测重建（各τ等身份合并，真τ细分见CSV/图）：\n\n'+table(pd.DataFrame([{'方法':labels[m],'成功/计划':f'{int(r.completed)}/{int(r.expected)}','EEG':r.nrmse_EEG,'HbO':r.nrmse_HbO,'HbR':r.nrmse_HbR} for m,r in part.iterrows()]))+'\n\n'
            body+='原坐标状态/参数恢复（成功条件；DC/AC不可替代原driver误差）：\n\n'+table(pd.DataFrame([{'方法':labels[m],'τ RMSE':r.tau_rmse,'r RMSE':r.truth_rmse_r,'f RMSE':r.truth_rmse_f,'p RMSE':r.truth_rmse_p,'v RMSE':r.truth_rmse_v,'初态log比RMSE':r.initial_logratio_rmse} for m,r in part.iterrows()]))+'\n\n'
    body+='各初态层的driver DC误差RMS与AC RMSE另列于initial_truth_strata.csv及initial_truth_tau_strata.csv，保留模式与真τ，不将DC偏差用去均值消除。\n\n'
    invariant_summary=[]
    for (kind,method),part in invariance.groupby(['kind','method'],sort=False):
        invariant_summary.append({'数据':kind,'方法':labels[method],'有限轨迹数':len(part),'成功数':int(part.converged.sum()),'全路径最大绝对p−v':f'{part.path_max_abs_p_minus_v.max():.3e}'})
    body+='### 全路径约束核验\n\n当前方程下p0=v0应保持p=v；下表包含所有有限末有效轨迹，失败仍不进入主评分。tied轨迹最大绝对差超过1e−8会拒绝生成报告，不将偏离静默当成约束已执行。\n\n'+table(pd.DataFrame(invariant_summary))+'\n\n'
    full=detailed[detailed['mode'].eq('full')]
    for fields,name in [(['tau_rmse','truth_rmse_r','truth_rmse_f'],'initial_tau_driver_flow_recovery'),(['truth_rmse_p','truth_rmse_v'],'initial_p_v_recovery')]:
        fig,axes=plt.subplots(3,len(fields),figsize=(12,10),squeeze=False)
        for row,ratio in enumerate(config['synthetic']['initial_log_p_over_v']):
            for col,field in enumerate(fields):
                ax=axes[row,col]
                for method in config['methods']:
                    part=full[full.true_initial_log_p_over_v.eq(ratio)&full.method.eq(method)].sort_values('true_tau')
                    ax.plot(part.true_tau,part[field],'o-',label=labels[method],lw=1,ms=4)
                ax.set_title(f'真log(p0/v0)={ratio:+.2f} / {field}',fontsize=9);ax.set_xlabel('真τ / 秒');ax.set_ylabel('成功条件RMSE');ax.legend(fontsize=6)
        save(fig,name,'合成full按初态真值与τ同时分层，展示匹配层与错设层；失败完整分母见表/CSV。不跨层挑方法，不以总体平均掩盖偏差。')
    for tau in config['synthetic']['taus_s']:
        fig,axes=plt.subplots(4,3,figsize=(13,11),sharex='col')
        for col,ratio in enumerate(config['synthetic']['initial_log_p_over_v']):
            group=f'synthetic_tau{tau:g}_slow_r0_logpv{ratio:+.2f}';trial=0;path=run/'prepared'/f'{group}.npz'
            axes[0,col].set_title(f'τ={tau:g} / log(p0/v0)={ratio:+.2f}',fontsize=9)
            if not path.exists():
                for ax in axes[:,col]:ax.text(.5,.5,'固定身份缺失，不换样',ha='center',transform=ax.transAxes)
                continue
            with np.load(path,allow_pickle=False) as a:
                t=np.arange(a['target'].shape[1])*config['tensor']['dt_s'];true_state=a['truth_states'][trial]
                axes[0,col].plot(t,a['target'][trial,:,1]+a['target'][trial,:,2],color='.65',lw=.7,label='含噪HbT目标')
                values=[a['clean'][trial,:,1]+a['clean'][trial,:,2],true_state[:,0],true_state[:,2],true_state[:,4]-true_state[:,3]]
                for ax,value in zip(axes[:,col],values):ax.plot(t,value,'k',lw=1.4,label='真值')
            for method in config['methods']:
                path=run/'cells'/f'{group}__{method}__full'/'trajectories.npz'
                record=data[data.group.eq(group)&data.method.eq(method)&data['mode'].eq('full')&data.trial.eq(0)]
                status=record.status.iloc[0] if len(record) else 'missing'
                if not path.exists():
                    for ax in axes[:,col]:ax.plot([],[],':',label=labels[method]+' 未估计')
                    continue
                with np.load(path,allow_pickle=False) as a:
                    index=np.flatnonzero(a['trial_indices']==0)
                    if len(index)!=1:raise ValueError('Initial truth example trial0 missing')
                    k=index[0];state=a['states'][k];values=[a['prediction'][k,:,1]+a['prediction'][k,:,2],state[:,0],state[:,2],state[:,4]-state[:,3]]
                    for ax,value in zip(axes[:,col],values):ax.plot(t,value,'-' if status=='completed' else ':',lw=.85,label=labels[method]+('' if status=='completed' else ' 失败'))
            for ax,label in zip(axes[:,col],['处理HbT','driver r','flow f','p−v']):ax.set_ylabel(label);ax.legend(fontsize=5)
            axes[-1,col].set_xlabel('秒')
        save(fig,f'initial_truth_tau{tau:g}','预声明slow/rep0/trial0，三初态真值层分别展示，四方法与真实driver/flow/p−v比较；无幅度、符号、时间或基线对齐。失败保留，不按效果选样。')
    return body


def render_nonlinear_fit(run, out, *, gain_prior=False, fixed_roi_tau=False, fixed_roi_flow=False, fixed_roi_initial=False):
    """Read only this run's declared cells; keep failed and missing denominators."""
    if out.exists() and any(out.iterdir()):
        raise ValueError("Export directory is not empty; choose a new versioned --output")
    summary = json.loads((run / "summary.json").read_text())
    data = pd.read_csv(run / "metrics.csv")
    required = {"kind", "group", "subject", "session", "method", "mode", "trial", "status", "converged"}
    if data.empty or not required.issubset(data):
        raise ValueError("Missing nonlinear fit row identities/status")
    keys = ["kind", "group", "method", "mode", "trial"]
    if data.duplicated(keys).any() or len(data) != summary["observed_rows"]:
        raise ValueError("Duplicate identities or stale nonlinear summary")
    methods = ["fixed_tau_100", "fixed_tau_0p01", "linear_trained_tau_0p01", "linear_prior_tau_0p01"]
    labels = dict(zip(methods, ["固定τ/λ100", "固定τ/λ0.01", "训练τ/λ0.01", "先验τ/λ0.01"]))
    config = {}
    fixed_roi_flow = fixed_roi_flow or fixed_roi_initial
    fixed_roi_tau = fixed_roi_tau or fixed_roi_flow
    extended = gain_prior or fixed_roi_tau
    if extended:
        import yaml
        config = yaml.safe_load((run / "resolved_config.yaml").read_text())
        expected_schema = "shared_driver_fixed_roi_initial_v1" if fixed_roi_initial else "shared_driver_fixed_roi_flow_v1" if fixed_roi_flow else ("shared_driver_fixed_roi_tau_v1" if fixed_roi_tau else "shared_driver_gain_prior_v1")
        if config.get("schema") != expected_schema:
            raise ValueError("Report mode differs from resolved configuration schema")
        methods = list(config["methods"])
        labels = ({"fixed_roi_fixed_tau":"固定ROI/τ2", "fixed_roi_trained_tau":"固定ROI/训练τ", "fixed_roi_prior_tau":"固定ROI/先验τ"}
            if fixed_roi_tau else {"fixed_gain_no_amplitude":"A 固定β", "fixed_gain_amplitude":"B 固定β+幅度",
                                  "trained_gain_no_amplitude":"C 训练β", "trained_gain_amplitude":"D 训练β+幅度"})
        if fixed_roi_flow:
            labels = {"fixed_roi_fixed_tau":"A τ2/w0", "fixed_roi_trained_tau":"B 学τ/w0",
                      "fixed_tau_logflow1":"C τ2/w1", "trained_tau_logflow01":"D 学τ/w0.1",
                      "trained_tau_logflow1":"E 学τ/w1", "trained_tau_logflow4":"F 学τ/w4"}
        if fixed_roi_initial:
            labels={"fixed_tau_logflow1":"A 固定τ/free", "trained_tau_logflow1":"B 共享τ/free",
                    "fixed_tau_logflow1_tied":"C 固定τ/tied", "trained_tau_logflow1_tied":"D 共享τ/tied"}
        if methods != list(labels) or set(data.method) != set(methods):
            raise ValueError("Report methods differ from resolved configuration")
    parent_data = None
    if fixed_roi_flow:
        parent = Path(config['parent_run'])
        if not parent.is_absolute(): parent = Path(__file__).resolve().parents[2] / parent
        parent_data = pd.read_csv(parent/'metrics.csv')
        identity=['kind','group','subject','session','trial','method','mode']
        inherited=data[data.method.isin(methods[:2]) & (data.kind.eq('measured') if fixed_roi_initial else True)]
        previous=parent_data[parent_data.method.isin(methods[:2])&parent_data.kind.isin(data.kind.unique()) & (parent_data.kind.eq('measured') if fixed_roi_initial else True)]
        check=inherited.merge(previous,on=identity,how='outer',suffixes=('','_parent'),indicator=True,validate='one_to_one')
        if not check['_merge'].eq('both').all():raise ValueError('Inherited baseline identities differ from exact parent')
        for field in ['status','converged']+[c for c in data if c.startswith(('nmse_','nrmse_','last_valid_nmse_'))]:
            if not ((check[field]==check[field+'_parent'])|(check[field].isna()&check[field+'_parent'].isna())).all():
                raise ValueError('Inherited baseline changed '+field)
    out.mkdir(parents=True, exist_ok=True)
    (out / "figures").mkdir()
    font = Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc")
    if font.exists():
        font_manager.fontManager.addfont(str(font))
        plt.rcParams["font.family"] = font_manager.FontProperties(fname=str(font)).get_name()
    plt.rcParams.update({"font.size": 9, "axes.unicode_minus": False})
    figures = []

    def save(fig, name, caption):
        fig.tight_layout()
        fig.savefig(out / "figures" / f"{name}.png", dpi=240, bbox_inches="tight")
        plt.close(fig)
        figures.append((name, caption))

    aggregates = []
    for item in summary["aggregates"]:
        row = {k: item[k] for k in ("kind", "method", "mode", "expected", "observed_rows", "completed", "missing")}
        row["failed"] = item["observed_rows"] - item["completed"]
        row["complete_panel"] = item["completed"] == item["expected"]
        part = data[(data.kind == item["kind"]) & (data.method == item["method"]) & (data["mode"] == item["mode"])]
        success = part.status.eq("completed") & truth(part.converged)
        if len(part) != item["observed_rows"] or int(success.sum()) != item["completed"]:
            raise ValueError("Stale nonlinear aggregate counts")
        if item["mode"] == "full":
            errors = part.reindex(columns=["nrmse_"+c for c in COMPONENTS])
            row["joint_below_05"] = int((success & errors.lt(.5).all(axis=1)).sum())
        else:
            row["joint_below_05"] = "N/A"
        for c in COMPONENTS:
            row["NRMSE_" + c] = item.get("nrmse_" + c, np.nan)
        aggregates.append(row)
    agg = pd.DataFrame(aggregates)
    agg.to_csv(out / "nonlinear_summary.csv", index=False)
    failures = data.groupby(["kind", "method", "mode", "status"], dropna=False).size().rename("rows").reset_index()
    failures.to_csv(out / "failure_counts.csv", index=False)
    successful = data.status.eq("completed") & truth(data.converged)
    convergence_columns = [c for c in ["kind", "group", "method", "mode", "trial", "status", "converged", "evaluations", "selected_start", "gradient_inf_norm", "scaled_gradient_inf_norm", "rejected_steps", "domain_rejections", "derivative_rejections", "flow_prior_cost", "flow_prior_weight", "flow_prior_log_sd", "integration_max_difference_training_sd"] if c in data]
    data[convergence_columns].to_csv(out / "convergence.csv", index=False)

    kinds = list(agg.kind.unique())
    fig, axes = plt.subplots(len(kinds), 3, figsize=(13, 4 * len(kinds)), squeeze=False)
    for i, kind in enumerate(kinds):
        for j, mode in enumerate(["full", "center_EEG", "center_fNIRS"]):
            ax = axes[i, j]
            part = agg[(agg.kind == kind) & (agg["mode"] == mode)].set_index("method").reindex(methods)
            columns = COMPONENTS if mode == "full" else (["EEG"] if mode == "center_EEG" else ["HbO", "HbR"])
            x = np.arange(len(methods))
            for k, c in enumerate(columns):
                ax.bar(x + (k - (len(columns)-1)/2)*.22, part["NRMSE_"+c], .22, label=c)
            ax.set_xticks(x, [labels[m] + f"\n{int(part.loc[m, 'completed'])}/{int(part.loc[m, 'expected'])}" for m in methods], rotation=22, ha="right", fontsize=7)
            ax.axhline(.5, color="red", ls="--", lw=1)
            ax.set_title(f"{kind} / {mode}")
            ax.set_ylabel("成功子集 NRMSE"); ax.legend()
    save(fig, "nonlinear_errors", ("预声明四方法" if fixed_roi_initial else "预声明六方法" if fixed_roi_flow else ("预声明三方法" if fixed_roi_tau else "预声明四方法")) + "与相同模式比较；标注成功/计划分母。柱高仅成功子集，缺失或失败不能视作零误差。")

    # Baseline comparisons use the same successful identities, not different subsets.
    paired = []
    baseline = data[data.method == methods[0]]
    join = ["kind", "group", "mode", "trial"]
    for method in methods[1:]:
        merged = data[data.method == method].merge(baseline, on=join, suffixes=("", "_baseline"), validate="one_to_one")
        mask = merged.status.eq("completed") & truth(merged.converged) & merged.status_baseline.eq("completed") & truth(merged.converged_baseline)
        for (kind, mode), all_rows in merged.groupby(["kind", "mode"]):
            part = all_rows[mask.loc[all_rows.index]]
            expected = int(agg[(agg.kind == kind) & (agg.method == method) & (agg["mode"] == mode)].expected.iloc[0])
            row = dict(kind=kind, method=method, mode=mode, common_success=len(part), expected=expected)
            for c in COMPONENTS:
                field = "nmse_" + c
                if field in part and field+"_baseline" in part:
                    scores = part[["subject", "session"]].copy()
                    scores["delta"] = part[field] - part[field+"_baseline"]
                    row["delta_NMSE_"+c] = scores.groupby(["subject", "session"]).delta.mean().groupby("subject").mean().mean()
                    for suffix, label in [("", "candidate"), ("_baseline", "baseline")]:
                        value = part.groupby(["subject", "session"])[field+suffix].mean().groupby("subject").mean().mean()
                        row[label+"_NRMSE_"+c] = np.sqrt(value)
            paired.append(row)
    paired = pd.DataFrame(paired)
    paired.to_csv(out / "paired_baseline.csv", index=False)

    # Saved states include last-valid iterates of failed fits; label them explicitly.
    import importlib.util
    import sys
    source_root = run / "source_snapshot"
    if not source_root.is_dir():
        raise ValueError("Nonlinear report requires the run's frozen source_snapshot")
    operator_path = source_root / "src/inference/observation_baselines.py"
    sys.path.insert(0, str(source_root))
    operator_spec = importlib.util.spec_from_file_location("report_frozen_observation_baselines", operator_path)
    operator_module = importlib.util.module_from_spec(operator_spec)
    operator_spec.loader.exec_module(operator_module)
    native_feature_operators = operator_module.native_feature_operators
    visible_feature_interpolation = operator_module.visible_feature_interpolation
    state_rows, parameter_rows, profiles, own_rows, observation_rows = [], [], [], [], []
    for group in data.group.unique():
        preparation = json.loads((run / "prepared" / f"{group}.json").read_text())
        spec = preparation["spec"]
        immutable_target = immutable_normalizer = None
        if preparation["status"] == "completed":
            meta = preparation["metadata"]
            if spec["kind"] == "measured":
                projection = meta["projection"]
                observation_rows.append(dict(group=group, Hb_pair_index=projection["fnirs_pair"],
                    EEG_factor=projection["eeg_factor"], Hb_factor=projection["fnirs_factor"],
                    absolute_calibration=projection["observation_loading"]["absolute_calibration"]))
            with np.load(run / "prepared" / f"{group}.npz", allow_pickle=False) as a:
                if extended:
                    immutable_target, immutable_normalizer = a["target"].copy(), a["normalizer"].copy()
                n = a["target"].shape[1]; left, right = (n-16)//2, (n+16)//2
                op = native_feature_operators(n)
                trials = data.loc[data.group.eq(group), "trial"].unique()
                for mode in ("center_EEG", "center_fNIRS"):
                    eeg_gap = (left, right) if mode == "center_EEG" else (0, 0)
                    hb_gap = (int(2.5*left), int(2.5*right)) if mode == "center_fNIRS" else (0, 0)
                    pe = op["eeg"] @ visible_feature_interpolation(n, *eeg_gap)
                    ph = op["fnirs"] @ visible_feature_interpolation(int(2.5*n), *hb_gap)
                    columns = [0] if mode == "center_EEG" else [1, 2]
                    seen = np.r_[0:left, right:n]
                    for trial in trials:
                        pred = np.column_stack((pe @ a["feature_eeg"][trial], ph @ a["feature_fnirs"][trial]))
                        row = dict(kind=spec["kind"], group=group, mode=mode, trial=int(trial), subject=spec["subject"], session=meta["trials"][trial]["session"])
                        for col in columns:
                            pred[:, col] = np.interp(np.arange(n), seen, pred[seen, col])
                            row["nmse_"+COMPONENTS[col]] = float(np.mean(((pred[left:right,col]-a["target"][trial,left:right,col])/a["normalizer"][col])**2))
                        own_rows.append(row)
        parameter_values = preparation.get("selected_tau", {})
        if fixed_roi_tau:
            parameter_values = {methods[0]: config['reference_tau_s']}
            for method in methods[1:]:
                if method in ("fixed_tau_logflow1","fixed_tau_logflow1_tied"):
                    parameter_values[method] = config["reference_tau_s"]
                    continue
                selection_path = run/'training'/f'{group}__{method}'/'selection.json'
                selection = json.loads(selection_path.read_text()) if selection_path.exists() else {}
                parameter_values[method] = selection.get('parameter_value') if selection.get('status') == 'completed' else np.nan
        for method, tau in parameter_values.items():
            parameter_rows.append(dict(kind=spec["kind"], group=group, subject=spec["subject"], outer=spec["outer"], method=method, tau=tau, true_tau=spec.get("true_tau", np.nan)))
        for point in preparation.get("profiles", []):
            profiles.append(dict(group=group, kind=spec["kind"], **point))
        for method, mode in data[data.group == group][["method", "mode"]].drop_duplicates().itertuples(index=False, name=None):
            cell = run / "cells" / f"{group}__{method}__{mode}" / "trajectories.npz"
            if not cell.exists():
                continue
            with np.load(cell, allow_pickle=False) as arrays:
                if extended and immutable_target is not None:
                    if not np.array_equal(arrays["target"], immutable_target[arrays["trial_indices"]], equal_nan=True) or not np.array_equal(arrays["normalizer"], immutable_normalizer, equal_nan=True):
                        raise ValueError("Gain/prior cell changed frozen target or scoring scale")
                for index, trial in enumerate(arrays["trial_indices"]):
                    identity = data[(data.group == group) & (data.method == method) & (data["mode"] == mode) & (data.trial == trial)].iloc[0]
                    state = arrays["states"][index]
                    row = dict(kind=identity.kind, group=group, method=method, mode=mode, trial=int(trial), status=identity.status, true_tau=spec.get('true_tau'),true_initial_log_p_over_v=spec.get('true_initial_log_p_over_v'),converged=bool(successful.loc[identity.name]), finite_states=bool(np.isfinite(state).all()))
                    if row["finite_states"]:
                        for column, name in enumerate(["r", "s", "f", "v", "p", "q"]):
                            row[name+"_initial"] = state[0, column]
                            row[name+"_minimum"] = state[:, column].min()
                            row[name+"_maximum"] = state[:, column].max()
                        if extended:
                            r = state[:, 0]
                            if fixed_roi_flow:
                                dt=config['tensor']['dt_s']
                                row.update(driver_ac_rms=float(np.std(r)),flow_below_half_s=float(np.sum(state[:,2]<.5)*dt),
                                    flow_above_two_s=float(np.sum(state[:,2]>2)*dt),flow_below_01_s=float(np.sum(state[:,2]<.1)*dt))
                            center = slice((len(r)-16)//2, (len(r)+16)//2)
                            row.update(driver_mean=float(r.mean()), driver_rms=float(np.sqrt(np.mean(r*r))),
                                driver_center_rms=float(np.sqrt(np.mean(r[center]**2))),
                                driver_difference_rms=float(np.sqrt(np.mean(np.diff(r)**2))),
                                driver_center_difference_rms=float(np.sqrt(np.mean(np.diff(r[center])**2))))
                            for column, name in enumerate(["r", "s", "f", "v", "p", "q"]):
                                row[name+"_center_minimum"] = state[center,column].min()
                                row[name+"_center_maximum"] = state[center,column].max()
                        if fixed_roi_initial:
                            row['path_max_abs_p_minus_v']=float(np.max(np.abs(state[:,4]-state[:,3])))
                            row['initial_log_p_over_v']=float(np.log(state[0,4]/state[0,3]))
                        row["initial_max_fractional_excursion"] = np.max(abs(state[0, 2:] - 1))
                        row["path_max_fractional_excursion"] = np.max(abs(state[:, 2:] - 1))
                    if spec["kind"] == "synthetic" and row["finite_states"] and "truth_states" in arrays:
                        if fixed_roi_flow:
                            actual=arrays['truth_states'][index]; residual=state-actual
                            row.update(driver_dc_error=float(residual[:,0].mean()),flow_mean_error=float(residual[:,2].mean()),
                                driver_ac_rmse=float(np.std(residual[:,0])),initial_rmse=float(np.sqrt(np.mean(residual[0,1:]**2))))
                        for column, name in enumerate(["r", "s", "f", "v", "p", "q"]):
                            row["truth_rmse_"+name] = float(np.sqrt(np.mean((state[:,column]-arrays["truth_states"][index,:,column])**2)))
                    state_rows.append(row)
    states = pd.DataFrame(state_rows)
    states.to_csv(out / "state_ranges.csv", index=False)
    observation_map = pd.DataFrame(observation_rows)
    observation_map.to_csv(out / "observation_map.csv", index=False)
    state_summary = []
    state_kind = "measured" if "measured" in data.kind.unique() else "synthetic"
    for method in methods:
        part = states[(states.kind == state_kind) & (states.method == method)
                      & (states["mode"] == "full") & states.converged]
        state_summary.append(dict(method=method, successful=len(part), expected=int(agg.loc[
            (agg.kind == state_kind) & (agg.method == method) & (agg["mode"] == "full"), "expected"].sum()),
            f_min=part.f_minimum.min() if len(part) else np.nan,
            f_max=part.f_maximum.max() if len(part) else np.nan,
            f_below_01=int(part.f_minimum.lt(.1).sum()) if len(part) else 0,
            initial_over_05=int(part.initial_max_fractional_excursion.gt(.5).sum()) if len(part) else 0,
            path_over_05=int(part.path_max_fractional_excursion.gt(.5).sum()) if len(part) else 0))
    state_summary = pd.DataFrame(state_summary)
    state_summary.to_csv(out / "state_summary.csv", index=False)
    parameters = pd.DataFrame(parameter_rows, columns=["kind", "group", "subject", "outer", "method", "tau", "true_tau"])
    parameters.to_csv(out / "parameters.csv", index=False)
    pd.DataFrame(profiles).to_csv(out / "tau_profiles.csv", index=False)
    if not extended:
        fig, axes = plt.subplots(2, 1, figsize=(11, 7))
        measured_groups = list(parameters.loc[parameters.kind.eq("measured"), "group"].unique())
        for method in methods:
            p = parameters[(parameters.kind == "measured") & (parameters.method == method)].set_index("group")
            if len(p):
                axes[0].plot(range(len(measured_groups)), p.reindex(measured_groups).tau, "o-", label=labels[method])
        axes[0].set_xticks(range(len(measured_groups)), measured_groups, rotation=45, ha="right", fontsize=7)
        axes[0].set_ylabel("τ / s")
        if measured_groups: axes[0].legend()
        for group in measured_groups:
            points = [p for p in profiles if p["group"] == group]
            if points:
                points = sorted(points, key=lambda p:p["tau"])
                values = np.array([p["train_penalized"] for p in points])
                axes[1].plot([p["tau"] for p in points], values-values.min(), label=group)
        axes[1].set_xlabel("τ / s"); axes[1].set_ylabel("训练线性目标 − 最小值")
        if measured_groups: axes[1].legend(fontsize=6, ncol=3)
        save(fig, "nonlinear_tau", "τ由训练线性剖面冻结后用于非线性拟合；不是非线性τ估计或似然置信区间。")

    fig, axes = plt.subplots(1, 2, figsize=(11, 4))
    for ax, field in zip(axes, ["initial_max_fractional_excursion", "path_max_fractional_excursion"]):
        for i, method in enumerate(methods):
            part = states[(states.kind == state_kind) & (states.method == method) & (states["mode"] == "full")]
            if field not in part: continue
            for ok, marker in [(True, "o"), (False, "x")]:
                values = part.loc[part.converged.eq(ok), field].dropna().to_numpy()
                ax.scatter(i + np.linspace(-.12,.12,len(values)), values, marker=marker, s=12, label="成功" if i==0 and ok else ("失败最后有效迭代" if i==0 else None))
        ax.set_xticks(range(len(methods)), [labels[m] for m in methods], rotation=20, ha="right", fontsize=7)
        ax.set_ylabel(field); ax.legend(fontsize=7)
    save(fig, "nonlinear_states", "初态与全路径f/v/p/q相对静息偏离；失败的最后有效迭代用叉号保留。0.25小信号阈值不作为非线性生理合格界。")

    notes = []
    example_source = (parent_data[parent_data.kind.eq('measured')] if 'measured' in data.kind.unique() else data) if fixed_roi_initial else (parent_data[parent_data.kind.isin(data.kind.unique())] if fixed_roi_flow else data)
    examples = example_source[(example_source.kind == "measured") & (example_source.method == methods[0]) & (example_source["mode"] == "full")].copy()
    if examples.empty:
        examples = example_source[(example_source.method == methods[0]) & (example_source["mode"] == "full")].copy()
    error_fields = ["last_valid_nmse_"+c for c in COMPONENTS]
    examples["rank_error"] = examples.reindex(columns=error_fields).mean(axis=1)
    examples = examples.sort_values(["rank_error", "group", "trial"], kind="stable", na_position="last")
    selected = [("median", examples.iloc[(len(examples)-1)//2]), ("worst", examples.iloc[-1])] if len(examples) else []
    for label, example in selected:
        group, trial = example.group, int(example.trial)
        fig, axes = plt.subplots(3, 2, figsize=(12, 8), sharex=True)
        target_drawn = False
        for method in methods:
            path = run / "cells" / f"{group}__{method}__full" / "trajectories.npz"
            row = data[(data.group==group)&(data.trial==trial)&(data.method==method)&(data["mode"]=="full")]
            status = row.status.iloc[0] if len(row) else "missing"
            notes.append(f"{label} / {group} / trial {trial} / {method}: {status}")
            if not path.exists():
                if extended:
                    for ax in axes[:,0]:
                        ax.plot([], [], ":", label=labels[method]+f" 未估计 ({status})")
                    if method == methods[-1]:
                        for ax in axes[:,1]:
                            ax.text(.5, .5, f"{labels[method]}状态未估计\n{status} / 无轨迹文件",
                                ha="center", va="center", transform=ax.transAxes, fontsize=9)
                continue
            with np.load(path, allow_pickle=False) as a:
                indices = np.flatnonzero(a["trial_indices"] == trial)
                if not len(indices): raise ValueError("Representative trial missing from cell")
                index = int(indices[0]); t = np.arange(a["prediction"].shape[1])*.25
                if example.kind == "measured": t -= 5.
                if "target" in a and not target_drawn:
                    for j,c in enumerate(COMPONENTS): axes[j,0].plot(t,a["target"][index,:,j],"k",lw=1.5,label="实测／目标")
                    target_drawn = True
                style = "-" if status == "completed" else ":"
                for j,c in enumerate(COMPONENTS):
                    unavailable = extended and not np.isfinite(a["prediction"][index,:,j]).any()
                    suffix = f" 未估计 ({status})" if unavailable else (" (失败迭代)" if status!="completed" else "")
                    axes[j,0].plot(t,a["prediction"][index,:,j],style,lw=1,label=labels[method]+suffix)
                    axes[j,0].set_ylabel(c)
                if method == (methods[-1] if extended else "linear_trained_tau_0p01"):
                    if extended and not np.isfinite(a["states"][index]).any():
                        for ax in axes[:,1]:
                            ax.text(.5, .5, f"{labels[method]}状态未估计\n{status} / 无有限状态",
                                ha="center", va="center", transform=ax.transAxes, fontsize=9)
                    for j in range(3):
                        for k in [2*j,2*j+1]: axes[j,1].plot(t,a["states"][index,:,k],style,label=["r","s","f","v","p","q"][k])
                        axes[j,1].legend()
        axes[0,0].legend(fontsize=7)
        if example.kind == "measured":
            for ax in axes.flat: ax.axvline(0., color="grey", lw=.7, ls=":")
        for ax in axes[-1]: ax.set_xlabel("相对事件 / 秒" if example.kind == "measured" else "秒")
        fig.suptitle(f"{label}: {group} / trial {trial}; {labels[methods[0]]}排序")
        save(fig, "nonlinear_curves_"+label, f"{labels[methods[0]]}基线的{label}样例；右侧{labels[methods[-1]] if extended else '训练τ弱正则'}状态，失败最后有效迭代为虚线，不按候选性能换样。")

    own = pd.DataFrame(own_rows)
    own_comparisons = []
    if len(own):
        own.to_csv(out / "own_context_rows.csv", index=False)
        for method in methods:
            merged = data[(data.method == method) & successful].merge(own, on=["kind", "group", "mode", "trial", "subject", "session"], suffixes=("", "_own"), validate="one_to_one")
            for (kind, mode), part in merged.groupby(["kind", "mode"]):
                expected = int(agg[(agg.kind == kind)&(agg.method == method)&(agg["mode"] == mode)].expected.iloc[0])
                row = dict(kind=kind, method=method, mode=mode, common_success=len(part), expected=expected)
                for c in COMPONENTS:
                    for suffix, name in [("", "candidate"), ("_own", "own_context")]:
                        field = "nmse_"+c+suffix
                        if field in part:
                            value = part.groupby(["subject", "session"])[field].mean().groupby("subject").mean().mean()
                            row[name+"_NRMSE_"+c] = np.sqrt(value)
                own_comparisons.append(row)
    own_comparisons = pd.DataFrame(own_comparisons)
    own_comparisons.to_csv(out / "paired_own_context.csv", index=False)
    synthetic = data[data.kind.eq("synthetic")].copy()
    synthetic["successful"] = successful.loc[synthetic.index]
    recovery = []
    for (method, mode), part in synthetic.groupby(["method", "mode"]):
        good = part[part.successful]
        values = pd.to_numeric(good.get("driver_nrmse", pd.Series(dtype=float)), errors="coerce").dropna()
        recovery.append(dict(method=method, mode=mode, observed=len(part), completed=len(good), driver_scored=len(values), driver_mean=values.mean(), driver_max=values.max()))
    recovery = pd.DataFrame(recovery)
    recovery.to_csv(out / "synthetic_recovery.csv", index=False)
    # Communication tables are deliberately narrower than the retained CSV schema.
    kind_labels = {"synthetic": "合成", "measured": "实测"}
    mode_labels = {"full": "全观测", "center_EEG": "中心EEG隐藏", "center_fNIRS": "中心Hb隐藏"}

    def compact_group(value):
        return str(value).replace("subject_", "S").replace("synthetic_tau", "T").replace("mixed", "混合").replace("slow", "慢变")

    def section_tables(frame, builder):
        chunks = []
        for (kind, mode), part in frame.groupby(["kind", "mode"], sort=False):
            chunks.append(f"### {kind_labels.get(kind,kind)} · {mode_labels.get(mode,mode)}\n\n" + table(builder(part)))
        return "\n\n".join(chunks)

    def aggregate_display(part):
        rows = []
        for _, r in part.iterrows():
            rows.append({"方法": labels[r.method], "成功/计划": f"{r.completed}/{r.expected}",
                "失败/缺失": f"{r.failed}/{r.missing}", "三项<0.5": r.joint_below_05,
                **{c: r["NRMSE_"+c] for c in COMPONENTS}})
        return pd.DataFrame(rows)

    def paired_display(part, own=False):
        rows = []
        for _, r in part.iterrows():
            for c in COMPONENTS:
                candidate = r.get("candidate_NRMSE_"+c, np.nan)
                reference = r.get(("own_context" if own else "baseline")+"_NRMSE_"+c, np.nan)
                if not np.isfinite(candidate) and not np.isfinite(reference):
                    continue
                row = {"方法": labels[r.method], "分量": c, "共同/计划": f"{r.common_success}/{r.expected}",
                    "候选": candidate, "自身上下文" if own else (labels[methods[0]] if extended else "固定强正则"): reference}
                if not own: row["ΔNMSE"] = r.get("delta_NMSE_"+c, np.nan)
                rows.append(row)
        return pd.DataFrame(rows)

    aggregate_text = section_tables(agg, aggregate_display)
    paired_text = section_tables(paired, paired_display)
    own_text = section_tables(own_comparisons, lambda part: paired_display(part, own=True)) if len(own_comparisons) else ""
    failure_text = section_tables(failures, lambda part: part.assign(
        method=part.method.map(labels), status=part.status.map({"completed":"成功", "failed_numerical":"数值失败",
        "failed_preparation":"准备失败", "failed_training":"训练失败", "failed_exception":"异常", "failed_integration_check":"积分失败"}).fillna(part.status)
    )[["method","status","rows"]].rename(columns={"method":"方法","status":"状态","rows":"条数"}))
    recovery_text = ""
    for mode, part in recovery.groupby("mode", sort=False):
        display = pd.DataFrame([{"方法":labels[r.method], "成功/记录":f"{r.completed}/{r.observed}",
            "驱动评分数":r.driver_scored, "驱动均误差":r.driver_mean, "驱动最大误差":r.driver_max}
            for r in part.itertuples()])
        recovery_text += f"### 合成 · {mode_labels.get(mode,mode)}\n\n" + table(display) + "\n"
    parameter_text = ""
    for kind, part in parameters.groupby("kind", sort=False):
        rows = []
        for group, values in part.groupby("group", sort=False):
            r = {"被试/折或合成组":compact_group(group), "真τ":values.true_tau.iloc[0]}
            r.update({labels[v.method]:v.tau for v in values.itertuples()})
            rows.append(r)
        parameter_text += f"### {kind_labels.get(kind,kind)} · τ（秒）\n\n" + table(pd.DataFrame(rows)) + "\n"
    observation_display = observation_map.rename(columns={"group":"被试/折", "Hb_pair_index":"Hb对", "EEG_factor":"EEG因子", "Hb_factor":"Hb因子", "absolute_calibration":"绝对标定"}).copy()
    if len(observation_display):
        observation_display["被试/折"] = observation_display["被试/折"].map(compact_group)
        observation_display["绝对标定"] = observation_display["绝对标定"].map({True:"是",False:"否"})
    state_display = pd.DataFrame([{"方法":labels[r.method], "成功/计划":f"{r.successful}/{r.expected}",
        "f最小":r.f_min, "f最大":r.f_max, "f<0.1":r.f_below_01,
        "初态>0.5":r.initial_over_05, "路径>0.5":r.path_over_05} for r in state_summary.itertuples()])

    text = "# 共享驱动非线性拟合：正则敏感性与重建\n\n"
    text += f"来源运行：`{run.name}`。计划cell {summary['expected_cells']}，终态{summary['terminal_cells']}；计划评价{summary['expected_rows']}，已记录{summary['observed_rows']}，收敛成功{summary['completed_rows']}。cell完成不等于其内部拟合成功。\n\n"
    text += "各方法均预先指定；没有按留出结果选胜者。λ100与λ0.01比较使用相同初态惩罚100。两个适配τ方法只使用训练线性剖面，τ在非线性评价时冻结，不声称完成非线性生理参数学习。\n\n"
    text += "## 完整分母与误差\n\n" + aggregate_text + "\n"
    text += "EEG是处理后的log-power/PCA特征，不是原始电压；三分量使用冻结外折训练SD。NRMSE来自成功收敛且积分分辨率检查通过的子集：trial→session→subject等权NMSE后取根。失败/缺失始终保留在expected中；成功数小于计划数时不能将条件均值写成完整面板结果。三项<0.5是逐trial三分量均达阈值且成功的条数；不是由聚合均值推算。模式full为全观测重建，center模式只评分隐藏分量。数学域有效不等同于经验生理合理、独立共享信息或teacher资格。\n\n"
    text += "## 配对基线比较\n\n相对固定τ/λ100，以下NMSE差只使用双方共同成功的相同身份；负值为改善，完整分母同时保留。\n\n" + paired_text + "\n"
    if len(own_comparisons):
        text += "## 自身上下文对照\n\n先在原生特征可见支持插值，再执行原观测处理，最后在处理坐标内插补隐藏段；与既有线性诊断基线一致。算子从本run的source_snapshot加载；只比较候选成功的相同身份，未拟合生理参数。\n\n" + own_text + "\n"
    text += "## 失败与收敛\n\n" + failure_text + "\n完整梯度、预算、起点与积分差异见convergence.csv；每次起点结果仍由原cell/result.json持有。失败最后有效轨迹不进入主误差或合成恢复均值。\n\n"
    text += "## 合成恢复与参数\n\n" + recovery_text + "\n"
    text += parameter_text + "\n合成驱动误差不做事后幅度、符号或时间对齐；合成条件不等同实测生理真值。参数跨折差异受重叠训练集合与通道选择影响，不能直接解释为真实被试差异。\n\n"
    if len(observation_map):
        text += "## 观测映射的解释边界\n\n下表直接取自本run的prepared元数据。EEG与Hb幅度因子来自训练MAD和冻结参考模型的协方差规范，absolute_calibration为False；Hb对索引在部分被试内随折改变。它们没有提供独立跨模态生理标定。固定EEG loading定义驱动单位，再固定neurovascular gain仍会额外固定二者的有效幅度比。\n\n" + table(observation_display) + "\n"
        text += "原方程ds/dt=βr−κs−γ(f−1)、df/dt=s不含τ：固定r与s/f初态时，调τ不能直接修复异常血流；τ只改变下游v/p/q，重新拟合时可通过驱动与初态补偿。下一轮若研究有效增益，必须保留当前评分尺度，并将增益与绝对生理效能区别开。\n\n"
    text += "## 状态范围诊断\n\n" + table(state_display) + "\n统计限成功full轨迹，失败始终保留在完整分母中。f<0.1及相对静息偏离>0.5只是定位极端状态的描述阈值，并非文献验证的生理正常界限。\n\n"
    text += "## 状态与固定样例\n\nstate_ranges.csv逐条保留六状态初值、最小值、最大值及成功标志；合成另保留各状态原坐标真值RMSE。样例按固定τ/λ100最后有效三分量平均NMSE排序取中位和最差；无有限输出排在末尾并明确保留。右侧展示训练τ/λ0.01状态。\n\n" + "\n\n".join(notes) + "\n\n"
    if gain_prior:
        text = text.replace("共享驱动非线性拟合：正则敏感性与重建", "共享驱动非线性拟合：幅度先验与有效增益")
        text = text.replace("各方法均预先指定；没有按留出结果选胜者。λ100与λ0.01比较使用相同初态惩罚100。两个适配τ方法只使用训练线性剖面，τ在非线性评价时冻结，不声称完成非线性生理参数学习。",
            "四臂为预先指定的2×2比较：A固定β，B固定β加驱动幅度先验，C训练共享β，D训练共享β加幅度先验；不按验证结果选胜者。τ固定2秒，曲率惩罚0.01、初态惩罚100。β仅使用每组18条训练数据估计，多起点全部终态后按收敛训练目标选择，并在验证时冻结。")
        text = text.replace("相对固定τ/λ100", "相对A固定β无幅度先验基线")
        text = text.replace("固定EEG loading定义驱动单位，再固定neurovascular gain仍会额外固定二者的有效幅度比。",
            "所有臂的观测算子、目标、均值处理、训练SD与通道映射不变；没有为了β重新调整EEG或Hb观测尺度。β是这套任意观测单位中的有效耦合，不能解释为被试神经血管生理效能。")
        text = text.replace("下一轮若研究有效增益，必须保留当前评分尺度，并将增益与绝对生理效能区别开。", "本轮只释放有效增益，仍须区分动力学补偿与生理恢复。")
        text = text.replace("样例按固定τ/λ100", "样例按A固定β无幅度先验")
        text = text.replace("右侧展示训练τ/λ0.01状态。", "右侧展示D训练β加幅度先验状态。")
        text = text.replace(parameter_text, "τ在所有臂固定为2秒；训练β见下文。\n") if parameter_text else text
        if config.get("continuation"):
            parent_name = Path(config["continuation"]["parent_run"]).name
            text = text.replace("共享驱动非线性拟合：幅度先验与有效增益", "共享驱动非线性拟合：训练预算热启动对照")
            text = text.replace("## 完整分母与误差", (
                f"本轮延续父运行`{parent_name}`：仅符合合同的{config['continuation']['expected_measured_restarts']}条实测训练起点谱系从最后有效解热启动，"
                "LM阻尼重置，目标、观测尺度和成功阈值不变；这不是精确恢复优化器状态。其余实测训练终态、全部合成训练/验证以及实测A/B验证继承；"
                "实测C/D验证按新训练选择重算。继承结果不构成新增独立复制，下表分母仍是原有身份集合。\n\n## 完整分母与误差"))
        text += gain_prior_details(run, out, config, data, states, labels, save, selected)
    elif fixed_roi_tau:
        text = text.replace("共享驱动非线性拟合：正则敏感性与重建", "固定AF7Fp1位置：共享τ与轨迹恢复")
        text = text.replace("各方法均预先指定；没有按留出结果选胜者。λ100与λ0.01比较使用相同初态惩罚100。两个适配τ方法只使用训练线性剖面，τ在非线性评价时冻结，不声称完成非线性生理参数学习。",
            "三臂预先指定：同ROI固定τ2、18条训练轨迹共享非线性τ、带工程log prior的共享τ；β固定1，曲率惩罚0.01、初态惩罚100，无driver幅度ridge。训练起点全部终态后按收敛目标选择，验证τ冻结，不按留出成绩选择方法。")
        text = text.replace("三分量使用冻结外折训练SD", "EEG沿用父训练SD；固定ROI的Hb分量SD仅由当前18条训练轨迹计算")
        text = text.replace("相对固定τ/λ100", "相对同一新ROI目标的固定τ2基线")
        text = text.replace("参数跨折差异受重叠训练集合与通道选择影响", "固定ROI消除了本轮跨折选对变化，但参数跨折差异仍受重叠训练集合、观测幅度规范与初态补偿影响")
        text = text.replace("Hb对索引在部分被试内随折改变", "本轮各被试、各折均固定AF7Fp1（pair0）")
        text = text.replace("下一轮若研究有效增益，必须保留当前评分尺度，并将增益与绝对生理效能区别开。", "本轮β仍固定1；τ自由度不等于直接控制血流或已有生理参数资格。")
        text = text.replace("样例按固定τ/λ100", "样例按同ROI固定τ2基线")
        text = text.replace("右侧展示训练τ/λ0.01状态。", "右侧展示同ROI先验τ臂状态。")
        if parameter_text: text = text.replace(parameter_text, "τ由训练选择冻结，实际选择及失败见后文。\n")
        text += fixed_roi_tau_details(run, out, config, data, states, labels, save)
        if fixed_roi_flow:
            if set(data.kind)=={'synthetic'}:
                text=text.replace('## 完整分母与误差','本报告仅覆盖已终态合成阶段；实测阶段未纳入，未启动实测不记失败或零误差。状态范围及固定样例均来自合成身份，不作实测结论。\n\n## 完整分母与误差')
            text=text.replace('固定AF7Fp1位置：共享τ与轨迹恢复','固定AF7Fp1位置：log-flow软约束对照')
            text=text.replace('三臂预先指定：同ROI固定τ2、18条训练轨迹共享非线性τ、带工程log prior的共享τ；','六臂预先指定：继承父运行固定τ2和训练τ；新增固定τ2/w1及训练τ/w0.1、1、4，无τ先验；')
            text=text.replace('右侧展示同ROI先验τ臂状态。','右侧展示训练τ/w4臂状态，不是按性能选择。')
            text=text.replace('EEG投影保持父规范，Hb目标和训练SD因固定位置而改变；仅本轮同一新目标下重算的τ2作为配对基线。旧parent的τ和误差只可作为背景，未并入本轮paired scores。','prepared从确切固定ROI父运行原样继承，EEG/Hb目标、观测算子与SD均不改变；父两臂的原身份、状态和误差继承后作为配对基线，不是新的重复实验。')
            text += fixed_roi_flow_details(run,out,config,data,states,labels,save,selected)
            if fixed_roi_initial:
                text=text.replace('固定AF7Fp1位置：log-flow软约束对照','固定AF7Fp1位置：free/tied初态约束对照')
                text=text.replace('六臂预先指定：继承父运行固定τ2和训练τ；新增固定τ2/w1及训练τ/w0.1、1、4，无τ先验；','四臂预先指定：固定τ2和训练τ各自比较free/tied初态，所有flow权重为1，无τ先验；')
                text=text.replace('## 完整分母与误差','合成主解释必须按true log(p0/v0)=−0.05/0/+0.05分层；本节合并合成数仅作整体完整性摘要，详细匹配/错设偏差见后文。仅实测free两臂继承父结果，合成全部为新真值家族；继承不构成独立复制。\n\n## 完整分母与误差')
                text=text.replace('六臂及权重预声明，不按验证成绩选择w。','四臂预声明，flow权重全为1；不按验证成绩选择free或tied。')
                text=text.replace('不同w的总目标不可横向排名；多起点仅在同组同臂内选择。继承父运行的两臂及原prepared观测不是新独立复制。','不同初态模型的总目标不作主评分胜负；多起点仅在同组同臂内选择。仅实测free结果和prepared继承，合成全新生成/拟合。')
                text=text.replace('继承父运行的两臂及原prepared观测不是新独立复制。','仅实测free两臂继承；合成四臂为新真值家族。')
                text=text.replace('右侧展示训练τ/w4臂状态，不是按性能选择。','右侧展示共享τ/tied臂状态，不是按性能选择。')
                text=text.replace('prepared从确切固定ROI父运行原样继承，EEG/Hb目标、观测算子与SD均不改变；父两臂的原身份、状态和误差继承后作为配对基线，不是新的重复实验。','实测prepared与free两臂结果从确切flow父运行继承；合成采用本轮新初态真值家族，四臂共享同一目标。继承不是独立复制。')
                text += fixed_roi_initial_details(run,out,config,data,states,labels,save)

            if 'measured' in data.kind.unique():
                text += flow_hbt_decomposition(run,out,config,data,labels,save,selected)
                text += flow_full_curve_atlas(run,out,config,data,labels,[methods[1],methods[3]] if fixed_roi_initial else None)
    for name,caption in figures: text += f"![{name}](figures/{name}.png)\n\n{caption}\n\n"
    text += "## 交付验证\n\n图形以240 dpi PNG嵌入，PDF正文表格可搜索。程序核验图像对象与文本提取；WPS滚动缩放未检查。\n"
    export_pdf(text, out, figures, extra_css="table{font-size:8pt;width:100%}td,th{padding:4pt}h3{font-size:10pt}")


def render_fixed_roi_optical(run,out):
    """Completed retained evidence only; distinct optical targets never pooled."""
    import hashlib
    import yaml
    if out.exists() and any(out.iterdir()):raise ValueError('Choose a new empty versioned export directory')
    manifest=json.loads((run/'manifest.json').read_text())
    cfg=yaml.safe_load((run/'resolved_config.yaml').read_text())
    methods=['fixed_tau_logflow1','trained_tau_logflow1'];pipelines=['no_motion','mne_tddr'];modes=['full','center_EEG','center_fNIRS']
    if (manifest.get('execution')!='completed' or manifest.get('pilot') is not False
            or manifest.get('experiment_id')!='SSM-SHARED-DRIVER-FIXED-ROI-OPTICAL-v1'
            or manifest.get('synthetic_terminal') is not True or manifest.get('measured_terminal') is not True
            or cfg.get('schema')!='shared_driver_fixed_roi_optical_v1' or cfg.get('methods')!=methods
            or cfg.get('optical_pipelines')!=pipelines or cfg.get('modes')!=modes
            or cfg.get('subjects')!=['subject_01','subject_09','subject_18'] or cfg.get('outer_folds')!=[0,1,2,3]
            or cfg.get('parent_run')!='experiments/runs/physiology_semantic_tokenizer/shared_driver_reconstruction/20260926_fixed_roi_flow_v1'
            or cfg.get('optical_audit_run')!='experiments/runs/physiology_semantic_tokenizer/shared_driver_reconstruction/20260926_optical_motion_audit_v1'):
        raise ValueError('Optical report requires exact completed nonpilot optical contract')
    snapshot=Path(manifest.get('source_root',''))
    if snapshot.resolve()!=(run/'source_snapshot').resolve() or not snapshot.is_dir():raise ValueError('Missing owning frozen source snapshot identity')
    root=Path(manifest['project_root']);parent=root/cfg['parent_run'];audit=root/cfg['optical_audit_run']
    pm=json.loads((parent/'manifest.json').read_text());am=json.loads((audit/'manifest.json').read_text())
    control=json.loads((audit/'synthetic_summary.json').read_text());reproduction=json.loads((audit/'measured_summary.json').read_text())
    if (pm.get('execution')!='completed' or pm.get('experiment_id')!='SSM-SHARED-DRIVER-FIXED-ROI-FLOW-v1'
            or am.get('experiment_id')!='SSM-SHARED-DRIVER-OPTICAL-MOTION-AUDIT-v1'
            or am.get('execution')!='completed' or am.get('result_status')!='completed'
            or reproduction.get('current_pipeline_reproduced') is not True or reproduction.get('windows')!=72
            or len(reproduction.get('reproduction_checks',[]))!=216
            or any(c.get('matches') is not True for c in reproduction['reproduction_checks']) or control.get('status')!='completed'):
        raise ValueError('Completed parent and optical reproduction evidence required')
    checks=reproduction['reproduction_checks'];sample_ids={c.get('sample_id') for c in checks}
    if len(sample_ids)!=72 or None in sample_ids or {(c.get('sample_id'),c.get('stage')) for c in checks}!={(sample,stage) for sample in sample_ids for stage in ['od','hb_native','hb_4hz']}:
        raise ValueError('Optical source reproduction identities incomplete')
    data=pd.read_csv(run/'metrics.csv');summary=json.loads((run/'summary.json').read_text())
    keys=['kind','group','method','mode','trial'];expected=set();preps={}
    for kind in ['synthetic','measured']:
        groups=([f'synthetic_tau{tau:g}_{spectrum}_r{rep}' for tau in [1.,2.,4.] for spectrum in ['slow','mixed'] for rep in range(2)]
            if kind=='synthetic' else [f'{pipeline}__{subject}_o{fold}' for pipeline in pipelines for subject in cfg['subjects'] for fold in range(4)])
        for group in groups:
            prep=json.loads((run/'prepared'/f'{group}.json').read_text());preps[group]=prep;spec=prep['spec']
            if spec['group']!=group or spec['kind']!=kind:raise ValueError('Prepared identity mismatch')
            for method in methods:
                for mode in modes:
                    path=run/'cells'/f'{group}__{method}__{mode}'/'result.json';cell=json.loads(path.read_text())
                    if cell['cell']!=path.parent.name or len(cell['rows'])!=6:raise ValueError('Terminal cell identity/count mismatch')
                    if kind=='synthetic':
                        original=json.loads((parent/'cells'/path.parent.name/'result.json').read_text())
                        if cell['rows']!=original['rows'] or cell.get('inheritance',{}).get('status')!='inherited_nonindependent':
                            raise ValueError('Synthetic rows changed rather than inherited')
                    for row in cell['rows']:
                        trial=row['trial'];identity=(kind,group,method,mode,trial)
                        expected_row=dict(spec,method=method,mode=mode,session=cfg['sessions'][trial//8])
                        if any(row.get(k)!=v for k,v in expected_row.items()):raise ValueError('Cell row differs from prepared identity')
                        if trial not in [i for i in range(24) if i%4==spec['outer']] or identity in expected:
                            raise ValueError('Validation split/duplicate mismatch')
                        expected.add(identity)
                        match=data[(data.kind==kind)&(data.group==group)&(data.method==method)&(data['mode']==mode)&(data.trial==trial)]
                        if len(match)!=1 or match.iloc[0].status!=row['status'] or bool(truth(match.converged).iloc[0])!=bool(row.get('converged',False)):
                            raise ValueError('Metrics differ from retained cell status')
                        for component in COMPONENTS:
                            k='nmse_'+component
                            if row.get(k) is not None and not np.isclose(float(match.iloc[0][k]),row[k],rtol=1e-12,atol=1e-14):
                                raise ValueError('Metrics differ from retained cell errors')
    if len(data)!=1296 or data.duplicated(keys).any() or set(data[keys].itertuples(index=False,name=None))!=expected or summary.get('observed_rows')!=1296 or summary.get('terminal_cells')!=216:
        raise ValueError('Optical report requires complete 1296 rows and 216 cells')
    out.mkdir(parents=True,exist_ok=True);(out/'figures').mkdir()
    font=Path('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc')
    if font.exists():font_manager.fontManager.addfont(str(font));plt.rcParams['font.family']=font_manager.FontProperties(fname=str(font)).get_name()
    plt.rcParams.update({'font.size':9,'axes.unicode_minus':False})
    labels=dict(zip(methods,['固定τ=2','18训练trial共享τ']));figures=[];pages=[]
    def save(fig,name):
        fig.tight_layout();fig.savefig(out/'figures'/f'{name}.png',dpi=240);plt.close(fig);figures.append(name)
        return f'![{name}](figures/{name}.png)'
    def good(part):return part[part.status.eq('completed')&truth(part.converged)]
    def errors(part):
        cols=['nmse_'+c for c in COMPONENTS]
        if not len(part):return [np.nan]*3
        return np.sqrt(part.groupby(['subject','session'])[cols].mean().groupby('subject').mean().mean()).tolist()
    measured=data[data.kind.eq('measured')].copy();aggregates=[]
    for (pipeline,method,mode),part in measured.groupby(['optical_pipeline','method','mode'],sort=False):
        if len(part)!=72:raise ValueError('Each pipeline/method/mode needs all 72 rows')
        ok=good(part);err=errors(ok);aggregates.append(dict(pipeline=pipeline,method=method,mode=mode,expected=72,completed=len(ok),
            all_three_below05=int(ok[['nrmse_'+c for c in COMPONENTS]].lt(.5).all(axis=1).sum()) if mode=='full' else None,
            **dict(zip(COMPONENTS,err))))
    ag=pd.DataFrame(aggregates);ag.to_csv(out/'pipeline_metrics.csv',index=False)
    failure=measured.groupby(['optical_pipeline','method','mode','status'],dropna=False).size().reset_index(name='rows')
    failure.to_csv(out/'all_terminal_status_counts.csv',index=False)
    data.to_csv(out/'all_validation_rows.csv',index=False)
    pages.append('# 光学处理敏感性与共享生理轨迹\n\n同一3被试、72窗口、AF7Fp1位置与EEG，比较 no_motion / MNE TDDR 两种光学目标。每条管道分别报告固定τ与训练τ；初态五个自由坐标，flow工程软锚权重1。\n\n'
        '1296/1296验证行、216/216 cells保留；其中432行合成模型证据继承自父flow运行，不是独立重复。864行属于两种新实测目标。无新增拟合或评分。\n\n'
        +table(ag[ag['mode']=='full'].drop(columns='mode').replace({'method':labels}).rename(columns={'pipeline':'管道','method':'方法','expected':'计划','completed':'成功','all_three_below05':'三项<0.5'}))
        +'\nNRMSE按各管道18个训练窗口的固定SD归一化。不同管道改变了目标和尺度，不能视作同任务排名；成功条件误差不能抵消失败分母。联合重建是主终点，遮挡是独立检验。')
    for pipeline in pipelines:
        fig,axes=plt.subplots(1,3,figsize=(7.1,3.4))
        for ax,mode in zip(axes,modes):
            for j,m in enumerate(methods):
                row=ag[(ag.pipeline==pipeline)&(ag.method==m)&(ag['mode']==mode)].iloc[0]
                ax.bar(np.arange(3)+j*.32,[row[c] for c in COMPONENTS],width=.3,label=f'{labels[m]} {row.completed}/72')
            ax.set_xticks(np.arange(3)+.16,COMPONENTS);ax.axhline(.5,color='r',ls=':');ax.set_title(mode);ax.legend(fontsize=6)
        pages.append(f'## {pipeline}：联合重建与遮挡\n\n'+save(fig,'errors_'+pipeline)+'\n仅成功子集聚合，subject/session等权。遮挡列按原保留评分定义；完整终态和逐条误差见CSV。')
    paired=[]
    for (pipeline,mode),part in measured.groupby(['optical_pipeline','mode']):
        pair=part.pivot(index=['group','subject','session','trial'],columns='method',values=['status','converged']+['nmse_'+c for c in COMPONENTS])
        valid=np.ones(len(pair),bool)
        for m in methods:valid&=pair['status'][m].eq('completed')&truth(pair['converged'][m])
        p=pair[valid]
        for c in COMPONENTS:
            delta=p['nmse_'+c][methods[1]].astype(float)-p['nmse_'+c][methods[0]].astype(float)
            paired.append(dict(pipeline=pipeline,mode=mode,component=c,paired_success=len(p),planned=72,
                mean_delta_nmse=float(delta.mean()),median_delta_nmse=float(delta.median()),improved=int((delta<0).sum())))
    paired=pd.DataFrame(paired);paired.to_csv(out/'within_pipeline_paired.csv',index=False)
    pages.append('## 同管道固定τ→训练τ配对\n\n'+table(paired[paired['mode']=='full'].drop(columns='mode'))+'\n差值为训练−固定，只在双方都成功的同一身份上计算；负值表示误差降低。各管道单独比较，配对成功数不是完整72分母的替代。遮挡配对完整保留CSV。')
    starts=[];selections=[];scales=[];states=[]
    for group,prep in preps.items():
        spec=prep['spec'];folder=run/'training'/f'{group}__{methods[1]}'
        choice=json.loads((folder/'selection.json').read_text());local=[]
        for i,start in enumerate([1.,2.,4.]):
            r=json.loads((folder/f'start_{i}'/'result.json').read_text())
            if r.get('spec')!=spec or r.get('start_index')!=i or r.get('start_tau')!=start:raise ValueError('Training start identity mismatch')
            if spec['kind']=='synthetic':
                original=json.loads((parent/'training'/folder.name/f'start_{i}'/'result.json').read_text())
                if {k:v for k,v in r.items() if k!='inheritance'}!={k:v for k,v in original.items() if k!='inheritance'}:raise ValueError('Inherited training changed')
            row=dict(kind=spec['kind'],pipeline=spec.get('optical_pipeline','inherited'),group=group,subject=spec['subject'],fold=spec['outer'],start=i,
                start_tau=start,status=r['status'],converged=bool(r.get('converged',False)),tau=r.get('parameter_value'),objective=r.get('objective'),
                boundary=r.get('boundary_status'),gradient=r.get('projected_scaled_gradient_inf_norm'))
            starts.append(row);local.append(row)
        successful=[r['tau'] for r in local if r['status']=='completed' and r['converged']]
        selections.append(dict(kind=spec['kind'],pipeline=spec.get('optical_pipeline','inherited'),group=group,subject=spec['subject'],fold=spec['outer'],
            status=choice['status'],tau=choice.get('parameter_value'),selected_start=choice.get('selected_training_start'),successful_starts=len(successful),
            min_tau=min(successful,default=np.nan),max_tau=max(successful,default=np.nan)))
        if spec['kind']!='measured':continue
        if prep['status']=='completed':
            meta=prep['metadata'];opt=meta['optical_processing'];base=json.loads((parent/'prepared'/f"{spec['subject']}_o{spec['outer']}.json").read_text())['metadata']
            scales.append(dict(pipeline=spec['optical_pipeline'],subject=spec['subject'],fold=spec['outer'],MAD=opt['common_training_MAD'],factor=opt['fnirs_factor'],
                parent_MAD=base['fixed_roi']['common_training_MAD'],parent_factor=base['fixed_roi']['fnirs_factor'],gauge=opt['frozen_observation_loading'],
                **{c+'_SD':meta['normalization_sd'][i] for i,c in enumerate(COMPONENTS)},**{c+'_parent_SD':base['normalization_sd'][i] for i,c in enumerate(COMPONENTS)}))
        for m in methods:
            for mode in modes:
                path=run/'cells'/f'{group}__{m}__{mode}'/'trajectories.npz'
                rows=measured[(measured.group==group)&(measured.method==m)&(measured['mode']==mode)]
                arrays=None
                if path.exists():
                    with np.load(path,allow_pickle=False) as a:arrays={k:a[k] for k in ['states','trial_indices']}
                for r in rows.itertuples():
                    row=dict(pipeline=spec['optical_pipeline'],group=group,method=m,mode=mode,trial=r.trial,status=r.status,converged=bool(truth(pd.Series([r.converged])).iloc[0]),finite=False)
                    if arrays is not None:
                        idx=np.flatnonzero(arrays['trial_indices']==r.trial)
                        if len(idx)!=1:raise ValueError('State trial index mismatch')
                        state=arrays['states'][idx[0]];row['finite']=bool(state.shape==(120,6) and np.isfinite(state).all())
                        if row['finite']:
                            for j,name in enumerate(['r','s','f','v','p','q']):row[name+'_min']=float(state[:,j].min());row[name+'_max']=float(state[:,j].max())
                    states.append(row)
    starts=pd.DataFrame(starts);selections=pd.DataFrame(selections);scales=pd.DataFrame(scales);states=pd.DataFrame(states)
    if scales.empty:scales=pd.DataFrame(columns=['pipeline','factor','parent_factor','HbO_SD','HbO_parent_SD','HbR_SD','HbR_parent_SD'])
    for name,frame in [('training_starts',starts),('training_selection',selections),('optical_scale_changes',scales),('all_state_ranges',states)]:frame.to_csv(out/(name+'.csv'),index=False)
    selected=selections[selections.kind=='measured'];candidates=starts[starts.kind=='measured']
    fig,axes=plt.subplots(2,1,figsize=(7.1,5.2))
    for ax,pipeline in zip(axes,pipelines):
        part=candidates[candidates.pipeline==pipeline]
        for i in range(3):
            q=part[part.start==i];x=[cfg['subjects'].index(v)*4+f for v,f in zip(q.subject,q.fold)]
            ax.scatter(np.array(x)+(i-1)*.1,q.tau,label=f'起点{[1,2,4][i]}',s=15,marker=['o','x','+'][i])
        ax.axhline(.5,color='r',ls=':');ax.axhline(8.,color='r',ls=':');ax.set_xticks(range(12),[f'{s[-2:]}/{f}' for s in cfg['subjects'] for f in range(4)],fontsize=7)
        ax.set_ylabel('τ / 秒');ax.set_title(pipeline);ax.legend(fontsize=7)
    pages.append('## 训练τ多起点与被试/折\n\n'+save(fig,'tau_starts')+'\n点含失败末有效终点，不能当估计；状态详见training_starts.csv。虚线0.5–8秒是数值搜索边界，不是生理正常范围。没有τ先验，多起点差异不是置信区间。')
    pages.append('## 全部24折训练终态\n\n'+table(selected.drop(columns=['kind','group']).rename(columns={'pipeline':'管道','subject':'被试','fold':'折','status':'终态','successful_starts':'成功起点','selected_start':'选中起点'}))+'\n只有收敛成功且积分检查通过的起点参与选取。失败折仍产生完整验证失败行。')
    fig,axes=plt.subplots(2,4,figsize=(7.1,5.1))
    for row,pipeline in enumerate(pipelines):
        for col,name in enumerate(['f','v','p','q']):
            ax=axes[row,col]
            for j,m in enumerate(methods):
                part=states[(states.pipeline==pipeline)&(states.method==m)&(states['mode']=='full')&states.finite&states.converged&states.status.eq('completed')]
                if len(part):ax.scatter(np.full(len(part),j)-.08,part[name+'_min'],s=7);ax.scatter(np.full(len(part),j)+.08,part[name+'_max'],s=7)
            ax.set_xticks([0,1],['固定','训练']);ax.set_title(pipeline+' '+name,fontsize=8);ax.axhline(1,color='.7',ls=':')
    pages.append('## 状态范围与解释边界\n\n'+save(fig,'state_ranges')+'\n每点为成功full轨迹的最小或最大值；完整CSV还保留失败末有效状态与缺失标记。f/v/p/q为归一化模型状态，参照1不是正常范围。flow软锚提供未观测的绝对水平约束，不证明参数已识别；参数分离也可能来自尺度和优化补偿。')
    fig,axes=plt.subplots(1,3,figsize=(7.1,3.8))
    for pipeline in pipelines:
        part=scales[scales.pipeline==pipeline]
        for ax,key,parentkey in zip(axes,['factor','HbO_SD','HbR_SD'],['parent_factor','HbO_parent_SD','HbR_parent_SD']):
            ax.plot(range(len(part)),part[key]/part[parentkey],'o-',label=pipeline,ms=3);ax.axhline(1,color='.6',ls=':');ax.set_title(key+' / 旧目标');ax.set_xlabel('被试/折固定顺序')
    axes[0].legend(fontsize=7)
    pages.append('## 目标与归一化尺度变化\n\n'+save(fig,'scale_changes')+'\nMAD、factor、EEG/Hb分量SD及父值均保留于optical_scale_changes.csv。Hb共同MAD和分量SD只用18训练窗口；EEG与旧fold保持原样，loading gauge固定。上述变化是观测坐标变化，不能解释为生理浓度变化或跨管道胜负。')
    failure_audit_path=run/'training_failure_audit'/'no_motion_preliminary.json'
    failure_audit=json.loads(failure_audit_path.read_text()) if failure_audit_path.exists() else None
    failure_note=''
    if failure_audit is not None:
        failed=[r for r in failure_audit['rows'] if r.get('status')!='completed' or not r.get('converged',False)]
        margin=[r['minimum_absolute_HbO_model_margin'] for r in failed if r.get('minimum_absolute_HbO_model_margin') is not None]
        failure_note=('\n\n保留的[no_motion训练失败诊断](../training_failure_audit/no_motion_preliminary.json)只检查S09四折与S01折0参考，'
            f'其中{len(failed)}个失败起点的末接受状态最小HbO模型余额为{min(margin):.3g}至{max(margin):.3g}（若存在）。'
            '该量为p−0.35q，不是标定浓度；多条记录的最小余额位于全局trial21初始时刻，且梯度未达阈值、域拒绝较多。'
            '日志只保留末接受轨迹的时刻，未记录被拒提议的精确失败时刻，因此这是边界相关的数值诊断，不是已证明的完整因果链，也不推广到其他管道。')
    pages.append('## 所有实测失败与终态分母\n\n'+table(failure.rename(columns={'optical_pipeline':'管道','method':'方法','mode':'模式','status':'终态','rows':'行数'}).replace({'方法':labels}))+'\n所有状态计数纳入，不把未收敛、训练失败、准备失败或积分失败当作零误差。每管道×方法×模式计划均为72。'+failure_note)
    syn=data[data.kind=='synthetic'];synrows=[]
    for (m,mode),part in syn.groupby(['method','mode']):synrows.append(dict(method=labels[m],mode=mode,planned=len(part),successful=len(good(part)),**dict(zip(COMPONENTS,errors(good(part))))))
    pages.append('## 合成证据继承与上游控制\n\n'+table(pd.DataFrame(synrows))+'\n12组合成模型、432验证行及36训练起点逐项继承父flow证据，失败一并保留；这不是经过TDDR后的新SSM恢复实验。光学校正的零/常数/正负非对称周期/真实线性趋势/噪声控制来自独立audit。控制记录总数按原summary统计，不能将MNE或不校正当真值。上游当前处理重现216/216通过，因而可归因固定输入下的处理差异，不能推断真实人体趋势。')
    new_control_path=run/'optical_controls'/'results.json'
    new_control=json.loads(new_control_path.read_text());control_rows=new_control['metrics']
    if len(control_rows)!=168 or new_control['config'].get('methods')!=['no_motion','current','mne_tddr']:
        raise ValueError('Expected separate 168-record optical sensitivity controls')
    control_ids={(r['context_s'],r['case'],r['seed'],r['method'],r['region']) for r in control_rows}
    if len(control_ids)!=168 or {r['method'] for r in control_rows}!={'no_motion','current','mne_tddr'} or {r['region'] for r in control_rows}!={'full','common_roi'}:
        raise ValueError('Optical context control identities incomplete')
    flat=[]
    for r in control_rows:
        for wavelength in range(2):
            flat.append({k:(v[wavelength] if isinstance(v,list) else v) for k,v in dict(r,wavelength_index=wavelength).items()})
    controls=pd.DataFrame(flat);controls.to_csv(out/'optical_context_controls.csv',index=False)
    fig,axes=plt.subplots(1,2,figsize=(7.1,3.7))
    cases=['slow_clean','slow_noise','slow_noise_spike','slow_noise_step']
    for ax,context in zip(axes,[30,120]):
        part=controls[(controls.context_s==context)&(controls.region=='common_roi')&(controls.wavelength_index==0)]
        for j,m in enumerate(['no_motion','current','mne_tddr']):
            values=[part[(part['case']==case)&(part.method==m)].filtered_response_gain.mean() for case in cases]
            ax.plot(range(4),values,'o-',label=m,ms=3)
        ax.set_xticks(range(4),['clean','noise','spike','step'],rotation=25);ax.axhline(1,color='.6',ls=':');ax.set_title(f'{context}s context / common ROI');ax.set_ylabel('filtered response gain');ax.legend(fontsize=6)
    pages.append('## 新合成上下文压力检查：168条独立记录\n\n'+save(fig,'optical_context_controls')+'\n28个构造输入×3种处理×full/common ROI两区间=168记录，每记录含两波长；CSV展开为336行。本图是760nm、共同30秒ROI、噪声seed平均的已知慢响应gain，全部误差与运动残余指标保持原值。残余减小不等于真实响应保留。\n\n'
        '该压力输入为双gamma慢响应加白噪声，缺少生理低频背景；强衰减不能外推为所有实测数据都会丢失HRF，也不是复现原TDDR论文的10分钟重复任务评估。上下文长度同时影响运动估计和有限窗口滤波，不能独立归因其中一项。此168记录与旧audit六类输入、432继承SSM行是三个不同证据家族。')

    extension_path=run/'optical_controls'/'background_extension_v1'/'results.json'
    quiet_path=extension_path.parent/'quiet_paired_reference.json'
    extension=json.loads(extension_path.read_text());quiet=json.loads(quiet_path.read_text())
    if extension.get('status')!='completed' or extension.get('condition_count')!=216 or extension.get('metric_rows')!=432 or len(extension['rows'])!=432:
        raise ValueError('Expected completed 216-condition/432-row background extension')
    background=[]
    for context in [30,120]:
        values=[r['filtered_gain'] for r in quiet['rows'] if r['context_s']==context and r['method']=='mne_tddr']
        if len(values)!=3:raise ValueError('Quiet paired comparator needs exactly three seeds')
        median=np.median(values,axis=0)
        background.append(dict(context_s=context,background='white only',gain_760=median[0],gain_850=median[1],seeds=3))
        for family,level in [('heart_white','low'),('heart_white','high'),('heart_resp_mayer_white','low'),('heart_resp_mayer_white','high')]:
            values=[r['filtered_response']['gain'] for r in extension['rows'] if r['context_s']==context and r['method']=='mne_tddr'
                    and r['motion']=='none' and r['region']=='common_roi' and r['family']==family and r['level']==level]
            if len(values)!=3:raise ValueError('Background comparator needs exactly three seeds')
            median=np.median(values,axis=0)
            background.append(dict(context_s=context,background=family+'/'+level,gain_760=median[0],gain_850=median[1],seeds=3))
    background=pd.DataFrame(background);background.to_csv(out/'background_paired_HRF_gain.csv',index=False)
    pages.append('## 频率背景扩展：配对HRF保留率\n\n'+table(background)+'\n216个条件、432条full/common ROI记录为另一套纯合成扩展。此表是无运动条件、共同ROI、MNE两波长gain的3 seed中位数。所有行统一采用 F(HRF+背景)−F(背景)，与同样线性滤波的真实HRF比较；white only来自quiet_paired_reference.json，不能拿上一页168记录的raw-truth gain直接替代。\n\n'
        '加入心动及心动/呼吸/Mayer频率背景后，配对HRF增益明显提高，说明安静白噪声压力案例不能概括所有真实记录；保留仍不完美。low/high是预设工程幅度与HRF峰值的比值，不是人体标定或正常范围。该扩展不是原论文AR10/重复任务/GLM AUC复现，亦不构成实测MNE校准或方法排名。')

    # Independent optical controls are retained OD, not SSM truth/predictions.
    with np.load(audit/'synthetic_traces.npz',allow_pickle=False) as archive:
        control_arrays={k:archive[k] for k in archive.files}
    fig,axes=plt.subplots(1,2,figsize=(7.1,3.5))
    control_keys=list(control_arrays)
    for ax,case in zip(axes,['asymmetric_positive','true_linear_trend']):
        for method in ['no_motion','current','mne_tddr']:
            key=case+'__'+method+'__od'
            if key not in control_arrays:raise ValueError('Missing retained optical control OD')
            y=control_arrays[key][:,0];ax.plot(np.arange(len(y))/10,y-y[0],label=method,lw=1)
        ax.set_title(case);ax.set_xlabel('秒');ax.set_ylabel('OD−初值');ax.legend(fontsize=7)
    pages.append('## 光学负对照：漂移与真实趋势损失\n\n'+save(fig,'optical_controls')+'\n输入无运动伪迹时，当前方法仍可能把非对称周期积分成漂移；MNE能够消除真实线性趋势。两种反例必须一起看，均不能作为人体真值。完整控制记录保持在source optical audit，未重算。')
    atlas=fitz.open()
    for pipeline in pipelines:
        folder=out/pipeline;folder.mkdir()
        flow_full_curve_atlas(run,folder,cfg,measured[measured.optical_pipeline==pipeline],labels,methods,both_states=True,context_label=pipeline)
        with fitz.open(folder/'FULL_CURVE_ATLAS.pdf') as part:atlas.insert_pdf(part)
    if len(atlas)!=48 or sum(len(page.get_images()) for page in atlas)!=48:
        raise ValueError('Combined optical atlas must have 48 bitmap figure pages')
    atlas.save(out/'FULL_CURVE_ATLAS.pdf',garbage=4,deflate=True);atlas.close()
    pages.append('## 全部曲线与来源\n\n[FULL_CURVE_ATLAS.pdf](FULL_CURVE_ATLAS.pdf)共48页，每管道24页、每页3条：全部72窗口的EEG/HbO/HbR与两臂f/v/p/q。固定顺序不按效果选样，失败与末有效状态均标明。每管道另有atlas_index.csv。\n\n'
        f'运行：{run.name}。配置、manifest、source snapshot来源原样保留。父模型：{parent.name}；光学audit：{audit.name}。旧flow目标包含已定位的处理漂移，其旧NRMSE不可直接与新目标做效果排名。\n\n'
        '所有图为240dpi PNG，PDF正文/表格可搜索。WPS未检查。报告不改变缓存、评分、模型或运行状态。')
    text='\n\n---\n\n'.join(pages);(out/'REPORT.md').write_text(text)
    pdf=fitz.open();css='body{font-family:sans-serif;font-size:9pt;line-height:1.35}h1{font-size:19pt}h2{font-size:14pt}table{font-size:7pt;border-collapse:collapse;width:100%}td,th{border:.4pt solid #bbb;padding:3pt}img{width:510pt}'
    def rectfn(n,filled):
        page=fitz.paper_rect('a4');return page,page+(36,32,-36,-32),None
    for section in pages:
        html=markdown.markdown(section,extensions=['tables']).replace('<img ','<img width="510" ')
        part=fitz.Story(html,user_css=css,archive=fitz.Archive(str(out))).write_with_links(rectfn);pdf.insert_pdf(part);part.close()
    pdf.save(out/'REPORT.pdf',garbage=4,deflate=True)
    validation=dict(pages=len(pdf),figure_count=len(figures),image_objects=sum(len(p.get_images()) for p in pdf),searchable_text_chars=sum(len(p.get_text()) for p in pdf),
        bitmap_dpi=240,wps_checked=False,atlas_pages=48,source_report=str(Path(__file__).resolve()),source_report_sha256=hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
        training_failure_audit_source=str(failure_audit_path) if failure_audit is not None else None,
        training_failure_audit_sha256=hashlib.sha256(failure_audit_path.read_bytes()).hexdigest() if failure_audit is not None else None,
        background_extension_source=str(extension_path),background_extension_sha256=hashlib.sha256(extension_path.read_bytes()).hexdigest(),
        quiet_paired_source=str(quiet_path),quiet_paired_sha256=hashlib.sha256(quiet_path.read_bytes()).hexdigest(),
        optical_controls_source=str(new_control_path),optical_controls_sha256=hashlib.sha256(new_control_path.read_bytes()).hexdigest(),
        source_manifest=str(run/'manifest.json'),source_manifest_sha256=hashlib.sha256((run/'manifest.json').read_bytes()).hexdigest())
    if not 10<=len(pdf)<=16 or validation['image_objects']!=len(figures) or validation['searchable_text_chars']<1000:raise ValueError('Optical report pagination/bitmap validation failed')
    for i,page in enumerate(pdf):page.get_pixmap(matrix=fitz.Matrix(1.3,1.3)).save(out/f'preview_page_{i+1:02d}.png')
    pdf.close();(out/'export_validation.json').write_text(json.dumps(validation,indent=2));print(json.dumps(validation))


def render_conditional_optical_gain(run, out, *, pilot_truth_audit=None, terminal_input_snapshot=None):
    """Audit and render terminal retained arrays; never load native data or fit."""
    import hashlib
    import yaml
    if out.exists() and any(out.iterdir()):
        raise ValueError('Choose a new empty versioned export directory')
    input_root=Path(terminal_input_snapshot).resolve() if terminal_input_snapshot is not None else run
    manifest=json.loads((input_root/'manifest.json').read_text())
    cfg=yaml.safe_load((input_root/'resolved_config.yaml').read_text())
    methods=['fixed_tau_logflow1','conditional_fixed_gain','conditional_trained_gain']
    labels=dict(zip(methods,['A identity β1','B 条件映射 βref','C 条件映射 学β']))
    modes=['full','center_EEG','center_fNIRS'];pipelines=['no_motion','mne_tddr']
    full=manifest.get('execution')=='completed' and manifest.get('measured_terminal') is True
    if (manifest.get('pilot') is not False or manifest.get('synthetic_terminal') is not True
            or manifest.get('execution') not in ['synthetic_terminal','completed']
            or manifest.get('experiment_id')!='SSM-SHARED-DRIVER-CONDITIONAL-OPTICAL-GAIN-v1'
            or cfg.get('schema')!='shared_driver_conditional_optical_gain_v1'
            or cfg.get('methods')!=methods or cfg.get('modes')!=modes
            or cfg.get('optical_pipelines')!=pipelines
            or (manifest.get('execution')=='completed' and not full)):
        raise ValueError('Requires exact nonpilot terminal conditional gain evidence')
    if Path(manifest['source_root']).resolve()!=(run/'source_snapshot').resolve():
        raise ValueError('Missing owning source snapshot identity')
    root=Path(manifest['project_root']);parent=root/cfg['parent_run']
    groups=[f'synthetic_g{g:g}_{s}_r{r}' for g in [.5,1.,2.] for s in ['slow','mixed'] for r in range(2)]
    if full:groups += [f'{p}__{s}_o{f}' for p in pipelines for s in cfg['subjects'] for f in range(4)]
    data=pd.read_csv(input_root/'metrics.csv');summary=json.loads((input_root/'summary.json').read_text())
    keys=['kind','group','method','mode','trial'];expected=set();preps={};starts=[];choices=[];states=[];array_checks=[];costs=[]
    for group in groups:
        prep=json.loads((run/'prepared'/f'{group}.json').read_text());preps[group]=prep;spec=prep['spec']
        if spec['group']!=group or prep['status']!='completed':raise ValueError('Prepared identity/status mismatch')
        with np.load(run/'prepared'/f'{group}.npz',allow_pickle=False) as a:
            target=a['target'].copy();sd=a['normalizer'].copy()
        if 'normalization_sd' in prep['metadata'] and not np.array_equal(sd,np.asarray(prep['metadata']['normalization_sd'])):raise ValueError('Prepared SD differs from metadata')
        if spec['kind']=='measured':
            with np.load(parent/'prepared'/f'{group}.npz',allow_pickle=False) as a:
                if not np.array_equal(target,a['target']) or not np.array_equal(sd,a['normalizer']):raise ValueError('Parent target/SD changed')
        ref=prep['conditional_mapping']['beta_reference']
        for method in methods:
            for mode in modes:
                folder=run/'cells'/f'{group}__{method}__{mode}'
                cell=json.loads((folder/'result.json').read_text())
                if cell['cell']!=folder.name or len(cell['rows'])!=6:raise ValueError('Terminal cell identity/count mismatch')
                if spec['kind']=='measured' and method==methods[0]:
                    old=json.loads((parent/'cells'/folder.name/'result.json').read_text())
                    if cell['rows']!=old['rows']:raise ValueError('Inherited A rows changed')
                if mode=='full':
                    for result in cell.get('trial_results',[]):
                        if 'weighted_data_sse' not in result:continue
                        initial_cost=float(result['initial_penalty']*np.sum((np.asarray(result['initial_state'])-np.array([0.,1.,1.,1.,1.]))**2))
                        cost=dict(kind=spec['kind'],pipeline=spec.get('optical_pipeline','synthetic'),group=group,method=method,trial=result['trial'],status=result.get('scoring_status',result['status']),converged=result.get('converged',False),data_sse=result['data_sse'],weighted_data_sse=result['weighted_data_sse'],initial_state_cost=initial_cost,curvature_cost=result['penalty']*result['roughness'],flow_cost=result['flow_prior_cost'],amplitude_cost=result.get('driver_amplitude_cost',0.),objective=result['objective'])
                        total=sum(cost[k] for k in ['weighted_data_sse','initial_state_cost','curvature_cost','flow_cost','amplitude_cost'])
                        if not np.isclose(total,cost['objective'],rtol=1e-9,atol=1e-8):raise ValueError('Retained objective decomposition mismatch')
                        costs.append(cost)
                arrays=None
                if (folder/'trajectories.npz').exists():
                    with np.load(folder/'trajectories.npz',allow_pickle=False) as a:arrays={k:a[k].copy() for k in a.files}
                    idx=arrays['trial_indices']
                    if not np.array_equal(arrays['target'],target[idx]) or not np.array_equal(arrays['normalizer'],sd):raise ValueError('Cell target/SD differ')
                    array_checks.append(dict(group=group,method=method,mode=mode,target_exact=True,SD_exact=True))
                for row in cell['rows']:
                    trial=row['trial'];identity=(spec['kind'],group,method,mode,trial)
                    if any(row.get(k)!=v for k,v in dict(spec,method=method,mode=mode,session=cfg['sessions'][trial//8]).items()):raise ValueError('Cell row differs from prepared identity')
                    if trial not in [i for i in range(24) if i%4==spec['outer']] or identity in expected:raise ValueError('Split/duplicate mismatch')
                    expected.add(identity)
                    selected=data[(data.group==group)&(data.method==method)&(data['mode']==mode)&(data.trial==trial)]
                    if len(selected)!=1 or selected.iloc[0].status!=row['status'] or bool(truth(selected.converged).iloc[0])!=bool(row.get('converged',False)):raise ValueError('Metric status mismatch')
                    for c in COMPONENTS:
                        value=row.get('nmse_'+c);actual=selected.iloc[0].get('nmse_'+c)
                        if (value is None and pd.notna(actual)) or (value is not None and not np.isclose(value,actual,rtol=1e-12,atol=1e-14)):raise ValueError('Metric value mismatch')
                    if mode=='full':
                        record=dict(kind=spec['kind'],pipeline=spec.get('optical_pipeline','synthetic'),group=group,method=method,trial=trial,status=row['status'],converged=row.get('converged',False))
                        if arrays is not None:
                            where=np.flatnonzero(arrays['trial_indices']==trial)
                            if len(where)!=1:raise ValueError('State trial identity mismatch')
                            k=int(where[0]);state=arrays['states'][k]
                            if np.isfinite(state).all():
                                for j,name in enumerate(['r','s','f','v','p','q']):
                                    record[name+'_min']=float(state[:,j].min());record[name+'_max']=float(state[:,j].max())
                                record['initial_max_rest_deviation']=float(np.max(abs(state[0,2:]-1)))
                                if spec['kind']=='synthetic':
                                    for j,name in enumerate(['r','s','f','v','p','q']):record[name+'_truth_rmse']=float(np.sqrt(np.mean((state[:,j]-arrays['truth_states'][k,:,j])**2)))
                        states.append(record)
        folder=run/'training'/f'{group}__{methods[2]}'
        selection=json.loads((folder/'selection.json').read_text());local=[]
        for i,g in enumerate([.5,1.,2.]):
            result=json.loads((folder/f'start_{i}'/'result.json').read_text())
            if result['spec']!=spec or result['start_index']!=i or result['start_relative_gain']!=g:raise ValueError('Training identity mismatch')
            beta=result.get('parameter_value');relative=beta/ref if beta is not None else None
            item=dict(kind=spec['kind'],pipeline=spec.get('optical_pipeline','synthetic'),group=group,subject=spec['subject'],fold=spec['outer'],start=i,start_g=g,status=result['status'],converged=result.get('converged',False),beta=beta,g=relative,beta_reference=ref,boundary=result.get('boundary_status'),objective=result.get('objective'),gradient=result.get('projected_scaled_gradient_inf_norm'))
            starts.append(item);local.append(item)
        valid=[r['g'] for r in local if r['status']=='completed' and r['converged']]
        beta=selection.get('parameter_value');g=beta/ref if beta is not None else None
        choices.append(dict(kind=spec['kind'],pipeline=spec.get('optical_pipeline','synthetic'),group=group,subject=spec['subject'],fold=spec['outer'],status=selection['status'],beta=beta,g=g,beta_reference=ref,true_g=spec.get('true_relative_gain'),successful_starts=len(valid),start_span=max(valid)-min(valid) if valid else None,selected_start=selection.get('selected_training_start')))
    expected_rows=1944 if full else 648
    if len(data)!=expected_rows or data.duplicated(keys).any() or set(data[keys].itertuples(index=False,name=None))!=expected or summary['observed_rows']!=expected_rows or summary['terminal_cells']!=expected_rows//6:raise ValueError('Incomplete terminal denominator')
    out.mkdir(parents=True);(out/'figures').mkdir()
    font=Path('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc')
    if font.exists():font_manager.fontManager.addfont(str(font));plt.rcParams['font.family']=font_manager.FontProperties(fname=str(font)).get_name()
    plt.rcParams.update({'font.size':9,'axes.unicode_minus':False});figures=[];pages=[]
    def save(fig,name):
        fig.tight_layout();fig.savefig(out/'figures'/f'{name}.png',dpi=240);plt.close(fig);figures.append(name)
        return f'![{name}](figures/{name}.png)'
    def good(part):return part[part.status.eq('completed')&truth(part.converged)]
    data['pipeline']=data.optical_pipeline.fillna('synthetic') if 'optical_pipeline' in data else 'synthetic'
    ag=[];paired=[]
    for (kind,pipeline,method,mode),part in data.groupby(['kind','pipeline','method','mode'],sort=False):
        ok=good(part);cols=['nmse_'+c for c in COMPONENTS]
        error=np.sqrt(ok.groupby(['subject','session'])[cols].mean().groupby('subject').mean().mean()) if len(ok) else [np.nan]*3
        ag.append(dict(kind=kind,pipeline=pipeline,method=method,mode=mode,planned=len(part),success=len(ok),all3_below05=int(ok[['nrmse_'+c for c in COMPONENTS]].lt(.5).all(axis=1).sum()) if mode=='full' else None,**dict(zip(COMPONENTS,error))))
    for (kind,pipeline,mode),part in data.groupby(['kind','pipeline','mode']):
        for left,right in [(methods[0],methods[1]),(methods[1],methods[2]),(methods[0],methods[2])]:
            joined=good(part[part.method==left]).merge(good(part[part.method==right]),on=['group','trial','subject','session'],suffixes=('_left','_right'))
            for c in COMPONENTS:
                delta=joined['nmse_'+c+'_right']-joined['nmse_'+c+'_left']
                paired.append(dict(kind=kind,pipeline=pipeline,mode=mode,contrast=left+' → '+right,component=c,planned=len(part)//3,common_success=len(joined),mean_delta_nmse=delta.mean(),median_delta_nmse=delta.median(),improved=int((delta<0).sum())))
    ag=pd.DataFrame(ag);paired=pd.DataFrame(paired);starts=pd.DataFrame(starts);choices=pd.DataFrame(choices);states=pd.DataFrame(states)
    for name in ['r','s','f','v','p','q']:
        for suffix in ['_min','_max','_truth_rmse']:
            if name+suffix not in states:states[name+suffix]=np.nan
    if 'initial_max_rest_deviation' not in states:states['initial_max_rest_deviation']=np.nan
    if 'driver_nrmse' not in data:data['driver_nrmse']=np.nan
    status=data.groupby(['kind','pipeline','method','mode','status'],dropna=False).size().reset_index(name='rows')
    for name,frame in [('all_validation_rows',data),('metrics_summary',ag),('paired_metrics',paired),('training_starts',starts),('training_selections',choices),('full_state_ranges_and_recovery',states),('terminal_status_counts',status),('target_SD_audit',pd.DataFrame(array_checks))]:frame.to_csv(out/(name+'.csv'),index=False)
    pages.append('# 条件性光学观测映射与共享血管增益\n\n'+f'阶段：{"完整合成与实测" if full else "仅合成终态"}；保留 {expected_rows}/{expected_rows} 验证行和全部失败。联合 EEG/HbO/HbR 重建为主终点，遮挡预测为独立检验。EEG 是1–45 Hz log-power/PCA特征，不是原始电压。\n\n'
        +('本报告只覆盖合成终态；计划实测阶段未纳入，完整实验未完成。\n\n' if not full else '')+'A：原identity映射、β=1。B：条件2×2映射B、βref=1/√det(B)。C：相同映射、18训练trial共享估计β，验证冻结。τ全部固定2秒。A→B同时改变映射、参考β和正则项相对作用，不能称纯映射效应；B→C隔离训练增益的贡献。\n\n'
        '71 µM总Hb基线、DPF=6、3 cm路径及组织敏感度1均为文献条件假设，无个体独立标定。β与g=β/βref是当前EEG/光学坐标下的有效耦合，不是已校准的个体神经血管耦合。g的[0.1,10]为数值搜索界，非正常生理区间。[Friston原始模型](https://www.fil.ion.ucl.ac.uk/spm/doc/papers/karl_nonlinear.pdf)的效能同时包含输入诱发神经强度和血管耦合，不能等同独立NVC。')
    for kind,pipeline in ag[['kind','pipeline']].drop_duplicates().itertuples(index=False,name=None):
        part=ag[(ag.kind==kind)&(ag.pipeline==pipeline)]
        fig,axes=plt.subplots(1,3,figsize=(7.1,3.4))
        for ax,mode in zip(axes,modes):
            for j,m in enumerate(methods):
                row=part[(part.method==m)&(part['mode']==mode)].iloc[0]
                ax.bar(np.arange(3)+j*.25,[row[c] for c in COMPONENTS],width=.23,label=f'{"ABC"[j]} {row.success}/{row.planned}')
            ax.set_xticks(np.arange(3)+.25,COMPONENTS);ax.axhline(.5,c='r',ls=':');ax.set_title(mode);ax.legend(fontsize=6)
        pages.append(f'## {kind} / {pipeline} 重建与独立遮挡\n\n'+save(fig,kind+'_'+pipeline+'_errors')+'\n\n'+table(part[part['mode']=='full'][['method','planned','success','all3_below05']].replace({'method':labels}))+'\n误差为成功条件子集subject/session等权NMSE开根；失败不填零。全分母、遮挡及所有身份保留CSV。每管道目标和训练SD跨三臂完全一致；不同管道不作同任务排名。')
    pages.append('## 同身份配对与失败\n\n'+table(paired[(paired['mode']=='full')&(paired.contrast==methods[1]+' → '+methods[2])].drop(columns=['mode','contrast']))+'\n差值=C−B；只在双方成功的同一身份上计算，不能替代完整分母。A→B/A→C和遮挡配对均保留paired_metrics.csv。\n\n'+table(status.groupby(['kind','pipeline','method','status'],dropna=False)['rows'].sum().reset_index())+'\n失败轨迹不计入有效恢复；图册不把末有效迭代伪装成成功。')
    for kind,pipeline in choices[['kind','pipeline']].drop_duplicates().itertuples(index=False,name=None):
        selected=choices[(choices.kind==kind)&(choices.pipeline==pipeline)];names=selected.group.tolist()
        fig,ax=plt.subplots(figsize=(7.1,3.7))
        for start in range(3):
            part=starts[(starts.kind==kind)&(starts.pipeline==pipeline)&(starts.start==start)]
            ok=part.status.eq('completed')&truth(part.converged)
            ax.scatter([names.index(g)+(start-1)*.13 for g in part[ok].group],part[ok].g,s=19,label=f'起点 g={ [.5,1.,2.][start]}')
            bad=part[~ok];ax.scatter([names.index(g)+(start-1)*.13 for g in bad.group],[.12]*len(bad),c='red',marker='x',s=25)
        if kind=='synthetic':ax.scatter(range(len(selected)),selected.true_g,marker='_',c='k',s=100,label='真实g')
        ax.set_yscale('log');ax.set_ylim(.08,12);ax.axhline(.1,c='r',ls=':');ax.axhline(10,c='r',ls=':');ax.set_xticks(range(len(names)),[n.replace('synthetic_','').replace('subject_','').replace(pipeline+'__','') for n in names],rotation=60,fontsize=6);ax.set_ylabel('g = β / βref');ax.legend(fontsize=7)
        pages.append(f'## {kind} / {pipeline} 增益多起点\n\n'+save(fig,kind+'_'+pipeline+'_gain')+'\n红叉只表示失败，纵坐标0.12不是该失败的估计。实际β、βref、g、边界、梯度和起点跨度均保留CSV。跨折训练重叠，折间差异不是独立被试证据。\n\n'+table(selected[['subject','fold','status','beta','g','true_g','successful_starts','start_span']]))
    if costs:
        costs=pd.DataFrame(costs);costs.to_csv(out/'full_objective_components.csv',index=False)
        ok_costs=costs[costs.status.eq('completed')&truth(costs.converged)]
        cost_summary=ok_costs.groupby(['kind','pipeline','method'])[['weighted_data_sse','initial_state_cost','curvature_cost','flow_cost']].mean().reset_index()
        cost_summary.to_csv(out/'objective_component_means.csv',index=False)
        pages.append('## 数据残差与工程惩罚的相对作用\n\n'+table(cost_summary)+'\n表为成功full验证trial的算术均值，完整状态、原始SSE和所有分项保留CSV。初态项按保留物理初态及权重精确复算，并核对分项和等于原objective；未重新拟合。A→B缩小物理状态时初态/flow惩罚相对作用也会减弱，因此重建改善可能包含解除工程惩罚的效果。总objective不用于跨臂判胜；只有同臂多起点的原优化选择可比较。')
    screen=summary.get('synthetic_engineering_screen')
    if screen is not None:
        pages.append('## 合成工程筛查与阶段边界\n\n'+table(pd.DataFrame(screen['groups'])[['group','true_relative_gain','selected_relative_gain','successful_starts','selected_relative_error','successful_span_over_true','successful_starts_interior','passed']].rename(columns={'true_relative_gain':'true_g','selected_relative_gain':'selected_g','successful_starts':'成功起点','selected_relative_error':'相对误差','successful_span_over_true':'跨度/真值','successful_starts_interior':'界内','passed':'通过'}))+'\n\n'+('完整实测终态已保留。' if full else '当前仅合成终态；计划中的实测阶段未包含于本报告，不能判定完整实验完成。')+' 成功起点数见训练表；只有一个成功起点时跨度0不构成稳定性证据。筛查界是工程诊断，不是独立生理资格。')
    component_rows=[]
    for group,prep in preps.items():
        if prep['spec']['kind']!='synthetic':continue
        diagnostic=prep['metadata'].get('truth_component_diagnostic')
        if diagnostic is not None:
            component_rows.append(dict(group=group,**{k:json.dumps(v,ensure_ascii=False) for k,v in diagnostic.items() if k!='interpretation'}))
    if component_rows:
        pd.DataFrame(component_rows).to_csv(out/'synthetic_truth_components.csv',index=False)
        pages.append('## 合成初态与驱动的反事实检查\n\n'+table(pd.DataFrame(component_rows).map(lambda v: '['+', '.join(f'{x:.4f}' for x in json.loads(v))+']' if isinstance(v,str) and v.startswith('[') else v))+'\n初态独立输入和驱动独立输入分别重新经过非线性前向模型、条件矩阵和处理算子。它们不是可相加的方差分解，不能解释成解释率。完整反事实数组保留于prepared，未据此改变目标或评分。')
    pilot_provenance=None
    if pilot_truth_audit is not None:
        import shutil
        pilot_truth_audit=Path(pilot_truth_audit).resolve()
        pilot=json.loads(pilot_truth_audit.read_text())
        if pilot.get('schema')!='conditional_gain_pilot_truth_audit_v1':raise ValueError('Unexpected pilot truth audit schema')
        rows=pilot['rows'];crows=[r for r in rows if r['method']=='conditional_trained_gain']
        if len(rows)!=6 or {r['trial'] for r in crows}!={0,4}:raise ValueError('Pilot audit identities incomplete')
        background=out/'pilot_background';background.mkdir()
        shutil.copy2(pilot_truth_audit,background/pilot_truth_audit.name)
        original_png=pilot_truth_audit.parent/'driver_dc_and_flow.png'
        if original_png.exists():shutil.copy2(original_png,background/original_png.name)
        pilot_provenance=dict(source=str(pilot_truth_audit),sha256=hashlib.sha256(pilot_truth_audit.read_bytes()).hexdigest(),source_png=str(original_png) if original_png.exists() else None,source_png_sha256=hashlib.sha256(original_png.read_bytes()).hexdigest() if original_png.exists() else None)
        fig,axes=plt.subplots(1,3,figsize=(7.1,3.4));x=np.arange(2)
        axes[0].bar(x,[r['raw_driver_nrmse'] for r in crows]);axes[0].set_title('C 原坐标驱动NRMSE');axes[0].axhline(.5,color='red',ls=':')
        axes[1].bar(x,[100*r['dc_fraction_of_squared_error'] for r in crows]);axes[1].set_title('DC占驱动误差SSE / %');axes[1].set_ylim(0,105)
        axes[2].bar(x-.16,[r['flow_mean_fit']-r['flow_mean_truth'] for r in crows],width=.3,label='实际均值偏差');axes[2].bar(x+.16,[r['linear_rest_equilibrium_shift_from_driver_dc'][1] for r in crows],width=.3,label='βc/γ近似');axes[2].set_title('f均值偏差');axes[2].legend(fontsize=6)
        for ax in axes:ax.set_xticks(x,[f'trial {r["trial"]}' for r in crows])
        pages.append('## Pilot背景：驱动DC与初态补偿\n\n'+save(fig,'pilot_driver_dc_background')+'\n仅为先前两个pilot验证trial的背景证据，不计入本次完整合成分母。C的原坐标驱动NRMSE分别为 '+ '/'.join(f'{r["raw_driver_nrmse"]:.3f}' for r in crows)+'；其中 '+ '/'.join(f'{100*r["dc_fraction_of_squared_error"]:.2f}%' for r in crows)+' 的驱动误差SSE来自DC。去均值误差只能诊断，不能替换注册的raw指标。小信号附近，驱动DC偏差c与初态/静态flow偏移βc/γ相互补偿；完整非线性血管系统不具有一般加性平移不变性。f接近1或曲线小幅并不意味着生理状态已验证。\n原审计JSON和原PNG保留于pilot_background；本图以保留标量重新作图240dpi，未重新拟合。')
    synthetic=good(data[(data.kind=='synthetic')&(data['mode']=='full')])
    recovery=states[(states.kind=='synthetic')&states.status.eq('completed')&truth(states.converged)].groupby('method')[[n+'_truth_rmse' for n in ['r','s','f','v','p','q']]].mean().reset_index()
    driver=synthetic.groupby('method').driver_nrmse.mean().reset_index()
    pages.append('## 合成参数、驱动和状态恢复\n\n'+table(driver)+'\n'+table(recovery)+'\n驱动NRMSE以真实驱动SD归一化；状态为各自原始坐标RMSE，不能跨状态比较幅度。新合成真值已在原生Hb加噪之前施加条件矩阵，12组跨g配对驱动、初态和标准化噪声；不是人体标定，也不是从TDDR后信号恢复真实Hb的证明。')
    mechanism_path=run/'training_mechanism_audit'/'review.json'
    if full and mechanism_path.exists():
        mechanism=json.loads(mechanism_path.read_text())
        if mechanism['total_starts']!=72 or mechanism['successful_starts']!=int(((starts.kind=='measured')&starts.status.eq('completed')&truth(starts.converged)).sum()):raise ValueError('Training mechanism audit differs from retained training starts')
        import shutil
        shutil.copy2(mechanism_path,out/'training_mechanism_review.json')
        cross=pd.DataFrame(mechanism['crossfold']);cross['beta_max_over_min']=cross.beta_max/cross.beta_min
        selected_measured=choices[choices.kind=='measured'];selected_count=int(selected_measured.status.eq('completed').sum())
        pages.append('## 实测训练：起点稳定与跨折可识别性\n\n'+f'{mechanism["successful_starts"]}/72起点成功，{selected_count}/24组有成功选择；其余组完整保留训练失败并传递到验证分母。成功起点在组内最大β相对跨度为 {100*mechanism["maximum_within_group_relative_beta_span"]:.4f}%。这不构成跨折参数稳定性证据，只有一个成功起点的组也不能用跨度0宣称稳定。\n\n'+table(cross)+'\nS09跨折β可变化约5–6倍，即使各折内多起点近乎一致。因此当前有效耦合尚不能解释为稳定的被试生理特征。共同训练trial的惩罚、驱动DC和初态补偿审计保留于training_mechanism_review.json；训练集合重叠，跨折不是独立重复。九个失败起点仍是数值失败，不能用末有效参数补成成功。')
    example_specs=[(g,f'synthetic_g{g:g}_slow_r0',0) for g in [.5,1.,2.]]
    fig_observation,obs_axes=plt.subplots(3,3,figsize=(7.1,6.6),sharex=True)
    fig_state,state_axes=plt.subplots(3,3,figsize=(7.1,6.6),sharex=True)
    example_status=[];time=np.arange(cfg['tensor']['steps'])*cfg['tensor']['dt_s']-5.
    for column,(true_g,group,trial) in enumerate(example_specs):
        with np.load(run/'prepared'/f'{group}.npz',allow_pickle=False) as a:
            observed=a['target'][trial].copy();true_driver=a['truth'][trial].copy();true_states=a['truth_states'][trial].copy()
        for j,c in enumerate(COMPONENTS):obs_axes[j,column].plot(time,observed[:,j],color='black',lw=1.,label='观测')
        for j,values in enumerate([true_driver,true_states[:,2],true_states[:,4]-true_states[:,3]]):state_axes[j,column].plot(time,values,color='black',lw=1.,label='真实')
        for method,color,style in zip(methods,['.5','tab:blue','tab:orange'],['--','-',':']):
            record=data[(data.group==group)&(data.method==method)&(data['mode']=='full')&(data.trial==trial)].iloc[0]
            example_status.append(dict(group=group,trial=trial,method=method,status=record.status,converged=bool(truth(pd.Series([record.converged])).iloc[0])))
            if record.status!='completed' or not bool(truth(pd.Series([record.converged])).iloc[0]):continue
            with np.load(run/'cells'/f'{group}__{method}__full'/'trajectories.npz',allow_pickle=False) as a:
                idx=np.flatnonzero(a['trial_indices']==trial)
                if len(idx)!=1:raise ValueError('Fixed synthetic example identity missing')
                k=int(idx[0]);prediction=a['prediction'][k].copy();state=a['states'][k].copy()
            if not np.isfinite(prediction).all() or not np.isfinite(state).all():raise ValueError('Successful example has nonfinite arrays')
            for j,c in enumerate(COMPONENTS):obs_axes[j,column].plot(time,prediction[:,j],color=color,ls=style,lw=.9,label='ABC'[methods.index(method)])
            for j,values in enumerate([state[:,0],state[:,2],state[:,4]-state[:,3]]):state_axes[j,column].plot(time,values,color=color,ls=style,lw=.9,label='ABC'[methods.index(method)])
        obs_axes[0,column].set_title(f'真实g={true_g:g} / slow r0 trial0',fontsize=8)
        state_axes[0,column].set_title(f'真实g={true_g:g} / slow r0 trial0',fontsize=8)
        for j in range(3):
            if column==0:
                obs_axes[j,column].set_ylabel(COMPONENTS[j],fontsize=8)
                state_axes[j,column].set_ylabel(['原坐标驱动','f','p − v'][j],fontsize=8)
            for ax in [obs_axes[j,column],state_axes[j,column]]:
                ax.tick_params(labelsize=7);ax.axvline(0,color='.8',ls=':',lw=.6)
                if j==2:ax.set_xlabel('相对事件 / 秒',fontsize=8)
        obs_axes[0,column].legend(fontsize=6,ncol=2);state_axes[0,column].legend(fontsize=6,ncol=2)
    pd.DataFrame(example_status).to_csv(out/'fixed_synthetic_example_status.csv',index=False)
    pages.append('## 固定合成示例：共享状态联合重建\n\n'+save(fig_observation,'synthetic_fixed_observation_curves')+'\n固定身份为g=0.5/1/2、slow、rep0、trial0，不按误差选样。黑色为实际加噪观测；灰虚A、蓝实B、橙点C为保留full预测。列间使用各自原始观测尺度；未去均值、翻转符号、平移时间或调整幅度。失败身份保留于下一页状态表，不画成成功。')
    pages.append('## 固定合成示例：原驱动及生理状态\n\n'+save(fig_state,'synthetic_fixed_driver_state_curves')+'\n同三条固定身份。黑色是真值；三臂保留原驱动、f及p−v。未进行任何对齐或去均值；原驱动DC差异必须计入恢复误差。各子图独立纵轴以保留可读性，不意味着状态幅度相同。f接近静息1、p−v接近0仅为图像观察，不是生理资格。\n\n'+table(pd.DataFrame(example_status)))
    measured=states[(states.kind=='measured')&states.status.eq('completed')&truth(states.converged)]
    if full:
        rows=[]
        for (pipeline,method),part in measured.groupby(['pipeline','method']):
            rows.append(dict(pipeline=pipeline,method=method,success=len(part),**{name+'_min':part[name+'_min'].min() for name in ['f','v','p','q']},**{name+'_max':part[name+'_max'].max() for name in ['f','v','p','q']},initial_deviation_over05=int((part.initial_max_rest_deviation>.5).sum()),f_min_percent_from_rest=100*(part.f_min.min()-1),f_max_percent_from_rest=100*(part.f_max.max()-1),trial_f_min_q05=float(part.f_min.quantile(.05)),trial_f_max_q95=float(part.f_max.quantile(.95))))
        pd.DataFrame(rows).to_csv(out/'measured_state_tail_summary.csv',index=False)
        pages.append('## 实测状态与完整曲线\n\n'+table(pd.DataFrame(rows)[['pipeline','method','success','f_min_percent_from_rest','f_max_percent_from_rest','trial_f_min_q05','trial_f_max_q95','initial_deviation_over05']])+'\n初态偏差>0.5仅为描述性极端计数，不是已验证正常界。数学正值、Hb balance和积分复核是数学域条件，不是经验生理范围。若f解释为静息归一局部血流，f=0.1即−90%、f=1.9即+90%；这些是相对幅度，前者可描述为极端且未生理验证，不能据此临床诊断，也不能只凭单篇研究均值判后者不可能。新的映射和β缩小使f接近1亦不自动证明生理合理。跨trial极值的5%/95%分位数表示尾部分布，不是正常界。\n\n完整72身份×两管道在FULL_CURVE_ATLAS.pdf固定顺序展示，不按误差挑样；后三列三臂独立纵轴，便于辨认小幅和极端状态。')
        atlas=fitz.open()
        for pipeline in pipelines:
            folder=out/pipeline;folder.mkdir()
            flow_full_curve_atlas(run,folder,cfg,data[data.pipeline==pipeline],labels,methods,both_states=True,context_label=pipeline)
            with fitz.open(folder/'FULL_CURVE_ATLAS.pdf') as part:atlas.insert_pdf(part)
        if len(atlas)!=48 or sum(len(p.get_images()) for p in atlas)!=48:raise ValueError('Incomplete bitmap atlas')
        atlas.save(out/'FULL_CURVE_ATLAS.pdf',garbage=4,deflate=True);atlas.close()
    pages.append('## 来源与限制\n\n'+f'运行 {run.name}；只读取保留prepared/cells/training与父运行 {parent.name}，未拟合或重评分。目标/SD逐单元核验见target_SD_audit.csv。实测A保留父运行失败，属于继承对照；合成三臂全部新运行。\n\n所有完整图为240dpi PNG；PDF文本可搜。WPS未检查。需要结合完整图册与失败分母判断恢复，不以成功条件均值宣称全体达标。')
    export_pdf('\n\n'.join(pages),out,figures,split_images=False)
    audit=json.loads((out/'export_validation.json').read_text())
    audit.update(pilot_truth_audit=pilot_provenance,expected_rows=expected_rows,terminal_cells=expected_rows//6,target_SD_checks=len(array_checks),full_measured=full,source_report=str(Path(__file__).resolve()),source_report_sha256=hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),source_manifest_sha256=hashlib.sha256((input_root/'manifest.json').read_bytes()).hexdigest(),terminal_input_snapshot=str(input_root))
    with fitz.open(out/'REPORT.pdf') as pdf:
        if any(p.get_drawings() for p in pdf):
            # Table borders are vector body content; only figure rasterization is required.
            audit['body_vector_table_borders_allowed']=True
        for i,page in enumerate(pdf):page.get_pixmap(matrix=fitz.Matrix(1.3,1.3)).save(out/f'preview_page_{i+1:02d}.png')
    (out/'export_validation.json').write_text(json.dumps(audit,indent=2))


def step_control_audit(run):
    """Read every paired start, including failures, without fitting a model."""
    import yaml
    cfg=yaml.safe_load((run/'resolved_config.yaml').read_text())
    if cfg.get('schema')!='shared_driver_conditional_step_control_v2':
        raise ValueError('paired solver report requires the v2 contract')
    rows=[]
    for kind in ('synthetic','measured'):
        summary=json.loads((run/f'{kind}_paired_summary.json').read_text())
        expected=cfg[kind+'_starts']
        if not summary['complete'] or summary['terminal_pairs']!=expected:
            raise ValueError('paired solver report requires complete denominators')
        if len({(p['group'],p['start']) for p in summary['rows']})!=expected:
            raise ValueError('duplicate solver pair')
        for pair in summary['rows']:
            records={};arrays={}
            for arm in ('armijo','quadratic_interpolation'):
                folder=run/'arms'/arm/'training'/f"{pair['group']}__conditional_trained_gain"/f"start_{pair['start']}"
                records[arm]=json.loads((folder/'result.json').read_text())
                if (folder/'trajectories.npz').exists():
                    with np.load(folder/'trajectories.npz',allow_pickle=False) as a:arrays[arm]={k:a[k].copy() for k in a.files}
            old,new=records['armijo'],records['quadratic_interpolation']
            success=lambda r:r['status']=='completed' and r.get('converged',False)
            row=dict(kind=kind,group=pair['group'],start=pair['start'],
                subject=new['spec'].get('subject','synthetic'),pipeline=new['spec'].get('optical_pipeline','synthetic'),
                old_success=success(old),new_success=success(new),
                recovered=success(new) and not success(old),regressed=success(old) and not success(new))
            for prefix,record in [('old',old),('new',new)]:
                for key in ('status','objective','relative_gain','optimization_evaluations','seconds','cpu_seconds',
                            'projected_scaled_gradient_inf_norm','convergence_reason'):
                    row[prefix+'_'+key]=record.get(key)
                row[prefix+'_iterations']=record.get('starts',[{}])[0].get('iterations')
                row[prefix+'_integration_gap']=max((c['difference_training_sd'] for c in record.get('integration_checks',[])),default=np.nan)
                checked=np.asarray([record.get('optimization_evaluations',np.nan),
                    record.get('projected_scaled_gradient_inf_norm',np.nan),row[prefix+'_integration_gap']],dtype=float)
                if success(record) and (not np.isfinite(checked).all() or checked[0]>3600 or
                    checked[1]>1e-4 or checked[2]>.005):
                    raise ValueError('successful start violates frozen budget/gradient/integration check')
            counts=new.get('starts',[{}])[0].get('interpolation_counts',{})
            row.update(interpolation_attempts=counts.get('attempts',0),interpolation_selected=counts.get('selected',0))
            inherited=run/'arms/armijo/baseline_training'/f"{pair['group']}__conditional_trained_gain"/f"start_{pair['start']}/result.json"
            parent=json.loads(inherited.read_text())
            row['legacy_reproduced']=all(old.get(k)==parent.get(k) for k in
                ('status','converged','objective','parameter_value','optimization_evaluations'))
            if len(arrays)==2:
                a,b=arrays['armijo'],arrays['quadratic_interpolation']
                if any(not np.array_equal(a[k],b[k]) for k in ('target','normalizer','trial_indices')):
                    raise ValueError('paired solver target/SD/trial mismatch')
                row['prediction_max_difference_sd']=float(np.max(abs(a['prediction']-b['prediction'])/a['normalizer']))
                row['driver_max_difference']=float(np.max(abs(a['driver']-b['driver'])))
                row['flow_max_difference']=float(np.max(abs(a['states'][:,:,2]-b['states'][:,:,2])))
            if old.get('objective') is not None and new.get('objective') is not None:
                row['objective_relative_difference']=abs(new['objective']-old['objective'])/max(abs(old['objective']),1e-12)
                row['gain_relative_difference']=abs(new['relative_gain']/old['relative_gain']-1)
            rows.append(row)
    frame=pd.DataFrame(rows);aggregate=[]
    for (kind,subject,pipeline),part in frame.groupby(['kind','subject','pipeline'],sort=False):
        row=dict(kind=kind,subject=subject,pipeline=pipeline,starts=len(part),
            old_success=int(part.old_success.sum()),new_success=int(part.new_success.sum()),
            recovered=int(part.recovered.sum()),regressions=int(part.regressed.sum()))
        for metric in ('optimization_evaluations','seconds','cpu_seconds'):
            for prefix in ('old','new'):row[prefix+'_'+metric]=float(part[prefix+'_'+metric].sum())
            row[metric+'_reduction']=1-row['new_'+metric]/row['old_'+metric]
        aggregate.append(row)
    return frame,pd.DataFrame(aggregate)


def render_step_control(run,out):
    if out.exists() and any(out.iterdir()):raise ValueError('choose a new versioned report output')
    frame,groups=step_control_audit(run)
    out.mkdir(parents=True,exist_ok=True);(out/'figures').mkdir()
    frame.to_csv(out/'paired_starts.csv',index=False);groups.to_csv(out/'subject_pipeline_summary.csv',index=False)
    font=Path('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc')
    if font.exists():
        font_manager.fontManager.addfont(str(font));plt.rcParams['font.family']=font_manager.FontProperties(fname=str(font)).get_name()
    plt.rcParams.update({'font.size':9,'axes.unicode_minus':False});figures=[]
    def save(fig,name):
        fig.tight_layout();fig.savefig(out/'figures'/f'{name}.png',dpi=240);plt.close(fig);figures.append(name)
        return f'![{name}](figures/{name}.png)'
    syn=frame[frame.kind=='synthetic'];real=frame[frame.kind=='measured'];rows=[]
    for kind,part in [('合成',syn),('实测',real)]:
        rows.append(dict(数据=kind,起点对=len(part),旧成功=int(part.old_success.sum()),新成功=int(part.new_success.sum()),
            修复=int(part.recovered.sum()),退步=int(part.regressed.sum()),
            优化调用降幅=f'{100*(1-part.new_optimization_evaluations.sum()/part.old_optimization_evaluations.sum()):.2f}%',
            CPU降幅=f'{100*(1-part.new_cpu_seconds.sum()/part.old_cpu_seconds.sum()):.2f}%'))
    pages=['# 同目标求解器：收敛与计算量\n\n'+table(pd.DataFrame(rows))+
        '\n同一固定目标、起点、3600次试次评估、90次迭代、梯度阈值1e−4。只增加一次有条件的二次插值；不改模型、先验或积分器。\n\n'+
        '两臂均重新运行，CPU时间含最终重放、信息矩阵及积分复核。墙钟受共享机器负载影响；调用量是确定性计算指标。分母含失败起点。']
    part=groups[groups.kind=='measured'];x=np.arange(len(part));fig,axes=plt.subplots(1,2,figsize=(11,4.5))
    labels=[r.subject.replace('subject_','S')+'\n'+r.pipeline for r in part.itertuples()]
    for prefix,offset,color,label in [('old',-.18,'#2877ab','原Armijo'),('new',.18,'#dc8030','二次插值')]:
        bars=axes[0].bar(x+offset,part[prefix+'_success'],.36,color=color,label=label)
        axes[0].bar_label(bars,padding=2,fontsize=8)
        axes[1].bar(x+offset,part[prefix+'_optimization_evaluations'],.36,color=color,label=label)
    for ax in axes:ax.set_xticks(x,labels,fontsize=7);ax.legend(fontsize=8)
    axes[0].set_ylabel('成功起点 / 12');axes[0].set_ylim(0,13);axes[1].set_ylabel('优化期试次forward调用总数')
    pages.append('## 三被试、两条处理路径\n\n'+save(fig,'solver_summary')+
        '\n全部四个划分、每组三个起点均计入。三位被试来自同一公开数据集；两种光学处理使用同源窗口，不是六个独立数据集。')
    fig,axes=plt.subplots(1,3,figsize=(11,3.8))
    for ax,metric,label in zip(axes,['optimization_evaluations','cpu_seconds','seconds'],['优化调用','进程CPU秒','墙钟秒']):
        for success,marker,name in [(True,'o','两臂成功'),(False,'x','至少一臂失败')]:
            sub=real[(real.old_success&real.new_success)==success]
            ax.scatter(sub['old_'+metric],sub['new_'+metric],s=17,marker=marker,label=name)
        lim=max(real['old_'+metric].max(),real['new_'+metric].max())*1.05
        ax.plot([0,lim],[0,lim],'k--',lw=.8);ax.set(xlim=(0,lim),ylim=(0,lim),xlabel='原Armijo '+label,ylabel='二次插值 '+label)
    axes[0].legend(fontsize=7)
    call_changes=real.new_optimization_evaluations/real.old_optimization_evaluations-1
    increased_count=int((call_changes>0).sum());worst_increase=float(call_changes.max())
    slower_groups=part[part.optimization_evaluations_reduction<0]
    slower_text='；'.join(f"{r.subject.replace('subject_','S')}/{r.pipeline} +{-100*r.optimization_evaluations_reduction:.2f}%" for r in slower_groups.itertuples()) or '无分组总调用增加'
    pages.append('## 每个起点的计算代价\n\n'+save(fig,'paired_cost')+
        f'\n对角线下方表示减少，叉号保留数值失败。{increased_count}/72起点调用增加，最大+{100*worst_increase:.2f}%。分组增加：{slower_text}。相同调用数不保证相同耗时；额外插值本身也有成本。')
    affected=['no_motion__subject_01_o1','no_motion__subject_18_o0','mne_tddr__subject_01_o3','mne_tddr__subject_18_o3']
    fig,axes=plt.subplots(4,2,figsize=(10,10))
    for i,group in enumerate(affected):
        ref=json.loads((run/'arms/quadratic_interpolation/training'/f'{group}__conditional_trained_gain/start_0/result.json').read_text())['parameter_value']
        for arm,color,label in [('armijo','#2877ab','原Armijo'),('quadratic_interpolation','#dc8030','二次插值')]:
            record=json.loads((run/'arms'/arm/'training'/f'{group}__conditional_trained_gain/start_0/result.json').read_text())
            trace=record['trace'];tail=trace[-10:]
            axes[i,0].plot(np.arange(1-len(tail),1),[100*(t['parameter_value']/ref-1) for t in tail],'.-',color=color,label=label,ms=3)
            axes[i,1].semilogy([t['evaluations'] for t in trace],[max(t['projected_scaled_gradient_inf_norm'],1e-14) for t in trace],color=color,label=label)
        axes[i,0].axhline(0,color='.6',lw=.6)
        axes[i,0].set(title=group+' / start0',xlabel='距本臂停止的记录步（0为停止）',ylabel='β相对新算法终值的偏差 / %')
        axes[i,1].axhline(1e-4,color='k',ls='--',lw=.8);axes[i,1].set(xlabel='累计优化forward调用',ylabel='缩放投影梯度')
        axes[i,0].ticklabel_format(axis='y',style='plain',useOffset=False)
        for ax in axes[i]:ax.legend(fontsize=7)
    pages.append('## 原先四个困难训练组\n\n'+save(fig,'convergence_traces')+
        '\n固定展示预声明四组的start0。左图为各自停止前10个记录点，只比较尾段变化；横轴不代表两臂用了相同迭代数。右图以实际调用量展示完整收敛过程。目标变化很小并不等于梯度达到阈值。')
    common=real[real.old_success&real.new_success]
    synthetic_screen=json.loads((run/'arms/quadratic_interpolation/synthetic_training_summary.json').read_text())['engineering_screen']
    availability=real.groupby('group')[['old_success','new_success']].sum()
    gain_spans=[]
    for group,part in real[real.new_success].groupby('group'):
        values=part.new_relative_gain.to_numpy(float)
        gain_spans.append(dict(group=group,successful_starts=len(values),relative_gain_span=(values.max()-values.min())/np.median(values)))
    pd.DataFrame(gain_spans).to_csv(out/'multistart_gain_consistency.csv',index=False)
    agreement=dict(common_success_pairs=len(common),legacy_reproduced=int(frame.legacy_reproduced.sum()),total_pairs=len(frame),
        synthetic_recovery_all_passed=synthetic_screen['all_groups_passed'],
        synthetic_maximum_gain_recovery_error=max(r['selected_relative_error'] for r in synthetic_screen['groups']),
        old_available_training_groups=int((availability.old_success>0).sum()),
        new_available_training_groups=int((availability.new_success>0).sum()),
        new_maximum_multistart_relative_gain_span=max(r['relative_gain_span'] for r in gain_spans),
        regressions=int(real.regressed.sum()),recovered=int(real.recovered.sum()),
        maximum_relative_objective_difference=float(common.objective_relative_difference.max()),
        maximum_relative_gain_difference=float(common.gain_relative_difference.max()),
        maximum_prediction_difference_training_sd=float(common.prediction_max_difference_sd.max()),
        maximum_driver_difference=float(common.driver_max_difference.max()),maximum_flow_difference=float(common.flow_max_difference.max()),
        interpolation_attempts=int(real.interpolation_attempts.sum()),interpolation_selected=int(real.interpolation_selected.sum()),
        measured_starts_with_increased_calls=increased_count,maximum_call_increase_fraction=worst_increase,
        max_integration_difference_training_sd=float(real.new_integration_gap.max()))
    (out/'audit.json').write_text(json.dumps(agreement,indent=2))
    pages.append('## 是否改变了解\n\n'+table(pd.DataFrame([agreement]).T.reset_index().rename(columns={'index':'核验',0:'值'}))+
        '\n一致性只比较两臂均成功的起点；修复失败单独计数。配对目标、训练尺度、试次身份逐元素一致；原算法另与历史结果核对。')
    pages.append('## 范围与限制\n\n'+f'来源：{run.name}。全部108对训练起点，未运行新的验证重建，没有独立数据集泛化证据。数值收敛改善不等于NRMSE已改善，也不证明β为唯一可识别的生理量。\n\n'+
        '默认Armijo仍保留；二次插值为显式选项。旧证据不覆盖。完整指标见paired_starts.csv，分组汇总见subject_pipeline_summary.csv。图为240dpi位图，PDF正文可搜索；WPS未检查。')
    export_pdf('\n\n'.join(pages),out,figures,split_images=False)
    # A short editable presentation shares the audited tables and bitmap figures.
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches,Pt
    deck=Presentation();deck.slide_width=Inches(13.333);deck.slide_height=Inches(7.5)
    def box(slide,text,x,y,w,h,size=22,color='17364B',bold=False):
        shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=shape.text_frame;tf.word_wrap=True
        for i,line in enumerate(text.split('\n')):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.space_after=Pt(14)
            for r in p.runs:r.font.name='Noto Sans CJK SC';r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=RGBColor.from_string(color)
    def slide(title,subtitle):
        s=deck.slides.add_slide(deck.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor(246,248,251)
        box(s,title,.6,.35,12.1,.65,30,bold=True);box(s,subtitle,.65,1.10,12,.48,15,color='597184')
        box(s,f'2026-09-28 · 同目标求解器验证 · {len(deck.slides)}',.65,7.03,12,.3,11,color='597184')
        s.notes_slide.notes_text_frame.text='事实源：'+str(run)+'\n逐起点指标：'+str(out/'paired_starts.csv')+'\n训练数值诊断，不是独立数据集或生理参数资格结论。'
        return s
    reduction=1-real.new_optimization_evaluations.sum()/real.old_optimization_evaluations.sum()
    cpu_reduction=1-real.new_cpu_seconds.sum()/real.old_cpu_seconds.sum()
    s=slide('求解器开发：收敛与速度的完整对照','相同目标、相同起点、相同预算；两臂均重新运行')
    box(s,f"实测成功起点\n{int(real.old_success.sum())}/72 → {int(real.new_success.sum())}/72",.75,2.05,3.8,1.9,28,bold=True)
    box(s,f'优化调用减少\n{100*reduction:.1f}%',4.85,2.05,3.5,1.9,28,bold=True)
    box(s,f'总CPU时间减少\n{100*cpu_reduction:.1f}%',8.95,2.05,3.7,1.9,28,bold=True)
    syn_change=100*(syn.new_optimization_evaluations.sum()/syn.old_optimization_evaluations.sum()-1)
    box(s,f"可用训练组 {agreement['old_available_training_groups']}/24 → {agreement['new_available_training_groups']}/24；救回失败 {int(real.recovered.sum())} 个，退步 {int(real.regressed.sum())} 个。\n合成成功 {int(syn.old_success.sum())}/36 → {int(syn.new_success.sum())}/36；调用量变化 {syn_change:+.2f}%。",.8,4.65,11.8,1.85,23)
    s=slide('覆盖多个被试与处理条件，冻结评价规则','不是独立数据集泛化：所有实测来自同一公开数据集的保留目标')
    box(s,'12组合成条件 × 3起点 = 36对拟合\n3被试 × 2处理路径 × 4划分 × 3起点 = 72对拟合\n每次用18条训练试次；未重新拟合验证目标\n'+f"合成β最大恢复误差 {100*agreement['synthetic_maximum_gain_recovery_error']:.2f}%",.85,1.85,11.6,2.95,24)
    box(s,'原目标、90次迭代、3600次试次forward预算、梯度阈值1e−4均不变。\n采用相同积分器，保留物理域检查和4/8子步积分复核。',.85,5.0,11.6,1.35,21)
    s=slide('逐被试、逐处理路径看成功率与计算量','蓝：原Armijo；橙：有条件的二次插值；失败起点保留在分母中')
    s.shapes.add_picture(str(out/'figures/solver_summary.png'),Inches(.65),Inches(1.72),width=Inches(12.0))
    s=slide('收益不均匀：部分起点增加计算量','每个点是一对相同起点；低于对角线表示新方法减少计算代价')
    s.shapes.add_picture(str(out/'figures/paired_cost.png'),Inches(.65),Inches(1.78),width=Inches(12.0))
    box(s,f'{increased_count}/72起点调用增加，最大+{100*worst_increase:.1f}%。\n{slower_text}。',.75,6.02,11.8,.90,16)
    s=slide('速度改善是否改变了解？','只在两臂都成功的相同起点上比较，新增成功单独统计')
    box(s,f"共同成功起点：{len(common)}\n目标值最大相对差：{agreement['maximum_relative_objective_difference']:.3g}\nβ最大相对变化：{100*agreement['maximum_relative_gain_difference']:.4f}%\n预测曲线最大差：{agreement['maximum_prediction_difference_training_sd']:.4g} 个训练SD\n血流f最大绝对差：{agreement['maximum_flow_difference']:.4g}",.85,1.85,11.5,4.45,25)
    s=slide('适用范围与仍未解决的问题','数值优化更可靠，不自动等于模型或生理解释已经通过验收')
    box(s,'只在实际下降远小于GN预测时，额外尝试一个较短步长。\n只有原目标更低、物理域有效时才采用；不放宽成功条件。\n显式启用：step_control="quadratic_interpolation"。',.85,1.85,11.6,2.9,24)
    box(s,'本次未验证跨独立数据集、验证NRMSE改善或生理参数可辨识性。\n原Armijo默认路径与旧证据保留；困难波形仍需独立的模型诊断。',.85,5.1,11.6,1.2,21)
    deck.save(out/'SOLVER_RESULTS.pptx')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--run", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--replay", action="store_true", help="Render nonlinear fixed-driver replay against its parent")
    parser.add_argument("--nonlinear-fit", action="store_true", help="Render nonlinear optimization cells and complete denominators")
    parser.add_argument("--gain-prior-fit", action="store_true", help="Render four-arm effective gain and driver amplitude prior diagnostic")
    parser.add_argument("--terminal-input-snapshot", type=Path, help="Frozen terminal manifest/config/summary/metrics for report replay while measured phase proceeds")
    parser.add_argument("--pilot-truth-audit", type=Path, help="Optional retained pilot DC/initial compensation JSON for conditional report background")
    parser.add_argument("--conditional-optical-gain", action="store_true", help="Render terminal conditional optical mapping and gain evidence")
    parser.add_argument("--fixed-roi-optical", action="store_true", help="Render separate retained optical target sensitivity conditions")
    parser.add_argument("--fixed-roi-initial", action="store_true", help="Render free vs tied initial total-Hb/volume contrast")
    parser.add_argument("--fixed-roi-flow", action="store_true", help="Render fixed ROI log-flow soft regularization diagnostic")
    parser.add_argument("--fixed-roi-tau", action="store_true", help="Render fixed AF7Fp1 shared nonlinear tau diagnostic")
    parser.add_argument("--step-control", action="store_true", help="Audit and render the complete paired training-only solver comparison")
    args = parser.parse_args()
    if sum((args.replay, args.nonlinear_fit, args.gain_prior_fit, args.fixed_roi_tau, args.fixed_roi_flow, args.fixed_roi_initial, args.fixed_roi_optical, args.conditional_optical_gain, args.step_control)) > 1:
        parser.error("Choose one report mode")
    if (args.pilot_truth_audit or args.terminal_input_snapshot) and not args.conditional_optical_gain:
        parser.error("--pilot-truth-audit requires --conditional-optical-gain")
    if args.step_control:
        render_step_control(args.run.resolve(),args.output.resolve())
        return
    if args.conditional_optical_gain:
        render_conditional_optical_gain(args.run.resolve(),args.output.resolve(),pilot_truth_audit=args.pilot_truth_audit,terminal_input_snapshot=args.terminal_input_snapshot)
        return
    if args.fixed_roi_optical:
        render_fixed_roi_optical(args.run.resolve(),args.output.resolve())
        return
    if args.fixed_roi_initial:
        render_nonlinear_fit(args.run.resolve(), args.output.resolve(), fixed_roi_initial=True)
        return
    if args.fixed_roi_flow:
        render_nonlinear_fit(args.run.resolve(), args.output.resolve(), fixed_roi_flow=True)
        return
    if args.fixed_roi_tau:
        render_nonlinear_fit(args.run.resolve(), args.output.resolve(), fixed_roi_tau=True)
        return
    if args.gain_prior_fit:
        render_nonlinear_fit(args.run.resolve(), args.output.resolve(), gain_prior=True)
        return
    renderer = render_nonlinear_fit if args.nonlinear_fit else (render_replay if args.replay else render)
    renderer(args.run.resolve(), args.output.resolve())


if __name__ == "__main__":
    main()
