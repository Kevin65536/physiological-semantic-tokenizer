#!/usr/bin/env python3
"""Render retained shared-driver linear diagnostic evidence; never fit models."""
from __future__ import annotations

import argparse
import json
from pathlib import Path
import re

import fitz
import markdown
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
import numpy as np
import pandas as pd


COMPONENTS = ["EEG", "HbO", "HbR"]
METHODS = ["fixed_tau", "fitted_tau", "source_prior_tau", "oracle_lower_bound", "own_context"]
LABELS = {"fixed_tau": "固定 τ", "fitted_tau": "训练拟合 τ", "source_prior_tau": "来源先验 τ", "oracle_lower_bound": "全目标 oracle 下界", "own_context": "自身上下文"}


def truth(series):
    return series.astype(str).str.lower().isin(["true", "1", "1.0"])


def table(frame):
    def cell(value):
        if isinstance(value, (float, np.floating)):
            return f"{value:.4f}" if np.isfinite(value) else "—"
        return str(value).replace("|", "/")
    return "| " + " | ".join(frame.columns) + " |\n| " + " | ".join(["---"] * len(frame.columns)) + " |\n" + "\n".join("| " + " | ".join(map(cell, row)) + " |" for row in frame.itertuples(index=False, name=None)) + "\n"


def summarize(data):
    rows = []
    for (kind, method, mode), part in data.groupby(["kind", "method", "mode"], sort=False):
        completed = part.status.isin(["completed", "ok", "success"])
        scored = COMPONENTS if mode == "full" else (["EEG"] if mode in ["center_EEG", "center_EEG_shift", "fNIRS_only"] else ["HbO", "HbR"])
        finite = np.isfinite(part[[f"nrmse_{c}" for c in scored]].to_numpy(float)).all(axis=1)
        valid = completed & truth(part.physical_valid) & truth(part.small_signal_valid) & finite
        passed = valid & part[[f"nrmse_{c}" for c in COMPONENTS]].lt(.5).all(axis=1)
        row = dict(kind=kind, method=method, mode=mode, total=len(part), valid=int(valid.sum()), failed=int((~valid).sum()), all_three_below_05=int(passed.sum()), physical_valid=int(truth(part.physical_valid).sum()), small_signal_valid=int(truth(part.small_signal_valid).sum()))
        if mode != "full":
            row["all_three_below_05"] = "N/A"
        if method == "own_context":
            row.update(valid="N/A", failed="N/A", all_three_below_05="N/A", physical_valid="N/A", small_signal_valid="N/A")
        # Fixed full denominator is retained; finite errors are descriptive,
        # including physically invalid reconstructions, never a success count.
        for component in COMPONENTS:
            values = pd.to_numeric(part[f"nrmse_{component}"], errors="coerce").replace([np.inf, -np.inf], np.nan)
            scoring = part[["subject", "session", f"nmse_{component}", f"nrmse_{component}"]].copy()
            scoring.replace([np.inf, -np.inf], np.nan, inplace=True)
            means = scoring.groupby(["subject", "session"])[[f"nmse_{component}", f"nrmse_{component}"]].mean().groupby("subject").mean().mean()
            row[f"NRMSE_{component}"] = np.sqrt(means[f"nmse_{component}"])
            row[f"NMSE_{component}"] = means[f"nmse_{component}"]
            row[f"mean_trial_NRMSE_{component}"] = means[f"nrmse_{component}"]
            row[f"finite_{component}"] = int(values.notna().sum())
        rows.append(row)
    return pd.DataFrame(rows)


def render(run, out):
    if out.exists() and any(out.iterdir()):
        raise ValueError("Export directory is not empty; choose a new versioned --output")
    data = pd.read_csv(run / "metrics.csv")
    required = {"kind", "group", "subject", "outer", "trial", "session", "method", "mode", "status", "physical_valid", "small_signal_valid", *(f"nrmse_{c}" for c in COMPONENTS), *(f"nmse_{c}" for c in COMPONENTS)}
    if not required.issubset(data.columns):
        raise ValueError(f"Missing metrics columns: {sorted(required - set(data.columns))}")
    if data.empty:
        raise ValueError("No completed evidence rows")
    out.mkdir(parents=True, exist_ok=True)
    (out / "figures").mkdir()
    font = Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc")
    if font.exists():
        font_manager.fontManager.addfont(str(font))
        plt.rcParams["font.family"] = font_manager.FontProperties(fname=str(font)).get_name()
    plt.rcParams.update({"font.size": 9, "axes.unicode_minus": False})
    figures = []

    def save(fig, name, caption):
        fig.tight_layout()
        fig.savefig(out / "figures" / f"{name}.png", dpi=240, bbox_inches="tight")
        plt.close(fig)
        figures.append((name, caption))

    summary = summarize(data)
    summary.to_csv(out / "summary.csv", index=False)
    kinds = list(data.kind.unique())
    fig, axes = plt.subplots(1, len(kinds), figsize=(6 * len(kinds), 4.2), squeeze=False)
    for ax, kind in zip(axes.flat, kinds):
        part = summary[(summary.kind == kind) & (summary["mode"] == "full")]
        x = np.arange(len(part))
        for j, c in enumerate(COMPONENTS):
            ax.bar(x + (j - 1) * .23, part[f"NRMSE_{c}"], .23, label=c)
        ax.axhline(.5, color="red", ls="--", lw=1)
        ax.set_xticks(x, [LABELS.get(m, m) for m in part.method], rotation=25, ha="right")
        ax.set_title(f"{kind}：全观测重建")
        ax.set_ylabel("NRMSE = √被试／会话等权 NMSE")
        ax.legend()
    save(fig, "aggregate_nrmse", "全观测误差包含物理失效结果；按 trial→session→subject 聚合 NMSE 后取根，oracle 不参与方法选择。")

    group_records = []
    for group in data.group.astype(str).unique():
        path = run / "groups" / f"{group}.json"
        if path.exists():
            group_records.append((group, json.loads(path.read_text())))
    measured_groups = [(g, r) for g, r in group_records if g in set(data.loc[data.kind.eq("measured"), "group"].astype(str))]
    records = measured_groups or group_records
    fig, axes = plt.subplots(2, 1, figsize=(10, 7))
    for method in METHODS[:3]:
        values = [record.get("selected_tau", {}).get(method, np.nan) for _, record in records]
        axes[0].plot(range(len(records)), values, "o-", label=LABELS[method])
    axes[0].set_xticks(range(len(records)), [g for g, _ in records], rotation=45, ha="right", fontsize=6)
    axes[0].set_ylabel("τ（秒）")
    axes[0].legend()
    for group, record in records:
        profile = pd.DataFrame(record.get("profiles", []))
        if {"tau", "train_sse"}.issubset(profile.columns):
            profile = profile.sort_values("tau")
            values = profile.train_sse.to_numpy(float)
            axes[1].plot(profile.tau, values - np.nanmin(values), label=group)
    axes[1].set_xscale("log")
    axes[1].set_xlabel("τ（秒）")
    axes[1].set_ylabel("训练 SSE − 本折最小值")
    if records:
        axes[1].legend(fontsize=6, ncol=3)
    save(fig, "tau_profiles", "被试／折参数与训练残差剖面；此 SSE 剖面不是置信区间，折间变化不等同于真实个体差异。")

    # Select examples once by the fixed baseline; candidates cannot pick their
    # own easiest examples. Large or invalid fits remain eligible for display.
    examples = data[data.kind.eq("measured") & data.method.eq("fixed_tau") & data["mode"].eq("full")].copy()
    if examples.empty:
        examples = data[data.method.eq("fixed_tau") & data["mode"].eq("full")].copy()
    examples["rank_error"] = examples[[f"nrmse_{c}" for c in COMPONENTS]].mean(axis=1)
    examples = examples.sort_values(["rank_error", "group", "trial"], kind="stable", na_position="last")
    selected = []
    if len(examples):
        selected = [("median", examples.iloc[(len(examples) - 1) // 2]), ("worst", examples.iloc[-1])]
    for label, row in selected:
        path = run / "predictions" / f"{row.group}.npz"
        with np.load(path, allow_pickle=False) as arrays:
            indices = arrays["trial_indices"]
            matched = np.flatnonzero(indices.astype(str) == str(row.trial))
            if not len(matched):
                raise ValueError(f"Missing example trial {row.trial} in {path}")
            index = int(matched[0])
            target = arrays["target"][index]
            t = np.arange(len(target)) * .25
            fig, axes = plt.subplots(3, 2, figsize=(11, 7.5), sharex=True)
            for j, c in enumerate(COMPONENTS):
                axes[j, 0].plot(t, target[:, j], "k", lw=1.5, label="实测／目标")
                axes[j, 0].set_ylabel(c + "（原评分坐标）")
                for method in METHODS[:4]:
                    key = f"prediction_{method}_full"
                    if key in arrays:
                        axes[j, 0].plot(t, arrays[key][index, :, j], lw=1, label=LABELS[method], alpha=.85)
                state_key = "states_fitted_tau_full"
                if state_key in arrays:
                    states = arrays[state_key][index]
                    for k in [2 * j, 2 * j + 1]:
                        axes[j, 1].plot(t, states[:, k], label=["r", "s", "f", "v", "p", "q"][k])
                axes[j, 1].legend()
            axes[0, 0].legend(fontsize=7)
            axes[-1, 0].set_xlabel("秒")
            axes[-1, 1].set_xlabel("秒")
            fig.suptitle(f"{label}：{row.group} / trial {row.trial}；按固定 τ 误差选样")
            save(fig, f"curves_{label}", f"固定 τ 基线排序的 {label} 样例；右列为拟合 τ 的六状态，失效状态不裁剪。")

    full = summary[summary["mode"].eq("full")]
    text = "# 共享生理驱动重建：线性小信号诊断\n\n"
    text += f"证据来源：`{run}`。本报告包含 {len(data)} 条评价记录；不将本轮线性诊断称为非线性 SSM 完成或生理 teacher 资格。\n\n"
    text += "## 主要结果\n\n"
    for row in full.itertuples():
        text += f"- {row.kind}／{LABELS.get(row.method, row.method)}：EEG、HbO、HbR 聚合 NRMSE 分别为 {row.NRMSE_EEG:.3f}、{row.NRMSE_HbO:.3f}、{row.NRMSE_HbR:.3f}；物理及小信号有效 {row.valid}/{row.total}，失败 {row.failed}，三分量均低于 0.5 且有效 {row.all_three_below_05}/{row.total}。\n"
    text += "\n主 NRMSE 是逐 trial NMSE 在 session 内平均、session 在 subject 内等权平均、subject 等权平均后取平方根，不是逐 trial NRMSE 的算术平均。后者仅在 summary.csv 单列描述。有限误差输出不能代替成功比例；失败、评分分量非有限、物理无效或超出小信号范围均不算生理重建成功。遮挡模式只评分隐藏分量，其他分量空白不算失败，三分量共同阈值记 N/A。自身上下文对照没有生理状态，其物理有效性标为 N/A，不计为生理模型失败。\n\n"
    for kind in kinds:
        subset = full[full.kind.eq(kind)].set_index("method")
        if "fixed_tau" not in subset.index:
            continue
        baseline = subset.loc["fixed_tau"]
        for method in ["fitted_tau", "source_prior_tau"]:
            if method not in subset.index:
                continue
            candidate = subset.loc[method]
            differences = [candidate[f"NRMSE_{c}"] - baseline[f"NRMSE_{c}"] for c in COMPONENTS]
            text += f"{kind}／{LABELS[method]} 相对固定 τ 的三分量聚合 NRMSE 差为 {differences[0]:+.3f}、{differences[1]:+.3f}、{differences[2]:+.3f}（负数表示改善；不是配对置信区间）。"
            passed = all(candidate[f"NRMSE_{c}"] < .5 for c in COMPONENTS)
            text += ("有限输出的三分量聚合误差达到 0.5 以下。" if passed else "有限输出的三分量聚合误差尚未全部达到 0.5 以下。")
            if candidate.failed:
                text += f"仍有 {int(candidate.failed)} 条生理失效或数值失败，不能宣布生理重建达标。"
            text += "\n\n"
    measured_full = full[full.kind.eq("measured")].set_index("method")
    if "fitted_tau" in measured_full.index:
        fit = measured_full.loc["fitted_tau"]
        if fit.NRMSE_HbO < .5 and fit.NRMSE_HbR < .5 and fit.failed:
            text += "**本轮关键结论：两项 Hb 的残差能够低于目标阈值，但共享状态超出已声明的小信号／物理域。曲线贴合没有消除状态异常，因此线性化的生理解释不可信，不能宣布生理重建达标。** 下一步应以受约束非线性动力学检查这些状态是否真实可实现，并保留当前负结果；本报告不把拟合 τ 差异宣布为生理个体差异。\n\n"
    synthetic = data[data.kind.eq("synthetic") & data["mode"].eq("full")]
    if "driver_nrmse" in synthetic and len(synthetic):
        text += "合成驱动恢复（没有事后尺度或时间对齐）：\n\n"
        diagnostics = synthetic.groupby("method")["driver_nrmse"].agg(["count", "mean", "max"]).reset_index()
        text += table(diagnostics) + "\n"
    columns = ["kind", "method", "mode", "total", "valid", "failed", "all_three_below_05", "NRMSE_EEG", "NRMSE_HbO", "NRMSE_HbR"]
    text += "## 全部分母和遮挡对照\n\n" + table(summary[columns])
    text += "\n## 方法与解释边界\n\n"
    text += "仅一个共享驱动和五个血流初态；血流状态没有逐点自由创新。参数只在训练折选择，评分采用冻结训练 SD。全观测重建与中心特征遮挡补全分别报告；离线处理后的补全不是原始传感器遮挡或未来预测。EEG 坐标是 log-power PCA，不是原始电压。\n\n"
    text += "oracle 使用同一目标选择最小 SSE，仅为有限 4 Hz ZOH 驱动、现有测量插值算子和 τ 网格所定义线性族的乐观残差下界，不能参与模型选择、个体参数解释或预测优劣比较。仅当三分量聚合 NMSE 的平均值大于 0.25，才可在这一有限线性族中排除三分量聚合 NRMSE 同时低于 0.5；不能对逐 trial NRMSE 算术平均作此推论。先验 τ 的工程 SSE 惩罚不是校准后验。τ 的离散剖面和折间一致性只用于诊断，不能证明参数唯一可辨识。\n\n"
    text += "[Tak 等（2015）](https://www.fil.ion.ucl.ac.uk/~wpenny/publications/tak-penny15.pdf)区分血流参数与 HbO/HbR 光学权重，并采用参数先验；这不意味着相同先验已获当前数据验证。[Friston 等（2016）](https://pmc.ncbi.nlm.nih.gov/articles/PMC4767224/)提供保留个体差异的层次收缩方法；本轮尚未实现层次模型。[Raue 等（2009）](https://pubmed.ncbi.nlm.nih.gov/19505944/)支持通过重新优化其余参数的剖面分析评估可辨识性，不能用局部曲率或参数不贴边替代。\n\n"
    text += "## 图像核查\n\n样例按固定 τ 全观测平均 NRMSE 稳定排序，选择中位与最差记录，候选方法共享样例。图像用于检查幅度、峰时、双 Hb 联动及状态异常；不能仅凭曲线接近证明生理语义。\n\n"
    for name, caption in figures:
        text += f"![{name}](figures/{name}.png)\n\n{caption}\n\n"
    text += "## 交付验证\n\n全部图形以 240 dpi PNG 嵌入 PDF，正文与表格保持可搜索。程序检查图像对象及正文提取；WPS 滚动与缩放未检查。\n"
    export_pdf(text, out, figures)


def export_pdf(text, out, figures, *, extra_css="", split_images=True):
    """Keep report prose selectable and embed whole figures as bitmap images."""
    (out / "REPORT.md").write_text(text)
    body = markdown.markdown(text, extensions=["tables"])
    body = body.replace("<img ", '<img width="510" ')
    css = "body{font-family:sans-serif;font-size:9pt;line-height:1.4}h1{font-size:20pt}h2{font-size:13pt}table{font-size:6pt;border-collapse:collapse}td,th{border:0.4pt solid #bbb;padding:3pt}img{max-width:100%}"
    def rectfn(number, filled):
        if number > 80:
            raise RuntimeError("PDF pagination failed")
        page = fitz.paper_rect("a4")
        return page, page + (36, 32, -36, -32), None
    # Separate stories prevent a leading image page break from discarding the
    # preceding prose overflow in MuPDF. Each figure still starts a fresh page.
    pdf = fitz.open()
    for section in re.split(r"(?=<h2>|<p><img )" if split_images else r"(?=<h2>)", body):
        if not section.strip():
            continue
        story = fitz.Story(section, user_css=css + extra_css, archive=fitz.Archive(str(out)))
        part = story.write_with_links(rectfn)
        pdf.insert_pdf(part)
        part.close()
    pdf.save(out / "REPORT.pdf", garbage=4, deflate=True)
    validation = {"pages": len(pdf), "image_objects": sum(len(p.get_images()) for p in pdf), "figure_count": len(figures), "searchable_text_chars": sum(len(p.get_text()) for p in pdf), "wps_checked": False, "bitmap_dpi": 240}
    if validation["image_objects"] != len(figures) or validation["searchable_text_chars"] < 100:
        raise RuntimeError(f"PDF verification failed: {validation}")
    # Render representative pages for visual review without rasterizing delivery.
    for i in sorted({0, max(0, len(pdf) - len(figures)), len(pdf) - 1}):
        pdf[i].get_pixmap(matrix=fitz.Matrix(1.3, 1.3)).save(out / f"preview_page_{i+1:02d}.png")
    pdf.close()
    (out / "export_validation.json").write_text(json.dumps(validation, indent=2))
    print(json.dumps(validation))


def render_replay(run, out):
    """Compare a fixed-drive nonlinear replay against its exact parent fits."""
    if out.exists() and any(out.iterdir()):
        raise ValueError("Export directory is not empty; choose a new versioned --output")
    manifest = json.loads((run / "manifest.json").read_text())
    summary = json.loads((run / "summary.json").read_text())
    data = pd.read_csv(run / "metrics.csv")
    parent = Path(manifest["source_run"])
    if not parent.is_absolute():
        parent = Path(__file__).resolve().parents[2] / parent
    parent_data = pd.read_csv(parent / "metrics.csv")
    out.mkdir(parents=True, exist_ok=True)
    (out / "figures").mkdir()
    font = Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc")
    if font.exists():
        font_manager.fontManager.addfont(str(font))
        plt.rcParams["font.family"] = font_manager.FontProperties(fname=str(font)).get_name()
    plt.rcParams.update({"font.size": 9, "axes.unicode_minus": False})
    figures = []

    def save(fig, name, caption):
        fig.tight_layout()
        fig.savefig(out / "figures" / f"{name}.png", dpi=240, bbox_inches="tight")
        plt.close(fig)
        figures.append((name, caption))

    aggregates = summary["aggregates"]
    rows = []
    for record in aggregates:
        counts = record.get("failure_counts", {})
        row = {"kind": record["kind"], "method": record["method"], "expected": record["expected"], "observed": record["observed_rows"], "completed": record["completed"], "domain_fail": counts.get("failed_domain", 0), "integration_fail": counts.get("failed_integration_check", 0), "missing": record["expected"] - record["observed_rows"]}
        for c in COMPONENTS:
            row[f"linear_{c}"] = record.get(f"linear_nrmse_same_subset_{c}", np.nan)
            row[f"nonlinear_{c}"] = record.get(f"nrmse_{c}", np.nan)
        rows.append(row)
    frame = pd.DataFrame(rows)
    frame.to_csv(out / "replay_summary.csv", index=False)
    kinds = list(frame.kind.unique())
    fig, axes = plt.subplots(len(kinds), 3, figsize=(12, 3.8 * len(kinds)), squeeze=False)
    for i, kind in enumerate(kinds):
        subset = frame[frame.kind.eq(kind)]
        for j, c in enumerate(COMPONENTS):
            ax = axes[i, j]
            x = np.arange(len(subset))
            ax.bar(x - .18, subset[f"linear_{c}"], .36, label="线性：同成功子集")
            ax.bar(x + .18, subset[f"nonlinear_{c}"], .36, label="非线性固定驱动重放")
            ax.axhline(.5, color="red", ls="--", lw=1)
            ax.set_xticks(x, [f"{LABELS[m]}\n{n}/{d} 成功" for m, n, d in zip(subset.method, subset.completed, subset.expected)], rotation=25, ha="right", fontsize=7)
            ax.set_title(f"{kind} / {c}")
            ax.set_ylabel("√被试／会话等权 NMSE")
            if i == 0 and j == 0:
                ax.legend(fontsize=7)
    save(fig, "replay_aggregate", "只在非线性域及积分分辨率检查均通过的同一子集比较误差；完整失败分母单独列出，不将缺失条目算作成功。")

    examples = parent_data[parent_data.kind.eq("measured") & parent_data.method.eq("fixed_tau") & parent_data["mode"].eq("full")].copy()
    if examples.empty:
        examples = parent_data[parent_data.method.eq("fixed_tau") & parent_data["mode"].eq("full")].copy()
    examples["rank_error"] = examples[[f"nrmse_{c}" for c in COMPONENTS]].mean(axis=1)
    examples = examples.sort_values(["rank_error", "group", "trial"], kind="stable", na_position="last")
    if examples.empty:
        raise ValueError("Parent fixed-baseline examples unavailable")
    selections = [("median", examples.iloc[(len(examples) - 1) // 2]), ("worst", examples.iloc[-1])]
    selection_notes = []
    for label, selected in selections:
        replay_path = run / "predictions" / f"{selected.group}.npz"
        if not replay_path.exists():
            selection_notes.append(f"父级 {label} 样例 {selected.group}/trial {selected.trial} 的重放数组缺失；保留身份，不换样。")
            continue
        with np.load(parent / "predictions" / f"{selected.group}.npz", allow_pickle=False) as linear, np.load(replay_path, allow_pickle=False) as nonlinear:
            indices = np.flatnonzero(linear["trial_indices"].astype(str) == str(selected.trial))
            replay_indices = np.flatnonzero(nonlinear["trial_indices"].astype(str) == str(selected.trial))
            if len(indices) != 1 or len(replay_indices) != 1:
                raise ValueError("Parent/replay example identity mismatch")
            index, replay_index = int(indices[0]), int(replay_indices[0])
            target = linear["target"][index]
            if not np.array_equal(target, nonlinear["target"][replay_index]):
                raise ValueError("Parent/replay target mismatch")
            t = np.arange(len(target)) * .25
            fig, axes = plt.subplots(3, 2, figsize=(11, 7.5), sharex=True)
            selected_rows = data[data.group.eq(selected.group) & data.trial.eq(selected.trial)]
            statuses = dict(zip(selected_rows.method, selected_rows.status))
            for j, c in enumerate(COMPONENTS):
                axes[j, 0].plot(t, target[:, j], "k", lw=1.5, label="实测／目标")
                for method, color in [("fixed_tau", "#2877ab"), ("fitted_tau", "#dc8030"), ("source_prior_tau", "#358d53")]:
                    key = f"prediction_{method}_full"
                    axes[j, 0].plot(t, linear[key][index, :, j], "--", color=color, lw=1, label=f"{LABELS[method]} 线性")
                    if statuses.get(method) == "completed":
                        axes[j, 0].plot(t, nonlinear[key][replay_index, :, j], color=color, lw=1.2, label=f"{LABELS[method]} 非线性")
                axes[j, 0].set_ylabel(c + "（原评分坐标）")
                state_key = "states_fitted_tau_full"
                for k, color in zip([2 * j, 2 * j + 1], ["#2877ab", "#dc8030"]):
                    name = ["r", "s", "f", "v", "p", "q"][k]
                    axes[j, 1].plot(t, linear[state_key][index, :, k], "--", color=color, label=f"{name} 线性")
                    if statuses.get("fitted_tau") == "completed":
                        axes[j, 1].plot(t, nonlinear[state_key][replay_index, :, k], color=color, label=f"{name} 非线性")
                axes[j, 1].legend(fontsize=7)
            axes[0, 0].legend(fontsize=6, ncol=2)
            for ax in axes[-1]:
                ax.set_xlabel("秒")
            fig.suptitle(f"父级 {label} 样例：{selected.group} / trial {selected.trial}；不重新拟合")
            status_text = "；".join(f"{LABELS.get(m,m)}={s}" for m, s in statuses.items())
            save(fig, f"replay_{label}", f"父级固定 τ 选定的 {label} 样例。虚线=原线性，实线=非线性重放；右侧为拟合 τ 六状态。{status_text}。失败方法不画非线性曲线，不换样。")

    text = "# 非线性固定驱动重放\n\n"
    text += f"本轮来源：`{run}`；父级线性拟合：`{parent}`。保持父级 τ、驱动 r(t) 及血流初态固定，代入非线性动力学，**不是重新拟合**。\n\n"
    text += f"## 主要结果与完整分母\n\n预期 {summary['expected_groups']} 组，完成 {summary['completed_groups']} 组，组失败 {summary['failed_groups']}；方法×trial 成功 {summary['completed_method_trials']}/{summary['expected_method_trials']}。组作业完成不意味着其中所有轨迹通过。\n\n"
    for row in frame.itertuples():
        text += f"- {row.kind}／{LABELS.get(row.method,row.method)}：成功 {row.completed}/{row.expected}；域失败 {row.domain_fail}、积分检查失败 {row.integration_fail}、缺行 {row.missing}。同成功子集 EEG/HbO/HbR NRMSE：线性 {row.linear_EEG:.3f}/{row.linear_HbO:.3f}/{row.linear_HbR:.3f} → 非线性 {row.nonlinear_EEG:.3f}/{row.nonlinear_HbO:.3f}/{row.nonlinear_HbR:.3f}。\n"
    text += "\n" + table(frame[["kind", "method", "expected", "observed", "completed", "domain_fail", "integration_fail", "missing"]])
    text += "\n同成功子集的聚合误差：\n\n" + table(frame[["kind", "method", *[f"{prefix}_{c}" for prefix in ["linear", "nonlinear"] for c in COMPONENTS]]])
    if "integration_max_difference_training_sd" in data:
        precision = data.groupby("kind").integration_max_difference_training_sd.agg(["count", "median", "max"]).reset_index()
        text += "\n积分分辨率核查：下表比较两个子步分辨率的预测差，单位为冻结训练 SD；域失败没有可用积分误差，不填零。\n\n" + table(precision)
        for row in precision.itertuples():
            text += f"{row.kind}：最大差 {row.max:.6g} 个训练 SD，中位差 {row.median:.6g}；有效比较 {row.count} 条。\n\n"
    text += "\n## 结论与限制\n\n误差比较采用同成功子集，按 session、subject 等权聚合 NMSE 后取根；该条件性指标不代表失败轨迹的误差，也不等同于完整面板达标。各方法成功子集可能不同，不能用此表直接作跨方法胜负排名。\n\n"
    for row in frame[frame.kind.eq("measured")].itertuples():
        values = [row.nonlinear_EEG, row.nonlinear_HbO, row.nonlinear_HbR]
        reached = all(np.isfinite(v) and v < .5 for v in values)
        text += f"{LABELS.get(row.method,row.method)}：" + ("成功子集三分量误差均低于 0.5。" if reached else "成功子集三分量误差尚未全部低于 0.5。")
        if row.completed < row.expected:
            text += f"另有 {row.expected-row.completed} 条未通过，不能宣称全分母达标。"
        text += "\n\n"
    text += "本轮只能判断既有线性驱动／初态在非线性动力学中的可实现性与重放误差；不能将重放失败等同于所有非线性模型不可拟合，也不能把成功重放称为参数唯一可辨识或 teacher 资格。下一步若需要降低误差，应另行实现受约束非线性重拟合并保持共同训练／验证分母。oracle 继承全目标拟合驱动，只作诊断，不是可部署候选或选择依据。\n\n"
    oracle = frame[frame.kind.eq("measured") & frame.method.eq("oracle_lower_bound")]
    if len(oracle) and all(oracle.iloc[0][f"nonlinear_{c}"] < .5 for c in COMPONENTS):
        text += "oracle 同成功子集的非线性误差仍全部低于 0.5，提示完整非线性建模存在值得验证的空间，不能据此断言非线性重建不可能。但该驱动及逐 trial τ 由同一目标优化而来，并且仍有完整分母中的失败；它不能作为主方法成功或泛化资格。\n\n"
    text += "## 固定样例与图像核查\n\n中位与最差样例严格沿用父级 fixed_tau/full 的误差排序；失败时保留样例及失败说明，不挑换漂亮曲线。\n\n"
    text += "\n\n".join(selection_notes) + "\n\n"
    for name, caption in figures:
        text += f"![{name}](figures/{name}.png)\n\n{caption}\n\n"
    text += "## 交付验证\n\n全部图形作为 240 dpi PNG 嵌入 PDF；正文与表格可搜索。WPS 滚动／缩放未检查。\n"
    export_pdf(text, out, figures)


def continuation_parent_comparison(run, out, config, data, labels):
    """Compare exactly matched parent identities; newly available fits are separate."""
    manifest = json.loads((run / 'manifest.json').read_text())
    parent = Path(config['continuation']['parent_run'])
    if not parent.is_absolute():
        parent = Path(manifest['project_root']) / parent
    previous = pd.read_csv(parent / 'metrics.csv')
    identity = ['kind','group','subject','session','trial','method','mode']
    if data.duplicated(identity).any() or previous.duplicated(identity).any():
        raise ValueError('Continuation parent comparison has duplicate trial identities')
    joined = data.merge(previous, on=identity, suffixes=('_new','_parent'), how='outer', indicator=True, validate='one_to_one')
    if not joined['_merge'].eq('both').all():
        raise ValueError('Continuation and parent identity sets differ')
    expected = len(config['subjects'])*len(config['outer_folds'])*6
    if expected != 72:
        raise ValueError('Continuation comparison requires the frozen 72-trial panel')
    for (kind, method, mode), part in joined.groupby(['kind','method','mode']):
        if len(part) != expected or len(part.group.unique()) != 12 or not part.groupby('group').size().eq(6).all():
            raise ValueError('Continuation comparison does not contain all 72 identities')
    fixed = joined[joined.method.isin(['fixed_gain_no_amplitude','fixed_gain_amplitude'])]
    if (not fixed.status_new.eq(fixed.status_parent).all()
            or not truth(fixed.converged_new).eq(truth(fixed.converged_parent)).all()):
        raise ValueError('Inherited fixed-gain status changed from parent')
    for prefix in ['nmse_', 'nrmse_']:
        for component in COMPONENTS:
            field = prefix+component
            a = pd.to_numeric(fixed[field+'_new'], errors='coerce').to_numpy()
            b = pd.to_numeric(fixed[field+'_parent'], errors='coerce').to_numpy()
            if not np.array_equal(a, b, equal_nan=True):
                raise ValueError('Inherited fixed-gain score changed from parent')
    joined['success_new'] = joined.status_new.eq('completed') & truth(joined.converged_new)
    joined['success_parent'] = joined.status_parent.eq('completed') & truth(joined.converged_parent)
    transitions, scores = [], []
    for (kind, method, mode), part in joined.groupby(['kind','method','mode'], sort=False):
        groups = {'success_to_success':part.success_new & part.success_parent,
                  'new_success':part.success_new & ~part.success_parent,
                  'lost_success':~part.success_new & part.success_parent,
                  'both_failed':~part.success_new & ~part.success_parent}
        transitions.append(dict(kind=kind, method=method, mode=mode, expected=expected,
            **{name:int(mask.sum()) for name,mask in groups.items()}))
        for subset, mask in groups.items():
            selected = part[mask]
            if subset == 'both_failed': continue
            for component in COMPONENTS:
                field = 'nmse_'+component
                row = dict(kind=kind, method=method, mode=mode, subset=subset,
                    common_or_new_identities=len(selected), expected=expected, component=component)
                for suffix in ['parent','new']:
                    usable = selected if subset == 'success_to_success' else selected[selected['success_'+suffix]]
                    column = field+'_'+suffix
                    value = usable.groupby(['subject','session'])[column].mean().groupby('subject').mean().mean()
                    row[suffix+'_nrmse'] = float(np.sqrt(value))
                if subset == 'success_to_success':
                    delta = selected[['subject','session']].assign(delta=selected[field+'_new']-selected[field+'_parent'])
                    row['delta_nmse'] = delta.groupby(['subject','session']).delta.mean().groupby('subject').mean().mean()
                else:
                    row['delta_nmse'] = np.nan
                if np.isfinite(row['parent_nrmse']) or np.isfinite(row['new_nrmse']): scores.append(row)
    transitions = pd.DataFrame(transitions); scores = pd.DataFrame(scores)
    transitions.to_csv(out / 'continuation_parent_transitions.csv', index=False)
    scores.to_csv(out / 'continuation_parent_paired_scores.csv', index=False)
    body = '\n## 与父运行的同身份验证对照\n\n'
    body += (f'父证据：`{parent.name}/metrics.csv`。按kind/group/subject/session/trial/method/mode一对一核对，'
        '每种数据、方法和模式均保留原12组×6条=72个身份。固定A/B的status、converged及各分量NMSE/NRMSE与父证据逐值一致（含相同缺失值）。'
        '下面将成功状态迁移与误差比较分开；新增成功样本不混入共同成功的改善量。继承的合成与固定臂不是新独立复制。\n\n')
    for (kind, mode), part in transitions.groupby(['kind','mode'], sort=False):
        body += f'### {kind} · {mode} · 成功状态迁移\n\n'
        body += table(pd.DataFrame([{'臂':labels[r.method].split()[0], '原成功→成功':r.success_to_success,
            '新增成功':r.new_success, '失去成功':r.lost_success, '双方失败':r.both_failed, '完整分母':r.expected}
            for r in part.itertuples()])) + '\n\n'
    if len(scores):
        # The key scientific comparison is measured C/D; full CSV also keeps inherited arms.
        relevant = scores[scores.kind.eq('measured') & scores.method.str.startswith('trained_gain_')]
        names = {'success_to_success':'共同成功', 'new_success':'新增成功（无父成功分数）', 'lost_success':'失去成功（无新成功分数）'}
        for (mode, subset), part in relevant.groupby(['mode','subset'], sort=False):
            body += f'### 实测 · {mode} · {names[subset]}\n\n'
            body += table(pd.DataFrame([{'臂':labels[r.method].split()[0], '分量':r.component,
                '身份/计划':f'{r.common_or_new_identities}/{r.expected}', '父NRMSE':r.parent_nrmse,
                '本轮NRMSE':r.new_nrmse, 'ΔNMSE':r.delta_nmse} for r in part.itertuples()])) + '\n\n'
    body += ('NRMSE均按所列相同身份内trial→session→subject等权NMSE后取根；ΔNMSE仅在双方成功时计算。'
        '新增或失去成功只报告对应可用一侧，不拿失败最后有效解补分数，不用不同成功子集的整体均值声称预测改善。\n\n')
    return body


def gain_continuation_details(run, out, config, data, labels, save):
    """Summarize inherited evidence and declared warm-restart traces, without refitting."""
    rows, traces, cells = [], [], []
    for group in data.group.unique():
        kind = data.loc[data.group.eq(group), 'kind'].iloc[0]
        for amplitude in (0, 1):
            method = 'trained_gain_amplitude' if amplitude else 'trained_gain_no_amplitude'
            for start in range(len(config['gain_training']['start_values'])):
                path = run / 'training' / f'{group}__amplitude{amplitude}' / f'start_{start}' / 'result.json'
                result = json.loads(path.read_text()) if path.exists() else {}
                info = result.get('continuation', {})
                inherited = info.get('inherited', False)
                parent_total = result.get('evaluations') if inherited else info.get('parent_evaluations')
                parent_opt = result.get('optimization_evaluations') if inherited else info.get('parent_optimization_evaluations')
                trace = result.get('trace', [])
                first = trace[0] if trace else {}
                parent_objective = info.get('parent_objective')
                initial_objective = first.get('objective', info.get('initial_objective'))
                difference = (initial_objective-parent_objective
                    if initial_objective is not None and parent_objective is not None else np.nan)
                row = dict(kind=kind, group=group, method=method, start=start,
                    action=info.get('action', 'missing_provenance' if result else 'missing'), parent_record=info.get('parent_record'),
                    status=result.get('status', 'missing'), converged=bool(result.get('converged', False)),
                    parent_library_evaluations=parent_total, new_library_evaluations=0 if inherited else info.get('new_evaluations'),
                    cumulative_library_evaluations=parent_total if inherited else info.get('cumulative_evaluations'),
                    parent_optimization_evaluations=parent_opt,
                    new_optimization_evaluations=0 if inherited else info.get('new_optimization_evaluations'),
                    cumulative_optimization_evaluations=parent_opt if inherited else info.get('cumulative_optimization_evaluations'),
                    original_start_gain=result.get('start_gain'), parent_gain=info.get('parent_gain'),
                    initial_gain=first.get('parameter_value'), final_gain=result.get('parameter_value'),
                    parent_objective=parent_objective, initial_objective=initial_objective,
                    first_objective_difference=difference, objective_continuity=info.get('objective_continuity'),
                    continuity_rtol=info.get('objective_continuity_rtol'), continuity_atol=info.get('objective_continuity_atol'),
                    trace_points=len(trace),
                    new_domain_rejections=(None if inherited else result.get('starts', [{}])[0].get('domain_rejections') if result.get('starts') else None),
                    new_fine_replays=0 if inherited else len(result.get('integration_checks', [])))
                rows.append(row)
                for point in trace:
                    traces.append(dict(kind=kind, group=group, method=method, start=start, **point))
        for method, mode in data[data.group.eq(group)][['method','mode']].drop_duplicates().itertuples(index=False, name=None):
            path = run / 'cells' / f'{group}__{method}__{mode}' / 'result.json'
            result = json.loads(path.read_text()) if path.exists() else {}
            info = result.get('continuation', {})
            cells.append(dict(kind=kind, group=group, method=method, mode=mode,
                action=info.get('action', 'new_validation' if result else 'missing'), parent_record=info.get('parent_record')))
    records = pd.DataFrame(rows); trace_data = pd.DataFrame(traces); cell_data = pd.DataFrame(cells)
    records.to_csv(out / 'continuation_training.csv', index=False)
    trace_data.to_csv(out / 'continuation_traces.csv', index=False)
    cell_data.to_csv(out / 'continuation_cells.csv', index=False)
    body = '\n## 预算热启动与继承证据\n\n'
    body += ('保留原始起点谱系不等于从原始0.5/1/2重新初始化：eligible起点使用父运行最后有效driver、初态与β，'
        'LM阻尼重新设定。因此这是warm restart，不是exact optimizer resume；即使预算总量相同也不保证等于一次不中断优化。'
        '所有原身份、失败起点与成功阈值保留，未以热启动后较好的结果增加独立样本数。\n\n')
    action_labels = {'inherited_training':'继承训练', 'restarted_training':'热启动训练',
                     'inherited_cell':'继承验证', 'inherited_validation':'继承验证', 'new_validation':'重算验证'}
    for kind, part in records.groupby('kind', sort=False):
        display = []
        for (method, action), subset in part.groupby(['method','action'], sort=False):
            display.append({'臂':labels[method].split()[0], '来源':action_labels.get(action,action),
                '起点数':len(subset), '收敛成功':int((subset.status.eq('completed') & subset.converged).sum()),
                '新增优化调用':int(subset.new_optimization_evaluations.sum()) if subset.new_optimization_evaluations.notna().all() else '记录不完整',
                '累计优化调用':int(subset.cumulative_optimization_evaluations.sum()) if subset.cumulative_optimization_evaluations.notna().all() else '记录不完整'})
        body += f'### {kind} · 训练来源与调用数\n\n' + table(pd.DataFrame(display)) + '\n\n'
    counts = cell_data.groupby(['kind','action']).size().rename('cells').reset_index()
    body += table(counts.assign(action=counts.action.map(lambda x:action_labels.get(x,x))).rename(
        columns={'kind':'数据','action':'验证来源','cells':'cell数'})) + '\n\n'
    body += ('优化调用按单trial forward计数；CSV另列父/新增/累计的库reported total evaluations，'
        '该total=optimization+库内4子步独立核验（完整组18条）；runner随后另做18条8子步fine replay，不计入该字段。CSV用library_evaluations命名，另列new_fine_replays，绝不能将库调用数称为整个workflow的全部forward调用或墙钟时间。'
        '继承条目新增库调用和新增fine replay均为0，累计库数保持父值；这不抹去父阶段已发生的库调用与细积分历史成本，也不表示它们按本轮3600/90重新运行。\n\n')
    restarted = records[records.action.eq('restarted_training')]
    if len(restarted):
        continuity = restarted.first_objective_difference.abs().dropna()
        passed = int(restarted.objective_continuity.eq(True).sum())
        body += (f'热启动记录{len(restarted)}/{config["continuation"]["expected_measured_restarts"]}；'
            f'首目标与父末目标连续性通过{passed}/{len(restarted)}，可计算差值{len(continuity)}/{len(restarted)}，'
            f'绝对差中位{continuity.median():.3g}、最大{continuity.max():.3g}。'
            '通过标志及容差取自runner记录；首目标由trace首点独立读取，失败或缺失不补作通过。\n\n')
        domain_seen = restarted.new_domain_rejections.gt(0)
        domain_known = restarted.new_domain_rejections.notna()
        domain_failed = domain_seen & ~(restarted.status.eq('completed') & restarted.converged)
        body += (f'父运行“无域/导数拒绝”只用于热启动资格，不保证后续轨迹无拒绝。'
            f'本热启动段域拒绝记录可用{int(domain_known.sum())}/{len(restarted)}，'
            f'出现域拒绝{int(domain_seen.sum())}条，其中最终未成功{int(domain_failed.sum())}条；'
            '逐步计数保留在continuation_traces.csv，不能将父阶段无域拒绝标签沿用为本段结果。\n\n')
        short = lambda g: str(g).replace('subject_', 'S')
        body += table(pd.DataFrame([{'组':short(r.group), '臂':labels[r.method].split()[0],
            '原起点':r.original_start_gain, '首目标−父末':f'{r.first_objective_difference:.3g}',
            '连续':str(r.objective_continuity), '新增/累计优化':f'{r.new_optimization_evaluations}/{r.cumulative_optimization_evaluations}',
            '终态':r.status.replace('completed','成功').replace('failed_numerical','数值失败')}
            for r in restarted.itertuples()])) + '\n\n'
    if len(trace_data):
        fig, axes = plt.subplots(3, 2, figsize=(12, 9), sharex='col')
        for column, method in enumerate(['trained_gain_no_amplitude','trained_gain_amplitude']):
            part = trace_data[trace_data.method.eq(method)]
            groups = list(part.group.unique())
            colors = plt.get_cmap('tab20')
            for group_index, group in enumerate(groups):
                for start, points in part[part.group.eq(group)].groupby('start', sort=True):
                    points = points.sort_values(['evaluations','iteration'], kind='stable')
                    x = points.evaluations.to_numpy()
                    values = [points.objective-points.objective.iloc[0], points.projected_scaled_gradient_inf_norm, points.parameter_value]
                    for ax, value in zip(axes[:,column], values):
                        ax.plot(x, value, color=colors(group_index % 20), ls=['-', '--', ':'][int(start)%3],
                            lw=.9, alpha=.8, label=group.replace('subject_','S') if start == part[part.group.eq(group)].start.min() else None)
            axes[0,column].set_title(labels[method])
            axes[0,column].set_ylabel('目标 − 热启动首目标')
            axes[1,column].set_ylabel('projected scaled gradient')
            axes[1,column].set_yscale('log')
            axes[1,column].axhline(config['gain_training']['gradient_tolerance'], color='red', ls=':', lw=1)
            axes[2,column].set_ylabel('β'); axes[2,column].set_xlabel('本热启动段优化trial forward调用')
            if groups: axes[0,column].legend(fontsize=5, ncol=2)
        save(fig, 'gain_continuation_traces', '每条线是一条热启动谱系：颜色区分训练组，线型区分原起点。上行目标仅减去同条首目标以显示变化；中行保留实际梯度及原收敛阈值，末行为β。横轴为新增优化调用，父/累计调用另见表与CSV；继承谱系没有新trace，不伪造轨迹。')
    else:
        body += '本次记录没有可用热启动trace，不能展示优化轨迹。\n\n'
    return body


def gain_prior_details(run, out, config, data, states, labels, save, selected):
    """Read declared training starts only; no parameter estimation or winner search."""
    candidates, selections = [], []
    training = config["gain_training"]
    for group in data.group.unique():
        prep = json.loads((run / "prepared" / f"{group}.json").read_text())
        spec = prep["spec"]
        for amplitude in (0, 1):
            method = "trained_gain_amplitude" if amplitude else "trained_gain_no_amplitude"
            folder = run / "training" / f"{group}__amplitude{amplitude}"
            selection_path = folder / "selection.json"
            selection = json.loads(selection_path.read_text()) if selection_path.exists() else {}
            chosen = selection.get("selected_training_start")
            group_rows = []
            for index, start_gain in enumerate(training["start_values"]):
                path = folder / f"start_{index}" / "result.json"
                result = json.loads(path.read_text()) if path.exists() else {}
                info = result.get("information", {})
                attempts = result.get("starts", [])
                attempt = attempts[0] if len(attempts) == 1 else {}
                row = dict(kind=spec["kind"], group=group, method=method, start=index,
                    start_gain=start_gain, true_gain=spec.get("true_gain"),
                    selected=index == chosen, status=result.get("status", "missing"),
                    converged=bool(result.get("converged", False)),
                    gain=result.get("parameter_value"), objective=result.get("objective"),
                    boundary=result.get("boundary_status", "N/A"),
                    convergence_reason=result.get("convergence_reason", attempt.get("convergence_reason")),
                    projected_gradient=result.get("projected_scaled_gradient_inf_norm", attempt.get("projected_scaled_gradient_inf_norm")),
                    domain_rejections=attempt.get("domain_rejections"),
                    nondecreasing_rejections=attempt.get("nondecreasing_rejections"),
                    derivative_rejections=attempt.get("derivative_rejections"),
                    evaluations=result.get("optimization_evaluations"),
                    weighted_data_sse=result.get("weighted_data_sse"),
                    driver_amplitude_cost=result.get("driver_amplitude_cost"),
                    parameter_prior_cost=result.get("parameter_prior_cost"),
                    initial_state_cost=result.get("initial_state_penalty_cost"),
                    data_information=info.get("data_only", {}).get("value"),
                    regularized_information=info.get("nuisance_regularized", {}).get("value"),
                    prior_information=info.get("with_parameter_prior"))
                candidates.append(row); group_rows.append(row)
            chosen_rows = [r for r in group_rows if r["selected"]]
            if selection.get("status") == "completed" and len(chosen_rows) != 1:
                raise ValueError("Completed gain selection has no declared selected start")
            chosen_row = chosen_rows[0] if chosen_rows else {}
            completed = [r for r in group_rows if r["status"] == "completed" and r["converged"]]
            if chosen_row and (chosen_row["status"] != "completed" or not chosen_row["converged"]):
                raise ValueError("Gain selection points to nonconverged training start")
            selections.append(dict(kind=spec["kind"], group=group, method=method,
                status=selection.get("status", "missing"), gain=selection.get("parameter_value"),
                true_gain=spec.get("true_gain"), selected_start=chosen,
                completed_starts=len(completed), expected_starts=len(training["start_values"]),
                converged_gain_min=min([r["gain"] for r in completed], default=np.nan),
                converged_gain_max=max([r["gain"] for r in completed], default=np.nan),
                boundary=chosen_row.get("boundary", "N/A"),
                data_information=chosen_row.get("data_information"),
                regularized_information=chosen_row.get("regularized_information"),
                prior_information=chosen_row.get("prior_information")))
    candidates = pd.DataFrame(candidates); selections = pd.DataFrame(selections)
    candidates.to_csv(out / "gain_training_candidates.csv", index=False)
    selections.to_csv(out / "gain_training_selection.csv", index=False)
    short = lambda value: str(value).replace("subject_", "S").replace("synthetic_gain", "G").replace("mixed", "混").replace("slow", "慢")
    body = "\n## 有效增益训练与工程先验\n\n"
    body += (f"每个训练组和可训练臂有{len(training['start_values'])}个独立起点；只在全部起点终态后按收敛目标选择。"
        f"β范围{training['bounds']}、log prior SD={training['prior_log_sd']:.4g}是工程敏感性设置，不是正常生理范围。"
        "幅度先验的σr直接取既有冻结训练EEG normalizer/loading，未重新去均值估计尺度。"
        "其中心为零但不强制driver样本均值为零，也不代表事件前五秒是静息。"
        "下列局部信息在logβ坐标，data-only先消除driver和初态方向；regularized另包含这些变量的惩罚。"
        "加β先验的信息单列，不能将正则带来的信息解释成数据可辨识性或置信区间。\n\n")
    # Synthetic truth permits an explicit bias diagnostic; never infer physical calibration.
    expected_synthetic = len(config["synthetic"]["gains"])*len(config["synthetic"]["spectra"])*config["synthetic"]["replicates"]
    synthetic = selections[selections.kind.eq("synthetic")].copy()
    synthetic["relative_gain_error"] = pd.to_numeric(synthetic.gain, errors="coerce") / pd.to_numeric(synthetic.true_gain, errors="coerce") - 1.
    gain_recovery = []
    for method in ("trained_gain_no_amplitude", "trained_gain_amplitude"):
        part = synthetic[synthetic.method.eq(method)]
        good = part[part.status.eq("completed") & np.isfinite(part.relative_gain_error)]
        gain_recovery.append(dict(method=method, expected=expected_synthetic, terminal=int(part.status.isin(["completed", "failed_training"]).sum()),
            completed=len(good), relative_error_minimum=good.relative_gain_error.min(),
            relative_error_median=good.relative_gain_error.median(), relative_error_maximum=good.relative_gain_error.max()))
    pd.DataFrame(gain_recovery).to_csv(out / "synthetic_gain_recovery.csv", index=False)
    if all(r["terminal"] == expected_synthetic for r in gain_recovery):
        body += "### 合成真值偏差诊断\n\n"
        for r in gain_recovery:
            if not r["completed"]:
                body += f"{labels[r['method']]}：成功训练组0/{r['expected']}，β误差不可估计。"
                continue
            body += (f"{labels[r['method']]}：成功训练组{r['completed']}/{r['expected']}；"
                f"β相对真值误差(估计/真值−1)中位{r['relative_error_median']:+.1%}，"
                f"范围[{r['relative_error_minimum']:+.1%}, {r['relative_error_maximum']:+.1%}]。")
        no_amp = synthetic[synthetic.method.eq("trained_gain_no_amplitude") & synthetic.status.eq("completed")]
        with_amp = synthetic[synthetic.method.eq("trained_gain_amplitude") & synthetic.status.eq("completed")]
        matched = with_amp.merge(no_amp, on="group", suffixes=("_with", "_without"), validate="one_to_one")
        if len(matched):
            delta = matched.relative_gain_error_with - matched.relative_gain_error_without
            direction = "全部上移" if delta.gt(0).all() else "全部下移" if delta.lt(0).all() else "方向不一致"
            body += (f"同组配对{len(matched)}/{expected_synthetic}：加入幅度先验后β偏差{direction}，"
                f"误差变化中位{delta.median()*100:+.1f}个百分点。")
            if delta.gt(0).all():
                body += ("这项负结果表明幅度先验下的有效增益恢复出现系统上偏。"
                    "幅度惩罚压低r而提高β可部分维持血管输入β×r，是与该结果一致、受结果支持的补偿解释；"
                    "仅凭这些拟合不能证明它是唯一机制，也不能把更平稳的状态或更大的正则信息当作无偏生理参数恢复。")
        body += "这仍是给定生成模型与观测坐标的合成诊断，不是独立标定的被试生理估计；失败训练组不进入误差统计，但保留完整分母。\n\n"
    else:
        body += "合成训练选择尚未全部终态，暂不生成完整面板的β真值偏差结论。\n\n"
    steps = config["tensor"]["steps"]
    dt = config["tensor"]["dt_s"]
    amplitude = config["regularization"]["driver_amplitude_weight"]
    baseline_samples = int(round(5. / dt))
    special_squared = steps / baseline_samples
    ordinary_shrink = 1. / (1. + amplitude*dt)
    special_shrink = special_squared / (special_squared + amplitude*dt)
    body += ("简化算子解释：仅保留全观测EEG数据项与幅度ridge，忽略Hb、曲率和初态耦合时，"
        f"本合同基线算子B=I−1wᵀ使用首{baseline_samples}/{steps}点；"
        f"其奇异值为{steps-2}个1、1个√{special_squared:g}和1个0。"
        "由于driver prior SD来自同一EEG训练尺度/loading，各可观测奇异模态的收缩因子为"
        "σ²/(σ²+λ·dt)，这里σ指B的奇异值而非噪声SD。"
        f"本合同λ={amplitude:g}、dt={dt:g}秒，对应普通模态{ordinary_shrink:.3g}、"
        f"特殊模态{special_shrink:.3g}，DC由零中心ridge锚定为0。"
        f"若主要血管输入βr近似保持，普通模态的收缩可伴随β约放大{1./ordinary_shrink:.3g}倍，"
        "与上述合成偏差方向及量级相容；这是简化机制解释，非完整联合模型的精确预测或唯一原因。"
        "实际数据权重使用训练信号SD而非噪声SD，且未对时间协方差白化，故本目标应称标准化数据误差加工程ridge，"
        "不能将其当作校准likelihood，也不能据此断言生理先验本身失效。\n\n")
    for (kind, method), part in selections.groupby(["kind", "method"], sort=False):
        body += f"### {kind} · {labels[method]}\n\n"
        display = pd.DataFrame([{"组":short(r.group), "真β":r.true_gain, "选择β":r.gain,
            "成功起点":f"{r.completed_starts}/{r.expected_starts}", "边界":r.boundary,
            "起点β跨度":f"{r.converged_gain_min:.3g}–{r.converged_gain_max:.3g}"}
            for r in part.itertuples()])
        body += table(display) + "\n\n"
        information = pd.DataFrame([{"组":short(r.group), "数据":r.data_information,
            "加nuisance正则":r.regularized_information, "再加β先验":r.prior_information}
            for r in part.itertuples()])
        body += table(information) + "\n\n"
    body += ("全部起点的目标、β、收敛、边界、梯度、数据SSE和惩罚分解保留在gain_training_candidates.csv；"
        "同组起点跨度仅统计成功起点，失败数没有删除。完整库结果仍由training/start_*/result.json持有。\n\n")
    # Compact candidate tables retain all starts, split by group to avoid wide PDFs.
    for kind, part in candidates.groupby("kind", sort=False):
        body += f"### {kind} · 全部起点目标\n\n"
        display = []
        for (group, method), starts in part.groupby(["group", "method"], sort=False):
            row = {"组":short(group), "臂":labels[method].split()[0]}
            for r in starts.itertuples():
                objective = f"{r.objective:.4g}" if pd.notna(r.objective) else "—"
                status = "成功" if r.status == "completed" and r.converged else "失败"
                row[f"{'原起点' if config.get('continuation') else '起点'}{r.start_gain:g}"] = objective + "/" + status + ("*" if r.selected else "")
            display.append(row)
        body += table(pd.DataFrame(display)) + "\n\n"
    body += "星号标记所选起点；失败起点目标仅为审计信息，不能参与选择。\n\n"
    failed_starts = candidates[~(candidates.status.eq("completed") & candidates.converged)].copy()
    failed_starts["domain_category"] = np.select(
        [failed_starts.domain_rejections.eq(0), failed_starts.domain_rejections.gt(0)],
        ["无域拒绝", "有域拒绝"], default="域记录缺失")
    failure_diagnostics = []
    for (kind, method, category), part in failed_starts.groupby(["kind", "method", "domain_category"], sort=False):
        gradients = pd.to_numeric(part.projected_gradient, errors="coerce").dropna()
        failure_diagnostics.append(dict(kind=kind, method=method, domain_category=category,
            failed_starts=len(part), evaluation_budget=int(part.convergence_reason.eq("evaluation_budget").sum()),
            gradient_scored=len(gradients), gradient_minimum=gradients.min(),
            gradient_median=gradients.median(), gradient_maximum=gradients.max(),
            convergence_threshold=training["gradient_tolerance"]))
    failure_diagnostics = pd.DataFrame(failure_diagnostics)
    failure_diagnostics.to_csv(out / "gain_training_failures.csv", index=False)
    body += "### 训练起点失败的分类\n\n"
    if len(failure_diagnostics):
        for kind, part in failure_diagnostics.groupby("kind", sort=False):
            body += f"{kind}：\n\n" + table(pd.DataFrame([{
                "臂":labels[r.method].split()[0], "域拒绝":r.domain_category,
                "失败数":r.failed_starts, "预算耗尽":r.evaluation_budget,
                "梯度数":r.gradient_scored, "梯度中位":f"{r.gradient_median:.3g}",
                "梯度范围":f"{r.gradient_minimum:.3g}–{r.gradient_maximum:.3g}"}
                for r in part.itertuples()])) + "\n\n"
    else:
        body += "没有失败的训练起点。\n\n"
    batch = 18
    ordinary_steps = max(0, (training["max_evaluations"] - batch) // (2*batch))
    body += (f"梯度为库返回的projected_scaled_gradient_inf_norm，收敛阈值{training['gradient_tolerance']:.3g}；"
        "无梯度或域记录的起点不补零。domain_rejections读取每个独立任务内部starts[0]，表示优化提议曾被拒绝，"
        "不等于最终可行轨迹生理不合理；没有域拒绝也不证明已接近收敛。"
        f"本合同每起点max_evaluations={training['max_evaluations']}，按每组{batch}个trial计forward调用；"
        f"初始求值消耗{batch}次，若每个正常GN步只有一次完整提议与Jacobian复算，则每步{2*batch}次，"
        f"最多{ordinary_steps}个这样的步。max_iterations={training['max_iterations']}只是另一个上限，"
        "不保证能执行到该数；拒绝和额外线搜索会消耗预算。evaluation_budget不能一概解释成域失败或近收敛。\n\n")
    fig, axes = plt.subplots(1, 2, figsize=(12, 4))
    for ax, kind in zip(axes, ["synthetic", "measured"]):
        part = selections[selections.kind.eq(kind)]
        groups = list(part.group.unique())
        for method in part.method.unique():
            values = part[part.method.eq(method)].set_index("group").reindex(groups)
            ax.plot(range(len(groups)), pd.to_numeric(values.gain, errors="coerce"), "o-", label=labels[method])
        if kind == "synthetic" and len(part):
            actual = part.drop_duplicates("group").set_index("group").reindex(groups)
            ax.plot(range(len(groups)), pd.to_numeric(actual.true_gain, errors="coerce"), "k--", label="真β")
        ax.set_xticks(range(len(groups)), [short(g) for g in groups], rotation=55, ha="right", fontsize=6)
        ax.set_ylabel("训练冻结 β"); ax.set_title(kind)
        if len(part): ax.legend(fontsize=7)
    save(fig, "gain_training", "每组训练冻结的有效β；缺失/失败不补值。合成β为预设测试点，非人口生理范围。")
    body += "## 驱动与缺口状态诊断\n\n"
    diagnostics = []
    for (kind, method, mode), part in states.groupby(["kind", "method", "mode"], sort=False):
        good = part[part.converged & part.finite_states]
        denominator = len(data[(data.kind == kind) & (data.method == method) & (data['mode'] == mode)])
        row = dict(kind=kind, method=method, mode=mode, successful=len(good), observed=denominator)
        for field in ["driver_mean", "driver_rms", "driver_center_rms", "driver_difference_rms", "driver_center_difference_rms"]:
            values = pd.to_numeric(good.get(field, pd.Series(dtype=float)), errors="coerce")
            row[field+"_median"] = values.median(); row[field+"_maximum"] = values.max()
        for name in ["r", "s", "f", "v", "p", "q"]:
            for suffix, operation in [("minimum", "min"), ("maximum", "max"), ("center_minimum", "min"), ("center_maximum", "max")]:
                values = pd.to_numeric(good.get(name+"_"+suffix, pd.Series(dtype=float)), errors="coerce")
                row[name+"_"+suffix] = getattr(values, operation)()
        diagnostics.append(row)
    diagnostics = pd.DataFrame(diagnostics)
    diagnostics.to_csv(out / "gain_driver_state_diagnostics.csv", index=False)
    for (kind, mode), part in diagnostics.groupby(["kind", "mode"], sort=False):
        body += f"### {kind} · {mode}\n\n"
        body += table(pd.DataFrame([{"臂":labels[r.method].split()[0], "成功/记录":f"{r.successful}/{r.observed}",
            "r RMS中位":r.driver_rms_median, "缺口r RMS最大":r.driver_center_rms_maximum,
            "缺口Δr RMS最大":r.driver_center_difference_rms_maximum,
            "全路径f最小":r.f_minimum, "全路径f最大":r.f_maximum} for r in part.itertuples()])) + "\n\n"
    body += ("这里的Δr是相邻采样差分，RMS不是导数也不是生理正常范围；高差分RMS提示残余振荡。"
        "中心16点均以相同位置统计；full模式的中心统计不表示数据被隐藏。六状态全路径/中心极值和失败最后有效轨迹详见state_ranges.csv，"
        "成功子集汇总见gain_driver_state_diagnostics.csv。\n\n")
    # Factorial contrasts require the same four successful identities.
    identity = ["kind", "group", "mode", "trial", "subject", "session"]
    common = None
    for letter, method in zip("ABCD", config["methods"]):
        part = data[(data.method == method) & data.status.eq("completed") & truth(data.converged)]
        fields = ["nmse_"+c for c in COMPONENTS]
        part = part.reindex(columns=identity+fields).rename(columns={f:f+"_"+letter for f in fields})
        common = part if common is None else common.merge(part, on=identity, validate="one_to_one")
    factorial = []
    for (kind, mode), part in common.groupby(["kind", "mode"], sort=False):
        expected = len(data[(data.kind == kind) & (data['mode'] == mode) & (data.method == config['methods'][0])])
        for component in COMPONENTS:
            f = "nmse_"+component
            if part[f+"_A"].isna().all(): continue
            for contrast, values in [("B−A", part[f+"_B"]-part[f+"_A"]),
                ("C−A", part[f+"_C"]-part[f+"_A"]), ("D−B", part[f+"_D"]-part[f+"_B"]),
                ("D−C", part[f+"_D"]-part[f+"_C"]),
                ("交互 D−C−B+A", part[f+"_D"]-part[f+"_C"]-part[f+"_B"]+part[f+"_A"])]:
                frame = part[['subject','session']].assign(delta=values)
                value = frame.groupby(['subject','session']).delta.mean().groupby('subject').mean().mean()
                factorial.append(dict(kind=kind, mode=mode, component=component, contrast=contrast,
                    common_success=len(part), observed_reference=expected, delta_nmse=value))
    factorial = pd.DataFrame(factorial)
    factorial.to_csv(out / 'paired_factorial.csv', index=False)
    body += "## 四臂同身份比较\n\n以下全部使用四臂共同成功的相同身份；ΔNMSE按session及subject等权，负值表示该对比误差下降。B−A隔离固定β下幅度先验，D−B隔离有幅度先验时训练β，交互不能由各臂不同成功子集推断。无四臂共同成功的分组没有可估计对比；其失败仍保留在前文完整分母。\n\n"
    if len(factorial):
        for (kind, mode), part in factorial.groupby(['kind','mode'], sort=False):
            body += f"### {kind} · {mode}\n\n" + table(part[['component','contrast','common_success','observed_reference','delta_nmse']].rename(
                columns={'component':'分量','contrast':'比较','common_success':'共同成功','observed_reference':'A记录数','delta_nmse':'ΔNMSE'})) + "\n\n"
    # Hidden EEG states on the exact identities chosen from full-observation A.
    fig, axes = plt.subplots(len(selected), 3, figsize=(12, 3.6*len(selected)), squeeze=False)
    for i, (name, sample) in enumerate(selected):
        group, trial = sample.group, int(sample.trial)
        n = config['tensor']['steps']
        t = np.arange(n)*config['tensor']['dt_s'] - (5 if sample.kind == 'measured' else 0)
        target_drawn = False
        unavailable_notes = []
        for method in config["methods"]:
            path = run / "cells" / f"{group}__{method}__center_EEG" / "trajectories.npz"
            rows = data[(data.group == group) & (data.method == method) & (data['mode'] == "center_EEG") & (data.trial == trial)]
            status = rows.status.iloc[0] if len(rows) else "missing"
            if not path.exists():
                unavailable_notes.append(f"{labels[method].split()[0]}: {status} / 未估计")
                for ax in axes[i]: ax.plot([], [], ":", label=labels[method]+" 未估计")
                continue
            with np.load(path, allow_pickle=False) as arrays:
                matches = np.flatnonzero(arrays["trial_indices"] == trial)
                if not len(matches): raise ValueError("Fixed hidden example missing")
                j = int(matches[0]); n = arrays['prediction'].shape[1]
                t = np.arange(n)*config['tensor']['dt_s'] - (5 if sample.kind == 'measured' else 0)
                style = '-' if status == 'completed' else ':'
                if not target_drawn and 'target' in arrays:
                    axes[i,0].plot(t, arrays['target'][j,:,0], 'k', lw=1.3, label='目标')
                    target_drawn = True
                if not np.isfinite(arrays['states'][j]).any():
                    unavailable_notes.append(f"{labels[method].split()[0]}: {status} / 未估计")
                for ax, value in zip(axes[i], [arrays['prediction'][j,:,0], arrays['states'][j,:,0], arrays['states'][j,:,2]]):
                    suffix = ' 未估计' if not np.isfinite(value).any() else (' 失败迭代' if status != 'completed' else '')
                    ax.plot(t, value, style, lw=.9, label=labels[method]+suffix)
        if unavailable_notes:
            axes[i,1].text(.02, .98, "\n".join(unavailable_notes), transform=axes[i,1].transAxes,
                ha='left', va='top', fontsize=6, bbox=dict(facecolor='white', alpha=.85, edgecolor='none'))
        for ax, ylabel in zip(axes[i], ['EEG', 'driver r', 'flow f']):
            ax.axvspan(t[(n-16)//2], t[(n+16)//2-1], color='grey', alpha=.12)
            ax.set_ylabel(ylabel); ax.set_xlabel('相对事件 / 秒' if sample.kind == 'measured' else '秒')
        axes[i,0].set_title(f'{name}: {short(group)} / trial {trial}')
        axes[i,0].legend(fontsize=6)
    save(fig, 'gain_hidden_fixed_examples', '沿用full模式A基线选出的中位/最差身份，显示中心EEG隐藏时预测、驱动和血流；灰区为隐藏段，失败最后有效轨迹为虚线，未按候选结果重挑。')
    # These identities are fixed in advance, independent of reconstruction scores.
    truth_groups = ["synthetic_gain1_slow_r0", "synthetic_gain1_mixed_r0"]
    truth_methods = ["trained_gain_no_amplitude", "trained_gain_amplitude"]
    fig, axes = plt.subplots(3, 2, figsize=(11, 8), sharex="col")
    truth_notes = []
    for column, group in enumerate(truth_groups):
        trial = 0
        prepared_path = run / "prepared" / f"{group}.npz"
        axes[0,column].set_title(f"{group} / trial {trial}", fontsize=9)
        if not prepared_path.exists():
            for ax in axes[:,column]:
                ax.text(.5, .5, "准备数据缺失；固定身份不替换", ha="center", va="center", transform=ax.transAxes)
            truth_notes.append(f"{group}: prepared缺失")
            continue
        with np.load(prepared_path, allow_pickle=False) as prepared:
            n = prepared["target"].shape[1]
            t = np.arange(n)*config["tensor"]["dt_s"]
            axes[0,column].plot(t, prepared["target"][trial,:,0], "k", lw=1.15, label="含噪EEG目标")
            axes[0,column].plot(t, prepared["clean"][trial,:,0], color=".55", ls=":", lw=1., label="无噪处理EEG")
            axes[1,column].plot(t, prepared["truth"][trial], "k", lw=1.4, label="driver真值")
            axes[2,column].plot(t, prepared["truth_states"][trial,:,2], "k", lw=1.4, label="flow真值")
        for method, color in zip(truth_methods, ["tab:blue", "tab:orange"]):
            rows = data[(data.group == group) & (data.method == method) & (data['mode'] == "full") & (data.trial == trial)]
            status = rows.status.iloc[0] if len(rows) else "missing"
            truth_notes.append(f"{group} / {labels[method]}: {status}")
            path = run / "cells" / f"{group}__{method}__full" / "trajectories.npz"
            if not path.exists():
                for ax in axes[:,column]:
                    ax.plot([], [], color=color, ls=":", label=labels[method]+" 无轨迹")
                continue
            with np.load(path, allow_pickle=False) as arrays:
                matches = np.flatnonzero(arrays["trial_indices"] == trial)
                if not len(matches):
                    raise ValueError("Prespecified synthetic truth example missing trial 0")
                index = int(matches[0])
                values = [arrays["prediction"][index,:,0], arrays["states"][index,:,0], arrays["states"][index,:,2]]
                style = "-" if status == "completed" else ":"
                for ax, value in zip(axes[:,column], values):
                    suffix = " 未估计" if not np.isfinite(value).any() else ("" if status == "completed" else " 失败最后有效轨迹")
                    ax.plot(t, value, color=color, ls=style, lw=1., label=labels[method]+suffix)
        for ax, ylabel in zip(axes[:,column], ["处理坐标 EEG", "driver r", "flow f"]):
            ax.set_ylabel(ylabel); ax.legend(fontsize=6)
        axes[2,column].set_xlabel("秒")
    save(fig, "gain_synthetic_truth_examples", "事先固定β真值1、replicate0的slow/mixed条件，各trial0；首行黑线为含噪处理EEG目标，灰虚线为无噪处理EEG；后二行黑线为无噪driver/flow真值。仅叠加C/D全观测结果，失败最后有效轨迹为虚线，未做去均值、幅度或时间对齐。")
    body += ("## 固定合成真轨迹样例\n\n选样身份固定为synthetic_gain1_slow_r0和synthetic_gain1_mixed_r0，各trial0，"
        "与误差排序或候选效果无关。EEG同时显示含噪观测目标及无噪处理轨迹，r/f显示生成真值；仅对比C/D。"
        "没有重新拟合或事后去均值、幅度、时间对齐。失败、缺失保留原身份，不以较好样例替换。\n\n"
        + "；".join(truth_notes) + "。\n\n")
    if config.get("continuation"):
        body = body.replace("个独立起点；只在全部起点终态后按收敛目标选择。", "条原起点谱系；其中可能继承或从父解热启动，只在全部谱系终态后按同一收敛目标选择。")
        body = body.replace("本合同每起点max_evaluations=", "本轮仅热启动段每起点max_evaluations=")
        body += "上述新预算不适用于继承的终态起点；继承记录保留父运行的实际预算与调用数。合成偏差与固定合成曲线沿用父证据，不能当作本轮独立复现。\n\n"
        body += gain_continuation_details(run, out, config, data, labels, save)
        body += continuation_parent_comparison(run, out, config, data, labels)
    return body


def fixed_roi_training_failure_details(run, out, config, candidates, save):
    """Map local training failures to original trials; never score last-valid fits."""
    rows=[]
    failed=candidates[~(candidates.status.eq('completed')&candidates.converged)]
    for item in failed.itertuples():
        folder=run/'training'/f'{item.group}__{item.method}'/f'start_{item.start}'
        path=folder/'result.json'
        if not path.exists():continue
        result=json.loads(path.read_text());attempt=(result.get('starts') or [{}])[0]
        training=result.get('training_trials',[]);events=attempt.get('evaluation_failures',[])
        by_local={}
        for event in events:
            local=int(event['trial'])
            if not 0<=local<len(training):raise ValueError('Training failure lacks local-to-global trial mapping')
            by_local.setdefault(local,[]).append(event)
        arrays_path=folder/'trajectories.npz'
        saved_states=saved_indices=None
        if arrays_path.exists():
            with np.load(arrays_path,allow_pickle=False) as arrays:
                saved_states=arrays['states'].copy();saved_indices=arrays['trial_indices'].copy()
            if training and not np.array_equal(saved_indices,np.asarray(training)):
                raise ValueError('Failed training trajectory trial order differs from training_trials')
        for local,trial in enumerate(training):
            ev=by_local.get(local,[])
            row=dict(kind=item.kind,group=item.group,subject=item.subject,method=item.method,start=item.start,
                status=item.status,local_trial=local,global_trial=int(trial),evaluation_failures=len(ev),
                forward_failures=sum(not e.get('derivative',False) for e in ev),
                derivative_failures=sum(bool(e.get('derivative',False)) for e in ev),
                errors='; '.join(sorted(set(e.get('error','') for e in ev))),
                start_domain_rejections=attempt.get('domain_rejections'),trajectory=str(arrays_path),
                finite_last_state=False,final_flow_min=np.nan,minimum_time_from_window_start_s=np.nan,
                minimum_time_relative_event_s=np.nan)
            if saved_states is not None and np.isfinite(saved_states[local]).all():
                flow=saved_states[local,:,2];index=int(np.argmin(flow));time=index*config['tensor']['dt_s']
                row.update(finite_last_state=True,final_flow_min=float(flow[index]),minimum_time_from_window_start_s=time,
                    minimum_time_relative_event_s=time-(5. if item.kind=='measured' else 0.))
            rows.append(row)
    frame=pd.DataFrame(rows)
    frame.to_csv(out/'failure_by_trial.csv',index=False)
    body='\n## 失败训练的试次定位\n\n'
    body+=('evaluation_failures中的trial是训练集合局部索引，本表经该起点training_trials映射为原始global trial；'
        '不能把局部序号直接当作原试次。统计只覆盖失败训练起点；forward求值异常与derivative异常分列，'
        '重复求值次数不是独立失败试次数。末有效轨迹不是有效参数估计，不计入验证误差或合成恢复主分数。\n\n')
    if not len(frame):return body+'没有可映射的失败训练轨迹。\n\n'
    counted=frame[frame.evaluation_failures.gt(0)]
    if len(counted):
        summary=[]
        for (subject,trial),part in counted.groupby(['subject','global_trial'],sort=False):
            summary.append({'被试':str(subject).replace('subject_','S'),'原trial':trial,'涉及失败起点':len(part),
                'forward失败':int(part.forward_failures.sum()),'derivative失败':int(part.derivative_failures.sum()),
                '末有效f最小':f'{part.final_flow_min.min():.3g}'})
        body+=table(pd.DataFrame(summary))+'\n\n'
    else:body+='这些失败起点没有记录trial级求值异常；不能据此归因为某个试次域失败。\n\n'
    measured=frame[frame.kind.eq('measured')&frame.finite_last_state].sort_values(
        ['subject','final_flow_min','group','method','start','global_trial'],kind='stable')
    selected=measured.groupby('subject',sort=False).head(1)
    selected.to_csv(out/'failed_training_selected_examples.csv',index=False)
    if len(selected):
        fig,axes=plt.subplots(3,len(selected),figsize=(4.2*len(selected),8),squeeze=False,sharex='col')
        for column,item in enumerate(selected.itertuples()):
            with np.load(item.trajectory,allow_pickle=False) as arrays:
                matches=np.flatnonzero(arrays['trial_indices']==item.global_trial)
                if len(matches)!=1:raise ValueError('Failed training selected global trial is not unique')
                index=int(matches[0]);t=np.arange(arrays['states'].shape[1])*config['tensor']['dt_s']-5.
                for row,component in enumerate([1,2]):
                    axes[row,column].plot(t,arrays['target'][index,:,component],'k',lw=1.2,label='训练目标')
                    axes[row,column].plot(t,arrays['prediction'][index,:,component],color='tab:red',ls='--',lw=1.,label='失败末有效预测')
                    axes[row,column].set_ylabel(['HbO','HbR'][row]);axes[row,column].legend(fontsize=6)
                flow=arrays['states'][index,:,2]
                axes[2,column].plot(t,flow,color='tab:red',ls='--',lw=1.,label='失败末有效flow')
                axes[2,column].axhline(0,color='grey',lw=.7)
                axes[2,column].axvline(item.minimum_time_relative_event_s,color='grey',ls=':',lw=.7)
                axes[2,column].set_ylabel('flow f');axes[2,column].set_xlabel('相对事件 / 秒')
                axes[2,column].text(.03,.05,f'min f={item.final_flow_min:.3g}',transform=axes[2,column].transAxes,fontsize=8)
                axes[0,column].set_title(f"{str(item.subject).replace('subject_','S')} / trial{item.global_trial}\n{item.group} / {'先验τ' if item.method.endswith('prior_tau') else '训练τ'} / start{item.start}",fontsize=8)
                for ax in axes[:,column]:ax.axvline(0,color='grey',ls=':',lw=.6)
        save(fig,'fixed_roi_failed_training_states','每被试从全部失败训练起点的末有效轨迹中，按最小flow固定选一条（相同最小值按身份排序）。这里只展示病态失败迭代，不是有效估计，不计主评分，也未替换成功样例；灰竖线定位最小flow。')
        body+=table(pd.DataFrame([{'被试':str(r.subject).replace('subject_','S'),'组':str(r.group).replace('subject_','S'),
            '原trial':r.global_trial,'起点':r.start,'最小f':f'{r.final_flow_min:.3g}',
            '窗内时间s':r.minimum_time_from_window_start_s,'事件时间s':r.minimum_time_relative_event_s}
            for r in selected.itertuples()]))+'\n\n'
    body+=('共享组目标需要所有训练trial共同可行；上述集中求值异常和接近零的末有效flow支持“个别轨迹阻碍共享组优化”的解释，'
        '但不证明输入试次本身无效，也不证明删除它们即可修复模型。初态、驱动、固定β及未经绝对标定的幅度映射仍可能共同作用。'
        '本轮未删除、替换或重新拟合这些身份；failure_by_trial.csv同时保留没有求值异常但属于失败起点的训练trial。\n\n')
    return body


def fixed_roi_tau_details(run, out, config, data, states, labels, save):
    """Describe shared nonlinear tau estimates; optimizer traces are not profiles."""
    methods = config['methods']; initial = config['schema'] == 'shared_driver_fixed_roi_initial_v1'
    flow = config['schema'] in ('shared_driver_fixed_roi_flow_v1','shared_driver_fixed_roi_initial_v1')
    trainable = [m for m in methods if m not in ('fixed_roi_fixed_tau','fixed_tau_logflow1','fixed_tau_logflow1_tied')]
    options = config['tau_training']; candidates, selections, traces, mapping = [], [], [], []
    for group in data.group.unique():
        prep = json.loads((run/'prepared'/f'{group}.json').read_text()); spec = prep['spec']
        if spec['kind'] == 'measured' and prep['status'] == 'completed':
            projection = prep['metadata']['projection']; roi = prep['metadata']['fixed_roi']
            if projection['fnirs_pair'] != config['fixed_roi']['pair_index'] or roi['pair_name'] != config['fixed_roi']['pair_name']:
                raise ValueError('Fixed ROI metadata differs from declared AF7Fp1 target')
            mapping.append(dict(group=group, pair=roi['pair_name'], pair_index=projection['fnirs_pair'],
                parent_pair_index=roi.get('parent_pair_index'), EEG_factor=projection['eeg_factor'],
                Hb_factor=projection['fnirs_factor'], common_training_MAD=roi.get('common_training_MAD')))
        for method in trainable:
            folder = run/'training'/f'{group}__{method}'
            path = folder/'selection.json'; choice = json.loads(path.read_text()) if path.exists() else {}
            selected = choice.get('selected_training_start'); local = []
            for start, start_tau in enumerate(options['start_values']):
                path = folder/f'start_{start}'/'result.json'; result = json.loads(path.read_text()) if path.exists() else {}
                info = result.get('information', {}); attempt = (result.get('starts') or [{}])[0]
                row = dict(kind=spec['kind'], group=group, subject=spec['subject'], outer=spec['outer'],
                    method=method, start=start, start_tau=start_tau, true_tau=spec.get('true_tau'),
                    true_initial_log_p_over_v=spec.get('true_initial_log_p_over_v'),
                    initial_state_constraint=result.get('initial_state_constraint'),
                    nuisance_free_parameters_per_trial=result.get('nuisance_free_parameters_per_trial'),
                    selected=start == selected, status=result.get('status','missing'), converged=bool(result.get('converged',False)),
                    tau=result.get('parameter_value'), objective=result.get('objective'), boundary=result.get('boundary_status','N/A'),
                    gradient=result.get('projected_scaled_gradient_inf_norm',attempt.get('projected_scaled_gradient_inf_norm')),
                    reason=result.get('convergence_reason',attempt.get('convergence_reason')),
                    optimization_evaluations=result.get('optimization_evaluations'), domain_rejections=attempt.get('domain_rejections'),
                    weighted_data_sse=result.get('weighted_data_sse'), parameter_prior_cost=result.get('parameter_prior_cost'),
                    initial_state_cost=result.get('initial_state_penalty_cost'),
                    flow_prior_cost=result.get('flow_prior_cost'),flow_prior_weight=result.get('flow_prior_weight'),
                    flow_prior_log_sd=result.get('flow_prior_log_sd'),inheritance=json.dumps(result.get('inheritance',{})),
                    derivative_rejections=attempt.get('derivative_rejections'),integration_checks=json.dumps(result.get('integration_checks',[])),
                    data_information=info.get('data_only',{}).get('value'),
                    regularized_information=info.get('nuisance_regularized',{}).get('value'), total_information=info.get('with_parameter_prior'))
                candidates.append(row); local.append(row)
                for point in result.get('trace',[]): traces.append(dict(group=group,kind=spec['kind'],method=method,start=start,**point))
            chosen = [r for r in local if r['selected']]
            if choice.get('status') == 'completed' and (len(chosen) != 1 or chosen[0]['status'] != 'completed' or not chosen[0]['converged']):
                raise ValueError('Tau selection lacks one converged declared start')
            chosen = chosen[0] if chosen else {}
            good = [r for r in local if r['status'] == 'completed' and r['converged']]
            selections.append(dict(kind=spec['kind'], group=group, subject=spec['subject'], outer=spec['outer'], method=method,
                status=choice.get('status','missing'), tau=choice.get('parameter_value'), true_tau=spec.get('true_tau'),
                true_initial_log_p_over_v=spec.get('true_initial_log_p_over_v'),
                successful_starts=len(good), expected_starts=len(options['start_values']), selected_start=selected,
                tau_min=min([r['tau'] for r in good],default=np.nan), tau_max=max([r['tau'] for r in good],default=np.nan),
                **{key:chosen.get(key) for key in ['boundary','objective','gradient','data_information','regularized_information','total_information']}))
    candidates = pd.DataFrame(candidates); selections = pd.DataFrame(selections); traces = pd.DataFrame(traces)
    candidates.to_csv(out/'tau_training_candidates.csv',index=False)
    selections.to_csv(out/'tau_training_selection.csv',index=False)
    traces.to_csv(out/'tau_training_traces.csv',index=False)
    pd.DataFrame(mapping).to_csv(out/'fixed_roi_mapping.csv',index=False)
    short = lambda x:str(x).replace('subject_','S').replace('synthetic_tau','T').replace('slow','慢').replace('mixed','混')
    method_short = labels if flow else {methods[0]:'τ2',methods[1]:'训练τ',methods[2]:'先验τ'}
    body = ('\n## 固定观测位置与解释范围\n\nAF7Fp1（pair0）是预声明的共同帽布局位置，不是个体皮层共定位。'
        'EEG投影保持父规范，Hb目标和训练SD因固定位置而改变；仅本轮同一新目标下重算的τ2作为配对基线。'
        '旧parent的τ和误差只可作为背景，未并入本轮paired scores。目标和预测共用原生观测算子；初始前5秒baseline仍不是潜在静息约束。'
        '所有cell保存的目标/SD与prepared逐值核对，固定ROI元数据另行核对。\n\n')
    if mapping:
        body += table(pd.DataFrame([{'组':short(r['group']),'本轮Hb对':r['pair'],'旧对索引':r['parent_pair_index'],
            'EEG因子':r['EEG_factor'],'新Hb因子':r['Hb_factor']} for r in mapping]))+'\n\n'
    body += ('## 共享τ、多起点与局部信息\n\nτ仅在18条训练轨迹上共享拟合，验证时冻结；'
        f'搜索范围{options["bounds"]}秒，' + ('各新臂无τ先验；log-flow为工程软锚。' if flow else f'先验臂log SD={options["prior_log_sd"]:.4g}，均为工程设置。') +
        '多起点终点散点与迭代trace不是逐点重优化的profile；局部投影信息不是置信区间、校准后验或全局唯一可辨识性证明。'
        '分别列data-only、消除带正则nuisance后的信息、再加入参数先验的信息，不能把正则带来的信息冒充数据支持。\n\n')
    for (kind,method),part in selections.groupby(['kind','method'],sort=False):
        body += f'### {kind} · {method_short[method]}\n\n'
        body += table(pd.DataFrame([{'组':short(r.group),'真τ':r.true_tau,'选择τ':r.tau,
            '成功起点':f'{r.successful_starts}/{r.expected_starts}','起点τ范围':f'{r.tau_min:.3g}–{r.tau_max:.3g}' if r.successful_starts else '—',
            '边界':r.boundary or '—'} for r in part.itertuples()]))+'\n\n'
        body += table(pd.DataFrame([{'组':short(r.group),'训练目标':r.objective,'数据':r.data_information,
            '加nuisance正则':r.regularized_information,'再加τ先验':r.total_information} for r in part.itertuples()]))+'\n\n'
    failures = candidates[~(candidates.status.eq('completed')&candidates.converged)].copy()
    body += f'全部训练起点{len(candidates)}，未成功{len(failures)}；失败起点保留终点目标、梯度、域拒绝及预算原因，不参与参数选择。'
    body += f'收敛阈值{options["gradient_tolerance"]:.3g}；目标/梯度/预算明细见tau_training_candidates.csv。\n\n'
    if len(failures):
        display=failures.groupby(['kind','method','reason'],dropna=False).size().rename('起点数').reset_index()
        display['method']=display.method.map(method_short)
        body+=table(display.rename(columns={'kind':'数据','method':'方法','reason':'停止原因'}))+'\n\n'
    variation=[]
    for (subject,method),part in selections[selections.kind.eq('measured')].groupby(['subject','method'],sort=False):
        good=pd.to_numeric(part.loc[part.status.eq('completed'),'tau'],errors='coerce').dropna()
        variation.append(dict(subject=subject,method=method,successful_folds=len(good),expected_folds=len(config['outer_folds']),
            tau_min=good.min(),tau_max=good.max(),geometric_mean=float(np.exp(np.log(good).mean())),log_tau_sd=np.log(good).std()))
    pd.DataFrame(variation).to_csv(out/'tau_subject_fold_variation.csv',index=False)
    body+='### 被试与折间变化（描述性）\n\n'
    if variation:
        body+=table(pd.DataFrame([{'被试':short(r['subject']),'方法':method_short[r['method']],
            '成功折':f"{r['successful_folds']}/{r['expected_folds']}",'τ范围':f"{r['tau_min']:.3g}–{r['tau_max']:.3g}" if r['successful_folds'] else '—',
            '几何均值':r['geometric_mean'],'折间log SD':r['log_tau_sd']} for r in variation]))+'\n\n'
    body+='折间训练集重叠，四折不能视为四次独立重复；固定位置减少选对混杂，但不消除观测规范、初态或固定β的补偿。没有足够数据支持人口参数分布或唯一的被试生理差异。\n\n'
    fig,axes=plt.subplots(2,2,figsize=(12,8))
    measured=selections[selections.kind.eq('measured')]; groups=list(measured.group.unique())
    for method in trainable:
        part=measured[measured.method.eq(method)].set_index('group').reindex(groups)
        axes[0,0].plot(range(len(groups)),pd.to_numeric(part.tau,errors='coerce'),'o-',label=method_short[method])
        synthetic=selections[selections.kind.eq('synthetic')&selections.method.eq(method)&selections.status.eq('completed')]
        axes[0,1].scatter(synthetic.true_tau,synthetic.tau,label=method_short[method],s=25)
        good=selections[selections.method.eq(method)&selections.status.eq('completed')]
        axes[1,0].scatter(pd.to_numeric(good.data_information),pd.to_numeric(good.regularized_information),label=method_short[method],s=20)
    axes[0,0].set_xticks(range(len(groups)),[short(g) for g in groups],rotation=50,ha='right',fontsize=6)
    axes[0,0].set_ylabel('训练冻结τ / 秒');axes[0,0].set_title('同位置：被试×折');axes[0,0].legend(fontsize=7)
    if not len(measured):
        axes[0,0].text(.5,.5,'实测阶段未纳入本报告',ha='center',va='center',transform=axes[0,0].transAxes)
    axes[0,1].plot(config['synthetic']['taus_s'],config['synthetic']['taus_s'],'k:');axes[0,1].set_xlabel('合成真τ / 秒');axes[0,1].set_ylabel('训练估计τ / 秒');axes[0,1].legend(fontsize=7)
    axes[1,0].set_xscale('symlog',linthresh=1e-8);axes[1,0].set_yscale('symlog',linthresh=1e-8)
    axes[1,0].set_xlabel('data-only局部信息');axes[1,0].set_ylabel('nuisance正则化局部信息');axes[1,0].legend(fontsize=7)
    for (group,method),part in candidates.groupby(['group','method'],sort=False):
        values=pd.to_numeric(part.objective,errors='coerce');minimum=values.min()
        for ok,marker in [(True,'o'),(False,'x')]:
            take=part.converged.eq(ok)&part.objective.notna()
            axes[1,1].scatter(pd.to_numeric(part.loc[take,'tau']),values[take]-minimum,marker=marker,s=14,alpha=.65,
                color=plt.get_cmap('tab10')(trainable.index(method)))
    axes[1,1].set_xlabel('多起点终点τ / 秒');axes[1,1].set_ylabel('终点目标 − 同组同臂最小终点目标');axes[1,1].set_title('圆：收敛；叉：失败。不是profile')
    save(fig,'fixed_roi_tau_parameters','被试/折与合成τ恢复、局部信息和多起点终点；未收敛目标仅供审计。信息与终点散点均不构成profile或置信区间。')
    if len(traces):
        for batch in range(0,len(trainable),2):
            plotted=trainable[batch:batch+2]
            fig,axes=plt.subplots(3,len(plotted),figsize=(12,9),sharex='col',squeeze=False)
            for col,method in enumerate(plotted):
                trace_kind='measured' if 'measured' in data.kind.unique() else 'synthetic'
                part=traces[traces.kind.eq(trace_kind)&traces.method.eq(method)]
                for index,((group,start),points) in enumerate(part.groupby(['group','start'],sort=False)):
                    points=points.sort_values(['evaluations','iteration'],kind='stable');color=plt.get_cmap('tab20')(index//3%20)
                    for ax,field in zip(axes[:,col],['objective','projected_scaled_gradient_inf_norm','parameter_value']):
                        ax.plot(points.evaluations,points[field],color=color,ls=['-','--',':'][int(start)%3],lw=.8,alpha=.7)
                axes[0,col].set_title(method_short[method]);axes[0,col].set_ylabel('训练目标')
                axes[1,col].set_yscale('log');axes[1,col].axhline(options['gradient_tolerance'],color='red',ls=':');axes[1,col].set_ylabel('projected scaled gradient')
                axes[2,col].set_ylabel('τ / 秒');axes[2,col].set_xlabel('单起点优化trial forward调用')
            save(fig,'fixed_roi_tau_traces'+(f'_{batch//2+1}' if len(trainable)>2 else ''),('实测' if 'measured' in data.kind.unique() else '合成')+'训练迭代轨迹：目标、实际梯度和τ；红线为原收敛阈值。各线为起点，不是profile扫描；完整身份见tau_training_traces.csv。')
    # Truth recovery uses the original state coordinates, without alignment.
    recovered=[]
    for (method,mode),part in states[states.kind.eq('synthetic')].groupby(['method','mode'],sort=False):
        good=part[part.converged & part.finite_states]
        row=dict(method=method,mode=mode,completed=len(good),expected=int(((data.kind=='synthetic')&(data.method==method)&(data['mode']==mode)).sum()))
        for name in ['r','s','f','v','p','q']:
            values=pd.to_numeric(good.get('truth_rmse_'+name,pd.Series(dtype=float)),errors='coerce')
            row[name+'_rmse']=float(np.sqrt(np.mean(values**2))) if len(values) else np.nan
        recovered.append(row)
    pd.DataFrame(recovered).to_csv(out/'fixed_roi_synthetic_state_recovery.csv',index=False)
    body+='### 合成状态恢复\n\n原坐标状态RMSE在成功trial内平方平均后取根，不做平移、幅度、符号或时间对齐；不是对静息常数归一化的误差。\n\n'
    for mode in config['modes']:
        part=[r for r in recovered if r['mode']==mode]
        body+=f'{mode}：\n\n'+table(pd.DataFrame([{'方法':method_short[r['method']],'成功/计划':f"{r['completed']}/{r['expected']}",
            **{name:r[name+'_rmse'] for name in ['r','s','f','v','p','q']}} for r in part]))+'\n\n'
    fig,axes=plt.subplots(3,3,figsize=(13,9),sharex='col');example_notes=[]
    for col,tau in enumerate(config['synthetic']['taus_s']):
        group=f'synthetic_tau{tau:g}_slow_r0'+('_logpv+0.00' if initial else '');trial=0;path=run/'prepared'/f'{group}.npz'
        axes[0,col].set_title(f'真τ={tau:g}s / slow / r0 / trial0',fontsize=9)
        if not path.exists():
            for ax in axes[:,col]:ax.text(.5,.5,'固定样例缺失，不替换',transform=ax.transAxes,ha='center')
            continue
        with np.load(path,allow_pickle=False) as a:
            t=np.arange(a['target'].shape[1])*config['tensor']['dt_s']
            axes[0,col].plot(t,a['target'][trial,:,1],color='.65',lw=.8,label='含噪HbO目标')
            for ax,value in zip(axes[:,col],[a['clean'][trial,:,1],a['truth'][trial],a['truth_states'][trial,:,2]]):ax.plot(t,value,'k',lw=1.3,label='生成真值')
        for method in methods:
            path=run/'cells'/f'{group}__{method}__full'/'trajectories.npz'
            row=data[(data.group==group)&(data.method==method)&(data['mode']=='full')&(data.trial==trial)]
            status=row.status.iloc[0] if len(row) else 'missing';example_notes.append(f'{group}/{method_short[method]}:{status}')
            if not path.exists():
                for ax in axes[:,col]:ax.plot([],[],':',label=method_short[method]+' 未估计')
                continue
            with np.load(path,allow_pickle=False) as a:
                indices=np.flatnonzero(a['trial_indices']==trial)
                if not len(indices):raise ValueError('Fixed three-tau example lacks trial0')
                i=int(indices[0]);values=[a['prediction'][i,:,1],a['states'][i,:,0],a['states'][i,:,2]]
                for ax,value in zip(axes[:,col],values):
                    suffix=' 未估计' if not np.isfinite(value).any() else (' 失败迭代' if status!='completed' else '')
                    ax.plot(t,value,'-' if status=='completed' else ':',lw=.9,label=method_short[method]+suffix)
        for ax,name in zip(axes[:,col],['处理HbO','driver r','flow f']):ax.set_ylabel(name);ax.legend(fontsize=5)
        axes[2,col].set_xlabel('秒')
    save(fig,'fixed_roi_three_tau_truth','固定选择slow、replicate0、trial0的τ=1/2/4三组；相同driver/初态/标准噪声配对，展示真值与预声明各臂估计。flow真值随τ不必改变，因为τ不直接进入s/f方程；Hb动态承载下游τ差异。失败保留，不做事后对齐。')
    if initial:body+='该通用三τ图仅示log(p0/v0)=0匹配层；±0.05错设层必须结合后文专门分层图解释。\n\n'
    body+='## 固定三τ真轨迹与隐藏检验边界\n\n三τ图事先按slow/replicate0/trial0身份固定，不按恢复质量选样。'+ '；'.join(example_notes)+'。\n\n'
    body+='full是联合重建，隐藏模式只评分被隐藏分量；参数训练与验证目标分开，不能把full重建当作隐藏恢复证据。中心插补使用两侧上下文，仍不是因果预测或独立被试验证；完整目标仅用于评分，own-context与τ2都在同一新观测坐标比较。\n\n'
    body += fixed_roi_training_failure_details(run, out, config, candidates, save)
    check_path = run/'unit_reexpression_check.json'
    if check_path.exists():
        check = json.loads(check_path.read_text())
        if check.get('status') == 'passed' and check.get('case_count') == 2:
            body += '单位重表达software check为2 passed：同步缩放目标、SD与观测算子保持物理拟合不变；这不是独立幅度标定或生理参数恢复证据。\n\n'
    return body


def decompose_processed_hbt(state, prediction, operator, p_loading):
    """Exact additive posthoc identity; the normalized cross term is signed."""
    volume=operator@(p_loading*(state[:,3]-1.))
    initial_difference=operator@(p_loading*(state[:,4]-state[:,3]))
    total=prediction[:,1]+prediction[:,2]
    error=float(np.max(np.abs(volume+initial_difference-total)))
    if not np.allclose(volume+initial_difference,total,rtol=1e-8,atol=1e-9):
        raise ValueError('Processed HbT decomposition disagrees with saved prediction')
    ms=float(np.mean(total**2));vms=float(np.mean(volume**2));dms=float(np.mean(initial_difference**2))
    cross=float(2*np.mean(volume*initial_difference))
    row=dict(total_rms=np.sqrt(ms),volume_rms=np.sqrt(vms),initial_difference_rms=np.sqrt(dms),
        signed_cross_term=cross,additivity_max_abs_difference=error,
        volume_rms_over_total=np.sqrt(vms/ms) if ms>1e-24 else np.nan,
        initial_difference_rms_over_total=np.sqrt(dms/ms) if ms>1e-24 else np.nan,
        signed_cross_over_total_ms=cross/ms if ms>1e-24 else np.nan)
    return row,volume,initial_difference,total


def flow_full_curve_atlas(run,out,config,data,labels,comparison_methods=None,*,both_states=False,context_label=""):
    """All 72 measured identities, fixed order; no score selection or new fit."""
    methods=comparison_methods or ['fixed_tau_logflow1','trained_tau_logflow1']
    names=[labels[m] for m in methods] if comparison_methods else ['固定τ/w1','共享τ/w1']
    identities=['group','subject','outer','session','trial']
    subset=data[data.kind.eq('measured')&data['mode'].eq('full')&data.method.isin(methods)]
    panels=subset[identities].drop_duplicates().sort_values(['subject','outer','trial'],kind='stable')
    if len(panels)!=72 or len(subset)!=72*len(methods):raise ValueError('Full curve atlas requires all arms and exactly 72 identities')
    if subset.duplicated(identities+['method']).any():raise ValueError('Duplicate atlas identity')
    folder=out/'curve_atlas';folder.mkdir();pdf=fitz.open();index=[]
    cache={}
    for (group,method),part in subset.groupby(['group','method']):
        path=run/'cells'/f'{group}__{method}__full'/'trajectories.npz'
        if path.exists():
            with np.load(path,allow_pickle=False) as a:cache[(group,method)]={k:a[k].copy() for k in ['target','prediction','states','trial_indices']}
    for start in range(0,72,3):
        page_number=start//3+1;separate_states=len(methods)>2
        ncols=3+len(methods) if separate_states else 4
        fig,axes=plt.subplots(3,ncols,figsize=(14 if separate_states else 12,8),sharex=True)
        page_labels=[]
        for row,identity in enumerate(panels.iloc[start:start+3].itertuples(index=False)):
            records=subset[subset.group.eq(identity.group)&subset.trial.eq(identity.trial)].set_index('method')
            t=np.arange(config['tensor']['steps'])*config['tensor']['dt_s']-5.
            text_status=[];actual_drawn=False
            for method,color,style in zip(methods,['.5','tab:blue','tab:orange'],['--','-',':']):
                record=records.loc[method];errors='/'.join(f'{record.get("nrmse_"+c,np.nan):.3f}' for c in COMPONENTS)
                text_status.append(names[methods.index(method)]+f' {record.status}, NRMSE={errors}')
                arrays=cache.get((identity.group,method))
                if arrays is None:continue
                where=np.flatnonzero(arrays['trial_indices']==identity.trial)
                if len(where)!=1:raise ValueError('Atlas trial missing')
                k=int(where[0])
                if not actual_drawn:
                    for col in range(3):axes[row,col].plot(t,arrays['target'][k,:,col],'k',lw=.9,label='实测')
                    actual_drawn=True
                successful=record.status=='completed' and bool(truth(pd.Series([record.converged])).iloc[0])
                if separate_states and not successful:
                    axes[row,3+methods.index(method)].text(.5,.5,'FAILED',ha='center',fontsize=6,transform=axes[row,3+methods.index(method)].transAxes)
                    continue
                for col in range(3):axes[row,col].plot(t,arrays['prediction'][k,:,col],style,color=color,lw=.9,label=names[methods.index(method)])
                state_ax=axes[row,3+methods.index(method)] if separate_states else axes[row,3]
                if method==methods[1] or both_states:
                    if not np.isfinite(arrays['states'][k]).any():state_ax.text(.5,.5,'未估计',ha='center',transform=state_ax.transAxes)
                    for j,name in enumerate(['f','v','p','q'],start=2):state_ax.plot(t,arrays['states'][k,:,j],lw=.85,ls=style if both_states else '-',color=f'C{j-2}' if both_states else None,label=name if separate_states else name+(' 定τ' if method==methods[0] else ' 学τ') if both_states else name)
            identity_label=f'{identity.subject} / fold {identity.outer} / trial {identity.trial}'
            page_labels.append(identity_label+'；'+'；'.join(text_status))
            axes[row,0].set_ylabel(identity_label,fontsize=7)
            for col,name in enumerate(['EEG','HbO','HbR',('两臂状态' if both_states else names[1]+'状态')] if not separate_states else ['EEG','HbO','HbR']+[name+'状态' for name in names]):
                axes[row,col].set_title(name if row==0 else '',fontsize=8)
                axes[row,col].tick_params(labelsize=6);axes[row,col].axvline(0,color='.75',ls=':',lw=.5)
            for state_axis in axes[row,3:]:state_axis.axhline(1,color='.5',ls=':',lw=.7)
            if row==0:
                axes[row,0].legend(fontsize=6,ncol=1);axes[row,3].legend(fontsize=6,ncol=2 if separate_states else 4)
            for ax in axes[row]:ax.set_xlabel('相对事件 / 秒',fontsize=7)
            index.append(dict(page=page_number,**identity._asdict(),**({m+'_status':records.loc[m,'status'] for m in methods} if separate_states else dict(fixed_status=records.loc[methods[0],'status'],trained_status=records.loc[methods[1],'status']))))
        fig.tight_layout(rect=(0,.01,1,.97));png=folder/f'page_{page_number:02d}.png';fig.savefig(png,dpi=240);plt.close(fig)
        page=pdf.new_page(width=864,height=720)
        heading=f'{context_label+" · " if context_label else ""}全72条曲线核验 · {page_number}/24 · {names[0]} 对 {names[1]}'
        if separate_states:heading=f'{context_label} · 全72条曲线 · {page_number}/24 · 三臂分列状态'
        page.insert_htmlbox(fitz.Rect(24,8,840,36),'<b>'+heading+'</b>',css='*{font-size:11pt}')
        page.insert_image(fitz.Rect(24,38,840,582),filename=str(png))
        caption='<br>'.join(page_labels)+f'<br>黑：实测；灰虚：{names[0]}；蓝：{names[1]}。末列为{names[1]} f/v/p/q，基线1是静息参照，不是正常界。NRMSE顺序EEG/HbO/HbR；失败保留，不以图像选择方法。当前Hb为相对观测坐标。'
        if both_states:caption+=' 状态列同色为同一状态，虚线固定τ、实线训练τ；失败的末有效轨迹不代表有效估计。'
        if separate_states:caption='<br>'.join(page_labels)+'<br>黑：实测；灰虚：A；蓝实：B；橙点：C。仅成功轨迹绘线，失败保留状态标签。后三列分别为各臂f/v/p/q，独立纵轴；基线1不是正常界。NRMSE顺序EEG/HbO/HbR。'
        spare,scale=page.insert_htmlbox(fitz.Rect(24,588,840,710),caption,css='*{font-size:9pt;line-height:1.35}')
        if spare<0 or scale<.95:raise ValueError('Atlas caption failed to fit legibly')
    pdf.save(out/'FULL_CURVE_ATLAS.pdf',garbage=4,deflate=True)
    audit=dict(pages=len(pdf),image_objects=sum(len(p.get_images()) for p in pdf),searchable_text_chars=sum(len(p.get_text()) for p in pdf),bitmap_dpi=240,wps_checked=False,identities=72)
    if audit['pages']!=24 or audit['image_objects']!=24 or audit['searchable_text_chars']<1000:raise ValueError('Atlas verification failed')
    (out/'atlas_validation.json').write_text(json.dumps(audit,indent=2));pd.DataFrame(index).to_csv(out/'atlas_index.csv',index=False);pdf.close()
    return f'\n## 全72条曲线核验图册\n\n[FULL_CURVE_ATLAS.pdf](FULL_CURVE_ATLAS.pdf)按固定被试/折/trial顺序展示全部72条实测身份，每页3条，比较{names[0]}与{names[1]}；包含EEG/HbO/HbR、共享状态及逐条误差/状态。图册保留失败、不按效果挑样；静息参照1不是生理正常界。\n\n'


def flow_hbt_decomposition(run,out,config,data,labels,save,selected):
    """Read retained successful full fits; do not change scores or fit anything."""
    import importlib.util
    import sys
    root=run/'source_snapshot'
    name='report_frozen_hbt_core'
    spec=importlib.util.spec_from_file_location(name,root/'src/inference/t3a_balloon_robust_ssm.py')
    core=importlib.util.module_from_spec(spec);sys.modules[name]=core;spec.loader.exec_module(core)
    fixed={k:v for k,v in config['fixed'].items() if k!='kappa'}
    p=core.BalloonParameters(core.BalloonFixedParameters(**fixed),core.BalloonFreeParameters(kappa=config['fixed']['kappa'],tau=config['reference_tau_s']))
    rest=np.array([0.,0.,1.,1.,1.,1.]);base=core.observation_map(rest,p)[1:].sum()
    shifted=rest.copy();shifted[4]+=1.;loading=float(core.observation_map(shifted,p)[1:].sum()-base)
    shifted=rest.copy();shifted[5]+=.1
    if abs(base)>1e-12 or abs(core.observation_map(shifted,p)[1:].sum()-base)>1e-12:
        raise ValueError('Frozen optical map does not permit HbO+HbR=P-loading*(p-1)')
    spec=importlib.util.spec_from_file_location('report_hbt_operators',root/'src/inference/observation_baselines.py')
    owners=importlib.util.module_from_spec(spec);spec.loader.exec_module(owners)
    operators=owners.native_feature_operators(config['tensor']['steps'])
    operator=operators['fnirs']@operators['native_interpolation']
    rows=[];curves={}
    for identity in data[data['mode'].eq('full')&data.status.eq('completed')&truth(data.converged)].itertuples():
        path=run/'cells'/f'{identity.group}__{identity.method}__full'/'trajectories.npz'
        with np.load(path,allow_pickle=False) as arrays:
            index=np.flatnonzero(arrays['trial_indices']==identity.trial)
            if len(index)!=1:raise ValueError('HbT identity missing')
            k=int(index[0]);row,volume,difference,total=decompose_processed_hbt(arrays['states'][k],arrays['prediction'][k],operator,loading)
        rows.append(dict(kind=identity.kind,group=identity.group,subject=identity.subject,session=identity.session,outer=identity.outer,trial=identity.trial,method=identity.method,**row))
        if any(example.group==identity.group and example.trial==identity.trial for _,example in selected):curves[(identity.group,identity.trial,identity.method)]=(volume,difference,total)
    frame=pd.DataFrame(rows);frame.to_csv(out/'flow_hbt_decomposition.csv',index=False)
    parent_check=run/'parent_initial_difference_decomposition.json'
    parent_verified=0
    if parent_check.exists():
        reference=json.loads(parent_check.read_text())
        if reference['parent_run']!=config['parent_run']:raise ValueError('HbT audit references a different parent')
        previous=pd.DataFrame(reference['rows']);identity=['subject','outer','trial','method']
        inherited=frame[frame.kind.eq('measured')&frame.method.isin(config['methods'][:2])]
        paired=inherited.merge(previous,on=identity,how='outer',validate='one_to_one',indicator=True)
        if not paired['_merge'].eq('both').all():raise ValueError('Parent HbT audit identities differ')
        for field,ref in [('volume_rms_over_total','volume_rms_ratio'),('initial_difference_rms_over_total','initial_difference_rms_ratio'),('signed_cross_over_total_ms','cross_energy_ratio')]:
            if not np.allclose(paired[field],paired[ref],rtol=1e-8,atol=1e-9,equal_nan=True):raise ValueError('Parent HbT decomposition mismatch '+field)
        parent_verified=len(paired)
    display=[]
    for (kind,method),part in frame.groupby(['kind','method'],sort=False):
        expected=int((data.kind.eq(kind)&data.method.eq(method)&data['mode'].eq('full')).sum())
        display.append({'数据':kind,'方法':labels[method],'成功/计划':f'{len(part)}/{expected}',
            'V/总RMS中位':part.volume_rms_over_total.median(),'δ/总RMS中位':part.initial_difference_rms_over_total.median(),
            '有符号交叉项中位':part.signed_cross_over_total_ms.median(),'最大相加误差':part.additivity_max_abs_difference.max()})
    display.sort(key=lambda row: (['synthetic','measured'].index(row['数据']),list(labels.values()).index(row['方法'])))
    summary=pd.DataFrame(display);summary.to_csv(out/'flow_hbt_decomposition_summary.csv',index=False)
    displayed=summary.copy();displayed['最大相加误差']=displayed['最大相加误差'].map(lambda value:f'{value:.2e}')
    body=('\n## Posthoc HbT组成诊断\n\n'
        '此处是事后机制诊断，不是新增主评分或独立解释比例。当前p=v+δ，实际冻结光学映射下HbT=HbO+HbR；'
        f'从冻结core逐项核验的p系数为{loading:g}。采用与full拟合相同的L=原生Hb处理算子×模型时钟插值，'
        'V=L[P0(v−1)]、D=L[P0(p−v)]严格相加为已保存预测HbO+HbR（实际增益/坐标系数已纳入p系数）。'
        'D是初态差值衰减项：幅度来自p0−v0，衰减速率f_out/(τv)仍随驱动下的v变化；不是独立于driver的私有源，只有近静息时才近似时间常数τ的指数。不能仅凭范数高就认定非生理。'
        '每trial核验RMS(total)^2=RMS(V)^2+RMS(D)^2+2mean(V·D)。'
        '两项RMS/total与有符号交叉项2mean(V·D)/mean(total²)可能大于1或为负，反映抵消，不能称独立解释百分比；total RMS不超过1e−12时比值记为未定义。'
        '表内分别取逐trial比值中位，不能再把这些中位数代回同一trial恒等式；完整逐条RMS和交叉项见CSV。'
        '仅成功full进入该诊断，失败不以末有效迭代补入；原分母、主评分和预声明方法不变。\n\n'+table(displayed)+'\n\n')
    if parent_verified:
        body+=f'父两基线{parent_verified}条成功full身份与独立[父初态差值分解审计](../parent_initial_difference_decomposition.json)逐项核对通过；继承结果不是新增复制。\n\n'
    for rank,example in selected:
        fig,axes=plt.subplots(2,len(config['methods'])//2,figsize=(13,7),sharex=True)
        for ax,method in zip(axes.flat,config['methods']):
            curve=curves.get((example.group,example.trial,method))
            ax.set_title(labels[method],fontsize=10)
            if curve is None:
                ax.text(.5,.5,'未成功：不作该分解\n固定身份不替换',ha='center',va='center',transform=ax.transAxes)
                continue
            volume,difference,total=curve;t=np.arange(len(total))*config['tensor']['dt_s']-(5 if example.kind=='measured' else 0)
            ax.plot(t,total,'k',label='预测HbT',lw=1.3);ax.plot(t,volume,label='血容量V',lw=1);ax.plot(t,difference,label='初态差异D',lw=1)
            ax.axhline(0,color='.7',ls=':',lw=.6);ax.legend(fontsize=7);ax.set_ylabel('处理后HbT')
        for ax in axes[-1]:ax.set_xlabel('相对事件 / 秒' if example.kind=='measured' else '秒')
        fig.suptitle(f'Posthoc {rank}: {example.group} / trial {int(example.trial)}',fontsize=11)
        save(fig,'flow_hbt_'+rank,'父基线固定身份的各臂预测HbT分解；V+D严格等于总预测，成分不独立，可互相抵消。失败不换样；这不是新主评分或生理正常性判据。')
    return body


def fixed_roi_flow_details(run, out, config, data, states, labels, save, selected):
    """Conditional paired scores and flow diagnostics, without selecting an arm."""
    identities=['kind','group','subject','session','trial','mode']
    transitions=[]; score_rows=[]; identity_rows=[]
    comparisons=[(m,'fixed_roi_fixed_tau') for m in config['methods'][2:]]
    comparisons += [(m,'fixed_roi_trained_tau') for m in config['methods'][3:]]
    comparisons += [('trained_tau_logflow1','fixed_tau_logflow1')]
    initial=config['schema']=='shared_driver_fixed_roi_initial_v1'
    if initial:
        m=config['methods'];comparisons=[(m[2],m[0]),(m[3],m[1]),(m[1],m[0]),(m[3],m[2])]
    for method,baseline in comparisons:
        a=data[data.method.eq(method)];b=data[data.method.eq(baseline)]
        joined=a.merge(b,on=identities,suffixes=('','_base'),how='outer',validate='one_to_one',indicator=True)
        if not joined['_merge'].eq('both').all():raise ValueError('Flow comparison changed identity set')
        ok=joined.status.eq('completed')&truth(joined.converged)
        baseok=joined.status_base.eq('completed')&truth(joined.converged_base)
        joined['transition']=np.select([ok&baseok,ok&~baseok,~ok&baseok],['both_success','new_success','lost_success'],default='both_failed')
        for (kind,mode),part in joined.groupby(['kind','mode'],sort=False):
            row=dict(kind=kind,method=method,baseline=baseline,mode=mode,expected=len(part))
            row.update({key:int(part.transition.eq(key).sum()) for key in ['both_success','new_success','lost_success','both_failed']})
            transitions.append(row)
            for subset in ['both_success','new_success']:
                good=part[part.transition.eq(subset)]
                for component in COMPONENTS:
                    if mode=='center_EEG' and component!='EEG':continue
                    if mode=='center_fNIRS' and component=='EEG':continue
                    field='nmse_'+component
                    if field not in good:continue
                    result=dict(kind=kind,method=method,baseline=baseline,mode=mode,subset=subset,n=len(good),expected=len(part),component=component)
                    def average(values):
                        return values.groupby([good.subject,good.session]).mean().groupby(level=0).mean().mean()
                    result['candidate_nrmse']=np.sqrt(average(good[field]))
                    result['baseline_nrmse']=np.sqrt(average(good[field+'_base'])) if subset=='both_success' else np.nan
                    result['delta_nmse']=average(good[field]-good[field+'_base']) if subset=='both_success' else np.nan
                    score_rows.append(result)
        identity_rows.append(joined[identities+['transition']].assign(method=method,baseline=baseline))
    transition=pd.DataFrame(transitions);scores=pd.DataFrame(score_rows)
    transition.to_csv(out/'flow_paired_transitions.csv',index=False)
    scores.to_csv(out/'flow_paired_scores.csv',index=False)
    pd.concat(identity_rows,ignore_index=True).to_csv(out/'flow_paired_identities.csv',index=False)
    body='\n## Log-flow工程软约束与配对结果\n\n六臂及权重预声明，不按验证成绩选择w。log-flow软约束只是以f=1为中心的工程锚，log SD=ln2不是正常范围、人口分布或校准生理先验。不同w的总目标不可横向排名；多起点仅在同组同臂内选择。继承父运行的两臂及原prepared观测不是新独立复制。\n\n'
    equation_path=run/'model_equation_audit.json'
    if 'measured' in data.kind.unique() and equation_path.exists():
        audit=json.loads(equation_path.read_text())
        if audit['execution_effect']['running_flow_model_changed'] or audit['execution_effect']['frozen_snapshot_changed']:
            raise ValueError('Equation audit says the running/frozen model was changed')
        body+=('### 方程来源与初态自由度\n\n'
            '本轮运行的是合同定义的inlet-balance变体：τ·dp/dt=f−f_out·p/v，'
            '与Tak等2015文中式(3)的τ·dp/dt=(f−f_out)·p/v不同，不能在独立p0/v0初态下视作相同方程。'
            '当前δ=p−v满足dδ/dt=−f_out·δ/(τv)，近静息时衰减尺度约为τ；'
            '因此自由p0/v0允许未知窗前历史，同时提供与τ耦合的衰减Hb成分。原文方程则保持p/v恒定；两者在p0=v0时均保持p=v。'
            '该自由度尚无本实验独立浓度/幅度标定支持，也不能仅凭数学自由度断言生理不可能。'
            '来源审计没有修改本轮方程或冻结代码；当前各臂比较仍使用完全相同的原变体。'
            '详见[model_equation_audit.json](../model_equation_audit.json)及'
            f"[Tak et al. 2015]({audit['primary_source']['url']})。\n\n")
        initial=[]
        for method in config['methods']:
            all_rows=data[data.kind.eq('measured')&data.method.eq(method)&data['mode'].eq('full')]
            good=states[states.kind.eq('measured')&states.method.eq(method)&states['mode'].eq('full')&states.converged&states.finite_states]
            ratio=good.p_initial/good.v_initial;delta=good.p_initial-good.v_initial
            initial.append({'方法':labels[method],'成功/计划':f'{len(good)}/{len(all_rows)}',
                'p0/v0最小':ratio.min(),'p0/v0中位':ratio.median(),'p0/v0最大':ratio.max(),'最大绝对初态差':delta.abs().max()})
        initial=pd.DataFrame(initial);initial.to_csv(out/'flow_initial_pv_audit.csv',index=False)
        body+='实测full成功子集的初态审计；失败完整分母仍保留，上述比值不是独立测得浓度：\n\n'+table(initial)+'\n\n'
    null_path=run/'stationary_baseline_null_check.json'
    if null_path.exists():
        null=json.loads(null_path.read_text())
        if null.get('status')!='passed':raise ValueError('Stationary baseline null check did not pass')
        records=null['records'];maximum=max(r['full_processed_max_abs'] for r in records)
        drift=max(r['rollout_state_max_abs_drift'] for r in records)
        body+=('### 恒定水平负对照\n\n'
            f"独立确定性算子检查（非重新拟合）覆盖f={null['grid']['flows']}、τ={null['grid']['taus_s']}："
            '令外加常数r=γ(c−1)/β、s=0、v=p=c^α、q=vE(c)/E0，'
            f'末轨迹漂移最大{drift:.3g}，完整处理后的观测绝对最大{maximum:.4g}。'
            '不同的数学域有效恒定绝对水平被baseline处理消去，不能由零处理观测证明f≈1。'
            '原初态惩罚已提供绝对水平锚，flow正则再增加工程锚；相同数据残差不表示正则总目标相等。'
            '这是恒定轨迹族的负对照，不是任意动态轨迹的全局等价变换，也不意味着实测τ毫无信息。'
            '证据：[stationary_baseline_null_check.json](../stationary_baseline_null_check.json)。\n\n')
    counts=[]
    for kind,part in data.groupby('kind',sort=False):
        for inherited,arm_names in [(True,config['methods'][:2]),(False,config['methods'][2:])]:
            rows=part[part.method.isin(arm_names)]
            if initial and kind=='synthetic':
                if inherited:continue
                rows=part
            counts.append({'数据':kind,'来源':'父继承' if inherited else '本轮新拟合',
                'cell数':len(rows[['group','method','mode']].drop_duplicates()),'记录trial数':len(rows),
                '成功trial数':int((rows.status.eq('completed')&truth(rows.converged)).sum())})
    body+=table(pd.DataFrame(counts))+'\n\n'
    for (kind,mode),part in transition.groupby(['kind','mode'],sort=False):
        body+=f'### {kind} / {mode} 身份转移\n\n'+table(pd.DataFrame([{'候选':labels[r.method],'基线':labels[r.baseline],'共同成功':r.both_success,'新成功':r.new_success,'失去成功':r.lost_success,'双方失败':r.both_failed,'计划':r.expected} for r in part.itertuples()]))+'\n\n'
    body+='同w=1固定τ对训练τ的误差只在共同成功身份比较；新成功样本另列，不能用不同成功子集的总体均值声称预测改善。\n\n'
    part=scores[scores.baseline.eq('fixed_tau_logflow1')]
    if initial:
        m=config['methods'];part=scores[(scores.method.eq(m[2])&scores.baseline.eq(m[0]))|(scores.method.eq(m[3])&scores.baseline.eq(m[1]))]
    for (kind,mode,method,baseline),group in part.groupby(['kind','mode','method','baseline'],sort=False):
        body+=f'{kind} / {mode} / {labels[method]} vs {labels[baseline]}：\n\n'+table(group[['subset','component','n','expected','candidate_nrmse','baseline_nrmse','delta_nmse']].rename(columns={'subset':'子集','component':'分量','n':'成功数','expected':'计划','candidate_nrmse':'候选','baseline_nrmse':'基线','delta_nmse':'ΔNMSE'}))+'\n\n'
    recovery=[];extremes=[]
    for (kind,method,mode),part in states.groupby(['kind','method','mode'],sort=False):
        good=part[part.converged&part.finite_states]
        row=dict(kind=kind,method=method,mode=mode,completed=len(good),expected=int(((data.kind==kind)&(data.method==method)&(data['mode']==mode)).sum()))
        for name in ['flow_below_half_s','flow_above_two_s','flow_below_01_s','driver_mean','driver_ac_rms']:
            row[name+'_mean']=pd.to_numeric(good.get(name,pd.Series(dtype=float))).mean()
        extremes.append(row)
        if kind=='synthetic':
            rec=dict(method=method,mode=mode,completed=len(good),expected=row['expected'])
            for name in ['driver_dc_error','driver_ac_rmse','initial_rmse','truth_rmse_f']:
                values=pd.to_numeric(good.get(name,pd.Series(dtype=float)))
                rec[name]=np.sqrt(np.mean(values**2)) if len(values) else np.nan
            rec['driver_dc_signed_mean']=pd.to_numeric(good.get('driver_dc_error',pd.Series(dtype=float))).mean()
            rec['flow_signed_mean']=pd.to_numeric(good.get('flow_mean_error',pd.Series(dtype=float))).mean()
            recovery.append(rec)
    pd.DataFrame(extremes).to_csv(out/'flow_duration_summary.csv',index=False)
    pd.DataFrame(recovery).to_csv(out/'flow_synthetic_recovery.csv',index=False)
    candidates=pd.read_csv(out/'tau_training_candidates.csv')
    body+='## 训练求值、域拒绝与正则代价\n\n总目标、数据SSE和flow代价分开保留于tau_training_candidates.csv；不同w不作总目标胜负比较。域拒绝次数、终点梯度和停止原因反映优化路径，不能把预算失败统一称为域不合理。\n\n'
    diagnostics=[]
    for (kind,method),part in candidates.groupby(['kind','method'],sort=False):
        good=part.status.eq('completed')&truth(part.converged)
        gradient=pd.to_numeric(part.gradient,errors='coerce')
        diagnostics.append({'数据':kind,'方法':labels[method],'成功/起点':f'{good.sum()}/{len(part)}',
            '域拒绝':pd.to_numeric(part.domain_rejections,errors='coerce').sum(),
            '梯度中位':gradient.median(),'梯度最大':gradient.max(),
            'flow代价中位':pd.to_numeric(part.flow_prior_cost,errors='coerce').median()})
    body+=table(pd.DataFrame(diagnostics))+'\n\n'
    body+='## Driver DC/AC、初态和极端持续时间\n\nflow低于0.5、高于2和低于0.1的时间仅作预声明描述阈值（采样数×dt），不是正常范围或连续域边界。失败末有效状态仅留在逐条state_ranges，不混入成功条件均值。DC误差、去各自均值后的AC误差分列；后者不是校正预测或主评分。\n\n'
    for mode in data['mode'].unique():
        part=pd.DataFrame(recovery);part=part[part['mode'].eq(mode)] if len(part) else part
        if len(part):body+=f'合成 / {mode}：\n\n'+table(part.drop(columns='mode').replace({'method':labels}).rename(columns={'method':'方法','completed':'成功','expected':'计划','driver_dc_error':'DC RMSE','driver_ac_rmse':'AC RMSE','initial_rmse':'初态RMSE','truth_rmse_f':'f RMSE','driver_dc_signed_mean':'DC均偏差','flow_signed_mean':'f均偏差'}))+'\n\n'
    fig,axes=plt.subplots(1,3,figsize=(13,4),squeeze=False)
    for ax,field,title in zip(axes[0],['flow_below_half_s','flow_above_two_s','flow_below_01_s'],['f<0.5','f>2','f<0.1']):
        for x,method in enumerate(config['methods']):
            p=states[states.method.eq(method)&states['mode'].eq('full')]
            for success,marker in [(True,'o'),(False,'x')]:
                q=p[p.converged.eq(success)&p.finite_states]
                ax.scatter(np.full(len(q),x),q[field],marker=marker,s=12,alpha=.45)
        ax.set_xticks(range(len(labels)),list(labels.values()),rotation=35,ha='right',fontsize=7)
        ax.set_title(title);ax.set_ylabel('路径采样持续时间 / 秒')
    save(fig,'flow_extreme_durations','全观测路径极端持续时间；圆为成功、叉为失败末有效状态。固定阈值只作工程诊断，不是正常生理范围。')
    synthetic=states[states.kind.eq('synthetic')&states['mode'].eq('full')]
    if len(synthetic):
        fig,axes=plt.subplots(1,3,figsize=(13,4))
        for ax,field,title in zip(axes,['driver_dc_error','driver_ac_rmse','flow_mean_error'],['driver DC有符号误差','driver AC RMSE','flow有符号均误差']):
            for mi,method in enumerate(config['methods']):
                part=synthetic[synthetic.method.eq(method)]
                for ok,marker in [(True,'o'),(False,'x')]:
                    good=part[part.converged.eq(ok)&part.finite_states]
                    ax.scatter(mi+np.linspace(-.15,.15,len(good)),good[field],marker=marker,s=13,alpha=.55)
            ax.axhline(0,color='.5',ls=':',lw=.7);ax.set_title(title)
            ax.set_xticks(range(len(labels)),list(labels.values()),rotation=35,ha='right',fontsize=7)
        save(fig,'flow_synthetic_bias','合成全观测driver DC、AC与flow均值偏差；无幅度/时间对齐，所有固定身份保留。圆为成功、叉为失败末有效状态。DC与flow保留有符号误差，AC单列不能替代原坐标评分。')
    if selected:
        fig,axes=plt.subplots(2,len(selected),figsize=(12,7),squeeze=False,sharex='col')
        for col,(rank,example) in enumerate(selected):
            for mi,method in enumerate(config['methods']):
                path=run/'cells'/f'{example.group}__{method}__full'/'trajectories.npz'
                record=data[(data.group==example.group)&(data.trial==example.trial)&(data.method==method)&data['mode'].eq('full')]
                success=len(record) and record.status.iloc[0]=='completed' and truth(record.converged).iloc[0]
                if not path.exists():
                    for ax in axes[:,col]:ax.plot([],[],':',label=labels[method]+' 未估计')
                    continue
                with np.load(path,allow_pickle=False) as arrays:
                    index=np.flatnonzero(arrays['trial_indices']==int(example.trial))
                    if len(index)!=1:raise ValueError('Fixed parent example missing')
                    state=arrays['states'][index[0]];t=np.arange(len(state))*config['tensor']['dt_s']-(5 if example.kind=='measured' else 0)
                    suffix='' if success else (' 失败迭代' if np.isfinite(state).any() else ' 未估计')
                    for ax,k in zip(axes[:,col],[0,2]):ax.plot(t,state[:,k],color=plt.get_cmap('tab10')(mi),ls='-' if success else ':',lw=.9,label=labels[method]+suffix)
            axes[0,col].set_title(f'{rank}: {example.group} / trial {int(example.trial)}',fontsize=9)
            axes[0,col].set_ylabel('driver r');axes[1,col].set_ylabel('flow f');axes[1,col].set_xlabel('相对事件秒' if example.kind=='measured' else '秒')
            axes[0,col].legend(fontsize=7,ncol=2)
        save(fig,'flow_fixed_parent_examples','固定基线median/worst身份的各臂driver/flow横比，无对齐。失败末有效迭代虚线，缺失不换样。')
    return body


def fixed_roi_initial_details(run,out,config,data,states,labels,save):
    """Prespecified truth strata expose tied-model misspecification, not arm selection."""
    truth_by_group={}
    for group in data.group.unique():
        prep=json.loads((run/'prepared'/f'{group}.json').read_text())
        truth_by_group[group]=prep['spec']
    synthetic=data[data.kind.eq('synthetic')].copy()
    synthetic['true_ratio']=synthetic.group.map(lambda g:truth_by_group[g]['true_initial_log_p_over_v'])
    synthetic['true_tau']=synthetic.group.map(lambda g:truth_by_group[g]['true_tau'])
    parameters=pd.read_csv(out/'parameters.csv')
    parameter_map={(r.group,r.method):r.tau for r in parameters.itertuples()}
    synthetic['tau_estimate']=[parameter_map.get((r.group,r.method),np.nan) for r in synthetic.itertuples()]
    def summarize(part, ratio, method, mode, tau=None):
        good=part[part.status.eq('completed')&truth(part.converged)]
        state=states[states.kind.eq('synthetic')&states.method.eq(method)&states['mode'].eq(mode)&states.group.isin(part.group.unique())&states.converged&states.finite_states]
        row=dict(true_initial_log_p_over_v=ratio,true_tau=tau,method=method,mode=mode,expected=len(part),completed=len(good))
        for component in COMPONENTS:
            field='nmse_'+component
            row['nrmse_'+component]=np.sqrt(good.groupby(['subject','session'])[field].mean().groupby('subject').mean().mean()) if field in good else np.nan
        row['tau_rmse']=np.sqrt(np.mean((good.tau_estimate-good.true_tau)**2)) if len(good) else np.nan
        for field in ['truth_rmse_r','truth_rmse_f','truth_rmse_p','truth_rmse_v','driver_dc_error','driver_ac_rmse']:
            values=pd.to_numeric(state.get(field,pd.Series(dtype=float)),errors='coerce')
            row[field]=np.sqrt(np.mean(values**2)) if len(values) else np.nan
        row['initial_logratio_rmse']=np.sqrt(np.mean((state.initial_log_p_over_v-ratio)**2)) if len(state) else np.nan
        row['path_max_abs_p_minus_v']=state.path_max_abs_p_minus_v.max() if len(state) else np.nan
        return row
    strata=[];detailed=[]
    for (ratio,method,mode),part in synthetic.groupby(['true_ratio','method','mode'],sort=True):
        strata.append(summarize(part,ratio,method,mode))
        for tau,cell in part.groupby('true_tau',sort=True):detailed.append(summarize(cell,ratio,method,mode,tau))
    strata=pd.DataFrame(strata);detailed=pd.DataFrame(detailed)
    strata.to_csv(out/'initial_truth_strata.csv',index=False);detailed.to_csv(out/'initial_truth_tau_strata.csv',index=False)
    invariance=states[states.finite_states].copy();invariance['tied']=invariance.method.isin(config['initial_constraint']['tied_methods'])
    invariance[['kind','group','method','mode','trial','status','converged','tied','initial_log_p_over_v','path_max_abs_p_minus_v']].to_csv(out/'initial_path_invariance.csv',index=False)
    if invariance.loc[invariance.tied,'path_max_abs_p_minus_v'].gt(1e-8).any():raise ValueError('Tied p=v invariant violated by saved trajectory')
    body=('\n## 初态约束：预声明真值分层\n\n'
        '四臂均固定flow权重1；比较固定τ/共享τ各自的free与tied初态。tied在降维优化坐标中令p0=v0，'
        '每trial自由度125→124，仍保留原五个物理初态惩罚项（v与p两项都计入）；不是投影旧解，也不改变动力学。'
        '合成36组为本轮全新生成/拟合；仅实测free两臂继承，不构成独立复制。'
        'true log(p0/v0)=0是约束匹配层，±0.05是预声明错设压力层，不是人口正常范围。'
        '必须分别看这三层及τ=1/2/4，不能用合并合成均值掩盖错设偏差。'
        '以下表按初态层分开；initial_truth_tau_strata.csv进一步保留每个τ、模式、方法的误差与完整失败分母。'
        '局部信息来自各自有效降维参数空间，不是置信区间，也不能将减少自由度后的信息变化当作额外数据。\n\n')
    for ratio in config['synthetic']['initial_log_p_over_v']:
        body+=f'### 真log(p0/v0)={ratio:+.2f}：'+('匹配' if ratio==0 else '错设压力')+'\n\n'
        for mode in config['modes']:
            part=strata[strata.true_initial_log_p_over_v.eq(ratio)&strata['mode'].eq(mode)].set_index('method').reindex(config['methods'])
            body+=f'{mode}观测重建（各τ等身份合并，真τ细分见CSV/图）：\n\n'+table(pd.DataFrame([{'方法':labels[m],'成功/计划':f'{int(r.completed)}/{int(r.expected)}','EEG':r.nrmse_EEG,'HbO':r.nrmse_HbO,'HbR':r.nrmse_HbR} for m,r in part.iterrows()]))+'\n\n'
            body+='原坐标状态/参数恢复（成功条件；DC/AC不可替代原driver误差）：\n\n'+table(pd.DataFrame([{'方法':labels[m],'τ RMSE':r.tau_rmse,'r RMSE':r.truth_rmse_r,'f RMSE':r.truth_rmse_f,'p RMSE':r.truth_rmse_p,'v RMSE':r.truth_rmse_v,'初态log比RMSE':r.initial_logratio_rmse} for m,r in part.iterrows()]))+'\n\n'
    body+='各初态层的driver DC误差RMS与AC RMSE另列于initial_truth_strata.csv及initial_truth_tau_strata.csv，保留模式与真τ，不将DC偏差用去均值消除。\n\n'
    invariant_summary=[]
    for (kind,method),part in invariance.groupby(['kind','method'],sort=False):
        invariant_summary.append({'数据':kind,'方法':labels[method],'有限轨迹数':len(part),'成功数':int(part.converged.sum()),'全路径最大绝对p−v':f'{part.path_max_abs_p_minus_v.max():.3e}'})
    body+='### 全路径约束核验\n\n当前方程下p0=v0应保持p=v；下表包含所有有限末有效轨迹，失败仍不进入主评分。tied轨迹最大绝对差超过1e−8会拒绝生成报告，不将偏离静默当成约束已执行。\n\n'+table(pd.DataFrame(invariant_summary))+'\n\n'
    full=detailed[detailed['mode'].eq('full')]
    for fields,name in [(['tau_rmse','truth_rmse_r','truth_rmse_f'],'initial_tau_driver_flow_recovery'),(['truth_rmse_p','truth_rmse_v'],'initial_p_v_recovery')]:
        fig,axes=plt.subplots(3,len(fields),figsize=(12,10),squeeze=False)
        for row,ratio in enumerate(config['synthetic']['initial_log_p_over_v']):
            for col,field in enumerate(fields):
                ax=axes[row,col]
                for method in config['methods']:
                    part=full[full.true_initial_log_p_over_v.eq(ratio)&full.method.eq(method)].sort_values('true_tau')
                    ax.plot(part.true_tau,part[field],'o-',label=labels[method],lw=1,ms=4)
                ax.set_title(f'真log(p0/v0)={ratio:+.2f} / {field}',fontsize=9);ax.set_xlabel('真τ / 秒');ax.set_ylabel('成功条件RMSE');ax.legend(fontsize=6)
        save(fig,name,'合成full按初态真值与τ同时分层，展示匹配层与错设层；失败完整分母见表/CSV。不跨层挑方法，不以总体平均掩盖偏差。')
    for tau in config['synthetic']['taus_s']:
        fig,axes=plt.subplots(4,3,figsize=(13,11),sharex='col')
        for col,ratio in enumerate(config['synthetic']['initial_log_p_over_v']):
            group=f'synthetic_tau{tau:g}_slow_r0_logpv{ratio:+.2f}';trial=0;path=run/'prepared'/f'{group}.npz'
            axes[0,col].set_title(f'τ={tau:g} / log(p0/v0)={ratio:+.2f}',fontsize=9)
            if not path.exists():
                for ax in axes[:,col]:ax.text(.5,.5,'固定身份缺失，不换样',ha='center',transform=ax.transAxes)
                continue
            with np.load(path,allow_pickle=False) as a:
                t=np.arange(a['target'].shape[1])*config['tensor']['dt_s'];true_state=a['truth_states'][trial]
                axes[0,col].plot(t,a['target'][trial,:,1]+a['target'][trial,:,2],color='.65',lw=.7,label='含噪HbT目标')
                values=[a['clean'][trial,:,1]+a['clean'][trial,:,2],true_state[:,0],true_state[:,2],true_state[:,4]-true_state[:,3]]
                for ax,value in zip(axes[:,col],values):ax.plot(t,value,'k',lw=1.4,label='真值')
            for method in config['methods']:
                path=run/'cells'/f'{group}__{method}__full'/'trajectories.npz'
                record=data[data.group.eq(group)&data.method.eq(method)&data['mode'].eq('full')&data.trial.eq(0)]
                status=record.status.iloc[0] if len(record) else 'missing'
                if not path.exists():
                    for ax in axes[:,col]:ax.plot([],[],':',label=labels[method]+' 未估计')
                    continue
                with np.load(path,allow_pickle=False) as a:
                    index=np.flatnonzero(a['trial_indices']==0)
                    if len(index)!=1:raise ValueError('Initial truth example trial0 missing')
                    k=index[0];state=a['states'][k];values=[a['prediction'][k,:,1]+a['prediction'][k,:,2],state[:,0],state[:,2],state[:,4]-state[:,3]]
                    for ax,value in zip(axes[:,col],values):ax.plot(t,value,'-' if status=='completed' else ':',lw=.85,label=labels[method]+('' if status=='completed' else ' 失败'))
            for ax,label in zip(axes[:,col],['处理HbT','driver r','flow f','p−v']):ax.set_ylabel(label);ax.legend(fontsize=5)
            axes[-1,col].set_xlabel('秒')
        save(fig,f'initial_truth_tau{tau:g}','预声明slow/rep0/trial0，三初态真值层分别展示，四方法与真实driver/flow/p−v比较；无幅度、符号、时间或基线对齐。失败保留，不按效果选样。')
    return body


def render_nonlinear_fit(run, out, *, gain_prior=False, fixed_roi_tau=False, fixed_roi_flow=False, fixed_roi_initial=False):
    """Read only this run's declared cells; keep failed and missing denominators."""
    if out.exists() and any(out.iterdir()):
        raise ValueError("Export directory is not empty; choose a new versioned --output")
    summary = json.loads((run / "summary.json").read_text())
    data = pd.read_csv(run / "metrics.csv")
    required = {"kind", "group", "subject", "session", "method", "mode", "trial", "status", "converged"}
    if data.empty or not required.issubset(data):
        raise ValueError("Missing nonlinear fit row identities/status")
    keys = ["kind", "group", "method", "mode", "trial"]
    if data.duplicated(keys).any() or len(data) != summary["observed_rows"]:
        raise ValueError("Duplicate identities or stale nonlinear summary")
    methods = ["fixed_tau_100", "fixed_tau_0p01", "linear_trained_tau_0p01", "linear_prior_tau_0p01"]
    labels = dict(zip(methods, ["固定τ/λ100", "固定τ/λ0.01", "训练τ/λ0.01", "先验τ/λ0.01"]))
    config = {}
    fixed_roi_flow = fixed_roi_flow or fixed_roi_initial
    fixed_roi_tau = fixed_roi_tau or fixed_roi_flow
    extended = gain_prior or fixed_roi_tau
    if extended:
        import yaml
        config = yaml.safe_load((run / "resolved_config.yaml").read_text())
        expected_schema = "shared_driver_fixed_roi_initial_v1" if fixed_roi_initial else "shared_driver_fixed_roi_flow_v1" if fixed_roi_flow else ("shared_driver_fixed_roi_tau_v1" if fixed_roi_tau else "shared_driver_gain_prior_v1")
        if config.get("schema") != expected_schema:
            raise ValueError("Report mode differs from resolved configuration schema")
        methods = list(config["methods"])
        labels = ({"fixed_roi_fixed_tau":"固定ROI/τ2", "fixed_roi_trained_tau":"固定ROI/训练τ", "fixed_roi_prior_tau":"固定ROI/先验τ"}
            if fixed_roi_tau else {"fixed_gain_no_amplitude":"A 固定β", "fixed_gain_amplitude":"B 固定β+幅度",
                                  "trained_gain_no_amplitude":"C 训练β", "trained_gain_amplitude":"D 训练β+幅度"})
        if fixed_roi_flow:
            labels = {"fixed_roi_fixed_tau":"A τ2/w0", "fixed_roi_trained_tau":"B 学τ/w0",
                      "fixed_tau_logflow1":"C τ2/w1", "trained_tau_logflow01":"D 学τ/w0.1",
                      "trained_tau_logflow1":"E 学τ/w1", "trained_tau_logflow4":"F 学τ/w4"}
        if fixed_roi_initial:
            labels={"fixed_tau_logflow1":"A 固定τ/free", "trained_tau_logflow1":"B 共享τ/free",
                    "fixed_tau_logflow1_tied":"C 固定τ/tied", "trained_tau_logflow1_tied":"D 共享τ/tied"}
        if methods != list(labels) or set(data.method) != set(methods):
            raise ValueError("Report methods differ from resolved configuration")
    parent_data = None
    if fixed_roi_flow:
        parent = Path(config['parent_run'])
        if not parent.is_absolute(): parent = Path(__file__).resolve().parents[2] / parent
        parent_data = pd.read_csv(parent/'metrics.csv')
        identity=['kind','group','subject','session','trial','method','mode']
        inherited=data[data.method.isin(methods[:2]) & (data.kind.eq('measured') if fixed_roi_initial else True)]
        previous=parent_data[parent_data.method.isin(methods[:2])&parent_data.kind.isin(data.kind.unique()) & (parent_data.kind.eq('measured') if fixed_roi_initial else True)]
        check=inherited.merge(previous,on=identity,how='outer',suffixes=('','_parent'),indicator=True,validate='one_to_one')
        if not check['_merge'].eq('both').all():raise ValueError('Inherited baseline identities differ from exact parent')
        for field in ['status','converged']+[c for c in data if c.startswith(('nmse_','nrmse_','last_valid_nmse_'))]:
            if not ((check[field]==check[field+'_parent'])|(check[field].isna()&check[field+'_parent'].isna())).all():
                raise ValueError('Inherited baseline changed '+field)
    out.mkdir(parents=True, exist_ok=True)
    (out / "figures").mkdir()
    font = Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc")
    if font.exists():
        font_manager.fontManager.addfont(str(font))
        plt.rcParams["font.family"] = font_manager.FontProperties(fname=str(font)).get_name()
    plt.rcParams.update({"font.size": 9, "axes.unicode_minus": False})
    figures = []

    def save(fig, name, caption):
        fig.tight_layout()
        fig.savefig(out / "figures" / f"{name}.png", dpi=240, bbox_inches="tight")
        plt.close(fig)
        figures.append((name, caption))

    aggregates = []
    for item in summary["aggregates"]:
        row = {k: item[k] for k in ("kind", "method", "mode", "expected", "observed_rows", "completed", "missing")}
        row["failed"] = item["observed_rows"] - item["completed"]
        row["complete_panel"] = item["completed"] == item["expected"]
        part = data[(data.kind == item["kind"]) & (data.method == item["method"]) & (data["mode"] == item["mode"])]
        success = part.status.eq("completed") & truth(part.converged)
        if len(part) != item["observed_rows"] or int(success.sum()) != item["completed"]:
            raise ValueError("Stale nonlinear aggregate counts")
        if item["mode"] == "full":
            errors = part.reindex(columns=["nrmse_"+c for c in COMPONENTS])
            row["joint_below_05"] = int((success & errors.lt(.5).all(axis=1)).sum())
        else:
            row["joint_below_05"] = "N/A"
        for c in COMPONENTS:
            row["NRMSE_" + c] = item.get("nrmse_" + c, np.nan)
        aggregates.append(row)
    agg = pd.DataFrame(aggregates)
    agg.to_csv(out / "nonlinear_summary.csv", index=False)
    failures = data.groupby(["kind", "method", "mode", "status"], dropna=False).size().rename("rows").reset_index()
    failures.to_csv(out / "failure_counts.csv", index=False)
    successful = data.status.eq("completed") & truth(data.converged)
    convergence_columns = [c for c in ["kind", "group", "method", "mode", "trial", "status", "converged", "evaluations", "selected_start", "gradient_inf_norm", "scaled_gradient_inf_norm", "rejected_steps", "domain_rejections", "derivative_rejections", "flow_prior_cost", "flow_prior_weight", "flow_prior_log_sd", "integration_max_difference_training_sd"] if c in data]
    data[convergence_columns].to_csv(out / "convergence.csv", index=False)

    kinds = list(agg.kind.unique())
    fig, axes = plt.subplots(len(kinds), 3, figsize=(13, 4 * len(kinds)), squeeze=False)
    for i, kind in enumerate(kinds):
        for j, mode in enumerate(["full", "center_EEG", "center_fNIRS"]):
            ax = axes[i, j]
            part = agg[(agg.kind == kind) & (agg["mode"] == mode)].set_index("method").reindex(methods)
            columns = COMPONENTS if mode == "full" else (["EEG"] if mode == "center_EEG" else ["HbO", "HbR"])
            x = np.arange(len(methods))
            for k, c in enumerate(columns):
                ax.bar(x + (k - (len(columns)-1)/2)*.22, part["NRMSE_"+c], .22, label=c)
            ax.set_xticks(x, [labels[m] + f"\n{int(part.loc[m, 'completed'])}/{int(part.loc[m, 'expected'])}" for m in methods], rotation=22, ha="right", fontsize=7)
            ax.axhline(.5, color="red", ls="--", lw=1)
            ax.set_title(f"{kind} / {mode}")
            ax.set_ylabel("成功子集 NRMSE"); ax.legend()
    save(fig, "nonlinear_errors", ("预声明四方法" if fixed_roi_initial else "预声明六方法" if fixed_roi_flow else ("预声明三方法" if fixed_roi_tau else "预声明四方法")) + "与相同模式比较；标注成功/计划分母。柱高仅成功子集，缺失或失败不能视作零误差。")

    # Baseline comparisons use the same successful identities, not different subsets.
    paired = []
    baseline = data[data.method == methods[0]]
    join = ["kind", "group", "mode", "trial"]
    for method in methods[1:]:
        merged = data[data.method == method].merge(baseline, on=join, suffixes=("", "_baseline"), validate="one_to_one")
        mask = merged.status.eq("completed") & truth(merged.converged) & merged.status_baseline.eq("completed") & truth(merged.converged_baseline)
        for (kind, mode), all_rows in merged.groupby(["kind", "mode"]):
            part = all_rows[mask.loc[all_rows.index]]
            expected = int(agg[(agg.kind == kind) & (agg.method == method) & (agg["mode"] == mode)].expected.iloc[0])
            row = dict(kind=kind, method=method, mode=mode, common_success=len(part), expected=expected)
            for c in COMPONENTS:
                field = "nmse_" + c
                if field in part and field+"_baseline" in part:
                    scores = part[["subject", "session"]].copy()
                    scores["delta"] = part[field] - part[field+"_baseline"]
                    row["delta_NMSE_"+c] = scores.groupby(["subject", "session"]).delta.mean().groupby("subject").mean().mean()
                    for suffix, label in [("", "candidate"), ("_baseline", "baseline")]:
                        value = part.groupby(["subject", "session"])[field+suffix].mean().groupby("subject").mean().mean()
                        row[label+"_NRMSE_"+c] = np.sqrt(value)
            paired.append(row)
    paired = pd.DataFrame(paired)
    paired.to_csv(out / "paired_baseline.csv", index=False)

    # Saved states include last-valid iterates of failed fits; label them explicitly.
    import importlib.util
    import sys
    source_root = run / "source_snapshot"
    if not source_root.is_dir():
        raise ValueError("Nonlinear report requires the run's frozen source_snapshot")
    operator_path = source_root / "src/inference/observation_baselines.py"
    sys.path.insert(0, str(source_root))
    operator_spec = importlib.util.spec_from_file_location("report_frozen_observation_baselines", operator_path)
    operator_module = importlib.util.module_from_spec(operator_spec)
    operator_spec.loader.exec_module(operator_module)
    native_feature_operators = operator_module.native_feature_operators
    visible_feature_interpolation = operator_module.visible_feature_interpolation
    state_rows, parameter_rows, profiles, own_rows, observation_rows = [], [], [], [], []
    for group in data.group.unique():
        preparation = json.loads((run / "prepared" / f"{group}.json").read_text())
        spec = preparation["spec"]
        immutable_target = immutable_normalizer = None
        if preparation["status"] == "completed":
            meta = preparation["metadata"]
            if spec["kind"] == "measured":
                projection = meta["projection"]
                observation_rows.append(dict(group=group, Hb_pair_index=projection["fnirs_pair"],
                    EEG_factor=projection["eeg_factor"], Hb_factor=projection["fnirs_factor"],
                    absolute_calibration=projection["observation_loading"]["absolute_calibration"]))
            with np.load(run / "prepared" / f"{group}.npz", allow_pickle=False) as a:
                if extended:
                    immutable_target, immutable_normalizer = a["target"].copy(), a["normalizer"].copy()
                n = a["target"].shape[1]; left, right = (n-16)//2, (n+16)//2
                op = native_feature_operators(n)
                trials = data.loc[data.group.eq(group), "trial"].unique()
                for mode in ("center_EEG", "center_fNIRS"):
                    eeg_gap = (left, right) if mode == "center_EEG" else (0, 0)
                    hb_gap = (int(2.5*left), int(2.5*right)) if mode == "center_fNIRS" else (0, 0)
                    pe = op["eeg"] @ visible_feature_interpolation(n, *eeg_gap)
                    ph = op["fnirs"] @ visible_feature_interpolation(int(2.5*n), *hb_gap)
                    columns = [0] if mode == "center_EEG" else [1, 2]
                    seen = np.r_[0:left, right:n]
                    for trial in trials:
                        pred = np.column_stack((pe @ a["feature_eeg"][trial], ph @ a["feature_fnirs"][trial]))
                        row = dict(kind=spec["kind"], group=group, mode=mode, trial=int(trial), subject=spec["subject"], session=meta["trials"][trial]["session"])
                        for col in columns:
                            pred[:, col] = np.interp(np.arange(n), seen, pred[seen, col])
                            row["nmse_"+COMPONENTS[col]] = float(np.mean(((pred[left:right,col]-a["target"][trial,left:right,col])/a["normalizer"][col])**2))
                        own_rows.append(row)
        parameter_values = preparation.get("selected_tau", {})
        if fixed_roi_tau:
            parameter_values = {methods[0]: config['reference_tau_s']}
            for method in methods[1:]:
                if method in ("fixed_tau_logflow1","fixed_tau_logflow1_tied"):
                    parameter_values[method] = config["reference_tau_s"]
                    continue
                selection_path = run/'training'/f'{group}__{method}'/'selection.json'
                selection = json.loads(selection_path.read_text()) if selection_path.exists() else {}
                parameter_values[method] = selection.get('parameter_value') if selection.get('status') == 'completed' else np.nan
        for method, tau in parameter_values.items():
            parameter_rows.append(dict(kind=spec["kind"], group=group, subject=spec["subject"], outer=spec["outer"], method=method, tau=tau, true_tau=spec.get("true_tau", np.nan)))
        for point in preparation.get("profiles", []):
            profiles.append(dict(group=group, kind=spec["kind"], **point))
        for method, mode in data[data.group == group][["method", "mode"]].drop_duplicates().itertuples(index=False, name=None):
            cell = run / "cells" / f"{group}__{method}__{mode}" / "trajectories.npz"
            if not cell.exists():
                continue
            with np.load(cell, allow_pickle=False) as arrays:
                if extended and immutable_target is not None:
                    if not np.array_equal(arrays["target"], immutable_target[arrays["trial_indices"]], equal_nan=True) or not np.array_equal(arrays["normalizer"], immutable_normalizer, equal_nan=True):
                        raise ValueError("Gain/prior cell changed frozen target or scoring scale")
                for index, trial in enumerate(arrays["trial_indices"]):
                    identity = data[(data.group == group) & (data.method == method) & (data["mode"] == mode) & (data.trial == trial)].iloc[0]
                    state = arrays["states"][index]
                    row = dict(kind=identity.kind, group=group, method=method, mode=mode, trial=int(trial), status=identity.status, true_tau=spec.get('true_tau'),true_initial_log_p_over_v=spec.get('true_initial_log_p_over_v'),converged=bool(successful.loc[identity.name]), finite_states=bool(np.isfinite(state).all()))
                    if row["finite_states"]:
                        for column, name in enumerate(["r", "s", "f", "v", "p", "q"]):
                            row[name+"_initial"] = state[0, column]
                            row[name+"_minimum"] = state[:, column].min()
                            row[name+"_maximum"] = state[:, column].max()
                        if extended:
                            r = state[:, 0]
                            if fixed_roi_flow:
                                dt=config['tensor']['dt_s']
                                row.update(driver_ac_rms=float(np.std(r)),flow_below_half_s=float(np.sum(state[:,2]<.5)*dt),
                                    flow_above_two_s=float(np.sum(state[:,2]>2)*dt),flow_below_01_s=float(np.sum(state[:,2]<.1)*dt))
                            center = slice((len(r)-16)//2, (len(r)+16)//2)
                            row.update(driver_mean=float(r.mean()), driver_rms=float(np.sqrt(np.mean(r*r))),
                                driver_center_rms=float(np.sqrt(np.mean(r[center]**2))),
                                driver_difference_rms=float(np.sqrt(np.mean(np.diff(r)**2))),
                                driver_center_difference_rms=float(np.sqrt(np.mean(np.diff(r[center])**2))))
                            for column, name in enumerate(["r", "s", "f", "v", "p", "q"]):
                                row[name+"_center_minimum"] = state[center,column].min()
                                row[name+"_center_maximum"] = state[center,column].max()
                        if fixed_roi_initial:
                            row['path_max_abs_p_minus_v']=float(np.max(np.abs(state[:,4]-state[:,3])))
                            row['initial_log_p_over_v']=float(np.log(state[0,4]/state[0,3]))
                        row["initial_max_fractional_excursion"] = np.max(abs(state[0, 2:] - 1))
                        row["path_max_fractional_excursion"] = np.max(abs(state[:, 2:] - 1))
                    if spec["kind"] == "synthetic" and row["finite_states"] and "truth_states" in arrays:
                        if fixed_roi_flow:
                            actual=arrays['truth_states'][index]; residual=state-actual
                            row.update(driver_dc_error=float(residual[:,0].mean()),flow_mean_error=float(residual[:,2].mean()),
                                driver_ac_rmse=float(np.std(residual[:,0])),initial_rmse=float(np.sqrt(np.mean(residual[0,1:]**2))))
                        for column, name in enumerate(["r", "s", "f", "v", "p", "q"]):
                            row["truth_rmse_"+name] = float(np.sqrt(np.mean((state[:,column]-arrays["truth_states"][index,:,column])**2)))
                    state_rows.append(row)
    states = pd.DataFrame(state_rows)
    states.to_csv(out / "state_ranges.csv", index=False)
    observation_map = pd.DataFrame(observation_rows)
    observation_map.to_csv(out / "observation_map.csv", index=False)
    state_summary = []
    state_kind = "measured" if "measured" in data.kind.unique() else "synthetic"
    for method in methods:
        part = states[(states.kind == state_kind) & (states.method == method)
                      & (states["mode"] == "full") & states.converged]
        state_summary.append(dict(method=method, successful=len(part), expected=int(agg.loc[
            (agg.kind == state_kind) & (agg.method == method) & (agg["mode"] == "full"), "expected"].sum()),
            f_min=part.f_minimum.min() if len(part) else np.nan,
            f_max=part.f_maximum.max() if len(part) else np.nan,
            f_below_01=int(part.f_minimum.lt(.1).sum()) if len(part) else 0,
            initial_over_05=int(part.initial_max_fractional_excursion.gt(.5).sum()) if len(part) else 0,
            path_over_05=int(part.path_max_fractional_excursion.gt(.5).sum()) if len(part) else 0))
    state_summary = pd.DataFrame(state_summary)
    state_summary.to_csv(out / "state_summary.csv", index=False)
    parameters = pd.DataFrame(parameter_rows, columns=["kind", "group", "subject", "outer", "method", "tau", "true_tau"])
    parameters.to_csv(out / "parameters.csv", index=False)
    pd.DataFrame(profiles).to_csv(out / "tau_profiles.csv", index=False)
    if not extended:
        fig, axes = plt.subplots(2, 1, figsize=(11, 7))
        measured_groups = list(parameters.loc[parameters.kind.eq("measured"), "group"].unique())
        for method in methods:
            p = parameters[(parameters.kind == "measured") & (parameters.method == method)].set_index("group")
            if len(p):
                axes[0].plot(range(len(measured_groups)), p.reindex(measured_groups).tau, "o-", label=labels[method])
        axes[0].set_xticks(range(len(measured_groups)), measured_groups, rotation=45, ha="right", fontsize=7)
        axes[0].set_ylabel("τ / s")
        if measured_groups: axes[0].legend()
        for group in measured_groups:
            points = [p for p in profiles if p["group"] == group]
            if points:
                points = sorted(points, key=lambda p:p["tau"])
                values = np.array([p["train_penalized"] for p in points])
                axes[1].plot([p["tau"] for p in points], values-values.min(), label=group)
        axes[1].set_xlabel("τ / s"); axes[1].set_ylabel("训练线性目标 − 最小值")
        if measured_groups: axes[1].legend(fontsize=6, ncol=3)
        save(fig, "nonlinear_tau", "τ由训练线性剖面冻结后用于非线性拟合；不是非线性τ估计或似然置信区间。")

    fig, axes = plt.subplots(1, 2, figsize=(11, 4))
    for ax, field in zip(axes, ["initial_max_fractional_excursion", "path_max_fractional_excursion"]):
        for i, method in enumerate(methods):
            part = states[(states.kind == state_kind) & (states.method == method) & (states["mode"] == "full")]
            if field not in part: continue
            for ok, marker in [(True, "o"), (False, "x")]:
                values = part.loc[part.converged.eq(ok), field].dropna().to_numpy()
                ax.scatter(i + np.linspace(-.12,.12,len(values)), values, marker=marker, s=12, label="成功" if i==0 and ok else ("失败最后有效迭代" if i==0 else None))
        ax.set_xticks(range(len(methods)), [labels[m] for m in methods], rotation=20, ha="right", fontsize=7)
        ax.set_ylabel(field); ax.legend(fontsize=7)
    save(fig, "nonlinear_states", "初态与全路径f/v/p/q相对静息偏离；失败的最后有效迭代用叉号保留。0.25小信号阈值不作为非线性生理合格界。")

    notes = []
    example_source = (parent_data[parent_data.kind.eq('measured')] if 'measured' in data.kind.unique() else data) if fixed_roi_initial else (parent_data[parent_data.kind.isin(data.kind.unique())] if fixed_roi_flow else data)
    examples = example_source[(example_source.kind == "measured") & (example_source.method == methods[0]) & (example_source["mode"] == "full")].copy()
    if examples.empty:
        examples = example_source[(example_source.method == methods[0]) & (example_source["mode"] == "full")].copy()
    error_fields = ["last_valid_nmse_"+c for c in COMPONENTS]
    examples["rank_error"] = examples.reindex(columns=error_fields).mean(axis=1)
    examples = examples.sort_values(["rank_error", "group", "trial"], kind="stable", na_position="last")
    selected = [("median", examples.iloc[(len(examples)-1)//2]), ("worst", examples.iloc[-1])] if len(examples) else []
    for label, example in selected:
        group, trial = example.group, int(example.trial)
        fig, axes = plt.subplots(3, 2, figsize=(12, 8), sharex=True)
        target_drawn = False
        for method in methods:
            path = run / "cells" / f"{group}__{method}__full" / "trajectories.npz"
            row = data[(data.group==group)&(data.trial==trial)&(data.method==method)&(data["mode"]=="full")]
            status = row.status.iloc[0] if len(row) else "missing"
            notes.append(f"{label} / {group} / trial {trial} / {method}: {status}")
            if not path.exists():
                if extended:
                    for ax in axes[:,0]:
                        ax.plot([], [], ":", label=labels[method]+f" 未估计 ({status})")
                    if method == methods[-1]:
                        for ax in axes[:,1]:
                            ax.text(.5, .5, f"{labels[method]}状态未估计\n{status} / 无轨迹文件",
                                ha="center", va="center", transform=ax.transAxes, fontsize=9)
                continue
            with np.load(path, allow_pickle=False) as a:
                indices = np.flatnonzero(a["trial_indices"] == trial)
                if not len(indices): raise ValueError("Representative trial missing from cell")
                index = int(indices[0]); t = np.arange(a["prediction"].shape[1])*.25
                if example.kind == "measured": t -= 5.
                if "target" in a and not target_drawn:
                    for j,c in enumerate(COMPONENTS): axes[j,0].plot(t,a["target"][index,:,j],"k",lw=1.5,label="实测／目标")
                    target_drawn = True
                style = "-" if status == "completed" else ":"
                for j,c in enumerate(COMPONENTS):
                    unavailable = extended and not np.isfinite(a["prediction"][index,:,j]).any()
                    suffix = f" 未估计 ({status})" if unavailable else (" (失败迭代)" if status!="completed" else "")
                    axes[j,0].plot(t,a["prediction"][index,:,j],style,lw=1,label=labels[method]+suffix)
                    axes[j,0].set_ylabel(c)
                if method == (methods[-1] if extended else "linear_trained_tau_0p01"):
                    if extended and not np.isfinite(a["states"][index]).any():
                        for ax in axes[:,1]:
                            ax.text(.5, .5, f"{labels[method]}状态未估计\n{status} / 无有限状态",
                                ha="center", va="center", transform=ax.transAxes, fontsize=9)
                    for j in range(3):
                        for k in [2*j,2*j+1]: axes[j,1].plot(t,a["states"][index,:,k],style,label=["r","s","f","v","p","q"][k])
                        axes[j,1].legend()
        axes[0,0].legend(fontsize=7)
        if example.kind == "measured":
            for ax in axes.flat: ax.axvline(0., color="grey", lw=.7, ls=":")
        for ax in axes[-1]: ax.set_xlabel("相对事件 / 秒" if example.kind == "measured" else "秒")
        fig.suptitle(f"{label}: {group} / trial {trial}; {labels[methods[0]]}排序")
        save(fig, "nonlinear_curves_"+label, f"{labels[methods[0]]}基线的{label}样例；右侧{labels[methods[-1]] if extended else '训练τ弱正则'}状态，失败最后有效迭代为虚线，不按候选性能换样。")

    own = pd.DataFrame(own_rows)
    own_comparisons = []
    if len(own):
        own.to_csv(out / "own_context_rows.csv", index=False)
        for method in methods:
            merged = data[(data.method == method) & successful].merge(own, on=["kind", "group", "mode", "trial", "subject", "session"], suffixes=("", "_own"), validate="one_to_one")
            for (kind, mode), part in merged.groupby(["kind", "mode"]):
                expected = int(agg[(agg.kind == kind)&(agg.method == method)&(agg["mode"] == mode)].expected.iloc[0])
                row = dict(kind=kind, method=method, mode=mode, common_success=len(part), expected=expected)
                for c in COMPONENTS:
                    for suffix, name in [("", "candidate"), ("_own", "own_context")]:
                        field = "nmse_"+c+suffix
                        if field in part:
                            value = part.groupby(["subject", "session"])[field].mean().groupby("subject").mean().mean()
                            row[name+"_NRMSE_"+c] = np.sqrt(value)
                own_comparisons.append(row)
    own_comparisons = pd.DataFrame(own_comparisons)
    own_comparisons.to_csv(out / "paired_own_context.csv", index=False)
    synthetic = data[data.kind.eq("synthetic")].copy()
    synthetic["successful"] = successful.loc[synthetic.index]
    recovery = []
    for (method, mode), part in synthetic.groupby(["method", "mode"]):
        good = part[part.successful]
        values = pd.to_numeric(good.get("driver_nrmse", pd.Series(dtype=float)), errors="coerce").dropna()
        recovery.append(dict(method=method, mode=mode, observed=len(part), completed=len(good), driver_scored=len(values), driver_mean=values.mean(), driver_max=values.max()))
    recovery = pd.DataFrame(recovery)
    recovery.to_csv(out / "synthetic_recovery.csv", index=False)
    # Communication tables are deliberately narrower than the retained CSV schema.
    kind_labels = {"synthetic": "合成", "measured": "实测"}
    mode_labels = {"full": "全观测", "center_EEG": "中心EEG隐藏", "center_fNIRS": "中心Hb隐藏"}

    def compact_group(value):
        return str(value).replace("subject_", "S").replace("synthetic_tau", "T").replace("mixed", "混合").replace("slow", "慢变")

    def section_tables(frame, builder):
        chunks = []
        for (kind, mode), part in frame.groupby(["kind", "mode"], sort=False):
            chunks.append(f"### {kind_labels.get(kind,kind)} · {mode_labels.get(mode,mode)}\n\n" + table(builder(part)))
        return "\n\n".join(chunks)

    def aggregate_display(part):
        rows = []
        for _, r in part.iterrows():
            rows.append({"方法": labels[r.method], "成功/计划": f"{r.completed}/{r.expected}",
                "失败/缺失": f"{r.failed}/{r.missing}", "三项<0.5": r.joint_below_05,
                **{c: r["NRMSE_"+c] for c in COMPONENTS}})
        return pd.DataFrame(rows)

    def paired_display(part, own=False):
        rows = []
        for _, r in part.iterrows():
            for c in COMPONENTS:
                candidate = r.get("candidate_NRMSE_"+c, np.nan)
                reference = r.get(("own_context" if own else "baseline")+"_NRMSE_"+c, np.nan)
                if not np.isfinite(candidate) and not np.isfinite(reference):
                    continue
                row = {"方法": labels[r.method], "分量": c, "共同/计划": f"{r.common_success}/{r.expected}",
                    "候选": candidate, "自身上下文" if own else (labels[methods[0]] if extended else "固定强正则"): reference}
                if not own: row["ΔNMSE"] = r.get("delta_NMSE_"+c, np.nan)
                rows.append(row)
        return pd.DataFrame(rows)

    aggregate_text = section_tables(agg, aggregate_display)
    paired_text = section_tables(paired, paired_display)
    own_text = section_tables(own_comparisons, lambda part: paired_display(part, own=True)) if len(own_comparisons) else ""
    failure_text = section_tables(failures, lambda part: part.assign(
        method=part.method.map(labels), status=part.status.map({"completed":"成功", "failed_numerical":"数值失败",
        "failed_preparation":"准备失败", "failed_training":"训练失败", "failed_exception":"异常", "failed_integration_check":"积分失败"}).fillna(part.status)
    )[["method","status","rows"]].rename(columns={"method":"方法","status":"状态","rows":"条数"}))
    recovery_text = ""
    for mode, part in recovery.groupby("mode", sort=False):
        display = pd.DataFrame([{"方法":labels[r.method], "成功/记录":f"{r.completed}/{r.observed}",
            "驱动评分数":r.driver_scored, "驱动均误差":r.driver_mean, "驱动最大误差":r.driver_max}
            for r in part.itertuples()])
        recovery_text += f"### 合成 · {mode_labels.get(mode,mode)}\n\n" + table(display) + "\n"
    parameter_text = ""
    for kind, part in parameters.groupby("kind", sort=False):
        rows = []
        for group, values in part.groupby("group", sort=False):
            r = {"被试/折或合成组":compact_group(group), "真τ":values.true_tau.iloc[0]}
            r.update({labels[v.method]:v.tau for v in values.itertuples()})
            rows.append(r)
        parameter_text += f"### {kind_labels.get(kind,kind)} · τ（秒）\n\n" + table(pd.DataFrame(rows)) + "\n"
    observation_display = observation_map.rename(columns={"group":"被试/折", "Hb_pair_index":"Hb对", "EEG_factor":"EEG因子", "Hb_factor":"Hb因子", "absolute_calibration":"绝对标定"}).copy()
    if len(observation_display):
        observation_display["被试/折"] = observation_display["被试/折"].map(compact_group)
        observation_display["绝对标定"] = observation_display["绝对标定"].map({True:"是",False:"否"})
    state_display = pd.DataFrame([{"方法":labels[r.method], "成功/计划":f"{r.successful}/{r.expected}",
        "f最小":r.f_min, "f最大":r.f_max, "f<0.1":r.f_below_01,
        "初态>0.5":r.initial_over_05, "路径>0.5":r.path_over_05} for r in state_summary.itertuples()])

    text = "# 共享驱动非线性拟合：正则敏感性与重建\n\n"
    text += f"来源运行：`{run.name}`。计划cell {summary['expected_cells']}，终态{summary['terminal_cells']}；计划评价{summary['expected_rows']}，已记录{summary['observed_rows']}，收敛成功{summary['completed_rows']}。cell完成不等于其内部拟合成功。\n\n"
    text += "各方法均预先指定；没有按留出结果选胜者。λ100与λ0.01比较使用相同初态惩罚100。两个适配τ方法只使用训练线性剖面，τ在非线性评价时冻结，不声称完成非线性生理参数学习。\n\n"
    text += "## 完整分母与误差\n\n" + aggregate_text + "\n"
    text += "EEG是处理后的log-power/PCA特征，不是原始电压；三分量使用冻结外折训练SD。NRMSE来自成功收敛且积分分辨率检查通过的子集：trial→session→subject等权NMSE后取根。失败/缺失始终保留在expected中；成功数小于计划数时不能将条件均值写成完整面板结果。三项<0.5是逐trial三分量均达阈值且成功的条数；不是由聚合均值推算。模式full为全观测重建，center模式只评分隐藏分量。数学域有效不等同于经验生理合理、独立共享信息或teacher资格。\n\n"
    text += "## 配对基线比较\n\n相对固定τ/λ100，以下NMSE差只使用双方共同成功的相同身份；负值为改善，完整分母同时保留。\n\n" + paired_text + "\n"
    if len(own_comparisons):
        text += "## 自身上下文对照\n\n先在原生特征可见支持插值，再执行原观测处理，最后在处理坐标内插补隐藏段；与既有线性诊断基线一致。算子从本run的source_snapshot加载；只比较候选成功的相同身份，未拟合生理参数。\n\n" + own_text + "\n"
    text += "## 失败与收敛\n\n" + failure_text + "\n完整梯度、预算、起点与积分差异见convergence.csv；每次起点结果仍由原cell/result.json持有。失败最后有效轨迹不进入主误差或合成恢复均值。\n\n"
    text += "## 合成恢复与参数\n\n" + recovery_text + "\n"
    text += parameter_text + "\n合成驱动误差不做事后幅度、符号或时间对齐；合成条件不等同实测生理真值。参数跨折差异受重叠训练集合与通道选择影响，不能直接解释为真实被试差异。\n\n"
    if len(observation_map):
        text += "## 观测映射的解释边界\n\n下表直接取自本run的prepared元数据。EEG与Hb幅度因子来自训练MAD和冻结参考模型的协方差规范，absolute_calibration为False；Hb对索引在部分被试内随折改变。它们没有提供独立跨模态生理标定。固定EEG loading定义驱动单位，再固定neurovascular gain仍会额外固定二者的有效幅度比。\n\n" + table(observation_display) + "\n"
        text += "原方程ds/dt=βr−κs−γ(f−1)、df/dt=s不含τ：固定r与s/f初态时，调τ不能直接修复异常血流；τ只改变下游v/p/q，重新拟合时可通过驱动与初态补偿。下一轮若研究有效增益，必须保留当前评分尺度，并将增益与绝对生理效能区别开。\n\n"
    text += "## 状态范围诊断\n\n" + table(state_display) + "\n统计限成功full轨迹，失败始终保留在完整分母中。f<0.1及相对静息偏离>0.5只是定位极端状态的描述阈值，并非文献验证的生理正常界限。\n\n"
    text += "## 状态与固定样例\n\nstate_ranges.csv逐条保留六状态初值、最小值、最大值及成功标志；合成另保留各状态原坐标真值RMSE。样例按固定τ/λ100最后有效三分量平均NMSE排序取中位和最差；无有限输出排在末尾并明确保留。右侧展示训练τ/λ0.01状态。\n\n" + "\n\n".join(notes) + "\n\n"
    if gain_prior:
        text = text.replace("共享驱动非线性拟合：正则敏感性与重建", "共享驱动非线性拟合：幅度先验与有效增益")
        text = text.replace("各方法均预先指定；没有按留出结果选胜者。λ100与λ0.01比较使用相同初态惩罚100。两个适配τ方法只使用训练线性剖面，τ在非线性评价时冻结，不声称完成非线性生理参数学习。",
            "四臂为预先指定的2×2比较：A固定β，B固定β加驱动幅度先验，C训练共享β，D训练共享β加幅度先验；不按验证结果选胜者。τ固定2秒，曲率惩罚0.01、初态惩罚100。β仅使用每组18条训练数据估计，多起点全部终态后按收敛训练目标选择，并在验证时冻结。")
        text = text.replace("相对固定τ/λ100", "相对A固定β无幅度先验基线")
        text = text.replace("固定EEG loading定义驱动单位，再固定neurovascular gain仍会额外固定二者的有效幅度比。",
            "所有臂的观测算子、目标、均值处理、训练SD与通道映射不变；没有为了β重新调整EEG或Hb观测尺度。β是这套任意观测单位中的有效耦合，不能解释为被试神经血管生理效能。")
        text = text.replace("下一轮若研究有效增益，必须保留当前评分尺度，并将增益与绝对生理效能区别开。", "本轮只释放有效增益，仍须区分动力学补偿与生理恢复。")
        text = text.replace("样例按固定τ/λ100", "样例按A固定β无幅度先验")
        text = text.replace("右侧展示训练τ/λ0.01状态。", "右侧展示D训练β加幅度先验状态。")
        text = text.replace(parameter_text, "τ在所有臂固定为2秒；训练β见下文。\n") if parameter_text else text
        if config.get("continuation"):
            parent_name = Path(config["continuation"]["parent_run"]).name
            text = text.replace("共享驱动非线性拟合：幅度先验与有效增益", "共享驱动非线性拟合：训练预算热启动对照")
            text = text.replace("## 完整分母与误差", (
                f"本轮延续父运行`{parent_name}`：仅符合合同的{config['continuation']['expected_measured_restarts']}条实测训练起点谱系从最后有效解热启动，"
                "LM阻尼重置，目标、观测尺度和成功阈值不变；这不是精确恢复优化器状态。其余实测训练终态、全部合成训练/验证以及实测A/B验证继承；"
                "实测C/D验证按新训练选择重算。继承结果不构成新增独立复制，下表分母仍是原有身份集合。\n\n## 完整分母与误差"))
        text += gain_prior_details(run, out, config, data, states, labels, save, selected)
    elif fixed_roi_tau:
        text = text.replace("共享驱动非线性拟合：正则敏感性与重建", "固定AF7Fp1位置：共享τ与轨迹恢复")
        text = text.replace("各方法均预先指定；没有按留出结果选胜者。λ100与λ0.01比较使用相同初态惩罚100。两个适配τ方法只使用训练线性剖面，τ在非线性评价时冻结，不声称完成非线性生理参数学习。",
            "三臂预先指定：同ROI固定τ2、18条训练轨迹共享非线性τ、带工程log prior的共享τ；β固定1，曲率惩罚0.01、初态惩罚100，无driver幅度ridge。训练起点全部终态后按收敛目标选择，验证τ冻结，不按留出成绩选择方法。")
        text = text.replace("三分量使用冻结外折训练SD", "EEG沿用父训练SD；固定ROI的Hb分量SD仅由当前18条训练轨迹计算")
        text = text.replace("相对固定τ/λ100", "相对同一新ROI目标的固定τ2基线")
        text = text.replace("参数跨折差异受重叠训练集合与通道选择影响", "固定ROI消除了本轮跨折选对变化，但参数跨折差异仍受重叠训练集合、观测幅度规范与初态补偿影响")
        text = text.replace("Hb对索引在部分被试内随折改变", "本轮各被试、各折均固定AF7Fp1（pair0）")
        text = text.replace("下一轮若研究有效增益，必须保留当前评分尺度，并将增益与绝对生理效能区别开。", "本轮β仍固定1；τ自由度不等于直接控制血流或已有生理参数资格。")
        text = text.replace("样例按固定τ/λ100", "样例按同ROI固定τ2基线")
        text = text.replace("右侧展示训练τ/λ0.01状态。", "右侧展示同ROI先验τ臂状态。")
        if parameter_text: text = text.replace(parameter_text, "τ由训练选择冻结，实际选择及失败见后文。\n")
        text += fixed_roi_tau_details(run, out, config, data, states, labels, save)
        if fixed_roi_flow:
            if set(data.kind)=={'synthetic'}:
                text=text.replace('## 完整分母与误差','本报告仅覆盖已终态合成阶段；实测阶段未纳入，未启动实测不记失败或零误差。状态范围及固定样例均来自合成身份，不作实测结论。\n\n## 完整分母与误差')
            text=text.replace('固定AF7Fp1位置：共享τ与轨迹恢复','固定AF7Fp1位置：log-flow软约束对照')
            text=text.replace('三臂预先指定：同ROI固定τ2、18条训练轨迹共享非线性τ、带工程log prior的共享τ；','六臂预先指定：继承父运行固定τ2和训练τ；新增固定τ2/w1及训练τ/w0.1、1、4，无τ先验；')
            text=text.replace('右侧展示同ROI先验τ臂状态。','右侧展示训练τ/w4臂状态，不是按性能选择。')
            text=text.replace('EEG投影保持父规范，Hb目标和训练SD因固定位置而改变；仅本轮同一新目标下重算的τ2作为配对基线。旧parent的τ和误差只可作为背景，未并入本轮paired scores。','prepared从确切固定ROI父运行原样继承，EEG/Hb目标、观测算子与SD均不改变；父两臂的原身份、状态和误差继承后作为配对基线，不是新的重复实验。')
            text += fixed_roi_flow_details(run,out,config,data,states,labels,save,selected)
            if fixed_roi_initial:
                text=text.replace('固定AF7Fp1位置：log-flow软约束对照','固定AF7Fp1位置：free/tied初态约束对照')
                text=text.replace('六臂预先指定：继承父运行固定τ2和训练τ；新增固定τ2/w1及训练τ/w0.1、1、4，无τ先验；','四臂预先指定：固定τ2和训练τ各自比较free/tied初态，所有flow权重为1，无τ先验；')
                text=text.replace('## 完整分母与误差','合成主解释必须按true log(p0/v0)=−0.05/0/+0.05分层；本节合并合成数仅作整体完整性摘要，详细匹配/错设偏差见后文。仅实测free两臂继承父结果，合成全部为新真值家族；继承不构成独立复制。\n\n## 完整分母与误差')
                text=text.replace('六臂及权重预声明，不按验证成绩选择w。','四臂预声明，flow权重全为1；不按验证成绩选择free或tied。')
                text=text.replace('不同w的总目标不可横向排名；多起点仅在同组同臂内选择。继承父运行的两臂及原prepared观测不是新独立复制。','不同初态模型的总目标不作主评分胜负；多起点仅在同组同臂内选择。仅实测free结果和prepared继承，合成全新生成/拟合。')
                text=text.replace('继承父运行的两臂及原prepared观测不是新独立复制。','仅实测free两臂继承；合成四臂为新真值家族。')
                text=text.replace('右侧展示训练τ/w4臂状态，不是按性能选择。','右侧展示共享τ/tied臂状态，不是按性能选择。')
                text=text.replace('prepared从确切固定ROI父运行原样继承，EEG/Hb目标、观测算子与SD均不改变；父两臂的原身份、状态和误差继承后作为配对基线，不是新的重复实验。','实测prepared与free两臂结果从确切flow父运行继承；合成采用本轮新初态真值家族，四臂共享同一目标。继承不是独立复制。')
                text += fixed_roi_initial_details(run,out,config,data,states,labels,save)

            if 'measured' in data.kind.unique():
                text += flow_hbt_decomposition(run,out,config,data,labels,save,selected)
                text += flow_full_curve_atlas(run,out,config,data,labels,[methods[1],methods[3]] if fixed_roi_initial else None)
    for name,caption in figures: text += f"![{name}](figures/{name}.png)\n\n{caption}\n\n"
    text += "## 交付验证\n\n图形以240 dpi PNG嵌入，PDF正文表格可搜索。程序核验图像对象与文本提取；WPS滚动缩放未检查。\n"
    export_pdf(text, out, figures, extra_css="table{font-size:8pt;width:100%}td,th{padding:4pt}h3{font-size:10pt}")


def render_fixed_roi_optical(run,out):
    """Completed retained evidence only; distinct optical targets never pooled."""
    import hashlib
    import yaml
    if out.exists() and any(out.iterdir()):raise ValueError('Choose a new empty versioned export directory')
    manifest=json.loads((run/'manifest.json').read_text())
    cfg=yaml.safe_load((run/'resolved_config.yaml').read_text())
    methods=['fixed_tau_logflow1','trained_tau_logflow1'];pipelines=['no_motion','mne_tddr'];modes=['full','center_EEG','center_fNIRS']
    if (manifest.get('execution')!='completed' or manifest.get('pilot') is not False
            or manifest.get('experiment_id')!='SSM-SHARED-DRIVER-FIXED-ROI-OPTICAL-v1'
            or manifest.get('synthetic_terminal') is not True or manifest.get('measured_terminal') is not True
            or cfg.get('schema')!='shared_driver_fixed_roi_optical_v1' or cfg.get('methods')!=methods
            or cfg.get('optical_pipelines')!=pipelines or cfg.get('modes')!=modes
            or cfg.get('subjects')!=['subject_01','subject_09','subject_18'] or cfg.get('outer_folds')!=[0,1,2,3]
            or cfg.get('parent_run')!='experiments/runs/physiology_semantic_tokenizer/shared_driver_reconstruction/20260926_fixed_roi_flow_v1'
            or cfg.get('optical_audit_run')!='experiments/runs/physiology_semantic_tokenizer/shared_driver_reconstruction/20260926_optical_motion_audit_v1'):
        raise ValueError('Optical report requires exact completed nonpilot optical contract')
    snapshot=Path(manifest.get('source_root',''))
    if snapshot.resolve()!=(run/'source_snapshot').resolve() or not snapshot.is_dir():raise ValueError('Missing owning frozen source snapshot identity')
    root=Path(manifest['project_root']);parent=root/cfg['parent_run'];audit=root/cfg['optical_audit_run']
    pm=json.loads((parent/'manifest.json').read_text());am=json.loads((audit/'manifest.json').read_text())
    control=json.loads((audit/'synthetic_summary.json').read_text());reproduction=json.loads((audit/'measured_summary.json').read_text())
    if (pm.get('execution')!='completed' or pm.get('experiment_id')!='SSM-SHARED-DRIVER-FIXED-ROI-FLOW-v1'
            or am.get('experiment_id')!='SSM-SHARED-DRIVER-OPTICAL-MOTION-AUDIT-v1'
            or am.get('execution')!='completed' or am.get('result_status')!='completed'
            or reproduction.get('current_pipeline_reproduced') is not True or reproduction.get('windows')!=72
            or len(reproduction.get('reproduction_checks',[]))!=216
            or any(c.get('matches') is not True for c in reproduction['reproduction_checks']) or control.get('status')!='completed'):
        raise ValueError('Completed parent and optical reproduction evidence required')
    checks=reproduction['reproduction_checks'];sample_ids={c.get('sample_id') for c in checks}
    if len(sample_ids)!=72 or None in sample_ids or {(c.get('sample_id'),c.get('stage')) for c in checks}!={(sample,stage) for sample in sample_ids for stage in ['od','hb_native','hb_4hz']}:
        raise ValueError('Optical source reproduction identities incomplete')
    data=pd.read_csv(run/'metrics.csv');summary=json.loads((run/'summary.json').read_text())
    keys=['kind','group','method','mode','trial'];expected=set();preps={}
    for kind in ['synthetic','measured']:
        groups=([f'synthetic_tau{tau:g}_{spectrum}_r{rep}' for tau in [1.,2.,4.] for spectrum in ['slow','mixed'] for rep in range(2)]
            if kind=='synthetic' else [f'{pipeline}__{subject}_o{fold}' for pipeline in pipelines for subject in cfg['subjects'] for fold in range(4)])
        for group in groups:
            prep=json.loads((run/'prepared'/f'{group}.json').read_text());preps[group]=prep;spec=prep['spec']
            if spec['group']!=group or spec['kind']!=kind:raise ValueError('Prepared identity mismatch')
            for method in methods:
                for mode in modes:
                    path=run/'cells'/f'{group}__{method}__{mode}'/'result.json';cell=json.loads(path.read_text())
                    if cell['cell']!=path.parent.name or len(cell['rows'])!=6:raise ValueError('Terminal cell identity/count mismatch')
                    if kind=='synthetic':
                        original=json.loads((parent/'cells'/path.parent.name/'result.json').read_text())
                        if cell['rows']!=original['rows'] or cell.get('inheritance',{}).get('status')!='inherited_nonindependent':
                            raise ValueError('Synthetic rows changed rather than inherited')
                    for row in cell['rows']:
                        trial=row['trial'];identity=(kind,group,method,mode,trial)
                        expected_row=dict(spec,method=method,mode=mode,session=cfg['sessions'][trial//8])
                        if any(row.get(k)!=v for k,v in expected_row.items()):raise ValueError('Cell row differs from prepared identity')
                        if trial not in [i for i in range(24) if i%4==spec['outer']] or identity in expected:
                            raise ValueError('Validation split/duplicate mismatch')
                        expected.add(identity)
                        match=data[(data.kind==kind)&(data.group==group)&(data.method==method)&(data['mode']==mode)&(data.trial==trial)]
                        if len(match)!=1 or match.iloc[0].status!=row['status'] or bool(truth(match.converged).iloc[0])!=bool(row.get('converged',False)):
                            raise ValueError('Metrics differ from retained cell status')
                        for component in COMPONENTS:
                            k='nmse_'+component
                            if row.get(k) is not None and not np.isclose(float(match.iloc[0][k]),row[k],rtol=1e-12,atol=1e-14):
                                raise ValueError('Metrics differ from retained cell errors')
    if len(data)!=1296 or data.duplicated(keys).any() or set(data[keys].itertuples(index=False,name=None))!=expected or summary.get('observed_rows')!=1296 or summary.get('terminal_cells')!=216:
        raise ValueError('Optical report requires complete 1296 rows and 216 cells')
    out.mkdir(parents=True,exist_ok=True);(out/'figures').mkdir()
    font=Path('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc')
    if font.exists():font_manager.fontManager.addfont(str(font));plt.rcParams['font.family']=font_manager.FontProperties(fname=str(font)).get_name()
    plt.rcParams.update({'font.size':9,'axes.unicode_minus':False})
    labels=dict(zip(methods,['固定τ=2','18训练trial共享τ']));figures=[];pages=[]
    def save(fig,name):
        fig.tight_layout();fig.savefig(out/'figures'/f'{name}.png',dpi=240);plt.close(fig);figures.append(name)
        return f'![{name}](figures/{name}.png)'
    def good(part):return part[part.status.eq('completed')&truth(part.converged)]
    def errors(part):
        cols=['nmse_'+c for c in COMPONENTS]
        if not len(part):return [np.nan]*3
        return np.sqrt(part.groupby(['subject','session'])[cols].mean().groupby('subject').mean().mean()).tolist()
    measured=data[data.kind.eq('measured')].copy();aggregates=[]
    for (pipeline,method,mode),part in measured.groupby(['optical_pipeline','method','mode'],sort=False):
        if len(part)!=72:raise ValueError('Each pipeline/method/mode needs all 72 rows')
        ok=good(part);err=errors(ok);aggregates.append(dict(pipeline=pipeline,method=method,mode=mode,expected=72,completed=len(ok),
            all_three_below05=int(ok[['nrmse_'+c for c in COMPONENTS]].lt(.5).all(axis=1).sum()) if mode=='full' else None,
            **dict(zip(COMPONENTS,err))))
    ag=pd.DataFrame(aggregates);ag.to_csv(out/'pipeline_metrics.csv',index=False)
    failure=measured.groupby(['optical_pipeline','method','mode','status'],dropna=False).size().reset_index(name='rows')
    failure.to_csv(out/'all_terminal_status_counts.csv',index=False)
    data.to_csv(out/'all_validation_rows.csv',index=False)
    pages.append('# 光学处理敏感性与共享生理轨迹\n\n同一3被试、72窗口、AF7Fp1位置与EEG，比较 no_motion / MNE TDDR 两种光学目标。每条管道分别报告固定τ与训练τ；初态五个自由坐标，flow工程软锚权重1。\n\n'
        '1296/1296验证行、216/216 cells保留；其中432行合成模型证据继承自父flow运行，不是独立重复。864行属于两种新实测目标。无新增拟合或评分。\n\n'
        +table(ag[ag['mode']=='full'].drop(columns='mode').replace({'method':labels}).rename(columns={'pipeline':'管道','method':'方法','expected':'计划','completed':'成功','all_three_below05':'三项<0.5'}))
        +'\nNRMSE按各管道18个训练窗口的固定SD归一化。不同管道改变了目标和尺度，不能视作同任务排名；成功条件误差不能抵消失败分母。联合重建是主终点，遮挡是独立检验。')
    for pipeline in pipelines:
        fig,axes=plt.subplots(1,3,figsize=(7.1,3.4))
        for ax,mode in zip(axes,modes):
            for j,m in enumerate(methods):
                row=ag[(ag.pipeline==pipeline)&(ag.method==m)&(ag['mode']==mode)].iloc[0]
                ax.bar(np.arange(3)+j*.32,[row[c] for c in COMPONENTS],width=.3,label=f'{labels[m]} {row.completed}/72')
            ax.set_xticks(np.arange(3)+.16,COMPONENTS);ax.axhline(.5,color='r',ls=':');ax.set_title(mode);ax.legend(fontsize=6)
        pages.append(f'## {pipeline}：联合重建与遮挡\n\n'+save(fig,'errors_'+pipeline)+'\n仅成功子集聚合，subject/session等权。遮挡列按原保留评分定义；完整终态和逐条误差见CSV。')
    paired=[]
    for (pipeline,mode),part in measured.groupby(['optical_pipeline','mode']):
        pair=part.pivot(index=['group','subject','session','trial'],columns='method',values=['status','converged']+['nmse_'+c for c in COMPONENTS])
        valid=np.ones(len(pair),bool)
        for m in methods:valid&=pair['status'][m].eq('completed')&truth(pair['converged'][m])
        p=pair[valid]
        for c in COMPONENTS:
            delta=p['nmse_'+c][methods[1]].astype(float)-p['nmse_'+c][methods[0]].astype(float)
            paired.append(dict(pipeline=pipeline,mode=mode,component=c,paired_success=len(p),planned=72,
                mean_delta_nmse=float(delta.mean()),median_delta_nmse=float(delta.median()),improved=int((delta<0).sum())))
    paired=pd.DataFrame(paired);paired.to_csv(out/'within_pipeline_paired.csv',index=False)
    pages.append('## 同管道固定τ→训练τ配对\n\n'+table(paired[paired['mode']=='full'].drop(columns='mode'))+'\n差值为训练−固定，只在双方都成功的同一身份上计算；负值表示误差降低。各管道单独比较，配对成功数不是完整72分母的替代。遮挡配对完整保留CSV。')
    starts=[];selections=[];scales=[];states=[]
    for group,prep in preps.items():
        spec=prep['spec'];folder=run/'training'/f'{group}__{methods[1]}'
        choice=json.loads((folder/'selection.json').read_text());local=[]
        for i,start in enumerate([1.,2.,4.]):
            r=json.loads((folder/f'start_{i}'/'result.json').read_text())
            if r.get('spec')!=spec or r.get('start_index')!=i or r.get('start_tau')!=start:raise ValueError('Training start identity mismatch')
            if spec['kind']=='synthetic':
                original=json.loads((parent/'training'/folder.name/f'start_{i}'/'result.json').read_text())
                if {k:v for k,v in r.items() if k!='inheritance'}!={k:v for k,v in original.items() if k!='inheritance'}:raise ValueError('Inherited training changed')
            row=dict(kind=spec['kind'],pipeline=spec.get('optical_pipeline','inherited'),group=group,subject=spec['subject'],fold=spec['outer'],start=i,
                start_tau=start,status=r['status'],converged=bool(r.get('converged',False)),tau=r.get('parameter_value'),objective=r.get('objective'),
                boundary=r.get('boundary_status'),gradient=r.get('projected_scaled_gradient_inf_norm'))
            starts.append(row);local.append(row)
        successful=[r['tau'] for r in local if r['status']=='completed' and r['converged']]
        selections.append(dict(kind=spec['kind'],pipeline=spec.get('optical_pipeline','inherited'),group=group,subject=spec['subject'],fold=spec['outer'],
            status=choice['status'],tau=choice.get('parameter_value'),selected_start=choice.get('selected_training_start'),successful_starts=len(successful),
            min_tau=min(successful,default=np.nan),max_tau=max(successful,default=np.nan)))
        if spec['kind']!='measured':continue
        if prep['status']=='completed':
            meta=prep['metadata'];opt=meta['optical_processing'];base=json.loads((parent/'prepared'/f"{spec['subject']}_o{spec['outer']}.json").read_text())['metadata']
            scales.append(dict(pipeline=spec['optical_pipeline'],subject=spec['subject'],fold=spec['outer'],MAD=opt['common_training_MAD'],factor=opt['fnirs_factor'],
                parent_MAD=base['fixed_roi']['common_training_MAD'],parent_factor=base['fixed_roi']['fnirs_factor'],gauge=opt['frozen_observation_loading'],
                **{c+'_SD':meta['normalization_sd'][i] for i,c in enumerate(COMPONENTS)},**{c+'_parent_SD':base['normalization_sd'][i] for i,c in enumerate(COMPONENTS)}))
        for m in methods:
            for mode in modes:
                path=run/'cells'/f'{group}__{m}__{mode}'/'trajectories.npz'
                rows=measured[(measured.group==group)&(measured.method==m)&(measured['mode']==mode)]
                arrays=None
                if path.exists():
                    with np.load(path,allow_pickle=False) as a:arrays={k:a[k] for k in ['states','trial_indices']}
                for r in rows.itertuples():
                    row=dict(pipeline=spec['optical_pipeline'],group=group,method=m,mode=mode,trial=r.trial,status=r.status,converged=bool(truth(pd.Series([r.converged])).iloc[0]),finite=False)
                    if arrays is not None:
                        idx=np.flatnonzero(arrays['trial_indices']==r.trial)
                        if len(idx)!=1:raise ValueError('State trial index mismatch')
                        state=arrays['states'][idx[0]];row['finite']=bool(state.shape==(120,6) and np.isfinite(state).all())
                        if row['finite']:
                            for j,name in enumerate(['r','s','f','v','p','q']):row[name+'_min']=float(state[:,j].min());row[name+'_max']=float(state[:,j].max())
                    states.append(row)
    starts=pd.DataFrame(starts);selections=pd.DataFrame(selections);scales=pd.DataFrame(scales);states=pd.DataFrame(states)
    if scales.empty:scales=pd.DataFrame(columns=['pipeline','factor','parent_factor','HbO_SD','HbO_parent_SD','HbR_SD','HbR_parent_SD'])
    for name,frame in [('training_starts',starts),('training_selection',selections),('optical_scale_changes',scales),('all_state_ranges',states)]:frame.to_csv(out/(name+'.csv'),index=False)
    selected=selections[selections.kind=='measured'];candidates=starts[starts.kind=='measured']
    fig,axes=plt.subplots(2,1,figsize=(7.1,5.2))
    for ax,pipeline in zip(axes,pipelines):
        part=candidates[candidates.pipeline==pipeline]
        for i in range(3):
            q=part[part.start==i];x=[cfg['subjects'].index(v)*4+f for v,f in zip(q.subject,q.fold)]
            ax.scatter(np.array(x)+(i-1)*.1,q.tau,label=f'起点{[1,2,4][i]}',s=15,marker=['o','x','+'][i])
        ax.axhline(.5,color='r',ls=':');ax.axhline(8.,color='r',ls=':');ax.set_xticks(range(12),[f'{s[-2:]}/{f}' for s in cfg['subjects'] for f in range(4)],fontsize=7)
        ax.set_ylabel('τ / 秒');ax.set_title(pipeline);ax.legend(fontsize=7)
    pages.append('## 训练τ多起点与被试/折\n\n'+save(fig,'tau_starts')+'\n点含失败末有效终点，不能当估计；状态详见training_starts.csv。虚线0.5–8秒是数值搜索边界，不是生理正常范围。没有τ先验，多起点差异不是置信区间。')
    pages.append('## 全部24折训练终态\n\n'+table(selected.drop(columns=['kind','group']).rename(columns={'pipeline':'管道','subject':'被试','fold':'折','status':'终态','successful_starts':'成功起点','selected_start':'选中起点'}))+'\n只有收敛成功且积分检查通过的起点参与选取。失败折仍产生完整验证失败行。')
    fig,axes=plt.subplots(2,4,figsize=(7.1,5.1))
    for row,pipeline in enumerate(pipelines):
        for col,name in enumerate(['f','v','p','q']):
            ax=axes[row,col]
            for j,m in enumerate(methods):
                part=states[(states.pipeline==pipeline)&(states.method==m)&(states['mode']=='full')&states.finite&states.converged&states.status.eq('completed')]
                if len(part):ax.scatter(np.full(len(part),j)-.08,part[name+'_min'],s=7);ax.scatter(np.full(len(part),j)+.08,part[name+'_max'],s=7)
            ax.set_xticks([0,1],['固定','训练']);ax.set_title(pipeline+' '+name,fontsize=8);ax.axhline(1,color='.7',ls=':')
    pages.append('## 状态范围与解释边界\n\n'+save(fig,'state_ranges')+'\n每点为成功full轨迹的最小或最大值；完整CSV还保留失败末有效状态与缺失标记。f/v/p/q为归一化模型状态，参照1不是正常范围。flow软锚提供未观测的绝对水平约束，不证明参数已识别；参数分离也可能来自尺度和优化补偿。')
    fig,axes=plt.subplots(1,3,figsize=(7.1,3.8))
    for pipeline in pipelines:
        part=scales[scales.pipeline==pipeline]
        for ax,key,parentkey in zip(axes,['factor','HbO_SD','HbR_SD'],['parent_factor','HbO_parent_SD','HbR_parent_SD']):
            ax.plot(range(len(part)),part[key]/part[parentkey],'o-',label=pipeline,ms=3);ax.axhline(1,color='.6',ls=':');ax.set_title(key+' / 旧目标');ax.set_xlabel('被试/折固定顺序')
    axes[0].legend(fontsize=7)
    pages.append('## 目标与归一化尺度变化\n\n'+save(fig,'scale_changes')+'\nMAD、factor、EEG/Hb分量SD及父值均保留于optical_scale_changes.csv。Hb共同MAD和分量SD只用18训练窗口；EEG与旧fold保持原样，loading gauge固定。上述变化是观测坐标变化，不能解释为生理浓度变化或跨管道胜负。')
    failure_audit_path=run/'training_failure_audit'/'no_motion_preliminary.json'
    failure_audit=json.loads(failure_audit_path.read_text()) if failure_audit_path.exists() else None
    failure_note=''
    if failure_audit is not None:
        failed=[r for r in failure_audit['rows'] if r.get('status')!='completed' or not r.get('converged',False)]
        margin=[r['minimum_absolute_HbO_model_margin'] for r in failed if r.get('minimum_absolute_HbO_model_margin') is not None]
        failure_note=('\n\n保留的[no_motion训练失败诊断](../training_failure_audit/no_motion_preliminary.json)只检查S09四折与S01折0参考，'
            f'其中{len(failed)}个失败起点的末接受状态最小HbO模型余额为{min(margin):.3g}至{max(margin):.3g}（若存在）。'
            '该量为p−0.35q，不是标定浓度；多条记录的最小余额位于全局trial21初始时刻，且梯度未达阈值、域拒绝较多。'
            '日志只保留末接受轨迹的时刻，未记录被拒提议的精确失败时刻，因此这是边界相关的数值诊断，不是已证明的完整因果链，也不推广到其他管道。')
    pages.append('## 所有实测失败与终态分母\n\n'+table(failure.rename(columns={'optical_pipeline':'管道','method':'方法','mode':'模式','status':'终态','rows':'行数'}).replace({'方法':labels}))+'\n所有状态计数纳入，不把未收敛、训练失败、准备失败或积分失败当作零误差。每管道×方法×模式计划均为72。'+failure_note)
    syn=data[data.kind=='synthetic'];synrows=[]
    for (m,mode),part in syn.groupby(['method','mode']):synrows.append(dict(method=labels[m],mode=mode,planned=len(part),successful=len(good(part)),**dict(zip(COMPONENTS,errors(good(part))))))
    pages.append('## 合成证据继承与上游控制\n\n'+table(pd.DataFrame(synrows))+'\n12组合成模型、432验证行及36训练起点逐项继承父flow证据，失败一并保留；这不是经过TDDR后的新SSM恢复实验。光学校正的零/常数/正负非对称周期/真实线性趋势/噪声控制来自独立audit。控制记录总数按原summary统计，不能将MNE或不校正当真值。上游当前处理重现216/216通过，因而可归因固定输入下的处理差异，不能推断真实人体趋势。')
    new_control_path=run/'optical_controls'/'results.json'
    new_control=json.loads(new_control_path.read_text());control_rows=new_control['metrics']
    if len(control_rows)!=168 or new_control['config'].get('methods')!=['no_motion','current','mne_tddr']:
        raise ValueError('Expected separate 168-record optical sensitivity controls')
    control_ids={(r['context_s'],r['case'],r['seed'],r['method'],r['region']) for r in control_rows}
    if len(control_ids)!=168 or {r['method'] for r in control_rows}!={'no_motion','current','mne_tddr'} or {r['region'] for r in control_rows}!={'full','common_roi'}:
        raise ValueError('Optical context control identities incomplete')
    flat=[]
    for r in control_rows:
        for wavelength in range(2):
            flat.append({k:(v[wavelength] if isinstance(v,list) else v) for k,v in dict(r,wavelength_index=wavelength).items()})
    controls=pd.DataFrame(flat);controls.to_csv(out/'optical_context_controls.csv',index=False)
    fig,axes=plt.subplots(1,2,figsize=(7.1,3.7))
    cases=['slow_clean','slow_noise','slow_noise_spike','slow_noise_step']
    for ax,context in zip(axes,[30,120]):
        part=controls[(controls.context_s==context)&(controls.region=='common_roi')&(controls.wavelength_index==0)]
        for j,m in enumerate(['no_motion','current','mne_tddr']):
            values=[part[(part['case']==case)&(part.method==m)].filtered_response_gain.mean() for case in cases]
            ax.plot(range(4),values,'o-',label=m,ms=3)
        ax.set_xticks(range(4),['clean','noise','spike','step'],rotation=25);ax.axhline(1,color='.6',ls=':');ax.set_title(f'{context}s context / common ROI');ax.set_ylabel('filtered response gain');ax.legend(fontsize=6)
    pages.append('## 新合成上下文压力检查：168条独立记录\n\n'+save(fig,'optical_context_controls')+'\n28个构造输入×3种处理×full/common ROI两区间=168记录，每记录含两波长；CSV展开为336行。本图是760nm、共同30秒ROI、噪声seed平均的已知慢响应gain，全部误差与运动残余指标保持原值。残余减小不等于真实响应保留。\n\n'
        '该压力输入为双gamma慢响应加白噪声，缺少生理低频背景；强衰减不能外推为所有实测数据都会丢失HRF，也不是复现原TDDR论文的10分钟重复任务评估。上下文长度同时影响运动估计和有限窗口滤波，不能独立归因其中一项。此168记录与旧audit六类输入、432继承SSM行是三个不同证据家族。')

    extension_path=run/'optical_controls'/'background_extension_v1'/'results.json'
    quiet_path=extension_path.parent/'quiet_paired_reference.json'
    extension=json.loads(extension_path.read_text());quiet=json.loads(quiet_path.read_text())
    if extension.get('status')!='completed' or extension.get('condition_count')!=216 or extension.get('metric_rows')!=432 or len(extension['rows'])!=432:
        raise ValueError('Expected completed 216-condition/432-row background extension')
    background=[]
    for context in [30,120]:
        values=[r['filtered_gain'] for r in quiet['rows'] if r['context_s']==context and r['method']=='mne_tddr']
        if len(values)!=3:raise ValueError('Quiet paired comparator needs exactly three seeds')
        median=np.median(values,axis=0)
        background.append(dict(context_s=context,background='white only',gain_760=median[0],gain_850=median[1],seeds=3))
        for family,level in [('heart_white','low'),('heart_white','high'),('heart_resp_mayer_white','low'),('heart_resp_mayer_white','high')]:
            values=[r['filtered_response']['gain'] for r in extension['rows'] if r['context_s']==context and r['method']=='mne_tddr'
                    and r['motion']=='none' and r['region']=='common_roi' and r['family']==family and r['level']==level]
            if len(values)!=3:raise ValueError('Background comparator needs exactly three seeds')
            median=np.median(values,axis=0)
            background.append(dict(context_s=context,background=family+'/'+level,gain_760=median[0],gain_850=median[1],seeds=3))
    background=pd.DataFrame(background);background.to_csv(out/'background_paired_HRF_gain.csv',index=False)
    pages.append('## 频率背景扩展：配对HRF保留率\n\n'+table(background)+'\n216个条件、432条full/common ROI记录为另一套纯合成扩展。此表是无运动条件、共同ROI、MNE两波长gain的3 seed中位数。所有行统一采用 F(HRF+背景)−F(背景)，与同样线性滤波的真实HRF比较；white only来自quiet_paired_reference.json，不能拿上一页168记录的raw-truth gain直接替代。\n\n'
        '加入心动及心动/呼吸/Mayer频率背景后，配对HRF增益明显提高，说明安静白噪声压力案例不能概括所有真实记录；保留仍不完美。low/high是预设工程幅度与HRF峰值的比值，不是人体标定或正常范围。该扩展不是原论文AR10/重复任务/GLM AUC复现，亦不构成实测MNE校准或方法排名。')

    # Independent optical controls are retained OD, not SSM truth/predictions.
    with np.load(audit/'synthetic_traces.npz',allow_pickle=False) as archive:
        control_arrays={k:archive[k] for k in archive.files}
    fig,axes=plt.subplots(1,2,figsize=(7.1,3.5))
    control_keys=list(control_arrays)
    for ax,case in zip(axes,['asymmetric_positive','true_linear_trend']):
        for method in ['no_motion','current','mne_tddr']:
            key=case+'__'+method+'__od'
            if key not in control_arrays:raise ValueError('Missing retained optical control OD')
            y=control_arrays[key][:,0];ax.plot(np.arange(len(y))/10,y-y[0],label=method,lw=1)
        ax.set_title(case);ax.set_xlabel('秒');ax.set_ylabel('OD−初值');ax.legend(fontsize=7)
    pages.append('## 光学负对照：漂移与真实趋势损失\n\n'+save(fig,'optical_controls')+'\n输入无运动伪迹时，当前方法仍可能把非对称周期积分成漂移；MNE能够消除真实线性趋势。两种反例必须一起看，均不能作为人体真值。完整控制记录保持在source optical audit，未重算。')
    atlas=fitz.open()
    for pipeline in pipelines:
        folder=out/pipeline;folder.mkdir()
        flow_full_curve_atlas(run,folder,cfg,measured[measured.optical_pipeline==pipeline],labels,methods,both_states=True,context_label=pipeline)
        with fitz.open(folder/'FULL_CURVE_ATLAS.pdf') as part:atlas.insert_pdf(part)
    if len(atlas)!=48 or sum(len(page.get_images()) for page in atlas)!=48:
        raise ValueError('Combined optical atlas must have 48 bitmap figure pages')
    atlas.save(out/'FULL_CURVE_ATLAS.pdf',garbage=4,deflate=True);atlas.close()
    pages.append('## 全部曲线与来源\n\n[FULL_CURVE_ATLAS.pdf](FULL_CURVE_ATLAS.pdf)共48页，每管道24页、每页3条：全部72窗口的EEG/HbO/HbR与两臂f/v/p/q。固定顺序不按效果选样，失败与末有效状态均标明。每管道另有atlas_index.csv。\n\n'
        f'运行：{run.name}。配置、manifest、source snapshot来源原样保留。父模型：{parent.name}；光学audit：{audit.name}。旧flow目标包含已定位的处理漂移，其旧NRMSE不可直接与新目标做效果排名。\n\n'
        '所有图为240dpi PNG，PDF正文/表格可搜索。WPS未检查。报告不改变缓存、评分、模型或运行状态。')
    text='\n\n---\n\n'.join(pages);(out/'REPORT.md').write_text(text)
    pdf=fitz.open();css='body{font-family:sans-serif;font-size:9pt;line-height:1.35}h1{font-size:19pt}h2{font-size:14pt}table{font-size:7pt;border-collapse:collapse;width:100%}td,th{border:.4pt solid #bbb;padding:3pt}img{width:510pt}'
    def rectfn(n,filled):
        page=fitz.paper_rect('a4');return page,page+(36,32,-36,-32),None
    for section in pages:
        html=markdown.markdown(section,extensions=['tables']).replace('<img ','<img width="510" ')
        part=fitz.Story(html,user_css=css,archive=fitz.Archive(str(out))).write_with_links(rectfn);pdf.insert_pdf(part);part.close()
    pdf.save(out/'REPORT.pdf',garbage=4,deflate=True)
    validation=dict(pages=len(pdf),figure_count=len(figures),image_objects=sum(len(p.get_images()) for p in pdf),searchable_text_chars=sum(len(p.get_text()) for p in pdf),
        bitmap_dpi=240,wps_checked=False,atlas_pages=48,source_report=str(Path(__file__).resolve()),source_report_sha256=hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
        training_failure_audit_source=str(failure_audit_path) if failure_audit is not None else None,
        training_failure_audit_sha256=hashlib.sha256(failure_audit_path.read_bytes()).hexdigest() if failure_audit is not None else None,
        background_extension_source=str(extension_path),background_extension_sha256=hashlib.sha256(extension_path.read_bytes()).hexdigest(),
        quiet_paired_source=str(quiet_path),quiet_paired_sha256=hashlib.sha256(quiet_path.read_bytes()).hexdigest(),
        optical_controls_source=str(new_control_path),optical_controls_sha256=hashlib.sha256(new_control_path.read_bytes()).hexdigest(),
        source_manifest=str(run/'manifest.json'),source_manifest_sha256=hashlib.sha256((run/'manifest.json').read_bytes()).hexdigest())
    if not 10<=len(pdf)<=16 or validation['image_objects']!=len(figures) or validation['searchable_text_chars']<1000:raise ValueError('Optical report pagination/bitmap validation failed')
    for i,page in enumerate(pdf):page.get_pixmap(matrix=fitz.Matrix(1.3,1.3)).save(out/f'preview_page_{i+1:02d}.png')
    pdf.close();(out/'export_validation.json').write_text(json.dumps(validation,indent=2));print(json.dumps(validation))


def render_conditional_optical_gain(run, out, *, pilot_truth_audit=None, terminal_input_snapshot=None):
    """Audit and render terminal retained arrays; never load native data or fit."""
    import hashlib
    import yaml
    if out.exists() and any(out.iterdir()):
        raise ValueError('Choose a new empty versioned export directory')
    input_root=Path(terminal_input_snapshot).resolve() if terminal_input_snapshot is not None else run
    manifest=json.loads((input_root/'manifest.json').read_text())
    cfg=yaml.safe_load((input_root/'resolved_config.yaml').read_text())
    methods=['fixed_tau_logflow1','conditional_fixed_gain','conditional_trained_gain']
    labels=dict(zip(methods,['A identity β1','B 条件映射 βref','C 条件映射 学β']))
    modes=['full','center_EEG','center_fNIRS'];pipelines=['no_motion','mne_tddr']
    full=manifest.get('execution')=='completed' and manifest.get('measured_terminal') is True
    if (manifest.get('pilot') is not False or manifest.get('synthetic_terminal') is not True
            or manifest.get('execution') not in ['synthetic_terminal','completed']
            or manifest.get('experiment_id')!='SSM-SHARED-DRIVER-CONDITIONAL-OPTICAL-GAIN-v1'
            or cfg.get('schema')!='shared_driver_conditional_optical_gain_v1'
            or cfg.get('methods')!=methods or cfg.get('modes')!=modes
            or cfg.get('optical_pipelines')!=pipelines
            or (manifest.get('execution')=='completed' and not full)):
        raise ValueError('Requires exact nonpilot terminal conditional gain evidence')
    if Path(manifest['source_root']).resolve()!=(run/'source_snapshot').resolve():
        raise ValueError('Missing owning source snapshot identity')
    root=Path(manifest['project_root']);parent=root/cfg['parent_run']
    groups=[f'synthetic_g{g:g}_{s}_r{r}' for g in [.5,1.,2.] for s in ['slow','mixed'] for r in range(2)]
    if full:groups += [f'{p}__{s}_o{f}' for p in pipelines for s in cfg['subjects'] for f in range(4)]
    data=pd.read_csv(input_root/'metrics.csv');summary=json.loads((input_root/'summary.json').read_text())
    keys=['kind','group','method','mode','trial'];expected=set();preps={};starts=[];choices=[];states=[];array_checks=[];costs=[]
    for group in groups:
        prep=json.loads((run/'prepared'/f'{group}.json').read_text());preps[group]=prep;spec=prep['spec']
        if spec['group']!=group or prep['status']!='completed':raise ValueError('Prepared identity/status mismatch')
        with np.load(run/'prepared'/f'{group}.npz',allow_pickle=False) as a:
            target=a['target'].copy();sd=a['normalizer'].copy()
        if 'normalization_sd' in prep['metadata'] and not np.array_equal(sd,np.asarray(prep['metadata']['normalization_sd'])):raise ValueError('Prepared SD differs from metadata')
        if spec['kind']=='measured':
            with np.load(parent/'prepared'/f'{group}.npz',allow_pickle=False) as a:
                if not np.array_equal(target,a['target']) or not np.array_equal(sd,a['normalizer']):raise ValueError('Parent target/SD changed')
        ref=prep['conditional_mapping']['beta_reference']
        for method in methods:
            for mode in modes:
                folder=run/'cells'/f'{group}__{method}__{mode}'
                cell=json.loads((folder/'result.json').read_text())
                if cell['cell']!=folder.name or len(cell['rows'])!=6:raise ValueError('Terminal cell identity/count mismatch')
                if spec['kind']=='measured' and method==methods[0]:
                    old=json.loads((parent/'cells'/folder.name/'result.json').read_text())
                    if cell['rows']!=old['rows']:raise ValueError('Inherited A rows changed')
                if mode=='full':
                    for result in cell.get('trial_results',[]):
                        if 'weighted_data_sse' not in result:continue
                        initial_cost=float(result['initial_penalty']*np.sum((np.asarray(result['initial_state'])-np.array([0.,1.,1.,1.,1.]))**2))
                        cost=dict(kind=spec['kind'],pipeline=spec.get('optical_pipeline','synthetic'),group=group,method=method,trial=result['trial'],status=result.get('scoring_status',result['status']),converged=result.get('converged',False),data_sse=result['data_sse'],weighted_data_sse=result['weighted_data_sse'],initial_state_cost=initial_cost,curvature_cost=result['penalty']*result['roughness'],flow_cost=result['flow_prior_cost'],amplitude_cost=result.get('driver_amplitude_cost',0.),objective=result['objective'])
                        total=sum(cost[k] for k in ['weighted_data_sse','initial_state_cost','curvature_cost','flow_cost','amplitude_cost'])
                        if not np.isclose(total,cost['objective'],rtol=1e-9,atol=1e-8):raise ValueError('Retained objective decomposition mismatch')
                        costs.append(cost)
                arrays=None
                if (folder/'trajectories.npz').exists():
                    with np.load(folder/'trajectories.npz',allow_pickle=False) as a:arrays={k:a[k].copy() for k in a.files}
                    idx=arrays['trial_indices']
                    if not np.array_equal(arrays['target'],target[idx]) or not np.array_equal(arrays['normalizer'],sd):raise ValueError('Cell target/SD differ')
                    array_checks.append(dict(group=group,method=method,mode=mode,target_exact=True,SD_exact=True))
                for row in cell['rows']:
                    trial=row['trial'];identity=(spec['kind'],group,method,mode,trial)
                    if any(row.get(k)!=v for k,v in dict(spec,method=method,mode=mode,session=cfg['sessions'][trial//8]).items()):raise ValueError('Cell row differs from prepared identity')
                    if trial not in [i for i in range(24) if i%4==spec['outer']] or identity in expected:raise ValueError('Split/duplicate mismatch')
                    expected.add(identity)
                    selected=data[(data.group==group)&(data.method==method)&(data['mode']==mode)&(data.trial==trial)]
                    if len(selected)!=1 or selected.iloc[0].status!=row['status'] or bool(truth(selected.converged).iloc[0])!=bool(row.get('converged',False)):raise ValueError('Metric status mismatch')
                    for c in COMPONENTS:
                        value=row.get('nmse_'+c);actual=selected.iloc[0].get('nmse_'+c)
                        if (value is None and pd.notna(actual)) or (value is not None and not np.isclose(value,actual,rtol=1e-12,atol=1e-14)):raise ValueError('Metric value mismatch')
                    if mode=='full':
                        record=dict(kind=spec['kind'],pipeline=spec.get('optical_pipeline','synthetic'),group=group,method=method,trial=trial,status=row['status'],converged=row.get('converged',False))
                        if arrays is not None:
                            where=np.flatnonzero(arrays['trial_indices']==trial)
                            if len(where)!=1:raise ValueError('State trial identity mismatch')
                            k=int(where[0]);state=arrays['states'][k]
                            if np.isfinite(state).all():
                                for j,name in enumerate(['r','s','f','v','p','q']):
                                    record[name+'_min']=float(state[:,j].min());record[name+'_max']=float(state[:,j].max())
                                record['initial_max_rest_deviation']=float(np.max(abs(state[0,2:]-1)))
                                if spec['kind']=='synthetic':
                                    for j,name in enumerate(['r','s','f','v','p','q']):record[name+'_truth_rmse']=float(np.sqrt(np.mean((state[:,j]-arrays['truth_states'][k,:,j])**2)))
                        states.append(record)
        folder=run/'training'/f'{group}__{methods[2]}'
        selection=json.loads((folder/'selection.json').read_text());local=[]
        for i,g in enumerate([.5,1.,2.]):
            result=json.loads((folder/f'start_{i}'/'result.json').read_text())
            if result['spec']!=spec or result['start_index']!=i or result['start_relative_gain']!=g:raise ValueError('Training identity mismatch')
            beta=result.get('parameter_value');relative=beta/ref if beta is not None else None
            item=dict(kind=spec['kind'],pipeline=spec.get('optical_pipeline','synthetic'),group=group,subject=spec['subject'],fold=spec['outer'],start=i,start_g=g,status=result['status'],converged=result.get('converged',False),beta=beta,g=relative,beta_reference=ref,boundary=result.get('boundary_status'),objective=result.get('objective'),gradient=result.get('projected_scaled_gradient_inf_norm'))
            starts.append(item);local.append(item)
        valid=[r['g'] for r in local if r['status']=='completed' and r['converged']]
        beta=selection.get('parameter_value');g=beta/ref if beta is not None else None
        choices.append(dict(kind=spec['kind'],pipeline=spec.get('optical_pipeline','synthetic'),group=group,subject=spec['subject'],fold=spec['outer'],status=selection['status'],beta=beta,g=g,beta_reference=ref,true_g=spec.get('true_relative_gain'),successful_starts=len(valid),start_span=max(valid)-min(valid) if valid else None,selected_start=selection.get('selected_training_start')))
    expected_rows=1944 if full else 648
    if len(data)!=expected_rows or data.duplicated(keys).any() or set(data[keys].itertuples(index=False,name=None))!=expected or summary['observed_rows']!=expected_rows or summary['terminal_cells']!=expected_rows//6:raise ValueError('Incomplete terminal denominator')
    out.mkdir(parents=True);(out/'figures').mkdir()
    font=Path('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc')
    if font.exists():font_manager.fontManager.addfont(str(font));plt.rcParams['font.family']=font_manager.FontProperties(fname=str(font)).get_name()
    plt.rcParams.update({'font.size':9,'axes.unicode_minus':False});figures=[];pages=[]
    def save(fig,name):
        fig.tight_layout();fig.savefig(out/'figures'/f'{name}.png',dpi=240);plt.close(fig);figures.append(name)
        return f'![{name}](figures/{name}.png)'
    def good(part):return part[part.status.eq('completed')&truth(part.converged)]
    data['pipeline']=data.optical_pipeline.fillna('synthetic') if 'optical_pipeline' in data else 'synthetic'
    ag=[];paired=[]
    for (kind,pipeline,method,mode),part in data.groupby(['kind','pipeline','method','mode'],sort=False):
        ok=good(part);cols=['nmse_'+c for c in COMPONENTS]
        error=np.sqrt(ok.groupby(['subject','session'])[cols].mean().groupby('subject').mean().mean()) if len(ok) else [np.nan]*3
        ag.append(dict(kind=kind,pipeline=pipeline,method=method,mode=mode,planned=len(part),success=len(ok),all3_below05=int(ok[['nrmse_'+c for c in COMPONENTS]].lt(.5).all(axis=1).sum()) if mode=='full' else None,**dict(zip(COMPONENTS,error))))
    for (kind,pipeline,mode),part in data.groupby(['kind','pipeline','mode']):
        for left,right in [(methods[0],methods[1]),(methods[1],methods[2]),(methods[0],methods[2])]:
            joined=good(part[part.method==left]).merge(good(part[part.method==right]),on=['group','trial','subject','session'],suffixes=('_left','_right'))
            for c in COMPONENTS:
                delta=joined['nmse_'+c+'_right']-joined['nmse_'+c+'_left']
                paired.append(dict(kind=kind,pipeline=pipeline,mode=mode,contrast=left+' → '+right,component=c,planned=len(part)//3,common_success=len(joined),mean_delta_nmse=delta.mean(),median_delta_nmse=delta.median(),improved=int((delta<0).sum())))
    ag=pd.DataFrame(ag);paired=pd.DataFrame(paired);starts=pd.DataFrame(starts);choices=pd.DataFrame(choices);states=pd.DataFrame(states)
    for name in ['r','s','f','v','p','q']:
        for suffix in ['_min','_max','_truth_rmse']:
            if name+suffix not in states:states[name+suffix]=np.nan
    if 'initial_max_rest_deviation' not in states:states['initial_max_rest_deviation']=np.nan
    if 'driver_nrmse' not in data:data['driver_nrmse']=np.nan
    status=data.groupby(['kind','pipeline','method','mode','status'],dropna=False).size().reset_index(name='rows')
    for name,frame in [('all_validation_rows',data),('metrics_summary',ag),('paired_metrics',paired),('training_starts',starts),('training_selections',choices),('full_state_ranges_and_recovery',states),('terminal_status_counts',status),('target_SD_audit',pd.DataFrame(array_checks))]:frame.to_csv(out/(name+'.csv'),index=False)
    pages.append('# 条件性光学观测映射与共享血管增益\n\n'+f'阶段：{"完整合成与实测" if full else "仅合成终态"}；保留 {expected_rows}/{expected_rows} 验证行和全部失败。联合 EEG/HbO/HbR 重建为主终点，遮挡预测为独立检验。EEG 是1–45 Hz log-power/PCA特征，不是原始电压。\n\n'
        +('本报告只覆盖合成终态；计划实测阶段未纳入，完整实验未完成。\n\n' if not full else '')+'A：原identity映射、β=1。B：条件2×2映射B、βref=1/√det(B)。C：相同映射、18训练trial共享估计β，验证冻结。τ全部固定2秒。A→B同时改变映射、参考β和正则项相对作用，不能称纯映射效应；B→C隔离训练增益的贡献。\n\n'
        '71 µM总Hb基线、DPF=6、3 cm路径及组织敏感度1均为文献条件假设，无个体独立标定。β与g=β/βref是当前EEG/光学坐标下的有效耦合，不是已校准的个体神经血管耦合。g的[0.1,10]为数值搜索界，非正常生理区间。[Friston原始模型](https://www.fil.ion.ucl.ac.uk/spm/doc/papers/karl_nonlinear.pdf)的效能同时包含输入诱发神经强度和血管耦合，不能等同独立NVC。')
    for kind,pipeline in ag[['kind','pipeline']].drop_duplicates().itertuples(index=False,name=None):
        part=ag[(ag.kind==kind)&(ag.pipeline==pipeline)]
        fig,axes=plt.subplots(1,3,figsize=(7.1,3.4))
        for ax,mode in zip(axes,modes):
            for j,m in enumerate(methods):
                row=part[(part.method==m)&(part['mode']==mode)].iloc[0]
                ax.bar(np.arange(3)+j*.25,[row[c] for c in COMPONENTS],width=.23,label=f'{"ABC"[j]} {row.success}/{row.planned}')
            ax.set_xticks(np.arange(3)+.25,COMPONENTS);ax.axhline(.5,c='r',ls=':');ax.set_title(mode);ax.legend(fontsize=6)
        pages.append(f'## {kind} / {pipeline} 重建与独立遮挡\n\n'+save(fig,kind+'_'+pipeline+'_errors')+'\n\n'+table(part[part['mode']=='full'][['method','planned','success','all3_below05']].replace({'method':labels}))+'\n误差为成功条件子集subject/session等权NMSE开根；失败不填零。全分母、遮挡及所有身份保留CSV。每管道目标和训练SD跨三臂完全一致；不同管道不作同任务排名。')
    pages.append('## 同身份配对与失败\n\n'+table(paired[(paired['mode']=='full')&(paired.contrast==methods[1]+' → '+methods[2])].drop(columns=['mode','contrast']))+'\n差值=C−B；只在双方成功的同一身份上计算，不能替代完整分母。A→B/A→C和遮挡配对均保留paired_metrics.csv。\n\n'+table(status.groupby(['kind','pipeline','method','status'],dropna=False)['rows'].sum().reset_index())+'\n失败轨迹不计入有效恢复；图册不把末有效迭代伪装成成功。')
    for kind,pipeline in choices[['kind','pipeline']].drop_duplicates().itertuples(index=False,name=None):
        selected=choices[(choices.kind==kind)&(choices.pipeline==pipeline)];names=selected.group.tolist()
        fig,ax=plt.subplots(figsize=(7.1,3.7))
        for start in range(3):
            part=starts[(starts.kind==kind)&(starts.pipeline==pipeline)&(starts.start==start)]
            ok=part.status.eq('completed')&truth(part.converged)
            ax.scatter([names.index(g)+(start-1)*.13 for g in part[ok].group],part[ok].g,s=19,label=f'起点 g={ [.5,1.,2.][start]}')
            bad=part[~ok];ax.scatter([names.index(g)+(start-1)*.13 for g in bad.group],[.12]*len(bad),c='red',marker='x',s=25)
        if kind=='synthetic':ax.scatter(range(len(selected)),selected.true_g,marker='_',c='k',s=100,label='真实g')
        ax.set_yscale('log');ax.set_ylim(.08,12);ax.axhline(.1,c='r',ls=':');ax.axhline(10,c='r',ls=':');ax.set_xticks(range(len(names)),[n.replace('synthetic_','').replace('subject_','').replace(pipeline+'__','') for n in names],rotation=60,fontsize=6);ax.set_ylabel('g = β / βref');ax.legend(fontsize=7)
        pages.append(f'## {kind} / {pipeline} 增益多起点\n\n'+save(fig,kind+'_'+pipeline+'_gain')+'\n红叉只表示失败，纵坐标0.12不是该失败的估计。实际β、βref、g、边界、梯度和起点跨度均保留CSV。跨折训练重叠，折间差异不是独立被试证据。\n\n'+table(selected[['subject','fold','status','beta','g','true_g','successful_starts','start_span']]))
    if costs:
        costs=pd.DataFrame(costs);costs.to_csv(out/'full_objective_components.csv',index=False)
        ok_costs=costs[costs.status.eq('completed')&truth(costs.converged)]
        cost_summary=ok_costs.groupby(['kind','pipeline','method'])[['weighted_data_sse','initial_state_cost','curvature_cost','flow_cost']].mean().reset_index()
        cost_summary.to_csv(out/'objective_component_means.csv',index=False)
        pages.append('## 数据残差与工程惩罚的相对作用\n\n'+table(cost_summary)+'\n表为成功full验证trial的算术均值，完整状态、原始SSE和所有分项保留CSV。初态项按保留物理初态及权重精确复算，并核对分项和等于原objective；未重新拟合。A→B缩小物理状态时初态/flow惩罚相对作用也会减弱，因此重建改善可能包含解除工程惩罚的效果。总objective不用于跨臂判胜；只有同臂多起点的原优化选择可比较。')
    screen=summary.get('synthetic_engineering_screen')
    if screen is not None:
        pages.append('## 合成工程筛查与阶段边界\n\n'+table(pd.DataFrame(screen['groups'])[['group','true_relative_gain','selected_relative_gain','successful_starts','selected_relative_error','successful_span_over_true','successful_starts_interior','passed']].rename(columns={'true_relative_gain':'true_g','selected_relative_gain':'selected_g','successful_starts':'成功起点','selected_relative_error':'相对误差','successful_span_over_true':'跨度/真值','successful_starts_interior':'界内','passed':'通过'}))+'\n\n'+('完整实测终态已保留。' if full else '当前仅合成终态；计划中的实测阶段未包含于本报告，不能判定完整实验完成。')+' 成功起点数见训练表；只有一个成功起点时跨度0不构成稳定性证据。筛查界是工程诊断，不是独立生理资格。')
    component_rows=[]
    for group,prep in preps.items():
        if prep['spec']['kind']!='synthetic':continue
        diagnostic=prep['metadata'].get('truth_component_diagnostic')
        if diagnostic is not None:
            component_rows.append(dict(group=group,**{k:json.dumps(v,ensure_ascii=False) for k,v in diagnostic.items() if k!='interpretation'}))
    if component_rows:
        pd.DataFrame(component_rows).to_csv(out/'synthetic_truth_components.csv',index=False)
        pages.append('## 合成初态与驱动的反事实检查\n\n'+table(pd.DataFrame(component_rows).map(lambda v: '['+', '.join(f'{x:.4f}' for x in json.loads(v))+']' if isinstance(v,str) and v.startswith('[') else v))+'\n初态独立输入和驱动独立输入分别重新经过非线性前向模型、条件矩阵和处理算子。它们不是可相加的方差分解，不能解释成解释率。完整反事实数组保留于prepared，未据此改变目标或评分。')
    pilot_provenance=None
    if pilot_truth_audit is not None:
        import shutil
        pilot_truth_audit=Path(pilot_truth_audit).resolve()
        pilot=json.loads(pilot_truth_audit.read_text())
        if pilot.get('schema')!='conditional_gain_pilot_truth_audit_v1':raise ValueError('Unexpected pilot truth audit schema')
        rows=pilot['rows'];crows=[r for r in rows if r['method']=='conditional_trained_gain']
        if len(rows)!=6 or {r['trial'] for r in crows}!={0,4}:raise ValueError('Pilot audit identities incomplete')
        background=out/'pilot_background';background.mkdir()
        shutil.copy2(pilot_truth_audit,background/pilot_truth_audit.name)
        original_png=pilot_truth_audit.parent/'driver_dc_and_flow.png'
        if original_png.exists():shutil.copy2(original_png,background/original_png.name)
        pilot_provenance=dict(source=str(pilot_truth_audit),sha256=hashlib.sha256(pilot_truth_audit.read_bytes()).hexdigest(),source_png=str(original_png) if original_png.exists() else None,source_png_sha256=hashlib.sha256(original_png.read_bytes()).hexdigest() if original_png.exists() else None)
        fig,axes=plt.subplots(1,3,figsize=(7.1,3.4));x=np.arange(2)
        axes[0].bar(x,[r['raw_driver_nrmse'] for r in crows]);axes[0].set_title('C 原坐标驱动NRMSE');axes[0].axhline(.5,color='red',ls=':')
        axes[1].bar(x,[100*r['dc_fraction_of_squared_error'] for r in crows]);axes[1].set_title('DC占驱动误差SSE / %');axes[1].set_ylim(0,105)
        axes[2].bar(x-.16,[r['flow_mean_fit']-r['flow_mean_truth'] for r in crows],width=.3,label='实际均值偏差');axes[2].bar(x+.16,[r['linear_rest_equilibrium_shift_from_driver_dc'][1] for r in crows],width=.3,label='βc/γ近似');axes[2].set_title('f均值偏差');axes[2].legend(fontsize=6)
        for ax in axes:ax.set_xticks(x,[f'trial {r["trial"]}' for r in crows])
        pages.append('## Pilot背景：驱动DC与初态补偿\n\n'+save(fig,'pilot_driver_dc_background')+'\n仅为先前两个pilot验证trial的背景证据，不计入本次完整合成分母。C的原坐标驱动NRMSE分别为 '+ '/'.join(f'{r["raw_driver_nrmse"]:.3f}' for r in crows)+'；其中 '+ '/'.join(f'{100*r["dc_fraction_of_squared_error"]:.2f}%' for r in crows)+' 的驱动误差SSE来自DC。去均值误差只能诊断，不能替换注册的raw指标。小信号附近，驱动DC偏差c与初态/静态flow偏移βc/γ相互补偿；完整非线性血管系统不具有一般加性平移不变性。f接近1或曲线小幅并不意味着生理状态已验证。\n原审计JSON和原PNG保留于pilot_background；本图以保留标量重新作图240dpi，未重新拟合。')
    synthetic=good(data[(data.kind=='synthetic')&(data['mode']=='full')])
    recovery=states[(states.kind=='synthetic')&states.status.eq('completed')&truth(states.converged)].groupby('method')[[n+'_truth_rmse' for n in ['r','s','f','v','p','q']]].mean().reset_index()
    driver=synthetic.groupby('method').driver_nrmse.mean().reset_index()
    pages.append('## 合成参数、驱动和状态恢复\n\n'+table(driver)+'\n'+table(recovery)+'\n驱动NRMSE以真实驱动SD归一化；状态为各自原始坐标RMSE，不能跨状态比较幅度。新合成真值已在原生Hb加噪之前施加条件矩阵，12组跨g配对驱动、初态和标准化噪声；不是人体标定，也不是从TDDR后信号恢复真实Hb的证明。')
    mechanism_path=run/'training_mechanism_audit'/'review.json'
    if full and mechanism_path.exists():
        mechanism=json.loads(mechanism_path.read_text())
        if mechanism['total_starts']!=72 or mechanism['successful_starts']!=int(((starts.kind=='measured')&starts.status.eq('completed')&truth(starts.converged)).sum()):raise ValueError('Training mechanism audit differs from retained training starts')
        import shutil
        shutil.copy2(mechanism_path,out/'training_mechanism_review.json')
        cross=pd.DataFrame(mechanism['crossfold']);cross['beta_max_over_min']=cross.beta_max/cross.beta_min
        selected_measured=choices[choices.kind=='measured'];selected_count=int(selected_measured.status.eq('completed').sum())
        pages.append('## 实测训练：起点稳定与跨折可识别性\n\n'+f'{mechanism["successful_starts"]}/72起点成功，{selected_count}/24组有成功选择；其余组完整保留训练失败并传递到验证分母。成功起点在组内最大β相对跨度为 {100*mechanism["maximum_within_group_relative_beta_span"]:.4f}%。这不构成跨折参数稳定性证据，只有一个成功起点的组也不能用跨度0宣称稳定。\n\n'+table(cross)+'\nS09跨折β可变化约5–6倍，即使各折内多起点近乎一致。因此当前有效耦合尚不能解释为稳定的被试生理特征。共同训练trial的惩罚、驱动DC和初态补偿审计保留于training_mechanism_review.json；训练集合重叠，跨折不是独立重复。九个失败起点仍是数值失败，不能用末有效参数补成成功。')
    example_specs=[(g,f'synthetic_g{g:g}_slow_r0',0) for g in [.5,1.,2.]]
    fig_observation,obs_axes=plt.subplots(3,3,figsize=(7.1,6.6),sharex=True)
    fig_state,state_axes=plt.subplots(3,3,figsize=(7.1,6.6),sharex=True)
    example_status=[];time=np.arange(cfg['tensor']['steps'])*cfg['tensor']['dt_s']-5.
    for column,(true_g,group,trial) in enumerate(example_specs):
        with np.load(run/'prepared'/f'{group}.npz',allow_pickle=False) as a:
            observed=a['target'][trial].copy();true_driver=a['truth'][trial].copy();true_states=a['truth_states'][trial].copy()
        for j,c in enumerate(COMPONENTS):obs_axes[j,column].plot(time,observed[:,j],color='black',lw=1.,label='观测')
        for j,values in enumerate([true_driver,true_states[:,2],true_states[:,4]-true_states[:,3]]):state_axes[j,column].plot(time,values,color='black',lw=1.,label='真实')
        for method,color,style in zip(methods,['.5','tab:blue','tab:orange'],['--','-',':']):
            record=data[(data.group==group)&(data.method==method)&(data['mode']=='full')&(data.trial==trial)].iloc[0]
            example_status.append(dict(group=group,trial=trial,method=method,status=record.status,converged=bool(truth(pd.Series([record.converged])).iloc[0])))
            if record.status!='completed' or not bool(truth(pd.Series([record.converged])).iloc[0]):continue
            with np.load(run/'cells'/f'{group}__{method}__full'/'trajectories.npz',allow_pickle=False) as a:
                idx=np.flatnonzero(a['trial_indices']==trial)
                if len(idx)!=1:raise ValueError('Fixed synthetic example identity missing')
                k=int(idx[0]);prediction=a['prediction'][k].copy();state=a['states'][k].copy()
            if not np.isfinite(prediction).all() or not np.isfinite(state).all():raise ValueError('Successful example has nonfinite arrays')
            for j,c in enumerate(COMPONENTS):obs_axes[j,column].plot(time,prediction[:,j],color=color,ls=style,lw=.9,label='ABC'[methods.index(method)])
            for j,values in enumerate([state[:,0],state[:,2],state[:,4]-state[:,3]]):state_axes[j,column].plot(time,values,color=color,ls=style,lw=.9,label='ABC'[methods.index(method)])
        obs_axes[0,column].set_title(f'真实g={true_g:g} / slow r0 trial0',fontsize=8)
        state_axes[0,column].set_title(f'真实g={true_g:g} / slow r0 trial0',fontsize=8)
        for j in range(3):
            if column==0:
                obs_axes[j,column].set_ylabel(COMPONENTS[j],fontsize=8)
                state_axes[j,column].set_ylabel(['原坐标驱动','f','p − v'][j],fontsize=8)
            for ax in [obs_axes[j,column],state_axes[j,column]]:
                ax.tick_params(labelsize=7);ax.axvline(0,color='.8',ls=':',lw=.6)
                if j==2:ax.set_xlabel('相对事件 / 秒',fontsize=8)
        obs_axes[0,column].legend(fontsize=6,ncol=2);state_axes[0,column].legend(fontsize=6,ncol=2)
    pd.DataFrame(example_status).to_csv(out/'fixed_synthetic_example_status.csv',index=False)
    pages.append('## 固定合成示例：共享状态联合重建\n\n'+save(fig_observation,'synthetic_fixed_observation_curves')+'\n固定身份为g=0.5/1/2、slow、rep0、trial0，不按误差选样。黑色为实际加噪观测；灰虚A、蓝实B、橙点C为保留full预测。列间使用各自原始观测尺度；未去均值、翻转符号、平移时间或调整幅度。失败身份保留于下一页状态表，不画成成功。')
    pages.append('## 固定合成示例：原驱动及生理状态\n\n'+save(fig_state,'synthetic_fixed_driver_state_curves')+'\n同三条固定身份。黑色是真值；三臂保留原驱动、f及p−v。未进行任何对齐或去均值；原驱动DC差异必须计入恢复误差。各子图独立纵轴以保留可读性，不意味着状态幅度相同。f接近静息1、p−v接近0仅为图像观察，不是生理资格。\n\n'+table(pd.DataFrame(example_status)))
    measured=states[(states.kind=='measured')&states.status.eq('completed')&truth(states.converged)]
    if full:
        rows=[]
        for (pipeline,method),part in measured.groupby(['pipeline','method']):
            rows.append(dict(pipeline=pipeline,method=method,success=len(part),**{name+'_min':part[name+'_min'].min() for name in ['f','v','p','q']},**{name+'_max':part[name+'_max'].max() for name in ['f','v','p','q']},initial_deviation_over05=int((part.initial_max_rest_deviation>.5).sum()),f_min_percent_from_rest=100*(part.f_min.min()-1),f_max_percent_from_rest=100*(part.f_max.max()-1),trial_f_min_q05=float(part.f_min.quantile(.05)),trial_f_max_q95=float(part.f_max.quantile(.95))))
        pd.DataFrame(rows).to_csv(out/'measured_state_tail_summary.csv',index=False)
        pages.append('## 实测状态与完整曲线\n\n'+table(pd.DataFrame(rows)[['pipeline','method','success','f_min_percent_from_rest','f_max_percent_from_rest','trial_f_min_q05','trial_f_max_q95','initial_deviation_over05']])+'\n初态偏差>0.5仅为描述性极端计数，不是已验证正常界。数学正值、Hb balance和积分复核是数学域条件，不是经验生理范围。若f解释为静息归一局部血流，f=0.1即−90%、f=1.9即+90%；这些是相对幅度，前者可描述为极端且未生理验证，不能据此临床诊断，也不能只凭单篇研究均值判后者不可能。新的映射和β缩小使f接近1亦不自动证明生理合理。跨trial极值的5%/95%分位数表示尾部分布，不是正常界。\n\n完整72身份×两管道在FULL_CURVE_ATLAS.pdf固定顺序展示，不按误差挑样；后三列三臂独立纵轴，便于辨认小幅和极端状态。')
        atlas=fitz.open()
        for pipeline in pipelines:
            folder=out/pipeline;folder.mkdir()
            flow_full_curve_atlas(run,folder,cfg,data[data.pipeline==pipeline],labels,methods,both_states=True,context_label=pipeline)
            with fitz.open(folder/'FULL_CURVE_ATLAS.pdf') as part:atlas.insert_pdf(part)
        if len(atlas)!=48 or sum(len(p.get_images()) for p in atlas)!=48:raise ValueError('Incomplete bitmap atlas')
        atlas.save(out/'FULL_CURVE_ATLAS.pdf',garbage=4,deflate=True);atlas.close()
    pages.append('## 来源与限制\n\n'+f'运行 {run.name}；只读取保留prepared/cells/training与父运行 {parent.name}，未拟合或重评分。目标/SD逐单元核验见target_SD_audit.csv。实测A保留父运行失败，属于继承对照；合成三臂全部新运行。\n\n所有完整图为240dpi PNG；PDF文本可搜。WPS未检查。需要结合完整图册与失败分母判断恢复，不以成功条件均值宣称全体达标。')
    export_pdf('\n\n'.join(pages),out,figures,split_images=False)
    audit=json.loads((out/'export_validation.json').read_text())
    audit.update(pilot_truth_audit=pilot_provenance,expected_rows=expected_rows,terminal_cells=expected_rows//6,target_SD_checks=len(array_checks),full_measured=full,source_report=str(Path(__file__).resolve()),source_report_sha256=hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),source_manifest_sha256=hashlib.sha256((input_root/'manifest.json').read_bytes()).hexdigest(),terminal_input_snapshot=str(input_root))
    with fitz.open(out/'REPORT.pdf') as pdf:
        if any(p.get_drawings() for p in pdf):
            # Table borders are vector body content; only figure rasterization is required.
            audit['body_vector_table_borders_allowed']=True
        for i,page in enumerate(pdf):page.get_pixmap(matrix=fitz.Matrix(1.3,1.3)).save(out/f'preview_page_{i+1:02d}.png')
    (out/'export_validation.json').write_text(json.dumps(audit,indent=2))


def render_waveform_diagnostic(run,out,volume_run=None):
    import yaml
    cfg=yaml.safe_load((run/'resolved_config.yaml').read_text())
    if json.loads((run/'manifest.json').read_text()).get('execution')!='completed':raise ValueError('diagnostic must be terminal')
    if out.exists():raise ValueError('choose a fresh versioned report export')
    out.mkdir(parents=True);(out/'figures').mkdir()
    font=Path('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc')
    if font.exists():
        font_manager.fontManager.addfont(str(font));plt.rcParams['font.family']=font_manager.FontProperties(fname=str(font)).get_name()
    plt.rcParams.update({'font.size':10,'axes.unicode_minus':False})
    frame=pd.read_csv(run/'linear_metrics.csv');joint=frame[frame['mode']=='joint'].copy()
    nonlinear=json.loads((run/'nonlinear_summary.json').read_text())['results']
    retry=run/'integration_failure_recheck_v1/summary.json'
    if retry.exists():
        replacement=json.loads(retry.read_text())
        nonlinear=[replacement if r['group']==replacement['group'] and r['arm']==replacement['arm'] else r for r in nonlinear]
    volumes=[]
    if volume_run is not None:
        if json.loads((volume_run/'manifest.json').read_text()).get('execution')!='completed':raise ValueError('volume followup must be terminal')
        volumes=json.loads((volume_run/'summary.json').read_text())['results']
    focus=joint[(joint.group.str.contains('subject_09'))&(joint.trial==3)].copy()
    focus['pipeline']=focus.group.str.split('__').str[0]
    metrics=['nrmse_'+c for c in COMPONENTS];aggregate=[]
    for (group,method),part in joint.groupby([joint.group.str.split('__').str[0],'method']):
        aggregate.append(dict(pipeline=group,method=method,trials=len(part),**dict(zip(metrics,np.sqrt(np.mean(part[metrics].to_numpy()**2,axis=0))))))
    aggregate=pd.DataFrame(aggregate);aggregate.to_csv(out/'validation_aggregate.csv',index=False)
    nonrows=[]
    for r in nonlinear:
        nonrows.append(dict(pipeline=r['group'].split('__')[0],arm=r['arm'],status=r['status'],
            **dict(zip(metrics,r.get('nrmse',[np.nan]*3))),integration_difference=r.get('integration_max_difference_training_sd')))
    nonframe=pd.DataFrame(nonrows);nonframe.to_csv(out/'nonlinear_focus.csv',index=False)
    figures=[]
    def save(fig,name):
        fig.tight_layout();fig.savefig(out/'figures'/f'{name}.png',dpi=240,bbox_inches='tight');plt.close(fig)
        figures.append(out/'figures'/f'{name}.png')
        return f'![{name}](figures/{name}.png)'
    fig,axes=plt.subplots(3,2,figsize=(12,8.2),sharex=True)
    correlations=[];landmarks=[];volrows=[]
    for col,pipeline in enumerate(['no_motion','mne_tddr']):
        group=f'{pipeline}__subject_09_o3';folder=run/'nonlinear'/f'{group}__same_objective'
        with np.load(folder/'trajectory.npz') as z:
            sd=z['normalizer'];y=z['target']/sd;pred=z['prediction']/sd
        residual=y-pred
        correlations.append(dict(pipeline=pipeline,target_hb_correlation=np.corrcoef(y[:,1:].T)[0,1],residual_hb_correlation=np.corrcoef(residual[:,1:].T)[0,1]))
        with np.load(run/'linear'/group/'predictions.npz') as z:
            curves={'原非线性模型':pred,'线性：共同血容量':z['joint__slow_volume'][0]/sd,
                    '线性：氧交换':z['joint__slow_exchange'][0]/sd}
        if volume_run is not None:
            with np.load(volume_run/group/'predictions.npz') as z:curves['线性：训练选成分比例']=z['prediction'][0]/sd
        colors=['#dc654a','#168a89','#b99122','#745bb1']
        for j,c in enumerate(COMPONENTS):
            ax=axes[j,col];t=np.arange(120)*.25;ax.plot(t,y[:,j],color='#243f4a',lw=2.1,label='观测目标')
            for (label,curve),color in zip(curves.items(),colors):ax.plot(t,curve[:,j],color=color,lw=1.35,label=label,alpha=.95)
            ax.axhline(0,color='#999',lw=.5);ax.grid(alpha=.18);ax.set_ylabel(c+' / 训练SD')
            if j==0:ax.set_title(('无运动校正' if pipeline=='no_motion' else 'MNE TDDR')+' · S09 / fold 3 / trial 3')
            if j==2:ax.set_xlabel('窗内时间（秒；0秒为事件前5秒）')
        for a,b,label,sign in [(6,12,'第一峰',1),(19,26,'第二峰',1),(26,29.75,'末端谷',-1)]:
            idx=np.flatnonzero((t>=a)&(t<=b));i=idx[np.argmax(sign*y[idx,1])]
            for method,curve in {'观测目标':y,**curves}.items():
                landmarks.append(dict(pipeline=pipeline,feature=label,observed_extremum_time=t[i],method=method,
                    HbO_at_observed_time=curve[i,1],HbR_at_observed_time=curve[i,2]))
    axes[0,0].legend(fontsize=8,ncol=2,loc='upper left')
    fig.suptitle('固定困难身份：成分对照改善波形，但不能确定其生理来源',fontsize=14)
    fig.subplots_adjust(top=.94)
    mainfigure=save(fig,'waveform_mechanisms')
    pd.DataFrame(landmarks).to_csv(out/'focus_landmarks.csv',index=False)
    pd.DataFrame(correlations).to_csv(out/'focus_correlations.csv',index=False)
    fig,axes=plt.subplots(2,2,figsize=(11,7),sharex=True)
    for col,pipeline in enumerate(['no_motion','mne_tddr']):
        group=f'{pipeline}__subject_09_o3'
        for arm,label in [('same_objective','原模型'),('no_curvature','去曲率惩罚'),('no_initial_prior','去初态惩罚'),('no_flow_prior','去血流惩罚'),('no_penalties','去全部惩罚')]:
            with np.load(run/'nonlinear'/f'{group}__{arm}'/'trajectory.npz') as z:
                y=z['target']/z['normalizer'];pred=z['prediction']/z['normalizer']
            record=next(r for r in nonlinear if r['group']==group and r['arm']==arm)
            if record['status']!='completed':label+='（未收敛）'
            for j in range(2):axes[j,col].plot(np.arange(120)*.25,pred[:,j+1],label=label,lw=1.2,ls='--' if record['status']!='completed' else '-')
        for j in range(2):
            axes[j,col].plot(np.arange(120)*.25,y[:,j+1],color='#263d48',lw=2,label='观测目标');axes[j,col].grid(alpha=.2)
            axes[j,col].set_ylabel(COMPONENTS[j+1]+' / 训练SD')
        axes[0,col].set_title(pipeline);axes[1,col].set_xlabel('窗内时间（秒）')
    axes[0,0].legend(fontsize=8,ncol=2);ablationfigure=save(fig,'penalty_ablations')
    fraction_text=''
    if volumes:
        for r in volumes:
            v=np.array([x['nrmse'] for x in r['rows']]);volrows.append(dict(group=r['group'],selected_fraction=r['selected_fraction'],**dict(zip(metrics,np.sqrt(np.mean(v*v,axis=0))))))
        fraction_text='## 训练选择额外成分比例\n\n'+table(pd.DataFrame(volrows))+'\n该比例是新增慢成分的有效HbR/HbT观测loading，不是测得的血氧饱和度。六个慢系数在各验证窗口上重新拟合，因此这些结果属于重建，不是独立预测。\n'
        fraction_focus=[];fraction_aggregate=[]
        parent=Path(__file__).resolve().parents[2]/cfg['parent_run']
        for pipeline in ('no_motion','mne_tddr'):
            before=[];after=[]
            for r in volumes:
                if not r['group'].startswith(pipeline):continue
                old=json.loads((parent/'cells'/f"{r['group']}__conditional_trained_gain__full"/'result.json').read_text())['rows']
                if [x['trial'] for x in old]!=[x['trial'] for x in r['rows']]:raise ValueError('volume paired identities differ')
                before.extend([[x[k] for k in metrics] for x in old]);after.extend([x['nrmse'] for x in r['rows']])
                if 'subject_09' in r['group']:
                    fraction_focus.append(dict(pipeline=pipeline,selected_fraction=r['selected_fraction'],**dict(zip(metrics,r['rows'][0]['nrmse']))))
            for label,values in [('retained_nonlinear',before),('train_selected_volume',after)]:
                values=np.asarray(values)
                fraction_aggregate.append(dict(pipeline=pipeline,method=label,trials=len(values),all_three_below_05=int(np.all(values<.5,axis=1).sum()),**dict(zip(metrics,np.sqrt(np.mean(values**2,axis=0))))))
        fraction_text+='\nS09困难身份：\n\n'+table(pd.DataFrame(fraction_focus))+'\n与原非线性模型的同身份汇总（两个模型类别不同）：\n\n'+table(pd.DataFrame(fraction_aggregate))
        relevant=pd.DataFrame(landmarks)
        relevant=relevant[(relevant.pipeline=='mne_tddr')&relevant.method.isin(['观测目标','原非线性模型','线性：训练选成分比例'])]
        fraction_text+='\n在观测HbO极值时刻的振幅核查；不是重新挑选预测极值：\n\n'+table(relevant[['feature','observed_extremum_time','method','HbO_at_observed_time','HbR_at_observed_time']])+'\n第一峰仍偏低，末端HbO谷深也未完全跟随；NRMSE达标不能代替完整形态验收。\n'
        pd.DataFrame(fraction_aggregate).to_csv(out/'volume_fraction_paired_aggregate.csv',index=False)
        pd.DataFrame(volrows).to_csv(out/'volume_fraction_groups.csv',index=False)
    keep=['fixed','beta_factor','kappa','gamma','tau','alpha','E0','venous_viscoelastic_s','slow_volume','slow_exchange']
    text='# HbO/HbR 波形偏离：参数与缺失成分诊断\n\n2026-09-28。原始困难身份不替换；本报告汇总合成、训练选参线性筛查和精确非线性惩罚对照。\n\n'
    text+='主要结论：单独松开曲率、初态或血流约束不能恢复S09困难波形。残差以同向Hb变化为主，额外共同血容量方向比等维氧交换方向更能恢复HbO；来源仍不唯一。整体NRMSE与困难窗口形态分别报告。\n\n'
    text+='## 理论定位\n\n当前 s_dot=βr−κs−γ(f−1)，故β控制输入幅度，κ与γ控制血流滤波；τ和α控制血容量时序，E0控制流量与氧提取的关系。增加β不提供新的时间过程。\n\n'
    text+='由实现方程在静息点线性化，c=1+(1−E0)ln(1−E0)/E0，慢变化稳态 dq/dp=1−(1−c)/α。默认E0=α=0.32时约−1.561，表示慢血流输入倾向于使总Hb上升而HbR下降。这个稳态方向不是任意瞬态或滤波后曲线必须反相关的定律。\n\n'
    text+='困难窗口残差的HbO/HbR同向变化因此提示：单一流量—提取耦合不足，或EEG代理未观测到相应驱动。新增共同血容量项保持其自身饱和比例，氧交换项保持总Hb不变；它们仅比较表达能力。原Tak模型还包含静脉黏弹性与光学混合因素，当前核心无前者。本次只在独立线性诊断中加入黏弹性，保留原非线性模型不变。\n\n'
    text+=table(pd.DataFrame(correlations))+'\n'
    text+='## 范围与方法\n\n三被试同outer=3，每人18训练/6验证，共18个唯一验证窗口、两种处理路径、36个路径×窗口。两路径不是独立重复。七种参数方向分别按训练SSE选网格，不联合优化；目标和SD均沿用父运行。新增慢成分各使用相同六个预声明余弦基。额外成分降低同目标残差具有自由度优势，等维对照仅部分控制此问题，不建立生理因果归因。\n\n'
    text+='合成匹配、τ、κ、黏弹性、血容量和氧交换六种控制全部通过；体积分量比例的独立合成恢复另见配套运行。原理测试不能代替真实生理验证。\n\n'
    text+='## 困难样例的精确非线性对照\n\n'+table(nonframe)+'\n'
    text+='失败行的NRMSE是最后有效粗积分轨迹的描述，不是成功结果。MNE全去惩罚未达到梯度收敛；两个Hb-only全去惩罚分支都未通过完整数值要求，伴随巨大驱动及近零flow。原failed_worker记录保留，补充重放在integration_failure_recheck_v1记录优化预算失败和细积分物理域失败，未将其改判成功。\n\n'+ablationfigure+'\n\n'
    text+='## 困难样例：训练选参和等维成分对照\n\n'+table(focus[focus.method.isin(keep)][['pipeline','method','parameter_value',*metrics]])+'\n\n'+mainfigure+'\n\n'
    text+='S09的α与E0分别选到工程网格上端0.96和0.9，但仍未解决两个Hb波形。此结果不能解释为已测出该被试的α或E0；可能是在补偿模型或观测遗漏。其他参数组合与完整非线性参数扩展尚未检验。S09黏弹性选择0，其他被试可能选择非零，因此不全盘否定该过程。\n\n'
    text+='## 同折全部验证窗口\n\n'+table(aggregate[aggregate.method.isin(keep)])+'\n各分量为等被试、等session的NMSE均值开方（本面板各组同样6条、每session2条）。仍须查看每个窗口，不能用总体达标掩盖困难身份。\n\n'
    text+=fraction_text+'\n'
    text+='## 数值与解释限制\n\njoint线性SVD的1e−10/1e−12/1e−14截断结果一致；Hb-only无惩罚解可用巨大驱动拟合Hb，且明显离开小信号范围，不能当作生理可行性证明。部分单参数线性解也离开小信号范围；完整逐窗幅度见linear_metrics.csv。成分模型仍依赖条件光学映射与窗口滤波，不能区分皮层局部血容量、头皮系统性成分、空间混合、测量串扰或EEG代理遗漏。\n\n'
    text+='下一步建议先建立含独立额外血容量成分的候选，使用不同时间段或独立被试检验预测与参数稳定性；若要确定来源，再使用短距离通道、空间共同成分或已记录的心电/呼吸等独立信息。尚未执行新的原始数据处理、完整非线性扩展训练或受保护评价。\n\n'
    text+='## 来源与可复现性\n\n'+f'主运行：{run}。配套比例运行：{volume_run}。原结果不覆盖。源快照、冻结合同、监督器命令、资源检查、完整成功/失败结果均在运行目录。图为240dpi PNG嵌入PDF，文字和表格可搜索；WPS未检查。\n\n'
    text+='生理模型依据：[Tak et al., 2015](https://pmc.ncbi.nlm.nih.gov/articles/PMC4401444/)；黏弹性出流实现核对：[SPM官方spm_fx_fnirs](https://raw.githubusercontent.com/spm/spm12/main/toolbox/dcm_fnirs/spm_fx_fnirs.m)。头皮任务诱发成分是文献支持的候选来源，但不是本数据已证实的结论：[Kirilina et al., 2012](https://pmc.ncbi.nlm.nih.gov/articles/PMC3348501/)。\n'
    export_pdf(text,out,figures)


def step_control_audit(run):
    """Read every paired start, including failures, without fitting a model."""
    import yaml
    cfg=yaml.safe_load((run/'resolved_config.yaml').read_text())
    if cfg.get('schema')!='shared_driver_conditional_step_control_v2':
        raise ValueError('paired solver report requires the v2 contract')
    rows=[]
    for kind in ('synthetic','measured'):
        summary=json.loads((run/f'{kind}_paired_summary.json').read_text())
        expected=cfg[kind+'_starts']
        if not summary['complete'] or summary['terminal_pairs']!=expected:
            raise ValueError('paired solver report requires complete denominators')
        if len({(p['group'],p['start']) for p in summary['rows']})!=expected:
            raise ValueError('duplicate solver pair')
        for pair in summary['rows']:
            records={};arrays={}
            for arm in ('armijo','quadratic_interpolation'):
                folder=run/'arms'/arm/'training'/f"{pair['group']}__conditional_trained_gain"/f"start_{pair['start']}"
                records[arm]=json.loads((folder/'result.json').read_text())
                if (folder/'trajectories.npz').exists():
                    with np.load(folder/'trajectories.npz',allow_pickle=False) as a:arrays[arm]={k:a[k].copy() for k in a.files}
            old,new=records['armijo'],records['quadratic_interpolation']
            success=lambda r:r['status']=='completed' and r.get('converged',False)
            row=dict(kind=kind,group=pair['group'],start=pair['start'],
                subject=new['spec'].get('subject','synthetic'),pipeline=new['spec'].get('optical_pipeline','synthetic'),
                old_success=success(old),new_success=success(new),
                recovered=success(new) and not success(old),regressed=success(old) and not success(new))
            for prefix,record in [('old',old),('new',new)]:
                for key in ('status','objective','relative_gain','optimization_evaluations','seconds','cpu_seconds',
                            'projected_scaled_gradient_inf_norm','convergence_reason'):
                    row[prefix+'_'+key]=record.get(key)
                row[prefix+'_iterations']=record.get('starts',[{}])[0].get('iterations')
                row[prefix+'_integration_gap']=max((c['difference_training_sd'] for c in record.get('integration_checks',[])),default=np.nan)
                checked=np.asarray([record.get('optimization_evaluations',np.nan),
                    record.get('projected_scaled_gradient_inf_norm',np.nan),row[prefix+'_integration_gap']],dtype=float)
                if success(record) and (not np.isfinite(checked).all() or checked[0]>3600 or
                    checked[1]>1e-4 or checked[2]>.005):
                    raise ValueError('successful start violates frozen budget/gradient/integration check')
            counts=new.get('starts',[{}])[0].get('interpolation_counts',{})
            row.update(interpolation_attempts=counts.get('attempts',0),interpolation_selected=counts.get('selected',0))
            inherited=run/'arms/armijo/baseline_training'/f"{pair['group']}__conditional_trained_gain"/f"start_{pair['start']}/result.json"
            parent=json.loads(inherited.read_text())
            row['legacy_reproduced']=all(old.get(k)==parent.get(k) for k in
                ('status','converged','objective','parameter_value','optimization_evaluations'))
            if len(arrays)==2:
                a,b=arrays['armijo'],arrays['quadratic_interpolation']
                if any(not np.array_equal(a[k],b[k]) for k in ('target','normalizer','trial_indices')):
                    raise ValueError('paired solver target/SD/trial mismatch')
                row['prediction_max_difference_sd']=float(np.max(abs(a['prediction']-b['prediction'])/a['normalizer']))
                row['driver_max_difference']=float(np.max(abs(a['driver']-b['driver'])))
                row['flow_max_difference']=float(np.max(abs(a['states'][:,:,2]-b['states'][:,:,2])))
            if old.get('objective') is not None and new.get('objective') is not None:
                row['objective_relative_difference']=abs(new['objective']-old['objective'])/max(abs(old['objective']),1e-12)
                row['gain_relative_difference']=abs(new['relative_gain']/old['relative_gain']-1)
            rows.append(row)
    frame=pd.DataFrame(rows);aggregate=[]
    for (kind,subject,pipeline),part in frame.groupby(['kind','subject','pipeline'],sort=False):
        row=dict(kind=kind,subject=subject,pipeline=pipeline,starts=len(part),
            old_success=int(part.old_success.sum()),new_success=int(part.new_success.sum()),
            recovered=int(part.recovered.sum()),regressions=int(part.regressed.sum()))
        for metric in ('optimization_evaluations','seconds','cpu_seconds'):
            for prefix in ('old','new'):row[prefix+'_'+metric]=float(part[prefix+'_'+metric].sum())
            row[metric+'_reduction']=1-row['new_'+metric]/row['old_'+metric]
        aggregate.append(row)
    return frame,pd.DataFrame(aggregate)


def render_step_control(run,out):
    if out.exists() and any(out.iterdir()):raise ValueError('choose a new versioned report output')
    frame,groups=step_control_audit(run)
    out.mkdir(parents=True,exist_ok=True);(out/'figures').mkdir()
    frame.to_csv(out/'paired_starts.csv',index=False);groups.to_csv(out/'subject_pipeline_summary.csv',index=False)
    font=Path('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc')
    if font.exists():
        font_manager.fontManager.addfont(str(font));plt.rcParams['font.family']=font_manager.FontProperties(fname=str(font)).get_name()
    plt.rcParams.update({'font.size':9,'axes.unicode_minus':False});figures=[]
    def save(fig,name):
        fig.tight_layout();fig.savefig(out/'figures'/f'{name}.png',dpi=240);plt.close(fig);figures.append(name)
        return f'![{name}](figures/{name}.png)'
    syn=frame[frame.kind=='synthetic'];real=frame[frame.kind=='measured'];rows=[]
    for kind,part in [('合成',syn),('实测',real)]:
        rows.append(dict(数据=kind,起点对=len(part),旧成功=int(part.old_success.sum()),新成功=int(part.new_success.sum()),
            修复=int(part.recovered.sum()),退步=int(part.regressed.sum()),
            优化调用降幅=f'{100*(1-part.new_optimization_evaluations.sum()/part.old_optimization_evaluations.sum()):.2f}%',
            CPU降幅=f'{100*(1-part.new_cpu_seconds.sum()/part.old_cpu_seconds.sum()):.2f}%'))
    pages=['# 同目标求解器：收敛与计算量\n\n'+table(pd.DataFrame(rows))+
        '\n同一固定目标、起点、3600次试次评估、90次迭代、梯度阈值1e−4。只增加一次有条件的二次插值；不改模型、先验或积分器。\n\n'+
        '两臂均重新运行，CPU时间含最终重放、信息矩阵及积分复核。墙钟受共享机器负载影响；调用量是确定性计算指标。分母含失败起点。']
    part=groups[groups.kind=='measured'];x=np.arange(len(part));fig,axes=plt.subplots(1,2,figsize=(11,4.5))
    labels=[r.subject.replace('subject_','S')+'\n'+r.pipeline for r in part.itertuples()]
    for prefix,offset,color,label in [('old',-.18,'#2877ab','原Armijo'),('new',.18,'#dc8030','二次插值')]:
        bars=axes[0].bar(x+offset,part[prefix+'_success'],.36,color=color,label=label)
        axes[0].bar_label(bars,padding=2,fontsize=8)
        axes[1].bar(x+offset,part[prefix+'_optimization_evaluations'],.36,color=color,label=label)
    for ax in axes:ax.set_xticks(x,labels,fontsize=7);ax.legend(fontsize=8)
    axes[0].set_ylabel('成功起点 / 12');axes[0].set_ylim(0,13);axes[1].set_ylabel('优化期试次forward调用总数')
    pages.append('## 三被试、两条处理路径\n\n'+save(fig,'solver_summary')+
        '\n全部四个划分、每组三个起点均计入。三位被试来自同一公开数据集；两种光学处理使用同源窗口，不是六个独立数据集。')
    fig,axes=plt.subplots(1,3,figsize=(11,3.8))
    for ax,metric,label in zip(axes,['optimization_evaluations','cpu_seconds','seconds'],['优化调用','进程CPU秒','墙钟秒']):
        for success,marker,name in [(True,'o','两臂成功'),(False,'x','至少一臂失败')]:
            sub=real[(real.old_success&real.new_success)==success]
            ax.scatter(sub['old_'+metric],sub['new_'+metric],s=17,marker=marker,label=name)
        lim=max(real['old_'+metric].max(),real['new_'+metric].max())*1.05
        ax.plot([0,lim],[0,lim],'k--',lw=.8);ax.set(xlim=(0,lim),ylim=(0,lim),xlabel='原Armijo '+label,ylabel='二次插值 '+label)
    axes[0].legend(fontsize=7)
    call_changes=real.new_optimization_evaluations/real.old_optimization_evaluations-1
    increased_count=int((call_changes>0).sum());worst_increase=float(call_changes.max())
    slower_groups=part[part.optimization_evaluations_reduction<0]
    slower_text='；'.join(f"{r.subject.replace('subject_','S')}/{r.pipeline} +{-100*r.optimization_evaluations_reduction:.2f}%" for r in slower_groups.itertuples()) or '无分组总调用增加'
    pages.append('## 每个起点的计算代价\n\n'+save(fig,'paired_cost')+
        f'\n对角线下方表示减少，叉号保留数值失败。{increased_count}/72起点调用增加，最大+{100*worst_increase:.2f}%。分组增加：{slower_text}。相同调用数不保证相同耗时；额外插值本身也有成本。')
    affected=['no_motion__subject_01_o1','no_motion__subject_18_o0','mne_tddr__subject_01_o3','mne_tddr__subject_18_o3']
    fig,axes=plt.subplots(4,2,figsize=(10,10))
    for i,group in enumerate(affected):
        ref=json.loads((run/'arms/quadratic_interpolation/training'/f'{group}__conditional_trained_gain/start_0/result.json').read_text())['parameter_value']
        for arm,color,label in [('armijo','#2877ab','原Armijo'),('quadratic_interpolation','#dc8030','二次插值')]:
            record=json.loads((run/'arms'/arm/'training'/f'{group}__conditional_trained_gain/start_0/result.json').read_text())
            trace=record['trace'];tail=trace[-10:]
            axes[i,0].plot(np.arange(1-len(tail),1),[100*(t['parameter_value']/ref-1) for t in tail],'.-',color=color,label=label,ms=3)
            axes[i,1].semilogy([t['evaluations'] for t in trace],[max(t['projected_scaled_gradient_inf_norm'],1e-14) for t in trace],color=color,label=label)
        axes[i,0].axhline(0,color='.6',lw=.6)
        axes[i,0].set(title=group+' / start0',xlabel='距本臂停止的记录步（0为停止）',ylabel='β相对新算法终值的偏差 / %')
        axes[i,1].axhline(1e-4,color='k',ls='--',lw=.8);axes[i,1].set(xlabel='累计优化forward调用',ylabel='缩放投影梯度')
        axes[i,0].ticklabel_format(axis='y',style='plain',useOffset=False)
        for ax in axes[i]:ax.legend(fontsize=7)
    pages.append('## 原先四个困难训练组\n\n'+save(fig,'convergence_traces')+
        '\n固定展示预声明四组的start0。左图为各自停止前10个记录点，只比较尾段变化；横轴不代表两臂用了相同迭代数。右图以实际调用量展示完整收敛过程。目标变化很小并不等于梯度达到阈值。')
    common=real[real.old_success&real.new_success]
    synthetic_screen=json.loads((run/'arms/quadratic_interpolation/synthetic_training_summary.json').read_text())['engineering_screen']
    availability=real.groupby('group')[['old_success','new_success']].sum()
    gain_spans=[]
    for group,part in real[real.new_success].groupby('group'):
        values=part.new_relative_gain.to_numpy(float)
        gain_spans.append(dict(group=group,successful_starts=len(values),relative_gain_span=(values.max()-values.min())/np.median(values)))
    pd.DataFrame(gain_spans).to_csv(out/'multistart_gain_consistency.csv',index=False)
    agreement=dict(common_success_pairs=len(common),legacy_reproduced=int(frame.legacy_reproduced.sum()),total_pairs=len(frame),
        synthetic_recovery_all_passed=synthetic_screen['all_groups_passed'],
        synthetic_maximum_gain_recovery_error=max(r['selected_relative_error'] for r in synthetic_screen['groups']),
        old_available_training_groups=int((availability.old_success>0).sum()),
        new_available_training_groups=int((availability.new_success>0).sum()),
        new_maximum_multistart_relative_gain_span=max(r['relative_gain_span'] for r in gain_spans),
        regressions=int(real.regressed.sum()),recovered=int(real.recovered.sum()),
        maximum_relative_objective_difference=float(common.objective_relative_difference.max()),
        maximum_relative_gain_difference=float(common.gain_relative_difference.max()),
        maximum_prediction_difference_training_sd=float(common.prediction_max_difference_sd.max()),
        maximum_driver_difference=float(common.driver_max_difference.max()),maximum_flow_difference=float(common.flow_max_difference.max()),
        interpolation_attempts=int(real.interpolation_attempts.sum()),interpolation_selected=int(real.interpolation_selected.sum()),
        measured_starts_with_increased_calls=increased_count,maximum_call_increase_fraction=worst_increase,
        max_integration_difference_training_sd=float(real.new_integration_gap.max()))
    (out/'audit.json').write_text(json.dumps(agreement,indent=2))
    pages.append('## 是否改变了解\n\n'+table(pd.DataFrame([agreement]).T.reset_index().rename(columns={'index':'核验',0:'值'}))+
        '\n一致性只比较两臂均成功的起点；修复失败单独计数。配对目标、训练尺度、试次身份逐元素一致；原算法另与历史结果核对。')
    pages.append('## 范围与限制\n\n'+f'来源：{run.name}。全部108对训练起点，未运行新的验证重建，没有独立数据集泛化证据。数值收敛改善不等于NRMSE已改善，也不证明β为唯一可识别的生理量。\n\n'+
        '默认Armijo仍保留；二次插值为显式选项。旧证据不覆盖。完整指标见paired_starts.csv，分组汇总见subject_pipeline_summary.csv。图为240dpi位图，PDF正文可搜索；WPS未检查。')
    export_pdf('\n\n'.join(pages),out,figures,split_images=False)
    # A short editable presentation shares the audited tables and bitmap figures.
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches,Pt
    deck=Presentation();deck.slide_width=Inches(13.333);deck.slide_height=Inches(7.5)
    def box(slide,text,x,y,w,h,size=22,color='17364B',bold=False):
        shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=shape.text_frame;tf.word_wrap=True
        for i,line in enumerate(text.split('\n')):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.space_after=Pt(14)
            for r in p.runs:r.font.name='Noto Sans CJK SC';r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=RGBColor.from_string(color)
    def slide(title,subtitle):
        s=deck.slides.add_slide(deck.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor(246,248,251)
        box(s,title,.6,.35,12.1,.65,30,bold=True);box(s,subtitle,.65,1.10,12,.48,15,color='597184')
        box(s,f'2026-09-28 · 同目标求解器验证 · {len(deck.slides)}',.65,7.03,12,.3,11,color='597184')
        s.notes_slide.notes_text_frame.text='事实源：'+str(run)+'\n逐起点指标：'+str(out/'paired_starts.csv')+'\n训练数值诊断，不是独立数据集或生理参数资格结论。'
        return s
    reduction=1-real.new_optimization_evaluations.sum()/real.old_optimization_evaluations.sum()
    cpu_reduction=1-real.new_cpu_seconds.sum()/real.old_cpu_seconds.sum()
    s=slide('求解器开发：收敛与速度的完整对照','相同目标、相同起点、相同预算；两臂均重新运行')
    box(s,f"实测成功起点\n{int(real.old_success.sum())}/72 → {int(real.new_success.sum())}/72",.75,2.05,3.8,1.9,28,bold=True)
    box(s,f'优化调用减少\n{100*reduction:.1f}%',4.85,2.05,3.5,1.9,28,bold=True)
    box(s,f'总CPU时间减少\n{100*cpu_reduction:.1f}%',8.95,2.05,3.7,1.9,28,bold=True)
    syn_change=100*(syn.new_optimization_evaluations.sum()/syn.old_optimization_evaluations.sum()-1)
    box(s,f"可用训练组 {agreement['old_available_training_groups']}/24 → {agreement['new_available_training_groups']}/24；救回失败 {int(real.recovered.sum())} 个，退步 {int(real.regressed.sum())} 个。\n合成成功 {int(syn.old_success.sum())}/36 → {int(syn.new_success.sum())}/36；调用量变化 {syn_change:+.2f}%。",.8,4.65,11.8,1.85,23)
    s=slide('覆盖多个被试与处理条件，冻结评价规则','不是独立数据集泛化：所有实测来自同一公开数据集的保留目标')
    box(s,'12组合成条件 × 3起点 = 36对拟合\n3被试 × 2处理路径 × 4划分 × 3起点 = 72对拟合\n每次用18条训练试次；未重新拟合验证目标\n'+f"合成β最大恢复误差 {100*agreement['synthetic_maximum_gain_recovery_error']:.2f}%",.85,1.85,11.6,2.95,24)
    box(s,'原目标、90次迭代、3600次试次forward预算、梯度阈值1e−4均不变。\n采用相同积分器，保留物理域检查和4/8子步积分复核。',.85,5.0,11.6,1.35,21)
    s=slide('逐被试、逐处理路径看成功率与计算量','蓝：原Armijo；橙：有条件的二次插值；失败起点保留在分母中')
    s.shapes.add_picture(str(out/'figures/solver_summary.png'),Inches(.65),Inches(1.72),width=Inches(12.0))
    s=slide('收益不均匀：部分起点增加计算量','每个点是一对相同起点；低于对角线表示新方法减少计算代价')
    s.shapes.add_picture(str(out/'figures/paired_cost.png'),Inches(.65),Inches(1.78),width=Inches(12.0))
    box(s,f'{increased_count}/72起点调用增加，最大+{100*worst_increase:.1f}%。\n{slower_text}。',.75,6.02,11.8,.90,16)
    s=slide('速度改善是否改变了解？','只在两臂都成功的相同起点上比较，新增成功单独统计')
    box(s,f"共同成功起点：{len(common)}\n目标值最大相对差：{agreement['maximum_relative_objective_difference']:.3g}\nβ最大相对变化：{100*agreement['maximum_relative_gain_difference']:.4f}%\n预测曲线最大差：{agreement['maximum_prediction_difference_training_sd']:.4g} 个训练SD\n血流f最大绝对差：{agreement['maximum_flow_difference']:.4g}",.85,1.85,11.5,4.45,25)
    s=slide('适用范围与仍未解决的问题','数值优化更可靠，不自动等于模型或生理解释已经通过验收')
    box(s,'只在实际下降远小于GN预测时，额外尝试一个较短步长。\n只有原目标更低、物理域有效时才采用；不放宽成功条件。\n显式启用：step_control="quadratic_interpolation"。',.85,1.85,11.6,2.9,24)
    box(s,'本次未验证跨独立数据集、验证NRMSE改善或生理参数可辨识性。\n原Armijo默认路径与旧证据保留；困难波形仍需独立的模型诊断。',.85,5.1,11.6,1.2,21)
    deck.save(out/'SOLVER_RESULTS.pptx')


def render_teacher_robustness(run, out):
    """Editable Chinese presentation from terminal four-arm evidence only."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches,Pt
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    summary=json.loads((run/'summary.json').read_text())
    if summary['execution']!='completed':raise ValueError('report requires complete task records')
    if out.exists() and (out/'SSM_TEACHER_ROBUSTNESS.pptx').exists():raise ValueError('use a new versioned report export')
    out.mkdir(parents=True,exist_ok=True);figures=out/'figures';figures.mkdir(exist_ok=True)
    arms=['M0','M-observation','M-prior','M-combined'];labels=['M0','观测修正','驱动先验','组合']
    colors=['#70869A','#D99440','#387BA7','#198A83']
    datasets=['eeg_fnirs_single_trial','simultaneous_eeg_nirs','visual_cognitive_motivation']
    dslabels={'eeg_fnirs_single_trial':'Single-Trial','simultaneous_eeg_nirs':'Simultaneous','visual_cognitive_motivation':'Visual'}
    scenarios=['control','common_basis','common_correlated','common_colored','EEG_gain','Hb_gain','noise','lag_1s','lag_3s','driver_amplitude','driver_shape','physiology_tau','weak_coupling']
    scenario_labels=['匹配对照','基底内共同成分','与驱动相关的共同成分','基底外有色共同成分','EEG 增益 ×1.3','Hb 增益 ×1.3','白噪声 ×3','Hb 时间偏移 1 s','Hb 时间偏移 3 s','真实驱动幅度 ×1.5','真实驱动形状改变','真实 τ：2 → 4 s','真实近零耦合']
    syn=pd.read_csv(run/'synthetic_summary.csv');meas=pd.read_csv(run/'measured_summary.csv')
    raw=pd.read_csv(run/'measured_metrics.csv');paired=pd.read_csv(run/'measured_paired.csv')
    nulls=pd.read_csv(run/'null_comparisons.csv');teacher=pd.read_csv(run/'teacher_summary.csv')
    contrasts=pd.read_csv(run/'counterfactual_summary.csv');repeat=pd.read_csv(run/'repeatability_summary.csv')
    profile=pd.read_csv(run/'profile_summary.csv');sources=[]
    font=Path('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc')
    if not font.exists():raise ValueError('CJK figure font unavailable; do not export missing glyphs')
    font_manager.fontManager.addfont(str(font))
    plt.rcParams.update({'font.family':font_manager.FontProperties(fname=str(font)).get_name(),'font.size':12,'axes.spines.top':False,
        'axes.spines.right':False,'axes.titlesize':15,'figure.facecolor':'white','savefig.facecolor':'white'})
    deck=Presentation();deck.slide_width=Inches(13.333);deck.slide_height=Inches(7.5)
    def text(slide,value,x,y,w,h,size=21,color='23384D',bold=False):
        shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
        tf=shape.text_frame;tf.word_wrap=True;tf.margin_left=0;tf.margin_right=0
        for i,line in enumerate(str(value).split('\n')):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.space_after=Pt(13)
            for r in p.runs:
                r.font.name='Noto Sans CJK SC';r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=RGBColor.from_string(color)
        return shape
    def slide(title,subtitle,evidence):
        s=deck.slides.add_slide(deck.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string('F7F9FB')
        text(s,title,.55,.30,12.2,.70,30,bold=True)
        text(s,subtitle,.58,1.08,12.1,.62,16,color='5A7187')
        text(s,f'2026.10.01  /  SSM 条件语义 teacher  /  {len(deck.slides):02}',.58,7.12,11.8,.23,10,color='6F8293')
        s.notes_slide.notes_text_frame.text='事实源：'+str(run)+'\n证据文件：'+evidence+'\nNRMSE 分母为冻结训练尺度；误差统计注明成功子集，失败另列。稳健性指标不等于后验概率或唯一生理来源。'
        sources.append(dict(slide=len(deck.slides),title=title,evidence=evidence));return s
    def table_on(s,headers,rows,widths=None,y=1.95,size=17):
        n=len(rows);height=min(4.55,.52*(n+1));shape=s.shapes.add_table(n+1,len(headers),Inches(.65),Inches(y),Inches(12.0),Inches(height));table=shape.table
        if widths:
            for c,w in zip(table.columns,widths):c.width=Inches(w)
        for i,row in enumerate([headers]+rows):
            for j,value in enumerate(row):
                cell=table.cell(i,j);cell.text=str(value);cell.margin_left=Inches(.1);cell.margin_right=Inches(.07)
                cell.vertical_anchor=MSO_ANCHOR.MIDDLE;cell.fill.solid()
                cell.fill.fore_color.rgb=RGBColor.from_string('23384D' if i==0 else ('FFFFFF' if i%2 else 'EAF0F5'))
                for p in cell.text_frame.paragraphs:
                    for r in p.runs:r.font.name='Noto Sans CJK SC';r.font.size=Pt(size);r.font.bold=i==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '23384D')
        return shape
    def savefig(fig,name):
        path=figures/(name+'.png');fig.savefig(path,dpi=240,bbox_inches='tight');plt.close(fig);return path
    def picture(s,path,y=1.82,h=4.95):
        from PIL import Image
        with Image.open(path) as im:w0,h0=im.size
        width=min(12.0,h*w0/h0);height=width*h0/w0
        s.shapes.add_picture(str(path),Inches((13.333-width)/2),Inches(y),width=Inches(width),height=Inches(height))
    def fmt(x,digits=3):return f'{x:.{digits}f}' if pd.notna(x) else '—'
    def value(ds,mode,arm,metric='total_nrmse_median'):
        x=meas[(meas.dataset==ds)&(meas['mode']==mode)&(meas.pairing=='real')&(meas.arm==arm)]
        return float(x.iloc[0][metric]) if len(x) else np.nan
    def barfigure(modes,methods,name,log=False):
        fig,axes=plt.subplots(len(modes),3,figsize=(12,2.3*len(modes)),squeeze=False)
        for i,(mode,title) in enumerate(modes):
            for j,ds in enumerate(datasets):
                ax=axes[i,j];values=[value(ds,mode,a) for a in methods]
                ax.bar(np.arange(len(methods)),values,color=[colors[arms.index(a)] if a in arms else '#B9C2CA' for a in methods])
                ax.set_xticks(np.arange(len(methods)),[labels[arms.index(a)] if a in arms else {'training_template':'模板','ridge_cross':'跨模态 ridge','own_context':'插值'}.get(a,a) for a in methods],rotation=20,ha='right',fontsize=10)
                ax.set_title(dslabels[ds]+' / '+title);ax.grid(axis='y',alpha=.2);ax.set_axisbelow(True)
                if log:ax.set_yscale('log')
                if j==0:ax.set_ylabel('训练尺度 NRMSE 中位数')
        fig.tight_layout();return savefig(fig,name)
    s=slide('让共享状态更可靠：四臂实验结果','保留六状态生理核心 · 观测失配、proper 先验与条件 teacher 的完整检验','summary.json; measured_summary.csv')
    text(s,'主要问题',.75,2.0,3.8,.5,20,color='198A83',bold=True)
    text(s,'当观测改变时，\n共享状态能否保持正确，\n或明确显示不可靠？',.75,2.7,6.2,2.5,29,bold=True)
    text(s,f"{summary['synthetic']['fits']:,} 次合成拟合\n{summary['measured']['fits']:,} 次实测四臂拟合\n72 个窗口 / 18 名被试 / 3 个数据集",7.05,2.5,5.5,2.8,25)
    text(s,'结果不等同于个体生理参数恢复或 tokenizer 语义资格。',.75,6.15,11.8,.6,20,color='5A7187')
    s=slide('先看核心端点：整段 EEG 隐藏','同一窗口、同一冻结尺度；表中为收敛窗口的 NRMSE 中位数，失败见执行审计','measured_summary.csv; measured_paired.csv')
    table_on(s,['数据集','M0','观测修正','驱动先验','组合','训练模板','跨模态 ridge'],
        [[dslabels[d]]+[fmt(value(d,'EEG_hidden',a)) for a in arms+['training_template','ridge_cross']] for d in datasets],size=17)
    text(s,'减小逆推发散只解决稳定性问题；只有优于模板/自身上下文并通过错误配对检验，\n才支持额外的跨模态信息。原始重建变好不能替代这一端点。',.75,4.9,11.85,1.45,23)
    s=slide('结果：改善稳定性，尚不能放行通用语义 teacher','重建、独立补全和状态正确性给出了不同答案，必须分别保留','summary.json; measured_summary.csv; null_comparisons.csv')
    full0=sum(v['M0'] for v in summary['full_pass'].values());fullc=sum(v['M-combined'] for v in summary['full_pass'].values())
    text(s,f'完整重建\n{full0}/72 → {fullc}/72',.8,2.0,3.9,1.4,28,bold=True,color='198A83')
    text(s,'隐藏 EEG\n误差降至约 1.10',4.8,2.0,3.9,1.4,28,bold=True,color='387BA7')
    whole=nulls[(nulls.arm=='M-combined')&nulls['mode'].isin(['EEG_hidden','Hb_hidden'])]
    text(s,f'整模态配对 / 时移\n{int((whole.ci_low>0).sum())}/{len(whole)} 区间全为正',8.8,2.0,3.9,1.4,25,bold=True,color='A56932')
    text(s,'驱动先验显著限制了隐藏 EEG 时的发散，但三个数据集的中位误差仍高于训练模板。\n观测项改善完整重建，却在三个数据集都使隐藏 HbR 的中位补全误差变差。\n当前应保留条件目标与敏感性输出，限制其用途；不把观测残差吸收进“已验证生理语义”。',.8,4.25,11.8,2.3,23)
    s=slide('四臂只改变观测误差与驱动约束','六状态动力学、H0 参数、处理算子、样本身份和优化预算保持一致','resolved_config.yaml')
    table_on(s,['方案','共同 Hb 成分','proper 驱动先验','要回答的问题'],[
        ['M0','无','无','保留原逆推失败与误差'],['M-observation','4 模式、固定方向','无','减少观测失配对状态的污染？'],
        ['M-prior','无','训练冻结幅度约束','隐藏 EEG 时限制 DC / 低频发散？'],['M-combined','同观测臂','同先验臂','两者互补还是产生新补偿？']],widths=[2.1,2.8,2.8,4.3],size=17)
    text(s,'物理 H(x)、额外共同成分、完整预测分别保存。额外成分不预先归因为头皮、血容量或仪器噪声。',.75,5.7,11.8,.95,21)
    s=slide('模型表达：限制容量，保留分解歧义','额外项为观测层变量；不是第七个生理状态','source_snapshot/src/inference/shared_driver_reconstruction.py; resolved_config.yaml')
    fig,ax=plt.subplots(figsize=(12,3.4));ax.axis('off')
    boxes=[(.02,.4,.24,.42,'单一共享驱动 r(t)\n+ 五个血管初态'),(.37,.4,.25,.42,'原六状态动力学\nx = (r,s,f,v,p,q)'),(.75,.57,.23,.28,'物理预测 H(x)'),(.75,.10,.23,.28,'共同观测项\n[0, 0.65, 0.35] c(t)')]
    for x,y,w,h,content in boxes:
        ax.add_patch(plt.Rectangle((x,y),w,h,transform=ax.transAxes,fc='#E8F3F1',ec='#198A83',lw=1.5));ax.text(x+w/2,y+h/2,content,ha='center',va='center',transform=ax.transAxes,fontsize=16)
    for a,b in [((.26,.61),(.36,.61)),((.62,.61),(.74,.70))]:ax.annotate('',xy=b,xytext=a,xycoords='axes fraction',arrowprops=dict(arrowstyle='->',lw=2,color='#23384D'))
    picture(s,savefig(fig,'model'),y=1.9,h=3.5)
    text(s,'c(t)：4 个固定余弦模式，系数由可见 Hb 精确求解并带幅度惩罚。\nr(t)：曲率惩罚 + 可选 proper 幅度约束，补上常数/线性方向。',.75,5.35,11.8,1.15,22)
    s=slide('训练与评价分开，失败继续留在分母中','合成与实测分别冻结强度；组合臂复用单项选择，不在测试窗口重新选参','synthetic_calibration.json; calibration/*.json; measured_plan.json')
    table_on(s,['对象','尺度来源','强度选择','独立评价'],[
        ['合成','已知训练驱动 / 共同成分分布','独立 seed 流；5 种条件 × 2 遮挡','32 真值组 × 13 条件 × 3 遮挡'],
        ['实测驱动先验','外折训练 H0 潜驱动 RMS','被试不重叠的训练内遮挡窗口','parent 预选的 72 窗口'],
        ['实测观测项','外折训练 H0 的 Hb 残差系数','同一训练内窗口与尺度','同一 72 窗口及错误配对']],widths=[1.8,3.5,3.6,3.1],size=16)
    text(s,'训练信号尺度、观测噪声、潜驱动分布分别命名。本轮没有校准观测噪声或后验概率。\nPCA/SD 继承外折冻结坐标；内层选择和区间均条件于这些坐标。',.75,5.35,11.8,1.25,21)
    for mode,title in [('full','合成：全观测下的驱动恢复'),('EEG_hidden','合成：整段 EEG 隐藏后的驱动恢复')]:
        s=slide(title,'每格 32 组真值；收敛解 NRMSE 中位数 / 训练潜驱动尺度；未满额时另标成功数','synthetic_summary.csv; synthetic_metrics.csv')
        frame=syn[syn['mode']==mode].pivot(index='scenario',columns='arm',values='driver_nrmse_median').reindex(index=scenarios,columns=arms)
        support=syn[syn['mode']==mode].pivot(index='scenario',columns='arm',values='converged').reindex(index=scenarios,columns=arms)
        fig,ax=plt.subplots(figsize=(11,5.1));im=ax.imshow(np.log10(np.maximum(frame.to_numpy(),.01)),aspect='auto',cmap='YlOrRd',vmin=-2,vmax=1)
        ax.set_yticks(range(len(scenarios)),scenario_labels,fontsize=11);ax.set_xticks(range(4),labels,fontsize=13)
        for i in range(len(scenarios)):
            for j in range(4):
                label=fmt(frame.iloc[i,j])+(f'  ({int(support.iloc[i,j])}/32)' if support.iloc[i,j]<32 else '')
                ax.text(j,i,label,ha='center',va='center',fontsize=10,color='white' if frame.iloc[i,j]>2 else '#23384D')
        fig.colorbar(im,ax=ax,label='log10 NRMSE');fig.tight_layout();picture(s,savefig(fig,'synthetic_'+mode),h=4.95)
    s=slide('驱动的绝对水平与波形形状必须分开','首 5 s 参考只改变评分坐标；不回头修改拟合或主要端点','synthetic_summary.csv')
    rows=[]
    for scenario in ['control','common_correlated','lag_3s']:
        group=syn[(syn['mode']=='full')&(syn.scenario==scenario)].set_index('arm')
        for arm in ['M0','M-combined']:
            r=group.loc[arm];rows.append([scenario_labels[scenarios.index(scenario)],labels[arms.index(arm)],fmt(r.driver_nrmse_median),fmt(r.centered_driver_nrmse_median),fmt(r.driver_correlation_median)])
    table_on(s,['条件','方案','完整驱动误差','参考后驱动误差','形状相关'],rows,widths=[4.,1.6,2.2,2.4,1.8],size=17)
    text(s,'匹配条件下先验引入的误差主要在偏移方向；不能据此说驱动形状已经被压坏。\n可监督的候选功能应单独验证，不自动把完整潜驱动或全部六状态作为真值。',.75,6.05,11.8,.9,18)
    s=slide('真实驱动改变时，模型是否跟随？','保持观察条件不变；比较相对同 seed 对照的估计变化与真变化','counterfactual_summary.csv')
    fig,axes=plt.subplots(1,2,figsize=(12,4.3))
    for ax,scenario,title in zip(axes,['driver_amplitude','driver_shape'],['真实幅度 ×1.5','真实形状改变']):
        group=contrasts[(contrasts['mode']=='full')&(contrasts.scenario==scenario)].set_index('arm').reindex(arms)
        ax.bar(np.arange(4)-.15,group.estimated_change_SD_median,.3,color=colors,label='估计变化')
        ax.bar(np.arange(4)+.15,group.truth_change_SD_median,.3,color='#CBD4DC',label='真实变化')
        ax.set_xticks(range(4),labels);ax.set_title(title);ax.set_ylabel('变化 RMS / 训练潜驱动尺度');ax.legend();ax.grid(axis='y',alpha=.2)
    fig.tight_layout();picture(s,savefig(fig,'true_changes'))
    s=slide('合成：直接看误差归属，而不只看完整重建','物理成分与已知物理真值比较；共同成分与注入成分比较','synthetic_summary.csv')
    selected=['control','common_basis','common_correlated','common_colored','lag_3s']
    rows=[]
    for scenario in selected:
        a=syn[(syn['mode']=='full')&(syn.scenario==scenario)].set_index('arm')
        rows.append([scenario_labels[scenarios.index(scenario)]]+[fmt(a.loc[x,'physical_truth_nrmse_median']) for x in arms])
    table_on(s,['条件','M0 物理误差','观测修正','驱动先验','组合'],rows,widths=[4.,2.,2.,2.,2.],size=17)
    text(s,'此处误差的目标是已知 H(x)；实测没有这份真值，不能把加残差后的低误差报告成物理状态恢复。',.75,6.0,11.8,.75,20)
    s=slide('负结果：稳定的输出仍然可能是错误的','同真值反事实集合上的敏感性筛查；大误差定义为任一观测条件下逐点误差 >0.5 个潜驱动尺度','synthetic_reliability.csv; synthetic_metrics.csv')
    reliability=pd.read_csv(run/'synthetic_reliability.csv').groupby('arm')[['robust_fraction','high_error_fraction','false_reassurance_fraction']].mean().reindex(arms)
    rows=[]
    for arm,label in zip(arms,labels):
        row=reliability.loc[arm];conditional=row.false_reassurance_fraction/max(row.robust_fraction,1e-12)
        rows.append([label,f'{100*row.robust_fraction:.1f}%',f'{100*row.high_error_fraction:.1f}%',f'{100*row.false_reassurance_fraction:.1f}%',f'{100*conditional:.1f}%' if row.robust_fraction else '—'])
    table_on(s,['方案','稳定时间占比','大误差时间占比','稳定且错 / 全部','稳定集合内错率'],rows,size=17)
    text(s,'这是本轮不能直接放行 semantic teacher 的关键证据。\n先验可以压低扰动响应，同时保留系统性偏差；需要带真值的校准与额外独立证据。',.75,5.25,11.8,1.4,23,color='A56932')
    s=slide('实测范围与时间处理边界','与上一轮使用完全相同的 QC 预选窗口；不是全体被试的新普查','measured_plan.json; processing_timing_audit.json')
    table_on(s,['数据集','被试','窗口','观测与重复边界'],[
        ['Single-Trial','6','24','相对 Hb + EEG log-power PCA；真实 session'],
        ['Simultaneous','6','24','发布 Hb 浓度；原 task/session/block 身份'],
        ['Visual','6','24','发布 Hb 单位未核定；Part 或记录内不重叠块']],widths=[2.5,1.,1.,7.5],size=18)
    text(s,'30 s / 120 点 / 4 Hz；首 5 s 为坐标参考，不统一称为静息。\n统一窗口内离线算子；12 s 移位对照不环绕，真实配对匹配相同支持。\nREFED 缺少可信 EEG–Hb 共同几何，本轮不补造跨模态对应。',.75,4.65,11.8,1.85,21)
    s=slide('实测：完整重建与物理成分分开','四臂都对可见全窗重新拟合驱动与初态；此页是重建，不是预测','measured_summary.csv')
    fig,axes=plt.subplots(1,3,figsize=(12,4.0))
    for ax,ds in zip(axes,datasets):
        ax.bar(np.arange(4)-.17,[value(ds,'full',a) for a in arms],.34,color=colors,label='完整预测')
        ax.bar(np.arange(4)+.17,[value(ds,'full',a,'physical_nrmse_median') for a in arms],.34,color=colors,alpha=.35,label='仅物理成分')
        ax.set_xticks(range(4),labels);ax.set_title(dslabels[ds]);ax.set_ylabel('训练尺度 NRMSE 中位数');ax.legend(fontsize=10);ax.grid(axis='y',alpha=.2)
    fig.tight_layout();picture(s,savefig(fig,'physical_vs_total'))
    s=slide('训练尺度合格，并不等于窗口波形准确','完整重建三分量均 <0.5 的计数；右列展示局部尺度误差','summary.json; measured_summary.csv')
    rows=[]
    for ds in datasets:
        for arm in ['M0','M-combined']:
            rows.append([dslabels[ds],labels[arms.index(arm)],str(summary['full_pass'][ds][arm])+'/24',fmt(value(ds,'full',arm)),fmt(value(ds,'full',arm,'total_local_nrmse_median'))])
    table_on(s,['数据集','方案','全分母达标','训练 SD NRMSE','窗口 SD NRMSE'],rows,widths=[2.8,1.8,2.2,2.6,2.6],size=17)
    s=slide('中心 4 s 遮挡：先与自身上下文比较','同一隐藏区间；完整测试结果与失败状态均保留','measured_summary.csv; measured_paired.csv')
    picture(s,barfigure([('center_EEG','中心 EEG'),('center_Hb','中心 Hb')],arms+['own_context','ridge_cross'],'center_completion'),h=4.95)
    s=slide('整模态缺失：先验是否改善不稳定逆推？','纵轴采用对数尺度；训练模板与跨模态 ridge 同时展示','measured_summary.csv')
    picture(s,barfigure([('EEG_hidden','整段 EEG'),('Hb_hidden','整段 Hb')],arms+['training_template','ridge_cross'],'whole_completion',log=True),h=4.95)
    s=slide('单条 Hb 隐藏：共同项是否帮助独立补全？','HbO/HbR 单列；不能用已观测另一分量的重建收益代替隐藏目标误差','measured_summary.csv')
    picture(s,barfigure([('HbO_hidden','HbO 隐藏'),('HbR_hidden','HbR 隐藏')],arms+['training_template','ridge_cross'],'single_Hb_completion'),h=4.95)
    s=slide('与 M0 的配对比较：给出被试层区间','正值表示候选误差更低；仅共同收敛窗口，按被试聚合后 bootstrap','measured_paired.csv')
    selection=paired[(paired.control=='M0')&(paired.arm=='M-combined')&paired['mode'].isin(['EEG_hidden','Hb_hidden','center_Hb'])]
    rows=[[dslabels[r.dataset],r.mode,f'{r.common_success}/{r.denominator}',fmt(r.mean_gain),f'[{fmt(r.ci_low)}, {fmt(r.ci_high)}]'] for r in selection.itertuples()]
    table_on(s,['数据集','隐藏目标','共同成功','组合改善均值','95% 被试区间'],rows,size=15)
    for mode,title in [('EEG_hidden','隐藏 EEG 的真实配对信息'),('Hb_hidden','隐藏 Hb 的真实配对信息')]:
        s=slide(title,'正值：错误配对误差更大；区间跨 0 时不主张真实配对优势','null_comparisons.csv')
        rows=[]
        for ds in datasets:
            for arm in ['M0','M-combined']:
                parts=[]
                for null in ['wrong_subject','nonwrapping_shift_12s']:
                    group=nulls[(nulls.dataset==ds)&(nulls['mode']==mode)&(nulls.arm==arm)&(nulls['null']==null)]
                    parts.append('不可评估' if group.empty else f"{fmt(group.iloc[0].true_pair_gain)} [{fmt(group.iloc[0].ci_low)}, {fmt(group.iloc[0].ci_high)}]")
                rows.append([dslabels[ds],labels[arms.index(arm)]]+parts)
        table_on(s,['数据集','方案','错被试：均值 [95% CI]','时间移位：均值 [95% CI]'],rows,widths=[2.4,1.6,4.,4.],size=16)
    s=slide('teacher 输出：目标值及其适用范围','逐点稳健性掩码来自明确扰动集合；不是后验置信度','teacher_summary.csv; measured/*/result.npz')
    fig,axes=plt.subplots(1,3,figsize=(12,4.0))
    for ax,ds in zip(axes,datasets):
        group=teacher[teacher.dataset==ds].groupby(['subject','arm'])[['robust_fraction','missingness_robust_fraction']].mean().groupby('arm').mean().reindex(arms)
        ax.bar(np.arange(4)-.17,100*group.robust_fraction,.34,color=colors,label='小扰动稳定')
        ax.bar(np.arange(4)+.17,100*group.missingness_robust_fraction,.34,color=colors,alpha=.35,label='遮挡变化稳定')
        ax.set_xticks(range(4),labels);ax.set_ylim(0,105);ax.set_title(dslabels[ds]);ax.set_ylabel('被试等权稳定时间比例 (%)');ax.legend(fontsize=10)
    fig.tight_layout();picture(s,savefig(fig,'teacher_masks'),h=4.1)
    text(s,'小扰动稳定不能替代正确性；先验很强时也可能稳定地输出错误轨迹。应联合查看合成真值和配对对照。',.75,6.1,11.8,.8,18)
    s=slide('实例：隐藏 EEG 之后，驱动如何变化？','固定选择 Single-Trial 中 M0 隐藏 EEG 误差最大的窗口；这是诊断图，不用于选参','measured_metrics.csv; measured/*/result.npz')
    candidates=raw[(raw.dataset==datasets[0])&(raw['mode']=='EEG_hidden')&(raw.pairing=='real')&(raw.arm=='M0')]
    refid=candidates.sort_values('total_nrmse',ascending=False).iloc[0]['id']
    with np.load(run/'measured'/refid/'result.npz',allow_pickle=False) as arrays:
        fig,axes=plt.subplots(1,2,figsize=(12,4.2))
        for mode,ax in zip(['full','EEG_hidden'],axes):
            ax.plot(np.arange(120)*.25,arrays['target'][:,0]/arrays['sd'][0],color='black',ls='--',label='观测 EEG 特征')
            for arm,label,color in zip(arms,labels,colors):
                k=f'{mode}__{arm}__prediction'
                if k in arrays:ax.plot(np.arange(120)*.25,arrays[k][:,0]/arrays['sd'][0],color=color,label=label)
            ax.set_title('全观测重建' if mode=='full' else '整段 EEG 隐藏');ax.set_xlabel('时间 (s)');ax.set_ylabel('冻结 EEG 训练 SD');ax.legend(fontsize=9);ax.grid(alpha=.2)
        fig.tight_layout();picture(s,savefig(fig,'hidden_EEG_example'),h=4.5)
    text(s,refid,.7,6.57,12,.3,10,color='6F8293')
    s=slide('独立记录 / 分块：检验功能稳定性','不同记录有不同真实神经活动；此处不是同一潜驱动的重复测量','repeatability_summary.csv')
    group=repeat[(repeat.metric=='driver_RMS_SD')&repeat.arm.isin(['M0','M-combined'])]
    rows=[[dslabels[r.dataset],labels[arms.index(r.arm)],'记录/session' if r.repeat_kind=='independent_record_or_native_session' else '记录内分块',r.subjects,fmt(r.symmetric_relative_difference_median),fmt(r.icc_absolute)] for r in group.itertuples()]
    table_on(s,['数据集','方案','重复类型','配对被试','RMS 对称相对差','绝对一致性 ICC'],rows[:9],widths=[2.2,1.4,2.4,1.5,2.4,2.1],size=15)
    text(s,'重复身份分开保留；不同任务、记录状态与坐标均影响数值，不据此发布个体参数。',.75,6.67,11.8,.3,12)
    for parameter,title in [('tau','τ'),('neurovascular_gain','增益')]:
        s=slide('参数—轨迹 profile：'+title,'每个网格点重新优化驱动、初态及观测项；5% 目标增量不是置信区间','profile_summary.csv; profile_metrics.csv')
        group=profile[(profile.arm=='M-combined')&(profile.parameter==parameter)].sort_values(['dataset','subject'])
        rows=[[dslabels[r.dataset]+'/'+r.subject,f'{r.best_value:.3g}',f'{r.near_min:g}–{r.near_max:g}',fmt(r.driver_spread_SD),f'{r.converged}/{r.denominator}'] for r in group.itertuples()]
        table_on(s,['身份','最优网格值','5% 近等价范围','驱动跨度 / SD','成功 / 网格'],rows,size=16)
        text(s,'表中展示组合方案；全部 6 窗口 × 4 臂 × 11 网格值均保留，包含 1e−6 的近零增益对照。',.7,6.5,12,.45,15)
    s=slide('数值与执行审计','运行完成不等于全部拟合成功；失败分母和资源证据一并交付','summary.json; resources.json; *_manifest.json; verification.json')
    resources=json.loads((run/'resources.json').read_text())
    table_on(s,['检查','结果'],[
        ['合成正式拟合',f"{summary['synthetic']['converged']} / {summary['synthetic']['fits']} 收敛"],
        ['实测四臂拟合',f"{summary['measured']['converged']} / {summary['measured']['fits']} 收敛"],
        ['敏感性 / profile',f"{summary['measured']['sensitivity_fits']} 条敏感性记录；{summary['measured']['profile_fits']} 次 profile 拟合"],
        ['计算',f"user systemd；最多 {resources['workers']} 进程 × 1 数值线程；独立源码快照"],
        ['误差与分母','未加权原坐标误差、训练 SD、局部 SD、全部失败身份均保留'],
        ['交付','PPT 正文与表格可编辑；图为 240 dpi PNG；WPS 未检查']],widths=[3.,9.],size=17)
    s=slide('本轮结论应如何使用','按端点决定使用范围，不以重建漂亮或参数离开边界替代验证','summary.json; measured_paired.csv; null_comparisons.csv; synthetic_summary.csv')
    text(s,'1  先验可作为缺失条件下的稳定化候选；本轮不能将其提升为更准确的通用 teacher。\n2  观测项保留为受约束失配表达；完整重建收益未稳定转为独立补全收益。\n3  驱动形状与偏移分开验证；小扰动稳定不等于正确，也不等于真实跨模态耦合。',.75,1.9,11.8,2.75,24)
    text(s,'下一步优先：检验可监督的参考后驱动功能、缺失条件先验的偏差与稳定性取舍，\n并用独立时钟 / 观测校准证据定位剩余失配。\n本轮没有启动 H1/H2 扩展、tokenizer 训练或受保护数据评估。',.75,5.05,11.8,1.55,21,color='198A83')
    s=slide('证据与复现入口','所有数字来自当前完成的运行；上一轮只提供被冻结的输入身份、坐标与窗口','resolved_config.yaml; launch.json; source_snapshot/; presentation_sources.json')
    text(s,'实验合同：shared_driver_teacher_robustness_v1.yaml\n入口：evaluate_shared_driver_reconstruction.py --teacher-robustness\n结果：synthetic_summary / measured_summary / measured_paired / null_comparisons\nteacher：measured/<window>/result.npz + result.json\n运行身份：'+run.name,.75,1.95,11.8,3.8,22)
    text(s,'全部结果仍为探索性条件证据。三数据集各 6 名被试，区间条件于冻结的训练对象；\n小样本重复性、未识别来源、无独立时钟标定等限制必须与数值一起解读。',.75,6.0,11.8,.8,18)
    path=out/'SSM_TEACHER_ROBUSTNESS.pptx';deck.save(path)
    (out/'presentation_sources.json').write_text(json.dumps(dict(run=str(run),slides=sources,
        figures='PNG at 240 dpi; text and tables editable',wps_checked=False),ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(dict(presentation=str(path),slides=len(deck.slides)),ensure_ascii=False))


def render_component_attribution(run, out):
    """Chinese slide report, with editable text and bitmap scientific figures."""
    import yaml
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt
    from pptx.enum.text import MSO_ANCHOR
    from PIL import Image
    from matplotlib.patches import FancyBboxPatch
    summary=json.loads((run/'summary.json').read_text())
    if summary['execution']!='completed':raise ValueError('report requires complete attribution task records')
    if json.loads((run/'verification.json').read_text())['status']!='passed':
        raise ValueError('attribution evidence verification must pass before export')
    filename='SSM_COMPONENT_ATTRIBUTION.pptx'
    if (out/filename).exists():raise ValueError('preserve old presentation; use a versioned export')
    out.mkdir(parents=True,exist_ok=True);figures=out/'figures';figures.mkdir(exist_ok=True)
    cfg=yaml.safe_load((run/'resolved_config.yaml').read_text())
    font=Path('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc')
    if not font.exists():raise ValueError('CJK font required')
    font_manager.fontManager.addfont(str(font))
    plt.rcParams.update({'font.family':font_manager.FontProperties(fname=str(font)).get_name(),
        'font.size':11,'axes.titlesize':13,'axes.spines.top':False,'axes.spines.right':False,
        'figure.facecolor':'white','savefig.facecolor':'white','axes.unicode_minus':False})
    datasets=['eeg_fnirs_single_trial','simultaneous_eeg_nirs','visual_cognitive_motivation']
    dsname=dict(zip(datasets,['Single-Trial','Simultaneous','Visual']))
    arms=['M0','D-fixed','D-trained','D-exchange']
    colors={'M0':'#718297','D-fixed':'#C88937','D-trained':'#008F87','D-exchange':'#7771AE',
        'D-combined':'#B74D5D','spatial_ridge':'#3880AC','training_template':'#B3BBC2','oracle':'#292D37'}
    labels={'M0':'原模型','D-fixed':'固定同向','D-trained':'训练同向','D-exchange':'交换样',
        'D-combined':'共同项+先验','spatial_ridge':'空间 ridge','training_template':'训练模板','oracle':'真值参数'}
    mode_labels={'full':'完整重建','HbO_hidden':'隐藏 HbO','HbR_hidden':'隐藏 HbR','center_Hb':'中心 Hb 缺失',
        'Hb_hidden':'整段 Hb 缺失','EEG_hidden':'整段 EEG 缺失','hidden_channel':'独立目标通道'}
    scenario_labels={'control':'匹配对照','common_basis':'基底内共同项','common_correlated':'驱动相关共同项',
        'common_colored':'有色共同项','common_rho070':'共同方向 ρ=0.70','common_exchange':'交换样共同项',
        'EEG_gain':'EEG 增益','Hb_gain':'Hb 增益','noise':'噪声增强','lag_1s':'Hb 偏移 1 s','lag_3s':'Hb 偏移 3 s',
        'driver_amplitude':'真实驱动幅度','driver_shape':'真实驱动形状','physiology_tau':'真实 τ 改变',
        'physiology_kappa':'真实 κ 改变','unmodeled_dynamics':'未建模动力学','weak_coupling':'近零耦合'}
    frames={name:pd.read_csv(run/(name+'.csv')) for name in ['audit_metrics','audit_associations','component_sensitivity',
        'jacobian_overlap','confounding_profiles','measured_metrics','measured_summary','measured_paired','spatial_nulls',
        'temporal_paired','synthetic_summary','counterfactual_summary','anchor_paired','anchor_metrics',
        'prototype_summary','prototype_intervention_summary']}
    deck=Presentation();deck.slide_width=Inches(13.333);deck.slide_height=Inches(7.5);sources=[]
    def text(slide,value,x,y,w,h,size=20,color='20384D',bold=False):
        shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=shape.text_frame
        tf.word_wrap=True;tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
        for i,line in enumerate(str(value).split('\n')):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.space_after=Pt(10)
            for r in p.runs:
                r.font.name='Noto Sans CJK SC';r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=RGBColor.from_string(color)
        return shape
    def slide(title,subtitle,evidence):
        s=deck.slides.add_slide(deck.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string('F6F8FA')
        text(s,title,.58,.30,12.2,.65,29,bold=True)
        text(s,subtitle,.60,1.10,12.05,.55,16,color='5C7286')
        text(s,f'2026.10.02  /  结构性成分归属与分级语义  /  {len(deck.slides):02}',.6,7.14,11.9,.22,10,color='7890A2')
        s.notes_slide.notes_text_frame.text=f'事实源：{run}\n证据：{evidence}\n所有实测端点为公共开发诊断，区间条件于冻结训练对象。失败保留分母。来源未定，不等于噪声或已识别生理机制。'
        sources.append(dict(slide=len(deck.slides),title=title,evidence=evidence));return s
    def table_on(s,headers,rows,widths=None,y=1.9,size=17):
        from PIL import ImageFont
        height=min(4.65,.50*(len(rows)+1))
        shape=s.shapes.add_table(len(rows)+1,len(headers),Inches(.65),Inches(y),Inches(12.0),Inches(height))
        table=shape.table
        if widths:
            for column,width in zip(table.columns,widths):column.width=Inches(width)
        for i,row in enumerate([headers]+rows):
            row_size=float(size)
            available_height=height*72/(len(rows)+1)-6.
            while row_size>11:
                pil_font=ImageFont.truetype(str(font),round(row_size*4))
                lines=[]
                for j,value in enumerate(row):
                    available_width=table.columns[j].width/914400*72-.18*72
                    count=0
                    for paragraph in str(value).split('\n'):
                        line='';count+=1
                        for character in paragraph:
                            if pil_font.getlength(line+character)/4>available_width and line:
                                count+=1;line=character
                            else:line+=character
                    lines.append(count)
                if max(lines)*row_size*1.12<=available_height:break
                row_size-=.5
            for j,value in enumerate(row):
                cell=table.cell(i,j);cell.text=str(value);cell.margin_left=Inches(.11);cell.margin_right=Inches(.07)
                cell.margin_top=Inches(.025);cell.margin_bottom=Inches(.025)
                cell.vertical_anchor=MSO_ANCHOR.MIDDLE;cell.fill.solid()
                cell.fill.fore_color.rgb=RGBColor.from_string('20384D' if i==0 else ('FFFFFF' if i%2 else 'E9F0F4'))
                for p in cell.text_frame.paragraphs:
                    p.space_before=Pt(0);p.space_after=Pt(0);p.line_spacing=1.
                    for r in p.runs:r.font.name='Noto Sans CJK SC';r.font.size=Pt(row_size);r.font.bold=i==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '20384D')
        return table
    def note(s,value):text(s,value,.7,6.68,11.95,.40,13,color='526E81')
    def picture(s,fig,name,y=1.82,h=4.67):
        path=figures/(name+'.png');fig.savefig(path,dpi=240,bbox_inches='tight');plt.close(fig)
        with Image.open(path) as im:w0,h0=im.size
        width=min(12.0,h*w0/h0);height=width*h0/w0
        s.shapes.add_picture(str(path),Inches((13.333-width)/2),Inches(y),width=Inches(width),height=Inches(height))
        return path
    def fmt(value,d=3):return f'{float(value):.{d}f}' if pd.notna(value) else '—'
    def ci(row):return f"{fmt(row.mean_gain)} [{fmt(row.ci_low)}, {fmt(row.ci_high)}]"
    def matched(frame,**filters):
        for key,value in filters.items():frame=frame[frame[key]==value]
        if len(frame)!=1:raise ValueError(f'expected one summary row: {filters}, found {len(frame)}')
        return frame.iloc[0]
    def gain_chart(frame,methods,panels,group_column='mode',control='M0'):
        fig,axes=plt.subplots(len(panels),3,figsize=(12,3.7 if len(panels)==1 else 2.2*len(panels)),squeeze=False)
        for i,(panel,title) in enumerate(panels):
            for j,ds in enumerate(datasets):
                ax=axes[i,j]
                for k,arm in enumerate(methods):
                    q=frame[(frame.dataset==ds)&(frame[group_column]==panel)&(frame.arm==arm)&(frame.control==control)]
                    if q.empty:continue
                    r=q.iloc[0]
                    ax.errorbar(k,r.mean_gain,yerr=np.array([[max(0,r.mean_gain-r.ci_low)],[max(0,r.ci_high-r.mean_gain)]]),
                        fmt='o',capsize=4,color=colors.get(arm,'#3880AC'),markersize=7)
                ax.axhline(0,color='#718297',lw=1);ax.set_xticks(range(len(methods)),[labels.get(a,a) for a in methods],rotation=15)
                ax.set_title(dsname[ds]+' / '+title);ax.grid(axis='y',alpha=.2)
                if j==0:ax.set_ylabel('配对 NRMSE 收益\n正数为改善')
                if panel=='hidden_channel':
                    q=matched(frame,dataset=ds,mode=panel,arm='D-trained',control=control)
                    ax.text(.03,.96,f'训练同向：{q.mean_gain:+.4f}\n[{q.ci_low:.4f}, {q.ci_high:.4f}]',
                        transform=ax.transAxes,va='top',fontsize=10,bbox=dict(fc='white',alpha=.9,ec='none'))
        fig.tight_layout();return fig
    paired=frames['measured_paired'];main=paired[(paired.control=='D-fixed')&(paired.arm=='D-trained')&paired['mode'].isin(['HbO_hidden','HbR_hidden'])]
    spatial=paired[(paired.control=='M0')&(paired.arm=='D-trained')&(paired['mode']=='hidden_channel')]
    s=slide('结构性成分应如何归属','四组实验结果：从条件拟合走向可检验的成分性质','summary.json; measured_paired.csv')
    text(s,'先保留结构，再按独立证据赋予语义',.8,2.02,11.7,.8,34,color='008F87',bold=True)
    text(s,f"方向选择：6 个单条 Hb 端点中，{int((main.ci_low>0).sum())} 个优于原固定方向，但仍未全面优于 M0。\n独立通道：受限共同项仅 1/3 小幅获益，直接空间 ridge 则在 2/3 获益。\n真实动力学仍可能进入额外项；原型在未见机制下仍会产生错误形态语义。",.82,3.23,11.65,2.45,24)
    note(s,f"实测 {summary['measured']['subjects']} 人 / {summary['measured']['windows']} 窗；合成 {summary['synthetic']['fits']} 次拟合；原型 {summary['prototype']['models']} 个模型。")
    s=slide('四组问题，四类独立证据','模型拟合、可预测性、来源和语义资格分别评价','resolved_config.yaml; summary.json')
    table_on(s,['实验组','检验对象','最关键的证据'],[
        ['1  归属审计','拟合收益、补全损害与分量稳定性','同窗配对；HbT/HbX；扰动与局部混淆'],
        ['2  方向与空间','同容量方向、目标 Hb 通道补全','训练选方向；其他区域预测；匹配 null'],
        ['3  真实变化','真实变化的成分归属','真值、成分变化、参数 oracle'],
        ['4  锚定与原型','EOG 独立关系、token 干预响应','来源状态；独立生成种子；同容量对照']],widths=[2.0,4.8,5.2],size=18)
    note(s,'所有实验均固定预算；没有根据外折结果继续调参或扩大面板。')
    s=slide('一个物理核心，一个额外 Hb 方向','额外项可以表达结构失配；空间共同性与来源语义仍须检验','resolved_config.yaml; src/inference/shared_driver_attribution.py')
    fig,ax=plt.subplots(figsize=(12,4));ax.set(xlim=(0,12),ylim=(0,4));ax.axis('off')
    boxes=[(.1,1.35,3.0,1.5,'H0 六状态物理核心\nr、s、f、v、p、q','#DFEAF1'),(4.2,2.0,3.2,1.1,'物理 Hb 预测\nH(x)','#DFEAF1'),
        (4.2,.35,3.2,1.1,'结构性成分\nb × 4 个时间基函数','#E5ECE1'),(8.5,1.35,3.2,1.5,'完整观测\nH(x) + b c(t) + 残差','#E3F1ED')]
    for x,y,w,h,label,color in boxes:
        ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle='round,pad=.1',fc=color,ec='none'));ax.text(x+w/2,y+h/2,label,ha='center',va='center',fontsize=17)
    for a,b in [((3.2,2.2),(4.0,2.55)),((7.55,2.55),(8.3,2.2)),((7.55,.9),(8.3,1.7))]:
        ax.annotate('',xy=b,xytext=a,arrowprops=dict(arrowstyle='->',color='#607D8B',lw=2))
    picture(s,fig,'decomposition');note(s,'固定方向 [0.65,0.35]；训练同向与反向交换样均保持四维时间容量和相同 loading 范数。')
    s=slide('实测比较的身份与边界','三公开数据集沿用同一 QC 预选面板；所有标定只来自训练身份','measured_plan.json; calibration/*.json; resolved_config.yaml')
    table_on(s,['数据集','面板','幅度 / 几何','独立单元'],[
        ['Single-Trial','6 人 / 24 窗','相对 Hb；原数据几何','被试；原 session'],
        ['Simultaneous','6 人 / 24 窗','发布浓度；30 mm 源探距','被试；原任务 / block'],
        ['Visual','6 人 / 24 窗','Hb 单位未核定；模板几何','被试；Part / Probe']],widths=[2.2,2.0,4.8,3.0],size=18)
    text(s,'主误差：冻结训练 SD 的 NRMSE，越低越好。\n区间：被试块 bootstrap，条件于固定训练坐标；每数据集只有 6 名评价被试。\n隐藏任务发生在离线处理后的特征层；完整观测拟合允许使用当前窗口。',.8,4.5,11.7,1.7,20)
    s=slide('第一组：重建收益与单条 Hb 损害同窗检查','每点一个窗口；横轴正数代表完整拟合改善，纵轴正数代表补全恶化','audit_metrics.csv')
    audit=frames['audit_metrics'];fig,axes=plt.subplots(1,3,figsize=(12,4))
    for ax,ds in zip(axes,datasets):
        group=audit[(audit.dataset==ds)&audit.converged& audit.HbO_hidden_pair_converged&audit.HbR_hidden_pair_converged]
        ax.scatter(group.full_gain,group.single_Hb_harm,c=group.component_strength,cmap='viridis',s=45,edgecolors='white')
        ax.axhline(0,color='#718297',lw=1);ax.axvline(0,color='#718297',lw=1)
        ax.set_title(dsname[ds]);ax.set_xlabel('完整重建 NRMSE 收益');ax.grid(alpha=.15)
    axes[0].set_ylabel('隐藏 HbO/HbR 平均损害');fig.tight_layout();picture(s,fig,'audit_gain_harm')
    note(s,'颜色表示共同项强度；这类关联是归属线索，不能单独证明伪迹或某种生理来源。')
    s=slide('归属关联的强弱与不确定性','Spearman 相关；重采样完整被试块，保留同被试窗口依赖','audit_associations.csv')
    assoc=frames['audit_associations'];features=cfg['audit']['associations'];feature_cols=['native_rho','processed_rho','HbT_energy_fraction','quality_burden','component_strength','single_Hb_harm']
    feature_names=['原生记录同步','处理窗内同步','HbT 能量占比','记录质量负担','共同项强度','单条 Hb 损害']
    rows=[]
    for name,label in zip(feature_cols,feature_names):
        row=[label]
        for ds in datasets:
            q=matched(assoc,dataset=ds,feature=name);row.append(f'{fmt(q.spearman,2)} [{fmt(q.ci_low,2)}, {fmt(q.ci_high,2)}]')
        rows.append(row)
    table_on(s,['与完整收益的关联']+[dsname[d] for d in datasets],rows,widths=[3.,3.,3.,3.],size=17)
    note(s,'native_rho 和质量统计来自记录级通道审计；不能当作逐窗真值或可靠性概率。')
    s=slide('HbT 与 HbX：收益落在哪个诊断坐标','HbT = HbO + HbR；HbX = 0.35 HbO − 0.65 HbR','audit_metrics.csv; audit/*/result.json')
    rows=[]
    for ds in datasets:
        g=audit[(audit.dataset==ds)&audit.converged]
        rows.append([dsname[ds],fmt(g['HbT_M0_nrmse'].mean()),fmt(g['HbT_M-observation_nrmse'].mean()),
            fmt(g['HbX_M0_nrmse'].mean()),fmt(g['HbX_M-observation_nrmse'].mean())])
    table_on(s,['数据集','HbT：M0','HbT：共同项','HbX：M0','HbX：共同项'],rows,widths=[3.,2.25,2.25,2.25,2.25],size=18)
    text(s,'共同项本身在 HbX 中严格相消，但重新拟合的物理部分仍可改变 HbX。\n坐标保留 HbO/HbR 相对幅度；分母由原校准训练窗计算。\n总 Hb 不自动等于皮层血容量，差分坐标不等于真实血氧饱和度。',.8,4.4,11.7,1.8,21)
    s=slide('驱动稳定，不能替代共同项稳定','原 12 种校准、时移、先验、方向与训练 bootstrap 扰动','component_sensitivity.csv; sensitivity_summary.csv')
    sensitivity=frames['component_sensitivity'];fig,axes=plt.subplots(1,3,figsize=(12,3.7))
    metrics=[('driver_change_SD','驱动变化 / 驱动训练尺度'),('physical_change_SD','物理 Hb 变化 / Hb 训练尺度'),('component_change_SD','共同项变化 / Hb 训练尺度')]
    for ax,(metric,title) in zip(axes,metrics):
        for j,arm in enumerate(['M-observation','M-combined']):
            vals=[]
            for ds in datasets:
                g=sensitivity[(sensitivity.dataset==ds)&(sensitivity.arm==arm)&sensitivity.converged];vals.append(g[metric].quantile(.9))
            ax.bar(np.arange(3)+(j-.5)*.34,vals,.32,label='共同项' if j==0 else '共同项+先验',color=['#C88937','#B74D5D'][j])
        ax.set_xticks(range(3),['Single','Sim','Visual']);ax.set_title(title);ax.grid(axis='y',alpha=.2)
    axes[0].legend(fontsize=10);fig.tight_layout();picture(s,fig,'component_sensitivity')
    note(s,f"图为成功配对扰动的 P90；旧合成 mask 的最大 false reassurance fraction = {fmt(summary['audit']['previous_false_reassurance_max'])}，仍不能解释为正确概率。")
    s=slide('局部混淆：物理变化与共同项可相互补偿','观察 Jacobian 子空间，再沿重叠方向固定共同项并重拟合物理部分','jacobian_overlap.csv; confounding_profiles.csv')
    overlap=frames['jacobian_overlap'];profiles=frames['confounding_profiles'];rows=[]
    for arm in arms[1:]:
        g=overlap[(overlap.arm==arm)&(overlap.family=='fixed_H0')&overlap.converged]
        q=profiles[(profiles.arm==arm)&profiles.near_equivalent&(profiles.offset_SD!=0)]
        rows.append([labels[arm],fmt(g.maximum_cosine.median(),4),fmt(g.minimum_cosine.median(),4),str(len(q)),
            fmt(q.driver_change_SD.max()) if len(q) else '—',fmt(q.component_change_SD.max()) if len(q) else '—'])
    table_on(s,['方向','最大重叠余弦','最小重叠余弦','非零近等价格','驱动最大变化','共同项最大变化'],rows,
        widths=[1.8,2.0,2.0,1.9,2.15,2.15],size=16)
    text(s,'6 个预选窗口；共同项偏移 ±0.25 / ±0.50 Hb 训练尺度。\n近等价 = 完整工程目标上升不超过 5%，含系数惩罚；不是统计置信域。\n局部几何不能唯一裁定真实来源；未使用硬正交去强制分解。',.8,4.3,11.7,1.8,20)
    s=slide('第二组：训练记录选择的同向 loading','四个余弦模式和 loading 范数固定；方向不随评价窗口改变','calibration/*.json; calibration_scores/*.json')
    calibrations=[json.loads(p.read_text()) for p in sorted((run/'calibration').glob('*.json'))]
    rows=[]
    for c in calibrations:
        ds=next(d for d in datasets if c['key'].startswith(d));suffix=c['key'][len(ds)+2:]
        rows.append([dsname[ds],suffix,fmt(c['selected_rho'],2),f"{len(c['training_ids'])} / {len(c['selection_ids'])}",
            fmt(c['spatial_loading']['D-trained'],2)])
    table_on(s,['数据集','区域 / 外折','选择 ρ','训练 / 内选择窗','独立通道 loading'],rows,widths=[2.25,3.6,1.5,2.3,2.35],size=13)
    note(s,'PCA/SD 复用 parent 外折训练坐标；方向选择条件于该坐标，不宣称全流程重新嵌套验证。')
    s=slide('单条 Hb 补全：方向改变是否带来独立收益','与固定同向方案配对比较；正值为 NRMSE 下降','measured_paired.csv')
    picture(s,gain_chart(paired,['D-trained','D-exchange'],[('HbO_hidden','隐藏 HbO'),('HbR_hidden','隐藏 HbR')],control='D-fixed'),'single_Hb_direction')
    note(s,'误差条为被试块 95% bootstrap 区间；比较的是隐藏观测，而非完整拟合误差。')
    s=slide('完整拟合与缺失任务必须分列','训练同向方案相对 M0；完整重建不能替代隐藏端点','measured_paired.csv')
    rows=[]
    for mode in cfg['measured']['modes']:
        row=[mode_labels[mode]]
        for ds in datasets:row.append(ci(matched(paired,dataset=ds,mode=mode,arm='D-trained',control='M0')))
        rows.append(row)
    table_on(s,['端点']+[dsname[d] for d in datasets],rows,widths=[2.4,3.2,3.2,3.2],size=15)
    note(s,'各格为被试平均配对收益及区间；整段 Hb 隐藏时，自身条件共同项没有独立输入。')
    panel=json.loads((run/'measured_plan.json').read_text())['windows']
    s=slide('隐藏 Hb 的实际曲线：固定、训练与竞争方向','每数据集按原面板身份排序取首窗；不按拟合好坏选择示例','measured_plan.json; measured/*/result.npz')
    fig,axes=plt.subplots(2,3,figsize=(12,5.0));time_s=np.arange(120)*.25
    for j,ds in enumerate(datasets):
        ref=min([r for r in panel if r['dataset']==ds],key=lambda r:r['id']);a=np.load(run/'measured'/ref['id']/'result.npz')
        for i,(mode,component) in enumerate([('HbO_hidden',1),('HbR_hidden',2)]):
            ax=axes[i,j];ax.plot(time_s,a['target'][:,component]/a['sd'][component],color='#172D3C',lw=2,label='隐藏真值')
            for arm in arms:ax.plot(time_s,a[f'{mode}__{arm}__prediction'][:,component]/a['sd'][component],color=colors[arm],lw=1.2,label=labels[arm])
            ax.set_title(f"{dsname[ds]} / {ref['subject']} / {'HbO' if i==0 else 'HbR'}");ax.grid(alpha=.15)
            if i==1:ax.set_xlabel('时间 / s')
            if j==0:ax.set_ylabel('训练尺度坐标')
    axes[0,0].legend(fontsize=8,ncol=2);fig.tight_layout();picture(s,fig,'hidden_Hb_examples')
    note(s,'模型只读取另一条 Hb 与 EEG；目标曲线仅用于评分。幅度分母对所有方案相同。')
    s=slide('独立通道补全的输入边界','目标 HbO/HbR 整段隐藏；目标 EEG 和其他区域的观测可见','resolved_config.yaml; measured/*/result.json')
    fig,ax=plt.subplots(figsize=(12,4));ax.set(xlim=(0,12),ylim=(0,4));ax.axis('off')
    for x,y,w,h,label,color in [(.15,2.3,3.3,1.1,'其他同步区域 Hb + EEG\n分别拟合 H0 残差','#DFEAF1'),
        (.15,.5,3.3,1.1,'目标区域 EEG\nHb 对完全隐藏','#DFEAF1'),(4.4,2.3,3.1,1.1,'投影到一个共享方向\n取其他区域均值','#E5ECE1'),
        (4.4,.5,3.1,1.1,'目标物理 Hb 预测\n仅由可见 EEG 约束','#E5ECE1'),(8.5,1.4,3.1,1.3,'训练冻结空间 loading\n生成目标 Hb 预测','#E3F1ED')]:
        ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle='round,pad=.08',fc=color,ec='none'));ax.text(x+w/2,y+h/2,label,ha='center',va='center',fontsize=15)
    for a,b in [((3.6,2.85),(4.2,2.85)),((3.6,1.05),(4.2,1.05)),((7.6,2.85),(8.3,2.25)),((7.6,1.05),(8.3,1.85))]:
        ax.annotate('',xy=b,xytext=a,arrowprops=dict(arrowstyle='->',lw=2,color='#607D8B'))
    picture(s,fig,'spatial_information');note(s,'目标 Hb 不参与当前窗口系数、尺度或 loading 拟合；不同区域的 outer 坐标使用同一被试折。')
    s=slide('其他区域能否帮助预测目标 Hb','与 EEG-only M0 配对；完整目标 Hb 对未进入推断','measured_paired.csv; measured_metrics.csv')
    picture(s,gain_chart(paired,['D-fixed','D-trained','D-exchange','spatial_ridge'],[('hidden_channel','独立通道')]),'spatial_gain',h=4.3)
    note(s,'空间 ridge 是直接从其他区域 Hb 预测的训练基线；共同项改善不自动证明其来源。')
    s=slide('空间信息存在，当前共同项的利用仍不足','直接空间 ridge 与受限共同项均使用训练对象，目标 Hb 同样整段隐藏','measured_paired.csv')
    rows=[]
    for ds in datasets:
        rows.append([dsname[ds],ci(matched(paired,dataset=ds,mode='hidden_channel',arm='D-trained',control='M0')),
            ci(matched(paired,dataset=ds,mode='hidden_channel',arm='spatial_ridge',control='M0'))])
    table_on(s,['数据集','训练同向共同项 vs M0','直接空间 ridge vs M0'],rows,widths=[2.4,4.8,4.8],size=17)
    text(s,'Single-Trial 与 Simultaneous 的其他区域包含可用预测信息。\n把其他区域压成“先拟合 H0 残差，再提取一个共同方向”会损失部分信息。\n这提示优先检验空间观测约束和成分提取方式；尚不能据此命名信号来源。',.8,4.4,11.7,1.8,21)
    s=slide('空间增量还必须胜过错配与时移','真实同步输入与任务匹配错被试 / 非环绕 12 s 移位比较','spatial_nulls.csv')
    nulls=frames['spatial_nulls'];rows=[]
    for ds in datasets:
        for method in ('D-trained','spatial_ridge'):
            row=[dsname[ds],labels[method]]
            for control in ('wrong_subject','nonwrapping_shift_12s'):
                g=nulls[(nulls.dataset==ds)&(nulls.method==method)&(nulls.control==control)]
                row.append(ci(g.iloc[0])+f"\n支持 {int(g.iloc[0].common_success)}/24 窗" if len(g) else '无匹配支持')
            rows.append(row)
    table_on(s,['数据集','方案','真实 − 错配优势','真实 − 时移优势'],rows,widths=[2.3,2.,3.85,3.85],size=14)
    note(s,'优势 = null 误差 − 真实输入误差；移位双方使用同一 18 s 评分支持，未环绕补值。')
    s=slide('受限 AR(1)：只作离线后缀补全检验','前 20 s Hb 可见，后 10 s 隐藏；EEG 保持可见','temporal_paired.csv; calibration/*.json')
    temporal=frames['temporal_paired'];rows=[]
    for ds in datasets:
        row=[dsname[ds]]
        for arm in ('AR1','persistence','cosine_extrapolation'):
            row.append(ci(matched(temporal,dataset=ds,method='D-trained',arm=arm,control='zero')))
        rows.append(row)
    table_on(s,['数据集','AR(1) vs 零共同项','保持 vs 零共同项','余弦外推 vs 零共同项'],rows,widths=[2.25,3.25,3.25,3.25],size=16)
    text(s,'AR 系数仅由训练记录估计，并限制在 [0,0.995]；不增加逐点过程噪声。\n本实验没有证明原生跨窗口连续状态：上游离线滤波和局部参考仍在。\n这一步衡量受约束时间延续是否有用，而不是给 AR 状态赋予解剖名称。',.8,4.3,11.7,1.8,20)
    s=slide('第三组：17 类配对反事实','同一驱动与噪声种子配对；额外成分、物理真值和残差分别记录','synthetic_plan.json; synthetic_summary.csv; resolved_config.yaml')
    table_on(s,['变化类别','干预','定位问题'],[
        ['结构性成分','同向 / 有色 / 相关 / 方向错配 / 交换样','是否保护物理状态，是否误定方向'],
        ['观测条件','EEG/Hb 增益、噪声、1/3 s 时间偏移','神经语义是否被观测变化污染'],
        ['真实驱动','幅度 ×1.5、形状改变','稳定化是否压缩真实变化'],
        ['真实动力学','τ、κ、Hb 快慢驱动混合','变化进入物理、额外项还是残差'],
        ['负对照','匹配对照、近零耦合','没有额外成分时是否产生误归属']],widths=[2.2,5.,4.8],size=17)
    note(s,f"32 条独立生成条件块；{summary['synthetic']['converged']}/{summary['synthetic']['fits']} 次拟合收敛，oracle 仅用于 τ/κ 合成定位。")
    syn=frames['synthetic_summary']
    for mode,title in [('full','完整观测下的驱动恢复'),('EEG_hidden','EEG 缺失时的驱动恢复')]:
        s=slide(title,'数值为驱动真值 NRMSE 中位数；颜色为 log10 误差，仅便于比较数量级','synthetic_summary.csv')
        scenarios=cfg['synthetic']['scenarios'];methods=cfg['synthetic']['arms']
        values=np.array([[matched(syn,scenario=sc,mode=mode,arm=a).driver_nrmse_median for a in methods] for sc in scenarios])
        fig,axes=plt.subplots(1,2,figsize=(12.3,4.7));logged=np.log10(np.maximum(values,1e-4))
        for ax,indices in zip(axes,[list(range(9)),list(range(9,len(scenarios)))]):
            ax.imshow(logged[indices],aspect='auto',cmap='YlOrRd',vmin=logged.min(),vmax=logged.max())
            ax.set_xticks(range(len(methods)),[labels[a] for a in methods],rotation=20,ha='right',fontsize=10)
            ax.set_yticks(range(len(indices)),[scenario_labels[scenarios[i]] for i in indices],fontsize=11)
            for row,i in enumerate(indices):
                for j in range(len(methods)):
                    dark=(logged[i,j]-logged.min())/max(logged.max()-logged.min(),1e-10)>.68
                    ax.text(j,row,fmt(values[i,j],2),ha='center',va='center',fontsize=10,color='white' if dark else 'black')
        fig.tight_layout();picture(s,fig,'synthetic_'+mode)
        note(s,'完整驱动含水平误差；参考后的形状误差在底层表中单列，不能用中心化误差替代绝对误差。')
    s=slide('真实变化被共同项吸收了多少','对真实观测变化的有符号投影；负值和大于 1 的值代表补偿，不是概率','counterfactual_summary.csv')
    contrast=frames['counterfactual_summary'];scenarios=['common_basis','common_correlated','driver_amplitude','driver_shape','physiology_tau','physiology_kappa','unmodeled_dynamics']
    methods=cfg['synthetic']['arms'];metric='observation_component_projection_fraction_median'
    values=np.array([[matched(contrast,scenario=sc,mode='full',arm=a)[metric] for a in methods] for sc in scenarios])
    fig,ax=plt.subplots(figsize=(11,4.4));limit=max(1.,float(np.nanmax(abs(values))));im=ax.imshow(values,aspect='auto',cmap='RdBu_r',vmin=-limit,vmax=limit)
    ax.set_xticks(range(len(methods)),[labels[a] for a in methods]);ax.set_yticks(range(len(scenarios)),[scenario_labels[x] for x in scenarios])
    for i in range(len(scenarios)):
        for j in range(len(methods)):ax.text(j,i,fmt(values[i,j],2),ha='center',va='center',fontsize=12,
            color='white' if abs(values[i,j])>.6*limit else 'black')
    fig.colorbar(im,ax=ax,label='共同项投影份额');fig.tight_layout();picture(s,fig,'component_absorption')
    note(s,'前两行的额外成分是真值；驱动、τ/κ 与动力学变化属于物理真值，进入共同项意味着归属偏差。')
    s=slide('保留真实驱动变化：稳定化与幅度偏差','比较真实变化大小和估计变化；驱动训练尺度固定','counterfactual_summary.csv')
    rows=[]
    for mode in ('full','EEG_hidden'):
        for arm in ('M0','D-fixed','D-combined'):
            r=matched(contrast,scenario='driver_amplitude',mode=mode,arm=arm)
            rows.append([mode_labels[mode],labels[arm],fmt(r.driver_truth_change_SD_median),
                fmt(r.driver_estimated_change_SD_median),fmt(r.driver_change_error_SD_median)])
    table_on(s,['可见条件','方案','真实变化 / SD','估计变化 / SD','变化误差 / SD'],rows,widths=[2.2,2.6,2.4,2.4,2.4],size=17)
    note(s,'大幅降低缺失条件的驱动波动，不等于保留了真实变化；两者须同时报告。')
    s=slide('真值参数 oracle 帮助定位动力学失配','oracle 使用生成器 τ/κ，不在实测端开放自由参数','synthetic_summary.csv')
    rows=[]
    for scenario in ('physiology_tau','physiology_kappa'):
        for arm in ('M0','D-fixed','D-trained','oracle'):
            r=matched(syn,scenario=scenario,mode='full',arm=arm)
            rows.append([scenario_labels[scenario],labels[arm],fmt(r.driver_nrmse_median),fmt(r.physical_truth_nrmse_median),fmt(r.component_truth_nrmse_median)])
    table_on(s,['生成变化','方案','驱动误差','物理真值误差','额外成分误差'],rows,widths=[2.5,2.5,2.3,2.35,2.35],size=16)
    note(s,'参数变化与结构性成分不是同一语义；良好重建仍可能对应错误分配。')
    s=slide('第四组：可用锚点与尚未获得的来源证据','核对原始文档、统一 loader 和本次实际读到的辅助记录','anchor_plan.json; auxiliary/*/record.json; docs/DATASETS_DESCRIPTION.md')
    auxrecords=[json.loads(p.read_text()) for p in sorted((run/'auxiliary').glob('*/record.json'))]
    names=sorted({name for r in auxrecords for name in r['auxiliary_names']})
    table_on(s,['数据集','本轮可用证据','来源解释边界'],[
        ['Single-Trial','空间几何；文档描述 ECG/呼吸','ECG/呼吸未读；采集描述不代表可用锚点'],
        ['Simultaneous',f"{len(auxrecords)} 条记录；"+'/'.join(names),'EOG 眼动参考；Hb 源探距 30 mm'],
        ['Visual','模板几何与同步其他区域','没有精确源探距或已核实短距离测量']],widths=[2.2,4.3,5.5],size=17)
    text(s,'空间广泛 ≠ 浅表来源；眼动相关 ≠ 系统性血容量。\n本轮可以验证关系语义，来源仍统一标记 unresolved。',.8,4.75,11.7,1.4,23,color='008F87')
    s=slide('EOG 锚定：预测条件成分与隐藏 Hb','训练 ridge 使用 0/1/2 s EOG envelope；评价被试不参与拟合','anchor_paired.csv; anchor_metrics.csv')
    anchor=frames['anchor_paired'];rows=[]
    for metric,title in [('component_nrmse','条件共同项'),('hidden_Hb_nrmse','隐藏 Hb')]:
        for control,label in [('zero','零共同项'),('training_mean','训练均值'),('wrong_subject','任务匹配错被试')]:
            q=anchor[(anchor.metric==metric)&(anchor.control==control)]
            if len(q):rows.append([title,label,ci(q.iloc[0]),f"{int(q.iloc[0].common_success)}/24 窗"])
    table_on(s,['评价目标','比较基线','真实 EOG 配对收益 [95% CI]','可用配对 / 全面板'],rows,widths=[2.0,2.9,4.6,2.5],size=16)
    note(s,'被预测的共同项仍是模型条件估计；即使可预测，也不能直接证明 Hb 变化由眼动或浅表组织产生。')
    s=slide('有限原型：神经语义与 Hb 形态分别拥有输入','连续 typed tokens；不训练 VQ，不把全部语义混称神经驱动','src/tokenizers/typed_component_prototype.py; prototype_data.json')
    fig,ax=plt.subplots(figsize=(12,4.2));ax.set(xlim=(0,12),ylim=(0,4.2));ax.axis('off')
    boxes=[(.1,2.65,2.0,.9,'EEG 特征','#DFEAF1'),(.1,.65,2.0,.9,'HbO/HbR 特征','#DFEAF1'),
        (3.2,2.65,4.0,.9,'12 个参考后驱动片段均值\nEEG-only semantic encoder','#E3F1ED'),
        (3.2,.65,4.0,.9,'4 个结构性 HbT 投影系数\nHb-only semantic encoder','#E3F1ED'),
        (8.35,1.6,3.3,1.2,'独立 observation 编码器\n+ 停止梯度语义条件\n成对重建 decoder','#EAE7F2')]
    for x,y,w,h,label,color in boxes:
        ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle='round,pad=.08',fc=color,ec='none'));ax.text(x+w/2,y+h/2,label,ha='center',va='center',fontsize=14)
    for a,b in [((2.2,3.1),(3.,3.1)),((2.2,1.1),(3.,1.1)),((7.3,3.1),(8.2,2.6)),((7.3,1.1),(8.2,1.8))]:
        ax.annotate('',xy=b,xytext=a,arrowprops=dict(arrowstyle='->',lw=2,color='#607D8B'))
    picture(s,fig,'typed_prototype');note(s,'神经 token 不受 Hb 改动是输入结构保证；对 EEG 扰动的稳定性仍必须实测，不能拿结构保证冒充学习证据。')
    s=slide('训练与评价严格分开','已知真值合成监督，对照为同容量重建训练与训练内线性探针','prototype_data.json; prototype/*/result.json')
    p=cfg['prototype'];rows=[['训练',str(p['train_seeds']),str(p['train_seeds']*len(p['train_interventions'])),'固定生成干预；训练统计量'],
        ['选择',str(p['validation_seeds']),str(p['validation_seeds']*len(p['train_interventions'])),'选择验证目标最小的 epoch'],
        ['评价',str(p['test_seeds']),str(p['test_seeds']*len(p['test_interventions'])),'包含未见方向、动力学与时移'],
        ['优化重复','3','2 方案 × 3 种子','固定 160 epochs；CPU 单线程']]
    table_on(s,['分区','独立生成种子','窗口 / 模型','用途'],rows,widths=[2.,2.2,2.8,5.],size=18)
    text(s,'监督目标：生成器真驱动与真实结构性 HbT 的投影。\n元数据：semantic_type / source_status / coordinate_id / support_mask。\n这些目标的可学习性不等于实测 teacher、来源分离或下游任务资格。',.8,4.9,11.7,1.3,20)
    proto=frames['prototype_summary'];interventions=frames['prototype_intervention_summary']
    s=slide('原型保留的语义是否与干预方向一致','值为 3 个优化种子的平均变化 / 变化误差；语义尺度来自训练真值','prototype_intervention_summary.csv')
    rows=[]
    for scenario in ['EEG_gain','Hb_gain','noise','driver_amplitude','driver_shape','common_basis','common_correlated']:
        g=interventions[(interventions.arm=='typed_supervision')&(interventions.scenario==scenario)].mean(numeric_only=True)
        rows.append([scenario_labels[scenario],fmt(g.neural_truth_change_SD),fmt(g.neural_change_SD),
            fmt(g.morphology_truth_change_SD),fmt(g.morphology_change_SD),fmt(g.observation_change_SD)])
    table_on(s,['干预','神经真变化','神经估计变化','形态真变化','形态估计变化','观测表示变化'],rows,
        widths=[2.8,1.8,1.95,1.8,1.95,1.7],size=15)
    note(s,'observation 数值用自身训练 latent SD 归一化，与语义误差不是同一物理量；不能直接当作成分能量份额。')
    s=slide('原型泛化：已见干预与未见机制分开看','typed supervision 与同容量 reconstruction-only + 训练探针比较','prototype_summary.csv')
    scenarios=['control','EEG_gain','Hb_gain','common_basis','common_correlated','common_colored','common_exchange','physiology_tau','unmodeled_dynamics']
    rows=[]
    for scenario in scenarios:
        row=[scenario_labels[scenario]]
        for arm in ('typed_supervision','reconstruction_only'):
            g=proto[(proto.scenario==scenario)&(proto.arm==arm)]
            row.extend([fmt(g.neural_nrmse.mean()),fmt(g.morphology_nrmse.mean())])
        rows.append(row)
    table_on(s,['评价条件','监督：神经','监督：形态','重建探针：神经','重建探针：形态'],rows,
        widths=[3.,2.2,2.2,2.3,2.3],size=15)
    note(s,'交换样成分的 HbT 真值为零；本原型只声明 HbT 形态 token，不覆盖全部氧合分配语义。')
    s=slide('原型的失败面需要进入下一轮目标','已知真值能改善学习；仍不能保证对未见机制正确归属','prototype_intervention_summary.csv; prototype_summary.csv')
    rows=[]
    for scenario,title in [('EEG_gain','EEG 增益 → 神经语义'),('physiology_tau','真实 τ 改变 → HbT 形态'),('common_exchange','交换样成分 → HbT 形态')]:
        g=interventions[(interventions.arm=='typed_supervision')&(interventions.scenario==scenario)].mean(numeric_only=True)
        field='neural' if scenario=='EEG_gain' else 'morphology'
        rows.append([title,fmt(g[field+'_truth_change_SD']),fmt(g[field+'_change_SD']),
            '应不变但出现变化' if g[field+'_truth_change_SD']<1e-8 else '需检查幅度偏差'])
    table_on(s,['干预与被污染的语义','真实变化 / SD','估计变化 / SD','判读'],rows,widths=[4.6,2.15,2.15,3.1],size=17)
    text(s,'模型对“已见过的成分定义”学得更好，不代表它已识别所有竞争机制。\n下一版监督需要加入这些失败反事实，再用新的独立生成种子评价；\n本轮保留负结果，不以补训后成绩覆盖当前评价。',.8,4.4,11.7,1.8,21)
    s=slide('按证据层级使用本轮结果','形态、关系、机制一致性和来源四层分别准入','measured_paired.csv; spatial_nulls.csv; counterfactual_summary.csv; anchor_paired.csv')
    table_on(s,['层级','本轮可交付对象','仍需保留的限制'],[
        ['形态','低频同向模式、HbT/HbX 与 loading','依赖幅度合同和观察条件；未定部分继续保留'],
        ['关系','其他区域 / EOG 的独立预测与 null 结果','按数据集和端点使用；不能统一升级所有窗口'],
        ['机制一致性','参数、真实驱动与未建模响应的分配反事实','合成生成机制有条件；实测没有真分量标签'],
        ['来源','统一 unresolved 标记','没有独立短距离 / 外周生理因果锚定'],
        ['tokenizer','合成连续 typed-token 原型及干预验证','未训练实测模型；未获得通用生理语义资格']],widths=[1.8,5.3,4.9],size=16)
    s=slide('执行完整性与失败证据','计算完成、数值收敛和科学支持是三个不同判断','summary.json; verification.json; resources.json; *_manifest.json')
    resources=json.loads((run/'resources.json').read_text())
    table_on(s,['检查','结果'],[
        ['合成拟合',f"{summary['synthetic']['converged']} / {summary['synthetic']['fits']} 收敛；失败保留原身份"],
        ['实测评分记录',f"{summary['measured']['converged']} / {summary['measured']['rows']} 可用；含空间和时间比较"],
        ['敏感性 / 混淆 profile',f"{summary['audit']['sensitivity_fits']} 次扰动；{summary['confounding']['profile_fits']} 次条件重拟合"],
        ['原型训练',f"{summary['prototype']['models']} 个模型；{summary['prototype']['test_rows']} 条评价记录"],
        ['资源与运行',f"user systemd；最多 {resources['workers']} 进程；源码快照与阶段日志留存"],
        ['交付格式','中文可编辑 PPT；图为 240 dpi PNG；附 PDF 与逐页事实源']],widths=[3.2,8.8],size=18)
    s=slide('下一步由失败面决定','保留未解释信息，同时收紧每一种 token 的解释范围','summary.json; measured_paired.csv; counterfactual_summary.csv; prototype_summary.csv')
    text(s,'1  对方向选择有效的端点继续验证；对独立补全失败的成分保留条件补偿身份。\n2  对参数 / 动力学被共同项吸收的情形，先引入竞争生成机制或独立观测约束。\n3  对原型中的神经、Hb 形态和 observation 分别检查干预响应，保留来源未定标记。\n4  若要升级来源语义，优先补足短距离或外周生理记录及独立重复，避免只扩大模型。',.8,1.98,11.7,3.9,24)
    note(s,'本轮没有启动新的受保护 campaign、实测 tokenizer 训练、VQ 或外部发布。')
    s=slide('事实源与复现入口','实验结果由运行记录持有；本 PPT 是可追溯的沟通导出','resolved_config.yaml; launch.json; source_snapshot/; presentation_sources.json')
    text(s,'合同：shared_driver_attribution_v1.yaml\n入口：evaluate_shared_driver_reconstruction.py --component-attribution\n结果：归属审计、隐藏补全、反事实、EOG、原型五类表\n逐窗分解与合成真值：各阶段 result.json / result.npz\n运行身份：'+run.name,.8,1.98,11.7,3.1,23)
    text(s,'参考：Gagnon et al., NeuroImage 2011（独立短距离观测约束）\nhttps://pubmed.ncbi.nlm.nih.gov/21385616/\nLocatello et al., ICML 2019（重建本身不足以指定解耦语义）',.8,5.35,11.7,1.05,15,color='526E81')
    s.notes_slide.notes_text_frame.text+='\nhttps://research.google/pubs/challenging-common-assumptions-in-the-unsupervised-learning-of-disentangled-representations/'
    deck.save(out/filename)
    (out/'presentation_sources.json').write_text(json.dumps(dict(run=str(run),slides=sources,
        figures='PNG 240 dpi; editable body text and tables',wps_checked=False),ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(dict(presentation=str(out/filename),slides=len(deck.slides)),ensure_ascii=False))



# Response dynamics and continuous semantics: retained-evidence report export.

"""Retained-record quantitative figures for the 2026-10-02 response suite.

The callable reads metric/record CSVs and JSONs only. It never loads raw signals,
checkpoints or prediction arrays, and performs no fitting or new inference.
"""


def _quant_json_clean(value):
    import math
    from pathlib import Path
    import numpy as np

    if isinstance(value, dict):
        return {str(k): _quant_json_clean(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [_quant_json_clean(v) for v in value]
    if isinstance(value, np.ndarray):
        return _quant_json_clean(value.tolist())
    if isinstance(value, np.generic):
        return _quant_json_clean(value.item())
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value


def build_response_quantitative_figures(repo, run, out):
    """Return ``dict(figures=[...], audit=..., tables=...)`` and write PNGs/JSON.

    ``repo``, ``run`` and ``out`` are paths; relative run/out paths resolve below
    repo. All plotting imports are local so integration in another renderer does
    not import a runner or change global data roots.
    """
    import json
    from pathlib import Path

    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib import font_manager
    from matplotlib.patches import Rectangle
    import numpy as np
    import pandas as pd

    repo = Path(repo).resolve()
    run = Path(run)
    out = Path(out)
    run = (repo / run).resolve() if not run.is_absolute() else run.resolve()
    out = (repo / out).resolve() if not out.is_absolute() else out.resolve()
    figdir = out / "figures"
    figdir.mkdir(parents=True, exist_ok=True)
    sources = []
    figures = []
    tables = {}
    colors = ["#147D92", "#E69A3B", "#7560A5", "#40546A"]
    datasets = ["eeg_fnirs_single_trial", "simultaneous_eeg_nirs", "visual_cognitive_motivation"]
    ds_names = dict(zip(datasets, ["单试次 EEG-fNIRS", "同时 EEG-NIRS", "视觉认知动机"]))
    ds_short = dict(zip(datasets, ["单试次", "同时", "视觉"]))
    regions = ["prefrontal", "motor", "posterior"]
    region_names = dict(zip(regions, ["前额", "运动", "后部"]))
    arms = ["observation_only_capacity_matched", "point_restricted_generator",
            "point_randomized_generator", "distribution_randomized_generator"]
    arm_names = dict(zip(arms, ["等容量纯观测", "受限点原型", "随机点原型", "随机分布原型"]))
    arm_short = dict(zip(arms, ["纯观测", "受限点", "随机点", "分布"]))
    semantic_arms = arms[1:]
    seeds = [101, 202, 303]
    scenario_order = ["baseline", "true_a_amplitude", "true_a_shape", "feature_gain_only",
                      "Hb_gain_only", "common_Hb", "tau_change", "initial_change",
                      "non_dct_impulse", "ou_drive", "time_varying_feature_gain",
                      "rotated_spectral_mixture", "correlated_Hb_component", "input_lowpass",
                      "viscoelastic_outflow", "alternative_hrf", "independent_Hb_drive",
                      "no_coupling", "exact_eeg_alias", "exact_joint_alias"]
    scenario_names = dict(zip(scenario_order, ["基准", "真实驱动幅值", "真实驱动形状", "EEG特征增益",
                         "Hb坐标增益", "公共Hb成分", "时间常数改变", "初始态改变", "未见脉冲", "OU驱动",
                         "时变EEG增益", "旋转谱混合", "相关Hb成分", "输入低通", "粘弹出流", "替代HRF",
                         "独立Hb驱动", "无耦合", "精确EEG别名", "精确联合别名"]))
    methods = ["context_raw", "context_plus_O", "context_plus_O_plus_S",
               "context_plus_capacity_matched_observation_only"]
    method_names = ["原始上下文", "+观测码O", "+观测O+语义S", "+等容量纯观测"]

    def _quant_read_json(relative):
        path = run / relative
        sources.append(str(path.relative_to(repo)))
        return json.loads(path.read_text())

    def _quant_read_csv(relative):
        path = run / relative
        sources.append(str(path.relative_to(repo)))
        return pd.read_csv(path)

    def _quant_records(frame):
        return _quant_json_clean(frame.to_dict("records"))

    def _quant_one(frame, **matches):
        sub = frame
        for key, value in matches.items():
            sub = sub.loc[sub[key].eq(value)]
        if len(sub) != 1:
            raise ValueError(f"Expected one retained row for {matches}, got {len(sub)}")
        return sub.iloc[0]

    def _quant_new(title, nrows=1, ncols=1, **kwargs):
        fig, axs = plt.subplots(nrows, ncols, figsize=(12, 5), squeeze=False, **kwargs)
        fig.suptitle(title, fontsize=19, color="#17334D", y=.975)
        fig.subplots_adjust(left=.10, right=.97, top=.85, bottom=.18, wspace=.42, hspace=.70)
        return fig, axs

    def _quant_save(fig, key, title, caption, source, stats, note, *, bitmap_notes=True, filename=None):
        if bitmap_notes:
            fig.text(.03, .035, note, fontsize=10, color="#4C5D6F", va="bottom")
        fig.canvas.draw()
        renderer=fig.canvas.get_renderer()
        clipped=[]
        for artist in fig.findobj(matplotlib.text.Text):
            if not artist.get_visible() or not artist.get_text() or artist.get_clip_on():
                continue
            bbox=artist.get_window_extent(renderer)
            if bbox.x0 < -1 or bbox.y0 < -1 or bbox.x1 > fig.bbox.width+1 or bbox.y1 > fig.bbox.height+1:
                clipped.append(artist.get_text())
        if clipped:
            raise ValueError(f"Figure text outside bitmap bounds ({key}): {clipped}")
        path = figdir / (filename or f"{key}.png")
        axis_text=[artist for ax in fig.axes for artist in ax.findobj(matplotlib.text.Text)
                   if artist.get_visible() and artist.get_text()]
        minimum_axis_label_pt=min(artist.get_fontsize() for artist in axis_text)
        fig.savefig(path, dpi=220, facecolor="white")
        plt.close(fig)
        figures.append(dict(key=key, path=str(path), title=title, caption=caption,
                            conclusion=caption, source=source, stats=_quant_json_clean(stats),
                            width_inches=12, height_inches=5, dpi=220,
                            text_inside_bitmap_bounds=True,
                            minimum_axis_cell_label_pt=minimum_axis_label_pt,
                            bitmap_notes=bitmap_notes,
                            uncertainty_note=note))

    def _quant_heat(ax, data, row_labels, col_labels, title, vmin=None, vmax=None,
                    cmap="YlGnBu", fmt=".3f", colorbar=True, label_fontsize=10, title_pad=7):
        data = np.asarray(data, dtype=float)
        im = ax.imshow(data, aspect="auto", vmin=vmin, vmax=vmax, cmap=cmap)
        ax.set_title(title, fontsize=12, pad=title_pad)
        ax.set_yticks(range(len(row_labels)), row_labels, fontsize=label_fontsize)
        ax.set_xticks(range(len(col_labels)), col_labels, fontsize=label_fontsize)
        ax.tick_params(length=0)
        norm = im.norm
        for iy in range(data.shape[0]):
            for ix in range(data.shape[1]):
                value = data[iy, ix]
                text = format(value, fmt) if np.isfinite(value) else "缺失"
                ax.text(ix, iy, text, ha="center", va="center", fontsize=label_fontsize,
                        color="white" if np.isfinite(value) and norm(value) > .68 else "#102A43")
        if colorbar:
            bar = ax.figure.colorbar(im, ax=ax, fraction=.034, pad=.025)
            bar.ax.tick_params(labelsize=label_fontsize)
            bar.ax.yaxis.get_offset_text().set_fontsize(label_fontsize)
        return im

    def _quant_forest(ax, rows, labels, title, color=colors[0]):
        ys = np.arange(len(rows))
        for y, row in zip(ys, rows):
            ax.plot([row["ci_low"], row["ci_high"]], [y, y], color=color, lw=2)
            ax.scatter(row["gain"], y, color=color, s=48, zorder=3)
        ax.axvline(0, color="#8795A5", lw=1, ls="--")
        ax.set_yticks(ys, labels, fontsize=11)
        ax.invert_yaxis()
        ax.set_title(title, fontsize=13)
        ax.set_xlabel("对照 − 候选：正值为候选改善", fontsize=11)
        ax.grid(axis="x", alpha=.16)
        ax.margins(y=.18)

    def _quant_subject_aggregate(frame, columns, keys):
        # Exact owner aggregation: regions/windows within subject, then subjects
        # within dataset. Dataset-equal reduction is applied by the caller.
        subject = frame.groupby(keys + ["dataset", "subject"], dropna=False)[columns].mean()
        return subject.groupby(level=list(range(len(keys) + 1))).mean().reset_index()

    font = Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc")
    if font.exists():
        font_manager.fontManager.addfont(str(font))
        font_family = font_manager.FontProperties(fname=str(font)).get_name()
    else:
        font_family = "DejaVu Sans"
    style = {"font.family": font_family, "font.size": 11, "axes.labelsize": 11,
             "axes.titlesize": 13, "xtick.labelsize": 10, "ytick.labelsize": 10,
             "legend.fontsize": 10, "axes.spines.top": False, "axes.spines.right": False,
             "axes.unicode_minus": False, "figure.facecolor": "white"}

    summary = _quant_read_json("summary.json")
    verified = _quant_read_json("verification.json")
    token = _quant_read_json("tokenizer/summary.json")
    pairs = _quant_read_csv("dataset_equal_comparisons.csv")
    ds_pairs = _quant_read_csv("paired_comparisons.csv")
    measured = _quant_read_csv("measured_cells.csv")
    measured_summary = _quant_read_csv("measured_summary.csv")
    nulls = _quant_read_csv("primary_null_specificity.csv")
    cross = _quant_read_csv("cross_mechanism_fit_matrix.csv")
    synth = _quant_read_csv("synthetic_summary.csv")
    metrics = _quant_read_csv("tokenizer/scenario_variant_seed_summary.csv")
    calibs = [_quant_read_json(f"calibration/{dataset}__{region}.json")
              for dataset in datasets for region in regions]
    audit = dict(retained_summary=summary, verification_passed=verified["passed"],
                 tokenizer_stage_statuses=token["stage_statuses"],
                 tokenizer_synthetic_expected=token["synthetic_expected_records"],
                 tokenizer_synthetic_completed=token["synthetic_completed_records"],
                 tokenizer_synthetic_nonfinite=token["synthetic_nonfinite_records"],
                 public_probe_planned=token["public_probe_planned_rows"],
                 public_probe_completed=token["public_probe_completed_rows"],
                 public_probe_failed=token["public_probe_failed_rows"],
                 evaluation_subjects_by_dataset={dataset: measured.loc[measured.dataset.eq(dataset), "subject"].nunique()
                                                 for dataset in datasets},
                 paired_planned_min=int(pairs.planned_pairs.min()),
                 paired_common_success_min=int(pairs.common_success_pairs.min()),
                 paired_failed_or_missing_max=int(pairs.failed_or_missing_pairs.max()),
                 synthetic_scenario_summary_groups=len(metrics),
                 synthetic_scenario_seed_counts=sorted(metrics.optimization_seed_count.unique().tolist()),
                 synthetic_scenario_unique_base_ids=sorted(metrics.unique_evaluation_base_ids.unique().tolist()),
                 synthetic_scenario_cases_per_seed_min=int(metrics.cases_per_seed_min.min()),
                 synthetic_scenario_cases_per_seed_max=int(metrics.cases_per_seed_max.max()),
                 observation_only_semantic_head="raw S head has no semantic supervision; retained synthetic semantic_predictions use training-truth-fitted linear readouts of frozen all_code_features",
                 observation_only_semantic_readout="frozen reconstruction representation plus train-only supervised ridge for shape/logA; not the raw unsupervised S head",
                 seed_interpretation="optimization repeats on shared base IDs/subjects; no independent-seed CI",
                 historical_exposure="all 71 public subjects previously exposed; evaluation is development evidence",
                 reads="retained metric and record CSV/JSON only; no fitting/raw/checkpoint/NPZ access")
    if not verified["passed"] or summary["execution"] != "completed" or token["status"] != "complete":
        raise ValueError("Retained owning execution/verification is not complete")
    if len(measured) != summary["measured"]["recorded"] or not measured.success.all():
        raise ValueError("Measured CSV denominator or success mismatch")
    if pairs.failed_or_missing_pairs.max() != 0 or nulls.failed_or_missing_pairs.max() != 0:
        raise ValueError("Incomplete paired support requires explicit plotting adaptation")

    with plt.rc_context(style):
        # 01: paired primary response contrasts and observation controls.
        comparisons = [("gamma_selected", "gamma_attenuation", "谱混合 vs 匹配衰减"),
                       ("gamma_selected", "gain_selected", "谱混合 vs 独立增益"),
                       ("gain_selected", "gamma_attenuation", "独立增益 vs 匹配衰减"),
                       ("best_stable", "gain_selected", "稳定动力学 vs 独立增益"),
                       ("best_stable", "own_Hb_prefix_ridge", "稳定动力学 vs 自身Hb历史"),
                       ("best_stable", "spatial_Hb_ridge", "稳定动力学 vs 其他区域Hb"),
                       ("spatial_Hb_local_EEG_ridge", "spatial_Hb_ridge", "其他Hb+局部EEG vs 其他Hb")]
        selected = [_quant_one(pairs, prefix_steps=40, arm=a, control=c, metric="Hb_nrmse") for a,c,_ in comparisons]
        fig, axs = _quant_new("SSM主对照：谱混合未超过匹配衰减，稳定响应仍落后于观测基线", ncols=2)
        fig.subplots_adjust(left=.25, wspace=1.0)
        _quant_forest(axs[0,0], selected[:4], [x[2] for x in comparisons[:4]], "响应扩展对照")
        _quant_forest(axs[0,1], selected[4:], [x[2] for x in comparisons[4:]], "观测条件基线", colors[1])
        for ax, rows in zip(axs[0], [selected[:4], selected[4:]]):
            lim = ax.get_xlim(); ax.set_xlim(lim[0], lim[1] + (lim[1]-lim[0])*.18)
            for y,row in enumerate(rows):
                ax.annotate(f'{row.gain:+.4f}', (row.gain,y), xytext=(7,8), textcoords="offset points", fontsize=10)
        key="quant_01_ssm_primary_forest"
        gamma=selected[0]; stable=selected[3]
        cap=(f"40点前缀下，谱混合对匹配衰减的NRMSE改善为{gamma.gain:.5f}，保留区间[{gamma.ci_low:.5f}, {gamma.ci_high:.5f}]跨零；"
             f"稳定动力学对独立增益改善{stable.gain:.5f}，但对自身Hb历史和其他区域Hb基线的改善分别为{selected[4].gain:.5f}、{selected[5].gain:.5f}，均为负。")
        _quant_save(fig,key,"SSM主对照及观测基线",cap,["dataset_equal_comparisons.csv"],
                    [dict(label=label,**row.to_dict()) for (_,_,label),row in zip(comparisons,selected)],
                    "来源：dataset_equal_comparisons.csv；15被试、360配对区域窗；区间为固定训练/选择下的探索性被试块bootstrap。")

        # 02: dataset stratification preserves the same paired endpoint.
        comps=comparisons[:4]
        fig,axs=_quant_new("数据集分层：响应改善的方向与幅度并不一致",ncols=3)
        fig.subplots_adjust(left=.20,wspace=.70)
        data=[]
        for i,dataset in enumerate(datasets):
            rows=[_quant_one(ds_pairs,dataset=dataset,prefix_steps=40,arm=a,control=c,metric="Hb_nrmse") for a,c,_ in comps]
            _quant_forest(axs[0,i],rows,[x[2] for x in comps] if i==0 else ["谱混合/衰减","谱混合/增益","增益/衰减","稳定/增益"],ds_names[dataset],colors[i])
            data.extend([dict(label=label,**row.to_dict()) for (_,_,label),row in zip(comps,rows)])
        stable_by_dataset=[_quant_one(ds_pairs,dataset=dataset,prefix_steps=40,arm="best_stable",control="gain_selected",metric="Hb_nrmse").gain for dataset in datasets]
        _quant_save(fig,"quant_02_ssm_dataset_strata","SSM配对主对照的数据集分层",
                    "每个数据集先对同一被试的窗口与区域取均值，再对被试等权汇总；三个区域不作为独立样本。"
                    "稳定动力学对独立增益的NRMSE改善在三个数据集依次为"+
                    "、".join(f"{x:+.5f}" for x in stable_by_dataset)+"；跨数据集平均不能代替方向一致性。",
                    ["paired_comparisons.csv"],data,
                    "来源：paired_comparisons.csv；评价被试数依次6、5、4；探索性被试块bootstrap，非优化种子CI。")

        # 03: every candidate score, with selected cells outlined.
        fig,axs=_quant_new("参数选择：6/9区域选择零增益，时间扩展仅在选择集冻结",nrows=2,ncols=2)
        fig.subplots_adjust(left=.14,right=.94,top=.87,bottom=.06,wspace=.62,hspace=.36)
        param_values={"gamma":[-1,-.5,0,.5,1],"gain":[0,.5,2**-.5,(.8)**.5,1],"tau_n":[0,.5,1,2],"tau_v":[0,1,3,6]}
        calibration_table=[]
        for ax,(param,candidates) in zip(axs.flat,param_values.items()):
            mat=[]
            for c in calibs:
                ranks=c["selection_rankings"][param]
                values=[next(float(x["score"]) for x in ranks if np.isclose(x[param],v)) for v in candidates]
                mat.append(np.array(values)-min(values))
                calibration_table.append(dict(dataset=c["dataset"],region=c["region"],parameter=param,
                                              selected=c["selected"]["40"][param],candidates=ranks))
            _quant_heat(ax,mat,[ds_short[c["dataset"]]+"·"+region_names[c["region"]] for c in calibs],
                        [f"{x:.3g}" for x in candidates],f"{param}：选择NRMSE − 本行最小值",0,max(.001,float(np.max(mat))),fmt=".3f",label_fontsize=11.5,title_pad=3)
            for iy,c in enumerate(calibs):
                ix=next(j for j,v in enumerate(candidates) if np.isclose(c["selected"]["40"][param],v))
                ax.add_patch(Rectangle((ix-.47,iy-.47),.94,.94,fill=False,ec="#D56A19",lw=2))
        tables["calibration_candidates"]=calibration_table
        _quant_save(fig,"quant_03_parameter_selection","谱混合、增益与时间常数的选择热图",
                    "9个数据集×区域单元中，6个选择gain=0；其余3个选择gain=0.5。热图是各参数候选的选择集NRMSE相对本行最小值，"
                    "橙框标示冻结选择；0点前缀复用40点选择，未重新选择。",
                    ["calibration/*__*.json"],calibration_table,
                    "来源：9个calibration记录；橙框=冻结候选；分数仅来自选择集；时间常数单位s；不同参数按合同顺序选择。",
                    bitmap_notes=False,filename="quant_03_parameter_selection_readable.png")

        # 04: prefix absence/presence, dataset-equal marginal descriptions.
        display_arms=["broad_a","gamma_selected","gamma_attenuation","gain_selected","best_stable",
                      "own_Hb_prefix_ridge","spatial_Hb_ridge","spatial_Hb_local_EEG_ridge"]
        display_names=["宽带驱动","谱混合","匹配衰减","独立增益","稳定动力学","自身Hb历史","其他区域Hb","其他Hb+局部EEG"]
        desc=measured_summary.query('pairing=="real"').groupby(["prefix_steps","arm"])[["Hb_nrmse","Gaussian_CRPS_normalized"]].mean()
        fig,axs=_quant_new("前缀0 vs 40：可见Hb历史显著改变任务难度与方法排序",ncols=2)
        fig.subplots_adjust(left=.15,wspace=.40)
        prefix_stats=[]
        for ax,metric,label in zip(axs[0],["Hb_nrmse","Gaussian_CRPS_normalized"],["Hb NRMSE","归一化Gaussian CRPS"]):
            for y,arm in enumerate(display_arms):
                zero=desc.loc[(0,arm),metric]; forty=desc.loc[(40,arm),metric]
                ax.plot([zero,forty],[y,y],color="#BBC5CF",lw=2)
                ax.scatter(zero,y,s=40,color=colors[1],label="0点前缀" if y==0 else None)
                ax.scatter(forty,y,s=40,color=colors[0],label="40点前缀" if y==0 else None)
                prefix_stats.append(dict(arm=arm,metric=metric,prefix_0=zero,prefix_40=forty))
            ax.set_yticks(range(len(display_arms)),display_names,fontsize=10);ax.invert_yaxis()
            ax.set_xlabel(label+"（越低越好）");ax.grid(axis="x",alpha=.16);ax.legend(loc="best")
        _quant_save(fig,"quant_04_prefix_zero_vs_forty","有无目标Hb前缀的描述性分数",
                    "0点前缀完全排除目标Hb输入；40点前缀提供10s可见历史，评价仍在20–30s。"
                    "两种前缀使用同一冻结超参数。图为被试等权后再数据集等权的边际分数，不能当作有区间的配对增益。",
                    ["measured_summary.csv"],prefix_stats,
                    "来源：measured_summary.csv；每方法/前缀360成功区域窗、15被试；离线特征补全，非原生信号因果预测。")

        # 05: conditional EEG increment beyond other Hb/history.
        spatial=[("spatial_Hb_local_EEG_ridge","spatial_Hb_ridge","局部EEG | 其他Hb"),
                 ("spatial_Hb_filterbank_EEG_ridge","spatial_Hb_ridge","滤波EEG | 其他Hb"),
                 ("spatial_Hb_filterbank_EEG_ridge","spatial_filterbank_train_permutation","真实EEG vs 训练置换")]
        fig,axs=_quant_new("条件EEG增量：加入EEG未稳定改善其他区域Hb基线",ncols=2)
        fig.subplots_adjust(left=.23,wspace=.82)
        spatial_rows=[]
        for ax,metric,label in zip(axs[0],["Hb_nrmse","Gaussian_CRPS_normalized"],["NRMSE增益","归一化CRPS增益"]):
            rows=[];labels=[]
            for prefix in [40,0]:
                for a,c,l in spatial:
                    row=_quant_one(pairs,prefix_steps=prefix,arm=a,control=c,metric=metric)
                    rows.append(row);labels.append(f"{prefix}点·{l}");spatial_rows.append(dict(label=l,**row.to_dict()))
            _quant_forest(ax,rows,labels,label)
        local=_quant_one(pairs,prefix_steps=40,arm=spatial[0][0],control=spatial[0][1],metric="Hb_nrmse")
        _quant_save(fig,"quant_05_spatial_conditional_increment","其他区域Hb与历史之上的EEG增量",
                    f"40点前缀下，其他区域Hb加入局部EEG后的NRMSE增益为{local.gain:.5f}，区间[{local.ci_low:.5f}, {local.ci_high:.5f}]为负；"
                    "滤波EEG及训练配对置换对照同样未提供可晋级的稳定增量。",
                    ["dataset_equal_comparisons.csv"],spatial_rows,
                    "来源：dataset_equal_comparisons.csv；正值=候选改善；目标Hb只允许可见前缀，空间上下文只能读其他区域Hb。")

        # 06: complete multiplicity family, no selective null reporting.
        fig,axs=_quant_new("九项配对null：Holm校正后0/9通过特异性检验",ncols=2)
        fig.subplots_adjust(left=.25,wspace=.70)
        null_names={"wrong_training_subject":"错误训练被试","nonwrapping_shift_12s":"非环绕12s移位","other_region_EEG":"其他区域EEG"}
        null_arm_names={"gamma_selected":"谱混合","gain_selected":"独立增益","best_stable":"稳定动力学"}
        labels=[null_arm_names[x.arm]+"·"+null_names[x.null_pairing] for _,x in nulls.iterrows()]
        ys=np.arange(9)
        axs[0,0].barh(ys,nulls.dataset_equal_gain,color=[colors[0] if x>0 else colors[1] for x in nulls.dataset_equal_gain])
        axs[0,0].axvline(0,color="#8795A5",ls="--");axs[0,0].set_yticks(ys,labels);axs[0,0].invert_yaxis()
        axs[0,0].set_xlim(nulls.dataset_equal_gain.min()-.010,nulls.dataset_equal_gain.max()+.007)
        axs[0,0].set_xlabel("null − 真实配对 NRMSE（正值支持真实配对）")
        for y,val in enumerate(nulls.dataset_equal_gain):
            axs[0,0].annotate(f"{val:+.4f}",(val,y),xytext=(5 if val>=0 else -5,0),textcoords="offset points",ha="left" if val>=0 else "right",va="center",fontsize=10)
        _quant_heat(axs[0,1],nulls[["p_one_sided_sign_flip","p_Holm_nine"]].values,labels,["单侧原始p","Holm(9) p"],"完整登记null家族",0,1,fmt=".3f")
        _quant_save(fig,"quant_06_nine_nulls_holm","全部九项配对null与Holm校正",
                    f"九项检验全部具有360/360公共成功配对、15个被试，Holm校正后{int(nulls.null_specificity_pass.sum())}/9通过。"
                    f"最小原始p={nulls.p_one_sided_sign_flip.min():.5f}；检验条件于固定训练与选择对象，结论仍为探索性。",
                    ["primary_null_specificity.csv"],_quant_records(nulls.drop(columns=["subject_differences"])),
                    "来源：primary_null_specificity.csv；三臂×三null；精确被试符号翻转32768模式，家族内Holm校正。")

        # 07: proper score plus calibration of observed prediction intervals.
        proper_arms=["gamma_selected","gain_selected","best_stable","own_Hb_prefix_ridge","spatial_Hb_ridge","spatial_Hb_local_EEG_ridge"]
        proper_names=["谱混合","独立增益","稳定动力学","自身Hb历史","其他Hb","其他Hb+EEG"]
        proper=measured_summary.query('pairing=="real" and prefix_steps==40').groupby("arm")[["Gaussian_CRPS_normalized","Gaussian_NLL","coverage_50","coverage_80","coverage_95"]].mean().loc[proper_arms]
        fig,axs=_quant_new("SSM预测分布：proper score与覆盖率必须一起阅读",ncols=2)
        fig.subplots_adjust(left=.09,wspace=.33)
        axs[0,0].bar(range(6),proper.Gaussian_CRPS_normalized,color=[colors[i%4] for i in range(6)])
        axs[0,0].set_xticks(range(6),proper_names,rotation=20,ha="right");axs[0,0].set_ylabel("归一化Gaussian CRPS（越低越好）")
        for i,val in enumerate(proper.Gaussian_CRPS_normalized):axs[0,0].text(i,val+.009,f"{val:.3f}",ha="center",fontsize=10)
        for i,arm in enumerate(proper_arms):
            axs[0,1].plot([.5,.8,.95],proper.loc[arm,["coverage_50","coverage_80","coverage_95"]],marker="o",color=colors[i%4],ls="-" if i<4 else "--",label=proper_names[i])
        axs[0,1].plot([.5,.95],[.5,.95],color="#8795A5",ls=":",label="名义覆盖")
        axs[0,1].set_xticks([.5,.8,.95],["50%","80%","95%"]);axs[0,1].set_ylim(.35,1.01)
        axs[0,1].set_xlabel("名义区间");axs[0,1].set_ylabel("实际覆盖率");axs[0,1].legend(ncols=2,loc="lower right")
        _quant_save(fig,"quant_07_ssm_proper_score_coverage","SSM proper score与50/80/95%覆盖",
                    "CRPS及区间覆盖均在冻结选择集残差尺度下评价。覆盖率接近名义水平只反映处理后Hb的预测校准，"
                    "不能把残差Gaussian分布解释为生理状态后验；原坐标NLL另在数值附件保留。",
                    ["measured_summary.csv"],_quant_records(proper.reset_index()),
                    "来源：measured_summary.csv；40点前缀；被试内窗口/区域→被试等权→数据集等权；残差尺度不是生理后验。")

        # 08: cross-mechanism performance of a frozen, collapsed selection policy.
        synthetic_calibs={family:_quant_read_json(f"synthetic_calibration/{family}.json")
                          for family in ["restricted","randomized"]}
        if any(c["selected"]["40"][parameter]!=0.0 for c in synthetic_calibs.values()
               for parameter in ["tau_n","tau_v"]):
            raise ValueError("Figure 08 collapse interpretation requires both families to select zero time extensions")
        scenes=["standard","nonrest","input_lowpass","viscoelastic_outflow","common_Hb","no_coupling"]
        scene_labels=["标准响应","非静息初态","输入低通","粘弹出流","公共Hb","无耦合"]
        syn_arms=["gain_selected","input_lowpass","viscoelastic_outflow"]
        sub=cross.query('family=="randomized" and prefix_steps==40')
        fig,axs=_quant_new("冻结选择策略的跨机制表现",ncols=2)
        fig.subplots_adjust(left=.15,right=.94,wspace=.48)
        for ax,metric,label in zip(axs[0],["Hb_nrmse","response_function_RMSE"],["Hb NRMSE","响应函数绝对RMSE"]):
            mat=[[_quant_one(sub,scenario=sc,arm=arm)[metric] for arm in syn_arms] for sc in scenes]
            _quant_heat(ax,mat,scene_labels,["独立增益","输入低通","粘弹出流"],label,0,max(np.nanmax(mat),.0001),fmt=".3f" if metric=="Hb_nrmse" else ".4f")
            if metric=="response_function_RMSE":
                fig.axes[-1].yaxis.set_major_formatter(matplotlib.ticker.FormatStrFormatter("%.3f"))
        _quant_save(fig,"quant_08_ssm_synthetic_mechanisms","冻结选择策略的跨机制表现",
                    "restricted与randomized两个family均选择τn=τv=0，三个拟合臂实际为同一核；本图不提供机制识别准确率。"
                    "图中保留随机合成族、40点前缀、六种已知场景的干预评价结果；每格12个生成重复且全部成功。"
                    "Hb NRMSE与响应函数绝对RMSE描述冻结选择策略在不同生成机制下的误差，不是机制分类或非零扩展的比较。"
                    "无耦合truth响应接近零，相对响应误差分母退化，因此图用绝对误差；原始相对值保留在附件。",
                    ["cross_mechanism_fit_matrix.csv","synthetic_calibration/restricted.json","synthetic_calibration/randomized.json"],_quant_records(sub),
                    "来源：cross_mechanism_fit_matrix.csv、synthetic_calibration/{restricted,randomized}.json。\n"
                    "两生成族均选择τn=τv=0，三个拟合臂实际为同一核；随机族40点前缀，每格12/12；不提供机制识别准确率。")

        # 09: losses are arm-specific objectives, not cross-arm semantic scores.
        fig,axs=_quant_new("连续原型训练：四臂×三优化种子，选择终点在各臂内部使用",nrows=2,ncols=2)
        fig.subplots_adjust(left=.08,right=.96,top=.83,bottom=.16,wspace=.28,hspace=.75)
        fit_stats=[]
        for ax,arm in zip(axs.flat,arms):
            for i,seed in enumerate(seeds):
                fit=next(x for x in token["fits"] if x["arm"]==arm and x["seed"]==seed)
                hist=pd.DataFrame(fit["history"])
                ax.plot(hist.epoch,hist.selection_loss,color=colors[i],lw=1.5,label=f"选择·{seed}")
                ax.plot(hist.epoch,hist.train_loss,color=colors[i],lw=.8,ls=":",alpha=.65)
                best=hist.loc[hist.epoch.eq(fit["best_epoch"])].iloc[0]
                ax.scatter(best.epoch,best.selection_loss,color=colors[i],s=22,zorder=3)
                fit_stats.append({k:fit[k] for k in ["arm","seed","completed_epochs","best_epoch","best_selection_loss","parameter_count","train_base_ids","selection_base_ids"]})
            ax.set_title(arm_names[arm],fontsize=12);ax.set_xlabel("epoch");ax.set_ylabel("该臂目标loss")
            ax.grid(alpha=.15);ax.legend(ncols=3,loc="best",fontsize=10)
        _quant_save(fig,"quant_09_tokenizer_training_curves","四臂三种子的训练与选择曲线",
                    "12次拟合全部完成，每模型458168参数，512个训练base与128个选择base。实线为选择loss，虚线为训练loss，点为选定epoch。"
                    "不同臂的损失定义不同，图用于检查训练和选择过程，不据此比较跨臂的语义优越性。",
                    ["tokenizer/summary.json:fits[].history"],fit_stats,
                    "来源：tokenizer/summary.json；101/202/303为优化重复；四臂容量相同；训练上限160 epochs，patience=20。")

        # 10/11: all twenty scenarios, split ten per panel at a common scale.
        heat_stats=[]
        max_shape=float(metrics.loc[metrics.arm.isin(semantic_arms),"shape_rmse_seed_mean"].max())
        for number,variant in [(10,"control"),(11,"intervention")]:
            variant_name="控制" if variant=="control" else "干预"
            fig,axs=_quant_new(f"合成{variant_name}形状恢复：20场景、两个模态、三个监督原型",nrows=2,ncols=2)
            fig.subplots_adjust(left=.15,right=.94,top=.87,bottom=.05,wspace=.65,hspace=.30)
            local_stats=[]
            for row,modality in enumerate(["eeg","hb"]):
                for col,scenes in enumerate([scenario_order[:10],scenario_order[10:]]):
                    mat=[]
                    for sc in scenes:
                        rows=[_quant_one(metrics,arm=arm,modality=modality,scenario=sc,variant=variant) for arm in semantic_arms]
                        mat.append([x.shape_rmse_seed_mean for x in rows]);local_stats.extend([x.to_dict() for x in rows])
                    _quant_heat(axs[row,col],mat,[scenario_names[x] for x in scenes],["受限点","随机点","随机分布"],
                                f'{modality.upper()} · {variant_name} · 场景{1 if col==0 else 11}–{10 if col==0 else 20}',0,max_shape,fmt=".2f",label_fontsize=11.5,title_pad=3)
            cap=("每格是同一256个base身份上三个优化种子的形状RMSE均值；两个变体分开汇总、全图与配套图共用色标。"
                 "图聚焦三个端到端语义监督原型；纯观测臂另以冻结重建表征+训练真值拟合线性readout评价，见四臂图，并非其原始未监督S头。")
            scene_for_claim="baseline" if variant=="control" else "non_dct_impulse"
            vals=[_quant_one(metrics,arm=arms[3],modality=m,scenario=scene_for_claim,variant=variant).shape_rmse_seed_mean for m in ["eeg","hb"]]
            cap += ("随机分布原型的基准EEG/Hb形状RMSE为" if variant=="control" else "随机分布原型的未见脉冲干预EEG/Hb形状RMSE升至")+f"{vals[0]:.4f}/{vals[1]:.4f}。"
            _quant_save(fig,f"quant_{number:02d}_tokenizer_{variant}_heatmaps",f"连续原型20场景{variant_name}形状恢复",cap,
                        ["tokenizer/scenario_variant_seed_summary.csv"],local_stats,
                        "来源：scenario_variant_seed_summary.csv；每格256 base×3优化重复；均值不附独立种子CI；控制/干预图色标相同。",
                        bitmap_notes=False,filename=f"quant_{number:02d}_tokenizer_{variant}_heatmaps_readable.png")
            heat_stats.extend(local_stats)
        tables["all_synthetic_scenario_variants"]=_quant_records(metrics)

        # 12: distribution proper scores against the matched randomized point arm.
        fig,axs=_quant_new("随机分布 vs 随机点：均值恢复与proper score的收益并不等价",nrows=2,ncols=2)
        fig.subplots_adjust(left=.10,right=.96,top=.84,bottom=.17,wspace=.30,hspace=.70)
        dist_comp=[]
        for row,modality in enumerate(["eeg","hb"]):
            for col,metric in enumerate(["shape_rmse_seed_mean","log_amplitude_crps_seed_mean"]):
                ax=axs[row,col];xs=[];ys=[]
                for sc in scenario_order:
                    point=_quant_one(metrics,arm=arms[2],modality=modality,scenario=sc,variant="intervention")
                    dist=_quant_one(metrics,arm=arms[3],modality=modality,scenario=sc,variant="intervention")
                    x=point[metric];y=dist[metric];xs.append(x);ys.append(y)
                    dist_comp.append(dict(modality=modality,scenario=sc,metric=metric,point=x,distribution=y,distribution_minus_point=y-x))
                    ax.scatter(x,y,color=colors[0] if y<x else colors[1],s=35)
                    if sc in ["baseline","non_dct_impulse"] or (col==1 and sc=="exact_joint_alias"):
                        near_right=sc!="baseline"
                        ax.annotate(scenario_names[sc],(x,y),xytext=(-5 if near_right else 5,5),
                                    textcoords="offset points",ha="right" if near_right else "left",fontsize=10)
                low=min(xs+ys)*.85;high=max(xs+ys)*1.12
                ax.plot([low,high],[low,high],ls="--",color="#8795A5");ax.set_xlim(low,high);ax.set_ylim(low,high)
                ax.set_xlabel("随机点原型");ax.set_ylabel("随机分布原型")
                ax.set_title(modality.upper()+" · "+("形状RMSE" if col==0 else "log幅值CRPS"),fontsize=12)
                ax.grid(alpha=.15)
        _quant_save(fig,"quant_12_distribution_vs_random_point","随机分布与同生成族随机点的proper-score比较",
                    "每点为一个干预场景的种子均值；虚线下方表示分布原型分数较低。形状RMSE检验点预测，log幅值CRPS同时惩罚偏差与预测宽度。"
                    "该对照保持随机生成族与网络容量相同，只检验分布目标的作用；所有数值及差值保存在附件。",
                    ["tokenizer/scenario_variant_seed_summary.csv"],dist_comp,
                    "来源：scenario_variant_seed_summary.csv；20干预场景；3优化重复同256 base；颜色仅标示描述性分数方向。")

        # 13: alias coverage and width, control/intervention retained separately.
        alias_groups=[("baseline","control"),("exact_eeg_alias","control"),("exact_eeg_alias","intervention"),
                      ("exact_joint_alias","control"),("exact_joint_alias","intervention")]
        alias_labels=["基准\n控制","EEG别名\n控制","EEG别名\n干预","联合别名\n控制","联合别名\n干预"]
        fig,axs=_quant_new("别名压力测试：区间很窄，联合别名干预幅值95%覆盖崩溃",nrows=2,ncols=2)
        fig.subplots_adjust(left=.08,right=.96,top=.84,bottom=.25,wspace=.30,hspace=.75)
        alias_stats=[]
        for ax,metric,label,iscoverage in zip(axs.flat,["shape_coverage_95","log_amplitude_coverage_95","shape_width_95","log_amplitude_width_95"],
                                              ["形状95%覆盖","log幅值95%覆盖","形状95%区间宽","log幅值95%区间宽"],[True,True,False,False]):
            for i,modality in enumerate(["eeg","hb"]):
                vals=[]
                for sc,variant in alias_groups:
                    x=_quant_one(metrics,arm=arms[3],modality=modality,scenario=sc,variant=variant)
                    vals.append(x[metric+"_seed_mean"])
                    alias_stats.append(dict(modality=modality,scenario=sc,variant=variant,metric=metric,
                                            mean=x[metric+"_seed_mean"],minimum=x[metric+"_seed_min"],maximum=x[metric+"_seed_max"]))
                ax.plot(range(5),vals,marker="o",color=colors[i],label=modality.upper())
            ax.set_xticks(range(5),alias_labels);ax.set_ylabel(label)
            if iscoverage:ax.axhline(.95,color="#8795A5",ls="--");ax.set_ylim(0,1.05)
            ax.grid(axis="y",alpha=.15)
        handles,labels=axs[0,0].get_legend_handles_labels()
        fig.legend(handles,labels,ncols=2,loc="lower center",bbox_to_anchor=(.5,.09),frameon=False)
        covs=[_quant_one(metrics,arm=arms[3],modality=m,scenario="exact_joint_alias",variant="intervention").log_amplitude_coverage_95_seed_mean for m in ["eeg","hb"]]
        _quant_save(fig,"quant_13_alias_coverage_width","精确别名的控制/干预覆盖与区间宽度",
                    f"分布原型在精确联合别名干预下，EEG/Hb的log幅值95%覆盖仅{covs[0]*100:.2f}%/{covs[1]*100:.2f}%。"
                    "控制与干预的可见输入相同，语义幅值truth却不同；窄区间不能消除此不可识别性。图保留两变体及形状、幅值宽度。",
                    ["tokenizer/scenario_variant_seed_summary.csv"],alias_stats,
                    "来源：scenario_variant_seed_summary.csv；仅随机分布原型；95%温度校准边际区间；优化种子均值，范围见附件。")

        # 14: retained fraction is selective evaluation, not a new full denominator.
        risk_parts=[];paired_parts=[];public_parts=[]
        for arm in arms:
            for seed in seeds:
                pp=_quant_read_csv(f"tokenizer/evaluation/{arm}/seed_{seed}/paired_changes.csv")
                pp["arm"]=arm;pp["optimization_seed"]=seed;paired_parts.append(pp)
                pub=_quant_read_csv(f"tokenizer/public_probes/{arm}/seed_{seed}/records.csv")
                pub["arm"]=arm;pub["optimization_seed"]=seed;public_parts.append(pub)
                if arm==arms[3]:
                    risk=_quant_read_csv(f"tokenizer/evaluation/{arm}/seed_{seed}/selective_risk.csv")
                    risk["optimization_seed"]=seed;risk_parts.append(risk)
        risk=pd.concat(risk_parts,ignore_index=True)
        paired=pd.concat(paired_parts,ignore_index=True)
        public=pd.concat(public_parts,ignore_index=True)
        if len(public)!=token["public_probe_completed_rows"] or not public.success.all():
            raise ValueError("Retained public probe CSV denominator mismatch")
        if not paired.complete_pair.all():
            raise ValueError("Incomplete synthetic counterfactual pair")
        risk_scenes=["baseline","non_dct_impulse","ou_drive","rotated_spectral_mixture","exact_joint_alias","no_coupling"]
        risk_summary=risk.groupby(["modality","scenario","requested_retained_fraction"])[["retained_fraction","shape_rmse","retained_records","total_records"]].agg(["mean","min","max"])
        fig,axs=_quant_new("OOD风险—保留曲线：按不确定性筛选不能修复生成族偏移",ncols=2)
        fig.subplots_adjust(left=.09,right=.97,top=.84,bottom=.30,wspace=.30)
        risk_colors=["#147D92","#D56A19","#7560A5","#62A27D","#A64D64","#40546A"]
        risk_stats=[]
        for ax,modality in zip(axs[0],["eeg","hb"]):
            for i,sc in enumerate(risk_scenes):
                sub=risk_summary.loc[(modality,sc)]
                ax.plot(sub["retained_fraction"]["mean"],sub["shape_rmse"]["mean"],marker="o",ms=3,color=risk_colors[i],label=scenario_names[sc])
                for fraction,row in sub.iterrows():
                    risk_stats.append(dict(modality=modality,scenario=sc,requested_retained_fraction=fraction,
                                           retained_fraction=row[("retained_fraction","mean")],shape_rmse_mean=row[("shape_rmse","mean")],
                                           shape_rmse_min=row[("shape_rmse","min")],shape_rmse_max=row[("shape_rmse","max")],
                                           retained_records_mean=row[("retained_records","mean")],total_records=row[("total_records","mean")]))
            ax.set_title(modality.upper());ax.set_xlabel("保留评价记录比例");ax.set_ylabel("保留样本形状RMSE")
            ax.grid(alpha=.15);ax.set_xlim(.08,1.02)
        handles,labels=axs[0,0].get_legend_handles_labels();fig.legend(handles,labels,ncols=6,loc="lower center",bbox_to_anchor=(.5,.09),frameon=False)
        _quant_save(fig,"quant_14_ood_risk_retention","OOD场景的选择性风险与保留比例",
                    "每条曲线把同场景的控制与干预记录合并，按冻结不确定性排序保留一部分记录，再计算形状RMSE；三个优化重复取均值。"
                    "筛除的记录仍计入完整评价分母。未见脉冲、别名等场景不能仅靠提高拒绝比例获得已证明的语义有效性。",
                    ["tokenizer/evaluation/distribution_randomized_generator/seed_*/selective_risk.csv"],risk_stats,
                    "来源：selective_risk.csv；仅分布原型；每场景/模态512控制+干预记录、256 base；曲线均值非独立种子CI。")

        # 15: synthetic paired true response and nuisance false response.
        paired_means=paired.groupby(["arm","optimization_seed","modality","scenario"])[["true_change_rms","predicted_change_rms","change_rmse","false_change_rms"]].mean().reset_index()
        paired_seeds=paired_means.groupby(["arm","modality","scenario"])[["true_change_rms","predicted_change_rms","change_rmse","false_change_rms"]].mean().reset_index()
        fig,axs=_quant_new("反事实检验：真实驱动变化有欠响应，观测扰动产生假语义变化",nrows=2,ncols=2)
        fig.subplots_adjust(left=.08,right=.96,top=.84,bottom=.25,wspace=.33,hspace=.75)
        true_scenes=["true_a_amplitude","true_a_shape","non_dct_impulse","exact_joint_alias"]
        false_scenes=["feature_gain_only","Hb_gain_only","common_Hb","rotated_spectral_mixture","no_coupling"]
        counter_stats=[]
        for row,modality in enumerate(["eeg","hb"]):
            for col,scenes in enumerate([true_scenes,false_scenes]):
                ax=axs[row,col];x=np.arange(len(scenes));width=.22
                for i,arm in enumerate(semantic_arms):
                    subrows=[_quant_one(paired_seeds,arm=arm,modality=modality,scenario=sc) for sc in scenes]
                    vals=[z.predicted_change_rms if col==0 else z.false_change_rms for z in subrows]
                    ax.bar(x+(i-1)*width,vals,width,color=colors[i],label=arm_short[arm])
                    counter_stats.extend([z.to_dict() for z in subrows])
                if col==0:
                    truth=[_quant_one(paired_seeds,arm=arms[3],modality=modality,scenario=sc).true_change_rms for sc in scenes]
                    ax.plot(x,truth,"k_",ms=18,mew=2,label="真实变化")
                ax.set_xticks(x,[scenario_names[sc] for sc in scenes],fontsize=10)
                ax.set_title(modality.upper()+(" · 真实驱动反事实" if col==0 else " · truth=0的观测扰动"),fontsize=12)
                ax.set_ylabel("驱动变化RMS");ax.grid(axis="y",alpha=.15)
                ax.margins(y=.15)
        handles,labels=axs[0,0].get_legend_handles_labels()
        fig.legend(handles,labels,ncols=4,loc="lower center",bbox_to_anchor=(.5,.09),frameon=False)
        alias_delta=_quant_one(paired_seeds,arm=arms[3],modality="eeg",scenario="exact_joint_alias")
        _quant_save(fig,"quant_15_counterfactual_true_false_changes","真实与假语义变化的反事实响应",
                    f"联合别名的真实驱动变化RMS为{alias_delta.true_change_rms:.5f}，分布原型预测变化为{alias_delta.predicted_change_rms:.5f}。"
                    "对truth不变的坐标增益、公共Hb、谱旋转等扰动，非零预测变化为假语义变化；图只比较三种端到端语义监督的原型。",
                    ["tokenizer/evaluation/*/seed_*/paired_changes.csv"],counter_stats,
                    "来源：paired_changes.csv；同base控制−干预配对，每臂/模态/场景256对×3优化重复；均值非独立种子CI。")
        tables["counterfactual_all_scenarios"]=_quant_records(paired_seeds)

        # Frozen public probes: retain primary methods, all rows are complete.
        public_ds=_quant_subject_aggregate(public,["NRMSE","CRPS_training_SD"],["arm","optimization_seed","prefix_steps","method"])
        public_equal=public_ds.groupby(["arm","optimization_seed","prefix_steps","method"])[["NRMSE","CRPS_training_SD"]].mean().reset_index()
        tables["public_probe_dataset_subject_equal"]=_quant_records(public_ds)
        tables["public_probe_dataset_equal"]=_quant_records(public_equal)
        primary=public_equal.loc[public_equal.arm.eq(arms[3]) & public_equal.method.isin(methods)]
        # Verify this derived retained-record aggregation against its owning record.
        for p in token["public_probes"]:
            for z in p["summaries"]:
                own=_quant_one(public_equal,arm=p["arm"],optimization_seed=p["seed"],prefix_steps=z["prefix_steps"],method=z["method"])
                for col in ["NRMSE","CRPS_training_SD"]:
                    if not np.isclose(own[col],z[col],rtol=1e-10,atol=1e-12):
                        raise ValueError(f"Public aggregation differs from owning record: {p['arm']} {col}")
        audit["public_probe_aggregation_matches_owning_records"]=True

        # 16: seed lines, not a statistical confidence interval.
        fig,axs=_quant_new("公开冻结探针：40点前缀，O有收益，S没有进一步增量",ncols=2)
        fig.subplots_adjust(left=.08,right=.97,top=.84,bottom=.22,wspace=.25)
        public_primary_stats=[]
        for ax,col,label in zip(axs[0],["NRMSE","CRPS_training_SD"],["Hb NRMSE","训练SD归一化CRPS"]):
            for i,seed in enumerate(seeds):
                vals=[_quant_one(primary,optimization_seed=seed,prefix_steps=40,method=m)[col] for m in methods]
                ax.plot(range(4),vals,marker="o",lw=1.2,color=colors[i],alpha=.75,label=f"优化种子{seed}")
            mean=primary.query("prefix_steps==40").groupby("method")[col].mean().loc[methods]
            ax.plot(range(4),mean,color="#17334D",lw=2.5,marker="D",ms=5,label="优化重复均值")
            for x,val in enumerate(mean):ax.annotate(f"{val:.4f}",(x,val),xytext=(0,10),textcoords="offset points",ha="center",fontsize=10)
            ax.set_xticks(range(4),method_names,rotation=12,ha="right");ax.set_ylabel(label+"（低为好）")
            ax.grid(axis="y",alpha=.16);ax.legend(loc="best",fontsize=10)
        mean40=primary.query("prefix_steps==40").groupby("method")[["NRMSE","CRPS_training_SD"]].mean().loc[methods]
        _quant_save(fig,"quant_16_public_probe_seed_lines","公开探针四方法的NRMSE与CRPS优化种子线",
                    f"40点前缀NRMSE种子均值依次为{mean40.NRMSE.iloc[0]:.4f}、{mean40.NRMSE.iloc[1]:.4f}、{mean40.NRMSE.iloc[2]:.4f}、{mean40.NRMSE.iloc[3]:.4f}。"
                    "O改善上下文基线；加入S后NRMSE与CRPS均未进一步改善，等容量纯观测表征更好。三种子共享同15评价被试。",
                    ["tokenizer/public_probes/distribution_randomized_generator/seed_*/records.csv","tokenizer/summary.json:public_probes"],
                    _quant_records(primary.query("prefix_steps==40")),
                    "来源：public probe records.csv/record.json；冻结合成编码器；公开train拟合、selection选ridge/残差SD；种子为优化重复。")

        # 17: each dataset/prefix, common score definitions and four methods.
        fig,axs=_quant_new("公开探针分层：零前缀收益有限，40点收益主要来自观测表征",ncols=2)
        fig.subplots_adjust(left=.15,right=.96,top=.84,bottom=.20,wspace=.45)
        matrix_stats=[]
        for ax,col,label in zip(axs[0],["NRMSE","CRPS_training_SD"],["Hb NRMSE","训练SD归一化CRPS"]):
            mat=[];rowlabels=[]
            for prefix in [0,40]:
                for dataset in datasets:
                    sub=public_ds.loc[public_ds.arm.eq(arms[3]) & public_ds.prefix_steps.eq(prefix) & public_ds.dataset.eq(dataset)]
                    means=sub.groupby("method")[col].mean().loc[methods]
                    mat.append(means.values);rowlabels.append(f"{prefix}点·{ds_short[dataset]}")
                    for method,val in means.items():matrix_stats.append(dict(prefix_steps=prefix,dataset=dataset,metric=col,method=method,optimization_seed_mean=val))
            _quant_heat(ax,mat,rowlabels,["原始","+O","+O+S","纯观测"],label,min(np.ravel(mat)),max(np.ravel(mat)),fmt=".3f")
        _quant_save(fig,"quant_17_public_probe_dataset_prefix","公开四方法的数据集与前缀分层",
                    "每行固定同一数据集与目标Hb前缀；每格为先被试等权后取三优化种子均值。"
                    "0点前缀完全排除目标Hb tokens，40点前缀保留可见目标Hb；表征收益依赖前缀与数据集，不能把总平均推广到所有数据边界。",
                    ["tokenizer/public_probes/distribution_randomized_generator/seed_*/records.csv"],matrix_stats,
                    "来源：public probes records.csv；原始/+O/+O+S/等容量纯观测共用上下文；15开发评价被试，无独立确认。")

        # 18: conditional source increment across all three supervised arms.
        fig,axs=_quant_new("冻结S的条件增量：三个监督原型×两种前缀×三个优化种子",ncols=2)
        fig.subplots_adjust(left=.16,right=.97,top=.84,bottom=.20,wspace=.42)
        increments=[]
        inc_labels=[]
        for arm in semantic_arms:
            for prefix in [40,0]:inc_labels.append(arm_short[arm]+f"·{prefix}点")
        for ax,col,label in zip(axs[0],["NRMSE","CRPS_training_SD"],["Δ NRMSE：O+S − O","Δ CRPS：O+S − O"]):
            y=0
            for arm in semantic_arms:
                for prefix in [40,0]:
                    vals=[]
                    for i,seed in enumerate(seeds):
                        before=_quant_one(public_equal,arm=arm,optimization_seed=seed,prefix_steps=prefix,method="context_plus_O")[col]
                        after=_quant_one(public_equal,arm=arm,optimization_seed=seed,prefix_steps=prefix,method="context_plus_O_plus_S")[col]
                        delta=after-before;vals.append(delta)
                        ax.scatter(delta,y+(i-1)*.13,color=colors[i],s=36,label=str(seed) if y==0 else None)
                        increments.append(dict(arm=arm,prefix_steps=prefix,optimization_seed=seed,metric=col,O=before,O_plus_S=after,delta=delta))
                    ax.plot([min(vals),max(vals)],[y,y],color="#A5B1BC",lw=2)
                    ax.scatter(np.mean(vals),y,marker="D",color="#17334D",s=30,zorder=3)
                    y+=1
            ax.axvline(0,color="#8795A5",ls="--");ax.set_yticks(range(6),inc_labels);ax.invert_yaxis()
            ax.set_xlabel(label+"（负值为S增量）");ax.grid(axis="x",alpha=.16);ax.legend(title="优化种子",loc="best",ncols=3)
        dist40=[x for x in increments if x["arm"]==arms[3] and x["prefix_steps"]==40 and x["metric"]=="NRMSE"]
        _quant_save(fig,"quant_18_public_source_increment_all_arms","全部监督原型的O→O+S条件公开增量",
                    f"随机分布原型40点前缀的O→O+S ΔNRMSE三种子分别为"+
                    "、".join(f'{x["delta"]:+.5f}' for x in dist40)+"，方向一致为变差。图同时保留随机点、受限点及0点前缀，"
                    "横线是优化重复的最小—最大范围，菱形为均值，不是独立统计置信区间。",
                    ["tokenizer/public_probes/*/seed_*/records.csv"],increments,
                    "来源：public probes records.csv；Δ=O+S−O，负为改善；同一冻结监督模型的O与O+S；范围非CI。")

        # 19: observation-only semantics are a trained linear readout of frozen
        # reconstruction features, not the raw unsupervised source head.
        probe_scenes=["baseline","non_dct_impulse","no_coupling","exact_joint_alias"]
        probe_labels=["基准","未见脉冲干预","无耦合干预","精确联合别名干预"]
        probe_arm_labels=["纯观测 + 监督线性readout","受限点（S头）","随机点（S头）","分布（S头）"]
        fig,axs=_quant_new("四臂合成评价：纯观测表征使用监督线性readout",ncols=2)
        fig.subplots_adjust(left=.15,right=.96,top=.86,bottom=.24,wspace=.45)
        probe_stats=[]
        probe_colors=["#768391",colors[0],colors[1],colors[2]]
        for ax,modality in zip(axs[0],["eeg","hb"]):
            largest=0.0
            for group,scene in enumerate(probe_scenes):
                ax.axhspan(group*5-.5,group*5+3.5,color="#EDF2F5" if group%2==0 else "white",zorder=0)
                for index,arm in enumerate(arms):
                    row=_quant_one(metrics,arm=arm,modality=modality,scenario=scene,variant="intervention")
                    mean=float(row.shape_rmse_seed_mean)
                    low=float(row.shape_rmse_seed_min);high=float(row.shape_rmse_seed_max)
                    y=group*5+index
                    ax.barh(y,mean,height=.72,color=probe_colors[index],label=probe_arm_labels[index] if group==0 else None,
                            hatch="///" if index==0 else None,edgecolor="white",linewidth=.5,zorder=2)
                    ax.plot([low,high],[y,y],color="#17334D",lw=1.2,zorder=3)
                    ax.annotate(f"{mean:.3f}",(high,y),xytext=(4,0),textcoords="offset points",va="center",fontsize=11.5,color="#17334D")
                    probe_stats.append(dict(row.to_dict(),semantic_prediction_route="frozen_all_code_features_to_training_truth_ridge" if index==0 else "direct_semantically_supervised_shape_logA_head"))
                    largest=max(largest,high)
            ax.set_yticks([group*5+1.5 for group in range(4)],probe_labels,fontsize=11.5)
            ax.set_ylim(18.8,-.8);ax.set_xlim(0,largest*1.19)
            ax.tick_params(axis="x",labelsize=11.5)
            ax.set_xlabel("形状RMSE（越低越好）",fontsize=11.5)
            ax.set_title(modality.upper()+" · 相同随机评价族",fontsize=13)
            ax.grid(axis="x",alpha=.15,zorder=1)
        handles,labels=axs[0,0].get_legend_handles_labels()
        fig.legend(handles,labels,ncols=4,loc="lower center",bbox_to_anchor=(.5,.08),frameon=False,fontsize=11.5,
                   handlelength=1.5,columnspacing=1.1)
        baseline_obs=[_quant_one(metrics,arm=arms[0],modality=m,scenario="baseline",variant="intervention").shape_rmse_seed_mean for m in ["eeg","hb"]]
        baseline_dist=[_quant_one(metrics,arm=arms[3],modality=m,scenario="baseline",variant="intervention").shape_rmse_seed_mean for m in ["eeg","hb"]]
        uncoupled_obs=_quant_one(metrics,arm=arms[0],modality="hb",scenario="no_coupling",variant="intervention").shape_rmse_seed_mean
        uncoupled_dist=_quant_one(metrics,arm=arms[3],modality="hb",scenario="no_coupling",variant="intervention").shape_rmse_seed_mean
        _quant_save(fig,"quant_19_observation_linear_probe","纯观测监督线性readout与三种语义S头的四臂合成比较",
                    "纯观测臂先冻结重建编码器，再用训练集真值拟合all_code_features→shape/logA ridge（α=1）并随checkpoint保存probes；"
                    "保留评价调用semantic_predictions读取此线性readout，并非原始未监督S头。其余三臂直接评价语义监督的shape/logA头。"
                    "458168等参数量只指基础网络，纯观测臂另附ridge参数。"
                    f"基准EEG/Hb形状RMSE：纯观测线性readout为{baseline_obs[0]:.4f}/{baseline_obs[1]:.4f}，分布S头为{baseline_dist[0]:.4f}/{baseline_dist[1]:.4f}。"
                    f"无耦合干预Hb的误差分别为{uncoupled_obs:.4f}/{uncoupled_dist:.4f}，不能声称S头在所有场景普遍优越。"
                    "柱为三个优化重复均值、横线为种子最小—最大范围，非独立种子CI；四场景统一展示干预变体，log幅值误差另附JSON。",
                    ["tokenizer/scenario_variant_seed_summary.csv",
                     "source_snapshot_training_v3/experiments/scripts/train_semantic_response_tokenizer.py:semantic_predictions/fit"],probe_stats,
                    "来源：scenario_variant_seed_summary.csv及冻结训练源码；256共享base×3优化重复；O-only=冻结重建表征+训练真值线性readout，非原始未监督S头。",
                    bitmap_notes=False)
        figures[-1]["section"]="tokenizer"
        figures[-1]["readout_contract"]="O-only: train-only supervised ridge alpha=1 on frozen reconstruction features, with additional ridge parameters beyond the 458168-parameter base network; other arms: direct semantically supervised shape/logA heads"

    manifest = dict(schema="semantic_response_quantitative_report_v1", figures=figures,
                    audit=audit, tables=tables, source_files=sorted(set(sources)),
                    reproduction="build_response_quantitative_figures(repo, run, out)",
                    output_boundary="bitmap PNG figures and descriptive JSON only; retained evidence unchanged")
    clean = _quant_json_clean(manifest)
    (out / "quantitative_provenance.json").write_text(json.dumps(clean,ensure_ascii=False,indent=2)+"\n")
    return clean

"""Report-only native/prepared/retained-response evidence; no fitting or cache writes."""


def build_response_waveform_figures(repo, run, out):
    """Export adjacent native, Hb, driver/state and EEG PNGs for dataset cases.

    Case selection is descriptive, from evaluation regional rows at the fixed
    40-point Hb prefix. Raw reads use the registered native readers. The only
    replay is deterministic feature preparation and array reconstruction.
    """
    import csv
    import gc
    import json
    from pathlib import Path
    import time

    import matplotlib.pyplot as plt
    from matplotlib.text import Text
    import numpy as np
    import yaml

    from src.data.clean_physiology_cache import CleanPhysiologyCacheIndex
    from src.data.registry import REGISTERED_DATASETS
    from src.data.unified_physiology import load_native_eeg_record, load_native_fnirs_record
    from src.inference.observation_baselines import eeg_band_power, native_feature_operators
    from src.inference.shared_driver_attribution import equal_capacity_hb_basis
    from src.inference.shared_driver_modes import spectral_loadings

    repo, run, out = (Path(p).resolve() for p in (repo, run, out))
    figroot = out / "figures"
    figroot.mkdir(parents=True, exist_ok=True)
    started = time.monotonic()

    def _wave_read(path):
        return json.loads(Path(path).read_text(encoding="utf-8"))

    def _wave_arrays(path):
        with np.load(path, allow_pickle=False) as arrays:
            return {k: arrays[k].copy() for k in arrays.files}

    def _wave_error(actual, expected, label, atol=1e-10):
        actual, expected = np.asarray(actual), np.asarray(expected)
        if actual.shape != expected.shape or not np.isfinite(actual).all() or not np.isfinite(expected).all():
            raise ValueError(f"Waveform replay shape/finite failure: {label}")
        error = float(np.max(np.abs(actual - expected)))
        if not np.allclose(actual, expected, rtol=1e-10, atol=atol):
            raise ValueError(f"Waveform replay mismatch: {label}: {error}")
        return dict(max_absolute_error=error, rtol=1e-10, atol=atol, passed=True)

    def _wave_path(path, root):
        path = Path(path).resolve()
        if not path.is_relative_to(root.resolve()):
            raise ValueError("Waveform source outside retained owning root")
        return path

    cfg = _wave_read(run / "resolved_config.json")
    plan = _wave_read(run / "measured_plan.json")
    parent = (repo / cfg["source_run"]).resolve()
    if (cfg.get("schema") != "semantic_response_discovery_v1"
            or cfg.get("protected_data") != "forbidden"
            or Path(plan["parent"]).resolve() != parent
            or cfg["tensor"]["steps"] != 120 or cfg["tensor"]["dt_s"] != .25):
        raise ValueError("Unexpected waveform evidence contract")
    source_cfg = yaml.safe_load((repo / cfg["source_config"]).read_text())
    evaluation = {x["ref"]["id"]: x for x in plan["windows"] if x["split"] == "evaluation"}
    if len(evaluation) != 360 or len({x["native_id"] for x in evaluation.values()}) != 120:
        raise ValueError("Expected exact 360 regional / 120 native evaluation inventory")
    datasets = ("eeg_fnirs_single_trial", "simultaneous_eeg_nirs", "visual_cognitive_motivation")
    if {x["ref"]["dataset"] for x in evaluation.values()} != set(datasets):
        raise ValueError("Unexpected waveform dataset before raw read")
    with (run / "measured_cells.csv").open(newline="") as handle:
        rows = [r for r in csv.DictReader(handle)
                if r["prefix_steps"] == "40" and r["pairing"] == "real"
                and r["arm"] in ("best_stable", "gain_selected")]
    best = {r["id"]: r for r in rows if r["arm"] == "best_stable"}
    gain = {r["id"]: r for r in rows if r["arm"] == "gain_selected"}
    choices = []
    for ds in datasets:
        pairs = [(float(gain[k]["Hb_nrmse"]) - float(r["Hb_nrmse"]), k)
                 for k, r in best.items() if r["dataset"] == ds and k in evaluation and k in gain
                 and r["success"] == "True" and gain[k]["success"] == "True"]
        expected = sum(x["ref"]["dataset"] == ds for x in evaluation.values())
        if len(pairs) != expected:
            raise ValueError("Case selection requires the complete paired evaluation denominator")
        median = float(np.median([d for d, _ in pairs]))
        chosen = (min(pairs, key=lambda p: (abs(p[0] - median), p[1])), min(pairs))
        for label, (delta, identity) in zip(("median", "worst"), chosen):
            choices.append(dict(id=identity, dataset=ds, role=label, delta=delta,
                                paired_population=len(pairs), population_median_delta=median))

    # This index reads registered metadata, never signal arrays or protected splits.
    index = CleanPhysiologyCacheIndex(repo / source_cfg["data"]["cache_root"])
    native_op = native_feature_operators(120)
    hb_model_op = native_op["fnirs"] @ native_op["native_interpolation"]
    block = np.zeros((360, 360))
    block[0::3, 0::3] = native_op["eeg"]
    block[1::3, 1::3] = block[2::3, 2::3] = hb_model_op
    common_basis = equal_capacity_hb_basis(block, modes=4, rho=.35).reshape(120, 3, 4)[:, 1:]
    t = np.arange(120) * .25
    prefix_s, score_start_s, score_stop_s = 10., cfg["measured"]["score_start"] * .25, 30.
    names = {"eeg_fnirs_single_trial": "Single-Trial", "simultaneous_eeg_nirs": "Simultaneous",
             "visual_cognitive_motivation": "Visual"}
    palette = dict(observed="#17212b", physical="#386cb0", component="#d98b26",
                   prediction="#7b4ab5", gain="#7a7a7a", HbO="#d1495b", HbR="#2374ab")
    figures, cases, loaded = [], [], None
    raw_eeg, raw_hb = None, None

    def _wave_style(ax, shade=True):
        if shade:
            ax.axvspan(0, prefix_s, color="#80bfa3", alpha=.15, zorder=0)
            ax.axvspan(score_start_s, score_stop_s, color="#e6bd65", alpha=.20, zorder=0)
        ax.axvline(prefix_s, color="#548a70", ls="--", lw=.8)
        ax.axvline(score_start_s, color="#aa8130", ls="--", lw=.8)
        ax.axvline(5., color="#acb4bc", ls=":", lw=.7)
        ax.set_xlim(0, 30)
        ax.spines[["top", "right"]].set_visible(False)
        ax.grid(axis="y", alpha=.18)
        ax.tick_params(labelsize=11.5)

    def _wave_save(fig, key, title, caption, source, stats):
        path = figroot / f"{key}.png"
        fig.canvas.draw()
        renderer = fig.canvas.get_renderer()
        visible_text = [artist for artist in fig.findobj(match=Text)
                        if artist.get_visible() and artist.get_text()]
        outside = []
        for artist in visible_text:
            bounds = artist.get_window_extent(renderer)
            if (bounds.x0 < -.5 or bounds.y0 < -.5
                    or bounds.x1 > fig.bbox.width + .5 or bounds.y1 > fig.bbox.height + .5):
                outside.append(dict(text=artist.get_text(), bounds=list(bounds.extents)))
        if outside:
            raise ValueError(f"Waveform text outside image: {key}: {outside}")
        min_font = min(a.get_fontsize() for a in visible_text)
        fig.savefig(path, dpi=250, facecolor="white")
        plt.close(fig)
        item = dict(key=key, path=str(path), title=title, caption=caption,
                    source=source, stats=stats, export_version="wide_final",
                    figure_inches=list(fig.get_size_inches()), minimum_font_size_pt=min_font,
                    minimum_scaled_font_size_pt=min_font * min(12. / fig.get_figwidth(), 4.72 / fig.get_figheight()),
                    text_bounds_checked=len(visible_text), text_clipping=False, legends_outside_data=True)
        figures.append(item)
        return item

    with plt.rc_context({"font.size": 12, "axes.titlesize": 12, "axes.labelsize": 12,
                         "legend.fontsize": 11.5, "font.family": "DejaVu Sans"}):
        for number, chosen in enumerate(choices, 1):
            item = evaluation[chosen["id"]]
            ref = item["ref"]
            ds = ref["dataset"]
            if (ds not in datasets or (ds == "visual_cognitive_motivation"
                                      and ref["subject"] == "S06" and "Part1" in ref["record"])):
                raise ValueError("Excluded dataset/window before native read")
            expected_key = f'{ds}__{ref["subject"]}__{ref["record"]}'
            if ref["key"] != expected_key or ref["array_index"] != ref["window"]:
                raise ValueError("Native identity drift")
            array_path = _wave_path(ref["array_path"], parent / "prepared")
            if array_path != parent / "prepared" / expected_key / (ref["region"] + ".npz"):
                raise ValueError("Prepared path does not match native identity")
            prepared_record = _wave_read(array_path.parent / "record.json")
            spec = prepared_record["spec"]
            join = f'{ds}|{ref["subject"]}|{ref["record"]}'
            if spec["join_key"] != join:
                raise ValueError("Prepared/native join mismatch")
            anchor = next(a for a in prepared_record["prepared"] if a["region"] == ref["region"])
            if anchor["eeg_channels"] != ref["eeg_channels"] or anchor["hb_channel"] != ref["hb_channel"]:
                raise ValueError("Prepared/native channel mismatch")
            windows = [w for w in anchor["windows"] if w["window"] == ref["window"]]
            if (len(windows) != 1 or any(windows[0][k] != ref[k] for k in
                                        ("eeg_start_s", "hb_start_s", "event_id", "task", "condition"))):
                raise ValueError("Prepared/native time/event mismatch")
            records = index.records_by_join_key.get(join, [])
            if len(records) != 1:
                raise ValueError("Registered native identity is not unique")
            record = records[0]
            if loaded != join:
                raw_eeg, raw_hb = None, None
                gc.collect()
                raw_eeg = load_native_eeg_record(repo, record)
                raw_hb = load_native_fnirs_record(repo, record)
                loaded = join
            registered_root = (repo / REGISTERED_DATASETS[ds].default_root).resolve()
            _wave_path(raw_eeg.source_path, registered_root)
            for p in raw_hb["provenance"]["source_paths"]:
                _wave_path(p, registered_root)
            lookup = {str(n).upper(): i for i, n in enumerate(raw_eeg.channel_names)}
            eeg_indices = [lookup[n.upper()] for n in ref["eeg_channels"]]
            eeg_start = round(ref["eeg_start_s"] * raw_eeg.sample_rate_hz)
            eeg_length = round(30 * raw_eeg.sample_rate_hz)
            eeg_values = raw_eeg.values[eeg_start:eeg_start + eeg_length, eeg_indices]
            eeg_time = np.arange(eeg_length) / raw_eeg.sample_rate_hz
            if eeg_values.shape != (eeg_length, 6) or not np.isfinite(eeg_values).all():
                raise ValueError("Incomplete native EEG window")
            hb_pair = int(anchor["hb_pair"])
            if raw_hb["channel_names"][2 * hb_pair:2 * hb_pair + 2] != [
                    ref["hb_channel"] + "_HbO", ref["hb_channel"] + "_HbR"]:
                raise ValueError("Native Hb pair identity drift")
            native_hb_select = (raw_hb["time_s"] >= ref["hb_start_s"]) & (
                raw_hb["time_s"] < ref["hb_start_s"] + 30)
            native_hb_time = raw_hb["time_s"][native_hb_select] - ref["hb_start_s"]
            native_hb_values = raw_hb["values"][native_hb_select, hb_pair]
            hb_grid = ref["hb_start_s"] + np.arange(300) / 10
            if hb_grid[0] < raw_hb["time_s"][0] or hb_grid[-1] > raw_hb["time_s"][-1]:
                raise ValueError("Native Hb window outside actual support")
            hb_input = np.column_stack([np.interp(hb_grid, raw_hb["time_s"],
                                                  raw_hb["values"][:, hb_pair, j]) for j in range(2)])
            prepared = _wave_arrays(array_path)
            eeg_prepared = prepared["eeg_features"][ref["array_index"]]
            hb_prepared = prepared["hb"][ref["array_index"]]
            power = eeg_band_power(eeg_values, bands=cfg["tensor"]["eeg_bands_hz"],
                                   sample_rate=raw_eeg.sample_rate_hz, target_rate=4.)
            if power.shape != (120, 6, 5) or np.any(power <= 0):
                raise ValueError("Invalid raw-to-prepared EEG power")
            checks = dict(
                raw_to_prepared_eeg=_wave_error(native_op["eeg"] @ np.log(power).reshape(120, 30),
                                                eeg_prepared, "raw to prepared EEG"),
                raw_to_prepared_hb=_wave_error(native_op["fnirs"] @ hb_input, hb_prepared,
                                               "raw to prepared Hb"))
            feature_path = run / "features" / "evaluation" / (ref["id"] + ".npz")
            feature = _wave_arrays(feature_path)
            feature_meta = _wave_read(feature_path.with_suffix(".json"))
            if feature_meta["ref"] != ref or feature_meta["split"] != "evaluation":
                raise ValueError("Retained feature identity drift")
            cal = _wave_read(run / "calibration" / (item["calibration_key"] + ".json"))
            checks["prepared_to_scaled_eeg"] = _wave_error(eeg_prepared * cal["eeg_factor"], feature["eeg"], "scaled EEG")
            checks["prepared_to_scaled_hb"] = _wave_error(hb_prepared * cal["hb_factor"], feature["hb"], "scaled Hb")
            row = _wave_read(run / "measured" / ref["id"] / "best_stable" / "real" / "prefix_40.json")
            gain_row = _wave_read(run / "measured" / ref["id"] / "gain_selected" / "real" / "prefix_40.json")
            response_path = _wave_path(run / row["arrays"], run / "measured")
            gain_path = _wave_path(run / gain_row["arrays"], run / "measured")
            response, control = _wave_arrays(response_path), _wave_arrays(gain_path)
            if row["id"] != ref["id"] or gain_row["id"] != ref["id"]:
                raise ValueError("Response identity mismatch")
            prediction, physical, component = (response[k] for k in ("prediction", "physical_prediction", "observation_component"))
            checks["physical_plus_common_equals_prediction"] = _wave_error(physical + component, prediction, "component sum")
            checks["common_basis_reconstruction"] = _wave_error(common_basis @ response["component_coefficients"], component, "common basis")
            states = response["states"]
            fixed = source_cfg["fixed"]
            canonical = np.column_stack((fixed["P0"] * (states[:, 4] - 1) - fixed["Q0"] * (states[:, 5] - 1),
                                         fixed["Q0"] * (states[:, 5] - 1)))
            checks["states_to_physical_hb"] = _wave_error(hb_model_op @ canonical, physical, "state to Hb")
            modes, selected_spec = feature["eeg_modes"], row["selected_spec"]
            modes_path = run / "eeg_modes" / item["calibration_key"] / ref["id"] / "real__fixed.npz"
            mode_arrays = _wave_arrays(modes_path)
            checks["retained_feature_modes_match_EEG_cache"] = _wave_error(modes, mode_arrays["r"], "retained modes")
            loading = spectral_loadings()
            eeg_a = native_op["eeg"] @ np.outer(modes[:, 0], loading[:, 0]) / cal["eeg_factor"]
            eeg_b = native_op["eeg"] @ np.outer(modes[:, 1], loading[:, 1]) / cal["eeg_factor"]
            eeg_prediction = mode_arrays["prediction"] / cal["eeg_factor"]
            eeg_residual = eeg_prepared - eeg_prediction
            checks["EEG_a_plus_b_equals_retained_prediction"] = _wave_error(eeg_a + eeg_b, eeg_prediction, "EEG loading contributions")
            checks["prepared_EEG_equals_a_plus_b_plus_residual"] = _wave_error(eeg_a + eeg_b + eeg_residual, eeg_prepared, "EEG decomposition")
            g = selected_spec["gamma"]
            expected_driver = ((modes[:, 0] + g * modes[:, 1]) / np.sqrt(1 + g * g)
                               if g else selected_spec["gain"] * modes[:, 0])
            checks["eeg_modes_to_driver"] = _wave_error(expected_driver, response["driver"], "EEG-only driver")
            endpoint = response["endpoint_mask"]
            if endpoint.shape != (120, 2) or not np.array_equal(endpoint, control["endpoint_mask"]):
                raise ValueError("Paired score support mismatch")
            nrmse = float(np.sqrt(np.mean(((prediction - feature["hb"]) / np.asarray(cal["sd_hb"]))[endpoint] ** 2)))
            checks["retained_nrmse_recomputed"] = _wave_error(np.array(nrmse), np.array(row["Hb_nrmse"]), "NRMSE")
            hb_factor = float(cal["hb_factor"])
            if hb_factor <= 0:
                raise ValueError("Hb inverse scale must be common positive")
            prediction, physical, component, gain_prediction = (a / hb_factor for a in
                (prediction, physical, component, control["prediction"]))
            residual = hb_prepared - prediction
            checks["prepared_equals_physical_plus_common_plus_observation_residual"] = _wave_error(
                physical + component + residual, hb_prepared, "prepared decomposition")
            hb_unit = ("relative MBLL units" if ds == "eeg_fnirs_single_trial" else
                       "mmol/L" if ds == "simultaneous_eeg_nirs" else "released units (unreported)")
            hb_axis_unit = ("relative Hb" if ds == "eeg_fnirs_single_trial" else
                            "mmol/L" if ds == "simultaneous_eeg_nirs" else "released units")
            eeg_unit = (raw_eeg.channel_units[eeg_indices[0]] if raw_eeg.channel_units else raw_eeg.native_unit)
            raw_label = "raw released Hb" if ds != "eeg_fnirs_single_trial" else "raw photodetector intensity"
            short = f'{names[ds]} · {ref["subject"]} · {ref["record"]} · {ref["region"]} · window {ref["window"]}'
            detail = f'EEG clock start {ref["eeg_start_s"]:.3f}s; Hb clock start {ref["hb_start_s"]:.3f}s; {ref["condition"]}'
            role = "Closest to dataset median" if chosen["role"] == "median" else "Worst paired change"
            statline = f'{role}: gain NRMSE {float(gain_row["Hb_nrmse"]):.3f} → stable {nrmse:.3f}; Δ(gain−stable)={chosen["delta"]:+.3f}'
            stats = dict(chosen, best_stable_nrmse=nrmse, gain_nrmse=float(gain_row["Hb_nrmse"]),
                         chosen_arm=row["chosen_arm"], spec=selected_spec, numerical_checks=checks)

            compact = f'{names[ds]} {ref["subject"]} · {ref["region"]} · w{ref["window"]}'
            fig = plt.figure(figsize=(13.4, 5.3))
            grid = fig.add_gridspec(2, 2, height_ratios=[1., 1.5], left=.07, right=.935,
                                   top=.82, bottom=.22, hspace=.70, wspace=.24)
            axes = [fig.add_subplot(grid[0, 0]), fig.add_subplot(grid[0, 1]), fig.add_subplot(grid[1, :])]
            fig.suptitle("Native signals → prepared EEG bands | " + compact,
                         x=.07, y=.982, ha="left", fontsize=16, weight="bold")
            fig.text(.07, .915, detail + "   |   " + role, fontsize=12, color="#4e5965")
            axes[0].plot(eeg_time, eeg_values[:, 0], color="#386cb0", lw=.55)
            axes[0].set_ylabel(f'EEG {eeg_unit}')
            axes[0].set_title(f'Native voltage: {ref["eeg_channels"][0]} (1 of 6 channels)', loc="left")
            _wave_style(axes[0])
            zoom = axes[0].inset_axes([.63, .47, .36, .51])
            zoom_mask = (eeg_time >= 20) & (eeg_time < 21)
            zoom.plot(eeg_time[zoom_mask], eeg_values[zoom_mask, 0], color="#386cb0", lw=.6)
            zoom.text(.5, .94, "20–21 s", transform=zoom.transAxes, ha="center", va="top", fontsize=11.5)
            zoom.set_xticks([])
            zoom.set_yticks([])
            zoom.set_facecolor("white")
            if ds == "eeg_fnirs_single_trial":
                values = raw_hb["optical_intensity"][native_hb_select, hb_pair]
                labels, colors, unit = ("760 nm", "850 nm"), ("#795da8", "#128c8d"), "optical V"
            else:
                values, labels, colors, unit = native_hb_values, ("HbO", "HbR"), (palette["HbO"], palette["HbR"]), hb_axis_unit
            for j, (label, color) in enumerate(zip(labels, colors)):
                axes[1].plot(native_hb_time, values[:, j], color=color, label=label, lw=1.3)
            raw_title = ("Raw optical" if ds == "eeg_fnirs_single_trial" else "Raw Hb")
            unknown_unit = " [unit unreported]" if ds == "visual_cognitive_motivation" else ""
            axes[1].set_title(f'{raw_title}: {ref["hb_channel"]}{unknown_unit}', loc="left")
            axes[1].set_ylabel(unit)
            axes[1].legend(loc="lower right", bbox_to_anchor=(1, 1.045), ncol=2,
                           frameon=False, borderaxespad=0)
            _wave_style(axes[1])
            extent = float(np.max(np.abs(eeg_prepared)))
            image = axes[2].imshow(eeg_prepared.T, aspect="auto", cmap="RdBu_r", vmin=-extent, vmax=extent,
                                   extent=(0, 30, 30, 0), interpolation="nearest")
            axes[2].set_yticks(np.arange(6) * 5 + 2.5, ref["eeg_channels"])
            for boundary in range(5, 30, 5):
                axes[2].axhline(boundary, color="white", lw=.7, alpha=.85)
            axes[2].set_title("Prepared EEG: 6 channels × 5 bands; log power relative to first 5 s", loc="left")
            axes[2].set_xlabel("Seconds from each registered modality window start")
            _wave_style(axes[2], shade=False)
            cax = fig.add_axes([.948, .22, .012, .26])
            colorbar = fig.colorbar(image, cax=cax)
            colorbar.ax.tick_params(labelsize=11.5)
            colorbar.ax.set_title("Δ log\npower", fontsize=11.5, pad=6)
            footer = "Bands: 1–4 / 4–8 / 8–13 / 13–30 / 30–45 Hz. Green: Hb input 0–10s; gold: score 20–30s."
            footer2 = ("Single-Trial optical V → natural-log OD → approximate relative MBLL Hb; the derived Hb is not a raw concentration."
                       if ds == "eeg_fnirs_single_trial" else
                       "Visual Hb has unreported physical units; it is not labelled µM." if ds == "visual_cognitive_motivation" else
                       "Released Hb remains in verified mmol/L; this plot applies no concentration-unit conversion.")
            fig.text(.07, .072, footer, fontsize=11.5, color="#4e5965")
            fig.text(.07, .024, footer2, fontsize=11.5, color="#4e5965")
            source = dict(raw_eeg=str(raw_eeg.source_path), raw_hb=raw_hb["provenance"]["source_paths"],
                          prepared=str(array_path), feature=str(feature_path), modes=str(modes_path),
                          response=str(response_path), control=str(gain_path))
            raw_fig = _wave_save(fig, f"waveform_{number:02d}_{chosen['role']}_native_wide", "原始信号与频带特征：" + short,
                "精确回溯本轮评价窗口。原始电压与发布Hb/光强保持各自单位；频带图展示已有局部特征变换，不把prepared称为raw。", source, stats)

            fig, axes = plt.subplots(3, 1, figsize=(12., 5.), sharex=True)
            fig.subplots_adjust(left=.12, right=.985, top=.77, bottom=.23, hspace=.27)
            fig.suptitle("Prepared Hb components | " + compact,
                         x=.12, y=.982, ha="left", fontsize=16, weight="bold")
            fig.text(.12, .915, f'Gain NRMSE {float(gain_row["Hb_nrmse"]):.3f} → stable {nrmse:.3f}; '
                     f'Δ={chosen["delta"]:+.3f} · {hb_unit} · {role}', fontsize=11.5, color="#4e5965")
            for j in range(2):
                ax = axes[j]
                for values, label, color, style, width in (
                    (hb_prepared, "Observed", palette["observed"], "-", 1.7),
                    (physical, "Physical", palette["physical"], "-", 1.25),
                    (component, "Common", palette["component"], "-", 1.15),
                    (prediction, "Phy + common", palette["prediction"], "-", 1.3),
                    (gain_prediction, "Gain ctrl", palette["gain"], "--", 1.1)):
                    ax.plot(t, values[:, j], label=label, color=color, ls=style, lw=width)
                ax.set_ylabel("HbO" if j == 0 else "HbR")
                ax.ticklabel_format(axis="y", style="sci", scilimits=(-2, 2))
                _wave_style(ax)
            handles, labels = axes[0].get_legend_handles_labels()
            for j, label in enumerate(("HbO", "HbR")):
                axes[2].plot(t, residual[:, j], label=label + " residual", color=palette[label], lw=1.2)
            axes[2].axhline(0, color="#9ba4ae", lw=.6)
            axes[2].set_ylabel("Residual")
            axes[2].ticklabel_format(axis="y", style="sci", scilimits=(-2, 2))
            residual_handles, residual_labels = axes[2].get_legend_handles_labels()
            fig.legend(handles + residual_handles, labels + residual_labels, loc="upper left",
                       bbox_to_anchor=(.12, .87), ncol=7, frameon=False, fontsize=11.5,
                       columnspacing=.8, handlelength=1.5, handletextpad=.4)
            _wave_style(axes[2])
            axes[2].set_xlabel("Seconds from the registered window start")
            fig.text(.12, .063, "Residual = observed − (physical + common). Green: Hb input 0–10s; gold: score 20–30s.", fontsize=11.5, color="#4e5965")
            fig.text(.12, .020, "Common/residual terms do not certify physiological sources or sensor noise.", fontsize=11.5, color="#4e5965")
            component_fig = _wave_save(fig, f"waveform_{number:02d}_{chosen['role']}_components_wide", "Hb成分分离：" + short,
                "同一prepared Hb坐标内展示物理模型项、附加共同观测项和观测残差，并对照冻结增益控制；绿色可见前缀和金色评分区显式分开。成分命名不等于来源确证。", source, stats)

            fig, axes = plt.subplots(2, 1, figsize=(12., 5.), sharex=True)
            fig.subplots_adjust(left=.09, right=.985, top=.80, bottom=.23, hspace=.35)
            fig.suptitle("EEG-owned driver → vascular model states | " + compact,
                         x=.09, y=.982, ha="left", fontsize=16, weight="bold")
            fig.text(.09, .915, f'Frozen {row["chosen_arm"]}: gain={selected_spec["gain"]:g}, '
                     f'τn={selected_spec["tau_n"]:g}s, τv={selected_spec["tau_v"]:g}s · {role}', fontsize=12, color="#4e5965")
            axes[0].plot(t, modes[:, 0], color="#989fb0", lw=1.1, label="EEG mode a")
            axes[0].plot(t, modes[:, 1], color="#68aaa0", lw=1.1, label="EEG mode b")
            axes[0].plot(t, response["driver"], color="#7b4ab5", lw=1.6, label="Fixed driver d")
            axes[0].plot(t, states[:, 0], color="#cf8750", lw=1.4, ls="--", label="Vascular input u")
            axes[0].set_ylabel("Driver / modes\nmodel coord.")
            axes[0].legend(loc="lower left", bbox_to_anchor=(0, 1.025), ncol=4,
                           frameon=False, borderaxespad=0)
            _wave_style(axes[0])
            for j, label, color in ((2, "f−1", "#d1495b"), (3, "v−1", "#386cb0"),
                                    (4, "p−1", "#7b4ab5"), (5, "q−1", "#128c8d")):
                axes[1].plot(t, states[:, j] - 1, label=label, color=color, lw=1.3)
            axes[1].plot(t, states[:, 1], label="s", color="#7e7e7e", lw=1.1, ls=":")
            axes[1].set_ylabel("States\nrelative coord.")
            axes[1].legend(loc="lower left", bbox_to_anchor=(0, 1.025), ncol=5,
                           frameon=False, borderaxespad=0)
            axes[1].set_xlabel("Seconds from the registered window start")
            _wave_style(axes[1])
            fig.text(.09, .060, "Green: Hb input 0–10s; gold: score 20–30s. Driver and states use model coordinates, separate from Hb units.", fontsize=11.5, color="#4e5965")
            fig.text(.09, .017, "Full-window processing is offline; these trajectories do not establish a causal forecast or unique physiological source.", fontsize=11.5, color="#4e5965")
            state_fig = _wave_save(fig, f"waveform_{number:02d}_{chosen['role']}_states_wide", "EEG驱动与血管状态：" + short,
                "固定EEG-only驱动d与实际血管输入u显式区分；相对状态f/v/p/q以及s单列，不能与Hb直接相加。保留前缀/评分区，采用保存的拟合状态，不新增拟合。", source, stats)

            fig, gridaxes = plt.subplots(2, 3, figsize=(12., 5.))
            fig.subplots_adjust(left=.06, right=.985, top=.80, bottom=.235, hspace=.58, wspace=.27)
            axes = gridaxes.ravel()
            fig.suptitle("EEG log-power components | " + compact,
                         x=.06, y=.982, ha="left", fontsize=15.5, weight="bold")
            fig.text(.06, .915, f'Channel {ref["eeg_channels"][0]} · a: broad band; b: α/β contrast · observed = a + b + residual',
                     fontsize=12, color="#4e5965")
            for band, ax in enumerate(axes[:5]):
                for values, label, color, style, width in (
                    (eeg_prepared, "Observed", palette["observed"], "-", 1.5),
                    (eeg_a, "a contribution", palette["physical"], "-", 1.05),
                    (eeg_b, "b contribution", "#128c8d", "-", 1.05),
                    (eeg_prediction, "a + b prediction", palette["prediction"], "-", 1.15),
                    (eeg_residual, "Observed − predicted", palette["component"], "--", 1.05)):
                    ax.plot(t, values[:, band], color=color, label=label, ls=style, lw=width)
                low, high = cfg["tensor"]["eeg_bands_hz"][band]
                ax.set_title(f'{low}–{high} Hz', loc="left", pad=3)
                ax.set_ylabel("Δ log power" if band in (0, 3) else "")
                _wave_style(ax)
                ax.set_xticks([0, 10, 20, 30])
                if band >= 3:
                    ax.set_xlabel("Window time (s)")
            axes[5].axis("off")
            handles, labels = axes[0].get_legend_handles_labels()
            axes[5].legend(handles, labels, loc="upper left", frameon=False, fontsize=11.5, borderaxespad=0)
            fig.text(.06, .064, "All 30 features verified. Green: paired Hb input 0–10s; gold: score 20–30s. EEG uses the full processed window.", fontsize=11.5, color="#4e5965")
            fig.text(.06, .018, "These contributions add in prepared log-power coordinates. They do not add in native voltage; residual is not certified noise.",
                     fontsize=11.5, color="#4e5965")
            eeg_fig = _wave_save(fig, f"waveform_{number:02d}_{chosen['role']}_eeg_components_wide", "EEG频带成分分离：" + short,
                "从保存的EEG-only a/b与固定谱载荷，在prepared相对log-power坐标恢复贡献；观测=a贡献+b贡献+残差。示图为与原始电压相同的一条通道的五频带，全30维重构均核验。电压波形不能与这些特征贡献相加。", source, stats)
            transforms = [
                dict(modality="EEG", operation="native reader decoding; no artifact cleanup added", unit=eeg_unit,
                     reference=raw_eeg.reference, evidence=raw_eeg.unit_evidence),
                dict(modality="EEG", operation="30s local fourth-order Butterworth SOS forward/backward bandpass; mean squared power in nonoverlapping 0.25s bins; natural log; first20-bin mean removed",
                     bands_hz=cfg["tensor"]["eeg_bands_hz"], feature_order="channel_major_band_minor", power_floor="none; all powers strictly positive"),
                dict(modality="Hb", operation=raw_hb["provenance"]["transform"], raw_provenance=raw_hb["provenance"]),
                dict(modality="Hb", operation="record-relative linear interpolation to 300 samples at10Hz; window-local 0.01–0.2Hz third-order SOS forward/backward filter; polyphase resample2/5 to4Hz; first20-bin mean removed", unit=hb_unit),
                dict(modality="EEG/Hb", operation="retained new-training common positive scalar; report Hb inverse-scaled to parent prepared coordinates",
                     eeg_factor=cal["eeg_factor"], hb_factor=hb_factor, cal_path=str(run / "calibration" / (item["calibration_key"] + ".json")), fitted_now=False),
                dict(modality="model", operation="retained EEG-only modes; retained fixed-driver response, prefix-fit initial state and common term; no new inference or fitting", state_order=["u", "s", "f", "v", "p", "q"], spec=selected_spec),
                dict(modality="EEG_components", operation="native baseline operator times retained a/b outer products with fixed unit-norm spectral loading; inverse retained EEG scale; residual=prepared−predicted",
                     loading=loading.tolist(), displayed_feature_indices=list(range(5)),
                     saved_model_residual_convention="prediction−observed; report displays observed−prediction", all30features_verified=True),
            ]
            cases.append(dict(**chosen, ref=ref, split="evaluation", native_id=item["native_id"],
                chosen_reason="dataset regional-row gain_selected minus best_stable NRMSE at40-pointprefix; nearest median or minimum delta; lexicographic ID tie break; descriptive illustrations selected after evaluation",
                identity_owner=str(run / "measured_plan.json"), sources=source, source_record_join=join,
                eeg=dict(channel_names=ref["eeg_channels"], displayed_native_channel=ref["eeg_channels"][0],
                         native_rate_hz=raw_eeg.sample_rate_hz, source_unit=raw_eeg.native_unit,
                         displayed_unit=eeg_unit, unit_evidence=raw_eeg.unit_evidence,
                         requested_start_s=ref["eeg_start_s"], actual_start_s=eeg_start / raw_eeg.sample_rate_hz,
                         native_sample_indices=[eeg_start, eeg_start + eeg_length], plotted_native_samples=eeg_length),
                hb=dict(pair=ref["hb_channel"], native_rate_hz=raw_hb["sample_rate_hz"],
                        raw_unit=raw_hb["provenance"]["native_unit"], prepared_unit=hb_unit,
                        clock_start_s=ref["hb_start_s"], actual_native_support_s=[float(native_hb_time[0]), float(native_hb_time[-1])],
                        plotted_native_samples=len(native_hb_time), raw_semantics=raw_label),
                transforms=transforms, checks=checks, metrics=stats,
                visible_hb_interval_s=[0, prefix_s], scored_hb_interval_s=[score_start_s, score_stop_s],
                event_window_time_s=5., figures=[raw_fig["key"], component_fig["key"], state_fig["key"], eeg_fig["key"]]))
    result = dict(schema="semantic_response_discovery_waveform_evidence_v1", figures=figures, cases=cases,
                  run=str(run), parent=str(parent), native_windows=120, evaluation_subjects=15,
                  regional_evaluation_rows=360, example_cases=len(cases), fitted_now=False,
                  raw_loader_owner="src/data/unified_physiology.py",
                  data_contract="docs/DATA_CONTRACT.md", export_version="wide_final",
                  retained_draft=dict(source=str(out / "build_sources" / "waveform_section_draft.py"),
                                      provenance=str(out / "waveform_provenance_draft.json"), figures_suffix="without _wide"),
                  slide_frame_inches=[12., 4.72], minimum_scaled_font_size_pt=min(f["minimum_scaled_font_size_pt"] for f in figures),
                  seconds=time.monotonic() - started)
    (out / "waveform_provenance.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return result

"""Read-only illustrations of retained synthetic cases and a frozen encoder.

The report function deliberately selects evaluation base 0 before reading scores.
It uses a fresh, isolated interpreter for the retained reader and generator.
No fitting, normalization update, or write to the evidence run is performed.
"""


def build_synthetic_waveform_figures(repo, run, out):
    """Return readable bitmap figures and their case-level audit manifest."""
    import json
    import os
    from pathlib import Path
    import subprocess
    import sys
    import tempfile

    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.font_manager import FontProperties
    from matplotlib.lines import Line2D
    from matplotlib.patches import Patch
    import numpy as np

    repo, run, out = (Path(value).resolve() for value in (repo, run, out))
    frozen = run / "source_snapshot_analysis_v1"
    tokenizer = run / "tokenizer"
    figures_dir = out / "figures"
    figures_dir.mkdir(parents=True, exist_ok=True)
    synw_prior_manifest_path = out / "synthetic_waveform_provenance.json"
    synw_prior_figures = {}
    if synw_prior_manifest_path.exists():
        synw_prior_figures = {item["key"]: item for item in json.loads(synw_prior_manifest_path.read_text())["figures"]}
    scenarios = ["baseline", "true_a_shape", "feature_gain_only", "true_a_amplitude",
                 "Hb_gain_only", "exact_joint_alias", "non_dct_impulse", "tau_change",
                 "correlated_Hb_component"]
    # The isolated interpreter prevents another report section's live src imports
    # from changing the namespace used by the frozen reader.
    synw_script = r'''
import importlib.util, json, sys
from pathlib import Path
import numpy as np
import torch
run, target = Path(sys.argv[1]), Path(sys.argv[2])
scenarios = json.loads(sys.argv[3])
frozen = run / "source_snapshot_analysis_v1"
sys.path.insert(0, str(frozen))
reader_path = frozen / "experiments/scripts/train_semantic_response_tokenizer.py"
spec = importlib.util.spec_from_file_location("synw_retained_reader", reader_path)
reader = importlib.util.module_from_spec(spec)
spec.loader.exec_module(reader)
from src.inference.semantic_response_synthetic import generate_response_case
import src.tokenizers.semantic_response_tokenizer as model_source
torch.set_num_threads(1)
torch.set_num_interop_threads(1)
root = run / "tokenizer"
resolved = json.loads((root / "resolved_tokenizer_config.json").read_text())
arm = "distribution_randomized_generator"
data = reader.load_prepared(root, arm, "evaluation", resolved["config_identity"])
base = str(data["base_id"][0])
indices = np.flatnonzero((data["base_id"] == base) & np.isin(data["scenario"], scenarios))
selected = {key: value[indices].copy() for key, value in data.items()}
del data
model, ckpt = reader.selected_model(root, arm, 101, resolved["options"], resolved["config_identity"], "cpu")
if set(selected["base_id"]) & set(ckpt["training_base_ids"] + ckpt["selection_base_ids"]):
    raise ValueError("selected illustration overlaps fitting/selection base IDs")
prediction = reader.predictions(model, selected, "cpu", len(indices), ckpt["amplitude_mean"], ckpt["amplitude_scale"])
arrays = {"data_" + key: value for key, value in selected.items()}
arrays["data_retained_row_index"] = indices
for modality, values in prediction.items():
    shape_sd, amplitude_sd = reader.uncertainty(arm, values, ckpt["calibration"][modality])
    values["shape_sd"] = shape_sd
    values["log_amplitude_sd"] = amplitude_sd
    branch = getattr(model, modality + "_branch")
    values["reconstruction_observation_coordinate"] = (values["reconstruction"] * branch.normalizer.scale.detach().numpy()
                                                       + branch.normalizer.mean.detach().numpy())
    for key, value in values.items():
        arrays["pred_" + modality + "_" + key] = value
for key in ("physical_hb", "observation_component", "hb_noise", "eeg_noise", "driver", "time_s"):
    arrays["truth_" + key] = []
checks = []
generated = {}
base_index = int(base.rsplit("/", 1)[-1])
generator_seed = int(resolved["options"]["generator_seed"])
for row, (scenario, variant) in enumerate(zip(selected["scenario"], selected["variant"])):
    if str(scenario) not in generated:
        generated[str(scenario)] = generate_response_case(generator_seed, str(scenario), "randomized",
                                                          partition="evaluation", base_id=base_index)
    truth = generated[str(scenario)][str(variant)]
    comparison = {}
    for key in ("eeg", "hb", "source_shape", "a"):
        rebuilt = np.asarray(truth[key], dtype=selected[key].dtype)
        delta = float(np.max(np.abs(rebuilt - selected[key][row])))
        comparison[key] = delta
        if not np.array_equal(rebuilt, selected[key][row]):
            raise ValueError(f"frozen generator no longer reproduces retained {scenario}/{variant}/{key}: {delta}")
    if float(truth["log_amplitude"]) != float(selected["log_amplitude"][row]):
        raise ValueError("retained amplitude identity differs")
    for key in ("physical_hb", "observation_component", "hb_noise", "eeg_noise", "driver", "time_s"):
        arrays["truth_" + key].append(truth[key])
    reconstructed_truth = (truth["metadata"]["hb_coordinate_gain"] * truth["physical_hb"]
                           + truth["observation_component"] + truth["hb_noise"])
    checks.append({"scenario": str(scenario), "variant": str(variant), "retained_row_index": int(indices[row]),
                   "bitwise_float32_replay_max_absolute_difference": comparison,
                   "generative_Hb_additive_closure_max_abs": float(np.max(np.abs(reconstructed_truth - truth["hb"])) )})
meta = {"base_id": base, "base_index": base_index, "arm": arm, "seed": 101, "generator_seed": generator_seed,
        "config_identity": resolved["config_identity"], "calibration": ckpt["calibration"],
        "amplitude_training_mean": ckpt["amplitude_mean"], "amplitude_training_scale": ckpt["amplitude_scale"],
        "reader_source": str(reader_path), "model_runtime_source": model_source.__file__,
        "generator_source": str(frozen / "src/inference/semantic_response_synthetic.py"),
        "selected_checkpoint": str(root / "fits" / arm / "seed_101/selected.pt"),
        "checkpoint_options": ckpt["options"], "replay_checks": checks,
        "inference_rows": len(indices), "inference_device": "cpu", "inference_threads": 1}
arrays["audit_json"] = np.asarray(json.dumps(meta, ensure_ascii=False))
np.savez_compressed(target, **{key: np.asarray(value) for key, value in arrays.items()})
'''
    with tempfile.TemporaryDirectory(prefix="semantic_waveforms_") as synw_tmp:
        synw_payload = Path(synw_tmp) / "frozen_subset.npz"
        synw_env = dict(os.environ, OMP_NUM_THREADS="1", OPENBLAS_NUM_THREADS="1", MKL_NUM_THREADS="1")
        subprocess.run([sys.executable, "-c", synw_script, str(run), str(synw_payload), json.dumps(scenarios)],
                       cwd=frozen, env=synw_env, check=True, capture_output=True, text=True)
        with np.load(synw_payload, allow_pickle=False) as synw_handle:
            synw_arrays = {key: synw_handle[key].copy() for key in synw_handle.files}

    synw_audit = json.loads(str(synw_arrays.pop("audit_json")))
    synw_data = {key[5:]: value for key, value in synw_arrays.items() if key.startswith("data_")}
    synw_truth = {key[6:]: value for key, value in synw_arrays.items() if key.startswith("truth_")}
    synw_pred = {modality: {key[len("pred_" + modality + "_"):]: value
                           for key, value in synw_arrays.items() if key.startswith("pred_" + modality + "_")}
                 for modality in ("eeg", "hb")}
    synw_metadata = [json.loads(str(value)) for value in synw_data["metadata_json"]]
    synw_time = synw_truth["time_s"][0]
    synw_colors = {"control": "#68788a", "intervention": "#8054b3", "truth": "#182c42",
                   "eeg": "#0084a4", "hb": "#d45050", "component": "#dc8b27", "noise": "#a5aab2"}
    synw_font_path = Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc")
    synw_font_name = FontProperties(fname=str(synw_font_path)).get_name() if synw_font_path.exists() else "DejaVu Sans"
    synw_previous_rc = plt.rcParams.copy()
    plt.rcParams.update({"font.family": synw_font_name, "font.size": 12, "axes.titlesize": 12,
                         "axes.labelsize": 12, "xtick.labelsize": 12, "ytick.labelsize": 12,
                         "legend.fontsize": 12, "axes.unicode_minus": False, "axes.spines.top": False,
                         "axes.spines.right": False, "figure.facecolor": "white", "axes.facecolor": "#fbfcfe"})
    synw_sources = {
        "prepared": str(tokenizer / "prepared/randomized/evaluation.npz"),
        "prepared_metadata": str(tokenizer / "prepared/randomized/evaluation.json"),
        "resolved_config": str(tokenizer / "resolved_tokenizer_config.json"),
        "checkpoint": synw_audit["selected_checkpoint"],
        "reader": synw_audit["reader_source"],
        "generator": synw_audit["generator_source"],
        "model": synw_audit["model_runtime_source"],
        "existing_evaluation_record": str(tokenizer / "evaluation/distribution_randomized_generator/seed_101/record.json"),
    }
    synw_figures = []

    def synw_pair(scenario):
        found = {}
        for variant in ("control", "intervention"):
            rows = np.flatnonzero((synw_data["scenario"] == scenario) & (synw_data["variant"] == variant))
            if len(rows) != 1:
                raise ValueError(f"illustration needs exactly one {scenario}/{variant}")
            found[variant] = int(rows[0])
        return found["control"], found["intervention"]

    def synw_style(ax, title, ylabel=None, time=True):
        ax.set_title(title, loc="left", pad=6)
        ax.grid(True, linewidth=.4, alpha=.24)
        if ylabel:
            ax.set_ylabel(ylabel)
        if time:
            ax.set_xlim(synw_time[0], synw_time[-1])
            ax.set_xlabel("时间 / s")
            ax.axvspan(0, 4.75, color="#aab4c0", alpha=.10, lw=0)

    def synw_figure(title, nrows=2, ncols=3, height=5.1, identity_footer=False):
        fig, axes = plt.subplots(nrows, ncols, figsize=(12, height), squeeze=False)
        fig.suptitle(title, x=.04, ha="left", fontsize=14, fontweight="bold", y=.988)
        fig.subplots_adjust(left=.09, right=.985, top=.75, bottom=.15, hspace=.72, wspace=.42)
        if identity_footer:
            fig.text(.04, .025, "预定首例 evaluation/0  ·  distribution_randomized  ·  训练种子101  ·  灰底=前20点参考段；所有结果仅描述该个案",
                     fontsize=12, color="#566678")
        return fig, axes

    def synw_interval(ax, row, modality, title):
        truth = synw_data["source_shape"][row]
        mean, sd = synw_pred[modality]["shape_mean"][row], synw_pred[modality]["shape_sd"][row]
        ax.fill_between(synw_time, mean-1.959963984540054*sd, mean+1.959963984540054*sd,
                        color=synw_colors[modality], alpha=.16, label="95% 点边际预测区间")
        ax.plot(synw_time, truth, color=synw_colors["truth"], lw=1.6, label="生成器真形状")
        ax.plot(synw_time, mean, color=synw_colors[modality], lw=1.5, label=f"{modality.upper()}语义均值")
        synw_style(ax, title, "单位 RMS 形状")
        ax.legend(loc="best", frameon=False)

    def synw_amplitudes(ax, rows, title, label_variants=True):
        labels, truth_ys = [], []
        for item, row in enumerate(rows):
            role = "控制" if str(synw_data["variant"][row]) == "control" else "干预"
            y = float(item)
            truth = float(synw_data["log_amplitude"][row])
            ax.scatter(truth, y, c=synw_colors["truth"], marker="|", s=190, linewidths=1.8,
                       label="真 log幅值" if item == 0 else None, zorder=5)
            for delta, modality in ((-.15, "eeg"), (.15, "hb")):
                mu = float(synw_pred[modality]["log_amplitude_mean"][row])
                half = 1.959963984540054*float(synw_pred[modality]["log_amplitude_sd"][row])
                ax.errorbar(mu, y+delta, xerr=half, fmt="o", markersize=3.5, capsize=2,
                            color=synw_colors[modality], linewidth=1.1,
                            label=f"{modality.upper()}均值±95%区间" if item == 0 else None)
            labels.append(role if label_variants else str(item+1))
            truth_ys.append(y)
        ax.set_yticks(truth_ys, labels)
        ax.set_ylim(-.5, len(rows)-.5)
        ax.invert_yaxis()
        ax.set_xlabel("ln(RMS(a − 前20点均值))")
        synw_style(ax, title, time=False)
        ax.legend(loc="best", frameon=False)

    def synw_eeg_bands(ax, row, title, reconstruction=False):
        bands = synw_data["eeg"][row].reshape(120, 6, 5).mean(axis=1)
        palettes = ["#607d8b", "#bb9554", "#0084a4", "#d45050", "#8054b3"]
        # Only bands 2/3 are named alpha/beta by the retained generator.
        for band, name in enumerate(("坐标0", "坐标1", "α", "β", "坐标4")):
            ax.plot(synw_time, bands[:, band], color=palettes[band], lw=1.15, label=name)
            if reconstruction:
                rec = synw_pred["eeg"]["reconstruction_observation_coordinate"][row].reshape(120, 6, 5).mean(axis=1)
                ax.plot(synw_time, rec[:, band], color=palettes[band], lw=.8, ls="--", alpha=.70)
        synw_style(ax, title, "有效 Δlog-power")
        ax.legend(loc="upper right", ncol=3, frameon=False)

    def synw_observed_hb(ax, row, title, reconstruction=False):
        for channel, name, color in ((0, "HbO", "#d45050"), (1, "HbR", "#0084a4")):
            ax.plot(synw_time, synw_data["hb"][row, :, channel], lw=1.35, color=color, label=f"{name}观测")
            if reconstruction:
                ax.plot(synw_time, synw_pred["hb"]["reconstruction_observation_coordinate"][row, :, channel],
                        color=color, ls="--", lw=1.2, label=f"{name}完整重建")
        synw_style(ax, title, "合成 Hb 观测坐标")
        ax.legend(loc="best", frameon=False, ncol=2)

    def synw_delta(ax, control, intervention, title):
        ax.plot(synw_time, synw_data["source_shape"][intervention]-synw_data["source_shape"][control],
                color=synw_colors["truth"], lw=1.8, label="真 Δshape")
        for modality in ("eeg", "hb"):
            ax.plot(synw_time, synw_pred[modality]["shape_mean"][intervention]-synw_pred[modality]["shape_mean"][control],
                    color=synw_colors[modality], lw=1.35, label=f"{modality.upper()}估计变化")
        synw_style(ax, title, "干预 − 控制")
        ax.legend(frameon=False, loc="best")

    def synw_observed_pair(ax, control, intervention, modality, title, channels=None):
        if modality == "eeg":
            values = synw_data["eeg"].reshape(-1, 120, 6, 5).mean(axis=2)
            channels = [(2, "α"), (3, "β")] if channels is None else channels
            ylabel = "有效 Δlog-power"
        else:
            values = synw_data["hb"]
            channels = [(0, "HbO"), (1, "HbR")] if channels is None else channels
            ylabel = "合成 Hb 观测坐标"
        for channel, name in channels:
            for row, role, color in ((control, "控制", synw_colors["control"]),
                                     (intervention, "干预", synw_colors["intervention"])):
                ax.plot(synw_time, values[row, :, channel], lw=1.2, color=color,
                        ls="-" if channel == channels[0][0] else "--", label=f"{name}{role}")
        synw_style(ax, title, ylabel)
        ax.legend(frameon=False, ncol=2, loc="best")

    def synw_case_stats(scenario):
        control, intervention = synw_pair(scenario)
        stats = dict(scenario=scenario, base_id=synw_audit["base_id"],
                     case_ids=[f"{synw_audit['base_id']}/{scenario}/{role}" for role in ("control", "intervention")],
                     retained_row_indices=[int(synw_data["retained_row_index"][row]) for row in (control, intervention)],
                     mean_absolute_input_difference={modality: float(np.mean(np.abs(synw_data[modality][intervention]
                                                                                   -synw_data[modality][control])))
                                                      for modality in ("eeg", "hb")},
                     bitwise_identical_inputs={modality: bool(np.array_equal(synw_data[modality][intervention], synw_data[modality][control]))
                                              for modality in ("eeg", "hb")},
                     bitwise_identical_input_masks={modality: bool(np.array_equal(synw_data[modality+"_mask"][intervention], synw_data[modality+"_mask"][control]))
                                                    for modality in ("eeg", "hb")},
                     truth_shape_delta_RMS=float(np.sqrt(np.mean((synw_data["source_shape"][intervention]
                                                                -synw_data["source_shape"][control])**2))),
                     truth_log_amplitude_delta=float(synw_data["log_amplitude"][intervention]-synw_data["log_amplitude"][control]),
                     truth_parameter_control=synw_metadata[control]["parameters"]["free"],
                     truth_parameter_intervention=synw_metadata[intervention]["parameters"]["free"],
                     interpretation="preselected_single_case_not_scenario_aggregate", modalities={})
        for modality in ("eeg", "hb"):
            pred = synw_pred[modality]
            bitwise = {key: bool(np.array_equal(value[control], value[intervention]))
                       for key, value in pred.items()}
            stats["modalities"][modality] = {
                "estimated_shape_delta_RMS": float(np.sqrt(np.mean((pred["shape_mean"][intervention]-pred["shape_mean"][control])**2))),
                "estimated_log_amplitude_delta": float(pred["log_amplitude_mean"][intervention]-pred["log_amplitude_mean"][control]),
                "all_exported_outputs_bitwise_identical": all(bitwise.values()),
                "outputs_bitwise_equal_by_key": bitwise,
                "shape_MAE_control_intervention": [float(np.mean(np.abs(pred["shape_mean"][row]-synw_data["source_shape"][row])))
                                                   for row in (control, intervention)],
                "shape_95_point_marginal_case_coverage": [float(np.mean(np.abs(pred["shape_mean"][row]-synw_data["source_shape"][row])
                                                                        <=1.959963984540054*pred["shape_sd"][row]))
                                                         for row in (control, intervention)],
                "log_amplitude_95_intervals_control_intervention": [[float(pred["log_amplitude_mean"][row]-1.959963984540054*pred["log_amplitude_sd"][row]),
                                                                      float(pred["log_amplitude_mean"][row]+1.959963984540054*pred["log_amplitude_sd"][row])]
                                                                     for row in (control, intervention)],
                "log_amplitude_95_case_covered_control_intervention": [bool(abs(float(pred["log_amplitude_mean"][row]-synw_data["log_amplitude"][row]))
                                                                          <=1.959963984540054*float(pred["log_amplitude_sd"][row]))
                                                                      for row in (control, intervention)],
            }
        return stats

    synw_all_case_stats = {scenario: synw_case_stats(scenario) for scenario in scenarios}
    synw_manifest_path = out / "synthetic_waveform_provenance.json"

    def synw_shared_legend(fig, key):
        def synw_line(color, label, style="-"):
            return Line2D([], [], color=color, lw=1.7, ls=style), label

        truth = synw_line(synw_colors["truth"], "真shape / log幅值")
        eeg = synw_line(synw_colors["eeg"], "EEG语义")
        hb = synw_line(synw_colors["hb"], "Hb语义")
        interval = (Patch(facecolor="#8495a8", alpha=.20), "95%边际区间")
        reconstruction = synw_line("#566678", "完整重建（虚线）", "--")
        control = synw_line(synw_colors["control"], "控制")
        intervention = synw_line(synw_colors["intervention"], "干预")
        channel_styles = [synw_line("#566678", "α / HbO（实线）"),
                          synw_line("#566678", "β / HbR（虚线）", "--")]
        if key in ("01_baseline", "05_ood_impulse"):
            entries = [synw_line(color, label) for color, label in zip(
                ["#607d8b", "#bb9554", "#0084a4", "#d45050", "#8054b3"],
                ["坐标0", "坐标1", "α", "β", "坐标4"])]
            entries += [synw_line("#d45050", "HbO观测"), synw_line("#0084a4", "HbR观测"),
                        reconstruction, truth, eeg, hb, interval]
            columns = 6
        elif key == "02_shape_vs_eeg_gain":
            entries = [control, intervention, *channel_styles,
                       synw_line(synw_colors["truth"], "真Δshape"),
                       synw_line(synw_colors["eeg"], "EEG估计变化"),
                       synw_line(synw_colors["hb"], "Hb估计变化")]
            columns = 4
        elif key == "04_exact_joint_alias":
            entries = [control, intervention, *channel_styles,
                       synw_line(synw_colors["truth"], "真log幅值"), eeg, hb, interval]
            columns = 4
        elif key == "06_tau_response":
            entries = [control, synw_line(synw_colors["intervention"], "τ改变"), reconstruction,
                       truth, eeg, hb, interval,
                       synw_line("#d45050", "HbO响应Δ"), synw_line("#0084a4", "HbR响应Δ"),
                       synw_line("#566678", "观测Δ（虚线）", "--"),
                       synw_line("#566678", "α（实线）"), synw_line("#566678", "β（虚线）", "--")]
            columns = 6
        elif key == "07_correlated_hb":
            entries = [synw_line(synw_colors["truth"], "物理响应×gain"),
                       synw_line(synw_colors["component"], "真共同Hb项"),
                       synw_line(synw_colors["noise"], "配对噪声"),
                       synw_line(synw_colors["hb"], "分量和=观测", "--"),
                       synw_line("#d45050", "HbO观测"), synw_line("#0084a4", "HbR观测"),
                       reconstruction, truth, eeg, hb, interval]
            columns = 6
        else:
            return
        title_abbreviations = {
            "EEG：6通道平均的5个频带坐标": "EEG：6通道平均频带坐标",
            "Hb完整重建：D(S均值, O代码)": "Hb完整重建：D(S, O)",
            "EEG-only：连续 shape 与95%区间": "EEG-only：shape与95%区间",
            "Hb-only：连续 shape 与95%区间": "Hb-only：shape与95%区间",
            "幅值：两模态独立估计 ln(RMS)": "两模态独立log幅值估计",
            "语义真Δshape=0，估计是否跟着变？": "真Δshape=0：估计变化",
            "真幅值不变：log幅值估计与95%区间": "真幅值不变：估计95%区间",
            "真驱动与共同项：参考RMS内积=0.85": "真驱动与共同项：内积0.85",
            "真幅值不同，预测分布完全相同": "不同真幅值，同一预测分布",
            "生成器定位：变化由τ响应产生": "生成器定位：τ响应改变",
            "参考RMS幅值固定：控制与脉冲": "RMS幅值固定：控制vs脉冲",
            "网络：完整重建 D(S均值,O)": "网络完整重建：D(S, O)",
        }
        for ax in fig.axes:
            if ax.get_legend() is not None:
                ax.get_legend().remove()
            old_title = ax.get_title(loc="left")
            ax.set_title(title_abbreviations.get(old_title, old_title), loc="left", fontsize=12)
        legend = fig.legend([item[0] for item in entries], [item[1] for item in entries],
                            loc="upper center", bbox_to_anchor=(.535, .932), ncol=columns,
                            frameon=False, fontsize=12, columnspacing=1.15, handletextpad=.45)
        fig.canvas.draw()
        if legend.get_window_extent().x0 < 0 or legend.get_window_extent().x1 > fig.bbox.width:
            raise ValueError(f"wide_v2 legend exceeds figure width: {key}")

    def synw_save(fig, key, title, caption, cases):
        if not key.startswith("03"):
            synw_shared_legend(fig, key)
            key += "_wide_v2"
        path = figures_dir / f"syn_wave_{key}.png"
        fig.savefig(path, dpi=220, facecolor="white")
        plt.close(fig)
        record = dict(key=f"syn_wave_{key}", path=str(path), title=title, caption=caption,
                      source=synw_sources, stats={scenario: synw_all_case_stats[scenario] for scenario in cases},
                      statsmanifest=str(synw_manifest_path))
        synw_figures.append(record)

    control, intervention = synw_pair("baseline")
    fig, axes = synw_figure("合成基线：真实输入、完整重建与两模态独立语义输出")
    synw_eeg_bands(axes[0, 0], control, "EEG：6通道平均的5个频带坐标")
    synw_observed_hb(axes[0, 1], control, "HbO / HbR：实际合成输入")
    synw_observed_hb(axes[0, 2], control, "Hb完整重建：D(S均值, O代码)", reconstruction=True)
    synw_interval(axes[1, 0], control, "eeg", "EEG-only：连续 shape 与95%区间")
    synw_interval(axes[1, 1], control, "hb", "Hb-only：连续 shape 与95%区间")
    synw_amplitudes(axes[1, 2], [control], "幅值：两模态独立估计 ln(RMS)")
    synw_save(fig, "01_baseline", "合成基线：输入与独立语义",
              "预定 evaluation/0 基线，未按评分挑选。频带是有效基线相对 log-power 坐标，不是原始电压；源语义为首20点参考后的单位RMS形状与ln(RMS幅值)。各编码器只读自身模态。阴影为在selection冻结温度后给出的95%点边际预测区间，不是整条轨迹的联合可信带。Hb完整重建来自非线性D(S均值,O)，不能拆成可加的semantic Hb与observation Hb。",
              ["baseline"])

    fig, axes = synw_figure("反事实形状：真实驱动改变 vs. EEG读出增益改变")
    for nrow, (scenario, name) in enumerate((("true_a_shape", "真 shape 改变"), ("feature_gain_only", "EEG增益×1.8，真驱动不变"))):
        control, intervention = synw_pair(scenario)
        synw_observed_pair(axes[nrow, 0], control, intervention, "eeg",
                           "真shape改变：α / β输入" if nrow == 0 else "仅EEG增益×1.8：α / β输入")
        synw_observed_pair(axes[nrow, 1], control, intervention, "hb", "同一反事实：HbO / HbR输入")
        synw_delta(axes[nrow, 2], control, intervention, "语义变化：真值与估计的差")
    synw_save(fig, "02_shape_vs_eeg_gain", "形状变化还是读出尺度变化？",
              "两行使用同一基础轨迹与配对噪声。上行真shape改变但参考RMS幅值固定；下行仅EEG共同log-power坐标增益乘1.8，Hb和真shape均保持不变。右列直接画干预减控制的shape：非零估计且真值为零表示该个案的虚假语义变化。图中没有对差分画95%区间，因为模型未提供两次预测或时间点之间的协方差。",
              ["true_a_shape", "feature_gain_only"])

    # The original 12x5.9 draft is retained. Two 12x5.1 exports keep all labels
    # 12 pt before embedding: 12 * (4.72/5.1) > 11 pt in the slide frame.
    synw_amplitude_cases = [("true_a_amplitude", "真驱动幅值×2"),
                            ("feature_gain_only", "仅EEG增益×1.8"),
                            ("Hb_gain_only", "仅Hb增益×1.7")]
    synw_amplitude_caption = "真驱动×2使ln(RMS)改变ln2≈0.693；两种读出增益干预的真幅值均不变。图用实际α与HbO/HbR输入展示变化，再比较控制/干预的真log幅值及两个独立编码器的95%边际预测区间。EEG gain只乘有效log-power共同读出项，Hb gain只乘物理Hb均值响应；都不是人体绝对增益恢复。"
    synw_amplitude_keys = ("syn_wave_03a_amplitude_inputs_wide", "syn_wave_03b_amplitude_intervals_wide")
    if all(key in synw_prior_figures and Path(synw_prior_figures[key]["path"]).exists()
           for key in synw_amplitude_keys):
        # Approved amplitude exports are reused byte-for-byte, without redrawing.
        synw_figures.extend(synw_prior_figures[key] for key in synw_amplitude_keys)
    else:
        with plt.rc_context({"font.size": 12, "axes.titlesize": 12, "axes.labelsize": 12,
                             "xtick.labelsize": 12, "ytick.labelsize": 12, "legend.fontsize": 12}):
            fig, axes = synw_figure("反事实幅值①：三种干预的实际EEG与Hb输入", identity_footer=True)
            fig.subplots_adjust(left=.085, top=.86, bottom=.18, hspace=.65, wspace=.35)
            fig.texts[-1].set_text("预定 evaluation/0  ·  冻结训练种子101  ·  单个案（非汇总）；灰底=前20点参考段")
            fig.texts[-1].set_fontsize(12)
            for col, (scenario, name) in enumerate(synw_amplitude_cases):
                control, intervention = synw_pair(scenario)
                synw_observed_pair(axes[0, col], control, intervention, "eeg", name+"：EEG α", channels=[(2, "α")])
                synw_observed_pair(axes[1, col], control, intervention, "hb", "实际HbO / HbR输入")
            synw_save(fig, "03a_amplitude_inputs_wide", "幅值与增益：实际输入波形",
                      synw_amplitude_caption, [case[0] for case in synw_amplitude_cases])

            fig, axes = synw_figure("反事实幅值②：真log幅值与两个独立编码器的95%区间", nrows=1, identity_footer=True)
            fig.subplots_adjust(left=.067, top=.71, bottom=.16, wspace=.35)
            fig.texts[-1].set_text("预定 evaluation/0  ·  冻结训练种子101  ·  单个案（非汇总）；95%边际预测区间")
            fig.texts[-1].set_fontsize(12)
            for col, (scenario, name) in enumerate(synw_amplitude_cases):
                control, intervention = synw_pair(scenario)
                synw_amplitudes(axes[0, col], [control, intervention], name)
                if col == 0:
                    synw_handles, synw_labels = axes[0, col].get_legend_handles_labels()
                axes[0, col].get_legend().remove()
            fig.legend(synw_handles, synw_labels, loc="upper center", bbox_to_anchor=(.53, .86),
                       ncol=3, frameon=False, fontsize=12)
            synw_save(fig, "03b_amplitude_intervals_wide", "幅值与增益：语义输出区间",
                      synw_amplitude_caption, [case[0] for case in synw_amplitude_cases])

    control, intervention = synw_pair("exact_joint_alias")
    alias = synw_all_case_stats["exact_joint_alias"]
    if (not all(alias["bitwise_identical_inputs"].values())
            or not all(alias["bitwise_identical_input_masks"].values())
            or not all(value["all_exported_outputs_bitwise_identical"] for value in alias["modalities"].values())):
        raise ValueError("exact joint alias illustration violated identical input/output contract")
    fig, axes = synw_figure("精确联合 alias：同一EEG/Hb观测，对应两个不同真实幅值")
    synw_observed_pair(axes[0, 0], control, intervention, "eeg", "EEG输入逐位相同：重叠曲线")
    synw_observed_pair(axes[0, 1], control, intervention, "hb", "Hb输入逐位相同：重叠曲线")
    axes[0, 2].plot(synw_time, synw_data["a"][control], color=synw_colors["control"], label="控制真a")
    axes[0, 2].plot(synw_time, synw_data["a"][intervention], color=synw_colors["intervention"], label="干预真a=2×控制")
    synw_style(axes[0, 2], "生成真值不同：a(t) 幅值×2", "规定的合成驱动单位")
    axes[0, 2].legend(frameon=False)
    for col, modality in enumerate(("eeg", "hb")):
        for row, role, color, ls in ((control, "控制输出", synw_colors["control"], "-"), (intervention, "干预输出", synw_colors["intervention"], "--")):
            axes[1, col].plot(synw_time, synw_pred[modality]["shape_mean"][row], color=color, ls=ls, label=role)
        synw_style(axes[1, col], f"{modality.upper()} shape输出逐位相同", "单位RMS形状")
        axes[1, col].legend(frameon=False)
    synw_amplitudes(axes[1, 2], [control, intervention], "真幅值不同，预测分布完全相同")
    synw_save(fig, "04_exact_joint_alias", "完全相同输入的不可辨识反例",
              "冻结生成器显式构造a×2、EEG共同读出增益×1/2及Hb观测项精确补偿。全30个EEG通道和双Hb通道输入差MAE都为0，mask相同，两分支全部导出输出逐位相同，而真log幅值差为ln2。这是确定性的不可辨识构造；它不能通过网络容量解决，也不能因不符合普通噪声分布而从评价分母删除。该个案的区间是否覆盖两个幅值仅为局部演示。",
              ["exact_joint_alias"])

    control, intervention = synw_pair("non_dct_impulse")
    fig, axes = synw_figure("分布外驱动：慢DCT训练族之外的短脉冲")
    synw_eeg_bands(axes[0, 0], intervention, "短脉冲干预后的实际EEG频带输入")
    synw_observed_hb(axes[0, 1], intervention, "脉冲经血流动力学后的实际Hb")
    synw_observed_hb(axes[0, 2], intervention, "完整Hb重建与实际输入", reconstruction=True)
    synw_interval(axes[1, 0], intervention, "eeg", "EEG-only：尖峰真shape与预测")
    synw_interval(axes[1, 1], intervention, "hb", "Hb-only：尖峰真shape与预测")
    synw_amplitudes(axes[1, 2], [control, intervention], "参考RMS幅值固定：控制与脉冲")
    synw_save(fig, "05_ood_impulse", "分布外脉冲暴露形状失败",
              "保留的non_dct_impulse干预在索引29/43/64/83加入短脉冲，真源仍在token聚合前定义且RMS幅值保持固定。模型只有离线卷积上下文，训练source族为慢DCT脉冲；Hb动力学会平滑尖峰。完整观测重建较平滑或贴近输入，不保证尖峰shape恢复。阴影仅为selection校准后的点边际预测区间；本图的120点覆盖比例不能代表整个OOD场景或独立样本覆盖率。",
              ["non_dct_impulse"])

    control, intervention = synw_pair("tau_change")
    tau0 = synw_metadata[control]["parameters"]["free"]["tau"]
    tau1 = synw_metadata[intervention]["parameters"]["free"]["tau"]
    kappa = synw_metadata[control]["parameters"]["free"]["kappa"]
    fig, axes = synw_figure(f"响应参数改变：τ={tau0:.2f}→{tau1:.2f}s，κ={kappa:.3f}固定，真驱动不变")
    synw_observed_pair(axes[0, 0], control, intervention, "eeg", "EEG输入保持完全相同")
    for col, (channel, name) in enumerate(((0, "HbO"), (1, "HbR")), start=1):
        for row, role, color in ((control, "控制", synw_colors["control"]), (intervention, "τ变化", synw_colors["intervention"])):
            axes[0, col].plot(synw_time, synw_data["hb"][row, :, channel], color=color, lw=1.3, label=f"{role}观测")
            axes[0, col].plot(synw_time, synw_pred["hb"]["reconstruction_observation_coordinate"][row, :, channel],
                              color=color, ls="--", lw=1, label=f"{role}完整重建")
        synw_style(axes[0, col], name+"：响应形状改变", "合成Hb观测坐标")
        axes[0, col].legend(frameon=False, ncol=2)
    synw_delta(axes[1, 0], control, intervention, "语义真Δshape=0，估计是否跟着变？")
    synw_amplitudes(axes[1, 1], [control, intervention], "真幅值固定：估计分布的移动")
    for channel, name, color in ((0, "HbO", "#d45050"), (1, "HbR", "#0084a4")):
        delta_physical = synw_metadata[control]["hb_coordinate_gain"]*(synw_truth["physical_hb"][intervention, :, channel]
                                                                                    -synw_truth["physical_hb"][control, :, channel])
        axes[1, 2].plot(synw_time, delta_physical, color=color, label=name+"真物理响应变化")
        axes[1, 2].plot(synw_time, synw_data["hb"][intervention, :, channel]-synw_data["hb"][control, :, channel],
                       color=color, ls="--", label=name+"观测变化")
    synw_style(axes[1, 2], "生成器定位：变化由τ响应产生", "干预 − 控制")
    axes[1, 2].legend(frameon=False, ncol=2)
    synw_save(fig, "06_tau_response", "动力学变化会被归入驱动吗？",
              f"保留tau_change场景使τ从{tau0:.4f}s乘1.6到{tau1:.4f}s，κ={kappa:.4f}在该对中固定；本轮没有独立κ-change评价场景。真驱动、EEG、共同Hb项及配对噪声保持不变。右下由生成器真物理响应验证Hb差异来源，左下和中下显示编码器语义输出是否随动力学变化移动。网络没有显式τ/κ估计头，因此语义改变不能被解释成已分离恢复的血管参数。",
              ["tau_change"])

    control, intervention = synw_pair("correlated_Hb_component")
    fig, axes = synw_figure("相关共同Hb项：生成成分可定位，网络潜变量没有可加物理语义")
    for col, (channel, name) in enumerate(((0, "HbO"), (1, "HbR"))):
        gain = synw_metadata[intervention]["hb_coordinate_gain"]
        axes[0, col].plot(synw_time, gain*synw_truth["physical_hb"][intervention, :, channel],
                          color=synw_colors["truth"], lw=1.25, label="真物理响应×Hb gain")
        axes[0, col].plot(synw_time, synw_truth["observation_component"][intervention, :, channel],
                          color=synw_colors["component"], lw=1.25, label="真附加共同Hb项")
        axes[0, col].plot(synw_time, synw_truth["hb_noise"][intervention, :, channel],
                          color=synw_colors["noise"], lw=.7, label="配对噪声")
        axes[0, col].plot(synw_time, synw_data["hb"][intervention, :, channel],
                          color=synw_colors["hb"], ls="--", lw=1.1, label="三者相加=观测")
        synw_style(axes[0, col], name+"：生成器已知的可加真成分", "合成Hb观测坐标")
        axes[0, col].legend(frameon=False, ncol=2)
    synw_observed_hb(axes[0, 2], intervention, "网络：完整重建 D(S均值,O)", reconstruction=True)
    synw_delta(axes[1, 0], control, intervention, "真驱动不变：Hb语义是否误归属？")
    synw_amplitudes(axes[1, 1], [control, intervention], "真幅值不变：log幅值估计与95%区间")
    direction = np.asarray(synw_metadata[intervention]["component_direction"])
    component = synw_truth["observation_component"][intervention] @ direction
    component /= synw_metadata[intervention]["component_amplitude"]
    source = synw_data["source_shape"][intervention]
    corr = float(np.corrcoef(component, source)[0, 1])
    reference_rms_cosine = float(np.mean(component*source)/np.sqrt(np.mean(component**2)*np.mean(source**2)))
    if not np.isclose(reference_rms_cosine, .85, atol=1e-6):
        raise ValueError("correlated-Hb truth does not reproduce the reference-RMS inner product")
    axes[1, 2].plot(synw_time, source, color=synw_colors["truth"], lw=1.35, label="真驱动单位RMS shape")
    axes[1, 2].plot(synw_time, component, color=synw_colors["component"], lw=1.25, label="沿真loading的归一化共同Hb项")
    synw_style(axes[1, 2], "真驱动与共同项：参考RMS内积=0.85", "规定的归一化真成分")
    axes[1, 2].legend(frameon=False)
    synw_all_case_stats["correlated_Hb_component"]["true_component_source_Pearson_correlation"] = corr
    synw_all_case_stats["correlated_Hb_component"]["true_component_source_reference_RMS_inner_product"] = reference_rms_cosine
    synw_save(fig, "07_correlated_hb", "共同Hb项的语义归属边界",
              "干预把共同Hb项幅值增加0.008，并将其与规定驱动的参考RMS归一化内积设为0.85；该合同量不是全窗Pearson相关（本例Pearson为0.883）。真a、EEG与血管物理响应未变。上排左两图是生成器真值提供的物理响应、附加Hb项与噪声，三者在观测坐标闭合。上排右图是网络完整重建；非线性decoder拼接S均值与O代码，所以semantic decode + observation latent不是两条可相加的物理Hb成分。右下相关只是已知生成构造，不足以在实测中唯一识别来源；语义输出移动仅显示该个案的归属敏感性。",
              ["correlated_Hb_component"])

    synw_manifest = {
        "schema": "semantic_response_synthetic_waveform_report_v1", "report_date": "2026-10-03",
        "execution": "completed_read_only_frozen_inference", "figures": synw_figures,
        "source": synw_sources, "frozen_inference": synw_audit,
        "case_selection": {"rule": "first_base_in_retained_evaluation_order_before_scores",
                           "base_id": synw_audit["base_id"], "scenarios": scenarios,
                           "score_based_selection": False, "all_counterfactual_variants_same_base": True},
        "case_stats": synw_all_case_stats,
        "uncertainty_contract": {"continuous": "selection_calibrated_Gaussian_predictive_point_marginals",
                                 "level": .95, "joint_trajectory_band": False,
                                 "counterfactual_difference_intervals": False,
                                 "case_coverage_interpretation": "120_dependent_points_in_one_preselected_case_not_population_coverage"},
        "model_observation_contract": {"decoder": "D(concat(own_source_mean, own_observation_code))",
                                       "semantic_reconstruction_gradient": "detached_for_semantically_supervised_arms",
                                       "reconstruction_target": "complete_own_modality_observation_in_train_normalized_coordinate",
                                       "physical_additive_semantic_observation_decomposition": False,
                                       "displayed_reconstruction": "inverse_transform_of_same_frozen_training_normalizer",
                                       "generative_Hb_decomposition": "truth_Hb_gain_times_physical_response_plus_processed_common_component_plus_paired_noise"},
        "scientific_boundary": "generator_conditional_synthetic_cases_no_measured_or_unique_physiological_source_qualification",
        "writes_to_evidence_run": False, "fitting_or_training_performed": False,
        "raster_export": {"format": "PNG", "dpi": 220, "normal_figure_inches": [12, 5.1],
                          "amplitude_figure_inches": [12, 5.1], "amplitude_label_font_pt": 12,
                          "amplitude_label_font_pt_in_4_72_inch_slide_frame": 12*4.72/5.1,
                          "wide_v2_minimum_font_pt": 12,
                          "wide_v2_minimum_font_pt_in_4_72_inch_slide_frame": 12*4.72/5.1,
                          "wide_v2_source_annotation": "notes_and_manifest_only",
                          "wide_v2_legend": "figure_level_shared_top"},
        "retained_previous_draft_figures": [str(figures_dir / f"syn_wave_{key}.png") for key in (
            "01_baseline", "02_shape_vs_eeg_gain", "03_amplitude_vs_gains", "04_exact_joint_alias",
            "05_ood_impulse", "06_tau_response", "07_correlated_hb")],
    }
    synw_manifest_path.write_text(json.dumps(synw_manifest, ensure_ascii=False, indent=2, allow_nan=False)+"\n")
    plt.rcParams.update(synw_previous_rc)
    return synw_manifest

def build_response_method_figures(repo, run, out):
    """Diagrams of the implemented contract; no signals or model fitting."""
    import json
    from pathlib import Path
    import numpy as np
    import matplotlib.pyplot as plt
    from matplotlib import font_manager
    from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
    out = Path(out)
    folder = out / 'figures'
    folder.mkdir(parents=True, exist_ok=True)
    font = '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'
    font_manager.fontManager.addfont(font)
    plt.rcParams.update({'font.family': font_manager.FontProperties(fname=font).get_name(),
                         'font.size': 13, 'axes.unicode_minus': False})
    navy, teal, amber, red = '#15364A', '#008C88', '#D18B30', '#B94758'
    figures = []
    def canvas():
        fig, ax = plt.subplots(figsize=(12, 4.8))
        ax.set(xlim=(0, 12), ylim=(0, 4.8))
        ax.axis('off')
        return fig, ax
    def box(ax, x, y, w, h, text, color=navy, size=14):
        ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle='round,pad=.07,rounding_size=.10',
                                   facecolor=color, edgecolor='none'))
        ax.text(x+w/2, y+h/2, text, ha='center', va='center', color='white', fontsize=size,
                linespacing=1.55)
    def arrow(ax, p, q, label=None, color=navy):
        ax.add_patch(FancyArrowPatch(p, q, arrowstyle='-|>', mutation_scale=17, lw=1.8, color=color))
        if label:
            ax.text((p[0]+q[0])/2, (p[1]+q[1])/2+.18, label, ha='center', color=color, fontsize=11)
    def save(fig, key, title, caption, source):
        path = folder / (key+'.png')
        fig.savefig(path, dpi=240, bbox_inches='tight', facecolor='white')
        plt.close(fig)
        figures.append(dict(key=key, path=str(path.resolve()), title=title, caption=caption, source=source))

    fig, ax = canvas()
    for x, title, lines, color in [(0.25, '讨论与可证伪假设', '谱信息还是输入衰减？\n局部 EEG 还是共享 Hb？\n形状还是幅值／观测增益？', navy),
            (4.3, '冻结的实验对照', '匹配衰减 + 独立增益\n空间 / 历史基线 + null\n反事实 + alias + 容量控制', teal),
            (8.35, '分层解释结果', '条件形状可恢复\n幅值与跨机制校准失败\n实测 S 未提供增量', amber)]:
        box(ax, x, 2.65, 3.4, .65, title, color)
        ax.text(x+1.7, 1.65, lines, ha='center', va='center', fontsize=15, linespacing=1.9, color=navy)
    arrow(ax, (3.75, 2.98), (4.16, 2.98))
    arrow(ax, (7.8, 2.98), (8.2, 2.98))
    ax.text(6, .35, '完成计算 ≠ 通过机制检验；重建改善 ≠ 识别共享生理来源', ha='center', color=red, fontsize=16)
    save(fig, 'method_evidence_chain', '讨论如何转化为实验，而不是直接转化为结论',
         '本报告把讨论、具体实现、对照端点和科学判读逐一对应；历史结果只作为动机。',
         'docs/METHOD_RATIONALE.md; docs/EXPERIMENT_PLAN.md; frozen run summary.json')

    fig, ax = canvas()
    box(ax, .2, 3.45, 2.15, .75, '原生 EEG 电压\n通道 / 时钟 / 单位', navy, 13)
    box(ax, 3.0, 3.45, 2.25, .75, '频带功率与参考\n6 通道 × 5 频带', teal, 13)
    box(ax, 6.0, 3.45, 2.2, .75, 'prepared EEG\n120 × 30', teal, 13)
    box(ax, 9.1, 3.45, 2.45, .75, '固定谱模式 a、b\n与未解释残差', amber, 13)
    box(ax, .2, 1.65, 2.15, .75, '原生光强 / 发布 Hb\n数据集定义不同', navy, 13)
    box(ax, 3.0, 1.65, 2.25, .75, '声明的转换 / 处理\n时钟对齐与参考', teal, 13)
    box(ax, 6.0, 1.65, 2.2, .75, 'prepared Hb\n120 × 2', teal, 13)
    box(ax, 9.1, 1.65, 2.45, .75, '响应项 + 共同项\n+ 模型残差', amber, 13)
    for y in (3.825, 2.025):
        for p, q in [(2.4, 2.9), (5.35, 5.9), (8.3, 9.0)]: arrow(ax, (p, y), (q, y))
    ax.text(6, .45, '反变换只回到 parent prepared 坐标；不能恢复被预处理移除的原始信息', ha='center', color=red, fontsize=15)
    save(fig, 'method_raw_lineage', '成分分离必须回到同一窗口与同一观测坐标',
         'EEG 的模式贡献可在 log-power 特征中相加，不能直接与高频原始电压相加。',
         'docs/DATA_CONTRACT.md; parent prepared record identities; unified_physiology.py')

    fig, ax = canvas()
    for y, name, color in [(3.65, '目标 Hb：前缀 40', teal), (2.45, '目标 Hb：前缀 0', amber), (1.25, '完整 EEG / 其他区域 Hb', navy)]:
        ax.text(.15, y+.22, name, ha='left', va='center', color=navy, fontsize=13)
        x0, width = 3.25, 8.3
        ax.add_patch(plt.Rectangle((x0, y), width, .45, color='#E8EDF1'))
        if y == 3.65: ax.add_patch(plt.Rectangle((x0, y), width/3, .45, color=teal))
        if y == 1.25: ax.add_patch(plt.Rectangle((x0, y), width, .45, color=navy))
        if y != 1.25:
            ax.add_patch(plt.Rectangle((x0+2*width/3, y), width/3, .45, fill=False, hatch='///', edgecolor=red, lw=1.2))
        for sec in (0, 10, 20, 30):
            x = x0+width*sec/30
            ax.text(x, y-.18, str(sec)+' s', ha='center', va='top', color=navy, fontsize=11)
    ax.text(7.4, .35, '绿色：可用目标前缀     灰色：不可输入目标 Hb     斜线：隐藏评分区 20–30 s', ha='center', fontsize=12)
    save(fig, 'method_visibility', '同一时间轴上区分可见、隐藏和评分支持',
         '前缀 40 = 10 s；零前缀不用目标 Hb token。parent 非因果处理与全窗 EEG 仍使端点属于离线补全。',
         'semantic_response_discovery_v1.yaml: measured; public_encoding_view; public_raw_context')

    fig, ax = canvas()
    box(ax, .2, 2.4, 2.1, 1, 'EEG-only\n固定 a(t)、b(t)', navy)
    box(ax, 3.05, 3.4, 2.5, .8, '谱混合 / 匹配衰减\n或独立增益 g·a', teal, 13)
    box(ax, 3.05, 1.4, 2.5, .8, '初态与共同项\n仅由可见 Hb 前缀校准', amber, 13)
    box(ax, 6.25, 3.4, 2.35, .8, '固定 H0 响应核\n单独增加 τn 或 τv', teal, 13)
    box(ax, 9.3, 2.4, 2.35, 1, '隐藏 Hb 补全\n及 proper-score', navy, 13)
    arrow(ax, (2.3, 3.0), (2.94, 3.8))
    arrow(ax, (5.55, 3.8), (6.12, 3.8))
    arrow(ax, (8.65, 3.8), (9.22, 3.08))
    arrow(ax, (5.6, 1.8), (9.18, 2.65))
    ax.text(6.1, .40, 'Hb 后缀不反向重估驱动；零前缀使用初态先验与零共同项', ha='center', fontsize=15, color=red)
    save(fig, 'method_ssm_inference', '新推断制度：先定 EEG 驱动，再用 Hb 前缀校准',
         '与历史全窗联合拟合的可见信息不同，因此历史重建数字不能直接当作本轮同条件基线。',
         'src/inference/semantic_response_ssm.py; frozen evaluation runner')

    fig, ax = canvas()
    for y, label in [(3.55, 'EEG + mask'), (1.15, 'Hb + mask')]:
        box(ax, .1, y, 1.9, .8, label, navy, 14)
        box(ax, 2.65, y+.20, 2.3, .65, '语义编码器 / 头 S\n形状、log幅值、SD', teal, 12)
        box(ax, 2.65, y-.65, 2.3, .65, '独立观测编码器 → O', navy, 13)
        box(ax, 6.3, y-.2, 2.4, .9, '非线性本模态解码器\n接收 O 与 detach(S)', amber, 12)
        box(ax, 9.35, y-.2, 2.35, .9, '完整本模态\n观测重建输出', teal, 13)
        arrow(ax, (2, y+.4), (2.56, y+.52))
        arrow(ax, (2, y+.25), (2.56, y-.30))
        arrow(ax, (4.97, y+.52), (6.2, y+.30), 'detach')
        arrow(ax, (4.97, y-.30), (6.2, y+.03))
        arrow(ax, (8.75, y+.25), (9.24, y+.25))
    ax.text(6, .12, '推理各读本模态；监督臂重建梯度不改写 S；O-only 对照让全部均值坐标学习重建', ha='center', fontsize=12, color=red)
    save(fig, 'method_tokenizer', 'semantic–observation 原型：输入独立，输出职责明确',
         '解码器是非线性联合解码；S 与 O 不是可直接相加的“神经波形”和“噪声波形”。',
         'src/tokenizers/semantic_response_tokenizer.py: _ModalityBranch / response_loss')

    fig, ax = canvas()
    box(ax, .1, 2.5, 3.2, 1.4, '所有方法相同的上下文\nprepared 局部 EEG\n其他区域 Hb + 自身前缀', navy, 14)
    for y, title, color in [(4.0, '仅共同上下文', navy), (3.0, '+ 同模型 O', teal),
                            (2.0, '+ 同模型 O + S', amber), (1.0, '+ 独立等容量纯 O', teal)]:
        box(ax, 4.1, y-.15, 3.2, .60, title, color, 13)
        arrow(ax, (3.33, 3.15), (4.0, y+.15))
        arrow(ax, (7.38, y+.15), (8.14, 2.65))
    box(ax, 8.25, 2.15, 3.45, 1.05, '训练集拟合 ridge\n选择集定 α / 预测残差 SD\n评价集只评分', navy, 13)
    ax.text(6.0, .2, '编码器冻结；O→O+S 检验语义增量，等容量纯 O 检验监督是否优于容量收益', ha='center', fontsize=13, color=red)
    save(fig, 'method_public_probe', '实测探针同时控制可见上下文和表征总容量',
         'context_raw 只是“未编码的 prepared 上下文”，并不是原生 EEG 电压或原始光强。',
         'train_semantic_response_tokenizer.py: probe / fit_public_probe; public_probes/*/selection.json')
    result = dict(figures=figures, schema='semantic_response_method_diagrams_v1')
    (out/'method_diagram_provenance.json').write_text(json.dumps(result, ensure_ascii=False, indent=2)+'\n')
    return result


def render_semantic_response_report(run, out, draft_name='SSM_SEMANTIC_RESPONSE_REPORT.pptx'):
    """Detailed, editable Chinese report of retained evidence and frozen replay."""
    import json
    import math
    import hashlib
    import sys
    from pathlib import Path
    import numpy as np
    import pandas as pd
    from PIL import Image, ImageFont
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR
    from pptx.enum.shapes import MSO_SHAPE
    run, out = Path(run).resolve(), Path(out).resolve()
    repo = Path(__file__).resolve().parents[2]
    if str(repo) not in sys.path:
        sys.path.insert(0, str(repo))
    if (out/draft_name).exists():
        raise ValueError('Preserve an existing presentation; choose a fresh versioned output or draft name')
    summary = json.loads((run/'summary.json').read_text())
    verification = json.loads((run/'verification.json').read_text())
    if summary['execution'] != 'completed' or not verification['passed']:
        raise ValueError('Report needs completed and verified owning run evidence')
    out.mkdir(parents=True, exist_ok=True)
    cfg = json.loads((run/'resolved_config.json').read_text())
    metadata = json.loads((run/'metadata_plan_proposal.json').read_text())
    tok = json.loads((run/'tokenizer/summary.json').read_text())
    manifests = {}
    for name, builder in [('quantitative_provenance.json', build_response_quantitative_figures),
                          ('waveform_provenance.json', build_response_waveform_figures),
                          ('synthetic_waveform_provenance.json', build_synthetic_waveform_figures)]:
        path = out/name
        manifests[name] = json.loads(path.read_text()) if path.exists() else builder(repo, run, out)
    methods = build_response_method_figures(repo, run, out)
    qfigs = manifests['quantitative_provenance.json']['figures']
    wfigs = manifests['waveform_provenance.json']['figures']
    sfigs = manifests['synthetic_waveform_provenance.json']['figures']
    mfigs = {f['key']: f for f in methods['figures']}
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    font = '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'
    navy, teal, red, muted = '15364A', '008C88', 'B94758', '526C7D'
    slide_sources, shape_audit = [], []

    def wrap_count(value, width, size):
        pil = ImageFont.truetype(font, round(size*4))
        count = 0
        for line in str(value).split('\n'):
            if not line:
                count += 1
                continue
            current = ''
            count += 1
            for char in line:
                if current and pil.getlength(current+char)/4 > width*72:
                    count += 1
                    current = char
                else:
                    current += char
        return count

    def text(slide, value, x, y, w, h, size=22, color=navy, bold=False, minimum=14):
        size = float(size)
        while size > minimum and wrap_count(value, w, size)*size*1.17 > h*72:
            size -= .5
        if wrap_count(value, w, size)*size*1.1 > h*72+3:
            raise ValueError(f'Text does not fit: {str(value)[:80]}')
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for idx, value_line in enumerate(str(value).split('\n')):
            para = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            para.text, para.line_spacing = value_line, 1.12
            para.space_before = para.space_after = Pt(0)
            for rr in para.runs:
                rr.font.name = 'Noto Sans CJK SC'
                rr.font.size, rr.font.bold = Pt(size), bold
                rr.font.color.rgb = RGBColor.from_string(color)
                if value_line.startswith(('https://', 'http://')):
                    rr.hyperlink.address = value_line
        shape_audit.append(dict(slide=len(deck.slides), text=str(value), x=x, y=y, width=w, height=h, font_pt=size))
        return shape

    def new(title, subtitle='', source='', chapter=''):
        s = deck.slides.add_slide(deck.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = RGBColor.from_string('F7F9FB')
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(.10))
        bar.fill.solid(); bar.fill.fore_color.rgb = RGBColor.from_string(teal); bar.line.fill.background()
        text(s, title, .55, .28, 12.22, .57, 28, bold=True, minimum=22)
        text(s, subtitle, .58, .99, 12.15, .46, 15, color=muted, minimum=12)
        text(s, '2026.10.03  |  响应动力学与连续语义  |  '+chapter, .58, 7.14, 11.7, .20, 10, color=muted, minimum=10)
        text(s, f'{len(deck.slides):02}', 12.2, 7.13, .5, .22, 10, color=muted, minimum=10)
        evidence = source if isinstance(source, str) else json.dumps(source, ensure_ascii=False)
        notes = (f'实验事实源：{run}\n证据：{evidence}\n'
                 '公开开发端点；此前cohort已分析。前缀和评分均位于parent非因果处理后的30秒窗口。'
                 '优化种子共享评价基础身份，不能当独立被试或独立置信区间。\n')
        s.notes_slide.notes_text_frame.text = notes
        slide_sources.append(dict(slide=len(deck.slides), title=title, chapter=chapter, evidence=source))
        return s

    def caption(s, value):
        text(s, value, .65, 6.48, 12.02, .52, 13, color=muted, minimum=11)

    def body(title, subtitle, paragraphs, source, chapter='', takeaway=None):
        s = new(title, subtitle, source, chapter)
        if isinstance(paragraphs, str): paragraphs = [paragraphs]
        count = len(paragraphs)
        height = 4.55/max(count, 1)
        for i, paragraph in enumerate(paragraphs):
            text(s, paragraph, .8, 1.72+i*height, 11.8, height-.16, 23 if count<=4 else 20,
                 color=navy, minimum=16)
        if takeaway: caption(s, takeaway)
        return s

    def table_slide(title, subtitle, headers, rows, widths, source, chapter='', takeaway=None, size=18):
        s = new(title, subtitle, source, chapter)
        # Size each row from the longest actual wrapped cell, not a fixed guess.
        font_size = float(size)
        while True:
            heights = [max(.43, max(wrap_count(value, width-.20, font_size) for value, width in zip(row, widths))*font_size*1.17/72+.12)
                       for row in [headers]+rows]
            if sum(heights) <= 4.67 or font_size <= 13:
                break
            font_size -= .5
        if sum(heights) > 4.8:
            raise ValueError('Table does not fit: '+title)
        sh = s.shapes.add_table(len(rows)+1, len(headers), Inches(.65), Inches(1.68), Inches(12), Inches(sum(heights)))
        tab = sh.table
        for col, width in zip(tab.columns, widths): col.width = Inches(width)
        for i, row in enumerate([headers]+rows):
            tab.rows[i].height = Inches(heights[i])
            for j, value in enumerate(row):
                c = tab.cell(i,j)
                c.text = str(value)
                c.margin_left, c.margin_right = Inches(.1), Inches(.07)
                c.margin_top, c.margin_bottom = Inches(.04), Inches(.04)
                c.vertical_anchor = MSO_ANCHOR.MIDDLE
                c.fill.solid(); c.fill.fore_color.rgb = RGBColor.from_string(navy if i==0 else ('FFFFFF' if i%2 else 'E9F0F4'))
                for p in c.text_frame.paragraphs:
                    p.space_before = p.space_after = Pt(0); p.line_spacing = 1.1
                    for rr in p.runs:
                        rr.font.name = 'Noto Sans CJK SC'; rr.font.size = Pt(font_size); rr.font.bold = i==0
                        rr.font.color.rgb = RGBColor.from_string('FFFFFF' if i==0 else navy)
        if takeaway: caption(s, takeaway)
        return s

    def fig_slide(f, chapter='', subtitle=None):
        key = f.get('key', Path(f['path']).stem)
        figure_title = f['title']
        if key.startswith('waveform_') and '：' in figure_title:
            figure_title, case_identity = figure_title.split('：', 1)
            subtitle = case_identity + ' · ' + ('最差差值案例' if '_worst_' in key else '中位差值案例')
        s = new(figure_title, subtitle or f.get('subtitle', '读取保留证据；图形为位图，结果与来源可追溯'), f.get('source', ''), chapter)
        p = Path(f['path'])
        if not p.is_absolute():
            p = p if p.exists() else out/p
        with Image.open(p) as im: iw, ih = im.size
        width = min(12.0, 4.72*iw/ih)
        height = width*ih/iw
        s.shapes.add_picture(str(p), Inches((13.333-width)/2), Inches(1.60+(4.72-height)/2),
                             width=Inches(width), height=Inches(height))
        concise = {
            'quant_10_tokenizer_control_heatmaps': '每格为256个基础身份、三个优化种子的控制变体形状RMSE均值。此页比较三种端到端语义原型；O-only监督线性读出另页报告。',
            'quant_11_tokenizer_intervention_heatmaps': '干预变体与控制图共用色标。未见脉冲等机制使形状恢复显著恶化；三种语义原型及O-only线性读出的对照不能据单一基准泛化。',
            'quant_19_observation_linear_probe': 'O-only采用训练真值监督的线性读出。基准EEG/Hb形状RMSE为1.0541/0.1288，分布头为0.0433/0.0971；无耦合Hb反而是O-only更低（0.4224 vs 0.7512）。',
            'syn_wave_01_baseline': '首个评价基础轨迹、seed101：各模态独立恢复生成器条件形状。灰带为模型区间；本页是单例，覆盖结论以全256基础身份汇总为准。',
            'syn_wave_02_shape_vs_eeg_gain': '真实形状变化应被保留；仅改变EEG观测增益却也改变估计驱动。对比同一基础轨迹，不能把点估计的变化自动解释为真实生理变化。',
            'syn_wave_03_amplitude_vs_gains': '真幅值翻倍的ΔlogA=0.693，EEG/Hb估计仅约0.440/0.520；纯观测增益分别引起约0.375/0.367的假变化。本页为固定单例。',
            'syn_wave_04_exact_joint_alias': '两个模态的输入、mask及模型输出逐位相同，真log幅值相差0.693；不能仅从这些观测辨认两种真值。细窄区间反映生成假设的限制。',
            'syn_wave_05_ood_impulse': '未见非DCT脉冲使预测形状明显偏离真值，区间并未充分扩张。单例可视化与前面的全场景误差、覆盖统计分别解释。',
            'syn_wave_06_tau_response': '本例只改变τ，κ固定；输入EEG不变，Hb响应形状改变。图中响应真成分可加，但网络O/S由非线性解码器联合重建，不具有相同可加定义。',
            'syn_wave_07_correlated_hb': '生成器共同Hb项与驱动具有0.85的参考RMS归一化内积（并非Pearson相关）；它改变Hb推断的形状与幅值，说明来源仍会混淆。',
        }
        short_caption = next((value for prefix,value in concise.items() if key.startswith(prefix)), f.get('caption', ''))
        caption(s, short_caption)
        slide_sources[-1].update(figure=str(p.resolve()), figure_key=key, caption=short_caption, statistics=f.get('stats',{}))
        s.notes_slide.notes_text_frame.text += '\n图：'+str(p.resolve())+'\n完整图注：'+f.get('caption','')+'\n统计：'+json.dumps(f.get('stats',{}), ensure_ascii=False)
        return s

    s = new('从原始波形到条件语义—观测分解', 'SSM 改进讨论、冻结实验与完整结果报告 · 2026-10-02 run / 2026-10-03 export',
            ['docs/METHOD_RATIONALE.md','docs/EXPERIMENT_PLAN.md','summary.json','tokenizer/summary.json'], '研究报告')
    text(s, '响应动力学与连续\nsemantic–observation tokenizer', .8, 1.9, 11.8, 1.7, 36, bold=True)
    text(s, '可恢复的条件形状、仍混淆的幅值、未成立的实测语义增量', .85, 4.1, 11.6, .9, 25, color=teal)
    text(s, '三公开数据集  ·  三个同步区域  ·  四臂 × 三种子\n原始波形、模型分解、反事实、不确定性和容量控制逐层核验', .85, 5.3, 11.6, .85, 20, color=muted)
    table_slide('先给出三层结论', '有工程收益；共享生理语义的关键检验尚未通过', ['问题','本轮结果','允许的解释'], [
        ['SSM 机制','谱混合未超过匹配衰减；null 未通过','不能把拟合改善解释为更正确的 EEG→Hb 耦合'],
        ['合成语义','基准形状改善；未见机制和 alias 失效','生成器条件下的函数恢复，未建立通用幅值语义'],
        ['实测表征','O 有收益；加入 S 后误差增加','观测表示有用；当前 S 未提供独立增量']], [2.05,4.7,5.25],
        'summary.json; dataset_equal_comparisons.csv; tokenizer/summary.json', '结论')
    body('报告阅读路径', '围绕“波形怎样被解释、哪些解释尚不能成立”组织证据', [
        '01  讨论：原始叙事、竞争解释与实验取舍',
        '02  数据与模型：原生波形 → prepared 特征 → 条件成分',
        '03  SSM：合成、公开主比较、null 与逐窗分解',
        '04  Tokenizer：独立编码、配对反事实、alias、跨生成机制',
        '05  实测增量：完整观测与容量基线、最终判断和复现'], 'presentation_sources.json', '导读')
    body('项目最初的问题仍然不变', '分别观测 EEG 和 fNIRS，能否获得共享生理过程的有用表示？', [
        '构造一个联合坐标，不等于每个模态都能独立恢复该坐标。',
        '每个模态都可预测一个坐标，也不等于该坐标具有物理充分性。',
        '共享生理语义必须进一步解释反事实变化，并提供超过完整观测基线的可复现增量。'],
        'docs/METHOD_RATIONALE.md: Research question', '讨论', '本轮保持独立模态输入与完整 observation 信息；下游分类性能不是训练或选择端点。')
    table_slide('先前证据为何促使本轮改变设计', '仅作为讨论动机；不同推断制度的历史分数不混入本轮对照', ['先前观察','竞争解释','本轮回应'], [
        ['谱模式改善重建','可能只是减弱错误血管输入','匹配衰减、独立增益和随机方向'],
        ['附加 Hb 项可吸收误差','共同项与初态 / 动力学存在混淆','固定驱动、前缀校准与响应扩展分开'],
        ['空间 Hb 可补全目标','任务形状或共享 Hb 可能已足够','空间基线后再检验局部 EEG 增量'],
        ['typed 目标能被学会','生成先验可能制造语义确定性','幅值 alias、未见生成机制与分布输出']],
        [3.0,4.4,4.6], 'shared_driver_modes/20261002_v1; shared_driver_attribution/20261002_v1; docs/EXPERIMENT_PLAN.md', '讨论', size=17)
    table_slide('讨论覆盖六条改进方向', '科学论证与实现审核分别负责；将建议转化为有限、可检验的改动', ['讨论方向','形成的判断','本轮落实'], [
        ['谱 / 空间语义','有限区域与谱模式可描述，不等于源定位','保留 a、b；加入空间条件基线'],
        ['动力学','输入滞后与静脉响应需竞争检验','分别增加低通 τn、出流松弛 τv'],
        ['初态与归属','自由初态可制造响应而无 EEG 驱动','p=v proper prior；零前缀先验'],
        ['可证伪性','真实配对需优于错配与替代输入','错被试、非环绕时移、其他区域 EEG'],
        ['目标与不确定性','形状与幅值应分开；必须允许多解','单位 RMS 形状、log 幅值和 alias'],
        ['容量与泄漏','多几个坐标并不证明语义有效','O→O+S 和独立等容量 O-only']],
        [2.15,5.2,4.65], '讨论形成的冻结合同：docs/EXPERIMENT_PLAN.md; semantic_response_discovery_v1.yaml', '讨论', size=16)
    fig_slide(mfigs['method_evidence_chain'], '讨论')
    body('哪些建议没有直接进入本轮', '收紧目标和解释范围，避免把尚无锚点的概念写进 token 名称', [
        '未增加细胞类型、E/I、代谢率或精确脑源标签：当前频带特征和稀疏 Hb 区域没有足够来源证据。',
        '未把共同 Hb 项规定为噪声：广泛神经变化也可能进入该坐标。',
        '未同时自由拟合所有动力学：先用两个单独扩展检验有限响应策略。',
        '未启动 VQ、受保护评价或结果驱动追加调参：连续表示先接受语义与增量检验。'],
        'docs/EXPERIMENT_PLAN.md; semantic_response_discovery_v1.yaml', '讨论')

    fig_slide(mfigs['method_raw_lineage'], '数据与波形')
    table_slide('数据分组与实际独立单位', '所有变换按允许的训练分组拟合；评价被试与本轮训练 / 选择互斥',
        ['数据集','训练人数 / 原生窗','选择人数 / 原生窗','评价人数 / 原生窗'],
        [[short]+[str(metadata['counts'][ds][part]['subjects'])+' / '+str(metadata['counts'][ds][part]['native_windows']) for part in ('train','selection','evaluation')]
         for ds,short in [('eeg_fnirs_single_trial','Single-Trial'),('simultaneous_eeg_nirs','Simultaneous'),('visual_cognitive_motivation','Visual')]],
        [3.0,3.0,3.0,3.0], 'metadata_plan_proposal.json', '数据与波形',
        '评价共 15 人、120 原生窗口、360 区域行；三个区域不独立。全 cohort 此前已分析，不能称全新确认队列。')
    body('任务支持与外推范围', '分组互斥并不使小样本任务或已查看队列变为独立确认', [
        'Simultaneous 选择集没有 2-back；评价中的 DSR 仅 5 个原生窗口、2 位被试。',
        '后面的最差 DSR 窗口说明一个明确失败案例，不能代表广泛的 DSR 人群表现。',
        'Visual 只使用有三区域同步支持的 Probe1，并按 EEG Part / 起点去重。',
        '整个 cohort 此前已分析；本轮更换评价身份减少直接适配，仍只属于开发性新端点。'],
        'summary.json: limitations; metadata_plan_proposal.json', '数据边界')
    table_slide('原始单位与处理后坐标不能混用', '图中按各数据集的发布定义标注；未知单位保持未知', ['数据集','原生观测','本轮模型坐标'], [
        ['Single-Trial','EEG 电压；双波长光电强度','转换并处理后的相对 Hb；相对频带 log-power'],
        ['Simultaneous','EEG 电压；发布 HbO/HbR','保留单位来源，经 parent 处理与新训练尺度映射'],
        ['Visual','EDF 电压；发布 Hb（绝对物理单位未核实）','不冒称 μM；仅在声明的 prepared 坐标解释']],
        [2.4,4.8,4.8], 'docs/DATA_CONTRACT.md; docs/DATASETS_DESCRIPTION.md; waveform_provenance.json', '数据与波形',
        '探针表中的 context_raw 指 prepared 特征直接入 ridge；它不是此处展示的原生波形。', size=17)
    # Place a true data example before the mathematical model; remaining cases follow SSM results.
    if wfigs: fig_slide(wfigs[0], '原生波形示例')
    body('模型真正使用的张量与时间参考', '30 秒窗口；4 Hz 特征时间轴；原生 EEG 的采样率另外标注', [
        'EEG：[120, 30]，来自 6 个邻近电极 × 5 个频带（1–4、4–8、8–13、13–30、30–45 Hz）。',
        'Hb：[120, 2]，分别为 HbO、HbR；三个同步区域是前额、运动和后部。',
        '前 20 点 = 5 s 参考段；本轮重拟合训练尺度，不复用旧 outer PCA / SD。',
        'parent 处理和窗口内表示是非因果的；所有后缀成绩称离线处理后特征补全。'],
        'resolved_config.json: tensor/data; metadata_plan_proposal.json', '数据与波形')
    fig_slide(mfigs['method_visibility'], '可见性合同')
    fig_slide(mfigs['method_ssm_inference'], 'SSM 设计')
    table_slide('驱动比较：谱信息和衰减分开检验', '所有候选使用同一 EEG-only a、b 和同一 Hb 可见性', ['臂','输入定义','检验问题'], [
        ['宽频','u = a','最基本的固定谱坐标'],
        ['谱混合','u = (a + γb) / √(1+γ²)','谱对比 b 是否提供血管信息'],
        ['匹配衰减','u = a / √(1+γ²)','去掉 b，保留相同 a 衰减'],
        ['独立增益','u = g·a；含 g=0','单纯改变读出强度是否已足够'],
        ['随机 / 零输入','预定随机方向；u=0','特定谱方向与无 EEG 驱动对照']],
        [2.2,5.0,4.8], 'resolved_config.json: model; model_spec', 'SSM 设计',
        'γ 是候选读出坐标的权重；不能直接解释为已测生理耦合强度。', size=17)
    table_slide('响应扩展：只改变一个候选机制', '固定 parent H0 参数；不让新增自由度自动获得生理身份', ['部分','本轮方程 / 规则','解释限制'], [
        ['输入低通','du/dt = (r−u)/τn；τn=0 为即时输入','检验有限输入滞后 / 平滑策略'],
        ['出流松弛','fout = (τ·v^(1/α)+τv·f)/(τ+τv)','p、q 平衡均使用同一 fout'],
        ['体积响应','dv/dt = (f−v^(1/α))/(τ+τv)','不是只改体积而忽略质量平衡'],
        ['初态','p=v；log 坐标 proper prior','自由初态只作固定消融'],
        ['共同 Hb 项','4 个低频模式；方向 / 尺度按合同','可以补偿误差，来源仍未定']],
        [2.0,5.7,4.3], 'src/inference/semantic_response_ssm.py; Buxton et al. 2004, Eq.11', 'SSM 设计', size=16)
    body('初态与观测项为何必须单独讨论', '模型中存在可产生 Hb 响应、但并不携带 EEG 驱动的自由度', [
        '由同一质量平衡方程可得：d(p−v)/dt = −[fout/(τ·v)]·(p−v)，此处分母 v 为体积状态。',
        '因此 p−v 是衰减暂态；从 p=v 出发会保持该约束，自由 p 初态不代表新增持续过程。',
        '有 Hb 前缀时，初态与共同观测项可形成自由响应；g=0 的改善不能作为 EEG 耦合证据。',
        '零前缀时只用先验初态和零共同项，阻止隐藏目标 Hb 决定这些自由度。'],
        'semantic_response_ssm.py: _physical_response_slope / fit_fixed_driver_response', 'SSM 设计')
    table_slide('基线阶梯：复杂模型必须超过哪些信息', '所有训练 / 选择边界按被试分离；标签只进入命名明确的模板对照', ['基线','可利用的信息','意义'], [
        ['训练均值 / 任务模板','训练波形；任务模板有标签 oracle','共同时间形状可能已经足够'],
        ['自身 Hb 前缀 ridge','目标 Hb 的同一可见前缀','分离血管历史与 EEG 信息'],
        ['完整 EEG ridge','本地 EEG prepared 特征','物理约束是否超过直接回归'],
        ['空间 Hb ridge','其他两个区域 Hb，不含目标 Hb','广泛 Hb 信息是否已解释目标'],
        ['空间 Hb + EEG','即时 EEG / 固定稳定滤波器组','局部 EEG 是否有条件增量']],
        [3.2,4.75,4.05], 'resolved_config.json: baselines; fit_baselines; public_raw_context', 'SSM 设计', size=16)
    table_slide('主 null family 与分母', '40点前缀：谱混合、独立增益、最佳稳定响应 × 三种错配，共九项', ['null','对照支持','所能检验的特异性'], [
        ['任务匹配错训练被试','相同目标与评分时间点','真实跨被试配对是否有优势'],
        ['非环绕移位 12 s','仅对比 real_shift_support','避免循环拼接或不同有效支持'],
        ['同步其他区域 EEG','相同窗口，不换目标 Hb','局部性；不是普遍零耦合证明']],
        [3.2,4.5,4.3], 'primary_null_specificity.csv/json; docs/EXPERIMENT_PLAN.md', 'SSM 设计',
        '先被试内、再数据集内被试等权、最后三数据集等权；整被试符号枚举 2^15，九项统一 Holm。', size=17)
    body('描述区间与 null 检验的条件', '统计显著性、机制识别与独立确认分别解释', [
        '效应区间条件于冻结的训练 / 选择对象；三数据集等权，不按区域数或窗口数增加独立样本量。',
        '符号检验需要零假设下被试差值对称。完整枚举 2^15 种符号配置，不等于无条件随机化检验。',
        '九项 Holm 只覆盖注册的 primary null family；不覆盖 tokenizer、γ主对照或所有后续描述图。',
        '三优化种子重复使用同一评价身份，误差范围表示优化敏感性，不是独立被试置信区间。'],
        'docs/EXPERIMENT_PLAN.md; primary_null_specificity.json; tokenizer/scenario_variant_seed_summary.csv', '统计边界')
    table_slide('分数和不确定性各自回答什么', '误差、概率质量与来源解释分开', ['指标','较好方向','边界'], [
        ['训练 SD 标准化 NRMSE','更低','绝对差值不是百分比改善'],
        ['Gaussian CRPS / NLL','更低','selection 残差 SD，不是生理后验'],
        ['50 / 80 / 95% 覆盖与宽度','覆盖接近名义且区间有信息','宽区间也可能无用；必须分场景'],
        ['反事实响应 / 假变化','跟随真变化、抵抗观测改变','需要已知生成真值，不能由实测重建替代'],
        ['配对 null 与条件增量','真实配对和 S 有额外收益','不由视觉平滑或重建美观决定']],
        [3.3,3.6,5.1], 'resolved_config.json: endpoints/uncertainty; dataset_equal_comparisons.csv', '评价合同', size=16)
    table_slide('SSM 合成设计与选择限制', '冻结选择策略接受检验；没有用评价场景标签重选参数', ['设计','执行内容','解释'], [
        ['六场景','标准、非rest、输入低通、出流松弛、共同Hb、零耦合','受限与随机化两个 family'],
        ['幅值 / 持续时间','3 个幅值、3 个持续时间的平衡变化','不是完整全因子'],
        ['全局选择','所有场景的 control 与 intervention 汇总','重复 control 占一半权重'],
        ['选择结果','两个 family 均 γ=0、g=1、τn=τv=0','三个响应候选实际退化同一核']],
        [2.3,5.5,4.2], 'synthetic_calibration/*.json; synthetic_summary.csv; cross_mechanism_fit_matrix.csv', 'SSM 结果',
        '矩阵不能解释为机制识别准确率，也不能由相同预测直接证明低通与出流机制不可区分。', size=16)

    # Quantitative figures declare a section so report ordering does not depend on filename accidents.
    ssm_quant, token_quant, public_quant = [], [], []
    for f in qfigs:
        ordinal = int(f['key'].split('_')[1])
        if f.get('section') == 'tokenizer':
            token_quant.append(f)
        elif ordinal >= 16:
            public_quant.append(f)
        elif ordinal >= 9:
            token_quant.append(f)
        else: ssm_quant.append(f)
    for f in ssm_quant: fig_slide(f, 'SSM 定量结果')
    body('如何阅读接下来的原始—成分波形图', '展示典型表现与失败边界；案例只是总体结果的解释，不替代总体统计', [
        '从已完成评价结果按固定报告选图规则取两例：best_stable 相对 gain_selected 的逐窗误差差值最接近中位数，以及最差。',
        '原生波形保留通道、单位与真实时钟；prepared 波形显示本轮实际拟合坐标。',
        'Hb 同坐标分解：观测 = 条件物理响应 + 受限共同项 + 残差；曲线可能相互抵消。',
        '这些是事后描述性案例，不是预注册独立检验；任何“神经源 / 浅表 / 纯噪声”标签都需要额外证据。'],
        'waveform_provenance.json; measured_cells.csv; retained per-cell NPZ', '成分波形')
    for f in wfigs[1:]: fig_slide(f, '原始与成分波形')
    table_slide('看见分离曲线之后，仍不能跳过的解释边界', '数学分解与来源识别不是同一件事', ['曲线 / 坐标','可用名称','不能据此命名'], [
        ['a、b','有效宽频 / 谱对比连续模式','兴奋源、抑制源、细胞来源'],
        ['physical_prediction','给定驱动、参数与初态的 Hb 响应','已恢复的神经源 Hb'],
        ['observation_component','受限共同 Hb 观测项','已识别浅表血流、系统生理、运动伪迹'],
        ['观测减预测','该模型与观测坐标下的残差','纯噪声或全部私有神经信息']],
        [3.0,4.6,4.4], 'semantic_response_ssm.py; waveform_provenance.json; docs/METHOD_RATIONALE.md', '成分解释', size=17)

    fig_slide(mfigs['method_tokenizer'], 'Tokenizer 设计')
    body('语义目标：形状与幅值分别输出', '完整 observation 分支保留模型不能可靠命名的信息', [
        '参考中心化：a_c(t) = a(t) − mean[a(0:20)]；前 20 点对应 5 s。',
        '幅值 A = RMS(a_c)；形状 s(t) = a_c(t)/A；幅值目标 ℓ = log A，并只用训练分组标准化。',
        '分布臂输出 Gaussian 边缘均值 / SD；不表示完整时序后验或跨时协方差。',
        '连续推理后聚合为 12 个 patch；平均边缘 SD 只是完全正相关下的上界，不是已校准 patch 区间。'],
        'semantic_response_tokenizer.py: token_metadata / export_tokens; synthetic prepared manifest', 'Tokenizer 设计')
    table_slide('四臂 × 三种子：目标、随机化和容量分开比较', '基础神经网络均为 458,168 参数；O-only 的附加线性读出另计', ['臂','训练机制','语义与重建'], [
        ['等容量 O-only','随机化生成器','全部均值坐标只学习完整重建'],
        ['受限点模型','受限生成器','shape / log幅值点监督 + 完整重建'],
        ['随机化点模型','随机化生成器','相同点监督，分离随机化效应'],
        ['随机化分布模型','随机化生成器','Gaussian语义损失 + 完整重建']],
        [3.0,3.3,5.7], 'tokenizer/summary.json; resolved_config.json: tokenizer', 'Tokenizer 设计',
        'width=64、O维度=8、batch=64、Adam lr=0.001；最多160 epoch、patience=20；seeds=101/202/303。', size=17)
    body('O-only 合成评价与区间的实际来源', '网络监督、冻结表征可读出性和预测区间分别记录', [
        'O-only 网络本身只训练完整重建；合成评价另用训练集真值，对全部均值坐标拟合 shape / log幅值线性 ridge（α=1），再冻结到评价集。',
        '因此 O-only 合成误差衡量“重建表征 + 监督线性读出”；等容量指基础网络与均值瓶颈，不包括这个附加 ridge。',
        '点模型和 O-only 用选择集残差 RMS：形状和幅值各一个常数尺度。分布模型用网络 SD × 选择集温度；温度限制为[0.1,10]。',
        '这个合成可读出性端点与后面的实测条件 Hb 探针不同；两者都不能自动赋予表示生理来源身份。'],
        'train_semantic_response_tokenizer.py: semantic_predictions / fit / calibrate / uncertainty', 'Tokenizer 设计')
    table_slide('合成训练与评价身份严格分开', '同基础轨迹的全部干预变体放在同一个 split', ['分组','基础身份','场景与记录'], [
        ['训练','512 / family','8场景 × control/intervention = 8,192条'],
        ['选择','128 / family','8场景 × control/intervention = 2,048条'],
        ['共同评价','256','20场景 × control/intervention = 10,240条'],
        ['计数解释','所有四臂共享相同评价案例','两个模态 × 12模型 = 245,760条评分记录']],
        [2.6,3.0,6.4], 'tokenizer/prepared/*/*.json; tokenizer/summary.json', 'Tokenizer 设计',
        '三优化种子不增加评价基础身份或被试数；区间与误差必须分 control / intervention 汇报。', size=17)
    table_slide('评价场景 I：从真实变化到观测混淆', '先检查训练支持范围内的函数恢复及不应出现的语义变化', ['类别','场景','问题'], [
        ['真实驱动改变','true_a_amplitude / true_a_shape','语义是否跟随真实变化'],
        ['观测增益','feature_gain_only / Hb_gain_only','幅值是否被误认作驱动变化'],
        ['共同 Hb','common_Hb / correlated_Hb_component','广泛变化是否污染共享语义'],
        ['动力学 / 初态','tau_change / initial_change','观测形状变化会被写入哪一分支'],
        ['基准','baseline','在已定义生成机制中能否学会目标']],
        [2.35,5.8,3.85], 'resolved_config.json: tokenizer.evaluation_scenarios; semantic_response_synthetic.py', 'Tokenizer 设计', size=16)
    table_slide('评价场景 II：生成机制外推与识别反例', '随机参数范围更宽，不等于已经覆盖新的生成机制', ['类别','场景','问题'], [
        ['未见时间基底','non_dct_impulse / ou_drive','平滑训练先验是否压制真实变化'],
        ['观测映射改变','time_varying_feature_gain / rotated_spectral_mixture','固定谱/增益假设是否稳定'],
        ['响应机制改变','input_lowpass / viscoelastic_outflow / alternative_hrf','时间响应错配与状态归属'],
        ['关系改变','independent_Hb_drive / no_coupling','无共享信息时会否仍输出共享语义'],
        ['精确别名','exact_eeg_alias / exact_joint_alias','相同观测是否对应不止一个真值']],
        [2.3,6.25,3.45], 'resolved_config.json: tokenizer.evaluation_scenarios; semantic_response_synthetic.py', 'Tokenizer 设计', size=15)
    for f in token_quant: fig_slide(f, 'Tokenizer 合成结果')
    body('精确 alias 应怎样解释', '这是明确允许的竞争机制，不是要求网络从相同输入凭空辨认真值', [
        '将真实驱动幅值放大，同时反向改变 EEG 观测增益，可保留完全相同的 EEG。',
        'joint alias 进一步使用精确 Hb 补偿，使两个模态输入都相同，但真实驱动幅值不同。',
        '相同输入的确定编码器必然给出相同输出；本轮失败在于窄幅值区间不覆盖允许的两个机制。',
        '解决方向是识别集合、先验敏感性或明示幅值不受支持；只增加输入异常检测器无法识别精确 alias。'],
        'semantic_response_synthetic.py: exact aliases; tokenizer/scenario_variant_seed_summary.csv', '反事实波形')
    for f in sfigs: fig_slide(f, '合成反事实波形')
    fig_slide(mfigs['method_public_probe'], '实测冻结探针')
    for f in public_quant: fig_slide(f, '实测冻结探针')
    table_slide('实测主结果：语义分支没有提供额外收益', '随机化分布模型，前缀40；三优化种子均值；相同上下文与评分支持',
        ['输入表征','NRMSE ↓','标准化 CRPS ↓'], [
        ['prepared 原始上下文', '1.228612','0.846234'],
        ['上下文 + 同模型 O','1.151679','0.790067'],
        ['上下文 + 同模型 O + S','1.155386','0.792139'],
        ['上下文 + 独立等容量纯 O','1.139294','0.781601']], [6.2,2.9,2.9],
        'tokenizer/public_probes/distribution_randomized_generator/seed_*/record.json', '实测结论',
        'O→O+S 的 NRMSE 增加0.003707；相对等容量纯O增加0.016092，三个优化种子方向一致。', size=19)
    body('三部分证据共同限定当前结论', '不能由合成成功翻转实测负结果，也不能由当前负结果否定所有生理表示方案', [
        'SSM：谱方向、稳定动力学与配对特异性尚未构成一致的 EEG→Hb 机制证据。',
        '合成：可以恢复生成器支持的相对形状；幅值与观测增益仍混淆，跨机制不确定性失校准。',
        '实测：observation 表示在可见前缀下有工程价值；semantic 坐标未超过 observation 与容量对照。',
        '本轮交付的是可复现的连续接口、有限正结果和清晰失败边界；没有获得生理 teacher / VQ 晋级依据。'],
        'registry main.semantic_response; summary.json; tokenizer/summary.json', '综合判断')
    table_slide('执行完成与科学支持分别记录', '失败和修复保留；最终分母没有由成功子集替代', ['核验项目','结果'], [
        ['SSM 合成','5,760项成功；2,880实际前缀拟合全部收敛'],
        ['SSM 公开','23,040项成功；9,000实际拟合全部收敛'],
        ['Tokenizer','12训练 + 12合成评价 + 12公开探针全部完成'],
        ['评分分母','245,760合成记录；36,720公开探针记录，无非有限 / 失败行'],
        ['软件与证据','125项定向测试；29项证据核验；状态校验通过']], [3.55,8.45],
        'verification.json; tokenizer/summary.json; supervisor/completion.json', '执行与复现', size=18)
    body('运行过程中的修复与资源选择', '源码版本、失败日志和计算规则均保留，未用重训掩盖负结果', [
        '持久 user systemd 监督；CPU 单数值线程；SSM 按独立单元并行，GPU 训练按实测吞吐选择每卡两任务。',
        '恢复入口曾因 CUDA checkpoint 将 RNG 状态放入错误设备而失败；修复只把 RNG 状态恢复到 CPU，训练损失与网络不变。',
        '汇总修复恢复了反事实表；features/batch_v2 补充训练增益来源，保留原 batch_v1。',
        '全部训练与评分结束；原始失败记录、六套不可变源码快照、配置、监督日志及完成状态都可追溯。'],
        'training_resume_v3.json; training_launch_v2.json; source_snapshot_*_identity.json; analysis_launch.json', '执行与复现')
    table_slide('下一步应由失败面决定', '建议，不代表新实验已启动', ['未解决问题','最小下一步','必须保留的检验'], [
        ['动力学选择退化','固定真值驱动 / 参数的有限 oracle 合成','区分输入误差、参数失配与机制混淆'],
        ['幅值非唯一','识别集合或生成先验敏感性','保留 exact alias 与场景分层覆盖'],
        ['跨机制形状失效','扩展时间基底前先定义可恢复函数目标','未见机制评价不能参与重选'],
        ['S 无实测增量','重新检验目标本身；保留完整 O 基线','同容量、同上下文、同前缀与配对 null']],
        [2.7,4.75,4.55], 'registry main.semantic_response.next_step; scientific review of retained endpoints', '下一步',
        '不自动追加实测调参、受保护评价或 VQ；新的来源主张需要独立观测或机制锚定。', size=16)
    table_slide('附录：完整响应方程与 Hb 读出', '实现坐标为 [u,s,f,v,p,q]；状态曲线是模型条件推断结果', ['坐标','方程 / 定义'], [
        ['输入状态','τn>0：du/dt=(r−u)/τn，u(0)=0；τn=0：u=r'],
        ['信号与血流','ds/dt = β·u − κ·s − γ_feedback·(f−1)；df/dt = s'],
        ['体积','dv/dt = (f−fout)/τ；fout = (τ·v^(1/α)+τv·f)/(τ+τv)'],
        ['总 Hb 坐标','dp/dt = [f−(fout/v)·p]/τ'],
        ['去氧 Hb 坐标','dq/dt = [f·E(f)/E0−(fout/v)·q]/τ；E(f)=1−(1−E0)^(1/f)'],
        ['Hb 读出','HbO=P0·(p−1)−Q0·(q−1)；HbR=Q0·(q−1)'],
        ['观测坐标','读出经过同一 Hb 处理算子后，再加共同项；不能直接与原生光强相加']],
        [2.3,9.7], 'semantic_response_ssm.py: _physical_response_slope / fixed_driver_response', '附录',
        'γ_feedback 为固定生理模型参数；与谱读出候选权重 γ 不是同一参数。β和其余未声明扩展参数固定为parent H0。', size=16)
    table_slide('附录：冻结的 SSM 候选与数值预算', '参数选择只使用允许的训练 / 选择资料；不按评价机制标签重选', ['参数 / 规则','冻结内容'], [
        ['谱混合 γ','−1、−0.5、0、0.5、1'],
        ['独立增益 g','0、0.5、1/√2、2/√5、1'],
        ['输入低通 τn / 出流 τv','τn∈{0,0.5,1,2}s；τv∈{0,1,3,6}s'],
        ['EEG 先验与时间基底','24 temporal modes；state SD=0.025；AR φ=0.95、weight=0.1'],
        ['前缀与选择','40点主选择复用于0点；失败惩罚10；同分选择较简单候选'],
        ['求解器','最大120 evaluations；RK每步4 substeps；gradient tolerance=10⁻⁶']],
        [3.8,8.2], 'resolved_config.json: model/solver/calibration', '附录', size=17)
    body('附录：训练损失和表征容量', '同样的网络参数数量，明确不同坐标受到什么监督', [
        '每模态 L = L_shape + 0.25·L_logA + L_full_reconstruction，总损失为 EEG 与 Hb 之和。点模型语义用平方误差，分布模型用 Gaussian NLL；重建在本模态训练标准化坐标计算。',
        '受监督臂把语义均值 detach 后送入本模态解码器；完整重建不能反向改写语义头。',
        'O-only 臂不使用语义真值损失；全部均值坐标都可通过完整重建学习，因此它的原始 S 头误差不是公平的语义恢复基线。',
        '均值表征：每时间点8维 O、1维形状，加1个窗口幅值；尺度头所有臂均存在，但点模型不将未监督尺度头解释为区间。'],
        'semantic_response_tokenizer.py: response_loss / _ModalityBranch; resolved_config.json', '附录')
    epoch_rows = []
    for arm, label in [('observation_only_capacity_matched','等容量 O-only'),
                       ('point_restricted_generator','受限点'),('point_randomized_generator','随机化点'),
                       ('distribution_randomized_generator','随机化分布')]:
        row = [label]
        for seed in (101,202,303):
            rec = next(x for x in tok['fits'] if x['arm']==arm and x['seed']==seed)
            row.append(f"{rec['best_epoch']} / {rec['completed_epochs']}")
        epoch_rows.append(row)
    table_slide('附录：12 个训练实例的选择与停止', '每格为最佳选择 epoch / 实际停止 epoch；相同上限不等于实际训练步数相同',
        ['臂','seed101','seed202','seed303'], epoch_rows, [3.6,2.8,2.8,2.8],
        'tokenizer/fits/*/seed_*/record.json; tokenizer/summary.json', '附录',
        '全部≤160 epoch；早停patience=20。最佳模型由各臂注册的选择损失决定，不能直接横比不同损失的原始数值。', size=19)
    table_slide('文献提供动机，不替代本数据的证据', '本轮引用均用于明确模型假设或解释边界', ['文献','与本轮有关的内容','不能外推'], [
        ['Buxton et al., 2004','血流、血容量、氧合响应及出流松弛模型','不是本 EEG–fNIRS 参数的标定'],
        ['Hermes et al., 2017','不同电生理成分与血管响应的关系','ECoG–BOLD 结果不直接指定头皮谱读出'],
        ['Locatello et al., 2019','解耦依赖模型 / 数据的归纳假设','重建好并不赋予坐标来源名称'],
        ['Hermans et al., 2022','模拟推断可能过度自信，需覆盖审计','本轮 Gaussian 头不是完整 Bayesian SBI']],
        [3.1,5.0,3.9], [
          'https://fmri.ucsd.edu/tliu/pdf/buxton04_model.pdf',
          'https://journals.plos.org/plosbiology/article?id=10.1371/journal.pbio.2001461',
          'https://proceedings.mlr.press/v97/locatello19a.html',
          'https://arxiv.org/abs/2110.06581'], '参考', size=16)
    body('参考文献与可访问链接', '文献动机与本轮数据支持程度分开；不以引用替代实验验证', [
        'Buxton et al. (2004). Modeling the hemodynamic response to brain activation.\nhttps://fmri.ucsd.edu/tliu/pdf/buxton04_model.pdf',
        'Hermes et al. (2017). Neuronal synchrony and the relation between the blood-oxygen-level dependent response and the local field potential.\nhttps://journals.plos.org/plosbiology/article?id=10.1371/journal.pbio.2001461',
        'Locatello et al. (2019). Challenging Common Assumptions in the Unsupervised Learning of Disentangled Representations.\nhttps://proceedings.mlr.press/v97/locatello19a.html',
        'Hermans et al. (2022). A Trust Crisis In Simulation-Based Inference? Your Posterior Approximations Can Be Unfaithful.\nhttps://arxiv.org/abs/2110.06581'],
        'Primary sources checked 2026-10-03; see hyperlinks on slide', '参考')
    body('证据与复现导航', '正文、表格可编辑；所有科学图为 200–250 dpi PNG，并保留来源清单', [
        '实验合同：semantic_response_discovery_v1.yaml；状态 owner：research_state/registry.json。',
        'SSM：summary / dataset_equal_comparisons / primary_null_specificity / calibration 及逐单元 NPZ。',
        'Tokenizer：summary、scenario_variant_seed_summary，以及每模型 evaluation / public_probes。',
        '本报告：presentation_sources.json、三类图形 provenance、源代码快照与 PDF 导出核验。'],
        'experiments/RESULTS_INDEX.md; presentation_sources.json', '复现')
    deck.save(out/draft_name)
    record = dict(schema='semantic_response_presentation_sources_v1', run=str(run), presentation=draft_name,
                  slides=slide_sources, slide_count=len(deck.slides), bitmap_figures=True,
                  editable_body_text=True, shape_audit=shape_audit,
                  interpretation='dated communication export; owning evidence remains the registered run',
                  report_source_sha256=hashlib.sha256(Path(__file__).read_bytes()).hexdigest())
    (out/'presentation_sources.json').write_text(json.dumps(record, ensure_ascii=False, indent=2)+'\n')
    print(json.dumps(dict(presentation=str(out/draft_name), slides=len(deck.slides),
                          figures=sum('figure' in row for row in slide_sources)), ensure_ascii=False))
    return record


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--run", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--replay", action="store_true", help="Render nonlinear fixed-driver replay against its parent")
    parser.add_argument("--nonlinear-fit", action="store_true", help="Render nonlinear optimization cells and complete denominators")
    parser.add_argument("--gain-prior-fit", action="store_true", help="Render four-arm effective gain and driver amplitude prior diagnostic")
    parser.add_argument("--terminal-input-snapshot", type=Path, help="Frozen terminal manifest/config/summary/metrics for report replay while measured phase proceeds")
    parser.add_argument("--pilot-truth-audit", type=Path, help="Optional retained pilot DC/initial compensation JSON for conditional report background")
    parser.add_argument("--conditional-optical-gain", action="store_true", help="Render terminal conditional optical mapping and gain evidence")
    parser.add_argument("--step-control", action="store_true", help="Audit and render the complete paired training-only solver comparison")
    parser.add_argument("--fixed-roi-optical", action="store_true", help="Render separate retained optical target sensitivity conditions")
    parser.add_argument("--fixed-roi-initial", action="store_true", help="Render free vs tied initial total-Hb/volume contrast")
    parser.add_argument("--fixed-roi-flow", action="store_true", help="Render fixed ROI log-flow soft regularization diagnostic")
    parser.add_argument("--fixed-roi-tau", action="store_true", help="Render fixed AF7Fp1 shared nonlinear tau diagnostic")
    parser.add_argument("--waveform-diagnostic", action="store_true")
    parser.add_argument("--teacher-robustness", action="store_true")
    parser.add_argument("--component-attribution", action="store_true")
    parser.add_argument("--semantic-response", action="store_true", help="Render the complete response dynamics and continuous semantics PPT report")
    parser.add_argument("--presentation-name", default="SSM_SEMANTIC_RESPONSE_REPORT.pptx", help="Fresh semantic-response PPT filename; existing exports are preserved")
    parser.add_argument("--volume-fraction-run", type=Path)
    args = parser.parse_args()
    if args.semantic_response:
        render_semantic_response_report(args.run.resolve(), args.output.resolve(), draft_name=args.presentation_name)
        return
    if args.component_attribution:
        render_component_attribution(args.run.resolve(),args.output.resolve())
        return
    if args.teacher_robustness:
        render_teacher_robustness(args.run.resolve(),args.output.resolve())
        return
    if sum((args.replay, args.nonlinear_fit, args.gain_prior_fit, args.fixed_roi_tau, args.fixed_roi_flow, args.fixed_roi_initial, args.fixed_roi_optical, args.conditional_optical_gain, args.waveform_diagnostic,args.step_control)) > 1:
        parser.error("Choose one report mode")
    if (args.pilot_truth_audit or args.terminal_input_snapshot) and not args.conditional_optical_gain:
        parser.error("--pilot-truth-audit requires --conditional-optical-gain")
    if args.volume_fraction_run and not args.waveform_diagnostic:
        parser.error("--volume-fraction-run requires --waveform-diagnostic")
    if args.step_control:
        render_step_control(args.run.resolve(),args.output.resolve())
        return
    if args.waveform_diagnostic:
        render_waveform_diagnostic(args.run.resolve(),args.output.resolve(),args.volume_fraction_run.resolve() if args.volume_fraction_run else None)
        return
    if args.conditional_optical_gain:
        render_conditional_optical_gain(args.run.resolve(),args.output.resolve(),pilot_truth_audit=args.pilot_truth_audit,terminal_input_snapshot=args.terminal_input_snapshot)
        return
    if args.fixed_roi_optical:
        render_fixed_roi_optical(args.run.resolve(),args.output.resolve())
        return
    if args.fixed_roi_initial:
        render_nonlinear_fit(args.run.resolve(), args.output.resolve(), fixed_roi_initial=True)
        return
    if args.fixed_roi_flow:
        render_nonlinear_fit(args.run.resolve(), args.output.resolve(), fixed_roi_flow=True)
        return
    if args.fixed_roi_tau:
        render_nonlinear_fit(args.run.resolve(), args.output.resolve(), fixed_roi_tau=True)
        return
    if args.gain_prior_fit:
        render_nonlinear_fit(args.run.resolve(), args.output.resolve(), gain_prior=True)
        return
    renderer = render_nonlinear_fit if args.nonlinear_fit else (render_replay if args.replay else render)
    renderer(args.run.resolve(), args.output.resolve())


if __name__ == "__main__":
    main()
