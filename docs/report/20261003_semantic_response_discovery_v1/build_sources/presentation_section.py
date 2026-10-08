def build_response_method_figures(repo, run, out):
    """Diagrams of the implemented contract; no signals or model fitting."""
    import json
    from pathlib import Path
    import numpy as np
    import matplotlib.pyplot as plt
    from matplotlib import font_manager
    from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
    out = Path(out)
    folder = out / 'figures'
    folder.mkdir(parents=True, exist_ok=True)
    font = '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'
    font_manager.fontManager.addfont(font)
    plt.rcParams.update({'font.family': font_manager.FontProperties(fname=font).get_name(),
                         'font.size': 13, 'axes.unicode_minus': False})
    navy, teal, amber, red = '#15364A', '#008C88', '#D18B30', '#B94758'
    figures = []
    def canvas():
        fig, ax = plt.subplots(figsize=(12, 4.8))
        ax.set(xlim=(0, 12), ylim=(0, 4.8))
        ax.axis('off')
        return fig, ax
    def box(ax, x, y, w, h, text, color=navy, size=14):
        ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle='round,pad=.07,rounding_size=.10',
                                   facecolor=color, edgecolor='none'))
        ax.text(x+w/2, y+h/2, text, ha='center', va='center', color='white', fontsize=size,
                linespacing=1.55)
    def arrow(ax, p, q, label=None, color=navy):
        ax.add_patch(FancyArrowPatch(p, q, arrowstyle='-|>', mutation_scale=17, lw=1.8, color=color))
        if label:
            ax.text((p[0]+q[0])/2, (p[1]+q[1])/2+.18, label, ha='center', color=color, fontsize=11)
    def save(fig, key, title, caption, source):
        path = folder / (key+'.png')
        fig.savefig(path, dpi=240, bbox_inches='tight', facecolor='white')
        plt.close(fig)
        figures.append(dict(key=key, path=str(path.resolve()), title=title, caption=caption, source=source))

    fig, ax = canvas()
    for x, title, lines, color in [(0.25, '讨论与可证伪假设', '谱信息还是输入衰减？\n局部 EEG 还是共享 Hb？\n形状还是幅值／观测增益？', navy),
            (4.3, '冻结的实验对照', '匹配衰减 + 独立增益\n空间 / 历史基线 + null\n反事实 + alias + 容量控制', teal),
            (8.35, '分层解释结果', '条件形状可恢复\n幅值与跨机制校准失败\n实测 S 未提供增量', amber)]:
        box(ax, x, 2.65, 3.4, .65, title, color)
        ax.text(x+1.7, 1.65, lines, ha='center', va='center', fontsize=15, linespacing=1.9, color=navy)
    arrow(ax, (3.75, 2.98), (4.16, 2.98))
    arrow(ax, (7.8, 2.98), (8.2, 2.98))
    ax.text(6, .35, '完成计算 ≠ 通过机制检验；重建改善 ≠ 识别共享生理来源', ha='center', color=red, fontsize=16)
    save(fig, 'method_evidence_chain', '讨论如何转化为实验，而不是直接转化为结论',
         '本报告把讨论、具体实现、对照端点和科学判读逐一对应；历史结果只作为动机。',
         'docs/METHOD_RATIONALE.md; docs/EXPERIMENT_PLAN.md; frozen run summary.json')

    fig, ax = canvas()
    box(ax, .2, 3.45, 2.15, .75, '原生 EEG 电压\n通道 / 时钟 / 单位', navy, 13)
    box(ax, 3.0, 3.45, 2.25, .75, '频带功率与参考\n6 通道 × 5 频带', teal, 13)
    box(ax, 6.0, 3.45, 2.2, .75, 'prepared EEG\n120 × 30', teal, 13)
    box(ax, 9.1, 3.45, 2.45, .75, '固定谱模式 a、b\n与未解释残差', amber, 13)
    box(ax, .2, 1.65, 2.15, .75, '原生光强 / 发布 Hb\n数据集定义不同', navy, 13)
    box(ax, 3.0, 1.65, 2.25, .75, '声明的转换 / 处理\n时钟对齐与参考', teal, 13)
    box(ax, 6.0, 1.65, 2.2, .75, 'prepared Hb\n120 × 2', teal, 13)
    box(ax, 9.1, 1.65, 2.45, .75, '响应项 + 共同项\n+ 模型残差', amber, 13)
    for y in (3.825, 2.025):
        for p, q in [(2.4, 2.9), (5.35, 5.9), (8.3, 9.0)]: arrow(ax, (p, y), (q, y))
    ax.text(6, .45, '反变换只回到 parent prepared 坐标；不能恢复被预处理移除的原始信息', ha='center', color=red, fontsize=15)
    save(fig, 'method_raw_lineage', '成分分离必须回到同一窗口与同一观测坐标',
         'EEG 的模式贡献可在 log-power 特征中相加，不能直接与高频原始电压相加。',
         'docs/DATA_CONTRACT.md; parent prepared record identities; unified_physiology.py')

    fig, ax = canvas()
    for y, name, color in [(3.65, '目标 Hb：前缀 40', teal), (2.45, '目标 Hb：前缀 0', amber), (1.25, '完整 EEG / 其他区域 Hb', navy)]:
        ax.text(.15, y+.22, name, ha='left', va='center', color=navy, fontsize=13)
        x0, width = 3.25, 8.3
        ax.add_patch(plt.Rectangle((x0, y), width, .45, color='#E8EDF1'))
        if y == 3.65: ax.add_patch(plt.Rectangle((x0, y), width/3, .45, color=teal))
        if y == 1.25: ax.add_patch(plt.Rectangle((x0, y), width, .45, color=navy))
        if y != 1.25:
            ax.add_patch(plt.Rectangle((x0+2*width/3, y), width/3, .45, fill=False, hatch='///', edgecolor=red, lw=1.2))
        for sec in (0, 10, 20, 30):
            x = x0+width*sec/30
            ax.text(x, y-.18, str(sec)+' s', ha='center', va='top', color=navy, fontsize=11)
    ax.text(7.4, .35, '绿色：可用目标前缀     灰色：不可输入目标 Hb     斜线：隐藏评分区 20–30 s', ha='center', fontsize=12)
    save(fig, 'method_visibility', '同一时间轴上区分可见、隐藏和评分支持',
         '前缀 40 = 10 s；零前缀不用目标 Hb token。parent 非因果处理与全窗 EEG 仍使端点属于离线补全。',
         'semantic_response_discovery_v1.yaml: measured; public_encoding_view; public_raw_context')

    fig, ax = canvas()
    box(ax, .2, 2.4, 2.1, 1, 'EEG-only\n固定 a(t)、b(t)', navy)
    box(ax, 3.05, 3.4, 2.5, .8, '谱混合 / 匹配衰减\n或独立增益 g·a', teal, 13)
    box(ax, 3.05, 1.4, 2.5, .8, '初态与共同项\n仅由可见 Hb 前缀校准', amber, 13)
    box(ax, 6.25, 3.4, 2.35, .8, '固定 H0 响应核\n单独增加 τn 或 τv', teal, 13)
    box(ax, 9.3, 2.4, 2.35, 1, '隐藏 Hb 补全\n及 proper-score', navy, 13)
    arrow(ax, (2.3, 3.0), (2.94, 3.8))
    arrow(ax, (5.55, 3.8), (6.12, 3.8))
    arrow(ax, (8.65, 3.8), (9.22, 3.08))
    arrow(ax, (5.6, 1.8), (9.18, 2.65))
    ax.text(6.1, .40, 'Hb 后缀不反向重估驱动；零前缀使用初态先验与零共同项', ha='center', fontsize=15, color=red)
    save(fig, 'method_ssm_inference', '新推断制度：先定 EEG 驱动，再用 Hb 前缀校准',
         '与历史全窗联合拟合的可见信息不同，因此历史重建数字不能直接当作本轮同条件基线。',
         'src/inference/semantic_response_ssm.py; frozen evaluation runner')

    fig, ax = canvas()
    for y, label in [(3.55, 'EEG + mask'), (1.15, 'Hb + mask')]:
        box(ax, .1, y, 1.9, .8, label, navy, 14)
        box(ax, 2.65, y+.20, 2.3, .65, '语义编码器 / 头 S\n形状、log幅值、SD', teal, 12)
        box(ax, 2.65, y-.65, 2.3, .65, '独立观测编码器 → O', navy, 13)
        box(ax, 6.3, y-.2, 2.4, .9, '非线性本模态解码器\n接收 O 与 detach(S)', amber, 12)
        box(ax, 9.35, y-.2, 2.35, .9, '完整本模态\n观测重建输出', teal, 13)
        arrow(ax, (2, y+.4), (2.56, y+.52))
        arrow(ax, (2, y+.25), (2.56, y-.30))
        arrow(ax, (4.97, y+.52), (6.2, y+.30), 'detach')
        arrow(ax, (4.97, y-.30), (6.2, y+.03))
        arrow(ax, (8.75, y+.25), (9.24, y+.25))
    ax.text(6, .12, '推理各读本模态；监督臂重建梯度不改写 S；O-only 对照让全部均值坐标学习重建', ha='center', fontsize=12, color=red)
    save(fig, 'method_tokenizer', 'semantic–observation 原型：输入独立，输出职责明确',
         '解码器是非线性联合解码；S 与 O 不是可直接相加的“神经波形”和“噪声波形”。',
         'src/tokenizers/semantic_response_tokenizer.py: _ModalityBranch / response_loss')

    fig, ax = canvas()
    box(ax, .1, 2.5, 3.2, 1.4, '所有方法相同的上下文\nprepared 局部 EEG\n其他区域 Hb + 自身前缀', navy, 14)
    for y, title, color in [(4.0, '仅共同上下文', navy), (3.0, '+ 同模型 O', teal),
                            (2.0, '+ 同模型 O + S', amber), (1.0, '+ 独立等容量纯 O', teal)]:
        box(ax, 4.1, y-.15, 3.2, .60, title, color, 13)
        arrow(ax, (3.33, 3.15), (4.0, y+.15))
        arrow(ax, (7.38, y+.15), (8.14, 2.65))
    box(ax, 8.25, 2.15, 3.45, 1.05, '训练集拟合 ridge\n选择集定 α / 预测残差 SD\n评价集只评分', navy, 13)
    ax.text(6.0, .2, '编码器冻结；O→O+S 检验语义增量，等容量纯 O 检验监督是否优于容量收益', ha='center', fontsize=13, color=red)
    save(fig, 'method_public_probe', '实测探针同时控制可见上下文和表征总容量',
         'context_raw 只是“未编码的 prepared 上下文”，并不是原生 EEG 电压或原始光强。',
         'train_semantic_response_tokenizer.py: probe / fit_public_probe; public_probes/*/selection.json')
    result = dict(figures=figures, schema='semantic_response_method_diagrams_v1')
    (out/'method_diagram_provenance.json').write_text(json.dumps(result, ensure_ascii=False, indent=2)+'\n')
    return result


def render_semantic_response_report(run, out, draft_name='SSM_SEMANTIC_RESPONSE_REPORT.pptx'):
    """Detailed, editable Chinese report of retained evidence and frozen replay."""
    import json
    import math
    import hashlib
    import sys
    from pathlib import Path
    import numpy as np
    import pandas as pd
    from PIL import Image, ImageFont
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR
    from pptx.enum.shapes import MSO_SHAPE
    run, out = Path(run).resolve(), Path(out).resolve()
    repo = Path(__file__).resolve().parents[2]
    if str(repo) not in sys.path:
        sys.path.insert(0, str(repo))
    if (out/draft_name).exists():
        raise ValueError('Preserve an existing presentation; choose a fresh versioned output or draft name')
    summary = json.loads((run/'summary.json').read_text())
    verification = json.loads((run/'verification.json').read_text())
    if summary['execution'] != 'completed' or not verification['passed']:
        raise ValueError('Report needs completed and verified owning run evidence')
    out.mkdir(parents=True, exist_ok=True)
    cfg = json.loads((run/'resolved_config.json').read_text())
    metadata = json.loads((run/'metadata_plan_proposal.json').read_text())
    tok = json.loads((run/'tokenizer/summary.json').read_text())
    manifests = {}
    for name, builder in [('quantitative_provenance.json', build_response_quantitative_figures),
                          ('waveform_provenance.json', build_response_waveform_figures),
                          ('synthetic_waveform_provenance.json', build_synthetic_waveform_figures)]:
        path = out/name
        manifests[name] = json.loads(path.read_text()) if path.exists() else builder(repo, run, out)
    methods = build_response_method_figures(repo, run, out)
    qfigs = manifests['quantitative_provenance.json']['figures']
    wfigs = manifests['waveform_provenance.json']['figures']
    sfigs = manifests['synthetic_waveform_provenance.json']['figures']
    mfigs = {f['key']: f for f in methods['figures']}
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    font = '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'
    navy, teal, red, muted = '15364A', '008C88', 'B94758', '526C7D'
    slide_sources, shape_audit = [], []

    def wrap_count(value, width, size):
        pil = ImageFont.truetype(font, round(size*4))
        count = 0
        for line in str(value).split('\n'):
            if not line:
                count += 1
                continue
            current = ''
            count += 1
            for char in line:
                if current and pil.getlength(current+char)/4 > width*72:
                    count += 1
                    current = char
                else:
                    current += char
        return count

    def text(slide, value, x, y, w, h, size=22, color=navy, bold=False, minimum=14):
        size = float(size)
        while size > minimum and wrap_count(value, w, size)*size*1.17 > h*72:
            size -= .5
        if wrap_count(value, w, size)*size*1.1 > h*72+3:
            raise ValueError(f'Text does not fit: {str(value)[:80]}')
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for idx, value_line in enumerate(str(value).split('\n')):
            para = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            para.text, para.line_spacing = value_line, 1.12
            para.space_before = para.space_after = Pt(0)
            for rr in para.runs:
                rr.font.name = 'Noto Sans CJK SC'
                rr.font.size, rr.font.bold = Pt(size), bold
                rr.font.color.rgb = RGBColor.from_string(color)
                if value_line.startswith(('https://', 'http://')):
                    rr.hyperlink.address = value_line
        shape_audit.append(dict(slide=len(deck.slides), text=str(value), x=x, y=y, width=w, height=h, font_pt=size))
        return shape

    def new(title, subtitle='', source='', chapter=''):
        s = deck.slides.add_slide(deck.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = RGBColor.from_string('F7F9FB')
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(.10))
        bar.fill.solid(); bar.fill.fore_color.rgb = RGBColor.from_string(teal); bar.line.fill.background()
        text(s, title, .55, .28, 12.22, .57, 28, bold=True, minimum=22)
        text(s, subtitle, .58, .99, 12.15, .46, 15, color=muted, minimum=12)
        text(s, '2026.10.03  |  响应动力学与连续语义  |  '+chapter, .58, 7.14, 11.7, .20, 10, color=muted, minimum=10)
        text(s, f'{len(deck.slides):02}', 12.2, 7.13, .5, .22, 10, color=muted, minimum=10)
        evidence = source if isinstance(source, str) else json.dumps(source, ensure_ascii=False)
        notes = (f'实验事实源：{run}\n证据：{evidence}\n'
                 '公开开发端点；此前cohort已分析。前缀和评分均位于parent非因果处理后的30秒窗口。'
                 '优化种子共享评价基础身份，不能当独立被试或独立置信区间。\n')
        s.notes_slide.notes_text_frame.text = notes
        slide_sources.append(dict(slide=len(deck.slides), title=title, chapter=chapter, evidence=source))
        return s

    def caption(s, value):
        text(s, value, .65, 6.48, 12.02, .52, 13, color=muted, minimum=11)

    def body(title, subtitle, paragraphs, source, chapter='', takeaway=None):
        s = new(title, subtitle, source, chapter)
        if isinstance(paragraphs, str): paragraphs = [paragraphs]
        count = len(paragraphs)
        height = 4.55/max(count, 1)
        for i, paragraph in enumerate(paragraphs):
            text(s, paragraph, .8, 1.72+i*height, 11.8, height-.16, 23 if count<=4 else 20,
                 color=navy, minimum=16)
        if takeaway: caption(s, takeaway)
        return s

    def table_slide(title, subtitle, headers, rows, widths, source, chapter='', takeaway=None, size=18):
        s = new(title, subtitle, source, chapter)
        # Size each row from the longest actual wrapped cell, not a fixed guess.
        font_size = float(size)
        while True:
            heights = [max(.43, max(wrap_count(value, width-.20, font_size) for value, width in zip(row, widths))*font_size*1.17/72+.12)
                       for row in [headers]+rows]
            if sum(heights) <= 4.67 or font_size <= 13:
                break
            font_size -= .5
        if sum(heights) > 4.8:
            raise ValueError('Table does not fit: '+title)
        sh = s.shapes.add_table(len(rows)+1, len(headers), Inches(.65), Inches(1.68), Inches(12), Inches(sum(heights)))
        tab = sh.table
        for col, width in zip(tab.columns, widths): col.width = Inches(width)
        for i, row in enumerate([headers]+rows):
            tab.rows[i].height = Inches(heights[i])
            for j, value in enumerate(row):
                c = tab.cell(i,j)
                c.text = str(value)
                c.margin_left, c.margin_right = Inches(.1), Inches(.07)
                c.margin_top, c.margin_bottom = Inches(.04), Inches(.04)
                c.vertical_anchor = MSO_ANCHOR.MIDDLE
                c.fill.solid(); c.fill.fore_color.rgb = RGBColor.from_string(navy if i==0 else ('FFFFFF' if i%2 else 'E9F0F4'))
                for p in c.text_frame.paragraphs:
                    p.space_before = p.space_after = Pt(0); p.line_spacing = 1.1
                    for rr in p.runs:
                        rr.font.name = 'Noto Sans CJK SC'; rr.font.size = Pt(font_size); rr.font.bold = i==0
                        rr.font.color.rgb = RGBColor.from_string('FFFFFF' if i==0 else navy)
        if takeaway: caption(s, takeaway)
        return s

    def fig_slide(f, chapter='', subtitle=None):
        key = f.get('key', Path(f['path']).stem)
        figure_title = f['title']
        if key.startswith('waveform_') and '：' in figure_title:
            figure_title, case_identity = figure_title.split('：', 1)
            subtitle = case_identity + ' · ' + ('最差差值案例' if '_worst_' in key else '中位差值案例')
        s = new(figure_title, subtitle or f.get('subtitle', '读取保留证据；图形为位图，结果与来源可追溯'), f.get('source', ''), chapter)
        p = Path(f['path'])
        if not p.is_absolute():
            p = p if p.exists() else out/p
        with Image.open(p) as im: iw, ih = im.size
        width = min(12.0, 4.72*iw/ih)
        height = width*ih/iw
        s.shapes.add_picture(str(p), Inches((13.333-width)/2), Inches(1.60+(4.72-height)/2),
                             width=Inches(width), height=Inches(height))
        concise = {
            'quant_10_tokenizer_control_heatmaps': '每格为256个基础身份、三个优化种子的控制变体形状RMSE均值。此页比较三种端到端语义原型；O-only监督线性读出另页报告。',
            'quant_11_tokenizer_intervention_heatmaps': '干预变体与控制图共用色标。未见脉冲等机制使形状恢复显著恶化；三种语义原型及O-only线性读出的对照不能据单一基准泛化。',
            'quant_19_observation_linear_probe': 'O-only采用训练真值监督的线性读出。基准EEG/Hb形状RMSE为1.0541/0.1288，分布头为0.0433/0.0971；无耦合Hb反而是O-only更低（0.4224 vs 0.7512）。',
            'syn_wave_01_baseline': '首个评价基础轨迹、seed101：各模态独立恢复生成器条件形状。灰带为模型区间；本页是单例，覆盖结论以全256基础身份汇总为准。',
            'syn_wave_02_shape_vs_eeg_gain': '真实形状变化应被保留；仅改变EEG观测增益却也改变估计驱动。对比同一基础轨迹，不能把点估计的变化自动解释为真实生理变化。',
            'syn_wave_03_amplitude_vs_gains': '真幅值翻倍的ΔlogA=0.693，EEG/Hb估计仅约0.440/0.520；纯观测增益分别引起约0.375/0.367的假变化。本页为固定单例。',
            'syn_wave_04_exact_joint_alias': '两个模态的输入、mask及模型输出逐位相同，真log幅值相差0.693；不能仅从这些观测辨认两种真值。细窄区间反映生成假设的限制。',
            'syn_wave_05_ood_impulse': '未见非DCT脉冲使预测形状明显偏离真值，区间并未充分扩张。单例可视化与前面的全场景误差、覆盖统计分别解释。',
            'syn_wave_06_tau_response': '本例只改变τ，κ固定；输入EEG不变，Hb响应形状改变。图中响应真成分可加，但网络O/S由非线性解码器联合重建，不具有相同可加定义。',
            'syn_wave_07_correlated_hb': '生成器共同Hb项与驱动具有0.85的参考RMS归一化内积（并非Pearson相关）；它改变Hb推断的形状与幅值，说明来源仍会混淆。',
        }
        short_caption = next((value for prefix,value in concise.items() if key.startswith(prefix)), f.get('caption', ''))
        caption(s, short_caption)
        slide_sources[-1].update(figure=str(p.resolve()), figure_key=key, caption=short_caption, statistics=f.get('stats',{}))
        s.notes_slide.notes_text_frame.text += '\n图：'+str(p.resolve())+'\n完整图注：'+f.get('caption','')+'\n统计：'+json.dumps(f.get('stats',{}), ensure_ascii=False)
        return s

    s = new('从原始波形到条件语义—观测分解', 'SSM 改进讨论、冻结实验与完整结果报告 · 2026-10-02 run / 2026-10-03 export',
            ['docs/METHOD_RATIONALE.md','docs/EXPERIMENT_PLAN.md','summary.json','tokenizer/summary.json'], '研究报告')
    text(s, '响应动力学与连续\nsemantic–observation tokenizer', .8, 1.9, 11.8, 1.7, 36, bold=True)
    text(s, '可恢复的条件形状、仍混淆的幅值、未成立的实测语义增量', .85, 4.1, 11.6, .9, 25, color=teal)
    text(s, '三公开数据集  ·  三个同步区域  ·  四臂 × 三种子\n原始波形、模型分解、反事实、不确定性和容量控制逐层核验', .85, 5.3, 11.6, .85, 20, color=muted)
    table_slide('先给出三层结论', '有工程收益；共享生理语义的关键检验尚未通过', ['问题','本轮结果','允许的解释'], [
        ['SSM 机制','谱混合未超过匹配衰减；null 未通过','不能把拟合改善解释为更正确的 EEG→Hb 耦合'],
        ['合成语义','基准形状改善；未见机制和 alias 失效','生成器条件下的函数恢复，未建立通用幅值语义'],
        ['实测表征','O 有收益；加入 S 后误差增加','观测表示有用；当前 S 未提供独立增量']], [2.05,4.7,5.25],
        'summary.json; dataset_equal_comparisons.csv; tokenizer/summary.json', '结论')
    body('报告阅读路径', '围绕“波形怎样被解释、哪些解释尚不能成立”组织证据', [
        '01  讨论：原始叙事、竞争解释与实验取舍',
        '02  数据与模型：原生波形 → prepared 特征 → 条件成分',
        '03  SSM：合成、公开主比较、null 与逐窗分解',
        '04  Tokenizer：独立编码、配对反事实、alias、跨生成机制',
        '05  实测增量：完整观测与容量基线、最终判断和复现'], 'presentation_sources.json', '导读')
    body('项目最初的问题仍然不变', '分别观测 EEG 和 fNIRS，能否获得共享生理过程的有用表示？', [
        '构造一个联合坐标，不等于每个模态都能独立恢复该坐标。',
        '每个模态都可预测一个坐标，也不等于该坐标具有物理充分性。',
        '共享生理语义必须进一步解释反事实变化，并提供超过完整观测基线的可复现增量。'],
        'docs/METHOD_RATIONALE.md: Research question', '讨论', '本轮保持独立模态输入与完整 observation 信息；下游分类性能不是训练或选择端点。')
    table_slide('先前证据为何促使本轮改变设计', '仅作为讨论动机；不同推断制度的历史分数不混入本轮对照', ['先前观察','竞争解释','本轮回应'], [
        ['谱模式改善重建','可能只是减弱错误血管输入','匹配衰减、独立增益和随机方向'],
        ['附加 Hb 项可吸收误差','共同项与初态 / 动力学存在混淆','固定驱动、前缀校准与响应扩展分开'],
        ['空间 Hb 可补全目标','任务形状或共享 Hb 可能已足够','空间基线后再检验局部 EEG 增量'],
        ['typed 目标能被学会','生成先验可能制造语义确定性','幅值 alias、未见生成机制与分布输出']],
        [3.0,4.4,4.6], 'shared_driver_modes/20261002_v1; shared_driver_attribution/20261002_v1; docs/EXPERIMENT_PLAN.md', '讨论', size=17)
    table_slide('讨论覆盖六条改进方向', '科学论证与实现审核分别负责；将建议转化为有限、可检验的改动', ['讨论方向','形成的判断','本轮落实'], [
        ['谱 / 空间语义','有限区域与谱模式可描述，不等于源定位','保留 a、b；加入空间条件基线'],
        ['动力学','输入滞后与静脉响应需竞争检验','分别增加低通 τn、出流松弛 τv'],
        ['初态与归属','自由初态可制造响应而无 EEG 驱动','p=v proper prior；零前缀先验'],
        ['可证伪性','真实配对需优于错配与替代输入','错被试、非环绕时移、其他区域 EEG'],
        ['目标与不确定性','形状与幅值应分开；必须允许多解','单位 RMS 形状、log 幅值和 alias'],
        ['容量与泄漏','多几个坐标并不证明语义有效','O→O+S 和独立等容量 O-only']],
        [2.15,5.2,4.65], '讨论形成的冻结合同：docs/EXPERIMENT_PLAN.md; semantic_response_discovery_v1.yaml', '讨论', size=16)
    fig_slide(mfigs['method_evidence_chain'], '讨论')
    body('哪些建议没有直接进入本轮', '收紧目标和解释范围，避免把尚无锚点的概念写进 token 名称', [
        '未增加细胞类型、E/I、代谢率或精确脑源标签：当前频带特征和稀疏 Hb 区域没有足够来源证据。',
        '未把共同 Hb 项规定为噪声：广泛神经变化也可能进入该坐标。',
        '未同时自由拟合所有动力学：先用两个单独扩展检验有限响应策略。',
        '未启动 VQ、受保护评价或结果驱动追加调参：连续表示先接受语义与增量检验。'],
        'docs/EXPERIMENT_PLAN.md; semantic_response_discovery_v1.yaml', '讨论')

    fig_slide(mfigs['method_raw_lineage'], '数据与波形')
    table_slide('数据分组与实际独立单位', '所有变换按允许的训练分组拟合；评价被试与本轮训练 / 选择互斥',
        ['数据集','训练人数 / 原生窗','选择人数 / 原生窗','评价人数 / 原生窗'],
        [[short]+[str(metadata['counts'][ds][part]['subjects'])+' / '+str(metadata['counts'][ds][part]['native_windows']) for part in ('train','selection','evaluation')]
         for ds,short in [('eeg_fnirs_single_trial','Single-Trial'),('simultaneous_eeg_nirs','Simultaneous'),('visual_cognitive_motivation','Visual')]],
        [3.0,3.0,3.0,3.0], 'metadata_plan_proposal.json', '数据与波形',
        '评价共 15 人、120 原生窗口、360 区域行；三个区域不独立。全 cohort 此前已分析，不能称全新确认队列。')
    body('任务支持与外推范围', '分组互斥并不使小样本任务或已查看队列变为独立确认', [
        'Simultaneous 选择集没有 2-back；评价中的 DSR 仅 5 个原生窗口、2 位被试。',
        '后面的最差 DSR 窗口说明一个明确失败案例，不能代表广泛的 DSR 人群表现。',
        'Visual 只使用有三区域同步支持的 Probe1，并按 EEG Part / 起点去重。',
        '整个 cohort 此前已分析；本轮更换评价身份减少直接适配，仍只属于开发性新端点。'],
        'summary.json: limitations; metadata_plan_proposal.json', '数据边界')
    table_slide('原始单位与处理后坐标不能混用', '图中按各数据集的发布定义标注；未知单位保持未知', ['数据集','原生观测','本轮模型坐标'], [
        ['Single-Trial','EEG 电压；双波长光电强度','转换并处理后的相对 Hb；相对频带 log-power'],
        ['Simultaneous','EEG 电压；发布 HbO/HbR','保留单位来源，经 parent 处理与新训练尺度映射'],
        ['Visual','EDF 电压；发布 Hb（绝对物理单位未核实）','不冒称 μM；仅在声明的 prepared 坐标解释']],
        [2.4,4.8,4.8], 'docs/DATA_CONTRACT.md; docs/DATASETS_DESCRIPTION.md; waveform_provenance.json', '数据与波形',
        '探针表中的 context_raw 指 prepared 特征直接入 ridge；它不是此处展示的原生波形。', size=17)
    # Place a true data example before the mathematical model; remaining cases follow SSM results.
    if wfigs: fig_slide(wfigs[0], '原生波形示例')
    body('模型真正使用的张量与时间参考', '30 秒窗口；4 Hz 特征时间轴；原生 EEG 的采样率另外标注', [
        'EEG：[120, 30]，来自 6 个邻近电极 × 5 个频带（1–4、4–8、8–13、13–30、30–45 Hz）。',
        'Hb：[120, 2]，分别为 HbO、HbR；三个同步区域是前额、运动和后部。',
        '前 20 点 = 5 s 参考段；本轮重拟合训练尺度，不复用旧 outer PCA / SD。',
        'parent 处理和窗口内表示是非因果的；所有后缀成绩称离线处理后特征补全。'],
        'resolved_config.json: tensor/data; metadata_plan_proposal.json', '数据与波形')
    fig_slide(mfigs['method_visibility'], '可见性合同')
    fig_slide(mfigs['method_ssm_inference'], 'SSM 设计')
    table_slide('驱动比较：谱信息和衰减分开检验', '所有候选使用同一 EEG-only a、b 和同一 Hb 可见性', ['臂','输入定义','检验问题'], [
        ['宽频','u = a','最基本的固定谱坐标'],
        ['谱混合','u = (a + γb) / √(1+γ²)','谱对比 b 是否提供血管信息'],
        ['匹配衰减','u = a / √(1+γ²)','去掉 b，保留相同 a 衰减'],
        ['独立增益','u = g·a；含 g=0','单纯改变读出强度是否已足够'],
        ['随机 / 零输入','预定随机方向；u=0','特定谱方向与无 EEG 驱动对照']],
        [2.2,5.0,4.8], 'resolved_config.json: model; model_spec', 'SSM 设计',
        'γ 是候选读出坐标的权重；不能直接解释为已测生理耦合强度。', size=17)
    table_slide('响应扩展：只改变一个候选机制', '固定 parent H0 参数；不让新增自由度自动获得生理身份', ['部分','本轮方程 / 规则','解释限制'], [
        ['输入低通','du/dt = (r−u)/τn；τn=0 为即时输入','检验有限输入滞后 / 平滑策略'],
        ['出流松弛','fout = (τ·v^(1/α)+τv·f)/(τ+τv)','p、q 平衡均使用同一 fout'],
        ['体积响应','dv/dt = (f−v^(1/α))/(τ+τv)','不是只改体积而忽略质量平衡'],
        ['初态','p=v；log 坐标 proper prior','自由初态只作固定消融'],
        ['共同 Hb 项','4 个低频模式；方向 / 尺度按合同','可以补偿误差，来源仍未定']],
        [2.0,5.7,4.3], 'src/inference/semantic_response_ssm.py; Buxton et al. 2004, Eq.11', 'SSM 设计', size=16)
    body('初态与观测项为何必须单独讨论', '模型中存在可产生 Hb 响应、但并不携带 EEG 驱动的自由度', [
        '由同一质量平衡方程可得：d(p−v)/dt = −[fout/(τ·v)]·(p−v)，此处分母 v 为体积状态。',
        '因此 p−v 是衰减暂态；从 p=v 出发会保持该约束，自由 p 初态不代表新增持续过程。',
        '有 Hb 前缀时，初态与共同观测项可形成自由响应；g=0 的改善不能作为 EEG 耦合证据。',
        '零前缀时只用先验初态和零共同项，阻止隐藏目标 Hb 决定这些自由度。'],
        'semantic_response_ssm.py: _physical_response_slope / fit_fixed_driver_response', 'SSM 设计')
    table_slide('基线阶梯：复杂模型必须超过哪些信息', '所有训练 / 选择边界按被试分离；标签只进入命名明确的模板对照', ['基线','可利用的信息','意义'], [
        ['训练均值 / 任务模板','训练波形；任务模板有标签 oracle','共同时间形状可能已经足够'],
        ['自身 Hb 前缀 ridge','目标 Hb 的同一可见前缀','分离血管历史与 EEG 信息'],
        ['完整 EEG ridge','本地 EEG prepared 特征','物理约束是否超过直接回归'],
        ['空间 Hb ridge','其他两个区域 Hb，不含目标 Hb','广泛 Hb 信息是否已解释目标'],
        ['空间 Hb + EEG','即时 EEG / 固定稳定滤波器组','局部 EEG 是否有条件增量']],
        [3.2,4.75,4.05], 'resolved_config.json: baselines; fit_baselines; public_raw_context', 'SSM 设计', size=16)
    table_slide('主 null family 与分母', '40点前缀：谱混合、独立增益、最佳稳定响应 × 三种错配，共九项', ['null','对照支持','所能检验的特异性'], [
        ['任务匹配错训练被试','相同目标与评分时间点','真实跨被试配对是否有优势'],
        ['非环绕移位 12 s','仅对比 real_shift_support','避免循环拼接或不同有效支持'],
        ['同步其他区域 EEG','相同窗口，不换目标 Hb','局部性；不是普遍零耦合证明']],
        [3.2,4.5,4.3], 'primary_null_specificity.csv/json; docs/EXPERIMENT_PLAN.md', 'SSM 设计',
        '先被试内、再数据集内被试等权、最后三数据集等权；整被试符号枚举 2^15，九项统一 Holm。', size=17)
    body('描述区间与 null 检验的条件', '统计显著性、机制识别与独立确认分别解释', [
        '效应区间条件于冻结的训练 / 选择对象；三数据集等权，不按区域数或窗口数增加独立样本量。',
        '符号检验需要零假设下被试差值对称。完整枚举 2^15 种符号配置，不等于无条件随机化检验。',
        '九项 Holm 只覆盖注册的 primary null family；不覆盖 tokenizer、γ主对照或所有后续描述图。',
        '三优化种子重复使用同一评价身份，误差范围表示优化敏感性，不是独立被试置信区间。'],
        'docs/EXPERIMENT_PLAN.md; primary_null_specificity.json; tokenizer/scenario_variant_seed_summary.csv', '统计边界')
    table_slide('分数和不确定性各自回答什么', '误差、概率质量与来源解释分开', ['指标','较好方向','边界'], [
        ['训练 SD 标准化 NRMSE','更低','绝对差值不是百分比改善'],
        ['Gaussian CRPS / NLL','更低','selection 残差 SD，不是生理后验'],
        ['50 / 80 / 95% 覆盖与宽度','覆盖接近名义且区间有信息','宽区间也可能无用；必须分场景'],
        ['反事实响应 / 假变化','跟随真变化、抵抗观测改变','需要已知生成真值，不能由实测重建替代'],
        ['配对 null 与条件增量','真实配对和 S 有额外收益','不由视觉平滑或重建美观决定']],
        [3.3,3.6,5.1], 'resolved_config.json: endpoints/uncertainty; dataset_equal_comparisons.csv', '评价合同', size=16)
    table_slide('SSM 合成设计与选择限制', '冻结选择策略接受检验；没有用评价场景标签重选参数', ['设计','执行内容','解释'], [
        ['六场景','标准、非rest、输入低通、出流松弛、共同Hb、零耦合','受限与随机化两个 family'],
        ['幅值 / 持续时间','3 个幅值、3 个持续时间的平衡变化','不是完整全因子'],
        ['全局选择','所有场景的 control 与 intervention 汇总','重复 control 占一半权重'],
        ['选择结果','两个 family 均 γ=0、g=1、τn=τv=0','三个响应候选实际退化同一核']],
        [2.3,5.5,4.2], 'synthetic_calibration/*.json; synthetic_summary.csv; cross_mechanism_fit_matrix.csv', 'SSM 结果',
        '矩阵不能解释为机制识别准确率，也不能由相同预测直接证明低通与出流机制不可区分。', size=16)

    # Quantitative figures declare a section so report ordering does not depend on filename accidents.
    ssm_quant, token_quant, public_quant = [], [], []
    for f in qfigs:
        ordinal = int(f['key'].split('_')[1])
        if f.get('section') == 'tokenizer':
            token_quant.append(f)
        elif ordinal >= 16:
            public_quant.append(f)
        elif ordinal >= 9:
            token_quant.append(f)
        else: ssm_quant.append(f)
    for f in ssm_quant: fig_slide(f, 'SSM 定量结果')
    body('如何阅读接下来的原始—成分波形图', '展示典型表现与失败边界；案例只是总体结果的解释，不替代总体统计', [
        '从已完成评价结果按固定报告选图规则取两例：best_stable 相对 gain_selected 的逐窗误差差值最接近中位数，以及最差。',
        '原生波形保留通道、单位与真实时钟；prepared 波形显示本轮实际拟合坐标。',
        'Hb 同坐标分解：观测 = 条件物理响应 + 受限共同项 + 残差；曲线可能相互抵消。',
        '这些是事后描述性案例，不是预注册独立检验；任何“神经源 / 浅表 / 纯噪声”标签都需要额外证据。'],
        'waveform_provenance.json; measured_cells.csv; retained per-cell NPZ', '成分波形')
    for f in wfigs[1:]: fig_slide(f, '原始与成分波形')
    table_slide('看见分离曲线之后，仍不能跳过的解释边界', '数学分解与来源识别不是同一件事', ['曲线 / 坐标','可用名称','不能据此命名'], [
        ['a、b','有效宽频 / 谱对比连续模式','兴奋源、抑制源、细胞来源'],
        ['physical_prediction','给定驱动、参数与初态的 Hb 响应','已恢复的神经源 Hb'],
        ['observation_component','受限共同 Hb 观测项','已识别浅表血流、系统生理、运动伪迹'],
        ['观测减预测','该模型与观测坐标下的残差','纯噪声或全部私有神经信息']],
        [3.0,4.6,4.4], 'semantic_response_ssm.py; waveform_provenance.json; docs/METHOD_RATIONALE.md', '成分解释', size=17)

    fig_slide(mfigs['method_tokenizer'], 'Tokenizer 设计')
    body('语义目标：形状与幅值分别输出', '完整 observation 分支保留模型不能可靠命名的信息', [
        '参考中心化：a_c(t) = a(t) − mean[a(0:20)]；前 20 点对应 5 s。',
        '幅值 A = RMS(a_c)；形状 s(t) = a_c(t)/A；幅值目标 ℓ = log A，并只用训练分组标准化。',
        '分布臂输出 Gaussian 边缘均值 / SD；不表示完整时序后验或跨时协方差。',
        '连续推理后聚合为 12 个 patch；平均边缘 SD 只是完全正相关下的上界，不是已校准 patch 区间。'],
        'semantic_response_tokenizer.py: token_metadata / export_tokens; synthetic prepared manifest', 'Tokenizer 设计')
    table_slide('四臂 × 三种子：目标、随机化和容量分开比较', '基础神经网络均为 458,168 参数；O-only 的附加线性读出另计', ['臂','训练机制','语义与重建'], [
        ['等容量 O-only','随机化生成器','全部均值坐标只学习完整重建'],
        ['受限点模型','受限生成器','shape / log幅值点监督 + 完整重建'],
        ['随机化点模型','随机化生成器','相同点监督，分离随机化效应'],
        ['随机化分布模型','随机化生成器','Gaussian语义损失 + 完整重建']],
        [3.0,3.3,5.7], 'tokenizer/summary.json; resolved_config.json: tokenizer', 'Tokenizer 设计',
        'width=64、O维度=8、batch=64、Adam lr=0.001；最多160 epoch、patience=20；seeds=101/202/303。', size=17)
    body('O-only 合成评价与区间的实际来源', '网络监督、冻结表征可读出性和预测区间分别记录', [
        'O-only 网络本身只训练完整重建；合成评价另用训练集真值，对全部均值坐标拟合 shape / log幅值线性 ridge（α=1），再冻结到评价集。',
        '因此 O-only 合成误差衡量“重建表征 + 监督线性读出”；等容量指基础网络与均值瓶颈，不包括这个附加 ridge。',
        '点模型和 O-only 用选择集残差 RMS：形状和幅值各一个常数尺度。分布模型用网络 SD × 选择集温度；温度限制为[0.1,10]。',
        '这个合成可读出性端点与后面的实测条件 Hb 探针不同；两者都不能自动赋予表示生理来源身份。'],
        'train_semantic_response_tokenizer.py: semantic_predictions / fit / calibrate / uncertainty', 'Tokenizer 设计')
    table_slide('合成训练与评价身份严格分开', '同基础轨迹的全部干预变体放在同一个 split', ['分组','基础身份','场景与记录'], [
        ['训练','512 / family','8场景 × control/intervention = 8,192条'],
        ['选择','128 / family','8场景 × control/intervention = 2,048条'],
        ['共同评价','256','20场景 × control/intervention = 10,240条'],
        ['计数解释','所有四臂共享相同评价案例','两个模态 × 12模型 = 245,760条评分记录']],
        [2.6,3.0,6.4], 'tokenizer/prepared/*/*.json; tokenizer/summary.json', 'Tokenizer 设计',
        '三优化种子不增加评价基础身份或被试数；区间与误差必须分 control / intervention 汇报。', size=17)
    table_slide('评价场景 I：从真实变化到观测混淆', '先检查训练支持范围内的函数恢复及不应出现的语义变化', ['类别','场景','问题'], [
        ['真实驱动改变','true_a_amplitude / true_a_shape','语义是否跟随真实变化'],
        ['观测增益','feature_gain_only / Hb_gain_only','幅值是否被误认作驱动变化'],
        ['共同 Hb','common_Hb / correlated_Hb_component','广泛变化是否污染共享语义'],
        ['动力学 / 初态','tau_change / initial_change','观测形状变化会被写入哪一分支'],
        ['基准','baseline','在已定义生成机制中能否学会目标']],
        [2.35,5.8,3.85], 'resolved_config.json: tokenizer.evaluation_scenarios; semantic_response_synthetic.py', 'Tokenizer 设计', size=16)
    table_slide('评价场景 II：生成机制外推与识别反例', '随机参数范围更宽，不等于已经覆盖新的生成机制', ['类别','场景','问题'], [
        ['未见时间基底','non_dct_impulse / ou_drive','平滑训练先验是否压制真实变化'],
        ['观测映射改变','time_varying_feature_gain / rotated_spectral_mixture','固定谱/增益假设是否稳定'],
        ['响应机制改变','input_lowpass / viscoelastic_outflow / alternative_hrf','时间响应错配与状态归属'],
        ['关系改变','independent_Hb_drive / no_coupling','无共享信息时会否仍输出共享语义'],
        ['精确别名','exact_eeg_alias / exact_joint_alias','相同观测是否对应不止一个真值']],
        [2.3,6.25,3.45], 'resolved_config.json: tokenizer.evaluation_scenarios; semantic_response_synthetic.py', 'Tokenizer 设计', size=15)
    for f in token_quant: fig_slide(f, 'Tokenizer 合成结果')
    body('精确 alias 应怎样解释', '这是明确允许的竞争机制，不是要求网络从相同输入凭空辨认真值', [
        '将真实驱动幅值放大，同时反向改变 EEG 观测增益，可保留完全相同的 EEG。',
        'joint alias 进一步使用精确 Hb 补偿，使两个模态输入都相同，但真实驱动幅值不同。',
        '相同输入的确定编码器必然给出相同输出；本轮失败在于窄幅值区间不覆盖允许的两个机制。',
        '解决方向是识别集合、先验敏感性或明示幅值不受支持；只增加输入异常检测器无法识别精确 alias。'],
        'semantic_response_synthetic.py: exact aliases; tokenizer/scenario_variant_seed_summary.csv', '反事实波形')
    for f in sfigs: fig_slide(f, '合成反事实波形')
    fig_slide(mfigs['method_public_probe'], '实测冻结探针')
    for f in public_quant: fig_slide(f, '实测冻结探针')
    table_slide('实测主结果：语义分支没有提供额外收益', '随机化分布模型，前缀40；三优化种子均值；相同上下文与评分支持',
        ['输入表征','NRMSE ↓','标准化 CRPS ↓'], [
        ['prepared 原始上下文', '1.228612','0.846234'],
        ['上下文 + 同模型 O','1.151679','0.790067'],
        ['上下文 + 同模型 O + S','1.155386','0.792139'],
        ['上下文 + 独立等容量纯 O','1.139294','0.781601']], [6.2,2.9,2.9],
        'tokenizer/public_probes/distribution_randomized_generator/seed_*/record.json', '实测结论',
        'O→O+S 的 NRMSE 增加0.003707；相对等容量纯O增加0.016092，三个优化种子方向一致。', size=19)
    body('三部分证据共同限定当前结论', '不能由合成成功翻转实测负结果，也不能由当前负结果否定所有生理表示方案', [
        'SSM：谱方向、稳定动力学与配对特异性尚未构成一致的 EEG→Hb 机制证据。',
        '合成：可以恢复生成器支持的相对形状；幅值与观测增益仍混淆，跨机制不确定性失校准。',
        '实测：observation 表示在可见前缀下有工程价值；semantic 坐标未超过 observation 与容量对照。',
        '本轮交付的是可复现的连续接口、有限正结果和清晰失败边界；没有获得生理 teacher / VQ 晋级依据。'],
        'registry main.semantic_response; summary.json; tokenizer/summary.json', '综合判断')
    table_slide('执行完成与科学支持分别记录', '失败和修复保留；最终分母没有由成功子集替代', ['核验项目','结果'], [
        ['SSM 合成','5,760项成功；2,880实际前缀拟合全部收敛'],
        ['SSM 公开','23,040项成功；9,000实际拟合全部收敛'],
        ['Tokenizer','12训练 + 12合成评价 + 12公开探针全部完成'],
        ['评分分母','245,760合成记录；36,720公开探针记录，无非有限 / 失败行'],
        ['软件与证据','125项定向测试；29项证据核验；状态校验通过']], [3.55,8.45],
        'verification.json; tokenizer/summary.json; supervisor/completion.json', '执行与复现', size=18)
    body('运行过程中的修复与资源选择', '源码版本、失败日志和计算规则均保留，未用重训掩盖负结果', [
        '持久 user systemd 监督；CPU 单数值线程；SSM 按独立单元并行，GPU 训练按实测吞吐选择每卡两任务。',
        '恢复入口曾因 CUDA checkpoint 将 RNG 状态放入错误设备而失败；修复只把 RNG 状态恢复到 CPU，训练损失与网络不变。',
        '汇总修复恢复了反事实表；features/batch_v2 补充训练增益来源，保留原 batch_v1。',
        '全部训练与评分结束；原始失败记录、六套不可变源码快照、配置、监督日志及完成状态都可追溯。'],
        'training_resume_v3.json; training_launch_v2.json; source_snapshot_*_identity.json; analysis_launch.json', '执行与复现')
    table_slide('下一步应由失败面决定', '建议，不代表新实验已启动', ['未解决问题','最小下一步','必须保留的检验'], [
        ['动力学选择退化','固定真值驱动 / 参数的有限 oracle 合成','区分输入误差、参数失配与机制混淆'],
        ['幅值非唯一','识别集合或生成先验敏感性','保留 exact alias 与场景分层覆盖'],
        ['跨机制形状失效','扩展时间基底前先定义可恢复函数目标','未见机制评价不能参与重选'],
        ['S 无实测增量','重新检验目标本身；保留完整 O 基线','同容量、同上下文、同前缀与配对 null']],
        [2.7,4.75,4.55], 'registry main.semantic_response.next_step; scientific review of retained endpoints', '下一步',
        '不自动追加实测调参、受保护评价或 VQ；新的来源主张需要独立观测或机制锚定。', size=16)
    table_slide('附录：完整响应方程与 Hb 读出', '实现坐标为 [u,s,f,v,p,q]；状态曲线是模型条件推断结果', ['坐标','方程 / 定义'], [
        ['输入状态','τn>0：du/dt=(r−u)/τn，u(0)=0；τn=0：u=r'],
        ['信号与血流','ds/dt = β·u − κ·s − γ_feedback·(f−1)；df/dt = s'],
        ['体积','dv/dt = (f−fout)/τ；fout = (τ·v^(1/α)+τv·f)/(τ+τv)'],
        ['总 Hb 坐标','dp/dt = [f−(fout/v)·p]/τ'],
        ['去氧 Hb 坐标','dq/dt = [f·E(f)/E0−(fout/v)·q]/τ；E(f)=1−(1−E0)^(1/f)'],
        ['Hb 读出','HbO=P0·(p−1)−Q0·(q−1)；HbR=Q0·(q−1)'],
        ['观测坐标','读出经过同一 Hb 处理算子后，再加共同项；不能直接与原生光强相加']],
        [2.3,9.7], 'semantic_response_ssm.py: _physical_response_slope / fixed_driver_response', '附录',
        'γ_feedback 为固定生理模型参数；与谱读出候选权重 γ 不是同一参数。β和其余未声明扩展参数固定为parent H0。', size=16)
    table_slide('附录：冻结的 SSM 候选与数值预算', '参数选择只使用允许的训练 / 选择资料；不按评价机制标签重选', ['参数 / 规则','冻结内容'], [
        ['谱混合 γ','−1、−0.5、0、0.5、1'],
        ['独立增益 g','0、0.5、1/√2、2/√5、1'],
        ['输入低通 τn / 出流 τv','τn∈{0,0.5,1,2}s；τv∈{0,1,3,6}s'],
        ['EEG 先验与时间基底','24 temporal modes；state SD=0.025；AR φ=0.95、weight=0.1'],
        ['前缀与选择','40点主选择复用于0点；失败惩罚10；同分选择较简单候选'],
        ['求解器','最大120 evaluations；RK每步4 substeps；gradient tolerance=10⁻⁶']],
        [3.8,8.2], 'resolved_config.json: model/solver/calibration', '附录', size=17)
    body('附录：训练损失和表征容量', '同样的网络参数数量，明确不同坐标受到什么监督', [
        '每模态 L = L_shape + 0.25·L_logA + L_full_reconstruction，总损失为 EEG 与 Hb 之和。点模型语义用平方误差，分布模型用 Gaussian NLL；重建在本模态训练标准化坐标计算。',
        '受监督臂把语义均值 detach 后送入本模态解码器；完整重建不能反向改写语义头。',
        'O-only 臂不使用语义真值损失；全部均值坐标都可通过完整重建学习，因此它的原始 S 头误差不是公平的语义恢复基线。',
        '均值表征：每时间点8维 O、1维形状，加1个窗口幅值；尺度头所有臂均存在，但点模型不将未监督尺度头解释为区间。'],
        'semantic_response_tokenizer.py: response_loss / _ModalityBranch; resolved_config.json', '附录')
    epoch_rows = []
    for arm, label in [('observation_only_capacity_matched','等容量 O-only'),
                       ('point_restricted_generator','受限点'),('point_randomized_generator','随机化点'),
                       ('distribution_randomized_generator','随机化分布')]:
        row = [label]
        for seed in (101,202,303):
            rec = next(x for x in tok['fits'] if x['arm']==arm and x['seed']==seed)
            row.append(f"{rec['best_epoch']} / {rec['completed_epochs']}")
        epoch_rows.append(row)
    table_slide('附录：12 个训练实例的选择与停止', '每格为最佳选择 epoch / 实际停止 epoch；相同上限不等于实际训练步数相同',
        ['臂','seed101','seed202','seed303'], epoch_rows, [3.6,2.8,2.8,2.8],
        'tokenizer/fits/*/seed_*/record.json; tokenizer/summary.json', '附录',
        '全部≤160 epoch；早停patience=20。最佳模型由各臂注册的选择损失决定，不能直接横比不同损失的原始数值。', size=19)
    table_slide('文献提供动机，不替代本数据的证据', '本轮引用均用于明确模型假设或解释边界', ['文献','与本轮有关的内容','不能外推'], [
        ['Buxton et al., 2004','血流、血容量、氧合响应及出流松弛模型','不是本 EEG–fNIRS 参数的标定'],
        ['Hermes et al., 2017','不同电生理成分与血管响应的关系','ECoG–BOLD 结果不直接指定头皮谱读出'],
        ['Locatello et al., 2019','解耦依赖模型 / 数据的归纳假设','重建好并不赋予坐标来源名称'],
        ['Hermans et al., 2022','模拟推断可能过度自信，需覆盖审计','本轮 Gaussian 头不是完整 Bayesian SBI']],
        [3.1,5.0,3.9], [
          'https://fmri.ucsd.edu/tliu/pdf/buxton04_model.pdf',
          'https://journals.plos.org/plosbiology/article?id=10.1371/journal.pbio.2001461',
          'https://proceedings.mlr.press/v97/locatello19a.html',
          'https://arxiv.org/abs/2110.06581'], '参考', size=16)
    body('参考文献与可访问链接', '文献动机与本轮数据支持程度分开；不以引用替代实验验证', [
        'Buxton et al. (2004). Modeling the hemodynamic response to brain activation.\nhttps://fmri.ucsd.edu/tliu/pdf/buxton04_model.pdf',
        'Hermes et al. (2017). Neuronal synchrony and the relation between the blood-oxygen-level dependent response and the local field potential.\nhttps://journals.plos.org/plosbiology/article?id=10.1371/journal.pbio.2001461',
        'Locatello et al. (2019). Challenging Common Assumptions in the Unsupervised Learning of Disentangled Representations.\nhttps://proceedings.mlr.press/v97/locatello19a.html',
        'Hermans et al. (2022). A Trust Crisis In Simulation-Based Inference? Your Posterior Approximations Can Be Unfaithful.\nhttps://arxiv.org/abs/2110.06581'],
        'Primary sources checked 2026-10-03; see hyperlinks on slide', '参考')
    body('证据与复现导航', '正文、表格可编辑；所有科学图为 200–250 dpi PNG，并保留来源清单', [
        '实验合同：semantic_response_discovery_v1.yaml；状态 owner：research_state/registry.json。',
        'SSM：summary / dataset_equal_comparisons / primary_null_specificity / calibration 及逐单元 NPZ。',
        'Tokenizer：summary、scenario_variant_seed_summary，以及每模型 evaluation / public_probes。',
        '本报告：presentation_sources.json、三类图形 provenance、源代码快照与 PDF 导出核验。'],
        'experiments/RESULTS_INDEX.md; presentation_sources.json', '复现')
    deck.save(out/draft_name)
    record = dict(schema='semantic_response_presentation_sources_v1', run=str(run), presentation=draft_name,
                  slides=slide_sources, slide_count=len(deck.slides), bitmap_figures=True,
                  editable_body_text=True, shape_audit=shape_audit,
                  interpretation='dated communication export; owning evidence remains the registered run',
                  report_source_sha256=hashlib.sha256(Path(__file__).read_bytes()).hexdigest())
    (out/'presentation_sources.json').write_text(json.dumps(record, ensure_ascii=False, indent=2)+'\n')
    print(json.dumps(dict(presentation=str(out/draft_name), slides=len(deck.slides),
                          figures=sum('figure' in row for row in slide_sources)), ensure_ascii=False))
    return record
